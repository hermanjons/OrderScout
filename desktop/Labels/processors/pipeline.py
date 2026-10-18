# Labels/processors/pipeline.py
from __future__ import annotations

import gc
import json
import os
import zipfile
from copy import deepcopy
from functools import lru_cache
from typing import List, Dict, Any, Optional, Tuple

import tempfile
from pathlib import Path
from io import BytesIO
from datetime import datetime, timezone
import math
from lxml import etree

from Feedback.processors.pipeline import Result, map_error_to_message
from settings import DB_NAME

# ✅ Orders'dan: seçilen orderNumber'ların full detayları (header + data + items)
from Orders.processors.trendyol.pipeline import get_order_full_details_by_numbers

# ✅ yeni constants split sistemi
from Labels.constants import LABEL_IMAGES_DIR
from Labels.constants.template_constants import get_label_model_config

from sqlalchemy.orm.exc import DetachedInstanceError

_LABEL_EXPORT_MERGE_CHUNK_SIZE = 50
_LABEL_INLINE_IMAGE_CACHE_SIZE = 64
_BARCODE_CACHE_SIZE = 512
_WORDPROCESSINGML_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
_WORD_NS = {"w": _WORDPROCESSINGML_NS}


# ============================================================
# 0) VALIDATION + NORMALIZATION
# ============================================================
def _normalize_cargo_name(name: str) -> str:
    # casefold + whitespace normalize
    return " ".join((name or "").strip().casefold().split())


def _normalize_cargo_logo_map(raw: Dict[str, dict], aliases: Dict[str, list] | None = None) -> Dict[str, dict]:
    out: Dict[str, dict] = {}
    for k, v in (raw or {}).items():
        nk = _normalize_cargo_name(k)
        if nk and nk not in out:
            out[nk] = v
        if aliases and k in aliases:
            for a in aliases.get(k) or []:
                na = _normalize_cargo_name(a)
                if na and na not in out:
                    out[na] = v
    return out


def validate_label_model_config(brand_code: str, model_code: str) -> Result:
    try:
        cfg = get_label_model_config(brand_code, model_code)
        if not cfg:
            return Result.fail(f"Etiket konfigi bulunamadı: {brand_code}/{model_code}", close_dialog=False)

        # required config
        if int(cfg.get("labels_per_page", 0) or 0) <= 0:
            return Result.fail("labels_per_page geçersiz.", close_dialog=False)
        if int(cfg.get("max_items_per_label", 0) or 0) <= 0:
            return Result.fail("max_items_per_label geçersiz.", close_dialog=False)

        placeholders = cfg.get("placeholders")
        if not isinstance(placeholders, dict):
            return Result.fail("placeholders konfigi eksik.", close_dialog=False)

        required_base = {
            "ordernumber",
            "name",
            "surname",
            "address",
            "cargotrackingnumber",
            "cargoprovidername",
            "barcode",
            "storename",
            "platform",
            "sla_hours_left",
            "product",
            "qty",
        }
        missing = [k for k in required_base if k not in placeholders]
        if missing:
            return Result.fail(
                f"Template placeholders eksik: {', '.join(sorted(missing))}",
                close_dialog=False,
            )

        barcode_cfg = cfg.get("barcode")
        if not isinstance(barcode_cfg, dict):
            return Result.fail("Barkod config eksik.", close_dialog=False)
        for k in ("module_width", "module_height", "font_size", "text_distance", "quiet_zone", "image_width_mm", "attention_image_width_mm"):
            if k not in barcode_cfg:
                return Result.fail(f"Barkod config eksik: {k}", close_dialog=False)

        aliases = cfg.get("cargo_provider_aliases")
        if aliases is not None and not isinstance(aliases, dict):
            return Result.fail("cargo_provider_aliases formatı geçersiz.", close_dialog=False)

        tp = cfg.get("template_path")
        if not tp or not Path(tp).is_file():
            return Result.fail(f"Word şablonu bulunamadı: {tp}", close_dialog=False)

        return Result.ok("Label config OK.", close_dialog=False, data={"cfg": cfg})
    except Exception as e:
        return Result.fail(map_error_to_message(e), error=e, close_dialog=False)


@lru_cache(maxsize=1)
def _barcode_runtime():
    from barcode import Code128
    from barcode.writer import ImageWriter

    return Code128, ImageWriter


@lru_cache(maxsize=1)
def _docxtpl_runtime():
    from docxtpl import DocxTemplate, InlineImage, RichText
    from docx.shared import Mm

    return DocxTemplate, InlineImage, RichText, Mm


@lru_cache(maxsize=1)
def _docx_merge_runtime():
    from docx import Document
    from docxcompose.composer import Composer

    return Document, Composer


@lru_cache(maxsize=_LABEL_INLINE_IMAGE_CACHE_SIZE)
def _read_inline_image_bytes(path_text: str) -> bytes:
    return Path(path_text).read_bytes()


@lru_cache(maxsize=_BARCODE_CACHE_SIZE)
def _render_code128_png_cached(value: str, barcode_cfg_json: str, dpi: int) -> bytes:
    barcode_cfg = json.loads(barcode_cfg_json)
    Code128, ImageWriter = _barcode_runtime()
    writer_opts: Dict[str, Any] = {
        "module_width": float(barcode_cfg["module_width"]),
        "module_height": float(barcode_cfg["module_height"]),
        "font_size": int(barcode_cfg["font_size"]),
        "text_distance": float(barcode_cfg["text_distance"]),
        "quiet_zone": float(barcode_cfg["quiet_zone"]),
        "write_text": False,
        "dpi": int(dpi),
    }
    code = Code128(value, writer=ImageWriter())
    pil_img = code.render(writer_options=writer_opts)

    buf = BytesIO()
    pil_img.save(buf, format="PNG")
    return buf.getvalue()

# ============================================================
# 0) PAYLOAD SORT
# ============================================================
def sort_label_payload(label_payload: dict, mode: str) -> Result:
    """
    Aynı orderNumber grubu asla bölünmez.
    mode: none / product / quantity / optimal
    """
    try:
        if not isinstance(label_payload, dict):
            return Result.fail("label_payload dict olmalı.", close_dialog=False)

        mode = (mode or "none").strip().lower()
        if not label_payload or mode == "none":
            return Result.ok("Sıralama yok.", close_dialog=False, data={"label_payload": label_payload})

        pages = label_payload.get("pages") or []
        labels_per_page = int(label_payload.get("labels_per_page") or 24)
        max_items_per_label = int(label_payload.get("max_items_per_label") or 8)

        # Flatten
        all_labels: List[Dict[str, Any]] = [lbl for page in pages for lbl in (page or [])]
        if not all_labels:
            return Result.ok("Boş payload.", close_dialog=False, data={"label_payload": label_payload})

        # Prepare keys once
        prod_keys = [f"prod{i}" for i in range(1, max_items_per_label + 1)]
        qty_keys = [f"qty{i}" for i in range(1, max_items_per_label + 1)]

        # Group by orderNumber; keep insertion order
        groups: list[tuple[str, list[dict]]] = []
        group_map: dict[str, list[dict]] = {}

        for idx, lbl in enumerate(all_labels):
            order_no = (lbl.get("orderNumber") or "").strip()
            key = order_no if order_no else f"__SINGLE__{idx}"

            bucket = group_map.get(key)
            if bucket is None:
                bucket = []
                group_map[key] = bucket
                groups.append((key, bucket))
            bucket.append(lbl)

        metrics_cache: dict[str, tuple[str, int]] = {}

        def compute_metrics_for_group(glabels: list[dict]) -> tuple[str, int]:
            primary_product = ""
            total_qty = 0

            for g in glabels:
                if not primary_product:
                    for pk in prod_keys:
                        p = (g.get(pk, "") or "").strip()
                        if p:
                            primary_product = p
                            break

                for qk in qty_keys:
                    q = g.get(qk, 0)
                    try:
                        total_qty += int(q)
                    except (TypeError, ValueError):
                        pass

            return primary_product.strip().casefold(), total_qty

        def sort_key(group: tuple[str, list[dict]]) -> Tuple[Any, ...]:
            key, glabels = group
            if key not in metrics_cache:
                metrics_cache[key] = compute_metrics_for_group(glabels)

            prod_name, total_qty = metrics_cache[key]

            if mode == "product":
                return prod_name, -total_qty
            if mode == "quantity":
                return -total_qty, prod_name
            if mode == "optimal":
                return prod_name, -total_qty

            return (0,)

        sorted_groups = sorted(groups, key=sort_key)

        sorted_labels: list[dict] = []
        for _, glabels in sorted_groups:
            sorted_labels.extend(glabels)

        if labels_per_page <= 0:
            labels_per_page = 24

        new_pages: list[list[dict]] = [
            sorted_labels[i:i + labels_per_page]
            for i in range(0, len(sorted_labels), labels_per_page)
        ]

        new_payload = dict(label_payload)
        new_payload["pages"] = new_pages
        new_payload["total_labels"] = len(sorted_labels)
        new_payload["total_pages"] = len(new_pages)

        return Result.ok("Payload sıralandı.", close_dialog=False, data={"label_payload": new_payload})

    except Exception as e:
        return Result.fail(map_error_to_message(e), error=e, close_dialog=False)


# ============================================================
# 2) PAYLOAD ÜRETİMİ - helper'lar
# ============================================================
def _normalize_items(items: list, *, latest_snapshot: Any = None) -> Result:
    try:
        parsed_rows: List[Dict[str, Any]] = []
        normalized: List[Dict[str, Any]] = []

        for it in (items or []):
            def _v(name: str, default=None):
                return _safe_get_attr(it, name, default)

            qty_raw = _v("quantity", None)
            qty_int = _to_int(qty_raw, default=0)
            if qty_int < 0:
                qty_int = 0

            merchant_sku = str(_v("merchantSku", "") or "").strip()
            product_name = str(_v("productName", "") or "").strip()
            barcode = str(_v("barcode", "") or "").strip()
            status = str(_v("orderLineItemStatusName", "") or _v("status", "") or "").strip()
            status_norm = _normalize_status(status)
            task_date = _to_int(_v("taskDate", None), default=-1)
            order_data_id_raw = _v("order_data_id", None)
            if order_data_id_raw is None:
                order_data_id_raw = _v("orderDataId", None)

            row = {
                "item": it,
                "qty": qty_int,
                "merchant_sku": merchant_sku,
                "product_name": product_name,
                "barcode": barcode,
                "status": status,
                "status_norm": status_norm,
                "task_date": task_date,
                "order_data_id_text": str(order_data_id_raw).strip() if order_data_id_raw is not None else "",
                "order_data_id_int": _to_int(order_data_id_raw, default=-1),
                "pk_int": _to_int(_v("pk", None), default=-1),
            }
            parsed_rows.append(row)

        if not parsed_rows:
            return Result.ok("Items normalize edildi.", close_dialog=False, data={"items": normalized})

        latest_snapshot_id = _safe_get_attr(latest_snapshot, "id", None)
        latest_snapshot_id_text = str(latest_snapshot_id).strip() if latest_snapshot_id is not None else ""
        if not latest_snapshot_id_text:
            return Result.fail("Latest snapshot id bulunamadı.", close_dialog=False)

        selected_rows = [
            r for r in parsed_rows
            if (not _is_cancel_like_status(r.get("status_norm")) and r.get("order_data_id_text") == latest_snapshot_id_text)
        ]
        selected_rows = _dedupe_rows_by_identity(selected_rows)

        selected_rows = sorted(
            selected_rows,
            key=lambda r: (
                _to_int(r.get("task_date"), default=-1),
                -_to_int(r.get("pk_int"), default=-1),
                str(r.get("merchant_sku", "")).casefold(),
            ),
            reverse=True,
        )

        for row in selected_rows:
            qty_int = _to_int(row.get("qty"), default=0)
            if qty_int <= 0:
                continue

            name = str(row.get("merchant_sku", "") or "").strip()
            if not name:
                name = str(row.get("product_name", "") or "").strip()
            if not name:
                name = "HATALI ETİKET"

            normalized.append({"name": name, "qty": qty_int})

        return Result.ok("Items normalize edildi.", close_dialog=False, data={"items": normalized})

    except Exception as e:
        return Result.fail(map_error_to_message(e), error=e, close_dialog=False)


def _chunk_items(normalized: List[Dict[str, Any]], *, max_items_per_label: int) -> Result:
    try:
        if max_items_per_label <= 0:
            max_items_per_label = 8
        chunks = [normalized[i:i + max_items_per_label] for i in range(0, len(normalized), max_items_per_label)]
        return Result.ok("Chunk OK.", close_dialog=False, data={"chunks": chunks or [[]]})
    except Exception as e:
        return Result.fail(map_error_to_message(e), error=e, close_dialog=False)


def _safe_get_attr(obj, name: str, default=None):
    try:
        if obj is None:
            return default
        # dict support
        if isinstance(obj, dict):
            return obj.get(name, default)
        return getattr(obj, name, default)
    except (DetachedInstanceError, Exception):
        return default


def _to_int(value: Any, default: int = 0) -> int:
    try:
        return int(value)
    except Exception:
        return int(default)


def _normalize_status(status_value: Any) -> str:
    return str(status_value or "").strip().casefold()


def _is_cancel_like_status(status_value: Any) -> bool:
    status_text = _normalize_status(status_value)
    return status_text in {"cancelled", "canceled", "cancel"}


def _item_identity(row: Dict[str, Any]) -> tuple:
    item = row.get("item")
    line_id = (
        _safe_get_attr(item, "lineId", None)
        or _safe_get_attr(item, "line_item_id", None)
        or _safe_get_attr(item, "lineItemId", None)
        or _safe_get_attr(item, "orderLineItemId", None)
        or _safe_get_attr(item, "id", None)
    )
    product_code = _safe_get_attr(item, "productCode", None)
    merchant_sku = str(row.get("merchant_sku", "") or "").strip().casefold()
    barcode = str(row.get("barcode", "") or "").strip().casefold()
    return (
        str(line_id).strip() if line_id is not None else "",
        str(product_code).strip() if product_code is not None else "",
        merchant_sku,
        barcode,
    )


def _is_row_newer(candidate: Dict[str, Any], current: Dict[str, Any]) -> bool:
    c_task = _to_int(candidate.get("task_date"), default=-1)
    p_task = _to_int(current.get("task_date"), default=-1)
    if c_task != p_task:
        return c_task > p_task

    c_data = _to_int(candidate.get("order_data_id_int"), default=-1)
    p_data = _to_int(current.get("order_data_id_int"), default=-1)
    if c_data != p_data:
        return c_data > p_data

    c_pk = _to_int(candidate.get("pk_int"), default=-1)
    p_pk = _to_int(current.get("pk_int"), default=-1)
    return c_pk > p_pk


def _dedupe_rows_by_identity(rows: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
    grouped: Dict[tuple, Dict[str, Any]] = {}
    for row in rows or []:
        key = _item_identity(row)
        existing = grouped.get(key)
        if existing is None or _is_row_newer(row, existing):
            grouped[key] = row
    return list(grouped.values())


def _build_label_dict(
        *,
        latest_snapshot: Any,
        chunk: List[Dict[str, Any]],
        chunk_idx: int,
        max_items_per_label: int,
        debug_index: int,
) -> Result:
    """
    latest_snapshot: OrderData instance (mode="latest" ile 1 adet gelmeli).

    Not:
      - store_name/platform bu fonksiyonun işi değil.
      - Detached riskini azaltmak için relationship'e güvenmeyiz.
      - store/platform dışarıdan (pkg) label'a sonradan enjekte edilir.
    """
    try:
        if latest_snapshot is None:
            return Result.fail("Sipariş snapshot bulunamadı (latest_snapshot None).", close_dialog=False)

        if max_items_per_label <= 0:
            max_items_per_label = 8

        # -----------------------------
        # ORDER NUMBER
        # -----------------------------
        order_no = (_safe_get_attr(latest_snapshot, "orderNumber", "") or "").strip()
        if not order_no:
            return Result.fail("Latest snapshot orderNumber boş.", close_dialog=False)

        # -----------------------------
        # CUSTOMER + ADDRESS
        # -----------------------------
        first = (_safe_get_attr(latest_snapshot, "customerFirstName", "") or "").strip()
        last = (_safe_get_attr(latest_snapshot, "customerLastName", "") or "").strip()
        fullname = " ".join(p for p in (first, last) if p)

        address = ""
        addr_dict = (
                _safe_get_attr(latest_snapshot, "shipmentAddress", None)
                or _safe_get_attr(latest_snapshot, "invoiceAddress", None)
                or {}
        )
        if isinstance(addr_dict, dict):
            parts: list[str] = []
            for key in ("fullAddress", "address", "neighborhood", "district", "city", "postalCode"):
                v = addr_dict.get(key)
                if v:
                    parts.append(str(v).strip())
            address = ", ".join(parts)

        # -----------------------------
        # CARGO + SLA
        # -----------------------------
        cargo_tracking = (_safe_get_attr(latest_snapshot, "cargoTrackingNumber", "") or "").strip()
        cargo_provider_name = (_safe_get_attr(latest_snapshot, "cargoProviderName", "") or "").strip()

        agreed_delivery_ms = _safe_get_attr(latest_snapshot, "agreedDeliveryDate", None)
        try:
            agreed_delivery_ms = int(agreed_delivery_ms) if agreed_delivery_ms is not None else None
        except (TypeError, ValueError):
            agreed_delivery_ms = None

        # -----------------------------
        # BASE LABEL DICT
        # -----------------------------
        label_dict: Dict[str, Any] = {
            "orderNumber": order_no,
            "cargoTrackingNumber": cargo_tracking,
            "fullname": fullname,
            "address": address,
            "cargoProviderName": cargo_provider_name,

            # bunlar dışarıdan enjekte edilecek ama anahtarları şimdiden hazır olsun
            "storeName": "",
            "platform": "",

            "agreedDeliveryDate_ms": agreed_delivery_ms,
            "debug_index": debug_index,
            "is_primary_for_order": (chunk_idx == 0),
        }

        # -----------------------------
        # PRODUCTS / QTY SLOTS
        # -----------------------------
        for idx in range(max_items_per_label):
            p_key = f"prod{idx + 1}"
            q_key = f"qty{idx + 1}"
            if idx < len(chunk):
                label_dict[p_key] = chunk[idx].get("name", "")
                label_dict[q_key] = chunk[idx].get("qty", "")
            else:
                label_dict[p_key] = ""
                label_dict[q_key] = ""

        return Result.ok("Label dict OK.", close_dialog=False, data={"label": label_dict})

    except Exception as e:
        return Result.fail(map_error_to_message(e), error=e, close_dialog=False)


def _paginate(labels: List[Dict[str, Any]], *, labels_per_page: int) -> Result:
    try:
        if labels_per_page <= 0:
            labels_per_page = 24
        pages = [labels[i:i + labels_per_page] for i in range(0, len(labels), labels_per_page)]
        return Result.ok("Paginate OK.", close_dialog=False, data={"pages": pages})
    except Exception as e:
        return Result.fail(map_error_to_message(e), error=e, close_dialog=False)


# ============================================================
# 2.7) CREATE LABEL PAYLOAD
# ============================================================
def create_order_label_from_orders(
        order_numbers: list[str],
        *,
        brand_code: str = "TANEX",
        model_code: str = "TANEX_2736",
        order_refs: list[dict] | None = None,
        db_name: str | None = None,
) -> Result:
    """
    Pipeline (UI bağımsız).
    Input: order_numbers (list[str])
    Output: label_payload + order_numbers

    Not:
      - DetachedInstanceError yaşamamak için snapshot üzerinden relationship okumaya güvenmiyoruz.
      - store_name/platform bilgisi Orders pipeline'dan (pkg["store_name"], pkg["platform"]) gelir
        ve label dict'e sonradan enjekte edilir (✅ detached yok, plain string).
    """
    try:
        vres = validate_label_model_config(brand_code, model_code)
        if not vres.success:
            return vres
        cfg = vres.data.get("cfg") or {}

        max_items_per_label: int = int(cfg.get("max_items_per_label", 8))
        labels_per_page: int = int(cfg.get("labels_per_page", 24))

        normalized_numbers = [str(n).strip() for n in (order_numbers or []) if str(n).strip()]
        normalized_numbers = list(dict.fromkeys(normalized_numbers))  # uniq, order-preserving
        if not normalized_numbers:
            return Result.fail("Sipariş numarası listesi boş.", close_dialog=False)

        # ✅ Orders full details (LATEST snapshot mode)
        detail_res = get_order_full_details_by_numbers(
            normalized_numbers,
            mode="latest",
            order_refs=order_refs,
            db_name=db_name,
        )
        if not detail_res or not isinstance(detail_res, Result):
            return Result.fail("Sipariş detayları alınamadı.", close_dialog=False)
        if not detail_res.success:
            return detail_res

        orders = (detail_res.data or {}).get("orders", []) or []

        final_labels: List[Dict[str, Any]] = []
        order_refs: List[Dict[str, Any]] = []
        seen_refs: set[tuple[str, Optional[int]]] = set()
        debug_idx = 0

        for pkg in orders:
            header = pkg.get("header")
            snapshots = (pkg.get("data") or [])  # mode="latest" => 0/1
            items = pkg.get("items", []) or []

            latest_snapshot = snapshots[0] if snapshots else None
            if latest_snapshot is None:
                ord_no = (getattr(header, "orderNumber", "") or "").strip() if header else ""
                return Result.fail(
                    f"Latest snapshot bulunamadı. (orderNumber={ord_no or 'UNKNOWN'})",
                    close_dialog=False,
                )

            ref_order_no = (
                (getattr(header, "orderNumber", None) if header else None)
                or getattr(latest_snapshot, "orderNumber", None)
                or ""
            )
            ref_order_no = str(ref_order_no).strip()

            ref_api_account_id = (
                (getattr(header, "api_account_id", None) if header else None)
                or getattr(latest_snapshot, "api_account_id", None)
            )
            try:
                ref_api_account_id = int(ref_api_account_id) if ref_api_account_id is not None else None
            except Exception:
                ref_api_account_id = None

            if ref_order_no:
                ref_key = (ref_order_no, ref_api_account_id)
                if ref_key not in seen_refs:
                    seen_refs.add(ref_key)
                    order_refs.append(
                        {
                            "orderNumber": ref_order_no,
                            "api_account_id": ref_api_account_id,
                        }
                    )

            # ✅ store/platform: ApiAccount üzerinden hesaplanıp pkg'ye konmuş plain string
            store_name = (pkg.get("store_name") or "").strip()
            platform = (pkg.get("platform") or "").strip()

            norm_res = _normalize_items(items, latest_snapshot=latest_snapshot)
            if not norm_res.success:
                return norm_res
            normalized_items = norm_res.data["items"]

            chunk_res = _chunk_items(normalized_items, max_items_per_label=max_items_per_label)
            if not chunk_res.success:
                return chunk_res
            chunks = chunk_res.data["chunks"]

            for chunk_idx, chunk in enumerate(chunks):
                debug_idx += 1

                lbl_res = _build_label_dict(
                    latest_snapshot=latest_snapshot,
                    chunk=chunk,
                    chunk_idx=chunk_idx,
                    max_items_per_label=max_items_per_label,
                    debug_index=debug_idx,
                )
                if not lbl_res.success:
                    return lbl_res

                label = lbl_res.data["label"]

                # ✅ inject (detached yok, relationship yok)
                label["storeName"] = store_name
                label["platform"] = platform

                final_labels.append(label)

        pages_res = _paginate(final_labels, labels_per_page=labels_per_page)
        if not pages_res.success:
            return pages_res
        pages = pages_res.data["pages"]

        payload = {
            "brand_code": brand_code,
            "model_code": model_code,
            "max_items_per_label": max_items_per_label,
            "labels_per_page": labels_per_page,
            "total_labels": len(final_labels),
            "total_pages": len(pages),
            "pages": pages,
        }

        return Result.ok(
            f"{len(normalized_numbers)} sipariş için {len(final_labels)} label hazırlandı.",
            close_dialog=False,
            data={
                "label_payload": payload,
                "order_numbers": normalized_numbers,
                "order_refs": order_refs,
            },
        )

    except Exception as e:
        return Result.fail(map_error_to_message(e), error=e, close_dialog=False)


# ============================================================
# 3) BARKOD (STRICT TEMPLATE)
# ============================================================
def generate_code128_barcode(
        *,
        value: str,
        barcode_cfg: Dict[str, Any],
        dpi: int = 300,
        save_path: Optional[str] = None,
        return_pil: bool = False,
) -> Result:
    """
    Code128 barkod üretir.
    - TÜM render parametreleri barcode_cfg (template) üzerinden gelir.
    - Eksik parametre varsa üretim YAPMAZ.
    """
    try:
        if not value or not isinstance(value, str):
            return Result.fail("Geçersiz barkod değeri.", close_dialog=False)

        if not isinstance(barcode_cfg, dict):
            return Result.fail("barcode_cfg dict olmalı (template'ten gelmeli).", close_dialog=False)

        required_keys = {
            "module_width",
            "module_height",
            "font_size",
            "text_distance",
            "quiet_zone",
        }
        missing = required_keys - barcode_cfg.keys()
        if missing:
            return Result.fail(
                f"Barkod config eksik: {', '.join(sorted(missing))}",
                close_dialog=False,
            )

        try:
            _barcode_runtime()
        except Exception:
            return Result.fail(
                "Barkod kütüphaneleri eksik. 'python-barcode' ve 'Pillow' kurun.",
                close_dialog=False,
            )

        barcode_cfg_json = json.dumps(barcode_cfg, ensure_ascii=True, sort_keys=True)
        png_bytes = _render_code128_png_cached(str(value), barcode_cfg_json, int(dpi))
        pil_img = None
        if return_pil:
            try:
                from PIL import Image

                pil_img = Image.open(BytesIO(png_bytes)).copy()
            except Exception as image_error:
                return Result.fail("Barkod görseli çözümlenemedi.", error=image_error, close_dialog=False)

        if save_path:
            with open(save_path, "wb") as f:
                f.write(png_bytes)

        return Result.ok(
            "Barkod üretildi.",
            close_dialog=False,
            data={
                "png_bytes": png_bytes,
                "pil_image": pil_img if return_pil else None,
                "width": pil_img.size[0] if pil_img is not None else None,
                "height": pil_img.size[1] if pil_img is not None else None,
                "dpi": dpi,
                "path": save_path,
            },
        )

    except Exception as e:
        return Result.fail(map_error_to_message(e), error=e, close_dialog=False)


# ============================================================
# 4) DOCX RICH TEXT + STYLE
# ============================================================
def _make_rich_text(
        value: str,
        *,
        font_name: str | None = None,
        font_size: int | None = None,
        color: str | None = None,
        bold: bool | None = None,
) -> Result:
    try:
        if value in (None, ""):
            return Result.ok("Empty text.", close_dialog=False, data={"rt": ""})

        _, _, RichText, _ = _docxtpl_runtime()
        rt = RichText()
        kwargs = {}
        if font_name:
            kwargs["font"] = font_name
        if font_size:
            kwargs["size"] = font_size
        if color:
            kwargs["color"] = color
        if bold is not None:
            kwargs["bold"] = bold

        rt.add(str(value), **kwargs)
        return Result.ok("RichText OK.", close_dialog=False, data={"rt": rt})

    except Exception as e:
        return Result.fail(map_error_to_message(e), error=e, close_dialog=False)


def _style_text(value: str, st: dict) -> Result:
    try:
        if not value:
            return Result.ok("Empty style text.", close_dialog=False, data={"val": ""})

        rt_res = _make_rich_text(
            value,
            font_name=st.get("font_name"),
            font_size=st.get("font_size"),
            color=st.get("color"),
            bold=st.get("bold"),
        )
        if not rt_res.success:
            return rt_res
        return Result.ok("Style text OK.", close_dialog=False, data={"val": rt_res.data["rt"]})
    except Exception as e:
        return Result.fail(map_error_to_message(e), error=e, close_dialog=False)


def _get_styles(cfg: dict) -> Result:
    try:
        field_styles = cfg.get("fields", {}) or {}

        def fs(k: str):
            return field_styles.get(k, {}) or {}

        styles = {
            "ordernumber": fs("ordernumber"),
            "name": fs("name"),
            "surname": fs("surname"),
            "address": fs("address"),
            "cargotrackingnumber": fs("cargotrackingnumber"),
            "product": fs("product"),
            "qty": fs("qty"),
            "storename": fs("storename"),
            "platform": fs("platform"),
        }
        return Result.ok("Styles OK.", close_dialog=False, data={"styles": styles})
    except Exception as e:
        return Result.fail(map_error_to_message(e), error=e, close_dialog=False)


def _safe_inline_image(doc: DocxTemplate, *, path: Path, width_mm: float, height_mm: float | None = None) -> Result:
    try:
        if not path.is_file():
            return Result.ok("Image yok.", close_dialog=False, data={"image": None})

        _, InlineImage, _, Mm = _docxtpl_runtime()
        kwargs = {"width": Mm(width_mm)}
        if height_mm is not None:
            kwargs["height"] = Mm(height_mm)

        img = InlineImage(doc, BytesIO(_read_inline_image_bytes(str(path.resolve()))), **kwargs)
        return Result.ok("Image OK.", close_dialog=False, data={"image": img})

    except Exception as e:
        return Result.fail(map_error_to_message(e), error=e, close_dialog=False)


def _style_fitted_text(value: str, st: dict) -> Result:
    """Fit variable-length label text without changing the template geometry.

    `fit_chars_at_base` describes how much text the fixed template cell can hold
    at its normal font size. Longer values are reduced only as far as the
    configured minimum, then safely shortened so Word cannot wrap them outside
    the millimetric label cell.
    """

    try:
        text = " ".join(str(value or "").split())
        if not text:
            return Result.ok("Empty fitted text.", close_dialog=False, data={"val": ""})

        fitted_style = dict(st or {})
        try:
            base_size = max(1, int(round(float(fitted_style.get("font_size") or 0))))
            base_capacity = max(0, int(fitted_style.get("fit_chars_at_base") or 0))
            min_size = max(1, int(round(float(fitted_style.get("min_font_size") or base_size))))
        except (TypeError, ValueError):
            base_size = base_capacity = 0
            min_size = 1

        if base_size > 0 and base_capacity > 0 and len(text) > base_capacity:
            fitted_size = max(min_size, min(base_size, int((base_size * base_capacity) / len(text))))
            fitted_style["font_size"] = fitted_size
            fitted_capacity = max(base_capacity, int((base_capacity * base_size) / fitted_size))
            if len(text) > fitted_capacity:
                text = text[: max(1, fitted_capacity - 3)].rstrip() + "..."

        return _style_text(text, fitted_style)
    except Exception as e:
        return Result.fail(map_error_to_message(e), error=e, close_dialog=False)


# ============================================================
# 5) SLA (parametreli)
# ============================================================
def _compute_sla_bucket(*, agreed_ms: Optional[int], now_utc: Optional[datetime] = None) -> Result:
    try:
        if agreed_ms is None:
            return Result.ok("SLA yok.", close_dialog=False, data={"delta_hours": None})

        try:
            agreed_ms_int = int(agreed_ms)
        except (TypeError, ValueError):
            return Result.ok("SLA parse edilemedi.", close_dialog=False, data={"delta_hours": None})

        deadline_utc = datetime.fromtimestamp(agreed_ms_int / 1000.0, timezone.utc).replace(tzinfo=None)
        now_utc = now_utc or datetime.now(timezone.utc).replace(tzinfo=None)
        delta_hours = (deadline_utc - now_utc).total_seconds() / 3600.0

        return Result.ok("Delta hours OK.", close_dialog=False, data={"delta_hours": delta_hours})

    except Exception as e:
        return Result.fail(map_error_to_message(e), error=e, close_dialog=False)


def _compute_sla_image(
        *,
        doc: DocxTemplate,
        agreed_ms: Optional[int],
        sla_width_mm: float,
        finish_img_path: Path,
        enough_img_path: Path,
        threshold_hours: float,
        mode: str = "lte",
) -> Result:
    try:
        dh_res = _compute_sla_bucket(agreed_ms=agreed_ms)
        if not dh_res.success:
            return dh_res

        delta_hours = dh_res.data.get("delta_hours")
        if delta_hours is None:
            return Result.ok("SLA image yok.", close_dialog=False, data={"image": None, "delta_hours": None})

        if mode == "lt":
            use_finish = delta_hours < threshold_hours
        elif mode == "lte":
            use_finish = delta_hours <= threshold_hours
        elif mode == "gt":
            use_finish = delta_hours > threshold_hours
        elif mode == "gte":
            use_finish = delta_hours >= threshold_hours
        else:
            use_finish = delta_hours <= threshold_hours

        target = finish_img_path if use_finish else enough_img_path
        img_res = _safe_inline_image(doc, path=target, width_mm=sla_width_mm)
        if not img_res.success:
            return img_res

        return Result.ok(
            "SLA image OK.",
            close_dialog=False,
            data={"image": img_res.data["image"], "delta_hours": delta_hours, "use_finish": use_finish},
        )

    except Exception as e:
        return Result.fail(map_error_to_message(e), error=e, close_dialog=False)


# ============================================================
# 6) PAYLOAD UNWRAP + FLATTEN
# ============================================================
def _unwrap_label_payload(maybe_payload: Any) -> Result:
    try:
        if isinstance(maybe_payload, dict):
            return Result.ok("Payload dict OK.", close_dialog=False, data={"label_payload": maybe_payload})

        if isinstance(maybe_payload, Result):
            if not maybe_payload.success:
                return maybe_payload

            data = maybe_payload.data or {}
            if isinstance(data.get("label_payload"), dict):
                return Result.ok(
                    "Payload unwrap OK (Result.data['label_payload']).",
                    close_dialog=False,
                    data={"label_payload": data["label_payload"]},
                )
            if isinstance(data, dict) and "pages" in data:
                return Result.ok("Payload unwrap OK (Result.data).", close_dialog=False, data={"label_payload": data})

            return Result.fail("Result geldi ama label_payload bulunamadı.", close_dialog=False)

        return Result.fail(f"label_payload tipi geçersiz: {type(maybe_payload).__name__}", close_dialog=False)

    except Exception as e:
        return Result.fail(map_error_to_message(e), error=e, close_dialog=False)


def _flatten_labels(maybe_payload: Any) -> Result:
    try:
        uw = _unwrap_label_payload(maybe_payload)
        if not uw.success:
            return uw

        payload = uw.data["label_payload"]
        pages = payload.get("pages") or []
        labels = [lbl for page in pages for lbl in page]

        return Result.ok("Flatten OK.", close_dialog=False, data={"labels": labels, "payload": payload})

    except Exception as e:
        return Result.fail(map_error_to_message(e), error=e, close_dialog=False)


# ============================================================
# 7) CONTEXT BUILD (STRICT BARCODE CFG)
# ============================================================
def _build_ctx_for_label(
        *,
        doc: DocxTemplate,
        cfg: dict,
        styles: dict,
        lbl: dict,
        cargo_logo_cfg_norm: dict,
        n: int,
        sla_threshold_hours: float,
        sla_mode: str,
        qty_attention_threshold: int,
        qty_attention_color: str,
        qty_attention_bold: bool,
) -> Result:
    try:
        max_items = int(cfg.get("max_items_per_label", 8))
        cargo_logo_cfg = cargo_logo_cfg_norm or {}

        # -------------------------------------------------
        # ✅ BARCODE CFG (STRICT)
        # -------------------------------------------------
        barcode_cfg = cfg.get("barcode")
        if not isinstance(barcode_cfg, dict):
            return Result.fail("Template barkod config içermiyor (cfg['barcode'] yok).", close_dialog=False)

        # Render width'ler template contract
        if "image_width_mm" not in barcode_cfg:
            return Result.fail("Barkod config eksik: image_width_mm", close_dialog=False)
        if "attention_image_width_mm" not in barcode_cfg:
            return Result.fail("Barkod config eksik: attention_image_width_mm", close_dialog=False)

        barcode_width = float(barcode_cfg["image_width_mm"])
        attention_w = float(barcode_cfg["attention_image_width_mm"])

        attention_h = barcode_cfg.get("attention_image_height_mm")
        attention_h = float(attention_h) if attention_h is not None else None
        attention_path = LABEL_IMAGES_DIR / "split_order_attention_img.png"

        # -------------------------------------------------
        # SLA images
        # -------------------------------------------------
        sla_width_mm = float(cfg.get("sla_image_width_mm", 10))
        finish_img_path = LABEL_IMAGES_DIR / "finish_time_attention.png"
        enough_img_path = LABEL_IMAGES_DIR / "enough_time_info.png"

        # -------------------------------------------------
        # Label values
        # -------------------------------------------------
        order_no = (lbl.get("orderNumber") or "").strip()
        fullname = (lbl.get("fullname") or "").strip()
        address_val = (lbl.get("address") or "").strip()

        cargo_provider = (lbl.get("cargoProviderName") or "").strip()
        cargo_raw = (lbl.get("cargoTrackingNumber") or "").strip()

        store_name = (lbl.get("storeName") or "").strip()
        platform_val = (lbl.get("platform") or "").strip()

        agreed_ms = lbl.get("agreedDeliveryDate_ms")
        barcode_val = cargo_raw or order_no
        is_primary = bool(lbl.get("is_primary_for_order", True))

        ctx: Dict[str, Any] = {}

        s = _style_text(order_no, styles["ordernumber"])
        if not s.success:
            return s
        ctx[f"ordernumber_{n}"] = s.data["val"]

        s = _style_text(fullname, styles["name"])
        if not s.success:
            return s
        ctx[f"name_{n}"] = s.data["val"]

        ctx[f"surname_{n}"] = ""

        s = _style_fitted_text(address_val, styles["address"])
        if not s.success:
            return s
        ctx[f"address_{n}"] = s.data["val"]

        s = _style_text(barcode_val, styles["cargotrackingnumber"])
        if not s.success:
            return s
        ctx[f"cargotrackingnumber_{n}"] = s.data["val"]

        s = _style_text(store_name, styles["storename"])
        if not s.success:
            return s
        ctx[f"storename_{n}"] = s.data["val"]

        s = _style_text(platform_val, styles["platform"])
        if not s.success:
            return s
        ctx[f"platform_{n}"] = s.data["val"]

        sla_res = _compute_sla_image(
            doc=doc,
            agreed_ms=agreed_ms,
            sla_width_mm=sla_width_mm,
            finish_img_path=finish_img_path,
            enough_img_path=enough_img_path,
            threshold_hours=float(sla_threshold_hours),
            mode=sla_mode,
        )
        if not sla_res.success:
            return sla_res
        ctx[f"sla_hours_left_{n}"] = sla_res.data.get("image") or ""

        # -------------------------------------------------
        # Cargo logo
        # -------------------------------------------------
        cargo_logo_el = None
        if cargo_provider:
            info = cargo_logo_cfg.get(_normalize_cargo_name(cargo_provider))
            if info:
                logo_path = LABEL_IMAGES_DIR / info["filename"]
                img_res = _safe_inline_image(doc, path=logo_path, width_mm=float(info.get("width_mm", 12)))
                if not img_res.success:
                    return img_res
                cargo_logo_el = img_res.data.get("image")

        ctx[f"cargoprovidername_{n}"] = cargo_logo_el if cargo_logo_el else ""

        # -------------------------------------------------
        # Barcode OR attention image
        # -------------------------------------------------
        if is_primary:
            res_bar = generate_code128_barcode(value=barcode_val, barcode_cfg=barcode_cfg, dpi=300, save_path=None)
            if not res_bar.success:
                return Result.fail(
                    f"Barkod üretilemedi: {res_bar.message}",
                    error=res_bar.error,
                    close_dialog=False,
                )
            try:
                _, InlineImage, _, Mm = _docxtpl_runtime()
                ctx[f"barcode_{n}"] = InlineImage(
                    doc,
                    BytesIO(res_bar.data["png_bytes"]),
                    width=Mm(barcode_width),
                )
            except Exception as e:
                return Result.fail("Barkod görseli etikete yerleştirilemedi.", error=e, close_dialog=False)
        else:
            if attention_path.is_file():
                _, InlineImage, _, Mm = _docxtpl_runtime()
                kwargs = {"width": Mm(attention_w)}
                if attention_h is not None:
                    kwargs["height"] = Mm(attention_h)
                ctx[f"barcode_{n}"] = InlineImage(doc, str(attention_path), **kwargs)
            else:
                ctx[f"barcode_{n}"] = ""

        # -------------------------------------------------
        # Products + qty
        # -------------------------------------------------
        style_product = styles["product"]
        style_qty = styles["qty"]

        for i in range(1, max_items + 1):
            pkey = f"prod{i}"
            qkey = f"qty{i}"

            pv = (lbl.get(pkey, "") or "").strip()
            qv = lbl.get(qkey, "")

            stp = _style_fitted_text(pv, style_product)
            if not stp.success:
                return stp
            ctx[f"{pkey}_{n}"] = stp.data["val"]

            if not qv:
                ctx[f"{qkey}_{n}"] = ""
                continue

            try:
                qint = int(qv)
            except Exception:
                qint = 0

            if qint >= int(qty_attention_threshold):
                color = qty_attention_color
                bold = qty_attention_bold
            else:
                color = style_qty.get("color")
                bold = style_qty.get("bold")

            rt_res = _make_rich_text(
                str(qv),
                font_name=style_qty.get("font_name"),
                font_size=style_qty.get("font_size"),
                color=color,
                bold=bold,
            )
            if not rt_res.success:
                return rt_res
            ctx[f"{qkey}_{n}"] = rt_res.data["rt"]

        return Result.ok("Ctx OK.", close_dialog=False, data={"ctx": ctx})

    except Exception as e:
        return Result.fail(map_error_to_message(e), error=e, close_dialog=False)


def _fill_empty_slots(ctx: dict, *, start_n: int, labels_per_page: int, max_items: int) -> Result:
    try:
        for n in range(start_n, labels_per_page + 1):
            ctx[f"ordernumber_{n}"] = ""
            ctx[f"name_{n}"] = ""
            ctx[f"surname_{n}"] = ""
            ctx[f"address_{n}"] = ""
            ctx[f"cargotrackingnumber_{n}"] = ""
            ctx[f"cargoprovidername_{n}"] = ""
            ctx[f"barcode_{n}"] = ""
            ctx[f"storename_{n}"] = ""
            ctx[f"platform_{n}"] = ""
            ctx[f"sla_hours_left_{n}"] = ""
            for i in range(1, max_items + 1):
                ctx[f"prod{i}_{n}"] = ""
                ctx[f"qty{i}_{n}"] = ""
        return Result.ok("Empty slots OK.", close_dialog=False)
    except Exception as e:
        return Result.fail(map_error_to_message(e), error=e, close_dialog=False)


def _render_page(
        *,
        template_path: Path,
        page_labels: List[dict],
        cfg: dict,
        sla_threshold_hours: float,
        sla_mode: str,
        qty_attention_threshold: int,
        qty_attention_color: str,
        qty_attention_bold: bool,
) -> Result:
    try:
        try:
            DocxTemplate, _, _, _ = _docxtpl_runtime()
            doc = DocxTemplate(str(template_path))
        except Exception as e:
            return Result.fail("Şablon yüklenemedi.", error=e, close_dialog=False)

        styles_res = _get_styles(cfg)
        if not styles_res.success:
            return styles_res
        styles = styles_res.data["styles"]

        ctx: Dict[str, Any] = {}
        cargo_logo_cfg_norm = _normalize_cargo_logo_map(
            cfg.get("cargo_provider_logos", {}) or {},
            cfg.get("cargo_provider_aliases", {}) or {},
        )

        for slot, lbl in enumerate(page_labels, start=1):
            one_res = _build_ctx_for_label(
                doc=doc,
                cfg=cfg,
                styles=styles,
                lbl=lbl,
                cargo_logo_cfg_norm=cargo_logo_cfg_norm,
                n=slot,
                sla_threshold_hours=sla_threshold_hours,
                sla_mode=sla_mode,
                qty_attention_threshold=qty_attention_threshold,
                qty_attention_color=qty_attention_color,
                qty_attention_bold=qty_attention_bold,
            )
            if not one_res.success:
                return one_res
            ctx.update(one_res.data["ctx"])

        labels_per_page = int(cfg.get("labels_per_page", 24))
        max_items = int(cfg.get("max_items_per_label", 8))

        fill_res = _fill_empty_slots(
            ctx,
            start_n=len(page_labels) + 1,
            labels_per_page=labels_per_page,
            max_items=max_items,
        )
        if not fill_res.success:
            return fill_res

        try:
            doc.render(ctx)
            return Result.ok("Sayfa render edildi.", close_dialog=False, data={"doc": doc})
        except Exception as e:
            return Result.fail("Sayfa oluşturulamadı.", error=e, close_dialog=False)

    except Exception as e:
        return Result.fail(map_error_to_message(e), error=e, close_dialog=False)


def _merge_docx_files(page_files: List[str], output_path: str) -> Result:
    try:
        if not page_files:
            return Result.fail("Merge için sayfa yok.", close_dialog=False)

        Document, Composer = _docx_merge_runtime()
        main_doc = Document(page_files[0])
        comp = Composer(main_doc)
        for pf in page_files[1:]:
            comp.append(Document(pf))
        comp.save(output_path)

        return Result.ok("Merge tamam.", close_dialog=False, data={"output_path": output_path})
    except Exception as e:
        return Result.fail("Word dosyaları birleştirilirken hata oluştu.", error=e, close_dialog=False)


def _merge_pages(page_files: List[str], output_path: str, *, chunk_size: int = _LABEL_EXPORT_MERGE_CHUNK_SIZE) -> Result:
    try:
        files = [str(Path(path)) for path in list(page_files or []) if str(path or "").strip()]
        if not files:
            return Result.fail("Merge için sayfa yok.", close_dialog=False)

        safe_chunk_size = max(2, int(chunk_size or _LABEL_EXPORT_MERGE_CHUNK_SIZE))
        if len(files) <= safe_chunk_size:
            return _merge_docx_files(files, output_path)

        with tempfile.TemporaryDirectory(prefix="orderscout_label_merge_") as temp_root:
            temp_dir = Path(temp_root)
            current_files = files
            pass_index = 0
            while len(current_files) > safe_chunk_size:
                next_files: list[str] = []
                for chunk_index in range(0, len(current_files), safe_chunk_size):
                    chunk = current_files[chunk_index: chunk_index + safe_chunk_size]
                    merged_path = temp_dir / f"merge_pass_{pass_index}_{chunk_index // safe_chunk_size}.docx"
                    merged_res = _merge_docx_files(chunk, str(merged_path))
                    if not merged_res.success:
                        return merged_res
                    next_files.append(str(merged_path))
                    gc.collect()
                current_files = next_files
                pass_index += 1

            return _merge_docx_files(current_files, output_path)
    except Exception as e:
        return Result.fail("Word dosyaları chunk'lı birleştirilirken hata oluştu.", error=e, close_dialog=False)


def _restore_output_outer_table_geometry(*, template_path: Path, output_path: Path) -> Result:
    """Restore the template's exact outer tblGrid after docxtpl/python-docx serialization.

    Cell widths survive rendering, but python-docx may recalculate `tblGrid` to a
    different column count. Label sheets are millimetric forms, so every emitted page
    must carry the source template grid instead of relying on Word's repair heuristics.
    """

    try:
        with zipfile.ZipFile(template_path, "r") as source_package:
            template_xml = source_package.read("word/document.xml")
        with zipfile.ZipFile(output_path, "r") as output_package:
            output_xml = output_package.read("word/document.xml")

        template_root = etree.fromstring(template_xml)
        output_root = etree.fromstring(output_xml)
        template_tables = template_root.xpath("/w:document/w:body/w:tbl", namespaces=_WORD_NS)
        output_tables = output_root.xpath("/w:document/w:body/w:tbl", namespaces=_WORD_NS)
        if len(template_tables) != 1 or not output_tables:
            return Result.fail("Etiket dış tablo geometrisi doğrulanamadı.", close_dialog=False)

        template_grid = template_tables[0].find(f"{{{_WORDPROCESSINGML_NS}}}tblGrid")
        if template_grid is None:
            return Result.fail("Etiket şablonunda tblGrid bulunamadı.", close_dialog=False)

        expected_row_count = len(template_tables[0].findall(f"{{{_WORDPROCESSINGML_NS}}}tr"))
        expected_cell_count = len(
            template_tables[0].xpath("./w:tr/w:tc", namespaces=_WORD_NS)
        )
        for table in output_tables:
            if len(table.findall(f"{{{_WORDPROCESSINGML_NS}}}tr")) != expected_row_count:
                return Result.fail("Etiket çıktı satır geometrisi şablonla eşleşmiyor.", close_dialog=False)
            if len(table.xpath("./w:tr/w:tc", namespaces=_WORD_NS)) != expected_cell_count:
                return Result.fail("Etiket çıktı hücre geometrisi şablonla eşleşmiyor.", close_dialog=False)
            current_grid = table.find(f"{{{_WORDPROCESSINGML_NS}}}tblGrid")
            replacement = deepcopy(template_grid)
            if current_grid is None:
                table.insert(0, replacement)
            else:
                table.replace(current_grid, replacement)

        patched_xml = etree.tostring(
            output_root,
            encoding="UTF-8",
            xml_declaration=True,
            standalone=True,
        )
        temporary_path = output_path.with_suffix(output_path.suffix + ".geometry.tmp")
        with zipfile.ZipFile(output_path, "r") as source_package, zipfile.ZipFile(
            temporary_path,
            "w",
        ) as target_package:
            for info in source_package.infolist():
                payload = patched_xml if info.filename == "word/document.xml" else source_package.read(info.filename)
                target_package.writestr(info, payload)
        os.replace(temporary_path, output_path)

        expected_grid = etree.tostring(template_grid)
        with zipfile.ZipFile(output_path, "r") as verified_package:
            verified_root = etree.fromstring(verified_package.read("word/document.xml"))
        verified_tables = verified_root.xpath("/w:document/w:body/w:tbl", namespaces=_WORD_NS)
        if any(
            etree.tostring(table.find(f"{{{_WORDPROCESSINGML_NS}}}tblGrid")) != expected_grid
            for table in verified_tables
        ):
            return Result.fail("Etiket tblGrid read-back doğrulaması başarısız.", close_dialog=False)

        return Result.ok(
            "Etiket dış tablo geometrisi şablondan read-back ile doğrulandı.",
            close_dialog=False,
            data={"page_table_count": len(verified_tables), "geometry_verified": True},
        )
    except Exception as e:
        return Result.fail("Etiket çıktı geometrisi korunamadı.", error=e, close_dialog=False)


# ============================================================
# 8) WORD EXPORT
# ============================================================
def export_labels_to_word(
        label_payload: Any,
        brand_code: str | None = None,
        model_code: str | None = None,
        output_path: str | None = None,
        *,
        template_path: str | None = None,
        progress_cb=None,
        sla_threshold_hours: float | None = None,
        sla_mode: str = "lte",
        qty_attention_threshold: int = 2,
        qty_attention_color: str = "FF0000",
        qty_attention_bold: bool = True,
) -> Result:
    try:
        def report(p: int):
            if progress_cb:
                try:
                    progress_cb(max(0, min(100, int(p))))
                except Exception:
                    pass

        report(0)

        flat_res = _flatten_labels(label_payload)
        if not flat_res.success:
            return flat_res

        labels: List[dict] = flat_res.data["labels"]
        payload: dict = flat_res.data["payload"]

        if not labels:
            return Result.fail("Yazdırılacak etiket bulunamadı.", close_dialog=False)

        if brand_code is None:
            brand_code = payload.get("brand_code")
        if model_code is None:
            model_code = payload.get("model_code")

        vres = validate_label_model_config(brand_code, model_code)
        if not vres.success:
            return vres
        cfg = vres.data.get("cfg") or {}

        tp = Path(template_path or cfg.get("template_path") or "")
        if not tp.is_file():
            return Result.fail(f"Word şablonu bulunamadı: {tp}", close_dialog=False)

        if not output_path:
            return Result.fail("output_path verilmedi.", close_dialog=False)

        if sla_threshold_hours is None:
            sla_threshold_hours = float(cfg.get("sla_threshold_hours", 24))

        labels_per_page = int(cfg.get("labels_per_page", 24))
        total = len(labels)
        total_pages = math.ceil(total / labels_per_page)

        with tempfile.TemporaryDirectory(prefix="orderscout_label_pages_") as temp_root:
            tmp_dir = Path(temp_root)
            page_files: List[str] = []

            for pidx in range(total_pages):
                start = pidx * labels_per_page
                end = min(start + labels_per_page, total)
                page_labels = labels[start:end]

                page_res = _render_page(
                    template_path=tp,
                    page_labels=page_labels,
                    cfg=cfg,
                    sla_threshold_hours=float(sla_threshold_hours),
                    sla_mode=sla_mode,
                    qty_attention_threshold=int(qty_attention_threshold),
                    qty_attention_color=str(qty_attention_color),
                    qty_attention_bold=bool(qty_attention_bold),
                )
                if not page_res.success:
                    return Result.fail(
                        f"Sayfa render edilemedi. (sayfa={pidx + 1})",
                        error=page_res.error,
                        close_dialog=False,
                    )

                outp = tmp_dir / f"page_{pidx + 1}.docx"
                try:
                    page_res.data["doc"].save(outp)
                except Exception as e:
                    return Result.fail(
                        f"Sayfa kaydedilemedi. (sayfa={pidx + 1})",
                        error=e,
                        close_dialog=False,
                    )

                page_files.append(str(outp))
                del page_res
                gc.collect()
                report(5 + (pidx + 1) * 90 / total_pages)

            merge_res = _merge_pages(page_files, str(output_path))
            if not merge_res.success:
                return merge_res

            geometry_res = _restore_output_outer_table_geometry(
                template_path=tp,
                output_path=Path(output_path),
            )
            if not geometry_res.success:
                return geometry_res

        report(100)
        return Result.ok(
            f"{total} etiket başarıyla oluşturuldu.",
            close_dialog=False,
            data={
                "output_path": str(output_path),
                "geometry_verified": True,
                "page_table_count": int((geometry_res.data or {}).get("page_table_count", 0) or 0),
            },
        )

    except Exception as e:
        return Result.fail("Beklenmeyen hata.", error=e, close_dialog=False)
