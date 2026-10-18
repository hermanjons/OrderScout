from __future__ import annotations

import hashlib
import json
import math
import re
from dataclasses import dataclass
from datetime import date, datetime, timezone
from decimal import Decimal
from pathlib import Path
from typing import Any, Iterable, Mapping, Sequence
from uuid import uuid4


REPORT_SNAPSHOT_SCHEMA_VERSION = 1
_MANIFEST_SHEET = "_OrderScout"
_DANGEROUS_CELL_PREFIXES = ("=", "+", "-", "@")


def _text(value: Any) -> str:
    return str(value or "").strip()


def _json_value(value: Any) -> Any:
    if value is None or isinstance(value, (str, bool, int)):
        return value
    if isinstance(value, float):
        if not math.isfinite(value):
            raise ValueError("Report snapshot cannot contain non-finite numbers.")
        return value
    if isinstance(value, Decimal):
        return float(value)
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, Mapping):
        return {
            str(key): _json_value(child)
            for key, child in sorted(value.items(), key=lambda item: str(item[0]))
        }
    if isinstance(value, (list, tuple)):
        return [_json_value(child) for child in value]
    return str(value)


def _canonical_json(value: Any) -> str:
    return json.dumps(
        _json_value(value),
        ensure_ascii=False,
        sort_keys=True,
        separators=(",", ":"),
    )


def _content_hash(payload: Mapping[str, Any]) -> str:
    return hashlib.sha256(_canonical_json(payload).encode("utf-8")).hexdigest()


def _snapshot_cell_value(value: Any, column: "ReportColumn") -> Any:
    normalized = _json_value(value)
    if value is None:
        return None
    if column.data_type in {"string", "date", "datetime"}:
        return str(normalized)
    return normalized


def _safe_sheet_name(value: Any) -> str:
    name = _text(value) or "Rapor"
    for token in ("[", "]", ":", "*", "?", "/", "\\"):
        name = name.replace(token, "-")
    name = name[:31].strip() or "Rapor"
    if name == _MANIFEST_SHEET:
        raise ValueError(f"{_MANIFEST_SHEET} is reserved for the report manifest.")
    return name


@dataclass(frozen=True)
class ReportColumn:
    key: str
    label: str
    data_type: str = "string"
    number_format: str = "General"
    width: float = 16.0

    def __post_init__(self) -> None:
        if not _text(self.key):
            raise ValueError("Report column key is required.")
        if self.data_type not in {"string", "integer", "number", "boolean", "date", "datetime"}:
            raise ValueError(f"Unsupported report column type: {self.data_type}")
        object.__setattr__(self, "key", _text(self.key))
        object.__setattr__(self, "label", _text(self.label) or _text(self.key))
        object.__setattr__(self, "number_format", _text(self.number_format) or "General")
        object.__setattr__(self, "width", max(8.0, min(60.0, float(self.width or 16.0))))

    def to_dict(self) -> dict[str, Any]:
        return {
            "key": self.key,
            "label": self.label,
            "data_type": self.data_type,
            "number_format": self.number_format,
            "width": self.width,
        }

    @classmethod
    def from_dict(cls, payload: Mapping[str, Any]) -> "ReportColumn":
        return cls(
            key=_text(payload.get("key")),
            label=_text(payload.get("label")),
            data_type=_text(payload.get("data_type")) or "string",
            number_format=_text(payload.get("number_format")) or "General",
            width=float(payload.get("width") or 16.0),
        )


@dataclass(frozen=True)
class ReportSheetSnapshot:
    name: str
    columns: tuple[ReportColumn, ...]
    rows: tuple[tuple[Any, ...], ...]

    def __post_init__(self) -> None:
        object.__setattr__(self, "name", _safe_sheet_name(self.name))
        if not self.columns:
            raise ValueError("Report sheet must have at least one column.")
        keys = [column.key for column in self.columns]
        if len(keys) != len(set(keys)):
            raise ValueError("Report sheet column keys must be unique.")
        width = len(self.columns)
        normalized_rows: list[tuple[Any, ...]] = []
        for row in self.rows:
            if len(row) != width:
                raise ValueError("Report row width does not match its column schema.")
            normalized_rows.append(
                tuple(
                    _snapshot_cell_value(value, column)
                    for value, column in zip(row, self.columns)
                )
            )
        object.__setattr__(self, "rows", tuple(normalized_rows))

    def to_dict(self) -> dict[str, Any]:
        return {
            "name": self.name,
            "columns": [column.to_dict() for column in self.columns],
            "rows": [list(row) for row in self.rows],
        }

    @classmethod
    def from_dict(cls, payload: Mapping[str, Any]) -> "ReportSheetSnapshot":
        return cls(
            name=_text(payload.get("name")),
            columns=tuple(
                ReportColumn.from_dict(column)
                for column in list(payload.get("columns") or [])
                if isinstance(column, Mapping)
            ),
            rows=tuple(
                tuple(row)
                for row in list(payload.get("rows") or [])
                if isinstance(row, (list, tuple))
            ),
        )


@dataclass(frozen=True)
class ReportSnapshot:
    snapshot_id: str
    content_hash: str
    title: str
    source_domain: str
    report_type: str
    created_at: str
    filters_json: str
    metadata_json: str
    sheets: tuple[ReportSheetSnapshot, ...]
    schema_version: int = REPORT_SNAPSHOT_SCHEMA_VERSION

    @property
    def filters(self) -> dict[str, Any]:
        value = json.loads(self.filters_json or "{}")
        return dict(value) if isinstance(value, dict) else {}

    @property
    def metadata(self) -> dict[str, Any]:
        value = json.loads(self.metadata_json or "{}")
        return dict(value) if isinstance(value, dict) else {}

    def content_payload(self) -> dict[str, Any]:
        return {
            "schema_version": self.schema_version,
            "title": self.title,
            "source_domain": self.source_domain,
            "report_type": self.report_type,
            "filters": self.filters,
            "metadata": self.metadata,
            "sheets": [sheet.to_dict() for sheet in self.sheets],
        }

    def to_dict(self) -> dict[str, Any]:
        return {
            "snapshot_id": self.snapshot_id,
            "content_hash": self.content_hash,
            "created_at": self.created_at,
            **self.content_payload(),
        }

    @classmethod
    def create(
        cls,
        *,
        title: str,
        source_domain: str,
        report_type: str,
        filters: Mapping[str, Any] | None,
        metadata: Mapping[str, Any] | None,
        sheets: Sequence[ReportSheetSnapshot],
        created_at: str | None = None,
    ) -> "ReportSnapshot":
        clean_sheets = tuple(sheets)
        if not clean_sheets:
            raise ValueError("Report snapshot must have at least one sheet.")
        names = [sheet.name.casefold() for sheet in clean_sheets]
        if len(names) != len(set(names)):
            raise ValueError("Report snapshot sheet names must be unique.")
        content = {
            "schema_version": REPORT_SNAPSHOT_SCHEMA_VERSION,
            "title": _text(title) or "OrderScout Raporu",
            "source_domain": _text(source_domain).casefold(),
            "report_type": _text(report_type),
            "filters": _json_value(dict(filters or {})),
            "metadata": _json_value(dict(metadata or {})),
            "sheets": [sheet.to_dict() for sheet in clean_sheets],
        }
        digest = _content_hash(content)
        return cls(
            snapshot_id=f"rpt_{digest[:24]}",
            content_hash=digest,
            title=content["title"],
            source_domain=content["source_domain"],
            report_type=content["report_type"],
            created_at=_text(created_at) or datetime.now(timezone.utc).isoformat(),
            filters_json=_canonical_json(content["filters"]),
            metadata_json=_canonical_json(content["metadata"]),
            sheets=clean_sheets,
        )

    @classmethod
    def from_dict(cls, payload: Mapping[str, Any]) -> "ReportSnapshot":
        sheets = tuple(
            ReportSheetSnapshot.from_dict(sheet)
            for sheet in list(payload.get("sheets") or [])
            if isinstance(sheet, Mapping)
        )
        snapshot = cls.create(
            title=_text(payload.get("title")),
            source_domain=_text(payload.get("source_domain")),
            report_type=_text(payload.get("report_type")),
            filters=dict(payload.get("filters") or {}),
            metadata=dict(payload.get("metadata") or {}),
            sheets=sheets,
            created_at=_text(payload.get("created_at")),
        )
        supplied_hash = _text(payload.get("content_hash"))
        supplied_id = _text(payload.get("snapshot_id"))
        if supplied_hash and supplied_hash != snapshot.content_hash:
            raise ValueError("Report snapshot content hash mismatch.")
        if supplied_id and supplied_id != snapshot.snapshot_id:
            raise ValueError("Report snapshot id mismatch.")
        return snapshot


def make_report_sheet(
    *,
    name: str,
    columns: Sequence[ReportColumn],
    rows: Iterable[Mapping[str, Any] | Sequence[Any]],
) -> ReportSheetSnapshot:
    clean_columns = tuple(columns)
    normalized: list[tuple[Any, ...]] = []
    for row in rows:
        if isinstance(row, Mapping):
            normalized.append(tuple(row.get(column.key) for column in clean_columns))
        else:
            normalized.append(tuple(row))
    return ReportSheetSnapshot(name=name, columns=clean_columns, rows=tuple(normalized))


def _xlsx_cell_value(value: Any, column: ReportColumn) -> Any:
    if value is None:
        return None
    if column.data_type == "integer":
        return int(value)
    if column.data_type == "number":
        return float(value)
    if column.data_type == "boolean":
        return bool(value)
    text = str(value)
    if text.startswith(_DANGEROUS_CELL_PREFIXES):
        return "'" + text
    if column.data_type == "string" and re.fullmatch(r"\d{8,}", text):
        # Some spreadsheet renderers coerce even an OOXML shared string with
        # number format "@" into scientific notation.  An inline rich-text run
        # remains a real, copy-safe string (no hidden prefix) in Excel while
        # preserving every identifier digit in renderers and exact readback.
        from openpyxl.cell.rich_text import CellRichText, TextBlock
        from openpyxl.cell.text import InlineFont

        return CellRichText(TextBlock(InlineFont(), text))
    return text


def _read_cell_value(value: Any, column: ReportColumn) -> Any:
    if value is None:
        return "" if column.data_type in {"string", "date", "datetime"} else None
    if isinstance(value, str) and value.startswith("'") and value[1:].startswith(_DANGEROUS_CELL_PREFIXES):
        value = value[1:]
    if column.data_type == "integer":
        return int(value)
    if column.data_type == "number":
        return float(value)
    if column.data_type == "boolean":
        return bool(value)
    return str(value)


def _file_sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def write_report_snapshot_xlsx(
    snapshot_value: ReportSnapshot | Mapping[str, Any],
    output_path: str | Path,
) -> dict[str, Any]:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    snapshot = (
        snapshot_value
        if isinstance(snapshot_value, ReportSnapshot)
        else ReportSnapshot.from_dict(snapshot_value)
    )
    path = Path(output_path).expanduser()
    if path.suffix.casefold() != ".xlsx":
        raise ValueError("Verified report output must use the .xlsx extension.")
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary_path = path.with_name(
        f".{path.stem}.{uuid4().hex}.tmp{path.suffix}"
    )

    workbook = Workbook()
    workbook.remove(workbook.active)
    header_fill = PatternFill("solid", fgColor="183153")
    header_font = Font(color="FFFFFF", bold=True)
    alternate_fill = PatternFill("solid", fgColor="F5F8FC")
    thin_gray = Side(style="thin", color="D7DEE8")
    border = Border(bottom=thin_gray, right=thin_gray)

    try:
        for sheet_snapshot in snapshot.sheets:
            worksheet = workbook.create_sheet(sheet_snapshot.name)
            worksheet.freeze_panes = "A2"
            worksheet.sheet_view.showGridLines = False
            worksheet.append([column.label for column in sheet_snapshot.columns])
            for cell in worksheet[1]:
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(vertical="center")
            worksheet.row_dimensions[1].height = 24

            for row in sheet_snapshot.rows:
                worksheet.append(
                    [
                        _xlsx_cell_value(value, column)
                        for value, column in zip(row, sheet_snapshot.columns)
                    ]
                )
            for column_index, column in enumerate(sheet_snapshot.columns, start=1):
                letter = get_column_letter(column_index)
                worksheet.column_dimensions[letter].width = column.width
                for row_index in range(2, worksheet.max_row + 1):
                    cell = worksheet.cell(row=row_index, column=column_index)
                    # Excel must never reinterpret identifiers such as GTIN/barcodes,
                    # order numbers or tracking codes as numbers.  Snapshot string,
                    # date and datetime values are intentionally serialized text, so
                    # keep the cell format textual unless the report explicitly asks
                    # for a different presentation format.
                    cell.number_format = (
                        "@"
                        if column.data_type in {"string", "date", "datetime"}
                        and column.number_format == "General"
                        else column.number_format
                    )
                    cell.border = border
                    cell.alignment = Alignment(vertical="top", wrap_text=column.width >= 28)
                    if row_index % 2 == 1:
                        cell.fill = alternate_fill
            if worksheet.max_row >= 1:
                worksheet.auto_filter.ref = worksheet.dimensions

        manifest = workbook.create_sheet(_MANIFEST_SHEET)
        manifest.sheet_state = "hidden"
        manifest_rows = [
            ("schema_version", str(snapshot.schema_version)),
            ("snapshot_id", snapshot.snapshot_id),
            ("content_hash", snapshot.content_hash),
            ("source_domain", snapshot.source_domain),
            ("report_type", snapshot.report_type),
            ("created_at", snapshot.created_at),
        ]
        for key, value in manifest_rows:
            manifest.append([key, value])

        workbook.save(temporary_path)
        precommit_verification = verify_report_snapshot_xlsx(snapshot, temporary_path)
        if not precommit_verification.get("verified"):
            raise ValueError(
                "Report workbook staging verification failed; existing output was preserved: "
                + ", ".join(precommit_verification.get("problems") or [])
            )
        temporary_path.replace(path)
    finally:
        if temporary_path.exists():
            temporary_path.unlink()

    verification = verify_report_snapshot_xlsx(snapshot, path)
    if not verification.get("verified"):
        raise ValueError(
            "Report workbook was written but exact semantic verification failed: "
            + ", ".join(verification.get("problems") or [])
        )
    return verification


def export_report_snapshot_xlsx(
    *,
    report_snapshot: Mapping[str, Any],
    output_path: str,
):
    """Canonical local-effect action shared by every report-producing domain."""

    from Feedback.processors.pipeline import Result, map_error_to_message

    try:
        verification = write_report_snapshot_xlsx(
            dict(report_snapshot or {}),
            output_path,
        )
        return Result.ok(
            "Dogrulanmis Excel raporu olusturuldu.",
            close_dialog=False,
            data=dict(verification),
        )
    except Exception as exc:
        return Result.fail(
            map_error_to_message(exc),
            error=exc,
            close_dialog=False,
        )


def verify_report_snapshot_xlsx(
    snapshot_value: ReportSnapshot | Mapping[str, Any],
    output_path: str | Path,
) -> dict[str, Any]:
    from openpyxl import load_workbook

    snapshot = (
        snapshot_value
        if isinstance(snapshot_value, ReportSnapshot)
        else ReportSnapshot.from_dict(snapshot_value)
    )
    path = Path(output_path).expanduser()
    problems: list[str] = []
    if not path.exists() or not path.is_file():
        return {
            "verified": False,
            "snapshot_id": snapshot.snapshot_id,
            "output_path": str(path),
            "problems": ["output_file_missing"],
        }

    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        if _MANIFEST_SHEET not in workbook.sheetnames:
            problems.append("manifest_sheet_missing")
            manifest_values: dict[str, str] = {}
        else:
            manifest_sheet = workbook[_MANIFEST_SHEET]
            manifest_values = {
                _text(row[0]): _text(row[1])
                for row in manifest_sheet.iter_rows(values_only=True)
                if row and _text(row[0])
            }
        if manifest_values.get("snapshot_id") != snapshot.snapshot_id:
            problems.append("snapshot_id_mismatch")
        if manifest_values.get("content_hash") != snapshot.content_hash:
            problems.append("content_hash_mismatch")

        expected_names = [sheet.name for sheet in snapshot.sheets]
        actual_names = [name for name in workbook.sheetnames if name != _MANIFEST_SHEET]
        if actual_names != expected_names:
            problems.append("sheet_set_or_order_mismatch")

        readback_sheets: list[dict[str, Any]] = []
        for sheet_snapshot in snapshot.sheets:
            if sheet_snapshot.name not in workbook.sheetnames:
                problems.append(f"sheet_missing:{sheet_snapshot.name}")
                continue
            worksheet = workbook[sheet_snapshot.name]
            raw_rows = list(worksheet.iter_rows(values_only=True))
            headers = [str(value or "") for value in (raw_rows[0] if raw_rows else ())]
            expected_headers = [column.label for column in sheet_snapshot.columns]
            if headers != expected_headers:
                problems.append(f"headers_mismatch:{sheet_snapshot.name}")

            actual_rows: list[list[Any]] = []
            formula_found = False
            for raw_row in raw_rows[1:]:
                values: list[Any] = []
                for value, column in zip(raw_row, sheet_snapshot.columns):
                    if isinstance(value, str) and value.startswith("="):
                        formula_found = True
                    values.append(_read_cell_value(value, column))
                actual_rows.append(values)
            # XLSX stores null and empty text as a blank cell. Keep their
            # distinction in the immutable source snapshot and content hash.
            expected_rows = [["" if value is None and column.data_type in {"string", "date", "datetime"}
                              else value for value, column in zip(row, sheet_snapshot.columns)]
                             for row in sheet_snapshot.rows]
            if actual_rows != expected_rows:
                problems.append(f"rows_mismatch:{sheet_snapshot.name}")
            if formula_found:
                problems.append(f"unexpected_formula:{sheet_snapshot.name}")
            readback_sheets.append(
                {
                    "name": sheet_snapshot.name,
                    "row_count": len(actual_rows),
                    "headers": headers,
                }
            )
    finally:
        workbook.close()

    return {
        "verified": not problems,
        "verified_by": "xlsx_exact_snapshot_readback",
        "snapshot_id": snapshot.snapshot_id,
        "content_hash": snapshot.content_hash,
        "output_path": str(path),
        "size_bytes": path.stat().st_size,
        "file_sha256": _file_sha256(path),
        "sheet_count": len(snapshot.sheets),
        "sheets": readback_sheets,
        "problems": problems,
    }
