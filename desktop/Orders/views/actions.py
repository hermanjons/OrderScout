# Orders/views/actions.py
from __future__ import annotations

import asyncio
import json
import re
import unicodedata
from datetime import datetime, date, timedelta
from functools import lru_cache
from pathlib import Path
from typing import Optional, Callable, Any

from PyQt6.QtWidgets import QListWidgetItem
from PyQt6.QtCore import Qt, QSignalBlocker

from Core.views import SwitchButton, ListSmartItemWidget
from Core.threads.async_worker import AsyncWorker
from Core.threads.sync_worker import SyncWorker
from Core.processors.report_artifacts import export_report_snapshot_xlsx
from Core.utils.time_utils import coerce_to_date

from Core.process.process_runner import GenericProcessRunner, ProcessTarget

from Feedback.processors.pipeline import MessageHandler, Result, map_error_to_message
from Logistics.constants.catalog_constants import (
    cargo_provider_name_for_code as shared_cargo_provider_name_for_code,
    changeable_cargo_provider_records,
)
from Marketplaces.constants.catalog_constants import canonical_platform_code, platform_select_options
from Marketplaces.api.trendyol.base_api import (
    trendyol_account_client_kwargs,
    trendyol_account_payload_client_kwargs,
)
from settings import DB_NAME

from Orders.constants import ui_constants as orders_ui
from Orders.constants import (
    INITIAL_FETCH_LOOKBACK_HOURS,
    INCREMENTAL_FETCH_OVERLAP_MINUTES,
    MAX_FETCH_WINDOW_DAYS,
    MAX_FETCH_WINDOW_SAFETY_MINUTES,
    NONFINAL_REFRESH_INTERVAL_HOURS,
    TZ_GMT3,
)
from Orders.constants.paths import ORDER_FALLBACK_LOGO_PATH, ORDERS_RUNTIME_STATE_FILE
from Orders.constants.trendyol.constants import (
    TRENDYOL_API_STATUS_LIST,
    TRENDYOL_ORDER_STATUS_ALIASES,
    TRENDYOL_RETURN_STATUS_LABELS,
    TRENDYOL_RETURN_STATUS_LIST,
    TRENDYOL_STATUS_LIST,
    TRENDYOL_STATUS_LABELS,
)
from Account.views.actions import collect_selected_companies
from Account.processors.pipeline import get_company_by_id, get_active_companies
from Core.utils.model_utils import update_records
from Orders.signals.signals import order_signals
from Orders.api.trendyol.orders_api import TrendyolApi

_ORDER_PLATFORM_CODE_CACHE_ATTR = "_orders_platform_code_cache"
_ORDER_DATE_CACHE_ATTR = "_orders_date_cache"
ORDER_LEVEL_COST_KEYS = {"kargo", "platform-hizmet-bedeli"}
DEFAULT_SERVICE_VAT_RATE = 20.0


class _SignalBlockerFallback:
    def __init__(self, widget) -> None:
        self._widget = widget
        self._previous = False
        try:
            self._previous = bool(widget.blockSignals(True))
        except Exception:
            self._widget = None

    def release(self) -> None:
        widget = self._widget
        self._widget = None
        if widget is None:
            return
        try:
            widget.blockSignals(self._previous)
        except Exception:
            return

    def __del__(self) -> None:
        self.release()


def _signal_blocker(widget):
    try:
        return QSignalBlocker(widget)
    except TypeError:
        return _SignalBlockerFallback(widget)


def load_finance_cost_items(*args, **kwargs):
    from Finance.views.actions import load_finance_cost_items as _impl

    return _impl(*args, **kwargs)


def estimate_listing_profit(*args, **kwargs):
    from Listing.processors.pipeline import estimate_listing_profit as _impl

    return _impl(*args, **kwargs)


def get_listings_from_db(*args, **kwargs):
    from Listing.processors.pipeline import get_listings_from_db as _impl

    return _impl(*args, **kwargs)


def normalize_listing_cost_binding(*args, **kwargs):
    from Listing.processors.pipeline import normalize_listing_cost_binding as _impl

    return _impl(*args, **kwargs)


def resolve_listing_cost_source(*args, **kwargs):
    from Listing.processors.pipeline import resolve_listing_cost_source as _impl

    return _impl(*args, **kwargs)


def load_logistics_cargo_tariffs(*args, **kwargs):
    from Logistics.views.actions import load_logistics_cargo_tariffs as _impl

    return _impl(*args, **kwargs)


def action_list_item_balances(*args, **kwargs):
    from Stock.views.actions import action_list_item_balances as _impl

    return _impl(*args, **kwargs)


def action_list_items(*args, **kwargs):
    from Stock.views.actions import action_list_items as _impl

    return _impl(*args, **kwargs)


def _orders_pipeline(name: str):
    from Orders.processors.trendyol import pipeline as orders_pipeline

    return getattr(orders_pipeline, name)


def _classify_order_shipping_deadline(*args, **kwargs):
    return _orders_pipeline("classify_order_shipping_deadline")(*args, **kwargs)


def fetch_orders_all(*args, **kwargs):
    return _orders_pipeline("fetch_orders_all")(*args, **kwargs)


def filter_order_records_in_memory(*args, **kwargs):
    return _orders_pipeline("filter_order_records_in_memory")(*args, **kwargs)


def get_latest_orders(*args, **kwargs):
    return _orders_pipeline("get_latest_orders")(*args, **kwargs)


def get_latest_orders_filtered(*args, **kwargs):
    return _orders_pipeline("get_latest_orders_filtered")(*args, **kwargs)


def get_latest_order_items_by_numbers(*args, **kwargs):
    return _orders_pipeline("get_latest_order_items_by_numbers")(*args, **kwargs)


def get_latest_ready_to_ship_orders(*args, **kwargs):
    return _orders_pipeline("get_latest_ready_to_ship_orders")(*args, **kwargs)


def get_latest_return_claims(*args, **kwargs):
    return _orders_pipeline("get_latest_return_claims")(*args, **kwargs)


def get_latest_return_claims_filtered(*args, **kwargs):
    return _orders_pipeline("get_latest_return_claims_filtered")(*args, **kwargs)


def get_stale_return_claims(*args, **kwargs):
    return _orders_pipeline("get_stale_return_claims")(*args, **kwargs)


def get_nonfinal_order_numbers(*args, **kwargs):
    return _orders_pipeline("get_nonfinal_order_numbers")(*args, **kwargs)


def _get_order_record_date(*args, **kwargs):
    return _orders_pipeline("get_order_record_date")(*args, **kwargs)


def _resolve_order_record_platform_code(*args, **kwargs):
    return _orders_pipeline("resolve_order_record_platform_code")(*args, **kwargs)


def normalize_order_data(*args, **kwargs):
    return _orders_pipeline("normalize_order_data")(*args, **kwargs)


def normalize_return_claim(*args, **kwargs):
    return _orders_pipeline("normalize_return_claim")(*args, **kwargs)


def save_return_claims_to_db(*args, **kwargs):
    return _orders_pipeline("save_return_claims_to_db")(*args, **kwargs)


def mark_return_claims_stale(*args, **kwargs):
    return _orders_pipeline("mark_return_claims_stale")(*args, **kwargs)


def restore_return_claims_from_stale(*args, **kwargs):
    return _orders_pipeline("restore_return_claims_from_stale")(*args, **kwargs)


def _summarize_order_shipping_deadline(*args, **kwargs):
    return _orders_pipeline("summarize_order_shipping_deadline")(*args, **kwargs)


def update_last_used_at_for_accounts(*args, **kwargs):
    return _orders_pipeline("update_last_used_at_for_accounts")(*args, **kwargs)


def _trendyol_api(*args, **kwargs):
    return TrendyolApi(*args, **kwargs)


def _trendyol_api_for_account(account: Any):
    return _trendyol_api(**trendyol_account_client_kwargs(account))


def _trendyol_api_for_account_payload(payload: Any):
    return _trendyol_api(**trendyol_account_payload_client_kwargs(payload))


def _order_data_model():
    from Orders.models.trendyol.models import OrderData

    return OrderData


def _safe_str(value: Any) -> str:
    return str(value or "").strip()


def _optional_float(value: Any) -> float | None:
    try:
        text = str(value).strip().replace(",", ".")
        if not text:
            return None
        return float(text)
    except Exception:
        return None


def _record_pk(row: Any) -> int | None:
    if isinstance(row, dict):
        candidates = (row.get("pk"), row.get("id"))
    else:
        candidates = (getattr(row, "pk", None), getattr(row, "id", None))
    for value in candidates:
        try:
            key = int(value)
        except Exception:
            continue
        if key > 0:
            return key
    return None


def _amount_text(value: float | None) -> str:
    if value is None:
        return "-"
    return f"{value:.2f} TL"


def _delta_text(value: float | None) -> str:
    if value is None:
        return "-"
    if value > 0:
        return f"-{value:.2f} TL"
    if value < 0:
        return f"+{abs(value):.2f} TL"
    return "0.00 TL"


def _gross_amount(
    amount: Any,
    *,
    is_vat_included: bool = True,
    vat_rate: Any = None,
    default_vat_rate: Any = None,
) -> float | None:
    amount_val = _optional_float(amount)
    if amount_val is None:
        return None
    if is_vat_included:
        return amount_val
    rate = _optional_float(vat_rate)
    if rate is None:
        rate = _optional_float(default_vat_rate)
    if rate is None or rate < 0:
        return amount_val
    return amount_val * (1.0 + (rate / 100.0))


def _vat_amount_from_gross(amount: Any, *, vat_rate: Any = None, default_vat_rate: Any = None) -> float | None:
    amount_val = _optional_float(amount)
    if amount_val is None:
        return None
    rate = _optional_float(vat_rate)
    if rate is None:
        rate = _optional_float(default_vat_rate)
    if rate is None or rate < 0:
        return None
    if rate == 0:
        return 0.0
    return amount_val * (rate / (100.0 + rate))


def _dedupe_preserve_order(items: list[str]) -> list[str]:
    out: list[str] = []
    seen: set[str] = set()
    for item in items:
        text = _safe_str(item)
        if not text or text in seen:
            continue
        seen.add(text)
        out.append(text)
    return out


def _compose_linked_stock_label(
    stock_items: Any,
    *,
    stock_item_label_map: dict[int, str] | None,
    sale_quantity: float = 1.0,
    fallback: str = "",
) -> str:
    label_map = dict(stock_item_label_map or {})
    labels: list[str] = []
    for stock_item in list(stock_items or []):
        if not isinstance(stock_item, dict):
            continue
        try:
            stock_item_id = int(stock_item.get("stock_item_id"))
        except Exception:
            continue
        label = _safe_str(label_map.get(stock_item_id))
        if not label:
            continue
        try:
            qty_value = float(stock_item.get("qty") or 0) * float(sale_quantity)
        except Exception:
            qty_value = 0.0
        if qty_value > 0:
            labels.append(f"{label} x{qty_value:.1f}")
        else:
            labels.append(label)
    return ", ".join(labels) or _safe_str(fallback)


def _cost_items_for_platform(
    cost_items_cache: list[dict[str, Any]] | None,
    platform_code: Any,
) -> list[dict[str, Any]]:
    normalized_platform = canonical_platform_code(platform_code)
    if not normalized_platform:
        return [dict(row) for row in (cost_items_cache or []) if isinstance(row, dict)]
    out: list[dict[str, Any]] = []
    for row in (cost_items_cache or []):
        if not isinstance(row, dict):
            continue
        supported = [
            canonical_platform_code(value)
            for value in list(row.get("supportedPlatforms") or [])
            if canonical_platform_code(value)
        ]
        if supported and normalized_platform not in supported:
            continue
        out.append(dict(row))
    return out


def _build_cost_card_map(
    cost_items_cache: list[dict[str, Any]] | None,
    *,
    platform_code: Any = None,
) -> dict[str, dict]:
    return {
        _safe_str(row.get("costKey")).casefold(): dict(row)
        for row in _cost_items_for_platform(cost_items_cache, platform_code)
        if _safe_str(row.get("costKey"))
    }


def _cost_key(value: Any) -> str:
    return _safe_str(value).casefold()


def _is_order_level_cost(row: dict[str, Any]) -> bool:
    return _cost_key(row.get("cost_key") or row.get("costKey")) in ORDER_LEVEL_COST_KEYS


def _lookup_value_variants(value: Any) -> list[str]:
    text = _safe_str(value)
    if not text or text == "-":
        return []
    variants = [text.casefold()]
    compact = text.replace(" ", "").casefold()
    if compact and compact not in variants:
        variants.append(compact)
    return variants


def _append_listing_lookup(index_map: dict[tuple[int, str], dict], api_account_id: int, value: Any, listing: dict) -> None:
    for variant in _lookup_value_variants(value):
        index_map.setdefault((int(api_account_id), variant), listing)


def _build_listing_lookup_indexes(listings: list[dict[str, Any]]) -> dict[str, dict[tuple[int, str], dict]]:
    by_product_code: dict[tuple[int, str], dict] = {}
    by_barcode: dict[tuple[int, str], dict] = {}
    by_stock_code: dict[tuple[int, str], dict] = {}
    for listing in listings or []:
        if not isinstance(listing, dict):
            continue
        try:
            api_account_id = int(listing.get("api_account_id") or 0)
        except Exception:
            api_account_id = 0
        if api_account_id <= 0:
            continue
        _append_listing_lookup(
            by_product_code,
            api_account_id,
            listing.get("productCode") or listing.get("product_code") or listing.get("productId") or listing.get("product_id"),
            listing,
        )
        _append_listing_lookup(by_barcode, api_account_id, listing.get("barcode"), listing)
        _append_listing_lookup(
            by_stock_code,
            api_account_id,
            listing.get("stockCode") or listing.get("stock_code"),
            listing,
        )
    return {
        "by_product_code": by_product_code,
        "by_barcode": by_barcode,
        "by_stock_code": by_stock_code,
    }


def _match_listing_for_order_item(
    *,
    order_item: dict[str, Any],
    api_account_id: int,
    lookup: dict[str, dict[tuple[int, str], dict]],
) -> dict[str, Any] | None:
    if api_account_id <= 0:
        return None

    def _find(index_name: str, value: Any) -> dict[str, Any] | None:
        bucket = lookup.get(index_name) or {}
        for variant in _lookup_value_variants(value):
            matched = bucket.get((int(api_account_id), variant))
            if matched is not None:
                return matched
        return None

    matched = _find("by_product_code", order_item.get("productCode") or order_item.get("product_code"))
    if matched is not None:
        return matched
    matched = _find("by_barcode", order_item.get("barcode"))
    if matched is not None:
        return matched
    return _find(
        "by_stock_code",
        order_item.get("stockCode") or order_item.get("stock_code") or order_item.get("merchantSku") or order_item.get("sku"),
    )


def _component_payload(
    *,
    key: str,
    label: str,
    total: float,
    vat: float,
    kind: str,
    note: str = "",
    source_type: str = "",
    source_trace: dict[str, Any] | None = None,
) -> dict[str, Any]:
    return {
        "key": key,
        "label": label,
        "kind": kind,
        "gross": total,
        "vat": vat,
        "net": total - vat,
        "gross_text": _amount_text(total),
        "vat_text": _amount_text(vat),
        "net_text": _amount_text(total - vat),
        "note": note,
        "source_type": source_type or None,
        "source_trace": dict(source_trace or {}) or None,
    }


def _resolve_order_level_costs(
    *,
    bindings_by_key: dict[str, dict[str, Any]],
    cost_card_map: dict[str, dict[str, Any]],
    sale_price: float,
    order_desi: float | None,
    platform_code: str,
    cargo_tariffs: list[dict[str, Any]],
) -> tuple[list[dict[str, Any]], list[str], dict[str, float], bool]:
    applied: list[dict[str, Any]] = []
    missing: list[str] = []
    totals = {"cargo_total": 0.0, "other_total": 0.0, "cargo_vat": 0.0, "other_vat": 0.0}
    complete = True

    for key in sorted(bindings_by_key.keys(), key=lambda value: 0 if value == "kargo" else 1):
        binding = dict(bindings_by_key.get(key) or {})
        card = dict(cost_card_map.get(key) or {})
        if not card:
            complete = False
            missing.append(f"Maliyet karti bulunamadi: {key}")
            continue
        label = _safe_str(card.get("name") or card.get("label") or key) or key
        normalized = normalize_listing_cost_binding(binding, card=card)
        card_currency = _safe_str(card.get("currency")).upper()
        value_type = _safe_str(card.get("valueType") or card.get("calculationType")).casefold()
        tariff_binding = key == "kargo" and normalized.get("source_type") == "builtin_resolver"
        if value_type != "percentage" and not tariff_binding and card_currency not in {"", "TRY"}:
            complete = False
            missing.append(
                f"Maliyet kartı {label} {card_currency}; doğrulanmış kur dönüşümü "
                "olmadan TRY kârı hesaplanamaz."
            )
            continue
        resolution = resolve_listing_cost_source(
            binding=normalized,
            card=card,
            context={
                "item": {"salePrice": sale_price, "dimensionalWeight": order_desi},
                "sale_price": sale_price,
                "effective_dimensional_weight": order_desi,
                "platform_code": platform_code,
                "cargo_tariffs": cargo_tariffs,
            },
        )
        if not bool(resolution.get("success")):
            complete = False
            missing.append(_safe_str(resolution.get("missing_reason")) or f"Maliyet hesabi basarisiz: {label}")
            continue
        resolved_amount = _optional_float(resolution.get("amount"))
        if resolved_amount is None:
            complete = False
            missing.append(f"Maliyet hesabi basarisiz: {label}")
            continue
        meta = dict(resolution.get("meta") or {})
        matched_tariff = dict(meta.get("matchedTariff") or {})
        tariff_currency = _safe_str(matched_tariff.get("currency")).upper()
        if tariff_currency not in {"", "TRY"}:
            complete = False
            missing.append(
                f"Kargo tarifesi {tariff_currency}; doğrulanmış kur dönüşümü "
                "olmadan TRY kârı hesaplanamaz."
            )
            continue
        gross_amount = (
            resolved_amount
            if bool(resolution.get("amount_is_gross", False))
            else _gross_amount(
                resolved_amount,
                is_vat_included=bool(card.get("isVatIncluded", True)),
                vat_rate=card.get("vatRate"),
                default_vat_rate=DEFAULT_SERVICE_VAT_RATE,
            )
        )
        if gross_amount is None:
            complete = False
            missing.append(f"Maliyet hesabi basarisiz: {label}")
            continue

        vat_rate = matched_tariff.get("vatRate") if key == "kargo" and matched_tariff else card.get("vatRate")
        vat_amount = _vat_amount_from_gross(
            gross_amount,
            vat_rate=vat_rate,
            default_vat_rate=DEFAULT_SERVICE_VAT_RATE,
        ) or 0.0
        note_parts: list[str] = []
        tariff_name = _safe_str(matched_tariff.get("name") or matched_tariff.get("label") or meta.get("matchedTariffName"))
        if tariff_name:
            note_parts.append(tariff_name)
        note = _safe_str(resolution.get("note"))
        if note and note not in note_parts:
            note_parts.append(note)

        if key == "kargo":
            provider = _safe_str(meta.get("selectedCargoProviderName") or meta.get("cargoProviderName"))
            if provider and provider not in note_parts:
                note_parts.append(provider)
            totals["cargo_total"] += gross_amount
            totals["cargo_vat"] += vat_amount
            kind = "cargo"
        else:
            totals["other_total"] += gross_amount
            totals["other_vat"] += vat_amount
            kind = "other"

        applied.append(
            _component_payload(
                key=key,
                label=label,
                total=gross_amount,
                vat=vat_amount,
                kind=kind,
                note=" | ".join(note_parts),
                source_type=_safe_str(resolution.get("source_type")),
                source_trace=dict(resolution.get("trace") or {}),
            )
        )

    return applied, missing, totals, complete


def _now_epoch_ms() -> int:
    return int(datetime.now(TZ_GMT3).timestamp() * 1000)


def _parse_last_used_to_ms(last_used) -> Optional[int]:
    if last_used in (None, ""):
        return None
    try:
        if isinstance(last_used, datetime):
            dt = last_used
            if dt.tzinfo is None:
                dt = dt.replace(tzinfo=TZ_GMT3)
            return int(dt.timestamp() * 1000)
        if isinstance(last_used, (int, float)):
            v = float(last_used)
            return int(v) if v > 10 ** 12 else int(v * 1000)
        if isinstance(last_used, str):
            s = last_used.strip()
            if not s:
                return None
            try:
                dt = datetime.fromisoformat(s)
                if dt.tzinfo is None:
                    dt = dt.replace(tzinfo=TZ_GMT3)
                return int(dt.timestamp() * 1000)
            except ValueError:
                v = float(int(s))
                return int(v) if v > 10 ** 12 else int(v * 1000)
    except Exception:
        return None
    return None


def _extract_last_used_ms_list(comp_api_account_list: list) -> list[int]:
    out: list[int] = []
    for acc in comp_api_account_list or []:
        last_used = acc[4] if isinstance(acc, (list, tuple)) and len(acc) >= 5 else None
        ms_val = _parse_last_used_to_ms(last_used)
        if ms_val is not None:
            out.append(ms_val)
    return out


def _compute_main_fetch_window(
    *,
    now_ms: int,
    last_used_ms_list: list[int],
    initial_lookback_hours: int = INITIAL_FETCH_LOOKBACK_HOURS,
) -> tuple[int, int]:
    start_ep_time = int(now_ms)
    max_window_ms = max(1, int(MAX_FETCH_WINDOW_DAYS)) * 24 * 60 * 60 * 1000
    safety_ms = max(0, int(MAX_FETCH_WINDOW_SAFETY_MINUTES)) * 60 * 1000
    max_window_ms = max(1, max_window_ms - safety_ms)
    min_allowed_final_ep_time = int(start_ep_time - max_window_ms)
    if last_used_ms_list:
        overlap_ms = max(1, int(INCREMENTAL_FETCH_OVERLAP_MINUTES)) * 60 * 1000
        last_scan_ms = int(min(last_used_ms_list))
        final_ep_time = int(last_scan_ms - overlap_ms)
    else:
        hours = max(1, int(initial_lookback_hours))
        final_ep_time = int(now_ms - hours * 60 * 60 * 1000)
    if final_ep_time < min_allowed_final_ep_time:
        final_ep_time = min_allowed_final_ep_time
    if final_ep_time > start_ep_time:
        final_ep_time = start_ep_time
    return final_ep_time, start_ep_time


def _load_orders_runtime_state(*, state_path: Path = ORDERS_RUNTIME_STATE_FILE) -> dict:
    try:
        if not state_path.exists():
            return {}
        raw = state_path.read_text(encoding="utf-8")
        parsed = json.loads(raw)
        return parsed if isinstance(parsed, dict) else {}
    except Exception:
        return {}


def _save_orders_runtime_state(data: dict, *, state_path: Path = ORDERS_RUNTIME_STATE_FILE) -> None:
    payload = dict(data or {})
    state_path.parent.mkdir(parents=True, exist_ok=True)
    state_path.write_text(
        json.dumps(payload, ensure_ascii=False, indent=2, sort_keys=True),
        encoding="utf-8",
    )


def _should_run_nonfinal_refresh(
    *,
    now_ms: Optional[int] = None,
    interval_hours: int = NONFINAL_REFRESH_INTERVAL_HOURS,
    state_path: Path = ORDERS_RUNTIME_STATE_FILE,
) -> bool:
    now = int(_now_epoch_ms() if now_ms is None else now_ms)
    interval_ms = max(1, int(interval_hours)) * 60 * 60 * 1000
    state = _load_orders_runtime_state(state_path=state_path)
    last_ms = _parse_last_used_to_ms(state.get("last_nonfinal_refresh_ms"))
    if last_ms is None:
        return True
    if now < last_ms:
        return True
    return (now - last_ms) >= interval_ms


def _mark_nonfinal_refresh_completed(
    *,
    now_ms: Optional[int] = None,
    state_path: Path = ORDERS_RUNTIME_STATE_FILE,
) -> None:
    now = int(_now_epoch_ms() if now_ms is None else now_ms)
    state = _load_orders_runtime_state(state_path=state_path)
    state["last_nonfinal_refresh_ms"] = now
    state["last_nonfinal_refresh_iso"] = datetime.fromtimestamp(now / 1000, tz=TZ_GMT3).isoformat()
    _save_orders_runtime_state(state, state_path=state_path)


def _set_progress_lock(progress_target, locked: bool) -> None:
    try:
        setattr(progress_target, "_orders_progress_locked", bool(locked))
    except Exception:
        pass


def _is_progress_locked(progress_target) -> bool:
    return bool(getattr(progress_target, "_orders_progress_locked", False))


# ============================================================
# 1) OrdersListWidget helpers
# ============================================================

def resolve_order_logo_path(order) -> str:
    logo = getattr(getattr(order, "api_account", None), "logo_path", None)
    return logo or ORDER_FALLBACK_LOGO_PATH


def _coerce_epoch_ms(value: Any) -> int | None:
    try:
        if value in (None, "", 0):
            return None
        if isinstance(value, datetime):
            dt = value
            if dt.tzinfo is None:
                dt = dt.replace(tzinfo=TZ_GMT3)
            return int(dt.timestamp() * 1000)
        if isinstance(value, date):
            dt = datetime(value.year, value.month, value.day, tzinfo=TZ_GMT3)
            return int(dt.timestamp() * 1000)
        ms = int(value)
        if ms <= 0:
            return None
        if ms < 10 ** 12:
            ms *= 1000
        return ms
    except Exception:
        return None


def _format_shipping_delta_parts(delta_ms: int) -> str:
    total_seconds = max(0, int(abs(delta_ms) // 1000))
    days, rem = divmod(total_seconds, 24 * 60 * 60)
    hours, rem = divmod(rem, 60 * 60)
    minutes, _seconds = divmod(rem, 60)

    if days > 0:
        return f"{days} gün {hours} saat"
    if hours > 0:
        return f"{hours} saat" if minutes <= 0 else f"{hours} saat {minutes} dk"
    return f"{minutes} dk"


def _shipping_deadline_status_code(order) -> str:
    raw_status = getattr(order, "shipmentPackageStatus", None) or getattr(order, "status", None)
    code = str(raw_status or "").strip()
    if not code:
        return ""
    key = code.casefold()
    if key == "readytoship":
        return "Created"
    if key == "unsupplied":
        return "Cancelled"
    return code


def summarize_order_shipping_deadline(order, *, now_ms: int | None = None) -> dict:
    resolved_now_ms = _now_epoch_ms() if now_ms is None else now_ms
    return _summarize_order_shipping_deadline(order, now_ms=resolved_now_ms)


def classify_order_shipping_deadline(order, *, now_ms: int | None = None) -> str:
    resolved_now_ms = _now_epoch_ms() if now_ms is None else now_ms
    return _classify_order_shipping_deadline(order, now_ms=resolved_now_ms)


def format_order_summary(order) -> dict:
    total = getattr(order, "totalPrice", 0)
    try:
        total_fmt = f"{float(total):,.2f}".replace(",", ".")
    except Exception:
        total_fmt = str(total)

    date_part = (
        getattr(order, "claimDate", None)
        or getattr(order, "orderDate", None)
        or getattr(order, "createdDate", None)
    )
    date_str = date_part.strftime("%d.%m.%Y") if isinstance(date_part, datetime) else str(date_part or "-")
    is_return = bool(getattr(order, "_is_return", False))

    if is_return:
        claim_number = str(getattr(order, "claimNumber", "") or getattr(order, "claimId", "") or "-")
        order_number = str(getattr(order, "orderNumber", "") or "-")
        status_text = format_order_status(
            str(getattr(order, "claimItemStatus", "") or getattr(order, "status", "") or "")
        )
        return {
            "title": f"İade: {claim_number}",
            "subtitle": (
                f"Müşteri: {getattr(order, 'customerFirstName', '-')} "
                f"{getattr(order, 'customerLastName', '')} | Sipariş: {order_number}"
            ),
            "extra": (
                f"Durum: {status_text} | "
                f"Tarih: {date_str} | Tutar: {total_fmt} ₺"
            ),
            "identifier": claim_number,
            "logo_path": resolve_order_logo_path(order),
        }

    return {
        "title": f"Sipariş: {getattr(order, 'orderNumber', '-')}",
        "subtitle": f"Müşteri: {getattr(order, 'customerFirstName', '-')} {getattr(order, 'customerLastName', '')}",
        "extra": (
            f"Kargo: {getattr(order, 'cargoProviderName', '-')} | "
            f"Tarih: {date_str} | Tutar: {total_fmt} ₺"
        ),
        "identifier": getattr(order, "orderNumber", "-"),
        "logo_path": resolve_order_logo_path(order),
    }


def build_order_list(list_widget, orders: list, interaction_cb=None, selection_cb=None) -> Result:
    try:
        list_widget.setUpdatesEnabled(False)
        list_widget.clear()

        if not orders:
            info_item = QListWidgetItem("Gösterilecek sipariş bulunamadı.")
            info_item.setFlags(info_item.flags() & ~Qt.ItemFlag.ItemIsSelectable)
            list_widget.addItem(info_item)
            return Result.ok("Liste boş, sipariş bulunamadı.", data={"count": 0})

        added = 0
        for order in orders:
            display = format_order_summary(order)
            switch = SwitchButton()

            item_widget = ListSmartItemWidget(
                title=display["title"],
                subtitle=display["subtitle"],
                extra=display["extra"],
                identifier=display["identifier"],
                icon_path=display["logo_path"],
                optional_widget=switch,
            )

            if interaction_cb:
                item_widget.interaction.connect(interaction_cb)
            if selection_cb:
                item_widget.selectionRequested.connect(selection_cb)

            item = QListWidgetItem(list_widget)
            item.setSizeHint(item_widget.sizeHint())
            list_widget.setItemWidget(item, item_widget)
            added += 1

        return Result.ok(f"{added} sipariş başarıyla listelendi.", data={"count": added})

    except Exception as e:
        return Result.fail(map_error_to_message(e), error=e)
    finally:
        list_widget.setUpdatesEnabled(True)


def collect_selected_orders(list_widget) -> Result:
    """
    OrdersListWidget varsa: model tabanlı seçim (tüm sayfalar)
    yoksa: mevcut UI satırlarından checkbox seçimi
    """
    try:
        if hasattr(list_widget, "get_selected_orders"):
            selected_orders = list_widget.get_selected_orders() or []
            return Result.ok(
                "Seçili siparişler toplandı.",
                data={"selected_orders": selected_orders},
                close_dialog=False,
            )

        selected_orders = []
        for i in range(list_widget.count()):
            item = list_widget.item(i)
            widget = list_widget.itemWidget(item)
            if not widget:
                continue

            right = getattr(widget, "right_widget", None)
            order_obj = getattr(widget, "order_obj", None) or getattr(widget, "order", None)

            if right is not None and bool(right.isChecked()) and order_obj is not None:
                selected_orders.append(order_obj)

        return Result.ok(
            "Seçili siparişler toplandı.",
            data={"selected_orders": selected_orders},
            close_dialog=False,
        )

    except Exception as e:
        return Result.fail(map_error_to_message(e), error=e, close_dialog=False)


def extract_cargo_names(orders: list) -> list[str]:
    return sorted({
        str(getattr(o, "cargoProviderName", "") or "").strip()
        for o in orders
        if getattr(o, "cargoProviderName", None)
    })


def order_platform_filter_options(*, include_all: bool = False) -> list[tuple[str, str]]:
    return platform_select_options(include_all=include_all)


def _humanize_status_code(status_code: str) -> str:
    text = str(status_code or "").strip()
    if not text:
        return "-"
    text = text.replace("_", " ").replace("-", " ")
    text = re.sub(r"(?<!^)(?=[A-Z])", " ", text)
    text = re.sub(r"\s+", " ", text).strip()
    return text or "-"


def return_status_label_for_code(status_code: str) -> str | None:
    code = str(status_code or "").strip()
    if not code:
        return None
    for key, value in TRENDYOL_RETURN_STATUS_LABELS.items():
        if str(key).casefold() == code.casefold():
            return str(value)
    return None


def _return_status_label_for_code(status_code: str) -> str | None:
    return return_status_label_for_code(status_code)


def _normalize_filter_choice_key(value: object) -> str:
    text = str(value or "").strip()
    if not text:
        return ""
    text = text.replace("İ", "I").replace("ı", "i")
    normalized = unicodedata.normalize("NFKD", text.casefold())
    normalized = "".join(ch for ch in normalized if not unicodedata.combining(ch))
    cleaned = []
    for ch in normalized:
        if ch.isalnum():
            cleaned.append(ch)
        elif ch.isspace():
            cleaned.append(" ")
    return " ".join("".join(cleaned).split())


def _is_all_like_value(value: object, *, extra_labels: list[str] | tuple[str, ...] | None = None) -> bool:
    text_key = _normalize_filter_choice_key(value)
    if not text_key:
        return True
    compact_key = text_key.replace(" ", "")
    labels = {
        "",
        "all",
        "tumu",
        "tum",
        "tm",
        "seciniz",
        "seiniz",
        "select",
        "selection",
    }
    for row in extra_labels or ():
        label_key = _normalize_filter_choice_key(row)
        if label_key:
            labels.add(label_key)
            labels.add(label_key.replace(" ", ""))
    return text_key in labels or compact_key in labels


def _normalize_shipment_status_code(status_code: str) -> str:
    code = str(status_code or "").strip()
    if not code:
        return ""
    for alias_code, canonical_code in TRENDYOL_ORDER_STATUS_ALIASES.items():
        if str(alias_code).casefold() == code.casefold():
            return str(canonical_code)
    return code


def _status_code_from_label(value: object, label_map: dict[str, str]) -> str:
    key = _normalize_filter_choice_key(value)
    if not key:
        return ""
    for code, label in (label_map or {}).items():
        if _normalize_filter_choice_key(label) == key:
            return str(code or "").strip()
    return ""


def extract_order_statuses(orders: list) -> list[tuple[str, str]]:
    statuses: set[str] = set()
    for order in (orders or []):
        raw = (
            getattr(order, "shipmentPackageStatus", None)
            or getattr(order, "claimItemStatus", None)
            or getattr(order, "status", None)
        )
        text = _normalize_shipment_status_code(raw)
        if text:
            statuses.add(text)
    mapped: list[tuple[str, str]] = [
        (format_order_status(code), code)
        for code in statuses
    ]
    mapped.sort(key=lambda row: (str(row[0]).casefold(), str(row[1]).casefold()))
    return mapped


def order_status_label_for_code(status_code: str) -> str:
    code = _normalize_shipment_status_code(status_code)
    if not code:
        return "-"
    for key, value in TRENDYOL_STATUS_LABELS.items():
        if str(key).casefold() == code.casefold():
            return str(value)
    return _humanize_status_code(code)


def format_order_status(status_code: str) -> str:
    claim_label = _return_status_label_for_code(status_code)
    if claim_label:
        return claim_label
    return order_status_label_for_code(status_code)


def _extract_item_image(item) -> str:
    candidate_keys = (
        "imageUrl",
        "imageURL",
        "productImageUrl",
        "productImageURL",
        "image",
        "image_path",
        "thumbnail",
        "pictureUrl",
    )
    for key in candidate_keys:
        value = getattr(item, key, None)
        if isinstance(value, str) and value.strip():
            return value.strip()
        if isinstance(value, dict):
            nested = value.get("url") or value.get("imageUrl")
            if isinstance(nested, str) and nested.strip():
                return nested.strip()
    return ""


def _resolve_item_total_amount(item) -> Any:
    try:
        quantity = int(getattr(item, "quantity", 0) or 0)
    except Exception:
        quantity = 0

    def _coerce(raw: Any) -> float | None:
        try:
            if raw in (None, ""):
                return None
            return float(raw)
        except Exception:
            return None

    def _same(a: float | None, b: float | None) -> bool:
        return a is not None and b is not None and abs(a - b) < 0.0001

    unit_price = _coerce(getattr(item, "lineUnitPrice", None))
    if unit_price is None:
        unit_price = _coerce(getattr(item, "price", None))
    line_gross = _coerce(getattr(item, "lineGrossAmount", None))
    amount = _coerce(getattr(item, "amount", None))

    if unit_price is not None and quantity > 0:
        return unit_price * quantity

    if line_gross is not None:
        return line_gross * quantity if quantity > 1 else line_gross
    if amount is not None:
        return amount * quantity if quantity > 1 else amount
    return None


def _resolve_item_unit_price(item) -> float | None:
    def _coerce(raw: Any) -> float | None:
        try:
            if raw in (None, ""):
                return None
            return float(raw)
        except Exception:
            return None

    for key in ("lineUnitPrice", "price"):
        value = _coerce(getattr(item, key, None))
        if value is not None:
            return value

    total_amount = _coerce(_resolve_item_total_amount(item))
    try:
        quantity = int(getattr(item, "quantity", 0) or 0)
    except Exception:
        quantity = 0
    if total_amount is not None and quantity > 0:
        return total_amount / quantity
    return None


def _normalize_order_item_account_key(value: Any) -> int | None:
    try:
        if value in (None, ""):
            return None
        return int(value)
    except Exception:
        return None


def _order_item_group_key(order_number: Any, api_account_id: Any = None) -> tuple[int | None, str]:
    return (
        _normalize_order_item_account_key(api_account_id),
        str(order_number or "").strip(),
    )


def _order_item_to_view_row(item) -> dict:
    name = str(getattr(item, "productName", "") or "").strip() or "-"
    merchant_sku = str(getattr(item, "merchantSku", "") or "").strip()
    sku = merchant_sku or str(getattr(item, "sku", "") or "").strip() or "-"
    barcode = str(getattr(item, "barcode", "") or "").strip() or "-"
    qty = int(getattr(item, "quantity", 0) or 0)
    unit_price = _resolve_item_unit_price(item)
    amount = _resolve_item_total_amount(item)
    image_url = _extract_item_image(item)
    return {
        "orderNumber": str(getattr(item, "orderNumber", "") or "").strip(),
        "name": name,
        "sku": sku,
        "merchantSku": merchant_sku or sku,
        "stockCode": str(getattr(item, "stockCode", "") or "").strip() or merchant_sku or sku,
        "barcode": barcode,
        "qty": qty,
        "unit_price": unit_price,
        "amount": amount,
        "currencyCode": getattr(item, "currencyCode", None),
        "image": image_url,
        "productCode": getattr(item, "productCode", None),
        "vatRate": getattr(item, "vatRate", None),
        "api_account_id": getattr(item, "api_account_id", None),
        "lineId": getattr(item, "lineId", None),
    }


def load_order_items_for_order_numbers(
    order_numbers: list[str] | tuple[str, ...],
    *,
    api_account_ids: list[int] | tuple[int, ...] | None = None,
    db_name: str | None = None,
) -> Result:
    try:
        normalized_numbers = [
            str(number or "").strip()
            for number in (order_numbers or [])
            if str(number or "").strip()
        ]
        normalized_numbers = list(dict.fromkeys(normalized_numbers))
        if not normalized_numbers:
            return Result.ok(
                "Sipariş numarası boş.",
                close_dialog=False,
                data={"items": [], "items_by_order": {}},
            )

        normalized_account_ids = [
            int(value)
            for value in (api_account_ids or [])
            if _normalize_order_item_account_key(value) is not None
        ]
        normalized_account_ids = list(dict.fromkeys(normalized_account_ids))

        res = get_latest_order_items_by_numbers(
            normalized_numbers,
            api_account_ids=normalized_account_ids or None,
            db_name=db_name,
        )
        if not res.success:
            return res

        item_rows = (res.data or {}).get("items", []) or []
        grouped: dict[tuple[int | None, str], list[dict]] = {}
        flat_rows: list[dict] = []
        for item in item_rows:
            view_row = _order_item_to_view_row(item)
            flat_rows.append(view_row)
            key = _order_item_group_key(
                getattr(item, "orderNumber", None),
                getattr(item, "api_account_id", None),
            )
            grouped.setdefault(key, []).append(view_row)

        return Result.ok(
            f"Sipariş ürünleri yüklendi. ({len(flat_rows)})",
            close_dialog=False,
            data={"items": flat_rows, "items_by_order": grouped},
        )

    except Exception as e:
        return Result.fail(map_error_to_message(e), error=e, close_dialog=False)


def load_order_items_for_order_number(order_number: str, *, api_account_id: int | None = None) -> Result:
    try:
        order_no = str(order_number or "").strip()
        if not order_no:
            return Result.ok("Sipariş numarası boş.", close_dialog=False, data={"items": []})

        api_account_ids = [int(api_account_id)] if api_account_id is not None else None
        res = get_latest_order_items_by_numbers([order_no], api_account_ids=api_account_ids)
        if not res.success:
            return res

        item_rows = (res.data or {}).get("items", []) or []
        if not item_rows:
            return Result.ok("Sipariş detayı bulunamadı.", close_dialog=False, data={"items": []})


        rows: list[dict] = []
        for item in item_rows:
            name = str(getattr(item, "productName", "") or "").strip() or "-"
            merchant_sku = str(getattr(item, "merchantSku", "") or "").strip()
            sku = merchant_sku or str(getattr(item, "sku", "") or "").strip() or "-"
            barcode = str(getattr(item, "barcode", "") or "").strip() or "-"
            qty = int(getattr(item, "quantity", 0) or 0)
            unit_price = _resolve_item_unit_price(item)
            amount = _resolve_item_total_amount(item)
            image_url = _extract_item_image(item)
            rows.append(
                {
                    "orderNumber": str(getattr(item, "orderNumber", "") or "").strip(),
                    "name": name,
                    "sku": sku,
                    "merchantSku": merchant_sku or sku,
                    "stockCode": str(getattr(item, "stockCode", "") or "").strip() or merchant_sku or sku,
                    "barcode": barcode,
                    "qty": qty,
                    "unit_price": unit_price,
                    "amount": amount,
                    "currencyCode": getattr(item, "currencyCode", None),
                    "image": image_url,
                    "productCode": getattr(item, "productCode", None),
                    "vatRate": getattr(item, "vatRate", None),
                    "api_account_id": getattr(item, "api_account_id", None),
                    "lineId": getattr(item, "lineId", None),
                }
            )

        return Result.ok(
            f"Sipariş ürünleri yüklendi. ({len(rows)})",
            close_dialog=False,
            data={"items": rows},
        )

    except Exception as e:
        return Result.fail(map_error_to_message(e), error=e, close_dialog=False)


def load_order_profit_catalog(
    *,
    api_account_ids: list[int] | None = None,
    db_name: str = DB_NAME,
) -> Result:
    try:
        normalized_account_ids = sorted(
            {
                int(value)
                for value in (api_account_ids or [])
                if value is not None
            }
        )

        account_platform_map: dict[int, str] = {}
        accounts_res = get_active_companies(db_name=db_name)
        if accounts_res.success:
            for row in (accounts_res.data or {}).get("records", []) or []:
                try:
                    account_id = int(getattr(row, "pk", None))
                except Exception:
                    continue
                if normalized_account_ids and account_id not in normalized_account_ids:
                    continue
                account_platform_map[account_id] = canonical_platform_code(getattr(row, "platform", "") or "")

        listings_res = get_listings_from_db(
            api_account_ids=normalized_account_ids or None,
            db_name=db_name,
        )
        if not listings_res.success:
            return listings_res
        listings = [
            dict(row)
            for row in ((listings_res.data or {}).get("items", []) or [])
            if isinstance(row, dict)
        ]

        stock_item_label_map: dict[int, str] = {}
        stock_cost_map: dict[int, float] = {}
        stock_cost_meta_map: dict[int, dict[str, Any]] = {}
        stock_vat_rate_map: dict[int, float | None] = {}
        stock_dimension_map: dict[int, dict[str, float | None]] = {}

        stock_res = action_list_items(only_active=True, db_name=db_name)
        if stock_res.success:
            for row in (stock_res.data or {}).get("records", []) or []:
                item_id = _record_pk(row)
                if item_id is None:
                    continue
                stock_item_label_map[item_id] = f"{getattr(row, 'sku', '')} - {getattr(row, 'name', '')}"
                vat_rate = getattr(row, "vat_rate", None)
                try:
                    stock_vat_rate_map[item_id] = float(vat_rate) if vat_rate not in (None, "") else None
                except Exception:
                    stock_vat_rate_map[item_id] = None
                stock_dimension_map[item_id] = {
                    "weight_kg": getattr(row, "weight_kg", None),
                    "length_cm": getattr(row, "length_cm", None),
                    "width_cm": getattr(row, "width_cm", None),
                    "height_cm": getattr(row, "height_cm", None),
                }

        balances_res = action_list_item_balances(only_active=True, db_name=db_name)
        if balances_res.success:
            for row in (balances_res.data or {}).get("records", []) or []:
                try:
                    item = row.get("item")
                    item_id = _record_pk(item)
                    if item_id is None:
                        continue
                    avg_cost = float(row.get("avg_unit_cost"))
                except Exception:
                    continue
                stock_cost_map[item_id] = avg_cost
                stock_cost_meta_map[item_id] = {
                    "input_mode": row.get("avg_cost_input_mode"),
                    "currency": _safe_str(row.get("avg_cost_currency")).upper() or None,
                }

        cost_items_cache: list[dict[str, Any]] = []
        cost_res = load_finance_cost_items(active_state="Aktif", db_name=db_name)
        if cost_res.success:
            cost_items_cache = [
                {
                    "costKey": _safe_str(row.get("costKey")),
                    "name": _safe_str(row.get("name")),
                    "label": _safe_str(row.get("name")),
                    "valueType": _safe_str(row.get("valueType") or row.get("calculationType")),
                    "defaultMatchMode": _safe_str(row.get("defaultMatchMode")) or "manual",
                    "supportedSourceTypes": list(row.get("supportedSourceTypes") or []),
                    "defaultSourceType": _safe_str(row.get("defaultSourceType")),
                    "defaultSourceConfig": dict(row.get("defaultSourceConfig") or {}) or None,
                    "defaultAmount": row.get("defaultAmount"),
                    "currency": _safe_str(row.get("currency")) or "TRY",
                    "isVatIncluded": bool(row.get("isVatIncluded", True)),
                    "vatRate": row.get("vatRate"),
                    "supportedPlatforms": list(row.get("supportedPlatforms") or []),
                }
                for row in ((cost_res.data or {}).get("records", []) or [])
                if isinstance(row, dict)
                and _safe_str(row.get("costKey"))
                and (not bool(row.get("isBuiltin")) or bool(row.get("isBuiltinAccepted")))
            ]

        cargo_tariffs_cache: list[dict[str, Any]] = []
        tariff_res = load_logistics_cargo_tariffs(active_state="Aktif", db_name=db_name)
        if tariff_res.success:
            cargo_tariffs_cache = [
                {
                    "tariffKey": _safe_str(row.get("tariffKey")),
                    "name": _safe_str(row.get("name")),
                    "label": _safe_str(row.get("name") or row.get("cargoProviderName")),
                    "cargo_provider_id": int(row.get("cargo_provider_id") or 0) or None,
                    "platformCode": _safe_str(row.get("platformCode") or row.get("platform")),
                    "platform": _safe_str(row.get("platform")),
                    "cargoProviderCode": _safe_str(row.get("cargoProviderCode")),
                    "cargoProviderName": _safe_str(row.get("cargoProviderName")),
                    "minDesi": row.get("minDesi"),
                    "maxDesi": row.get("maxDesi"),
                    "minSalePrice": row.get("minSalePrice"),
                    "maxSalePrice": row.get("maxSalePrice"),
                    "amount": row.get("amount"),
                    "currency": _safe_str(row.get("currency")),
                    "isVatIncluded": bool(row.get("isVatIncluded", True)),
                    "vatRate": row.get("vatRate"),
                    "isActive": bool(row.get("isActive", True)),
                }
                for row in ((tariff_res.data or {}).get("records", []) or [])
                if isinstance(row, dict)
                and _safe_str(row.get("tariffKey"))
                and (not bool(row.get("isBuiltin")) or bool(row.get("isBuiltinAccepted")))
            ]

        return Result.ok(
            "Siparis kar katalogu hazirlandi.",
            close_dialog=False,
            data={
                "api_account_ids": normalized_account_ids,
                "account_platform_map": account_platform_map,
                "stock_item_label_map": stock_item_label_map,
                "stock_cost_map": stock_cost_map,
                "stock_cost_meta_map": stock_cost_meta_map,
                "stock_vat_rate_map": stock_vat_rate_map,
                "stock_dimension_map": stock_dimension_map,
                "cost_items_cache": cost_items_cache,
                "cargo_tariffs_cache": cargo_tariffs_cache,
                "listings": listings,
                "listing_lookup": _build_listing_lookup_indexes(listings),
            },
        )
    except Exception as e:
        return Result.fail(map_error_to_message(e), error=e, close_dialog=False)


def estimate_order_profit(
    *,
    order,
    order_items: list[dict[str, Any]] | None,
    catalog: dict[str, Any] | None,
) -> dict[str, Any]:
    order_items = [row for row in (order_items or []) if isinstance(row, dict)]
    catalog = catalog if isinstance(catalog, dict) else {}
    stock_item_label_map = catalog.get("stock_item_label_map") or {}
    stock_cost_map = catalog.get("stock_cost_map") or {}
    stock_cost_meta_map = catalog.get("stock_cost_meta_map") or {}
    stock_vat_rate_map = catalog.get("stock_vat_rate_map") or {}
    stock_dimension_map = catalog.get("stock_dimension_map") or {}
    cost_items_cache = catalog.get("cost_items_cache") or []
    cargo_tariffs_cache = catalog.get("cargo_tariffs_cache") or []
    account_platform_map = catalog.get("account_platform_map") or {}
    lookup = catalog.get("listing_lookup") or {}
    cost_card_map_cache = catalog.get("_cost_card_map_cache")
    if not isinstance(cost_card_map_cache, dict):
        cost_card_map_cache = {}
        catalog["_cost_card_map_cache"] = cost_card_map_cache

    def _cost_card_map_for(platform_code: Any) -> dict[str, dict]:
        cache_key = canonical_platform_code(platform_code)
        if cache_key not in cost_card_map_cache:
            cost_card_map_cache[cache_key] = _build_cost_card_map(
                cost_items_cache,
                platform_code=cache_key,
            )
        return cost_card_map_cache[cache_key]

    order_number = _safe_str(getattr(order, "orderNumber", ""))
    try:
        api_account_id = int(getattr(order, "api_account_id", None) or 0)
    except Exception:
        api_account_id = 0
    order_platform_code = canonical_platform_code(account_platform_map.get(api_account_id) or "")

    # The existing profit engine calculates TRY amounts and has no FX conversion
    # contract. Keep unlabelled stock costs and explicitly foreign monetary
    # inputs out of that calculation. Legacy sales/cost cards without a currency
    # retain their existing TRY behavior; stock balances now provide currency.
    currency_missing_reasons: list[str] = []
    order_currency = _safe_str(getattr(order, "currencyCode", None)).upper()
    if order_currency not in {"", "TRY"}:
        currency_missing_reasons.append(
            f"Sipariş satış para birimi {order_currency}; doğrulanmış kur dönüşümü "
            "olmadan TRY kârı hesaplanamaz."
        )
    for item in order_items:
        item_currency = _safe_str(item.get("currencyCode")).upper()
        if item_currency not in {"", "TRY"}:
            currency_missing_reasons.append(
                f"Sipariş kalemi satış para birimi {item_currency}; doğrulanmış kur "
                "dönüşümü olmadan TRY kârı hesaplanamaz."
            )
        listing = _match_listing_for_order_item(
            order_item=item, api_account_id=api_account_id, lookup=lookup,
        )
        if listing is None:
            continue
        if _optional_float(item.get("unit_price")) is None:
            listing_currency = _safe_str(listing.get("currencyCode") or listing.get("currency")).upper()
            if listing_currency not in {"", "TRY"}:
                currency_missing_reasons.append(
                    f"İlan satış para birimi {listing_currency}; doğrulanmış kur "
                    "dönüşümü olmadan TRY kârı hesaplanamaz."
                )
        listing_platform = canonical_platform_code(account_platform_map.get(listing.get("api_account_id")) or "")
        cards = _cost_card_map_for(listing_platform)
        for binding in listing.get("cost_items") or []:
            if not isinstance(binding, dict):
                continue
            key = _cost_key(binding.get("cost_key") or binding.get("costKey"))
            card = cards.get(key) or {}
            value_type = _safe_str(card.get("valueType") or card.get("calculationType")).casefold()
            normalized_binding = normalize_listing_cost_binding(binding, card=card)
            tariff_binding = key == "kargo" and normalized_binding.get("source_type") == "builtin_resolver"
            card_currency = _safe_str(card.get("currency")).upper()
            if value_type != "percentage" and not tariff_binding and card_currency not in {"", "TRY"}:
                label = _safe_str(card.get("name") or card.get("label") or key)
                currency_missing_reasons.append(
                    f"Maliyet kartı {label} {card_currency}; doğrulanmış kur dönüşümü "
                    "olmadan TRY kârı hesaplanamaz."
                )
        for stock_row in listing.get("stock_items") or []:
            if not isinstance(stock_row, dict):
                continue
            try:
                stock_item_id = int(stock_row.get("stock_item_id"))
            except (TypeError, ValueError):
                continue
            if (_optional_float(stock_row.get("qty")) or 0.0) <= 0:
                continue
            if _optional_float(stock_cost_map.get(stock_item_id)) is None:
                continue  # The existing estimator reports missing cost itself.
            cost_currency = _safe_str(
                (stock_cost_meta_map.get(stock_item_id) or {}).get("currency")
            ).upper()
            if not cost_currency:
                currency_missing_reasons.append(
                    f"Stok kartı maliyet para birimi eksik: #{stock_item_id}."
                )
            elif cost_currency != "TRY":
                currency_missing_reasons.append(
                    f"Stok kartı maliyeti {cost_currency}: #{stock_item_id}; "
                    "doğrulanmış kur dönüşümü olmadan TRY kârı hesaplanamaz."
                )
    if currency_missing_reasons:
        return {
            "estimated_profit": None,
            "estimated_profit_text": "-",
            "estimated_profit_ready": False,
            "estimated_profit_breakdown": [],
            "estimated_profit_detail": {
                "summary": {"order_number": order_number, "line_count": len(order_items),
                            "line_count_text": str(len(order_items))},
                "results": {},
                "lines": [],
            },
            "estimated_profit_missing_reasons": _dedupe_preserve_order(currency_missing_reasons),
        }

    if not order_items:
        return {
            "estimated_profit": None,
            "estimated_profit_text": "-",
            "estimated_profit_ready": False,
            "estimated_profit_breakdown": [],
            "estimated_profit_detail": {
                "summary": {"order_number": order_number, "line_count": 0, "line_count_text": "0"},
                "results": {},
                "lines": [],
            },
            "estimated_profit_missing_reasons": ["Siparis urunleri bulunamadi."],
        }

    line_entries: list[dict[str, Any]] = []
    combined_applied_costs: list[dict[str, Any]] = []
    order_level_bindings: dict[str, dict[str, Any]] = {}
    missing_reasons: list[str] = []
    complete = True
    matched_line_count = 0
    total_quantity = 0.0
    order_desi_total = 0.0
    has_order_desi = False
    totals = {
        "sale_price": 0.0,
        "payout_total": 0.0,
        "total_costs": 0.0,
        "input_vat_total": 0.0,
        "sale_vat": 0.0,
        "net_vat": 0.0,
        "estimated_profit": 0.0,
    }

    for item in order_items:
        qty_value = _optional_float(item.get("qty")) or 0.0
        if qty_value <= 0:
            qty_value = 1.0
        total_quantity += qty_value
        item_title = _safe_str(item.get("name")) or _safe_str(item.get("sku")) or _safe_str(item.get("barcode")) or "Urun"
        unit_sale_price = _optional_float(item.get("unit_price"))
        line_amount = (
            (unit_sale_price * qty_value)
            if unit_sale_price is not None
            else _optional_float(item.get("amount"))
        )
        matched_listing = _match_listing_for_order_item(order_item=item, api_account_id=api_account_id, lookup=lookup)
        if matched_listing is None:
            complete = False
            missing_reasons.append(f"Ilan eslesmesi bulunamadi: {item_title}")
            line_entries.append(
                {
                    "title": item_title,
                    "qty": qty_value,
                    "qty_text": f"{qty_value:.0f}" if float(qty_value).is_integer() else f"{qty_value:.2f}",
                    "unit_price": unit_sale_price,
                    "unit_price_text": _amount_text(unit_sale_price),
                    "amount": line_amount,
                    "amount_text": _amount_text(line_amount),
                    "listing_title": "",
                    "matched": False,
                    "estimated_profit": None,
                    "estimated_profit_text": "-",
                    "estimated_profit_ready": False,
                    "missing_reasons": ["Ilan eslesmesi bulunamadi."],
                }
            )
            continue

        matched_line_count += 1
        listing = dict(matched_listing)
        if unit_sale_price is not None:
            listing["salePrice"] = unit_sale_price
        order_vat_rate = _optional_float(item.get("vatRate"))
        if order_vat_rate is not None:
            listing["vatRate"] = order_vat_rate
        listing["linked_stock_label"] = _compose_linked_stock_label(
            listing.get("stock_items"),
            stock_item_label_map=stock_item_label_map,
            sale_quantity=qty_value,
            fallback=_safe_str(listing.get("linked_stock_label")),
        )
        listing_cost_items = [row for row in list(listing.get("cost_items") or []) if isinstance(row, dict)]
        for row in listing_cost_items:
            key = _cost_key(row.get("cost_key") or row.get("costKey"))
            if key in ORDER_LEVEL_COST_KEYS and key not in order_level_bindings:
                order_level_bindings[key] = dict(row)
        listing["cost_items"] = [row for row in listing_cost_items if not _is_order_level_cost(row)]

        listing_platform = canonical_platform_code(account_platform_map.get(listing.get("api_account_id")) or "")
        estimation = estimate_listing_profit(
            item=listing,
            stock_cost_map=stock_cost_map,
            stock_vat_rate_map=stock_vat_rate_map,
            stock_cost_meta_map=stock_cost_meta_map,
            stock_dimension_map=stock_dimension_map,
            cost_card_map=_cost_card_map_for(listing_platform),
            cargo_tariffs=cargo_tariffs_cache,
            platform_code=listing_platform,
            sale_quantity=qty_value,
        )

        line_ready = bool(estimation.get("estimated_profit_ready"))
        if not line_ready:
            complete = False
        for reason in list(estimation.get("estimated_profit_missing_reasons") or []):
            missing_reasons.append(f"{item_title}: {reason}")

        detail = dict(estimation.get("estimated_profit_detail") or {})
        results = dict(detail.get("results") or {})
        summary = dict(detail.get("summary") or {})
        for component in list(detail.get("applied_costs") or []):
            if not isinstance(component, dict):
                continue
            component_row = dict(component)
            note = _safe_str(component_row.get("note"))
            component_row["note"] = f"{item_title} | {note}" if note else item_title
            combined_applied_costs.append(component_row)
        line_desi = _optional_float(summary.get("stock_total_desi"))
        if line_desi is None:
            line_desi = _optional_float(summary.get("dimensional_weight"))
        if line_desi is not None and line_desi > 0:
            order_desi_total += line_desi
            has_order_desi = True
        sale_price_total = _optional_float(summary.get("sale_price"))
        if sale_price_total is None:
            sale_price_total = line_amount

        totals["sale_price"] += float(sale_price_total or 0.0)
        totals["payout_total"] += float(_optional_float(results.get("payout_total")) or 0.0)
        totals["total_costs"] += float(_optional_float(results.get("total_costs")) or 0.0)
        totals["input_vat_total"] += float(_optional_float(results.get("input_vat_total")) or 0.0)
        totals["sale_vat"] += float(_optional_float(results.get("sale_vat")) or 0.0)
        totals["net_vat"] += float(_optional_float(results.get("net_vat")) or 0.0)
        totals["estimated_profit"] += float(_optional_float(estimation.get("estimated_profit")) or 0.0)

        line_entries.append(
            {
                "title": item_title,
                "qty": qty_value,
                "qty_text": f"{qty_value:.0f}" if float(qty_value).is_integer() else f"{qty_value:.2f}",
                "unit_price": unit_sale_price,
                "unit_price_text": _amount_text(unit_sale_price),
                "amount": line_amount,
                "amount_text": _amount_text(line_amount),
                "listing_title": _safe_str(listing.get("title")) or item_title,
                "matched": True,
                "estimated_profit": estimation.get("estimated_profit"),
                "estimated_profit_text": _safe_str(estimation.get("estimated_profit_text")) or "-",
                "estimated_profit_ready": line_ready,
                "missing_reasons": list(estimation.get("estimated_profit_missing_reasons") or []),
                "detail": detail,
                "breakdown": list(estimation.get("estimated_profit_breakdown") or []),
            }
        )

    order_level_cost_map = _cost_card_map_for(order_platform_code)
    order_level_applied, order_level_missing, order_level_totals, order_level_complete = _resolve_order_level_costs(
        bindings_by_key=order_level_bindings,
        cost_card_map=order_level_cost_map,
        sale_price=totals["sale_price"],
        order_desi=order_desi_total if has_order_desi else None,
        platform_code=order_platform_code,
        cargo_tariffs=cargo_tariffs_cache,
    )
    combined_applied_costs.extend(order_level_applied)
    missing_reasons.extend(order_level_missing)
    if not order_level_complete:
        complete = False
    totals["total_costs"] += order_level_totals["cargo_total"] + order_level_totals["other_total"]
    totals["input_vat_total"] += order_level_totals["cargo_vat"] + order_level_totals["other_vat"]
    totals["payout_total"] -= order_level_totals["cargo_total"]
    totals["net_vat"] = totals["sale_vat"] - totals["input_vat_total"]
    totals["estimated_profit"] = totals["sale_price"] - totals["total_costs"] - totals["net_vat"]

    estimated_profit_value = totals["estimated_profit"] if line_entries and complete else None
    marketplace_deductions = totals["sale_price"] - totals["payout_total"]
    return {
        "estimated_profit": estimated_profit_value,
        "estimated_profit_text": _amount_text(estimated_profit_value) if complete and estimated_profit_value is not None else "-",
        "estimated_profit_ready": bool(complete and estimated_profit_value is not None),
        "estimated_profit_breakdown": [],
        "estimated_profit_detail": {
            "summary": {
                "order_number": order_number,
                "line_count": len(line_entries),
                "line_count_text": str(len(line_entries)),
                "matched_line_count": matched_line_count,
                "matched_line_count_text": str(matched_line_count),
                "item_quantity_total": total_quantity,
                "item_quantity_total_text": f"{total_quantity:.0f}" if float(total_quantity).is_integer() else f"{total_quantity:.2f}",
                "sale_price": totals["sale_price"],
                "sale_price_text": _amount_text(totals["sale_price"]),
            },
            "applied_costs": combined_applied_costs,
            "results": {
                "payout_total": totals["payout_total"],
                "payout_total_text": _amount_text(totals["payout_total"]),
                "marketplace_deductions": marketplace_deductions,
                "marketplace_deductions_text": _amount_text(marketplace_deductions),
                "total_costs": totals["total_costs"],
                "total_costs_text": _amount_text(totals["total_costs"]),
                "input_vat_total": totals["input_vat_total"],
                "input_vat_total_text": _amount_text(totals["input_vat_total"]),
                "sale_vat": totals["sale_vat"],
                "sale_vat_text": _amount_text(totals["sale_vat"]),
                "net_vat": totals["net_vat"],
                "net_vat_text": _delta_text(totals["net_vat"]),
                "estimated_profit": estimated_profit_value,
                "estimated_profit_text": _amount_text(estimated_profit_value),
            },
            "lines": line_entries,
        },
        "estimated_profit_missing_reasons": _dedupe_preserve_order(missing_reasons),
    }


def estimate_orders_profit_summary(
    *,
    orders: list[Any] | tuple[Any, ...],
    db_name: str = DB_NAME,
) -> Result:
    try:
        scoped_orders = [
            order
            for order in list(orders or [])
            if _safe_str(getattr(order, "orderNumber", ""))
        ]
        if not scoped_orders:
            return Result.ok(
                "Siparis olmadigi icin tahmini kar 0 TRY.",
                close_dialog=False,
                data={
                    "estimated_profit": 0.0,
                    "partial_estimated_profit": 0.0,
                    "estimated_profit_ready": True,
                    "revenue": 0.0,
                    "order_count": 0,
                    "complete_order_count": 0,
                    "incomplete_order_count": 0,
                    "missing_reasons": [],
                },
            )

        account_id_set: set[int] = set()
        for order in scoped_orders:
            account_id = _normalize_order_item_account_key(
                getattr(order, "api_account_id", None)
            )
            if isinstance(account_id, int) and account_id > 0:
                account_id_set.add(account_id)
        account_ids = sorted(account_id_set)
        catalog_res = load_order_profit_catalog(
            api_account_ids=account_ids or None,
            db_name=db_name,
        )
        if not catalog_res.success:
            return catalog_res
        catalog = dict(catalog_res.data or {})

        order_numbers = [
            _safe_str(getattr(order, "orderNumber", ""))
            for order in scoped_orders
        ]
        items_res = load_order_items_for_order_numbers(
            order_numbers,
            api_account_ids=account_ids or None,
            db_name=db_name,
        )
        if not items_res.success:
            return items_res
        items_by_order = dict((items_res.data or {}).get("items_by_order") or {})

        complete_count = 0
        partial_profit = 0.0
        revenue = 0.0
        missing_reasons: list[str] = []
        rows: list[dict[str, Any]] = []
        for order in scoped_orders:
            order_number = _safe_str(getattr(order, "orderNumber", ""))
            account_id = _normalize_order_item_account_key(getattr(order, "api_account_id", None))
            items = list(items_by_order.get((account_id, order_number), []) or [])
            estimation = estimate_order_profit(
                order=order,
                order_items=items,
                catalog=catalog,
            )
            ready = bool(estimation.get("estimated_profit_ready"))
            value = _optional_float(estimation.get("estimated_profit"))
            if ready and value is not None:
                complete_count += 1
                partial_profit += value
            for reason in list(estimation.get("estimated_profit_missing_reasons") or []):
                clean_reason = _safe_str(reason)
                if clean_reason:
                    missing_reasons.append(f"{order_number}: {clean_reason}")
            order_revenue = _optional_float(
                getattr(order, "packageTotalPrice", None)
                or getattr(order, "totalPrice", None)
            )
            revenue += float(order_revenue or 0.0)
            rows.append(
                {
                    "order_number": order_number,
                    "estimated_profit_ready": ready,
                    "estimated_profit": value if ready else None,
                    "missing_reasons": list(estimation.get("estimated_profit_missing_reasons") or []),
                }
            )

        total_count = len(scoped_orders)
        incomplete_count = total_count - complete_count
        ready = incomplete_count == 0
        estimated_profit = partial_profit if ready else None
        deduped_reasons = _dedupe_preserve_order(missing_reasons)
        if ready:
            message = f"{total_count} siparis icin tahmini kar {estimated_profit:.2f} TRY."
        else:
            message = (
                f"Tahmini kar hesaplanamadi: {incomplete_count}/{total_count} sipariste "
                "maliyet, para birimi veya ilan esleme verisi eksik ya da uyumsuz."
            )
        return Result.ok(
            message,
            close_dialog=False,
            data={
                "estimated_profit": estimated_profit,
                "partial_estimated_profit": partial_profit if complete_count else None,
                "estimated_profit_ready": ready,
                "revenue": revenue,
                "order_count": total_count,
                "complete_order_count": complete_count,
                "incomplete_order_count": incomplete_count,
                "missing_reasons": deduped_reasons[:50],
                "rows": rows,
            },
        )
    except Exception as exc:
        return Result.fail(map_error_to_message(exc), error=exc, close_dialog=False)


def write_order_profit_payload_to_excel(
    *,
    order_number: str,
    estimated_profit_payload: dict[str, Any],
    output_path: str,
) -> Result:
    try:
        payload = dict(estimated_profit_payload or {})
        detail = dict(payload.get("estimated_profit_detail") or {})
        results = dict(detail.get("results") or {})
        summary = dict(detail.get("summary") or {})
        lines = [row for row in list(detail.get("lines") or []) if isinstance(row, dict)]
        clean_order_number = _safe_str(order_number) or _safe_str(summary.get("order_number")) or "-"
        path = Path(str(output_path or "").strip()).expanduser()
        if not str(path):
            return Result.fail("Excel export icin output_path gerekir.", close_dialog=False)
        if path.suffix.casefold() != ".xlsx":
            path = path.with_suffix(".xlsx")
        path.parent.mkdir(parents=True, exist_ok=True)

        sale_text = _amount_text(_optional_float(results.get("sale_price") or summary.get("sale_price")))
        cost_text = _safe_str(results.get("total_costs_text")) or _amount_text(_optional_float(results.get("total_costs")))
        profit_text = _safe_str(payload.get("estimated_profit_text")) or _safe_str(results.get("estimated_profit_text")) or "-"

        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill

        wb = Workbook()
        ws = wb.active
        ws.title = "Ozet"
        ws.append(["Alan", "Değer"])
        for label, value in (
            ("Sipariş", clean_order_number),
            ("Toplam ürün adedi", _safe_str(summary.get("item_quantity_total_text")) or "-"),
            ("Toplam satış", sale_text),
            ("Toplam maliyet", cost_text),
            ("Tahmini kar", profit_text),
            ("Durum", "Hazır" if payload.get("estimated_profit_ready") else "Bekliyor"),
        ):
            ws.append([label, value])

        line_ws = wb.create_sheet("Satis")
        line_ws.append(["Ürün", "İlan", "Adet", "Birim", "Tutar"])
        for line in lines:
            line_ws.append(
                [
                    _safe_str(line.get("title")) or "Ürün",
                    _safe_str(line.get("listing_title")) or "",
                    _safe_str(line.get("qty_text")) or "-",
                    _safe_str(line.get("unit_price_text")) or "-",
                    _safe_str(line.get("amount_text")) or _amount_text(_optional_float(line.get("amount"))),
                ]
            )

        cost_ws = wb.create_sheet("Maliyetler")
        cost_ws.append(["Kalem", "Açıklama", "Tutar"])
        applied_costs = [row for row in list(detail.get("applied_costs") or []) if isinstance(row, dict)]
        if not applied_costs:
            for line in lines:
                line_detail = dict(line.get("detail") or {})
                for cost in list(line_detail.get("applied_costs") or []):
                    if isinstance(cost, dict):
                        applied_costs.append(dict(cost))
        for cost in applied_costs:
            amount = _optional_float(cost.get("gross"))
            if amount is None:
                amount = _optional_float(cost.get("value"))
            cost_ws.append(
                [
                    _safe_str(cost.get("label")) or _safe_str(cost.get("key")) or "Maliyet",
                    _safe_str(cost.get("note")) or "",
                    _amount_text(amount),
                ]
            )
        cost_ws.append(["Maliyet Toplamı", "", cost_text])

        for sheet in wb.worksheets:
            for cell in sheet[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="0F172A")
            sheet.column_dimensions["A"].width = 28
            sheet.column_dimensions["B"].width = 58
            sheet.column_dimensions["C"].width = 18
            sheet.column_dimensions["D"].width = 18
            sheet.column_dimensions["E"].width = 18
        wb.save(path)
        return Result.ok(
            "Siparis kar Excel dosyasi olusturuldu.",
            close_dialog=False,
            data={
                "order_number": clean_order_number,
                "output_path": str(path),
                "exists": path.exists(),
                "size_bytes": path.stat().st_size if path.exists() else 0,
            },
        )
    except Exception as exc:
        return Result.fail(map_error_to_message(exc), error=exc, close_dialog=False)


def action_export_order_profit_excel(
    *,
    order_number: str,
    output_path: str,
    estimated_profit_payload: dict[str, Any],
    api_account_id: int | None = None,
) -> Result:
    # api_account_id is an immutable scope/audit field. The workbook writer uses
    # only the already account-scoped payload prepared by the agentic read path.
    del api_account_id
    return write_order_profit_payload_to_excel(
        order_number=order_number,
        estimated_profit_payload=dict(estimated_profit_payload or {}),
        output_path=output_path,
    )


def action_load_order_snapshot(*, start_date_ms: int | None = None, end_date_ms: int | None = None,
                               api_account_ids: list[int] | None = None, filters: dict[str, Any] | None = None,
                               sort: str = "order_date_asc", db_name: str | None = None) -> Result:
    """Complete ordered dataset for manual exports and AI; never a page preview."""
    from Orders.processors.agentic_repository import query_orders
    options = {"db_name": db_name} if db_name else {}
    return query_orders(start_date_ms=start_date_ms, end_date_ms=end_date_ms, api_account_ids=api_account_ids,
                        filters=filters, sort=sort, mode="snapshot", **options)


def action_load_order_data_freshness(*, api_account_ids=None, db_name=None) -> Result:
    from Orders.processors.data_freshness import read_local_order_freshness
    options = {"db_name": db_name} if db_name else {}
    return read_local_order_freshness(api_account_ids=api_account_ids, **options)


def action_export_report_snapshot_xlsx(
    *,
    report_snapshot: dict[str, Any],
    output_path: str,
) -> Result:
    """Thin manual action wrapper around the shared Core artifact writer."""

    return export_report_snapshot_xlsx(
        report_snapshot=dict(report_snapshot or {}),
        output_path=output_path,
    )


def _normalize_cargo_provider_lookup_text(value: object) -> str:
    text = str(value or "").strip()
    if not text:
        return ""
    normalized = unicodedata.normalize("NFKD", text.casefold())
    cleaned = "".join(ch for ch in normalized if not unicodedata.combining(ch))
    cleaned = cleaned.replace("ı", "i")
    for token in ("(", ")", "-", "_", "/", ","):
        cleaned = cleaned.replace(token, " ")
    return " ".join(cleaned.split())


def cargo_provider_name_for_code(code: str, *, platform_code: str | None = None) -> str:
    key = str(code or "").strip().upper()
    return shared_cargo_provider_name_for_code(key, platform_code=platform_code) or key or "-"


@lru_cache(maxsize=32)
def _cached_cargo_provider_rows(platform_code: str) -> tuple[tuple[str, str, str, str, str], ...]:
    rows: list[tuple[str, str, str, str, str]] = []
    seen: set[str] = set()
    normalized_platform = canonical_platform_code(platform_code)
    if normalized_platform.casefold() in ("", "all", "tümü"):
        normalized_platform = ""
    candidate_rows = changeable_cargo_provider_records(normalized_platform or None)
    for row in candidate_rows:
        if not isinstance(row, dict):
            continue
        code = str(row.get("code") or "").strip().upper()
        name = str(row.get("name") or "").strip()
        if not code or not name or code in seen:
            continue
        seen.add(code)
        rows.append(
            (
                code,
                name,
                _normalize_cargo_provider_lookup_text(code),
                _normalize_cargo_provider_lookup_text(name),
                _normalize_cargo_provider_lookup_text(f"{name} ({code})"),
            )
        )
    return tuple(rows)


def list_changeable_cargo_providers(platform_code: str | None = None) -> list[dict[str, str]]:
    return [
        {"code": code, "name": name}
        for code, name, _normalized_code, _normalized_name, _normalized_label in _cached_cargo_provider_rows(
            str(platform_code or "")
        )
    ]


def cargo_provider_code_for_name(name_or_code: str, *, platform_code: str | None = None) -> str | None:
    needle = _normalize_cargo_provider_lookup_text(name_or_code)
    if not needle:
        return None
    for code, _name, normalized_code, normalized_name, normalized_label in _cached_cargo_provider_rows(str(platform_code or "")):
        if needle == normalized_code:
            return code
        if needle == normalized_name:
            return code
        if needle == normalized_label:
            return code
        if needle in normalized_name:
            return code
        if needle in normalized_label:
            return code
    return None


def _run_async_safely(coro):
    loop = asyncio.new_event_loop()
    try:
        return loop.run_until_complete(coro)
    finally:
        try:
            from Core.utils.request_utils import close_async_request_clients

            loop.run_until_complete(close_async_request_clients())
            # Windows Proactor transports finish SSL/socket close callbacks
            # on the next loop ticks; drain them before closing this private loop.
            loop.run_until_complete(asyncio.sleep(0.05))
        except Exception:
            pass
        try:
            loop.run_until_complete(loop.shutdown_asyncgens())
        except Exception:
            pass
        loop.close()


def _date_start_epoch_ms(day_value: date) -> int:
    dt = datetime(
        int(day_value.year),
        int(day_value.month),
        int(day_value.day),
        0,
        0,
        0,
        tzinfo=TZ_GMT3,
    )
    return int(dt.timestamp() * 1000)


def _date_end_epoch_ms(day_value: date) -> int:
    dt = datetime(
        int(day_value.year),
        int(day_value.month),
        int(day_value.day),
        23,
        59,
        59,
        tzinfo=TZ_GMT3,
    )
    return int(dt.timestamp() * 1000)


def _safe_int(value: Any, default_value: int = 0) -> int:
    try:
        return int(value)
    except Exception:
        return int(default_value)


def _safe_float(value: Any, default_value: float = 0.0) -> float:
    try:
        text = str(value).strip().replace(",", ".")
        return float(text)
    except Exception:
        return float(default_value)


def _parse_claim_datetime(value: Any) -> datetime | None:
    if value in (None, ""):
        return None
    try:
        if isinstance(value, datetime):
            dt = value
            if dt.tzinfo is not None:
                return dt.astimezone(TZ_GMT3).replace(tzinfo=None)
            return dt
        if isinstance(value, (int, float)):
            ts = float(value)
            if ts <= 0:
                return None
            if ts > 10 ** 11:
                ts = ts / 1000.0
            return datetime.fromtimestamp(ts, tz=TZ_GMT3).replace(tzinfo=None)
        if isinstance(value, str):
            text = value.strip()
            if not text:
                return None
            if text.isdigit():
                return _parse_claim_datetime(int(text))
            try:
                dt = datetime.fromisoformat(text.replace("Z", "+00:00"))
                if dt.tzinfo is not None:
                    return dt.astimezone(TZ_GMT3).replace(tzinfo=None)
                return dt
            except Exception:
                return None
    except Exception:
        return None
    return None


def _claim_date_from_rows(claim: dict, claim_item: dict) -> datetime | None:
    candidates = (
        claim_item.get("claimDate"),
        claim_item.get("createdDate"),
        claim_item.get("claimCreationDate"),
        claim.get("claimDate"),
        claim.get("createdDate"),
        claim.get("claimCreationDate"),
        claim.get("lastModifiedDate"),
    )
    for candidate in candidates:
        dt = _parse_claim_datetime(candidate)
        if dt is not None:
            return dt
    return None


def _extract_structured_text(value: Any, *, default_text: str = "") -> str:
    if isinstance(value, dict):
        for key in ("name", "code", "value", "title", "message"):
            text = str(value.get(key) or "").strip()
            if text:
                return text
        return str(default_text or "").strip()
    return str(value or default_text or "").strip()


def _extract_order_line_quantity(order_line: dict, claim_item: dict) -> int:
    direct_qty = order_line.get("quantity")
    if direct_qty not in (None, ""):
        return max(1, _safe_int(direct_qty, default_value=1))

    line_items = order_line.get("lineItems")
    if isinstance(line_items, list) and line_items:
        first_line = line_items[0] if isinstance(line_items[0], dict) else {}
        qty = first_line.get("quantity")
        if qty not in (None, ""):
            return max(1, _safe_int(qty, default_value=1))

    claim_qty = claim_item.get("quantity")
    if claim_qty not in (None, ""):
        return max(1, _safe_int(claim_qty, default_value=1))
    return 1


async def _fetch_claims_for_account(
    account: Any,
    *,
    start_ms: int,
    end_ms: int,
    size: int,
    claim_item_status: str | None = None,
    claim_ids: list[str] | None = None,
) -> tuple[list[dict], list[str]]:
    rows: list[dict] = []
    errors: list[str] = []
    seen_keys: set[tuple[Any, ...]] = set()

    seller_id = str(getattr(account, "account_id", "") or "").strip()
    comp_name = str(getattr(account, "comp_name", "") or seller_id or "-")

    try:
        api = _trendyol_api_for_account(account)
    except Exception as exc:
        errors.append(f"{comp_name}: {map_error_to_message(exc)}")
        return rows, errors
    page = 0
    total_pages = None
    while True:
        try:
            result = await api.find_claims(
                start_date=start_ms,
                end_date=end_ms,
                page=page,
                size=size,
                claim_item_status=claim_item_status,
                claim_ids=claim_ids,
            )
        except Exception as e:
            errors.append(f"{comp_name}: {map_error_to_message(e)}")
            break
        if not result.success:
            errors.append(f"{comp_name}: {result.message}")
            break

        payload = result.data or {}
        content = payload.get("content", []) or []
        for claim in content:
            if not isinstance(claim, dict):
                continue
            normalized_row = normalize_return_claim(claim, int(getattr(account, "pk", 0) or 0))
            if not str(normalized_row.get("claimId") or "").strip():
                continue
            if int(normalized_row.get("lastModifiedDate") or 0) <= 0:
                continue
            dedupe_key = (
                getattr(account, "pk", None),
                str(normalized_row.get("claimId") or "").strip(),
                int(normalized_row.get("lastModifiedDate") or 0),
            )
            if dedupe_key in seen_keys:
                continue
            seen_keys.add(dedupe_key)
            rows.append(normalized_row)

        if total_pages is None:
            total_pages = _safe_int(payload.get("totalPages"), default_value=0)
        if not content and total_pages and page + 1 < total_pages:
            errors.append(
                f"{comp_name}: claim pagination eksik döndü "
                f"(sayfa {page + 1}/{total_pages})."
            )
            break
        page += 1
        if total_pages and page >= total_pages:
            break
        if not content:
            break
        if len(content) < size and (not total_pages or page >= total_pages):
            break

    return rows, errors


async def _fetch_returns_from_active_companies_async(
    accounts: list[Any],
    *,
    start_ms: int,
    end_ms: int,
    size: int,
    claim_item_status: str | None = None,
    claim_ids: list[str] | None = None,
    progress_callback: Optional[Callable[[int, int], None]] = None,
) -> Result:
    all_rows: list[dict] = []
    all_errors: list[str] = []
    total_steps = max(len(accounts), 1)
    current_step = 0

    for account in accounts:
        try:
            rows, errors = await _fetch_claims_for_account(
                account,
                start_ms=start_ms,
                end_ms=end_ms,
                size=size,
                claim_item_status=claim_item_status,
                claim_ids=claim_ids,
            )
            if rows:
                all_rows.extend(rows)
            if errors:
                all_errors.extend(errors)
        finally:
            current_step += 1
            if progress_callback:
                progress_callback(current_step, total_steps)

    all_rows.sort(
        key=lambda row: (
            int(row.get("claimDate") or row.get("orderDate") or 0),
            str(row.get("claimNumber") or row.get("claimId") or ""),
        ),
        reverse=True,
    )

    if all_errors and not all_rows:
        return Result.fail(
            "İadeler alınamadı.",
            close_dialog=False,
            data={"return_claim_list": [], "errors": all_errors},
        )

    msg = f"İadeler yüklendi. ({len(all_rows)} kayıt)"
    if all_errors:
        msg = f"{msg} | Hata: {len(all_errors)} mağaza"
    return Result.ok(
        msg,
        close_dialog=False,
        data={
            "return_claim_list": all_rows,
            "errors": all_errors,
            "start_ms": start_ms,
            "end_ms": end_ms,
        },
    )


async def _refresh_active_return_claims_from_api_async(
    accounts: list[Any],
    *,
    active_status_codes: list[str] | None = None,
    size: int = 50,
    status_scan_days: int = 60,
    scan_statuses: bool = False,
) -> Result:
    try:
        statuses = list(active_status_codes or RETURN_ACTIVE_STATUS_CODES)
        account_by_id: dict[int, Any] = {}
        for account in accounts or []:
            try:
                account_id = int(getattr(account, "pk", None))
            except Exception:
                continue
            if account_id > 0:
                account_by_id[account_id] = account

        empty_counts = {"claims_inserted": 0, "claims_restored": 0, "claims_stale_marked": 0}
        if not account_by_id:
            return Result.ok(
                "Yenilenecek Trendyol iade hesabı bulunamadı.",
                close_dialog=False,
                data={"changed": False, "counts": empty_counts, "checked_claims": 0, "errors": []},
            )

        local_res = get_latest_return_claims_filtered(
            claim_statuses=statuses,
            api_account_ids=sorted(account_by_id.keys()),
        )
        if not local_res.success:
            return local_res

        claim_ids_by_account: dict[int, list[str]] = {}
        for claim in (local_res.data or {}).get("claims", []) or []:
            claim_id = str(getattr(claim, "claimId", "") or "").strip()
            if not claim_id:
                continue
            try:
                account_id = int(getattr(claim, "api_account_id", None))
            except Exception:
                continue
            if account_id not in account_by_id:
                continue
            claim_ids_by_account.setdefault(account_id, []).append(claim_id)

        for account_id, claim_ids in list(claim_ids_by_account.items()):
            claim_ids_by_account[account_id] = list(dict.fromkeys(claim_ids))

        checked_claims = sum(len(ids) for ids in claim_ids_by_account.values())
        if checked_claims <= 0:
            return Result.ok(
                "Açık iade claim yenilemesi için lokal kayıt bulunamadı.",
                close_dialog=False,
                data={"changed": False, "counts": empty_counts, "checked_claims": 0, "errors": []},
            )

        refreshed_rows: list[dict] = []
        fetched_keys: set[tuple[int, str]] = set()
        failed_account_ids: set[int] = set()
        errors: list[str] = []
        now_ms = _now_epoch_ms()
        scan_start_ms = now_ms - (max(1, int(status_scan_days)) * 24 * 60 * 60 * 1000)

        for status_code in (statuses if scan_statuses else []):
            for account_id, account in account_by_id.items():
                fetch_res = await _fetch_returns_from_active_companies_async(
                    [account],
                    start_ms=scan_start_ms,
                    end_ms=now_ms,
                    size=max(10, min(200, _safe_int(size, default_value=50))),
                    claim_item_status=status_code,
                )
                fetch_errors = (fetch_res.data or {}).get("errors", []) if isinstance(fetch_res.data, dict) else []
                if not fetch_res.success or fetch_errors:
                    if not fetch_res.success:
                        errors.append(
                            str(fetch_res.message or f"{account_id}: {status_code} iade status taraması başarısız.")
                        )
                    errors.extend([str(err) for err in (fetch_errors or []) if str(err).strip()])
                    continue
                for row in (fetch_res.data or {}).get("return_claim_list", []) or []:
                    if not isinstance(row, dict):
                        continue
                    claim_id = str(row.get("claimId") or "").strip()
                    if not claim_id:
                        continue
                    fetched_keys.add((account_id, claim_id))
                    refreshed_rows.append(row)

        for account_id, claim_ids in claim_ids_by_account.items():
            fetch_res = await _fetch_returns_from_active_companies_async(
                [account_by_id[account_id]],
                start_ms=0,
                end_ms=now_ms,
                size=max(10, min(200, _safe_int(size, default_value=50))),
                claim_ids=claim_ids,
            )
            fetch_errors = (fetch_res.data or {}).get("errors", []) if isinstance(fetch_res.data, dict) else []
            if not fetch_res.success or fetch_errors:
                failed_account_ids.add(account_id)
                if not fetch_res.success:
                    errors.append(str(fetch_res.message or f"{account_id}: iade claim refresh başarısız."))
                errors.extend([str(err) for err in (fetch_errors or []) if str(err).strip()])
                continue

            for row in (fetch_res.data or {}).get("return_claim_list", []) or []:
                if not isinstance(row, dict):
                    continue
                claim_id = str(row.get("claimId") or "").strip()
                if not claim_id:
                    continue
                fetched_keys.add((account_id, claim_id))
                refreshed_rows.append(row)

        changed = False
        counts = dict(empty_counts)
        save_payload: dict[str, Any] | None = None
        if refreshed_rows:
            save_res = save_return_claims_to_db(
                Result.ok(
                    "Açık iade claim refresh",
                    close_dialog=False,
                    data={"return_claim_list": refreshed_rows},
                )
            )
            if not save_res.success:
                return save_res
            save_payload = save_res.data or {}
            changed = changed or bool(save_payload.get("changed"))
            for key, value in (save_payload.get("counts") or {}).items():
                if key in counts:
                    try:
                        counts[key] += int(value or 0)
                    except Exception:
                        pass

        stale_keys: list[dict[str, Any]] = []
        for account_id, claim_ids in claim_ids_by_account.items():
            if account_id in failed_account_ids:
                continue
            for claim_id in claim_ids:
                if (account_id, claim_id) not in fetched_keys:
                    stale_keys.append({"api_account_id": account_id, "claim_id": claim_id})

        stale_payload: dict[str, Any] | None = None
        if stale_keys:
            stale_res = mark_return_claims_stale(
                claim_ids=[],
                claim_keys=stale_keys,
                reason="manual fetch active claimIds refresh returned no current record",
            )
            if not stale_res.success:
                return stale_res
            stale_payload = stale_res.data or {}
            stale_count = int(stale_payload.get("stale_marked", 0) or 0)
            counts["claims_stale_marked"] += stale_count
            changed = changed or stale_count > 0

        return Result.ok(
            f"Açık iade claim refresh tamamlandı. Kontrol edilen: {checked_claims}",
            close_dialog=False,
            data={
                "changed": changed,
                "counts": counts,
                "checked_claims": checked_claims,
                "refreshed_claims": len(fetched_keys),
                "stale_claims": len(stale_keys),
                "errors": errors,
                "save_result": save_payload,
                "stale_result": stale_payload,
            },
        )
    except Exception as e:
        return Result.fail(map_error_to_message(e), error=e, close_dialog=False)


def _resolve_returns_date_window(
    *,
    date_from: date | None = None,
    date_to: date | None = None,
) -> tuple[date, date, int, int]:
    resolved_from = coerce_to_date(date_from) if date_from is not None else None
    resolved_to = coerce_to_date(date_to) if date_to is not None else None
    today = datetime.now(TZ_GMT3).date()
    if resolved_from is None:
        resolved_from = today - timedelta(days=30)
    if resolved_to is None:
        resolved_to = today + timedelta(days=1)
    if resolved_to < resolved_from:
        resolved_to = resolved_from
    return (
        resolved_from,
        resolved_to,
        _date_start_epoch_ms(resolved_from),
        _date_end_epoch_ms(resolved_to),
    )


RETURN_ACTIVE_STATUS_FILTER_CODE = "__open_returns__"
RETURN_ACTIVE_STATUS_FILTER_LABEL = "Aktif Iadeler"
RETURN_ACTIVE_STATUS_CODES = [
    "Created",
    "WaitingInAction",
    "WaitingFraudCheck",
    "InAnalysis",
    "Unresolved",
]
RETURN_WAITING_ACTION_STATUS_CODES = ["WaitingInAction"]


def return_status_codes_for_load(status_code: str | None) -> list[str] | None:
    code = str(status_code or "").strip()
    if _is_all_like_value(code):
        return None
    key = _normalize_filter_choice_key(code)
    if code == RETURN_ACTIVE_STATUS_FILTER_CODE or key in {
        "aktif iadeler",
        "aktif iade",
        "acik iadeler",
        "acik iade",
        "open returns",
        "open return",
    }:
        return list(RETURN_ACTIVE_STATUS_CODES)
    label_code = _status_code_from_label(code, TRENDYOL_RETURN_STATUS_LABELS)
    if label_code:
        return [label_code]
    if key == "aksiyon bekliyor":
        return ["WaitingInAction"]
    return [code]


def _platform_codes_for_load(platform_code: str | None) -> list[str] | None:
    raw = str(platform_code or "").strip()
    if _is_all_like_value(raw):
        return None
    normalized = canonical_platform_code(raw)
    if not normalized:
        return None
    return [normalized]


def load_returns_for_source_filters(
    *,
    date_from: date | None = None,
    date_to: date | None = None,
    shipment_status_code: str | None = None,
    platform_code: str | None = None,
    sync_if_empty: bool = False,
) -> Result:
    try:
        _resolved_from, _resolved_to, start_ms, end_ms = _resolve_returns_date_window(
            date_from=date_from,
            date_to=date_to,
        )
        requested_platforms = _platform_codes_for_load(platform_code)
        read_res = get_latest_return_claims_filtered(
            start_ms=start_ms,
            end_ms=end_ms,
            claim_statuses=return_status_codes_for_load(shipment_status_code),
            platform_codes=requested_platforms,
        )
        if not read_res.success:
            return read_res

        records = (read_res.data or {}).get("claims", []) or []
        if records or not sync_if_empty:
            return Result.ok(
                "İadeler yüklendi.",
                close_dialog=False,
                data={"records": records},
            )

        if requested_platforms and "TRENDYOL" not in requested_platforms:
            return Result.ok(
                "İadeler yüklendi.",
                close_dialog=False,
                data={"records": records},
            )

        companies_res = get_active_companies()
        if not companies_res.success:
            return companies_res
        accounts = [
            account
            for account in ((companies_res.data or {}).get("records", []) or [])
            if str(getattr(account, "platform", "") or "").strip().upper() == "TRENDYOL"
        ]
        if not accounts:
            return Result.ok(
                "Aktif Trendyol mağazası bulunamadı.",
                close_dialog=False,
                data={"records": []},
            )

        fetch_res = _run_async_safely(
            _fetch_returns_from_active_companies_async(
                accounts,
                start_ms=start_ms,
                end_ms=end_ms,
                size=50,
            )
        )
        if not fetch_res.success:
            return fetch_res

        save_res = save_return_claims_to_db(fetch_res)
        if not save_res.success:
            return save_res

        reread_res = get_latest_return_claims_filtered(
            start_ms=start_ms,
            end_ms=end_ms,
            claim_statuses=return_status_codes_for_load(shipment_status_code),
            platform_codes=requested_platforms,
        )
        if not reread_res.success:
            return reread_res

        return Result.ok(
            "İadeler yüklendi.",
            close_dialog=False,
            data={
                "records": (reread_res.data or {}).get("claims", []) or [],
                "errors": (fetch_res.data or {}).get("errors", []) or [],
                "db_result": save_res.data or {},
            },
        )
    except Exception as e:
        return Result.fail(map_error_to_message(e), error=e, close_dialog=False)


def load_returns_from_db(
    *,
    date_from: date | None = None,
    date_to: date | None = None,
) -> Result:
    return load_returns_for_source_filters(date_from=date_from, date_to=date_to)
    try:
        _resolved_from, _resolved_to, start_ms, end_ms = _resolve_returns_date_window(
            date_from=date_from,
            date_to=date_to,
        )
        read_res = get_latest_return_claims(start_ms=start_ms, end_ms=end_ms)
        if not read_res.success:
            return read_res

        return Result.ok(
            "İadeler yüklendi.",
            close_dialog=False,
            data={"records": (read_res.data or {}).get("claims", []) or []},
        )
    except Exception as e:
        return Result.fail(map_error_to_message(e), error=e, close_dialog=False)


def load_returns_claims(
    *,
    date_from: date | None = None,
    date_to: date | None = None,
    size: int = 50,
) -> Result:
    try:
        _resolved_from, _resolved_to, start_ms, end_ms = _resolve_returns_date_window(
            date_from=date_from,
            date_to=date_to,
        )
        page_size = max(10, min(200, _safe_int(size, default_value=50)))
        companies_res = get_active_companies()
        if not companies_res.success:
            return companies_res
        accounts = [
            account
            for account in ((companies_res.data or {}).get("records", []) or [])
            if str(getattr(account, "platform", "") or "").strip().upper() == "TRENDYOL"
        ]
        if not accounts:
            return Result.ok(
                "Aktif Trendyol mağazası bulunamadı.",
                close_dialog=False,
                data={"records": []},
            )

        fetch_res = _run_async_safely(
            _fetch_returns_from_active_companies_async(
                accounts,
                start_ms=start_ms,
                end_ms=end_ms,
                size=page_size,
            )
        )
        if not fetch_res.success:
            return fetch_res

        save_res = save_return_claims_to_db(fetch_res)
        if not save_res.success:
            return save_res

        read_res = get_latest_return_claims(start_ms=start_ms, end_ms=end_ms)
        if not read_res.success:
            return read_res

        return Result.ok(
            "İadeler yüklendi.",
            close_dialog=False,
            data={
                "records": (read_res.data or {}).get("claims", []) or [],
                "errors": (fetch_res.data or {}).get("errors", []) or [],
                "db_result": save_res.data or {},
            },
        )
    except Exception as e:
        return Result.fail(map_error_to_message(e), error=e, close_dialog=False)


def approve_return_claim(
    order,
    *,
    date_from: date | None = None,
    date_to: date | None = None,
) -> Result:
    try:
        if order is None:
            return Result.fail("Önce bir iade seçmelisin.", close_dialog=False)

        claim_id = str(getattr(order, "claimId", "") or getattr(order, "claimNumber", "") or "").strip()
        if not claim_id:
            return Result.fail("İade claim id bilgisi eksik.", close_dialog=False)

        waiting_line_ids = [
            str(row or "").strip()
            for row in (getattr(order, "waitingClaimLineItemIds", None) or [])
            if str(row or "").strip()
        ]
        waiting_line_ids = list(dict.fromkeys(waiting_line_ids))
        if not waiting_line_ids:
            return Result.fail("Bu iade için onay bekleyen satır bulunamadı.", close_dialog=False)

        api_account = getattr(order, "api_account", None)
        if api_account is None:
            return Result.fail("İadenin bağlı olduğu mağaza bilgisi bulunamadı.", close_dialog=False)

        try:
            api = _trendyol_api_for_account(api_account)
        except Exception as exc:
            return Result.fail(map_error_to_message(exc), error=exc, close_dialog=False)
        approve_res = _run_async_safely(
            api.approve_claim_line_items(
                claim_id=claim_id,
                claim_line_item_ids=waiting_line_ids,
            )
        )
        if not approve_res.success:
            return approve_res

        resolved_from = coerce_to_date(date_from) if date_from is not None else None
        resolved_to = coerce_to_date(date_to) if date_to is not None else None
        today = datetime.now(TZ_GMT3).date()
        if resolved_from is None:
            resolved_from = today - timedelta(days=30)
        if resolved_to is None:
            resolved_to = today + timedelta(days=1)
        if resolved_to < resolved_from:
            resolved_to = resolved_from

        refresh_res = _run_async_safely(
            _fetch_returns_from_active_companies_async(
                [api_account],
                start_ms=_date_start_epoch_ms(resolved_from),
                end_ms=_date_end_epoch_ms(resolved_to),
                size=50,
                claim_ids=[claim_id],
            )
        )
        if refresh_res.success:
            save_res = save_return_claims_to_db(refresh_res)
            if save_res.success:
                return Result.ok(
                    "İade onaylandı ve veritabanı güncellendi.",
                    close_dialog=False,
                    data={
                        "claim_id": claim_id,
                        "approved_line_item_ids": waiting_line_ids,
                        "db_result": save_res.data or {},
                    },
                )

        return Result.ok(
            "İade onaylandı. Güncel durum veritabanına henüz yazılamadı.",
            close_dialog=False,
            data={
                "claim_id": claim_id,
                "approved_line_item_ids": waiting_line_ids,
            },
        )

    except Exception as e:
        return Result.fail(map_error_to_message(e), error=e, close_dialog=False)


def change_order_cargo_provider(order, *, cargo_provider_code: str) -> Result:
    try:
        if order is None:
            return Result.fail("Önce bir sipariş seçmelisin.", close_dialog=False)

        provider_code = str(cargo_provider_code or "").strip().upper()
        api_account = getattr(order, "api_account", None)
        if api_account is None:
            return Result.fail("Siparişin bağlı olduğu mağaza bilgisi bulunamadı.", close_dialog=False)

        platform_code = canonical_platform_code(getattr(api_account, "platform", None))
        provider_options = {row["code"] for row in list_changeable_cargo_providers(platform_code)}
        if provider_code not in provider_options:
            return Result.fail("Geçersiz kargo firması kodu.", close_dialog=False)

        package_id = getattr(order, "id", None)
        if package_id in (None, "", 0):
            return Result.fail("Sipariş paket id bilgisi eksik.", close_dialog=False)

        try:
            api = _trendyol_api_for_account(api_account)
        except Exception as exc:
            return Result.fail(map_error_to_message(exc), error=exc, close_dialog=False)
        api_res = _run_async_safely(
            api.change_cargo_provider(package_id=package_id, cargo_provider=provider_code)
        )
        if not api_res.success:
            return api_res

        provider_name = cargo_provider_name_for_code(provider_code, platform_code=platform_code)
        try:
            setattr(order, "cargoProviderName", provider_name)
        except Exception:
            pass

        order_pk = getattr(order, "pk", None)
        if order_pk not in (None, "", 0):
            db_res = update_records(
                model=_order_data_model(),
                filters={"pk": int(order_pk)},
                update_data={"cargoProviderName": provider_name},
            )
            if not db_res.success:
                return db_res

        return Result.ok(
            f"Kargo firması güncellendi: {provider_name}",
            close_dialog=False,
            data={
                "cargo_provider_code": provider_code,
                "cargo_provider_name": provider_name,
                "package_id": package_id,
            },
        )
    except Exception as e:
        return Result.fail(map_error_to_message(e), error=e, close_dialog=False)


# ============================================================
# 2) OrdersManagerWindow helpers
# ============================================================

def change_orders_cargo_provider_bulk(orders: list, *, cargo_provider_code: str) -> Result:
    try:
        rows = list(orders or [])
        if not rows:
            return Result.fail("Toplu güncelleme için seçili sipariş yok.", close_dialog=False)

        provider_code = str(cargo_provider_code or "").strip().upper()
        provider_options = {row["code"] for row in list_changeable_cargo_providers()}
        if provider_code not in provider_options:
            return Result.fail("Geçersiz kargo firması kodu.", close_dialog=False)

        unique_orders: list = []
        seen: set[tuple[str, str]] = set()
        for order in rows:
            order_no = str(getattr(order, "orderNumber", "") or "").strip()
            package_id = str(getattr(order, "id", "") or "").strip()
            key = (order_no, package_id)
            if key in seen:
                continue
            seen.add(key)
            unique_orders.append(order)

        success_count = 0
        failed: list[dict[str, str]] = []
        for order in unique_orders:
            one = change_order_cargo_provider(order, cargo_provider_code=provider_code)
            if one.success:
                success_count += 1
                continue
            failed.append(
                {
                    "order_number": str(getattr(order, "orderNumber", "") or "-"),
                    "message": str(one.message or "Bilinmeyen hata"),
                }
            )

        failed_count = len(failed)
        provider_name = cargo_provider_name_for_code(provider_code)
        if failed_count > 0:
            return Result.fail(
                f"Toplu kargo değişikliği tamamlandı. Başarılı: {success_count}, Hatalı: {failed_count}.",
                close_dialog=False,
                data={
                    "success_count": success_count,
                    "failed_count": failed_count,
                    "failed": failed,
                    "cargo_provider_code": provider_code,
                    "cargo_provider_name": provider_name,
                },
            )

        return Result.ok(
            f"Tüm seçili siparişlerin kargo firması güncellendi: {provider_name} ({success_count} sipariş).",
            close_dialog=False,
            data={
                "success_count": success_count,
                "failed_count": 0,
                "failed": [],
                "cargo_provider_code": provider_code,
                "cargo_provider_name": provider_name,
            },
        )
    except Exception as e:
        return Result.fail(map_error_to_message(e), error=e, close_dialog=False)


def load_ready_to_ship_orders() -> Result:
    try:
        result = get_latest_ready_to_ship_orders()
        if not result.success:
            return result

        orders = result.data.get("orders", []) or []
        return Result.ok("Siparişler yüklendi.", data={"records": orders})

    except Exception as e:
        return Result.fail(map_error_to_message(e), error=e)


def shipment_status_codes_for_load(status_code: str | None) -> list[str] | None:
    raw_code = str(status_code or "").strip()
    if _is_all_like_value(raw_code):
        return None
    label_code = _status_code_from_label(raw_code, TRENDYOL_STATUS_LABELS)
    normalized = _normalize_shipment_status_code(label_code or raw_code)
    if _is_all_like_value(normalized):
        return None
    if normalized == "Created":
        return ["Created", "ReadyToShip"]
    if normalized == "Cancelled":
        return ["Cancelled", "Unsupplied"]
    return [normalized]


def load_orders_for_source_filters(
    *,
    shipment_status_code: str | None = None,
    platform_code: str | None = None,
    date_from: date | None = None,
    date_to: date | None = None,
) -> Result:
    try:
        start_ms = end_ms = None
        if date_from is not None or date_to is not None:
            _resolved_from, _resolved_to, start_ms, end_ms = _resolve_returns_date_window(
                date_from=date_from,
                date_to=date_to,
            )
        result = get_latest_orders_filtered(
            shipment_statuses=shipment_status_codes_for_load(shipment_status_code),
            platform_codes=_platform_codes_for_load(platform_code),
            start_ms=start_ms,
            end_ms=end_ms,
        )
        if not result.success:
            return result

        orders = result.data.get("orders", []) or []
        return Result.ok("Siparişler yüklendi.", data={"records": orders})
    except Exception as e:
        return Result.fail(map_error_to_message(e), error=e)


def load_orders_for_status_filter(*, shipment_status_code: str | None = None) -> Result:
    return load_orders_for_source_filters(shipment_status_code=shipment_status_code)
    try:
        result = get_latest_orders(
            shipment_statuses=shipment_status_codes_for_load(shipment_status_code)
        )
        if not result.success:
            return result

        orders = result.data.get("orders", []) or []
        return Result.ok("Siparişler yüklendi.", data={"records": orders})

    except Exception as e:
        return Result.fail(map_error_to_message(e), error=e)


def load_orders_by_numbers(order_numbers: list[str]) -> Result:
    try:
        result = get_latest_orders(order_numbers=order_numbers)
        if not result.success:
            return result

        orders = result.data.get("orders", []) or []
        return Result.ok("Siparişler yüklendi.", data={"records": orders})

    except Exception as e:
        return Result.fail(map_error_to_message(e), error=e)


def refresh_cargo_filter(
    combo_box,
    orders: list,
    *,
    platform_code: str | None = None,
    preferred_cargo: str | None = None,
) -> Result:
    signal_blocker = None
    try:
        normalized_platform = canonical_platform_code(platform_code)
        if normalized_platform.casefold() in ("", "all", "tümü"):
            normalized_platform = ""
        signal_blocker = _signal_blocker(combo_box)
        if hasattr(combo_box, "clear"):
            combo_box.clear()

        try:
            combo_box.addItem(orders_ui.ORDERS_FILTER_ALL_TEXT, orders_ui.ORDERS_FILTER_ALL_VALUE)
        except Exception:
            combo_box.addItem(orders_ui.ORDERS_FILTER_ALL_TEXT)

        orders_for_filter = list(orders or [])
        if normalized_platform:
            orders_for_filter = [
                order
                for order in orders_for_filter
                if resolve_order_platform_code(order) == normalized_platform
            ]

        cargos = extract_cargo_names(orders_for_filter)
        if hasattr(combo_box, "addItems"):
            combo_box.addItems(cargos)
        else:
            for cargo in cargos:
                combo_box.addItem(cargo)

        target_cargo = str(preferred_cargo or "").strip() or orders_ui.ORDERS_FILTER_ALL_TEXT
        if target_cargo != orders_ui.ORDERS_FILTER_ALL_TEXT and target_cargo not in cargos:
            target_cargo = orders_ui.ORDERS_FILTER_ALL_TEXT
        if hasattr(combo_box, "setCurrentText"):
            combo_box.setCurrentText(target_cargo)
        elif hasattr(combo_box, "setCurrentIndex"):
            combo_box.setCurrentIndex(0)
        return Result.ok(f"{len(cargos)} kargo firması yüklendi.", close_dialog=False)

    except Exception as e:
        return Result.fail(map_error_to_message(e), error=e)
    finally:
        del signal_blocker


def refresh_order_status_filter(
    combo_box,
    orders: list,
    *,
    default_status_code: str | None = "Created",
    default_status_label: str | None = "Kargoya Hazır",
    include_known_statuses: bool = False,
    known_status_codes: list[str] | tuple[str, ...] | None = None,
) -> Result:
    signal_blocker = None
    try:
        signal_blocker = _signal_blocker(combo_box)
        if hasattr(combo_box, "clear"):
            combo_box.clear()

        try:
            combo_box.addItem(orders_ui.ORDERS_FILTER_ALL_TEXT, orders_ui.ORDERS_FILTER_ALL_VALUE)
        except Exception:
            combo_box.addItem(orders_ui.ORDERS_FILTER_ALL_TEXT)

        known_status_set = {str(row or "").strip() for row in (known_status_codes or [])}
        if "Created" in known_status_set and "WaitingInAction" in known_status_set:
            try:
                combo_box.addItem(RETURN_ACTIVE_STATUS_FILTER_LABEL, data=RETURN_ACTIVE_STATUS_FILTER_CODE)
            except TypeError:
                try:
                    combo_box.addItem(RETURN_ACTIVE_STATUS_FILTER_LABEL, RETURN_ACTIVE_STATUS_FILTER_CODE)
                except Exception:
                    combo_box.addItem(RETURN_ACTIVE_STATUS_FILTER_LABEL)
            except Exception:
                combo_box.addItem(RETURN_ACTIVE_STATUS_FILTER_LABEL)

        if known_status_codes:
            statuses = []
            seen_codes: set[str] = set()
            for status_code in known_status_codes:
                code = _normalize_shipment_status_code(status_code)
                key = code.casefold()
                if not code or key in seen_codes:
                    continue
                seen_codes.add(key)
                statuses.append((format_order_status(code), code))
        elif include_known_statuses:
            statuses = []
            seen_codes: set[str] = set()
            for status_code in TRENDYOL_STATUS_LIST:
                code = _normalize_shipment_status_code(status_code)
                key = code.casefold()
                if not code or key in seen_codes:
                    continue
                seen_codes.add(key)
                statuses.append((format_order_status(code), code))
        else:
            statuses = extract_order_statuses(orders)

        for label, code in statuses:
            try:
                combo_box.addItem(label, data=code)
            except TypeError:
                try:
                    combo_box.addItem(label, code)
                except Exception:
                    combo_box.addItem(label)
            except Exception:
                combo_box.addItem(label)

        default_idx = -1
        code_text = str(default_status_code or "").strip()
        if code_text and hasattr(combo_box, "findData"):
            try:
                default_idx = int(combo_box.findData(code_text))
            except Exception:
                default_idx = -1

        label_text = str(default_status_label or "").strip()
        if default_idx < 0 and label_text and hasattr(combo_box, "findText"):
            try:
                default_idx = int(combo_box.findText(label_text))
            except Exception:
                default_idx = -1

        if default_idx >= 0 and hasattr(combo_box, "setCurrentIndex"):
            combo_box.setCurrentIndex(default_idx)
        elif hasattr(combo_box, "setCurrentText"):
            combo_box.setCurrentText(orders_ui.ORDERS_FILTER_ALL_TEXT)
        elif hasattr(combo_box, "setCurrentIndex"):
            combo_box.setCurrentIndex(0)
        return Result.ok(f"{len(statuses)} sipariş durumu yüklendi.", close_dialog=False)

    except Exception as e:
        return Result.fail(map_error_to_message(e), error=e)
    finally:
        del signal_blocker


# ============================================================
# 3) Trendyol non-final refresh (async)
# ============================================================

async def _refresh_nonfinal_orders_async(
    order_numbers: list[str],
    comp_api_account_list: list,
    progress_callback: Optional[Callable[[int, int], None]] = None,
) -> Result:
    """
    Non-final siparişleri orderNumber üzerinden Trendyol'dan çekip normalize eder.

    Dönüş:
      Result.data = {"order_data_list":[...], "order_item_list":[...]}
    """
    try:
        if not order_numbers:
            return Result.ok(
                "Güncellenecek non-final sipariş bulunamadı.",
                close_dialog=False,
                data={"order_data_list": [], "order_item_list": []},
            )

        all_orders: list[dict] = []
        all_items: list[dict] = []

        total_steps = max(len(order_numbers) * len(comp_api_account_list), 1)
        current_step = 0

        for comp_api_account in comp_api_account_list:
            api_account_id = comp_api_account[0]
            api = _trendyol_api_for_account_payload(comp_api_account)

            for order_no in order_numbers:
                res = await api.get_order_by_number(order_no)

                # her döngüde progress ilerlet
                try:
                    if res and isinstance(res, Result) and res.success:
                        content = (res.data or {}).get("content", []) or []
                        for raw_order in content:
                            norm_orders, norm_items = normalize_order_data(raw_order, api_account_id)
                            all_orders.extend(norm_orders)
                            all_items.extend(norm_items)
                finally:
                    current_step += 1
                    if progress_callback:
                        progress_callback(current_step, total_steps)

        return Result.ok(
            f"{len(all_orders)} adet order_data, {len(all_items)} adet order_item normalize edildi.",
            close_dialog=False,
            data={"order_data_list": all_orders, "order_item_list": all_items},
        )

    except Exception as e:
        return Result.fail(map_error_to_message(e), error=e, close_dialog=False)


# ============================================================
# 3.1) DB Save Runner (NEW)
# ============================================================

def _start_db_save_worker(
    parent_widget,
    payload: dict,
    on_finished: Callable[[dict], None],
) -> None:
    """
    DB save işini ayrı process'te çalıştırır.
    - Dev: python Orders/processors/trendyol/db_save_worker.py
    - Frozen: OrderScout.exe --worker Orders/processors/trendyol/db_save_worker.py
    """
    target = ProcessTarget(identifier="Orders/processors/trendyol/db_save_worker.py")

    runner = GenericProcessRunner(
        target=target,
        initial_payload=payload,
        parent=parent_widget,
    )

    runner.finished.connect(on_finished)
    # stdout/stderr istersen UI log'a bağlarsın:
    # runner.stderr_text.connect(lambda t: print("[db_save STDERR]", t))
    runner.start()
    runner.close_stdin()

    # parent'ta referans tut (GC yemesin)
    return runner


# ============================================================
# 3.2) Progress helpers
# ============================================================

def update_progress(view_instance, current: int, total: int) -> Result:
    try:
        total = max(int(total), 1)
        current = max(int(current), 0)

        percent = int(current / total * 100)
        percent = max(0, min(100, percent))

        if hasattr(view_instance, "setProgress"):
            view_instance.setProgress(percent)
        elif hasattr(view_instance, "fetch_button") and hasattr(view_instance.fetch_button, "setProgress"):
            view_instance.fetch_button.setProgress(percent)

        return Result.ok(f"Progress {percent}% olarak güncellendi.", close_dialog=False)

    except Exception as e:
        return Result.fail(map_error_to_message(e), error=e, close_dialog=False)


def set_progress_stage_text(progress_target, text: str | None) -> None:
    try:
        if hasattr(progress_target, "setRunningText"):
            progress_target.setRunningText(str(text or "").strip())
    except Exception:
        pass


def _progress_within_range(
    progress_target,
    c: int,
    t: int,
    *,
    start_percent: int,
    end_percent: int,
) -> None:
    if _is_progress_locked(progress_target):
        return
    try:
        t = max(int(t), 1)
        start_percent = int(start_percent)
        end_percent = int(end_percent)
        if end_percent < start_percent:
            end_percent = start_percent
        span = end_percent - start_percent
        percent = start_percent + int(c * span / t)
        percent = max(start_percent, min(end_percent, percent))
    except Exception:
        percent = int(start_percent)
    update_progress(progress_target, percent, 100)


def _progress_main_to_59(progress_target, c: int, t: int) -> None:
    _progress_within_range(
        progress_target,
        c,
        t,
        start_percent=0,
        end_percent=59,
    )


def _progress_returns_60_to_79(progress_target, c: int, t: int) -> None:
    _progress_within_range(
        progress_target,
        c,
        t,
        start_percent=60,
        end_percent=79,
    )


def _progress_nonfinal_80_to_99(progress_target, c: int, t: int) -> None:
    _progress_within_range(
        progress_target,
        c,
        t,
        start_percent=80,
        end_percent=99,
    )


def _emit_orders_persisted_signal(
    *,
    source: str,
    payload: dict,
    db_result: dict,
) -> dict | None:
    """
    Emits a cross-domain signal only when new order-item rows are inserted.
    """
    try:
        counts = (db_result or {}).get("counts") or {}
        inserted_items = int(counts.get("items_inserted", 0) or 0)
    except Exception:
        inserted_items = 0

    if inserted_items <= 0:
        return None

    event_payload = {
        "source": str(source or "").strip() or "main",
        "order_item_list": (payload or {}).get("order_item_list", []) or [],
        "order_data_list": (payload or {}).get("order_data_list", []) or [],
        "db_result": db_result or {},
    }
    try:
        order_signals.orders_persisted.emit(event_payload)
    except Exception:
        pass
    return event_payload


def _start_order_stock_sync(
    parent_widget,
    *,
    source: str,
    payload: dict,
    db_result: dict,
) -> None:
    event_payload = _emit_orders_persisted_signal(
        source=source,
        payload=payload,
        db_result=db_result,
    )
    if not isinstance(event_payload, dict):
        return

    order_items = list(event_payload.get("order_item_list", []) or [])
    if not order_items:
        return

    from Stock.views.actions import action_auto_deduct_order_items

    worker_store = getattr(parent_widget, "_order_stock_sync_workers", None)
    if not isinstance(worker_store, list):
        worker_store = []
        setattr(parent_widget, "_order_stock_sync_workers", worker_store)

    worker = SyncWorker(
        action_auto_deduct_order_items,
        order_items,
        parent=parent_widget,
        created_by=f"orders:{str(source or '').strip() or 'auto'}",
    )
    worker_store.append(worker)

    def _cleanup() -> None:
        try:
            worker_store.remove(worker)
        except ValueError:
            pass

    def _handle_result(result: Result) -> None:
        if getattr(result, "success", False):
            return
        message = str(getattr(result, "message", "") or "").strip() or "Siparis stok dusumu basarisiz."
        try:
            if hasattr(parent_widget, "status_label") and parent_widget.status_label is not None:
                parent_widget.status_label.setText(message)
        except Exception:
            pass
        try:
            MessageHandler.show(parent_widget, Result.fail(message, close_dialog=False), only_errors=True)
        except Exception:
            pass

    worker.result_ready.connect(_handle_result)
    worker.finished.connect(_cleanup)
    worker.start()


def _finalize_pipeline(
    progress_target,
    comp_api_account_list: list,
    changed_state: bool,
) -> None:
    """
    Pipeline sonu: last_used güncelle, orders_changed emit (gerekirse), progress %100, buton reset+enable
    """
    try:
        update_last_used_at_for_accounts(comp_api_account_list, safety_minutes=0)
    except Exception:
        pass

    if changed_state:
        try:
            order_signals.orders_changed.emit()
        except Exception:
            pass

    set_progress_stage_text(progress_target, "")
    _set_progress_lock(progress_target, True)
    update_progress(progress_target, 100, 100)

    try:
        progress_target.setEnabled(True)
    except Exception:
        pass


# ============================================================
# 3.3) Main pipeline starter
# ============================================================
def get_orders_from_companies(parent_widget, company_list_widget, progress_target) -> Result:
    """
    OrdersTab.get_orders() burayı çağırır.
    """
    try:
        # 1) Seçilen şirketler
        res_sel = collect_selected_companies(company_list_widget)
        if not res_sel.success:
            return res_sel
        selected_company_pks = res_sel.data["selected_company_pks"]

        # 2) API credential'ları
        res_creds = get_company_by_id(selected_company_pks)
        if not res_creds.success:
            return res_creds

        comp_api_account_list = res_creds.data.get("accounts", []) or []
        if not comp_api_account_list:
            return Result.fail("Seçili şirketler için API bilgisi bulunamadı.", close_dialog=False)
        trendyol_accounts = [
            account
            for account in (res_creds.data.get("records", []) or [])
            if str(getattr(account, "platform", "") or "").strip().upper() == "TRENDYOL"
        ]

        state = {
            "changed": False,
            "counts": {
                "headers_inserted": 0,
                "data_inserted": 0,
                "items_inserted": 0,
                "claims_inserted": 0,
            },
            "warnings": [],
            "sync_diagnostics": {},
        }
        _set_progress_lock(progress_target, False)

        # 3) Zaman aralığı (start/end epoch ms)
        now_ms = _now_epoch_ms()
        last_used_ms_list = _extract_last_used_ms_list(comp_api_account_list)
        final_ep_time, start_ep_time = _compute_main_fetch_window(
            now_ms=now_ms,
            last_used_ms_list=last_used_ms_list,
            initial_lookback_hours=INITIAL_FETCH_LOOKBACK_HOURS,
        )
        # Manual "Veri Çek" must converge to the current marketplace state.
        # Throttling this step lets old Created/ReadyToShip snapshots survive
        # until the interval expires, so the manual pipeline always verifies
        # local non-final orders by orderNumber.
        run_nonfinal_refresh = True

        def _merge_counts(counts: dict | None) -> None:
            for key, value in (counts or {}).items():
                try:
                    state["counts"][key] = int(state["counts"].get(key, 0) or 0) + int(value or 0)
                except Exception:
                    continue

        def _append_warning(message: str) -> None:
            text = str(message or "").strip()
            if not text:
                return
            if text not in state["warnings"]:
                state["warnings"].append(text)

        def _set_stage_feedback(button_text: str, status_text: str | None = None) -> None:
            set_progress_stage_text(progress_target, button_text)
            if status_text is not None and hasattr(parent_widget, "status_label") and parent_widget.status_label is not None:
                try:
                    parent_widget.status_label.setText(str(status_text))
                except Exception:
                    pass

        def _complete_pipeline() -> None:
            parent_widget.on_orders_fetched(
                Result.ok(
                    "Sipariş ve iade senkronizasyonu tamamlandı.",
                    close_dialog=False,
                    data={
                        "changed": state["changed"],
                        "counts": dict(state["counts"]),
                        "warnings": list(state["warnings"]),
                        "sync_diagnostics": dict(state["sync_diagnostics"]),
                        "source_rows_complete": state["sync_diagnostics"].get("source_rows_complete", True),
                    },
                )
            )
            _finalize_pipeline(progress_target, comp_api_account_list, state["changed"])

        def _start_nonfinal_stage() -> None:
            if not run_nonfinal_refresh:
                _complete_pipeline()
                return

            res_nonfinal = get_nonfinal_order_numbers()
            if not res_nonfinal.success:
                _append_warning(res_nonfinal.message or "Non-final sipariş listesi alınamadı.")
                _complete_pipeline()
                return

            order_numbers = (res_nonfinal.data or {}).get("order_numbers", []) or []
            if not order_numbers:
                _mark_nonfinal_refresh_completed(now_ms=_now_epoch_ms())
                _complete_pipeline()
                return

            _set_stage_feedback(
                "Siparişler güncelleniyor...",
                "⏳ Non-final siparişler güncelleniyor...",
            )
            update_progress(progress_target, 80, 100)
            parent_widget.bg_api_worker = AsyncWorker(
                _refresh_nonfinal_orders_async,
                order_numbers,
                comp_api_account_list,
                parent=parent_widget,
            )

            parent_widget.bg_api_worker.progress_changed.connect(
                lambda c, t: _progress_nonfinal_80_to_99(progress_target, c, t)
            )

            def handle_bg_api(bg_res: Result):
                if not bg_res or not isinstance(bg_res, Result) or not bg_res.success:
                    _append_warning(
                        getattr(bg_res, "message", "") or "Non-final sipariş yenileme adımı başarısız oldu."
                    )
                    _complete_pipeline()
                    return

                payload_bg = {
                    "order_data_list": (bg_res.data or {}).get("order_data_list", []) or [],
                    "order_item_list": (bg_res.data or {}).get("order_item_list", []) or [],
                }

                def on_db_bg_finished(bg_db_payload: dict):
                    if bg_db_payload.get("success"):
                        bg_data_dict = bg_db_payload.get("data") or {}
                        if bool(bg_data_dict.get("changed")):
                            state["changed"] = True
                        _merge_counts(bg_data_dict.get("counts") or {})
                        _start_order_stock_sync(
                            parent_widget,
                            source="nonfinal_refresh",
                            payload=payload_bg,
                            db_result=bg_data_dict,
                        )
                        _mark_nonfinal_refresh_completed(now_ms=_now_epoch_ms())
                    else:
                        _append_warning(bg_db_payload.get("message", "Non-final siparişler DB'ye yazılamadı."))

                    _complete_pipeline()

                parent_widget.bg_db_runner = _start_db_save_worker(
                    parent_widget,
                    payload_bg,
                    on_db_bg_finished,
                )

            parent_widget.bg_api_worker.result_ready.connect(handle_bg_api)
            parent_widget.bg_api_worker.start()

        def _start_returns_stage() -> None:
            if not trendyol_accounts:
                _start_nonfinal_stage()
                return

            def _start_active_returns_refresh_stage() -> None:
                _set_stage_feedback(
                    "Açık iadeler doğrulanıyor...",
                    "Açık iade claim durumları doğrulanıyor...",
                )
                update_progress(progress_target, 76, 100)
                parent_widget.active_returns_refresh_worker = AsyncWorker(
                    _refresh_active_return_claims_from_api_async,
                    trendyol_accounts,
                    parent=parent_widget,
                )

                parent_widget.active_returns_refresh_worker.progress_changed.connect(
                    lambda c, t: _progress_returns_60_to_79(progress_target, c, t)
                )

                def handle_active_returns_refresh(refresh_res: Result):
                    if not refresh_res or not isinstance(refresh_res, Result) or not refresh_res.success:
                        _append_warning(
                            getattr(refresh_res, "message", "") or "Açık iade claim yenileme adımı başarısız oldu."
                        )
                        _start_nonfinal_stage()
                        return

                    refresh_data = refresh_res.data or {}
                    if bool(refresh_data.get("changed")):
                        state["changed"] = True
                    _merge_counts(refresh_data.get("counts") or {})
                    refresh_errors = refresh_data.get("errors") or []
                    if refresh_errors:
                        _append_warning(f"{len(refresh_errors)} açık iade claim yenileme uyarısı oluştu.")
                    _start_nonfinal_stage()

                parent_widget.active_returns_refresh_worker.result_ready.connect(handle_active_returns_refresh)
                parent_widget.active_returns_refresh_worker.start()

            _set_stage_feedback(
                "İadeler çekiliyor...",
                "⏳ İadeler çekiliyor...",
            )
            update_progress(progress_target, 60, 100)
            parent_widget.returns_api_worker = AsyncWorker(
                _fetch_returns_from_active_companies_async,
                trendyol_accounts,
                parent=parent_widget,
                kwargs={
                    "start_ms": start_ep_time,
                    "end_ms": final_ep_time,
                    "size": 50,
                },
            )

            parent_widget.returns_api_worker.progress_changed.connect(
                lambda c, t: _progress_returns_60_to_79(progress_target, c, t)
            )

            def handle_returns_api(returns_res: Result):
                if not returns_res or not isinstance(returns_res, Result) or not returns_res.success:
                    _append_warning(
                        getattr(returns_res, "message", "") or "İadeler alınamadı."
                    )
                    _start_active_returns_refresh_stage()
                    return

                fetch_errors = (returns_res.data or {}).get("errors", []) or []
                if fetch_errors:
                    _append_warning(f"{len(fetch_errors)} mağazada iade sorgu hatası oluştu.")

                payload_returns = {
                    "return_claim_list": (returns_res.data or {}).get("return_claim_list", []) or [],
                }

                if not payload_returns["return_claim_list"]:
                    _start_active_returns_refresh_stage()
                    return

                def on_db_returns_finished(returns_db_payload: dict):
                    if returns_db_payload.get("success"):
                        returns_data_dict = returns_db_payload.get("data") or {}
                        if bool(returns_data_dict.get("changed")):
                            state["changed"] = True
                        _merge_counts(returns_data_dict.get("counts") or {})
                    else:
                        _append_warning(returns_db_payload.get("message", "İadeler DB'ye yazılamadı."))

                    _start_nonfinal_stage()

                parent_widget.returns_db_runner = _start_db_save_worker(
                    parent_widget,
                    payload_returns,
                    on_db_returns_finished,
                )

            parent_widget.returns_api_worker.result_ready.connect(handle_returns_api)
            parent_widget.returns_api_worker.start()

        # 4) Ana fetch worker
        parent_widget.api_worker = AsyncWorker(
            fetch_orders_all,
            TRENDYOL_API_STATUS_LIST,
            final_ep_time,
            start_ep_time,
            comp_api_account_list,
            parent=parent_widget,
        )

        parent_widget.api_worker.progress_changed.connect(
            lambda c, t: _progress_main_to_59(progress_target, c, t)
        )
        _set_stage_feedback(
            "Siparişler çekiliyor...",
            "⏳ Siparişler çekiliyor...",
        )

        def handle_api_result(main_res: Result):
            if not main_res.success:
                _set_progress_lock(progress_target, True)
                parent_widget.on_orders_failed(main_res, progress_target)
                update_progress(progress_target, 0, 100)
                return

            main_sync_diagnostics = dict((main_res.data or {}).get("sync_diagnostics") or {})
            state["sync_diagnostics"] = main_sync_diagnostics
            if main_sync_diagnostics.get("source_rows_complete") is False:
                _append_warning(
                    _orders_pipeline("order_sync_diagnostics_warning")(main_sync_diagnostics)
                )

            payload_main = {
                "order_data_list": (main_res.data or {}).get("order_data_list", []) or [],
                "order_item_list": (main_res.data or {}).get("order_item_list", []) or [],
                "sync_diagnostics": main_sync_diagnostics,
            }

            # 5) DB save (process)
            def on_db_main_finished(db_payload: dict):
                if not db_payload.get("success"):
                    _set_progress_lock(progress_target, True)
                    parent_widget.on_orders_failed(
                        Result.fail(db_payload.get("message", "DB kayıt hatası.")),
                        progress_target,
                    )
                    update_progress(progress_target, 0, 100)
                    return

                data_dict = db_payload.get("data") or {}
                if bool(data_dict.get("changed")):
                    state["changed"] = True
                _merge_counts(data_dict.get("counts") or {})
                if data_dict.get("sync_diagnostics"):
                    state["sync_diagnostics"] = dict(data_dict.get("sync_diagnostics") or {})

                _start_order_stock_sync(
                    parent_widget,
                    source="main",
                    payload=payload_main,
                    db_result=data_dict,
                )
                _start_returns_stage()

            parent_widget.db_runner = _start_db_save_worker(
                parent_widget,
                payload_main,
                on_db_main_finished,
            )

        parent_widget.api_worker.result_ready.connect(handle_api_result)
        parent_widget.api_worker.start()

        return Result.ok("Worker başlatıldı.", close_dialog=False)

    except Exception as e:
        msg = map_error_to_message(e)
        parent_widget.on_orders_failed(Result.fail(msg, error=e), progress_target)
        return Result.fail(msg, error=e, close_dialog=False)


# ============================================================
# 4) Filtering
# ============================================================

def get_order_date(order) -> date | None:
    resolved = _get_order_record_date(order)
    if resolved:
        try:
            setattr(order, _ORDER_DATE_CACHE_ATTR, resolved)
        except Exception:
            pass
        return resolved
    cached = getattr(order, _ORDER_DATE_CACHE_ATTR, None)
    if isinstance(cached, date):
        return cached
    for attr in ("shipmentDate", "orderDate", "createdDate"):
        val = getattr(order, attr, None)
        coerced = coerce_to_date(val)
        if coerced:
            try:
                setattr(order, _ORDER_DATE_CACHE_ATTR, coerced)
            except Exception:
                pass
            return coerced
    return None


def resolve_order_platform_code(order) -> str:
    resolved = _resolve_order_record_platform_code(order)
    if resolved:
        try:
            setattr(order, _ORDER_PLATFORM_CODE_CACHE_ATTR, resolved)
        except Exception:
            pass
        return resolved
    cached = str(getattr(order, _ORDER_PLATFORM_CODE_CACHE_ATTR, "") or "").strip()
    if cached:
        return cached
    direct_code = canonical_platform_code(getattr(order, "platform", None))
    if direct_code:
        try:
            setattr(order, _ORDER_PLATFORM_CODE_CACHE_ATTR, direct_code)
        except Exception:
            pass
        return direct_code
    api_account = getattr(order, "api_account", None)
    resolved = canonical_platform_code(getattr(api_account, "platform", None))
    if resolved:
        try:
            setattr(order, _ORDER_PLATFORM_CODE_CACHE_ATTR, resolved)
        except Exception:
            pass
    return resolved


def _order_platform_code(order) -> str:
    return resolve_order_platform_code(order)


def _loaded_order_items(order) -> list:
    if order is None:
        return []
    try:
        loaded = vars(order).get("items", None)
    except Exception:
        loaded = None
    return list(loaded or []) if isinstance(loaded, list) else []


def filter_orders(orders: list, filters: dict) -> Result:
    runtime_filters = dict(filters or {})
    runtime_filters.setdefault("now_ms", _now_epoch_ms())
    return filter_order_records_in_memory(orders, runtime_filters)


def start_filter_worker(parent_widget, list_widget, filters: dict, *, apply_on_result: bool = True) -> SyncWorker:
    orders_snapshot = list(getattr(list_widget, "orders", []) or [])
    worker = SyncWorker(filter_orders, orders_snapshot, dict(filters or {}))

    def handle_result(result: Result):
        if not result.success:
            MessageHandler.show(parent_widget, result, only_errors=True)
            parent_widget.selected_count_label.setText(orders_ui.ORDERS_FILTER_FAILED_TEXT)
            return

        filtered = result.data.get("filtered", []) or []
        list_widget.apply_filter_result(filtered)
        list_widget.filtered_orders = filtered
        parent_widget._update_label()
        parent_widget.selected_count_label.setText(f"Filtre tamamlandı. (Kalan: {len(filtered)})")

    if apply_on_result:
        worker.result_ready.connect(handle_result)
    return worker


