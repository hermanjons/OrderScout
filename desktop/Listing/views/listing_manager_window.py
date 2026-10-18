# Listing/views/views.py
from __future__ import annotations

from typing import Any, Optional, Callable
from copy import deepcopy

from pathlib import Path
import os
import hashlib
import time
from collections import OrderedDict
from concurrent.futures import ThreadPoolExecutor, as_completed
from urllib.parse import urlsplit, urlunsplit

from PyQt6.QtCore import Qt, QTimer, QAbstractListModel, QModelIndex, QRect, QEvent, QSize, QPoint, QPointF, QRegularExpression, QSignalBlocker, pyqtSignal
from PyQt6.QtGui import QPixmap, QPainter, QColor, QPen, QFont, QConicalGradient, QBrush, QIcon, QIntValidator, QRegularExpressionValidator
from PyQt6.QtWidgets import (
    QWidget,
    QVBoxLayout,
    QHBoxLayout,
    QLabel,
    QFrame,
    QLineEdit,
    QListView,
    QStyledItemDelegate,
    QStyleOptionViewItem,
    QStyle,
    QPushButton,
    QComboBox,
    QDialog,
    QSplitter,
    QGraphicsDropShadowEffect,
    QSizePolicy,
    QGridLayout,
    QScrollArea,
    QFileDialog,
    QLayout,
)

from Core.views import (
    CircularProgressButton,
    CircularPercentageInput,
    PackageButton,
    compose_record_workspace_style,
    RecordWorkspaceBehaviorAdapter,
    RecordWorkspaceBehaviorController,
    RecordWorkspaceListDelegate,
    RecordWorkspaceListModel,
    RecordWorkspaceScaffold,
    SwitchSelectWidget,
    NumberStepper,
    BadgeButton,
    SwitchButton,
    ActionPulseButton,
    CurrencyAmountInput,
    NumericPulseLineEdit,
    LabeledActionInput,
    AppMessageCenter,
    SelectClearToggleButton,
    ResettableSplitter,
    SmoothHeightFrame,
)
from Listing.constants import ui_text
from Core.views.tab_widgets import DashboardHeroCard, SplitPanelSidebar
from Core.threads.async_worker import AsyncWorker
from Core.threads.sync_worker import SyncWorker
from Core.process.process_runner import GenericProcessRunner, ProcessTarget
from Feedback.processors.pipeline import MessageHandler, Result, map_error_to_message
from Account.views.company_list_widget import CompanyListWidget
from Account.signals.signals import account_signals
from Marketplaces.signals.signals import marketplace_signals
from Logistics.signals.signals import logistics_signals

from Listing.signals.signals import listing_signals
from Stock.signals.signals import stock_signals
from settings import ASSETS_DIR
from Listing.constants.paths import LISTING_IMAGE_CACHE_DIR
from Listing.constants.platform_constants import canonical_platform_code, platform_label_for_code
from Listing.constants.ui_constants import (
    COST_AMOUNT_PLACEHOLDER_BY_TYPE,
    COST_MATCH_CARGO_PROVIDER_PLACEHOLDER,
    COST_MATCH_CARGO_PROVIDER_TOOLTIP,
    COST_MATCH_DB_FIELD_TOOLTIP,
    COST_MATCH_MODES,
    COST_MATCH_PERCENT_DIAMETER,
    COST_MATCH_ROW_HEIGHT,
    COST_MATCH_SOURCE_TOOLTIP,
    COST_SOURCE_FIELD_PLACEHOLDER,
    DB_FIELD_EXCLUDED_NAMES,
    LISTING_BADGE_BTN_H,
    LISTING_BADGE_BTN_W,
    LISTING_BADGE_MIN_H,
    LISTING_BADGE_MIN_W,
    LISTING_BADGE_STACK_GAP,
    LISTING_COMPACT_BADGE_MIN_W,
    LISTING_COMPACT_PRICE_MIN_W,
    LISTING_COMPACT_QTY_MIN_W,
    LISTING_HIDE_BADGES_CLUSTER_W,
    LISTING_HOOK_BADGE_GAP,
    LISTING_HOOK_H,
    LISTING_HOOK_W,
    LISTING_IMAGE_DOWNLOAD_TIMEOUT_SECONDS,
    LISTING_MIN_RIGHT_SCALE,
    LISTING_MIN_TEXT_W,
    LISTING_LIST_LOADING_TEXT,
    LISTING_PRICE_H,
    LISTING_PRICE_W,
    LISTING_PROFIT_DETAIL_LOADING_TEXT,
    LISTING_PROFIT_H,
    LISTING_PROFIT_W,
    LISTING_QTY_H,
    LISTING_QTY_W,
    LISTING_STYLE,
    LISTING_TEXT_CLUSTER_GAP,
    LISTING_BADGE_VISIBLE_MIN_TEXT_W,
    PENDING_UI_MS,
    PREFERRED_DB_FIELD_LABELS,
    VERIFY_RETRY_MS,
)


def _listing_action(name: str):
    from Listing.views import actions

    return getattr(actions, name)


def _listing_pipeline(name: str):
    from Listing.processors import pipeline

    return getattr(pipeline, name)


def _stock_action(name: str):
    from Stock.views import actions

    return getattr(actions, name)


def _finance_cost_source_definitions() -> dict[str, Any]:
    from Finance.constants.domain_constants import FINANCE_COST_SOURCE_DEFINITIONS

    return FINANCE_COST_SOURCE_DEFINITIONS


def cost_source_type_label(*args, **kwargs):
    from Finance.processors.pipeline import cost_source_type_label as _cost_source_type_label

    return _cost_source_type_label(*args, **kwargs)


def normalize_cost_source_type(*args, **kwargs):
    from Finance.processors.pipeline import normalize_cost_source_type as _normalize_cost_source_type

    return _normalize_cost_source_type(*args, **kwargs)


def collect_selected_companies(*args, **kwargs):
    from Account.views.actions import collect_selected_companies as _collect_selected_companies

    return _collect_selected_companies(*args, **kwargs)


async def action_fetch_listings(*args, **kwargs):
    return await _listing_action("action_fetch_listings")(*args, **kwargs)


def action_get_all_companies(*args, **kwargs):
    return _listing_action("action_get_all_companies")(*args, **kwargs)


def action_get_company_credentials(*args, **kwargs):
    return _listing_action("action_get_company_credentials")(*args, **kwargs)


def action_get_listings_from_db(*args, **kwargs):
    return _listing_action("action_get_listings_from_db")(*args, **kwargs)


def action_get_listing_brands_from_db(*args, **kwargs):
    return _listing_action("action_get_listing_brands_from_db")(*args, **kwargs)


def action_load_logistics_cargo_providers(*args, **kwargs):
    return _listing_action("action_load_logistics_cargo_providers")(*args, **kwargs)


def action_load_logistics_cargo_tariffs(*args, **kwargs):
    return _listing_action("action_load_logistics_cargo_tariffs")(*args, **kwargs)


def action_load_finance_cost_items(*args, **kwargs):
    return _listing_action("action_load_finance_cost_items")(*args, **kwargs)


def action_set_listing_cost_items(*args, **kwargs):
    return _listing_action("action_set_listing_cost_items")(*args, **kwargs)


def action_set_listing_stock_items(*args, **kwargs):
    return _listing_action("action_set_listing_stock_items")(*args, **kwargs)


def action_sync_listing_cost_snapshots_for_item(*args, **kwargs):
    return _listing_action("action_sync_listing_cost_snapshots_for_item")(*args, **kwargs)


def action_update_listing_price_inventory_cache(*args, **kwargs):
    return _listing_action("action_update_listing_price_inventory_cache")(*args, **kwargs)


async def action_update_price_inventory_for_account(*args, **kwargs):
    return await _listing_action("action_update_price_inventory_for_account")(*args, **kwargs)


async def action_check_batch_request_for_account(*args, **kwargs):
    return await _listing_action("action_check_batch_request_for_account")(*args, **kwargs)


def estimate_listing_profit(*args, **kwargs):
    return _listing_pipeline("estimate_listing_profit")(*args, **kwargs)


def normalize_listing_cost_binding(*args, **kwargs):
    return _listing_pipeline("normalize_listing_cost_binding")(*args, **kwargs)


def action_list_items(*args, **kwargs):
    return _stock_action("action_list_items")(*args, **kwargs)


def action_list_item_balances(*args, **kwargs):
    return _stock_action("action_list_item_balances")(*args, **kwargs)


from PyQt6.QtCore import QThread, pyqtSignal


def _safe_str(val: Any) -> str:
    return str(val or "").strip()


def _listing_currency_code(item: Any, default: str = "TRY") -> str:
    if isinstance(item, dict):
        for key in ("currencyCode", "currency_code", "currency", "currencyType", "priceCurrency"):
            value = _safe_str(item.get(key))
            if value:
                return value.upper()
    fallback = _safe_str(default).upper()
    return fallback or "TRY"


def _listing_currency_label(code: Any) -> str:
    normalized = _safe_str(code).upper()
    if normalized in {"TRY", "TL", "?"}:
        return "TL"
    return normalized or "TL"


def _product_title(item: dict) -> str:
    return _safe_str(
        item.get("title")
        or item.get("name")
        or item.get("productName")
        or item.get("merchantSku")
        or item.get("stockCode")
        or item.get("barcode")
        or item.get("productId")
    )


def _product_id_text(item: dict[str, Any] | None) -> str:
    if not isinstance(item, dict):
        return ""
    return _safe_str(
        item.get("productId")
        or item.get("product_id")
        or item.get("platformListingId")
        or item.get("platform_listing_id")
        or item.get("productContentId")
        or item.get("product_content_id")
        or item.get("id")
    )


def _stock_item_identity(row: Any) -> int | None:
    if isinstance(row, dict):
        value = row.get("id")
        if value in (None, ""):
            value = row.get("pk")
    else:
        value = getattr(row, "id", None)
        if value in (None, ""):
            value = getattr(row, "pk", None)
    try:
        parsed = int(value)
    except Exception:
        return None
    return parsed if parsed > 0 else None


def _compose_linked_stock_label(
    stock_items: Any,
    *,
    stock_item_label_map: dict[int, str] | None,
    sale_quantity: float = 1.0,
    fallback: str = "",
) -> str:
    label_map = stock_item_label_map or {}
    labels: list[str] = []
    for stock_item in stock_items or ():
        if not isinstance(stock_item, dict):
            continue
        try:
            stock_item_id = int(stock_item.get("stock_item_id"))
        except Exception:
            continue
        label = _safe_str(label_map.get(stock_item_id))
        if not label:
            continue
        try:
            qty_value = float(stock_item.get("qty") or 0) * float(sale_quantity)
        except Exception:
            qty_value = 0.0
        if qty_value > 0:
            labels.append(f"{label} x{qty_value:.1f}")
        else:
            labels.append(label)
    return ", ".join(labels) or _safe_str(fallback)


def _normalized_account_platform(value: Any) -> str:
    return canonical_platform_code(value)


def _listing_badge_count_from_save_counts(counts: Any) -> int:
    if not isinstance(counts, dict):
        return 0
    try:
        inserted = max(0, int(counts.get("inserted") or 0))
    except Exception:
        inserted = 0
    try:
        updated = max(0, int(counts.get("updated") or 0))
    except Exception:
        updated = 0
    try:
        deleted = max(0, int(counts.get("deleted") or 0))
    except Exception:
        deleted = 0
    return inserted + updated + deleted


def _platform_filter_options_from_account_platforms(account_platform_map: dict[int, Any]) -> list[tuple[str, str]]:
    seen: set[str] = set()
    options: list[tuple[str, str]] = []
    for raw_value in account_platform_map.values():
        code = _normalized_account_platform(raw_value)
        if not code or code in seen:
            continue
        seen.add(code)
        label = platform_label_for_code(code) or ""
        if not label or label.casefold() == code.casefold():
            label = code.title()
        options.append((label, code))
    return sorted(options, key=lambda row: row[0].casefold())


def _cost_amount_placeholder(value_type: str) -> str:
    return COST_AMOUNT_PLACEHOLDER_BY_TYPE.get(value_type, "Tutar")


def _cost_source_field_placeholder(value_type: str) -> str:
    _ = value_type
    return COST_SOURCE_FIELD_PLACEHOLDER


def _cost_card_currency(card: dict[str, Any] | None, default: str = "TRY") -> str:
    code = _safe_str((card or {}).get("currency")).upper()
    fallback = _safe_str(default).upper()
    return code or fallback or "TRY"


def _numeric_field_candidate(value: Any) -> Optional[float]:
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    if isinstance(value, str):
        text = value.strip().replace(",", ".")
        if not text:
            return None
        try:
            return float(text)
        except Exception:
            return None
    return None


def _compact_numeric_text(value: Any) -> str:
    numeric = _numeric_field_candidate(value)
    if numeric is None:
        return _safe_str(value)
    text = f"{numeric:.4f}"
    return text.rstrip("0").rstrip(".") if "." in text else text


def _should_skip_db_field(path: str) -> bool:
    last = _safe_str(path.split(".")[-1]).casefold()
    if not last:
        return True
    if path in PREFERRED_DB_FIELD_LABELS or last in {key.casefold() for key in PREFERRED_DB_FIELD_LABELS}:
        return False
    if last in DB_FIELD_EXCLUDED_NAMES or last.endswith("id"):
        return True
    if any(token in last for token in ("price", "rate", "tax", "amount", "cost", "weight", "quantity", "discount")):
        return False
    return True


def _listing_db_field_options(item: dict[str, Any] | None) -> list[dict[str, str]]:
    payload = dict(item or {})
    options: "OrderedDict[str, dict[str, str]]" = OrderedDict()

    def remember(path: str, value: Any) -> None:
        clean_path = _safe_str(path)
        if not clean_path or clean_path in options:
            return
        note_label = PREFERRED_DB_FIELD_LABELS.get(clean_path) or clean_path
        display_label = note_label if note_label == clean_path else f"{note_label} ({clean_path})"
        display_value = _compact_numeric_text(value)
        if display_value:
            display_label = f"{display_label} • {display_value}"
        options[clean_path] = {
            "path": clean_path,
            "label": display_label,
            "noteLabel": note_label,
        }

    def walk(node: Any, *, prefix: str = "", depth: int = 0) -> None:
        if depth > 2 or not isinstance(node, dict):
            return
        for raw_key, raw_value in node.items():
            key = _safe_str(raw_key)
            if not key:
                continue
            path = f"{prefix}.{key}" if prefix else key
            if isinstance(raw_value, dict):
                walk(raw_value, prefix=path, depth=depth + 1)
                continue
            if isinstance(raw_value, (list, tuple, set)):
                continue
            if _numeric_field_candidate(raw_value) is None or _should_skip_db_field(path):
                continue
            remember(path, raw_value)

    for path in PREFERRED_DB_FIELD_LABELS:
        value = payload.get(path)
        if _numeric_field_candidate(value) is not None:
            remember(path, value)
    walk(payload)
    return list(options.values())


def _cost_row_is_complete(row: dict) -> bool:
    if not isinstance(row, dict):
        return False
    binding = normalize_listing_cost_binding(row)
    cost_key = _safe_str(binding.get("cost_key"))
    amount = binding.get("amount")
    source_type = _safe_str(binding.get("source_type"))
    source_config = dict(binding.get("source_config") or {})
    if not cost_key:
        return False
    if cost_key.casefold() == "kargo":
        try:
            provider_id = int(source_config.get("cargoProviderId") or source_config.get("cargo_provider_id") or 0)
        except Exception:
            provider_id = 0
        if provider_id <= 0:
            return False
    if source_type == "manual_value":
        return amount not in (None, "")
    if source_type == "db_field":
        return _safe_str(source_config.get("fieldPath")) != ""
    return True


def _compact_cost_row(row: dict) -> dict:
    binding = normalize_listing_cost_binding(row)
    return {
        "cost_key": _safe_str(binding.get("cost_key")),
        "calculation_type": _safe_str(binding.get("calculation_type")) or None,
        "source_type": _safe_str(binding.get("source_type")) or None,
        "source_config": dict(binding.get("source_config") or {}) or None,
        "amount": binding.get("amount"),
        "quantity": binding.get("quantity"),
        "tariff_key": binding.get("tariff_key"),
    }


def _remove_stock_match_row_at(rows_data: list[dict], idx: int) -> list[dict]:
    rows = [dict(row or {}) for row in (rows_data or [])]
    if 0 <= idx < len(rows):
        rows.pop(idx)
    return rows


def _first_image_url(item: dict) -> str:
    imgs = item.get("images") or []
    if isinstance(imgs, list) and imgs:
        first = imgs[0]
        if isinstance(first, dict):
            return _safe_str(first.get("url"))
        if isinstance(first, str):
            return _safe_str(first)
    return ""


def _normalize_image_url(url: str) -> str:
    if not url:
        return ""
    parts = urlsplit(url)
    # Remove CDN resize style to fetch higher quality original
    return urlunsplit((parts.scheme, parts.netloc, parts.path, "", ""))


def _image_cache_path(url: str) -> Optional[Path]:
    if not url:
        return None
    clean = url.split("?")[0]
    ext = Path(clean).suffix or ".jpg"
    name = hashlib.blake2b(url.encode("utf-8"), digest_size=16).hexdigest() + ext
    return LISTING_IMAGE_CACHE_DIR / name


def _cost_items_for_platform_from_cache(
    cost_items_cache: list[dict[str, Any]] | None,
    platform_code: Any,
) -> list[dict[str, Any]]:
    normalized_platform = _normalized_account_platform(platform_code)
    if not normalized_platform:
        return [dict(row) for row in (cost_items_cache or []) if isinstance(row, dict)]
    out: list[dict[str, Any]] = []
    for row in (cost_items_cache or []):
        if not isinstance(row, dict):
            continue
        supported = [
            _normalized_account_platform(value)
            for value in list(row.get("supportedPlatforms") or [])
            if _normalized_account_platform(value)
        ]
        if supported and normalized_platform not in supported:
            continue
        out.append(dict(row))
    return out


def _build_cost_card_map_from_cache(
    cost_items_cache: list[dict[str, Any]] | None,
    *,
    platform_code: Any = None,
) -> dict[str, dict]:
    return {
        _safe_str(row.get("costKey")).casefold(): dict(row)
        for row in _cost_items_for_platform_from_cache(cost_items_cache, platform_code)
        if _safe_str(row.get("costKey"))
    }


def _prepare_listing_item_for_view(
    product: dict[str, Any],
    *,
    account_name_map: dict[int, str] | None,
    account_logo_map: dict[int, str] | None,
    account_platform_map: dict[int, Any] | None,
    stock_item_label_map: dict[int, str] | None,
    stock_cost_map: dict[int, float] | None,
    stock_vat_rate_map: dict[int, float | None] | None,
    stock_cost_meta_map: dict[int, dict[str, Any]] | None,
    stock_dimension_map: dict[int, dict[str, float | None]] | None,
    cost_items_cache: list[dict[str, Any]] | None,
    cargo_tariffs: list[dict[str, Any]] | None,
    cost_card_map: Optional[dict[str, dict]] = None,
    sync_cost_snapshot: bool = False,
) -> dict[str, Any]:
    if not isinstance(product, dict):
        return {}

    prepared = dict(product)
    account_name_map = account_name_map or {}
    account_logo_map = account_logo_map or {}
    account_platform_map = account_platform_map or {}
    stock_item_label_map = stock_item_label_map or {}
    stock_cost_map = stock_cost_map or {}
    stock_vat_rate_map = stock_vat_rate_map or {}
    stock_cost_meta_map = stock_cost_meta_map or {}
    stock_dimension_map = stock_dimension_map or {}
    cargo_tariffs = cargo_tariffs or []

    acc_id = prepared.get("api_account_id")
    platform_code = account_platform_map.get(acc_id, "")
    resolved_cost_card_map = (
        cost_card_map
        if cost_card_map is not None
        else _build_cost_card_map_from_cache(cost_items_cache, platform_code=platform_code)
    )

    prepared["store_name"] = account_name_map.get(acc_id, "")
    prepared["store_logo"] = account_logo_map.get(acc_id, "")

    stock_items = prepared.get("stock_items") or []
    labels: list[str] = []
    for stock_item in stock_items:
        if not isinstance(stock_item, dict):
            continue
        stock_item_id = stock_item.get("stock_item_id")
        label = stock_item_label_map.get(stock_item_id, "")
        qty = stock_item.get("qty")
        if not label:
            continue
        if qty not in (None, ""):
            labels.append(f"{label} x{qty}")
        else:
            labels.append(label)
    prepared["linked_stock_label"] = ", ".join(labels)

    img_url = _normalize_image_url(_first_image_url(prepared))
    prepared["image_url"] = img_url
    img_path = None
    if img_url:
        cache_path = _image_cache_path(img_url)
        if cache_path:
            img_path = str(cache_path)
    prepared["image_path"] = img_path

    profit_payload = estimate_listing_profit(
        item=prepared,
        stock_cost_map=stock_cost_map,
        stock_vat_rate_map=stock_vat_rate_map,
        stock_cost_meta_map=stock_cost_meta_map,
        stock_dimension_map=stock_dimension_map,
        cost_card_map=resolved_cost_card_map,
        cargo_tariffs=cargo_tariffs,
        platform_code=platform_code,
    )
    prepared.update(profit_payload)
    if sync_cost_snapshot:
        action_sync_listing_cost_snapshots_for_item(item=prepared, estimation=profit_payload)
    return prepared


def _prepare_listing_items_for_view(
    items: list[dict[str, Any]] | None,
    *,
    account_name_map: dict[int, str] | None,
    account_logo_map: dict[int, str] | None,
    account_platform_map: dict[int, Any] | None,
    stock_item_label_map: dict[int, str] | None,
    stock_cost_map: dict[int, float] | None,
    stock_vat_rate_map: dict[int, float | None] | None,
    stock_cost_meta_map: dict[int, dict[str, Any]] | None,
    stock_dimension_map: dict[int, dict[str, float | None]] | None,
    cost_items_cache: list[dict[str, Any]] | None,
    cargo_tariffs: list[dict[str, Any]] | None,
    sync_cost_snapshot: bool = False,
) -> list[dict[str, Any]]:
    prepared_items: list[dict[str, Any]] = []
    platform_maps: dict[str, dict[str, dict]] = {}
    account_platform_map = account_platform_map or {}
    cargo_tariffs = cargo_tariffs or []

    for item in (items or []):
        if not isinstance(item, dict):
            continue
        platform_code = account_platform_map.get(item.get("api_account_id"), "")
        normalized_platform = _normalized_account_platform(platform_code)
        if normalized_platform not in platform_maps:
            platform_maps[normalized_platform] = _build_cost_card_map_from_cache(
                cost_items_cache,
                platform_code=normalized_platform,
            )
        prepared_items.append(
            _prepare_listing_item_for_view(
                item,
                account_name_map=account_name_map,
                account_logo_map=account_logo_map,
                account_platform_map=account_platform_map,
                stock_item_label_map=stock_item_label_map,
                stock_cost_map=stock_cost_map,
                stock_vat_rate_map=stock_vat_rate_map,
                stock_cost_meta_map=stock_cost_meta_map,
                stock_dimension_map=stock_dimension_map,
                cost_items_cache=cost_items_cache,
                cargo_tariffs=cargo_tariffs,
                cost_card_map=platform_maps[normalized_platform],
                sync_cost_snapshot=sync_cost_snapshot,
            )
        )
    return prepared_items


def _to_optional_number(value: Any) -> Optional[float]:
    if value is None or value == "":
        return None
    try:
        return float(value)
    except Exception:
        return None


def _apply_listing_view_filters(
    items: list[dict[str, Any]] | None,
    *,
    match_val: str,
    approval_val: str,
    campaign_val: str,
    lock_val: str,
    stock_min: Optional[float],
    stock_max: Optional[float],
    price_min: Optional[float],
    price_max: Optional[float],
) -> list[dict[str, Any]]:
    if not items:
        return []

    filter_match = match_val in ("matched", "unmatched")
    want_match = match_val == "matched"
    filter_approval = approval_val in ("approved", "unapproved")
    want_approved = approval_val == "approved"
    filter_campaign = campaign_val in ("campaign_on", "campaign_off")
    want_campaign_on = campaign_val == "campaign_on"
    filter_lock = lock_val in ("locked_on", "locked_off")
    want_locked = lock_val == "locked_on"
    filter_stock = stock_min is not None or stock_max is not None
    filter_price = price_min is not None or price_max is not None

    filtered_items: list[dict[str, Any]] = []
    for item in items:
        if not isinstance(item, dict):
            continue

        if filter_match:
            has_match = False
            stock_items = item.get("stock_items") or []
            if isinstance(stock_items, list):
                for stock_item in stock_items:
                    if isinstance(stock_item, dict) and stock_item.get("stock_item_id") is not None:
                        has_match = True
                        break
            if has_match != want_match:
                continue

        if filter_approval and ((item.get("approved") is True) != want_approved):
            continue

        if filter_campaign:
            campaign_state = item.get("hasActiveCampaign") if "hasActiveCampaign" in item else item.get("has_active_campaign")
            if bool(campaign_state) != want_campaign_on:
                continue

        if filter_lock and bool(item.get("locked")) != want_locked:
            continue

        if filter_stock:
            qty = _to_optional_number(item.get("quantity"))
            if qty is None:
                continue
            if stock_min is not None and qty < stock_min:
                continue
            if stock_max is not None and qty > stock_max:
                continue

        if filter_price:
            price = _to_optional_number(item.get("salePrice"))
            if price is None:
                price = _to_optional_number(item.get("listPrice"))
            if price is None:
                continue
            if price_min is not None and price < price_min:
                continue
            if price_max is not None and price > price_max:
                continue

        filtered_items.append(item)

    return filtered_items


def _prepare_filtered_listing_payload(
    *,
    query_params: dict[str, Any],
    filter_params: dict[str, Any],
    enrichment_context: dict[str, Any],
) -> Result:
    res = action_get_listings_from_db(**dict(query_params or {}))
    if not res.success:
        return res
    items = (res.data or {}).get("items", []) or []
    prepared_items = _prepare_listing_items_for_view(
        items,
        account_name_map=enrichment_context.get("account_name_map"),
        account_logo_map=enrichment_context.get("account_logo_map"),
        account_platform_map=enrichment_context.get("account_platform_map"),
        stock_item_label_map=enrichment_context.get("stock_item_label_map"),
        stock_cost_map=enrichment_context.get("stock_cost_map"),
        stock_vat_rate_map=enrichment_context.get("stock_vat_rate_map"),
        stock_cost_meta_map=enrichment_context.get("stock_cost_meta_map"),
        stock_dimension_map=enrichment_context.get("stock_dimension_map"),
        cost_items_cache=enrichment_context.get("cost_items_cache"),
        cargo_tariffs=enrichment_context.get("cargo_tariffs"),
        sync_cost_snapshot=False,
    )
    filtered_items = _apply_listing_view_filters(prepared_items, **dict(filter_params or {}))
    return Result.ok(
        "İlanlar hazırlandı.",
        close_dialog=False,
        data={"items": filtered_items},
    )


def _load_listing_cache_payload() -> Result:
    account_name_map: dict[int, str] = {}
    account_logo_map: dict[int, str] = {}
    account_platform_map: dict[int, str] = {}
    account_ids: list[int] = []

    account_res = action_get_all_companies()
    if account_res.success:
        records = (account_res.data or {}).get("records", []) or []
        account_name_map = {
            getattr(r, "pk", None): (getattr(r, "comp_name", "") or "").strip()
            for r in records
            if getattr(r, "pk", None) is not None
        }
        account_logo_map = {
            getattr(r, "pk", None): (getattr(r, "logo_path", "") or "")
            for r in records
            if getattr(r, "pk", None) is not None
        }
        account_platform_map = {
            getattr(r, "pk", None): _normalized_account_platform(getattr(r, "platform", "") or "")
            for r in records
            if getattr(r, "pk", None) is not None
        }
        account_ids = [rid for rid in account_name_map.keys()]

    stock_items_cache: list[dict] = []
    stock_item_label_map: dict[int, str] = {}
    stock_cost_map: dict[int, float] = {}
    stock_cost_meta_map: dict[int, dict[str, Any]] = {}
    stock_vat_rate_map: dict[int, float | None] = {}
    stock_dimension_map: dict[int, dict[str, float | None]] = {}
    stock_res = action_list_items(only_active=True)
    if stock_res.success:
        records = (stock_res.data or {}).get("records", []) or []
        stock_items_cache = [
            {
                "id": _stock_item_identity(it),
                "label": f"{getattr(it, 'sku', '')} - {getattr(it, 'name', '')}",
            }
            for it in records
            if _stock_item_identity(it) is not None
        ]
        stock_item_label_map = {
            it["id"]: it["label"]
            for it in stock_items_cache
        }
        for it in records:
            item_id = _stock_item_identity(it)
            if item_id is None:
                continue
            vat_rate = getattr(it, "vat_rate", None)
            try:
                stock_vat_rate_map[int(item_id)] = float(vat_rate) if vat_rate not in (None, "") else None
            except Exception:
                stock_vat_rate_map[int(item_id)] = None
            stock_dimension_map[int(item_id)] = {
                "weight_kg": getattr(it, "weight_kg", None),
                "length_cm": getattr(it, "length_cm", None),
                "width_cm": getattr(it, "width_cm", None),
                "height_cm": getattr(it, "height_cm", None),
            }
        balances_res = action_list_item_balances(only_active=True)
        if balances_res.success:
            for row in (balances_res.data or {}).get("records", []) or []:
                try:
                    item = row.get("item")
                    item_id = int(_stock_item_identity(item))
                except Exception:
                    continue
                try:
                    avg_cost = float(row.get("avg_unit_cost"))
                except Exception:
                    continue
                stock_cost_map[item_id] = avg_cost
                stock_cost_meta_map[item_id] = {
                    "input_mode": row.get("avg_cost_input_mode"),
                }

    cost_items_cache: list[dict] = []
    cost_res = action_load_finance_cost_items(active_state="Aktif")
    if cost_res.success:
        records = (cost_res.data or {}).get("records", []) or []
        cost_items_cache = [
            {
                "costKey": _safe_str(it.get("costKey")),
                "name": _safe_str(it.get("name")),
                "label": _safe_str(it.get("name")),
                "valueType": _safe_str(it.get("valueType") or it.get("calculationType")),
                "valueTypeLabel": _safe_str(it.get("valueTypeLabel") or it.get("calculationTypeLabel")),
                "defaultMatchMode": _safe_str(it.get("defaultMatchMode")) or "manual",
                "defaultMatchModeLabel": _safe_str(it.get("defaultMatchModeLabel")),
                "supportedSourceTypes": list(it.get("supportedSourceTypes") or []),
                "defaultSourceType": _safe_str(it.get("defaultSourceType")),
                "defaultSourceTypeLabel": _safe_str(it.get("defaultSourceTypeLabel")),
                "defaultSourceConfig": dict(it.get("defaultSourceConfig") or {}) or None,
                "defaultAmount": it.get("defaultAmount"),
                "currency": _safe_str(it.get("currency")) or "TRY",
                "isVatIncluded": bool(it.get("isVatIncluded", True)),
                "vatRate": it.get("vatRate"),
                "supportedPlatforms": list(it.get("supportedPlatforms") or []),
            }
            for it in records
            if _safe_str(it.get("costKey"))
            and (not bool(it.get("isBuiltin")) or bool(it.get("isBuiltinAccepted")))
        ]

    cargo_providers_cache: list[dict] = []
    provider_res = action_load_logistics_cargo_providers(source_filter="accepted")
    if provider_res.success:
        records = (provider_res.data or {}).get("records", []) or []
        cargo_providers_cache = [
            {
                "pk": int(it.get("pk") or 0),
                "platformCode": _safe_str(it.get("platformCode")),
                "cargoProviderCode": _safe_str(it.get("cargoProviderCode")).upper(),
                "cargoProviderName": _safe_str(it.get("cargoProviderName")),
            }
            for it in records
            if int(it.get("pk") or 0) > 0
            and _safe_str(it.get("cargoProviderCode"))
            and _safe_str(it.get("cargoProviderName"))
        ]

    cargo_tariffs_cache: list[dict] = []
    tariff_res = action_load_logistics_cargo_tariffs(active_state="Aktif")
    if tariff_res.success:
        records = (tariff_res.data or {}).get("records", []) or []
        cargo_tariffs_cache = [
            {
                "tariffKey": _safe_str(it.get("tariffKey")),
                "name": _safe_str(it.get("name")),
                "label": _safe_str(
                    it.get("name")
                    or f"{_safe_str(it.get('cargoProviderName'))} / {_safe_str(it.get('platform'))}"
                ),
                "cargo_provider_id": int(it.get("cargo_provider_id") or 0) or None,
                "platformCode": _safe_str(it.get("platformCode") or it.get("platform")),
                "platform": _safe_str(it.get("platform")),
                "cargoProviderCode": _safe_str(it.get("cargoProviderCode")),
                "cargoProviderName": _safe_str(it.get("cargoProviderName")),
                "minDesi": it.get("minDesi"),
                "maxDesi": it.get("maxDesi"),
                "minSalePrice": it.get("minSalePrice"),
                "maxSalePrice": it.get("maxSalePrice"),
                "amount": it.get("amount"),
                "currency": _safe_str(it.get("currency")),
                "isVatIncluded": bool(it.get("isVatIncluded", True)),
                "vatRate": it.get("vatRate"),
                "isActive": bool(it.get("isActive", True)),
            }
            for it in records
            if _safe_str(it.get("tariffKey"))
            and (not bool(it.get("isBuiltin")) or bool(it.get("isBuiltinAccepted")))
        ]

    return Result.ok(
        "İlan cache hazırlandı.",
        close_dialog=False,
        data={
            "account_name_map": account_name_map,
            "account_logo_map": account_logo_map,
            "account_platform_map": account_platform_map,
            "account_ids": account_ids,
            "stock_items_cache": stock_items_cache,
            "stock_item_label_map": stock_item_label_map,
            "stock_cost_map": stock_cost_map,
            "stock_cost_meta_map": stock_cost_meta_map,
            "stock_vat_rate_map": stock_vat_rate_map,
            "stock_dimension_map": stock_dimension_map,
            "cost_items_cache": cost_items_cache,
            "cargo_providers_cache": cargo_providers_cache,
            "cargo_tariffs_cache": cargo_tariffs_cache,
        },
    )


def _prepare_single_listing_item_payload(
    *,
    product: dict[str, Any],
    enrichment_context: dict[str, Any],
    platform_code: Any = None,
) -> Result:
    prepared = _prepare_listing_item_for_view(
        dict(product or {}),
        account_name_map=enrichment_context.get("account_name_map"),
        account_logo_map=enrichment_context.get("account_logo_map"),
        account_platform_map=enrichment_context.get("account_platform_map"),
        stock_item_label_map=enrichment_context.get("stock_item_label_map"),
        stock_cost_map=enrichment_context.get("stock_cost_map"),
        stock_vat_rate_map=enrichment_context.get("stock_vat_rate_map"),
        stock_cost_meta_map=enrichment_context.get("stock_cost_meta_map"),
        stock_dimension_map=enrichment_context.get("stock_dimension_map"),
        cost_items_cache=enrichment_context.get("cost_items_cache"),
        cargo_tariffs=enrichment_context.get("cargo_tariffs"),
        cost_card_map=_build_cost_card_map_from_cache(
            enrichment_context.get("cost_items_cache"),
            platform_code=platform_code,
        ),
        sync_cost_snapshot=False,
    )
    return Result.ok(
        "İlan hazırlandı.",
        close_dialog=False,
        data={"item": prepared},
    )


def _product_subtitle(item: dict) -> str:
    brand = _safe_str(item.get("brand") or item.get("brandName"))
    stock = _safe_str(item.get("stockCode") or item.get("stock_code"))
    barcode = _safe_str(item.get("barcode"))
    parts = []
    if brand:
        parts.append(f"Marka: {brand}")
    if stock:
        parts.append(f"Stok: {stock}")
    if barcode:
        parts.append(f"Barkod: {barcode}")
    return " | ".join(parts) if parts else ""


def _product_extra(item: dict) -> str:
    approved = item.get("approved")
    if isinstance(approved, bool):
        approved_text = "Onaylı" if approved else "Onaysız"
    else:
        approved_text = ""
    parts = []
    if approved_text:
        parts.append(approved_text)
    store = _safe_str(item.get("store_name"))
    if store:
        parts.append(f"Mağaza: {store}")
    return " | ".join(parts) if parts else ""


def _listing_right_side_rects(
    rect: QRect,
    view=None,
    *,
    hook_delegate: RecordWorkspaceListDelegate | None = None,
) -> dict[str, QRect]:
    delegate = hook_delegate or ListingListDelegate()
    effective_right = delegate._hook_content_right(rect, view=view, right_margin=8)

    content_left = rect.left() + 10 + 110 + 12
    visible_badges_left_reserved = content_left + LISTING_BADGE_VISIBLE_MIN_TEXT_W
    available_cluster_width = max(0, effective_right - visible_badges_left_reserved + 1)

    top_y = rect.top() + 18
    second_row_y = top_y + max(LISTING_QTY_H, LISTING_PRICE_H) + LISTING_BADGE_STACK_GAP
    gap = 14

    if available_cluster_width < LISTING_HIDE_BADGES_CLUSTER_W:
        qty_w = 116
        price_w = 172
        hook_rect = delegate._build_hook_rect(
            rect,
            view=view,
            content_right=effective_right,
            right_margin=8,
            vertical_rect=rect,
            hook_w=LISTING_HOOK_W,
            hook_h=LISTING_HOOK_H,
        )
        price_rect = QRect(hook_rect.left() - gap - price_w, top_y, price_w, LISTING_PRICE_H)
        qty_rect = QRect(price_rect.left() - gap - qty_w, top_y, qty_w, LISTING_QTY_H)
        profit_left = qty_rect.left()
        profit_width = max(188, hook_rect.left() - gap - profit_left)
        profit_rect = QRect(profit_left, second_row_y, profit_width, LISTING_PROFIT_H)
        return {
            "qty": qty_rect,
            "price": price_rect,
            "profit": profit_rect,
            "cost_badge": QRect(),
            "product_badge": QRect(),
            "hook": hook_rect,
            "compact": True,
            "badges_hidden": True,
        }

    base_cluster_width = (
        LISTING_QTY_W
        + 8
        + LISTING_PRICE_W
        + 10
        + LISTING_BADGE_BTN_W
        + LISTING_HOOK_BADGE_GAP
        + LISTING_HOOK_W
        + 10
    )
    if base_cluster_width > 0 and available_cluster_width > 0:
        scale = min(1.0, max(LISTING_MIN_RIGHT_SCALE, available_cluster_width / base_cluster_width))
    else:
        scale = 1.0

    if available_cluster_width < base_cluster_width:
        compression_span = max(1, base_cluster_width - LISTING_HIDE_BADGES_CLUSTER_W)
        compression = max(
            0.0,
            min(1.0, (available_cluster_width - LISTING_HIDE_BADGES_CLUSTER_W) / compression_span),
        )
        qty_w = int(LISTING_COMPACT_QTY_MIN_W + ((LISTING_QTY_W - LISTING_COMPACT_QTY_MIN_W) * compression))
        price_w = int(LISTING_COMPACT_PRICE_MIN_W + ((LISTING_PRICE_W - LISTING_COMPACT_PRICE_MIN_W) * compression))
        btn_w = int(LISTING_COMPACT_BADGE_MIN_W + ((LISTING_BADGE_BTN_W - LISTING_COMPACT_BADGE_MIN_W) * compression))
        btn_h = int(LISTING_BADGE_MIN_H + ((LISTING_BADGE_BTN_H - LISTING_BADGE_MIN_H) * compression))
        compact_mode = True
    else:
        qty_w = max(112, int(LISTING_QTY_W * scale))
        price_w = max(124, int(LISTING_PRICE_W * scale))
        btn_w = max(LISTING_BADGE_MIN_W, int(LISTING_BADGE_BTN_W * scale))
        btn_h = max(LISTING_BADGE_MIN_H, int(LISTING_BADGE_BTN_H * scale))
        compact_mode = False
    qty_h = LISTING_QTY_H
    price_h = LISTING_PRICE_H

    right_pad = max(8, int(10 * scale))
    between_price_badge = max(8, int(12 * scale))
    between_qty_price = max(8, int(10 * scale))
    hook_rect = delegate._build_hook_rect(
        rect,
        view=view,
        content_right=delegate._hook_content_right(rect, view=view, right_margin=right_pad),
        right_margin=right_pad,
        vertical_rect=rect,
        hook_w=LISTING_HOOK_W,
        hook_h=LISTING_HOOK_H,
    )
    badge_x = hook_rect.left() - max(16, int(LISTING_HOOK_BADGE_GAP * scale)) - btn_w

    cost_btn_rect = QRect(badge_x, top_y, btn_w, btn_h)
    product_btn_rect = QRect(badge_x, second_row_y, btn_w, btn_h)

    price_rect = QRect(badge_x - between_price_badge - price_w, top_y, price_w, price_h)
    qty_rect = QRect(price_rect.left() - between_qty_price - qty_w, top_y, qty_w, qty_h)
    profit_rect = QRect(
        qty_rect.left(),
        second_row_y,
        max(196, price_rect.right() - qty_rect.left() + 1),
        LISTING_PROFIT_H,
    )

    return {
        "qty": qty_rect,
        "price": price_rect,
        "profit": profit_rect,
        "cost_badge": cost_btn_rect,
        "product_badge": product_btn_rect,
        "hook": hook_rect,
        "compact": compact_mode,
        "badges_hidden": False,
    }

class ListingListModel(QAbstractListModel):
    TitleRole = RecordWorkspaceListModel.TitleRole
    SubtitleRole = RecordWorkspaceListModel.SubtitleRole
    ExtraRole = RecordWorkspaceListModel.ExtraRole
    IconRole = RecordWorkspaceListModel.IconRole
    ProductRole = RecordWorkspaceListModel.RecordRole
    SelectedRole = RecordWorkspaceListModel.SelectedRole
    ActiveRole = RecordWorkspaceListModel.ActiveRole
    PlaceholderRole = RecordWorkspaceListModel.PlaceholderRole
    ProfitTextRole = Qt.ItemDataRole.UserRole + 101
    ProfitReadyRole = Qt.ItemDataRole.UserRole + 102
    ProfitValueRole = Qt.ItemDataRole.UserRole + 103
    ProfitBreakdownRole = Qt.ItemDataRole.UserRole + 104

    def __init__(self, items: list[dict] | None = None, parent=None, *, on_selection_changed=None):
        super().__init__(parent)
        self._items: list[dict] = items if isinstance(items, list) else list(items or [])
        self._active_row: int | None = None
        self._on_selection_changed = on_selection_changed

    def rowCount(self, parent=QModelIndex()):
        if parent.isValid():
            return 0
        return len(self._items)

    def data(self, index: QModelIndex, role: int = Qt.ItemDataRole.DisplayRole):
        if not index.isValid():
            return None
        row = index.row()
        if row < 0 or row >= len(self._items):
            return None
        item = self._items[row]

        if role == Qt.ItemDataRole.DisplayRole:
            return _product_title(item)
        if role == self.TitleRole:
            return _product_title(item)
        if role == self.SubtitleRole:
            return _product_subtitle(item)
        if role == self.ExtraRole:
            return _product_extra(item)
        if role == self.IconRole:
            return item.get("image_path") or item.get("store_logo")
        if role == self.ProductRole:
            return item
        if role == self.PlaceholderRole:
            return False
        if role == self.SelectedRole:
            return bool(item.get("_selected"))
        if role == self.ActiveRole:
            return row == self._active_row
        if role == self.ProfitTextRole:
            return item.get("estimated_profit_text") or "-"
        if role == self.ProfitReadyRole:
            return bool(item.get("estimated_profit_ready"))
        if role == self.ProfitValueRole:
            return item.get("estimated_profit")
        if role == self.ProfitBreakdownRole:
            return item.get("estimated_profit_breakdown") or []
        return None

    def set_items(self, items: list[dict]):
        self.beginResetModel()
        self._items = items if isinstance(items, list) else list(items or [])
        self._active_row = None
        self.endResetModel()

    def item_at(self, row: int) -> dict | None:
        if 0 <= row < len(self._items):
            return self._items[row]
        return None

    def set_selected(self, row: int, selected: bool) -> None:
        if 0 <= row < len(self._items):
            self._items[row]["_selected"] = bool(selected)
            idx = self.index(row)
            self.dataChanged.emit(idx, idx, [self.SelectedRole])
            if self._on_selection_changed:
                self._on_selection_changed()

    def set_active_row(self, row: int | None) -> None:
        prev_row = self._active_row
        new_row = row if row is not None and 0 <= row < len(self._items) else None
        if prev_row == new_row:
            return
        self._active_row = new_row
        if prev_row is not None and 0 <= prev_row < len(self._items):
            idx_prev = self.index(prev_row)
            self.dataChanged.emit(idx_prev, idx_prev, [self.ActiveRole])
        if new_row is not None and 0 <= new_row < len(self._items):
            idx_new = self.index(new_row)
            self.dataChanged.emit(idx_new, idx_new, [self.ActiveRole])

    def clear_active_row(self) -> None:
        self.set_active_row(None)

    def active_row(self) -> int | None:
        return self._active_row


class ListingListDelegate(RecordWorkspaceListDelegate):
    def __init__(self, parent=None, *, on_select_button=None):
        super().__init__(parent, on_select_button=on_select_button)
        self._icon_cache: OrderedDict[tuple[str, int, float], QPixmap] = OrderedDict()
        self._icon_cache_limit = 256
        self._icon_size = 110
        self._row_height = 176

    def _load_icon(self, path: str, dpr: float) -> QPixmap | None:
        if not path:
            return None
        key = (path, self._icon_size, float(dpr))
        if key in self._icon_cache:
            pix_cached = self._icon_cache.pop(key)
            self._icon_cache[key] = pix_cached
            return pix_cached
        pix = QPixmap(path)
        if pix.isNull():
            return None
        target = int(self._icon_size * dpr)
        pix = pix.scaled(
            target,
            target,
            Qt.AspectRatioMode.KeepAspectRatio,
            Qt.TransformationMode.SmoothTransformation,
        )
        pix.setDevicePixelRatio(dpr)
        self._icon_cache[key] = pix
        while len(self._icon_cache) > self._icon_cache_limit:
            self._icon_cache.popitem(last=False)
        return pix

    def invalidate_icon(self, path: str) -> None:
        if not path:
            return
        dead_keys = [k for k in self._icon_cache if k[0] == path]
        for k in dead_keys:
            self._icon_cache.pop(k, None)

    def _item_rect(self, option: QStyleOptionViewItem) -> QRect:
        return option.rect.adjusted(6, 4, -6, -4)

    def _hook_rect(self, rect: QRect, view) -> QRect:
        return _listing_right_side_rects(rect, view, hook_delegate=self)["hook"]

    def _paint_card_content(self, painter: QPainter, rect: QRect, index: QModelIndex, hook_rect: QRect) -> None:
        title = index.data(ListingListModel.TitleRole) or ""
        subtitle = index.data(ListingListModel.SubtitleRole) or ""
        extra = index.data(ListingListModel.ExtraRole) or ""
        profit_text = index.data(ListingListModel.ProfitTextRole) or "-"
        profit_ready = bool(index.data(ListingListModel.ProfitReadyRole))
        profit_value = index.data(ListingListModel.ProfitValueRole)
        icon_path = index.data(ListingListModel.IconRole) or ""

        x = rect.left() + 10
        y = rect.top() + 10

        # icon
        icon_rect = QRect(x, y, self._icon_size, self._icon_size)
        dpr = painter.device().devicePixelRatioF()
        pix = self._load_icon(icon_path, dpr)
        if pix is not None:
            pr = pix.devicePixelRatio() or 1.0
            pw = int(pix.width() / pr)
            ph = int(pix.height() / pr)
            target = QRect(0, 0, pw, ph)
            target.moveCenter(icon_rect.center())
            painter.drawPixmap(target, pix)
        else:
            painter.setBrush(QColor("#E5E7EB"))
            painter.setPen(Qt.PenStyle.NoPen)
            painter.drawRoundedRect(icon_rect, 10, 10)
        x += self._icon_size + 12

        # right controls area
        side_rects = _listing_right_side_rects(rect, self.parent(), hook_delegate=self)
        hook_rect = side_rects["hook"]
        price_rect = side_rects["price"]
        qty_rect = side_rects["qty"]
        text_right = qty_rect.left() - LISTING_TEXT_CLUSTER_GAP

        # text
        painter.setPen(QColor("#111827"))
        title_font = QFont("Segoe UI", 10, QFont.Weight.Bold)
        painter.setFont(title_font)
        title_rect = QRect(x, rect.top() + 10, max(10, text_right - x), 22)
        fm = painter.fontMetrics()
        painter.drawText(title_rect, Qt.AlignmentFlag.AlignLeft | Qt.AlignmentFlag.AlignVCenter, fm.elidedText(title, Qt.TextElideMode.ElideRight, title_rect.width()))

        painter.setPen(QColor("#374151"))
        sub_font = QFont("Segoe UI", 8)
        painter.setFont(sub_font)
        sub_rect = QRect(x, title_rect.bottom() + 6, title_rect.width(), 18)
        fm = painter.fontMetrics()
        painter.drawText(sub_rect, Qt.AlignmentFlag.AlignLeft | Qt.AlignmentFlag.AlignVCenter, fm.elidedText(subtitle, Qt.TextElideMode.ElideRight, sub_rect.width()))

        painter.setPen(QColor("#6B7280"))
        extra_rect = QRect(x, sub_rect.bottom() + 4, title_rect.width(), 18)
        fm = painter.fontMetrics()
        painter.drawText(extra_rect, Qt.AlignmentFlag.AlignLeft | Qt.AlignmentFlag.AlignVCenter, fm.elidedText(extra, Qt.TextElideMode.ElideRight, extra_rect.width()))

    def editorEvent(self, event, model, option, index):
        rect = self._item_rect(option)
        x = rect.left() + 10
        y = rect.top() + 10
        icon_rect = QRect(x, y, self._icon_size, self._icon_size)
        view = option.widget
        side_rects = _listing_right_side_rects(rect, view, hook_delegate=self)
        hook_rect = side_rects["hook"]
        price_rect = side_rects["price"]
        qty_rect = side_rects["qty"]
        if event.type() in (QEvent.Type.MouseMove, QEvent.Type.HoverMove):
            pos = event.position().toPoint()
            hovering = hook_rect.contains(pos)
            if hovering and self._hover_row != index.row():
                self._hover_row = index.row()
                if view:
                    view.viewport().setCursor(Qt.CursorShape.PointingHandCursor)
                    view.viewport().update(option.rect.adjusted(-2, -2, 2, 2))
            elif not hovering and self._hover_row == index.row():
                self._hover_row = None
                if view:
                    view.viewport().setCursor(Qt.CursorShape.PointingHandCursor)
                    view.viewport().update(option.rect.adjusted(-2, -2, 2, 2))

        if event.type() in (QEvent.Type.Leave, QEvent.Type.HoverLeave):
            if self._hover_row is not None:
                self._hover_row = None
                if view:
                    view.viewport().setCursor(Qt.CursorShape.ArrowCursor)
                    view.viewport().update(option.rect.adjusted(-2, -2, 2, 2))

        if event.type() == QEvent.Type.MouseButtonPress and event.button() == Qt.MouseButton.LeftButton:
            pos = event.position().toPoint()
            if icon_rect.contains(pos):
                return True
            if qty_rect.contains(pos) or price_rect.contains(pos):
                return False
            if hook_rect.contains(pos):
                if view and hasattr(view, "set_active_row"):
                    view.set_active_row(index.row())
                return super().editorEvent(event, model, option, index)

        if event.type() == QEvent.Type.MouseButtonRelease and event.button() == Qt.MouseButton.LeftButton:
            pos = event.position().toPoint()
            if icon_rect.contains(pos):
                product = index.data(ListingListModel.ProductRole)
                if view and hasattr(view, "_open_image_cb") and view._open_image_cb:
                    view._open_image_cb(product)
                return True
            if hook_rect.contains(pos):
                if view and hasattr(view, "set_active_row"):
                    view.set_active_row(index.row())
                return super().editorEvent(event, model, option, index)
        return super().editorEvent(event, model, option, index)

    def sizeHint(self, option: QStyleOptionViewItem, index: QModelIndex) -> QSize:
        return QSize(option.rect.width(), self._row_height)


class InlinePriceControls(QWidget):
    def __init__(self, parent=None):
        super().__init__(parent)
        self.setAttribute(Qt.WidgetAttribute.WA_TransparentForMouseEvents, False)
        self.setAutoFillBackground(False)
        self.setMouseTracking(True)
        self.setAttribute(Qt.WidgetAttribute.WA_Hover, True)
        self._hook_view: Optional[QListView] = None
        self._hook_row_rect: Optional[QRect] = None
        self._price_currency_code = "TRY"
        self._price_currency_label = "TL"

        self.qty = LabeledActionInput(self, label="Stok", placeholder="0", button_text="OK", idle_glow=False)
        self.price = LabeledActionInput(self, label="Fiyat", placeholder="0.00", button_text="OK", idle_glow=False)
        self.qty.input.setObjectName("Input")
        self.price.input.setObjectName("Input")
        self.qty.input.setFixedHeight(30)
        self.price.input.setFixedHeight(30)
        if hasattr(self.qty, "_sync_sizes"):
            self.qty._sync_sizes()
        if hasattr(self.price, "_sync_sizes"):
            self.price._sync_sizes()
        self.qty.input.setValidator(QIntValidator(0, 10_000_000, self))
        price_rx = QRegularExpression(r"^\d*(?:[.,]\d{0,2})?$")
        self.price.input.setValidator(QRegularExpressionValidator(price_rx, self))

        self.installEventFilter(self)
        for child in self.findChildren(QWidget):
            child.setMouseTracking(True)
            child.setAttribute(Qt.WidgetAttribute.WA_Hover, True)
            child.installEventFilter(self)

    def set_hook_context(self, view: QListView, row_rect: QRect) -> None:
        self._hook_view = view
        self._hook_row_rect = QRect(row_rect)

    def _update_hook_cursor(self, global_pos) -> None:
        if self._hook_view is None or self._hook_row_rect is None or global_pos is None:
            return
        if hasattr(global_pos, "toPoint"):
            global_pos = global_pos.toPoint()
        pos = self._hook_view.viewport().mapFromGlobal(global_pos)
        if self._hook_row_rect.contains(pos):
            self._hook_view.viewport().setCursor(Qt.CursorShape.PointingHandCursor)
        else:
            self._hook_view.viewport().setCursor(Qt.CursorShape.ArrowCursor)

    def eventFilter(self, obj, event):
        if obj is self.price.input:
            if event.type() == QEvent.Type.FocusIn:
                self._prepare_price_for_edit()
            elif event.type() == QEvent.Type.FocusOut:
                self._apply_price_display_text()
        if event.type() in (QEvent.Type.MouseMove, QEvent.Type.HoverMove):
            gp = None
            if hasattr(event, "globalPosition"):
                gp = event.globalPosition()
            elif hasattr(event, "globalPos"):
                gp = event.globalPos()
            self._update_hook_cursor(gp)
        elif event.type() in (QEvent.Type.Leave, QEvent.Type.HoverLeave):
            if self._hook_view is not None:
                self._hook_view.viewport().setCursor(Qt.CursorShape.ArrowCursor)
        return super().eventFilter(obj, event)

    def mousePressEvent(self, event):
        try:
            if event.button() == Qt.MouseButton.LeftButton and self._hook_view is not None:
                global_pos = self.mapToGlobal(event.position().toPoint())
                view_pos = self._hook_view.viewport().mapFromGlobal(global_pos)
                idx = self._hook_view.indexAt(view_pos)
                if idx.isValid() and hasattr(self._hook_view, "set_active_row"):
                    self._hook_view.set_active_row(idx.row())
        except Exception:
            pass
        super().mousePressEvent(event)

    def place_controls(self, qty_rect: QRect, price_rect: QRect) -> None:
        qty_offset = 0
        price_offset = 0
        try:
            qty_offset = self.qty.input_center_offset()
        except Exception:
            qty_offset = 0
        try:
            price_offset = self.price.input_center_offset()
        except Exception:
            price_offset = 0
        if qty_offset:
            qty_rect = qty_rect.translated(0, -qty_offset)
        if price_offset:
            price_rect = price_rect.translated(0, -price_offset)
        self.qty.setGeometry(qty_rect)
        self.price.setGeometry(price_rect)

    @staticmethod
    def _format_price_number(price: float) -> str:
        text = f"{float(price):.2f}".rstrip("0").rstrip(".")
        return text or "0"

    def _price_suffix(self) -> str:
        return _listing_currency_label(self._price_currency_code)

    def _prepare_price_for_edit(self) -> None:
        text = _safe_str(self.price.input.text())
        suffix = self._price_suffix()
        if suffix and text.upper().endswith(f" {suffix}".upper()):
            numeric = text[: -(len(suffix) + 1)].rstrip()
            self.price.input.setText(numeric)

    def _apply_price_display_text(self) -> None:
        text = _safe_str(self.price.input.text())
        if not text:
            return
        suffix = self._price_suffix()
        if suffix and not text.upper().endswith(f" {suffix}".upper()):
            self.price.input.setText(f"{text} {suffix}")

    def set_currency_code(self, currency_code: Any) -> None:
        self._price_currency_code = _listing_currency_code({"currency": currency_code})
        self._price_currency_label = _listing_currency_label(self._price_currency_code)

    def set_values(self, qty: float | int, price: float, *, currency_code: Any = None) -> None:
        try:
            self.qty.setValue(str(int(qty)))
        except Exception:
            self.qty.setValue("0")
        try:
            val = float(price)
        except Exception:
            val = 0.0
        if currency_code is not None:
            self.set_currency_code(currency_code)
        text = self._format_price_number(val)
        self.price.setValue(text)
        if not self.price.input.hasFocus():
            self._apply_price_display_text()

    def qty_value(self) -> int:
        try:
            return int(self.qty.value() or "0")
        except Exception:
            return 0

    def price_value(self) -> float:
        text = _safe_str(self.price.value())
        suffix = self._price_suffix()
        if suffix and text.upper().endswith(f" {suffix}".upper()):
            text = text[: -(len(suffix) + 1)].rstrip()
        t = text.replace(",", ".")
        try:
            return float(t)
        except Exception:
            return 0.0

    def set_state(self, state: str) -> None:
        self.qty.setState(state)
        self.price.setState(state)

    def set_in_view(self, active: bool) -> None:
        for ctrl in (self.qty, self.price):
            inp = getattr(ctrl, "input", None)
            if inp is not None and hasattr(inp, "setInView"):
                try:
                    inp.setInView(active)
                except Exception:
                    pass


class ClickableFrame(QFrame):
    clicked = pyqtSignal()

    def __init__(self, parent=None):
        super().__init__(parent)
        self.setCursor(Qt.CursorShape.PointingHandCursor)

    def mouseReleaseEvent(self, event):
        if event.button() == Qt.MouseButton.LeftButton and self.rect().contains(event.position().toPoint()):
            self.clicked.emit()
            event.accept()
            return
        super().mouseReleaseEvent(event)


class ProfitMetricIcon(QWidget):
    def __init__(self, kind: str, accent: str, background: str, parent=None):
        super().__init__(parent)
        self._kind = kind
        self._accent = QColor(accent)
        self._background = QColor(background)
        self.setFixedSize(64, 64)

    def set_colors(self, *, accent: str, background: str) -> None:
        self._accent = QColor(accent)
        self._background = QColor(background)
        self.update()

    def paintEvent(self, event) -> None:
        painter = QPainter(self)
        painter.setRenderHint(QPainter.RenderHint.Antialiasing)
        rect = self.rect().adjusted(2, 2, -2, -2)
        painter.setBrush(QBrush(self._background))
        painter.setPen(QPen(self._accent, 1.2))
        painter.drawEllipse(rect)

        pen = QPen(self._accent, 3)
        pen.setCapStyle(Qt.PenCapStyle.RoundCap)
        pen.setJoinStyle(Qt.PenJoinStyle.RoundJoin)
        painter.setPen(pen)
        cx = self.width() // 2
        cy = self.height() // 2

        if self._kind == "sale":
            painter.drawLine(cx - 18, cy - 14, cx - 12, cy - 14)
            painter.drawLine(cx - 12, cy - 14, cx - 8, cy + 10)
            painter.drawLine(cx - 7, cy - 5, cx + 18, cy - 5)
            painter.drawLine(cx - 4, cy + 9, cx + 15, cy + 9)
            painter.drawLine(cx + 18, cy - 5, cx + 14, cy + 9)
            painter.setBrush(QBrush(self._accent))
            painter.drawEllipse(cx - 5, cy + 17, 5, 5)
            painter.drawEllipse(cx + 11, cy + 17, 5, 5)
            return

        if self._kind == "cost":
            painter.drawLine(cx, cy - 18, cx + 17, cy - 8)
            painter.drawLine(cx + 17, cy - 8, cx, cy + 2)
            painter.drawLine(cx, cy + 2, cx - 17, cy - 8)
            painter.drawLine(cx - 17, cy - 8, cx, cy - 18)
            painter.drawLine(cx - 17, cy - 8, cx - 17, cy + 12)
            painter.drawLine(cx + 17, cy - 8, cx + 17, cy + 12)
            painter.drawLine(cx - 17, cy + 12, cx, cy + 22)
            painter.drawLine(cx + 17, cy + 12, cx, cy + 22)
            painter.drawLine(cx, cy + 2, cx, cy + 22)
            return

        painter.drawLine(cx - 18, cy + 14, cx - 6, cy + 2)
        painter.drawLine(cx - 6, cy + 2, cx + 4, cy + 8)
        painter.drawLine(cx + 4, cy + 8, cx + 18, cy - 12)
        painter.drawLine(cx + 18, cy - 12, cx + 18, cy - 1)
        painter.drawLine(cx + 18, cy - 12, cx + 7, cy - 12)


class ProfitDetailControl(QWidget):
    detailRequested = pyqtSignal()

    def __init__(self, parent=None):
        super().__init__(parent)
        self._profit_text = "-"
        self._ready = False
        self._profit_value = None
        self.setMouseTracking(True)
        layout = QVBoxLayout(self)
        layout.setContentsMargins(0, 0, 0, 0)
        layout.setSpacing(2)

        self.caption = QLabel("Tahmini Kar")
        self.caption.setStyleSheet("font-size: 10px; color: #64748B;")
        self.caption.setFixedHeight(12)

        self.shell = QFrame(self)
        self.shell.setObjectName("ProfitField")
        shell_layout = QHBoxLayout(self.shell)
        shell_layout.setContentsMargins(0, 0, 0, 0)
        shell_layout.setSpacing(0)

        self.value_label = QLabel("-")
        self.value_label.setObjectName("ProfitValueLabel")
        self.value_label.setAlignment(Qt.AlignmentFlag.AlignVCenter | Qt.AlignmentFlag.AlignLeft)
        self.value_label.setMargin(0)
        self.value_label.setSizePolicy(QSizePolicy.Policy.Expanding, QSizePolicy.Policy.Expanding)

        self.detail_btn = QPushButton("Detay")
        self.detail_btn.setObjectName("ProfitDetailBtn")
        self.detail_btn.setCursor(Qt.CursorShape.PointingHandCursor)
        self.detail_btn.setMinimumWidth(92)
        self.detail_btn.setSizePolicy(QSizePolicy.Policy.Expanding, QSizePolicy.Policy.Expanding)
        self.detail_btn.clicked.connect(self.detailRequested.emit)

        shell_layout.addWidget(self.value_label, 1)
        shell_layout.addWidget(self.detail_btn, 1)
        shell_layout.setStretch(0, 11)
        shell_layout.setStretch(1, 5)

        layout.addWidget(self.caption)
        layout.addWidget(self.shell, 1)

        self.setMinimumHeight(LISTING_PROFIT_H)
        self.setStyleSheet(
            """
            QFrame#ProfitField {
                background: #FFFFFF;
                border: 1px solid #E5E7EB;
                border-radius: 12px;
            }
            QLabel#ProfitValueLabel {
                padding-left: 12px;
                padding-right: 8px;
                font-size: 11px;
                font-weight: 700;
                color: #0F172A;
            }
            QPushButton#ProfitDetailBtn {
                border: 0px;
                border-left: 1px solid #E5E7EB;
                background: #F8FAFC;
                color: #0F172A;
                font-size: 11px;
                font-weight: 800;
                border-top-right-radius: 12px;
                border-bottom-right-radius: 12px;
            }
            QPushButton#ProfitDetailBtn:hover {
                background: #EFF6FF;
                color: #2563EB;
            }
            QPushButton#ProfitDetailBtn:pressed {
                background: #DBEAFE;
            }
            """
        )

    def set_profit(self, text: str, ready: bool, value: Any) -> None:
        self._profit_text = _safe_str(text) or "-"
        self._ready = bool(ready)
        self._profit_value = value
        self.value_label.setText(self._profit_text)
        if not self._ready:
            color = "#94A3B8"
        else:
            try:
                color = "#DC2626" if float(value or 0) < 0 else "#15803D"
            except Exception:
                color = "#0F172A"
        self.value_label.setStyleSheet(
            f"padding-left: 12px; padding-right: 8px; font-size: 11px; font-weight: 700; color: {color};"
        )


class ProfitBreakdownDialog(QWidget):
    _DETAIL_ROWS_PER_PAGE = 4
    _MISSING_ROWS_PER_PAGE = 4

    def __init__(
        self,
        parent: QWidget,
        *,
        product_title: str,
        product: dict,
        estimation_context: Optional[dict[str, Any]] = None,
    ):
        super().__init__(parent, Qt.WindowType.Window)
        self._product_title = product_title
        self._base_product = deepcopy(product or {})
        self._estimation_context = dict(estimation_context or {})
        self._estimation_cache: dict[int, dict[str, Any]] = {}
        self._summary_expanded = False
        self._active_detail_tabs: set[str] = set()
        self._summary_cards: dict[str, dict[str, Any]] = {}
        self._export_drawer_visible = False
        self._missing_expanded = False
        self._missing_page_index = 0
        self._detail_page_indexes = {
            "sale": 0,
            "cost": 0,
        }
        self._detail_scroll_value: int | None = None
        self._detail_viewport_lock_generation = 0
        self.setObjectName("ListingProfitWindow")
        self.setWindowTitle("Tahmini Kar Detayı")
        self.resize(1180, 820)
        self.setMinimumSize(1040, 740)
        flags = self.windowFlags()
        flags |= Qt.WindowType.WindowMinMaxButtonsHint
        flags |= Qt.WindowType.WindowCloseButtonHint
        flags &= ~Qt.WindowType.WindowContextHelpButtonHint
        self.setWindowFlags(flags)
        self.setStyleSheet(
            """
            QWidget#ListingProfitWindow {
                background: #F7FAFD;
            }
            QLabel {
                background: transparent;
            }
            QFrame#HeroPanel, QFrame#EquationCard, QFrame#RevealPanel {
                background: transparent;
                border: 0px;
                border-radius: 0px;
            }
            QFrame#ControlCard, QFrame#SideCard, QFrame#FooterCard {
                background: #FFFFFF;
                border: 1px solid #E2E8F0;
                border-radius: 18px;
            }
            QFrame#SideStatRow {
                background: #F8FAFC;
                border: 1px solid #E8EEF6;
                border-radius: 16px;
            }
            QFrame#MissingPanel {
                background: #FFFBEB;
                border: 2px solid #F59E0B;
                border-radius: 22px;
            }
            QFrame#MissingPanel:hover {
                background: #FFF7ED;
                border-color: #D97706;
            }
            QFrame#MissingReasonRow {
                background: #FFFFFF;
                border: 1px solid #FDE68A;
                border-radius: 16px;
            }
            QFrame#ProfitHero {
                background: #ECFDF5;
                border: 2px solid #16A34A;
                border-radius: 28px;
            }
            QFrame#ProfitHero[danger="true"] {
                background: #FEF2F2;
                border-color: #DC2626;
            }
            QFrame#MetricCard {
                background: #FFFFFF;
                border: 1px solid #DDE7F2;
                border-radius: 24px;
            }
            QFrame#MetricCard[active="true"] {
                border-width: 2px;
            }
            QFrame#InnerDrawer {
                background: rgba(255, 255, 255, 0.84);
                border: 1px solid #E6EDF5;
                border-radius: 18px;
            }
            QFrame#InnerDrawer[kind="sale"] {
                background: #F7FFF9;
                border-top: 3px solid #16A34A;
            }
            QFrame#InnerDrawer[kind="cost"] {
                background: #FFF8F8;
                border-top: 3px solid #DC2626;
            }
            QFrame#LedgerRow {
                background: #FFFFFF;
                border: 1px solid #E5EDF6;
                border-radius: 14px;
            }
            QFrame#LedgerTotalRow {
                background: #F8FAFC;
                border: 1px solid #DDE7F2;
                border-radius: 14px;
            }
            QFrame#OperatorBadge {
                background: #FFFFFF;
                border: 1px solid #D9E2EC;
                border-radius: 24px;
            }
            QFrame#ExportDrawer {
                background: #FFFFFF;
                border: 1px solid #D9E2EC;
                border-radius: 16px;
            }
            QLabel#DialogTitle { font-size: 24px; font-weight: 900; color: #0F172A; }
            QLabel#DialogSub { font-size: 12px; color: #64748B; }
            QLabel#SectionTitle { font-size: 13px; font-weight: 800; color: #0F172A; }
            QLabel#BreakdownLabel, QLabel#CardTitle, QLabel#CostTitle { font-size: 11px; font-weight: 800; color: #475569; }
            QLabel#BreakdownValue, QLabel#CostRowValue { font-size: 11px; font-weight: 800; color: #0F172A; }
            QLabel#CardValue { font-size: 28px; font-weight: 900; color: #0F172A; }
            QLabel#CardNote, QLabel#CostRowLabel { font-size: 11px; color: #64748B; }
            QLabel#ProfitValue { font-size: 54px; font-weight: 950; }
            QLabel#FormulaText { font-size: 14px; font-weight: 800; color: #64748B; }
            QLabel#EquationLabel { font-size: 12px; font-weight: 900; color: #475569; }
            QLabel#EquationValue { font-size: 34px; font-weight: 950; }
            QLabel#EquationHint { font-size: 12px; font-weight: 700; color: #64748B; }
            QLabel#MissingValue { font-size: 42px; font-weight: 950; color: #B45309; }
            QLabel#MissingHint { font-size: 12px; font-weight: 750; color: #64748B; }
            QLabel#MissingChevron { font-size: 28px; font-weight: 950; color: #92400E; }
            QLabel#MissingReasonText { font-size: 12px; font-weight: 650; color: #475569; }
            QLabel#SideStatLabel { font-size: 11px; font-weight: 700; color: #64748B; }
            QLabel#SideValue { font-size: 13px; font-weight: 900; color: #0F172A; }
            QPushButton#PrimaryBtn {
                padding: 9px 16px;
                border-radius: 12px;
                border: 1px solid #0EA5E9;
                background: #0EA5E9;
                color: white;
                font-size: 12px;
                font-weight: 800;
            }
            QPushButton#PrimaryBtn:hover {
                background: #0284C7;
                border-color: #0284C7;
            }
            QPushButton#GhostBtn {
                padding: 9px 14px;
                border-radius: 12px;
                border: 1px solid #D9E2EC;
                background: #FFFFFF;
                color: #0F172A;
                font-size: 12px;
                font-weight: 700;
            }
            QPushButton#GhostBtn:hover {
                background: #F8FAFC;
            }
            QFrame#DetailPaginationBar {
                background: #F8FAFC;
                border: 1px solid #DDE7F2;
                border-radius: 18px;
            }
            QPushButton#PageNavButton {
                min-width: 42px;
                max-width: 42px;
                min-height: 42px;
                max-height: 42px;
                border-radius: 21px;
                border: 1px solid #CBD5E1;
                background: #FFFFFF;
                color: #0F172A;
                font-size: 22px;
                font-weight: 950;
            }
            QPushButton#PageNavButton:hover {
                background: #EFF6FF;
                border-color: #93C5FD;
            }
            QPushButton#PageNavButton:disabled {
                background: #F1F5F9;
                color: #94A3B8;
                border-color: #E2E8F0;
            }
            QLabel#PageIndicator {
                color: #334155;
                font-size: 12px;
                font-weight: 900;
            }
            QScrollArea {
                border: 0px;
                background: transparent;
            }
            """
        )

        root = QVBoxLayout(self)
        root.setContentsMargins(28, 24, 28, 20)
        root.setSpacing(18)

        header = QHBoxLayout()
        header.setContentsMargins(0, 0, 0, 0)
        header.setSpacing(12)
        title_wrap = QVBoxLayout()
        title_wrap.setContentsMargins(0, 0, 0, 0)
        title_wrap.setSpacing(6)
        title = QLabel("Tahmini Kar Detayı")
        title.setObjectName("DialogTitle")
        listing = QLabel(f"İlan: {product_title}")
        listing.setObjectName("DialogSub")
        listing.setWordWrap(True)
        title_wrap.addWidget(title)
        title_wrap.addWidget(listing)
        self.btn_export = QPushButton("Dışa Aktar")
        self.btn_export.setObjectName("GhostBtn")
        self.btn_export.setCursor(Qt.CursorShape.PointingHandCursor)
        self.btn_export.clicked.connect(self._toggle_export_drawer)
        header.addLayout(title_wrap, 1)
        header.addWidget(self.btn_export, 0, Qt.AlignmentFlag.AlignTop)
        root.addLayout(header)

        self.export_drawer = QFrame(self)
        self.export_drawer.setObjectName("ExportDrawer")
        export_layout = QHBoxLayout(self.export_drawer)
        export_layout.setContentsMargins(14, 12, 14, 12)
        export_layout.setSpacing(10)
        export_title = QLabel("Dışa aktarım")
        export_title.setObjectName("SectionTitle")
        export_hint = QLabel("Çıktı otomatik isimlendirilir; kaydetme yerini sen seçersin.")
        export_hint.setObjectName("DialogSub")
        export_hint.setWordWrap(True)
        export_excel = QPushButton("Excel olarak dışa aktar")
        export_excel.setObjectName("PrimaryBtn")
        export_excel.setCursor(Qt.CursorShape.PointingHandCursor)
        export_excel.clicked.connect(self._export_to_excel)
        export_layout.addWidget(export_title, 0)
        export_layout.addWidget(export_hint, 1)
        export_layout.addWidget(export_excel, 0)
        self.export_drawer.setVisible(False)
        root.addWidget(self.export_drawer)

        control_card = QFrame(self)
        control_card.setObjectName("ControlCard")
        control_layout = QHBoxLayout(control_card)
        control_layout.setContentsMargins(14, 10, 14, 10)
        control_layout.setSpacing(10)
        control_title = QLabel("Satılacak Ürün Adedi")
        control_title.setObjectName("CardTitle")
        control_layout.addWidget(control_title, 0)
        self.sale_qty = NumberStepper(minimum=1, maximum=10_000, step=1)
        initial_qty = 1
        try:
            initial_qty = max(
                1,
                int(
                    round(
                        float(
                            (
                                (self._base_product.get("estimated_profit_detail") or {})
                                .get("summary", {})
                                .get("sale_quantity")
                            )
                            or 1
                        )
                    )
                ),
            )
        except Exception:
            initial_qty = 1
        self.sale_qty.setValue(initial_qty)
        self.sale_qty.setMinimumWidth(120)
        control_layout.addWidget(self.sale_qty, 0)
        control_hint = QLabel("Kar hesabı seçtiğin toplam satış adedine göre yeniden hesaplanır.")
        control_hint.setObjectName("DialogSub")
        control_hint.setWordWrap(True)
        control_layout.addWidget(control_hint, 1)
        root.addWidget(control_card)

        body_row = QHBoxLayout()
        body_row.setContentsMargins(0, 0, 0, 0)
        body_row.setSpacing(12)

        body = QScrollArea(self)
        body.setWidgetResizable(True)
        body.setFrameShape(QFrame.Shape.NoFrame)
        body.setHorizontalScrollBarPolicy(Qt.ScrollBarPolicy.ScrollBarAlwaysOff)
        body.setVerticalScrollBarPolicy(Qt.ScrollBarPolicy.ScrollBarAlwaysOn)
        body.viewport().setStyleSheet("background: transparent;")
        body.verticalScrollBar().valueChanged.connect(self._position_detail_overlays)
        self.body_scroll = body
        body_content = QWidget(body)
        body_content.setObjectName("ProfitBodyContent")
        body_content.setStyleSheet("background: transparent;")
        self._body_content = body_content
        body_layout = QVBoxLayout(body_content)
        body_layout.setContentsMargins(0, 0, 18, 0)
        body_layout.setSpacing(12)
        self._body_layout = body_layout
        body.setWidget(body_content)
        body_row.addWidget(body, 1)

        side_host = QWidget(self)
        side_host.setObjectName("ProfitSideHost")
        side_host.setStyleSheet("background: transparent;")
        side_host.setFixedWidth(280)
        side_layout = QVBoxLayout(side_host)
        side_layout.setContentsMargins(0, 0, 0, 0)
        side_layout.setSpacing(10)
        self.health_card, self.health_body_layout = self._create_side_card("Veri Sağlığı")
        self.logic_card, self.logic_body_layout = self._create_side_card("Hesap Mantığı")
        side_layout.addWidget(self.health_card)
        side_layout.addWidget(self.logic_card)
        side_layout.addStretch(1)
        body_row.addWidget(side_host, 0)
        root.addLayout(body_row, 1)
        self.sale_qty.valueChanged.connect(self._on_sale_quantity_changed)
        if self._base_product.get("estimated_profit_detail") or self._base_product.get("estimated_profit_breakdown"):
            self._estimation_cache[int(initial_qty)] = deepcopy(self._base_product)

        close_row = QHBoxLayout()
        close_row.addStretch(1)
        btn_close = QPushButton("Kapat")
        btn_close.setObjectName("PrimaryBtn")
        btn_close.setCursor(Qt.CursorShape.PointingHandCursor)
        btn_close.clicked.connect(self.close)
        close_row.addWidget(btn_close)
        root.addLayout(close_row)
        self._rebuild_body()

    def exec(self) -> int:
        self.showMaximized()
        self.raise_()
        self.activateWindow()
        return 0

    @staticmethod
    def _clear_layout(layout) -> None:
        while layout.count():
            item = layout.takeAt(0)
            widget = item.widget()
            child_layout = item.layout()
            if widget is not None:
                widget.deleteLater()
            elif child_layout is not None:
                ProfitBreakdownDialog._clear_layout(child_layout)

    def _create_side_card(self, title: str) -> tuple[QFrame, QVBoxLayout]:
        card = QFrame(self)
        card.setObjectName("SideCard")
        layout = QVBoxLayout(card)
        layout.setContentsMargins(14, 14, 14, 14)
        layout.setSpacing(10)
        title_label = QLabel(title)
        title_label.setObjectName("SectionTitle")
        body = QVBoxLayout()
        body.setContentsMargins(0, 0, 0, 0)
        body.setSpacing(8)
        layout.addWidget(title_label)
        layout.addLayout(body)
        return card, body

    def _build_side_stat_row(self, label: str, value: str, *, tone: str = "neutral") -> QFrame:
        row = QFrame(self)
        row.setObjectName("SideStatRow")
        layout = QHBoxLayout(row)
        layout.setContentsMargins(12, 10, 10, 10)
        layout.setSpacing(12)
        color_map = {"neutral": "#475569", "ok": "#15803D", "danger": "#B91C1C", "warn": "#B45309"}
        fill_map = {"neutral": "#EEF2F7", "ok": "#DCFCE7", "danger": "#FEE2E2", "warn": "#FEF3C7"}
        accent = color_map.get(tone, "#475569")
        fill = fill_map.get(tone, "#EEF2F7")
        dot = QLabel("")
        dot.setFixedSize(10, 10)
        dot.setStyleSheet(f"background: {accent}; border-radius: 5px;")
        lbl = QLabel(label)
        lbl.setObjectName("SideStatLabel")
        val = QLabel(value)
        val.setObjectName("SideValue")
        val.setAlignment(Qt.AlignmentFlag.AlignCenter)
        val.setMinimumWidth(46)
        val.setStyleSheet(
            f"padding: 5px 10px; border-radius: 11px; background: {fill}; color: {accent}; font-size: 13px; font-weight: 900;"
        )
        layout.addWidget(dot, 0, Qt.AlignmentFlag.AlignVCenter)
        layout.addWidget(lbl, 1)
        layout.addWidget(val, 0)
        return row

    def _build_logic_note(self, text: str) -> QLabel:
        lbl = QLabel(f"• {text}")
        lbl.setObjectName("CardNote")
        lbl.setWordWrap(True)
        return lbl

    @staticmethod
    def _safe_export_filename(label: str, suffix: str) -> str:
        raw = f"{label}_{suffix}".strip() or "tahmini_kar_detayi"
        cleaned = "".join(ch if ch.isalnum() or ch in ("-", "_") else "_" for ch in raw)
        while "__" in cleaned:
            cleaned = cleaned.replace("__", "_")
        return cleaned.strip("_")[:120] or "tahmini_kar_detayi"

    def _toggle_export_drawer(self) -> None:
        self._export_drawer_visible = not self._export_drawer_visible
        self.export_drawer.setVisible(self._export_drawer_visible)

    def _export_to_excel(self) -> None:
        product = self._build_product_for_quantity(max(1, int(self.sale_qty.value())))
        detail = dict(product.get("estimated_profit_detail") or {})
        results = dict(detail.get("results") or {})
        summary = dict(detail.get("summary") or {})
        default_name = self._safe_export_filename(
            _safe_str(self._product_title) or "ilan",
            f"tahmini_kar_{time.strftime('%Y%m%d_%H%M')}.xlsx",
        )
        path, _ = QFileDialog.getSaveFileName(
            self,
            "Excel olarak dışa aktar",
            os.path.join(os.path.expanduser("~"), "Desktop", default_name),
            "Excel Dosyası (*.xlsx)",
        )
        if not path:
            return
        if not path.lower().endswith(".xlsx"):
            path = f"{path}.xlsx"

        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill

        wb = Workbook()
        ws = wb.active
        ws.title = "Ozet"
        rows = [
            ("İlan", self._product_title),
            ("Satış adedi", _safe_str(summary.get("sale_quantity_text")) or str(self.sale_qty.value())),
            ("Toplam satış", _safe_str(summary.get("sale_price_text")) or self._amount_text(results.get("sale_price"))),
            ("Toplam maliyet", _safe_str(results.get("total_costs_text")) or "-"),
            ("Tahmini kar", _safe_str(product.get("estimated_profit_text")) or "-"),
            ("Durum", "Hazır" if product.get("estimated_profit_ready") else "Bekliyor"),
        ]
        ws.append(["Alan", "Değer"])
        for row in rows:
            ws.append(list(row))
        ws["A1"].font = ws["B1"].font = Font(bold=True, color="FFFFFF")
        ws["A1"].fill = ws["B1"].fill = PatternFill("solid", fgColor="0F172A")
        ws.column_dimensions["A"].width = 22
        ws.column_dimensions["B"].width = 80

        sale_ws = wb.create_sheet("Satis")
        sale_ws.append(["İlan", "Adet", "Birim", "Tutar"])
        sale_ws.append([
            self._product_title,
            _safe_str(summary.get("sale_quantity_text")) or str(self.sale_qty.value()),
            _safe_str(summary.get("unit_sale_price_text")) or "-",
            _safe_str(summary.get("sale_price_text")) or self._amount_text(results.get("sale_price")),
        ])

        cost_ws = wb.create_sheet("Maliyetler")
        cost_ws.append(["Kalem", "Açıklama", "Tutar"])
        applied_costs = [row for row in list(detail.get("applied_costs") or []) if isinstance(row, dict)]
        for component in applied_costs:
            amount = self._optional_float(component.get("gross"))
            if amount is None:
                amount = self._optional_float(component.get("value"))
            cost_ws.append([
                _safe_str(component.get("label")) or "Maliyet",
                _safe_str(component.get("note")) or "",
                self._amount_text(amount),
            ])
        cost_ws.append(["Maliyet Toplamı", "", _safe_str(results.get("total_costs_text")) or "-"])
        for sheet in (sale_ws, cost_ws):
            for cell in sheet[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="0F172A")
            sheet.column_dimensions["A"].width = 34
            sheet.column_dimensions["B"].width = 58
            sheet.column_dimensions["C"].width = 20
            if "D" in sheet.column_dimensions:
                sheet.column_dimensions["D"].width = 20
        wb.save(path)
        self.btn_export.setText("Excel hazır")
        self.export_drawer.setVisible(False)
        self._export_drawer_visible = False
        QTimer.singleShot(1400, lambda: self.btn_export.setText("Dışa Aktar"))

    def _toggle_summary_expanded(self) -> None:
        self._summary_expanded = not self._summary_expanded
        if not self._summary_expanded:
            self._active_detail_tabs.clear()
        self._rebuild_body()

    def _toggle_detail_tab(self, key: str) -> None:
        self._capture_detail_scroll_anchor()
        normalized = "sale" if key == "sale" else "cost"
        if not self._summary_expanded:
            self._summary_expanded = True
            self._active_detail_tabs.add(normalized)
            self._rebuild_body()
            return
        if normalized in self._active_detail_tabs:
            self._active_detail_tabs.remove(normalized)
        else:
            self._active_detail_tabs.add(normalized)
        if self._summary_cards:
            self._render_detail_drawers(animated=True)
            return
        self._rebuild_body()

    def _capture_detail_scroll_anchor(self) -> None:
        scroll = getattr(self, "body_scroll", None)
        if not isinstance(scroll, QScrollArea):
            self._detail_scroll_value = None
            return
        self._detail_scroll_value = int(scroll.verticalScrollBar().value())
        content = getattr(self, "_body_content", None)
        if isinstance(content, QWidget) and content.width() > 0:
            locked_width = int(content.width())
            content.setMinimumWidth(locked_width)
            content.setMaximumWidth(locked_width)
            self._detail_viewport_lock_generation += 1
            generation = self._detail_viewport_lock_generation
            QTimer.singleShot(1400, lambda: self._release_detail_viewport_lock(generation))

    def _restore_detail_scroll_anchor(self) -> None:
        scroll_value = getattr(self, "_detail_scroll_value", None)
        scroll = getattr(self, "body_scroll", None)
        if scroll_value is None or not isinstance(scroll, QScrollArea):
            return
        bar = scroll.verticalScrollBar()
        if bar.value() != scroll_value:
            bar.setValue(scroll_value)

    def _position_detail_overlays(self) -> None:
        scroll = getattr(self, "body_scroll", None)
        if not isinstance(scroll, QScrollArea):
            return
        viewport = scroll.viewport()
        for meta in self._summary_cards.values():
            if not isinstance(meta, dict):
                continue
            button = meta.get("button")
            drawer = meta.get("drawer")
            if not isinstance(button, QWidget) or not isinstance(drawer, SmoothHeightFrame):
                continue
            height = max(0, int(drawer.maximumHeight()))
            if height <= 0 and not drawer.isAnimating() and not drawer.isExpanded():
                drawer.setGeometry(0, 0, 0, 0)
                continue
            top_left = button.mapTo(viewport, QPoint(0, button.height() + 8))
            width = max(0, int(button.width()))
            drawer.setGeometry(top_left.x(), top_left.y(), width, height)
            drawer.raise_()

    def _active_detail_drawer_height(self) -> int:
        height = 0
        for key in ("sale", "cost"):
            drawer = self._summary_cards.get(key, {}).get("drawer")
            if not isinstance(drawer, SmoothHeightFrame):
                continue
            if self._summary_expanded and key in self._active_detail_tabs:
                height = max(height, int(drawer.expandedHeight()), int(drawer.maximumHeight()))
                continue
            if drawer.isHidden() and not drawer.isExpanded() and not drawer.isAnimating():
                continue
            if drawer.isAnimating():
                height = max(height, int(drawer.expandedHeight()), int(drawer.maximumHeight()))
                continue
            height = max(height, int(drawer.maximumHeight()))
        return max(0, height)

    def _sync_detail_drawer_spacer(self) -> None:
        spacer = getattr(self, "detail_drawer_spacer", None)
        if not isinstance(spacer, QWidget):
            return
        height = self._active_detail_drawer_height()
        target = height + 24 if height > 0 else 0
        if spacer.maximumHeight() == target:
            return
        spacer.setVisible(target > 0)
        spacer.setMinimumHeight(target)
        spacer.setMaximumHeight(target)
        spacer.updateGeometry()
        content = getattr(self, "_body_content", None)
        if isinstance(content, QWidget):
            layout = content.layout()
            if layout is not None:
                layout.activate()
            content.updateGeometry()
            scroll = getattr(self, "body_scroll", None)
            if isinstance(scroll, QScrollArea):
                target_width = max(1, int(scroll.viewport().width()))
                target_height = max(int(content.height()), int(content.sizeHint().height()))
                content.resize(target_width, target_height)

    def _dispose_detail_drawers(self) -> None:
        for meta in list(self._summary_cards.values()):
            if not isinstance(meta, dict):
                continue
            drawer = meta.get("drawer")
            if not isinstance(drawer, SmoothHeightFrame):
                continue
            try:
                drawer.setExpanded(False, animated=False)
                drawer.setGeometry(0, 0, 0, 0)
                drawer.setParent(None)
                drawer.deleteLater()
            except RuntimeError:
                continue

    def _release_detail_viewport_lock(self, generation: int) -> None:
        if generation != self._detail_viewport_lock_generation:
            return
        content = getattr(self, "_body_content", None)
        try:
            if isinstance(content, QWidget):
                content.setMinimumWidth(0)
                content.setMaximumWidth(16777215)
                content.updateGeometry()
        except RuntimeError:
            return

    def resizeEvent(self, event) -> None:
        super().resizeEvent(event)
        self._position_detail_overlays()

    def _on_sale_quantity_changed(self, *_args) -> None:
        self._detail_page_indexes = {
            "sale": 0,
            "cost": 0,
        }
        self._missing_page_index = 0
        self._rebuild_body()

    def _detail_page_count(self, total_rows: int) -> int:
        per_page = max(1, int(self._DETAIL_ROWS_PER_PAGE))
        return max(1, (max(0, int(total_rows)) + per_page - 1) // per_page)

    def _detail_page_bounds(self, key: str, total_rows: int) -> tuple[int, int, int, int]:
        normalized = "sale" if key == "sale" else "cost"
        per_page = max(1, int(self._DETAIL_ROWS_PER_PAGE))
        page_count = self._detail_page_count(total_rows)
        current = max(0, min(int(self._detail_page_indexes.get(normalized, 0)), page_count - 1))
        self._detail_page_indexes[normalized] = current
        start = current * per_page
        end = start + per_page
        return current, page_count, start, end

    def _change_detail_page(self, key: str, delta: int) -> None:
        normalized = "sale" if key == "sale" else "cost"
        current = max(0, int(self._detail_page_indexes.get(normalized, 0)))
        self._detail_page_indexes[normalized] = current + int(delta or 0)
        self._rebuild_body()

    def _build_detail_pagination_bar(self, key: str, total_rows: int) -> QFrame:
        current, page_count, _start, _end = self._detail_page_bounds(key, total_rows)
        bar = QFrame(self)
        bar.setObjectName("DetailPaginationBar")
        layout = QHBoxLayout(bar)
        layout.setContentsMargins(10, 8, 10, 8)
        layout.setSpacing(12)

        prev_btn = QPushButton("<")
        prev_btn.setObjectName("PageNavButton")
        prev_btn.setEnabled(current > 0)
        prev_btn.clicked.connect(lambda _checked=False, k=key: self._change_detail_page(k, -1))

        page_label = QLabel(f"Sayfa {current + 1} / {page_count}")
        page_label.setObjectName("PageIndicator")
        page_label.setAlignment(Qt.AlignmentFlag.AlignCenter)

        next_btn = QPushButton(">")
        next_btn.setObjectName("PageNavButton")
        next_btn.setEnabled(current < page_count - 1)
        next_btn.clicked.connect(lambda _checked=False, k=key: self._change_detail_page(k, 1))

        layout.addWidget(prev_btn, 0)
        layout.addStretch(1)
        layout.addWidget(page_label, 0)
        layout.addStretch(1)
        layout.addWidget(next_btn, 0)
        return bar

    def _toggle_missing_expanded(self) -> None:
        self._missing_expanded = not self._missing_expanded
        self._rebuild_body()

    def _missing_page_count(self, total_rows: int) -> int:
        per_page = max(1, int(self._MISSING_ROWS_PER_PAGE))
        return max(1, (max(0, int(total_rows)) + per_page - 1) // per_page)

    def _missing_page_bounds(self, total_rows: int) -> tuple[int, int, int, int]:
        per_page = max(1, int(self._MISSING_ROWS_PER_PAGE))
        page_count = self._missing_page_count(total_rows)
        current = max(0, min(int(self._missing_page_index), page_count - 1))
        self._missing_page_index = current
        start = current * per_page
        end = start + per_page
        return current, page_count, start, end

    def _change_missing_page(self, delta: int) -> None:
        self._missing_page_index = max(0, int(self._missing_page_index) + int(delta or 0))
        self._rebuild_body()

    def _build_missing_pagination_bar(self, total_rows: int) -> QFrame:
        current, page_count, _start, _end = self._missing_page_bounds(total_rows)
        bar = QFrame(self)
        bar.setObjectName("DetailPaginationBar")
        layout = QHBoxLayout(bar)
        layout.setContentsMargins(10, 8, 10, 8)
        layout.setSpacing(12)

        prev_btn = QPushButton("<")
        prev_btn.setObjectName("PageNavButton")
        prev_btn.setEnabled(current > 0)
        prev_btn.clicked.connect(lambda _checked=False: self._change_missing_page(-1))

        page_label = QLabel(f"Sayfa {current + 1} / {page_count}")
        page_label.setObjectName("PageIndicator")
        page_label.setAlignment(Qt.AlignmentFlag.AlignCenter)

        next_btn = QPushButton(">")
        next_btn.setObjectName("PageNavButton")
        next_btn.setEnabled(current < page_count - 1)
        next_btn.clicked.connect(lambda _checked=False: self._change_missing_page(1))

        layout.addWidget(prev_btn, 0)
        layout.addStretch(1)
        layout.addWidget(page_label, 0)
        layout.addStretch(1)
        layout.addWidget(next_btn, 0)
        return bar

    def _build_missing_reason_row(self, text: str) -> QFrame:
        row = QFrame(self)
        row.setObjectName("MissingReasonRow")
        layout = QHBoxLayout(row)
        layout.setContentsMargins(12, 10, 12, 10)
        layout.setSpacing(10)

        icon = QLabel("!")
        icon.setAlignment(Qt.AlignmentFlag.AlignCenter)
        icon.setFixedSize(26, 26)
        icon.setStyleSheet(
            "background: #FEF3C7; border: 1px solid #FBBF24; border-radius: 13px; "
            "color: #B45309; font-size: 14px; font-weight: 950;"
        )
        label = QLabel(text)
        label.setObjectName("MissingReasonText")
        label.setWordWrap(True)
        layout.addWidget(icon, 0, Qt.AlignmentFlag.AlignTop)
        layout.addWidget(label, 1)
        return row

    def _build_missing_panel(self, missing_reasons: list[Any]) -> QFrame:
        reasons = [_safe_str(reason) for reason in missing_reasons if _safe_str(reason)]
        panel = ClickableFrame(self)
        panel.setObjectName("MissingPanel")
        panel.clicked.connect(self._toggle_missing_expanded)
        layout = QVBoxLayout(panel)
        layout.setContentsMargins(22, 18, 22, 18)
        layout.setSpacing(14)

        header = QHBoxLayout()
        header.setContentsMargins(0, 0, 0, 0)
        header.setSpacing(18)
        icon = QLabel("!")
        icon.setAlignment(Qt.AlignmentFlag.AlignCenter)
        icon.setFixedSize(64, 64)
        icon.setStyleSheet(
            "background: #FEF3C7; border: 2px solid #F59E0B; border-radius: 32px; "
            "color: #B45309; font-size: 28px; font-weight: 950;"
        )
        text = QVBoxLayout()
        text.setContentsMargins(0, 0, 0, 0)
        text.setSpacing(7)
        title = QLabel("EKSİK / BEKLEYEN VERİLER")
        title.setObjectName("EquationLabel")
        value = QLabel(str(len(reasons)))
        value.setObjectName("MissingValue")
        hint = QLabel(
            "Eksik kayıt detaylarını kapatmak için tıkla"
            if self._missing_expanded
            else "Eksik kayıt detaylarını görmek için tıkla"
        )
        hint.setObjectName("MissingHint")
        hint.setWordWrap(True)
        text.addWidget(title)
        text.addWidget(value)
        text.addWidget(hint)
        chevron = QLabel("^" if self._missing_expanded else "v")
        chevron.setObjectName("MissingChevron")
        chevron.setAlignment(Qt.AlignmentFlag.AlignCenter)
        chevron.setFixedSize(44, 44)
        header.addWidget(icon, 0, Qt.AlignmentFlag.AlignVCenter)
        header.addLayout(text, 1)
        header.addWidget(chevron, 0, Qt.AlignmentFlag.AlignVCenter)
        layout.addLayout(header)

        if self._missing_expanded and reasons:
            _current, page_count, start, end = self._missing_page_bounds(len(reasons))
            for reason in reasons[start:end]:
                layout.addWidget(self._build_missing_reason_row(reason))
            if page_count > 1:
                layout.addWidget(self._build_missing_pagination_bar(len(reasons)))
            note = QLabel(f"Her sayfada {self._MISSING_ROWS_PER_PAGE} kayıt gösterilir.")
            note.setObjectName("DialogSub")
            layout.addWidget(note)
        return panel

    def _metric_card(
        self,
        key: str,
        title: str,
        value: str,
        note: str,
        *,
        tone: str,
        drawer: QFrame | None = None,
    ) -> ClickableFrame:
        accent = "#16A34A" if tone == "ok" else "#DC2626" if tone == "danger" else "#475569"
        active = key in self._active_detail_tabs
        card = ClickableFrame(self)
        card.setObjectName("MetricCard")
        card.setMinimumWidth(0)
        card.setMinimumHeight(142)
        card.setProperty("active", "true" if active else "false")
        card.clicked.connect(lambda: self._toggle_detail_tab(key))
        self._style_metric_card(card, tone=tone, active=active)
        layout = QVBoxLayout(card)
        layout.setContentsMargins(18, 18, 18, 16)
        layout.setSpacing(0)
        head_host = QWidget(card)
        head_host.setFixedHeight(72)
        head = QHBoxLayout(head_host)
        head.setContentsMargins(0, 0, 0, 0)
        head.setSpacing(16)
        icon = ProfitMetricIcon(
            "sale" if key == "sale" else "cost",
            accent,
            "#EAF8EF" if key == "sale" else "#FFF1F2",
            card,
        )
        text = QVBoxLayout()
        text.setContentsMargins(0, 0, 0, 0)
        text.setSpacing(6)
        title_label = QLabel(title)
        title_label.setObjectName("EquationLabel")
        value_label = QLabel(value)
        value_label.setObjectName("EquationValue")
        value_label.setStyleSheet(f"font-size: 34px; font-weight: 950; color: {accent};")
        note_label = QLabel(note)
        note_label.setObjectName("EquationHint")
        note_label.setWordWrap(True)
        text.addWidget(title_label)
        text.addWidget(value_label)
        text.addWidget(note_label)
        chevron = QLabel("v" if active else ">")
        chevron.setStyleSheet("font-size: 28px; font-weight: 900; color: #64748B;")
        head.addWidget(icon, 0, Qt.AlignmentFlag.AlignVCenter)
        head.addLayout(text, 1)
        head.addWidget(chevron, 0, Qt.AlignmentFlag.AlignVCenter)
        layout.addWidget(head_host)
        if drawer is not None:
            if isinstance(drawer, SmoothHeightFrame):
                drawer.setParent(self.body_scroll.viewport())
                drawer.heightProgressChanged.connect(self._position_detail_overlays)
                drawer.heightProgressChanged.connect(self._sync_detail_drawer_spacer)
                drawer.setExpanded(active, animated=False)
            else:
                drawer.setVisible(active)
        layout.addStretch(1)
        self._summary_cards[key] = {
            "button": card,
            "chevron": chevron,
            "drawer": drawer,
            "header": head_host,
            "tone": tone,
        }
        return card

    def _style_metric_card(self, card: QFrame, *, tone: str, active: bool) -> None:
        accent = "#16A34A" if tone == "ok" else "#DC2626" if tone == "danger" else "#475569"
        bg = "#FFFFFF"
        if active:
            bg = "#F0FDF4" if tone == "ok" else "#FFF1F2" if tone == "danger" else "#F8FAFC"
        card.setProperty("active", "true" if active else "false")
        card.setStyleSheet(
            "QFrame#MetricCard {"
            f"background: {bg}; border: 2px solid {accent}; border-radius: 24px;"
            "}"
        )

    def _render_detail_drawers(self, *, animated: bool = False) -> None:
        self._sync_detail_drawer_spacer()
        for key in ("sale", "cost"):
            meta = self._summary_cards.get(key) or {}
            active = self._summary_expanded and key in self._active_detail_tabs
            button = meta.get("button")
            if isinstance(button, QFrame):
                self._style_metric_card(button, tone=str(meta.get("tone") or "neutral"), active=active)
            chevron = meta.get("chevron")
            if isinstance(chevron, QLabel):
                chevron.setText("v" if active else ">")
            drawer = meta.get("drawer")
            if isinstance(drawer, SmoothHeightFrame):
                drawer.setExpanded(active, animated=animated)
            elif isinstance(drawer, QFrame):
                drawer.setVisible(active)
        self._sync_detail_drawer_spacer()
        self._position_detail_overlays()

    def _sync_metric_card_height(self, key: str) -> None:
        meta = self._summary_cards.get(key) or {}
        button = meta.get("button")
        drawer = meta.get("drawer")
        header = meta.get("header")
        if not isinstance(button, QWidget) or not isinstance(drawer, SmoothHeightFrame):
            return
        target = 142
        if not drawer.isHidden() or drawer.isExpanded() or drawer.isAnimating():
            layout = button.layout()
            header_height = int(header.height()) if isinstance(header, QWidget) else 72
            if layout is not None:
                margins = layout.contentsMargins()
                target = max(target, int(margins.top() + margins.bottom() + header_height + drawer.maximumHeight() + 8))
        button.setMinimumHeight(target)
        button.updateGeometry()

    def _ledger_row(
        self,
        title: str,
        value: str,
        *,
        note: str = "",
        tone: str = "neutral",
        summary: bool = False,
        meta: tuple[str, str] | None = None,
    ) -> QFrame:
        row = QFrame(self)
        row.setObjectName("LedgerTotalRow" if summary else "LedgerRow")
        layout = QHBoxLayout(row)
        layout.setContentsMargins(13, 11, 13, 11)
        layout.setSpacing(10)
        color_map = {"neutral": "#0F172A", "ok": "#15803D", "danger": "#B91C1C", "warn": "#B45309"}
        fill_map = {"neutral": "#E2E8F0", "ok": "#BBF7D0", "danger": "#FECACA", "warn": "#FED7AA"}
        accent = color_map.get(tone, "#0F172A")
        if not summary:
            marker = QLabel("")
            marker.setFixedSize(14, 14)
            marker.setStyleSheet(
                f"background: {fill_map.get(tone, '#E2E8F0')}; border: 2px solid {accent}; border-radius: 7px;"
            )
            layout.addWidget(marker, 0, Qt.AlignmentFlag.AlignTop)
        text = QVBoxLayout()
        text.setContentsMargins(0, 0, 0, 0)
        text.setSpacing(4)
        lbl = QLabel(title)
        lbl.setObjectName("CostTitle")
        lbl.setWordWrap(True)
        if summary:
            lbl.setStyleSheet("font-size: 14px; font-weight: 950; color: #0F172A;")
        text.addWidget(lbl)
        if note:
            note_label = QLabel(note)
            note_label.setObjectName("CardNote")
            note_label.setWordWrap(True)
            text.addWidget(note_label)
        if meta:
            chips = QHBoxLayout()
            chips.setContentsMargins(0, 2, 0, 0)
            chips.setSpacing(6)
            for label, val in (("Adet", meta[0]), ("Birim", meta[1])):
                chip = QLabel(f"{label} {val}")
                chip.setStyleSheet(
                    "padding: 4px 8px; border-radius: 9px; border: 1px solid #DDE7F2; "
                    "background: #F8FAFC; color: #475569; font-size: 10px; font-weight: 800;"
                )
                chips.addWidget(chip)
            chips.addStretch(1)
            text.addLayout(chips)
        amount = QLabel(value)
        amount.setObjectName("CostRowValue")
        amount.setStyleSheet(f"font-size: {'16px' if summary else '13px'}; font-weight: 950; color: {accent};")
        layout.addLayout(text, 1)
        layout.addWidget(amount, 0, Qt.AlignmentFlag.AlignVCenter)
        return row

    def _compose_linked_stock_label(self, sale_quantity: int) -> str:
        label_map = self._estimation_context.get("stock_item_label_map") or {}
        if not label_map:
            return _safe_str(self._base_product.get("linked_stock_label"))

        labels: list[str] = []
        for row in (self._base_product.get("stock_items") or []):
            if not isinstance(row, dict):
                continue
            try:
                stock_item_id = int(row.get("stock_item_id"))
            except Exception:
                continue
            label = _safe_str(label_map.get(stock_item_id))
            if not label:
                continue
            try:
                qty_val = float(row.get("qty") or 0) * float(sale_quantity)
            except Exception:
                qty_val = 0.0
            if qty_val > 0:
                labels.append(f"{label} x{qty_val:.1f}")
            else:
                labels.append(label)
        return ", ".join(labels) or _safe_str(self._base_product.get("linked_stock_label"))

    def _build_product_for_quantity(self, sale_quantity: int) -> dict:
        cached = self._estimation_cache.get(int(sale_quantity))
        if cached is not None:
            return deepcopy(cached)
        product = deepcopy(self._base_product)
        product["linked_stock_label"] = self._compose_linked_stock_label(sale_quantity)

        context = dict(self._estimation_context or {})
        if not context:
            return product

        profit_payload = estimate_listing_profit(
            item=product,
            stock_cost_map=context.get("stock_cost_map"),
            stock_vat_rate_map=context.get("stock_vat_rate_map"),
            stock_cost_meta_map=context.get("stock_cost_meta_map"),
            stock_dimension_map=context.get("stock_dimension_map"),
            cost_card_map=context.get("cost_card_map"),
            cargo_tariffs=context.get("cargo_tariffs"),
            platform_code=context.get("platform_code", ""),
            sale_quantity=sale_quantity,
        )
        product.update(profit_payload)
        self._estimation_cache[int(sale_quantity)] = deepcopy(product)
        return product

    @staticmethod
    def _optional_float(value: Any) -> float | None:
        try:
            return float(value)
        except Exception:
            return None

    def _amount_text(self, value: Any) -> str:
        parsed = self._optional_float(value)
        if parsed is None:
            return "-"
        return f"{parsed:.2f} TL"

    def _rebuild_body(self, *_args) -> None:
        body_layout = getattr(self, "_body_layout", None)
        if body_layout is None:
            return
        self._dispose_detail_drawers()
        self._clear_layout(body_layout)
        self._summary_cards = {}

        try:
            sale_quantity = max(1, int(self.sale_qty.value()))
        except Exception:
            sale_quantity = 1

        product = self._build_product_for_quantity(sale_quantity)
        detail = dict(product.get("estimated_profit_detail") or {})
        summary = dict(detail.get("summary") or {})
        applied_costs = [row for row in list(detail.get("applied_costs") or []) if isinstance(row, dict)]
        results = dict(detail.get("results") or {})
        missing_reasons = list(product.get("estimated_profit_missing_reasons") or [])
        ready = bool(product.get("estimated_profit_ready"))

        sale_text = _safe_str(summary.get("sale_price_text")) or self._amount_text(results.get("sale_price"))
        cost_text = _safe_str(results.get("total_costs_text")) or "-"
        profit_text = _safe_str(product.get("estimated_profit_text")) or _safe_str(results.get("estimated_profit_text")) or "-"
        profit_value = self._optional_float(product.get("estimated_profit"))
        profit_color = "#94A3B8"
        if ready and profit_value is not None:
            profit_color = "#B91C1C" if profit_value < 0 else "#15803D"

        if detail:
            hero = ClickableFrame(self)
            hero.setObjectName("ProfitHero")
            hero.setProperty("danger", "true" if ready and profit_value is not None and profit_value < 0 else "false")
            hero.clicked.connect(self._toggle_summary_expanded)
            hero.setMinimumHeight(170)
            hero_layout = QHBoxLayout(hero)
            hero_layout.setContentsMargins(26, 24, 26, 22)
            hero_layout.setSpacing(20)
            profit_icon = ProfitMetricIcon(
                "profit",
                profit_color if profit_color != "#94A3B8" else "#64748B",
                "#FEE2E2" if profit_color == "#B91C1C" else "#DCFCE7" if profit_color == "#15803D" else "#F1F5F9",
                hero,
            )
            text = QVBoxLayout()
            text.setContentsMargins(0, 0, 0, 0)
            text.setSpacing(7)
            label = QLabel("TAHMİNİ KAR")
            label.setObjectName("EquationLabel")
            value = QLabel(profit_text)
            value.setObjectName("ProfitValue")
            value.setStyleSheet(f"font-size: 54px; font-weight: 950; color: {profit_color};")
            formula = QLabel(f"{sale_text} - {cost_text} = {profit_text}")
            formula.setObjectName("FormulaText")
            hint = QLabel("Satış ve maliyet kırılımı açık" if self._summary_expanded else "Detayları açmak için tıkla")
            hint.setObjectName("EquationHint")
            text.addWidget(label)
            text.addWidget(value)
            text.addWidget(formula)
            text.addWidget(hint)
            chevron = QLabel("^" if self._summary_expanded else "v")
            chevron.setAlignment(Qt.AlignmentFlag.AlignCenter)
            chevron.setFixedSize(44, 44)
            chevron.setStyleSheet(
                "border-radius: 22px; background: rgba(255,255,255,0.72); color: #475569; font-size: 24px; font-weight: 900;"
            )
            hero_layout.addWidget(profit_icon, 0, Qt.AlignmentFlag.AlignVCenter)
            hero_layout.addLayout(text, 1)
            hero_layout.addWidget(chevron, 0, Qt.AlignmentFlag.AlignVCenter)
            body_layout.addWidget(hero)

            if self._summary_expanded:
                row = QHBoxLayout()
                row.setContentsMargins(0, 0, 0, 0)
                row.setSpacing(12)
                sale_panel = self._build_sale_ledger(summary, sale_text, sale_quantity)
                cost_panel = self._build_cost_ledger(applied_costs, results)
                sale_card = self._metric_card(
                    "sale",
                    "TOPLAM SATIŞ",
                    sale_text,
                    "Satış satırlarını görmek için tıkla",
                    tone="ok",
                    drawer=sale_panel,
                )
                cost_card = self._metric_card(
                    "cost",
                    "TOPLAM MALİYET",
                    cost_text,
                    "Maliyet kalemlerini görmek için tıkla",
                    tone="danger" if self._optional_float(results.get("total_costs")) else "neutral",
                    drawer=cost_panel,
                )
                operator = QFrame(self)
                operator.setObjectName("OperatorBadge")
                operator.setFixedSize(48, 48)
                op_layout = QVBoxLayout(operator)
                op_layout.setContentsMargins(0, 0, 0, 2)
                op_label = QLabel("-")
                op_label.setAlignment(Qt.AlignmentFlag.AlignCenter)
                op_label.setStyleSheet("font-size: 34px; font-weight: 900; color: #475569;")
                op_layout.addWidget(op_label)
                op_host = QWidget(self)
                op_host.setStyleSheet("background: transparent;")
                op_host.setFixedWidth(64)
                op_host_layout = QVBoxLayout(op_host)
                op_host_layout.setContentsMargins(8, 47, 8, 0)
                op_host_layout.setSpacing(0)
                op_host_layout.addWidget(operator, 0, Qt.AlignmentFlag.AlignHCenter)
                op_host_layout.addStretch(1)
                row.addWidget(sale_card, 1, Qt.AlignmentFlag.AlignTop)
                row.addWidget(op_host, 0)
                row.addWidget(cost_card, 1, Qt.AlignmentFlag.AlignTop)
                body_layout.addLayout(row)
                self.detail_drawer_spacer = QWidget(self)
                self.detail_drawer_spacer.setStyleSheet("background: transparent;")
                self.detail_drawer_spacer.setSizePolicy(QSizePolicy.Policy.Expanding, QSizePolicy.Policy.Fixed)
                self.detail_drawer_spacer.setFixedHeight(0)
                self.detail_drawer_spacer.setVisible(False)
                body_layout.addWidget(self.detail_drawer_spacer)

        else:
            placeholder = QLabel("Bu ilan için henüz hesaplanmış karlılık detayı bulunmuyor.")
            placeholder.setObjectName("DialogSub")
            placeholder.setWordWrap(True)
            body_layout.addWidget(placeholder)

        if missing_reasons:
            body_layout.addWidget(self._build_missing_panel(missing_reasons))

        self._render_side(summary, missing_reasons, ready)
        body_layout.addStretch(1)
        QTimer.singleShot(0, self._position_detail_overlays)

    def _build_sale_ledger(self, summary: dict[str, Any], sale_text: str, sale_quantity: int) -> QFrame:
        panel = SmoothHeightFrame(self.body_scroll.viewport(), duration_ms=467)
        panel.setObjectName("InnerDrawer")
        panel.setProperty("kind", "sale")
        panel.setAttribute(Qt.WidgetAttribute.WA_StyledBackground, True)
        panel.heightProgressChanged.connect(self._restore_detail_scroll_anchor)
        layout = QVBoxLayout(panel)
        layout.setSizeConstraint(QLayout.SizeConstraint.SetNoConstraint)
        layout.setContentsMargins(14, 12, 14, 12)
        layout.setSpacing(8)
        title = _safe_str(self._product_title) or "İlan"
        unit = _safe_str(summary.get("unit_sale_price_text")) or "-"
        qty = _safe_str(summary.get("sale_quantity_text")) or str(sale_quantity)
        sale_rows = [
            {
                "title": title,
                "value": sale_text,
                "note": "İlan satış fiyatı üzerinden hesaplandı.",
                "meta": (qty, unit),
            }
        ]
        _current, page_count, start, end = self._detail_page_bounds("sale", len(sale_rows))
        for row in sale_rows[start:end]:
            layout.addWidget(
                self._ledger_row(
                    _safe_str(row.get("title")) or "İlan",
                    _safe_str(row.get("value")) or "-",
                    note=_safe_str(row.get("note")),
                    tone="ok",
                    meta=row.get("meta"),
                )
            )
        if page_count > 1:
            layout.addWidget(self._build_detail_pagination_bar("sale", len(sale_rows)))
        layout.addWidget(
            self._ledger_row("Satış Toplamı", sale_text, note="Seçilen satış adedinden toplandı.", tone="ok", summary=True)
        )
        return panel

    def _build_cost_ledger(self, applied_costs: list[dict[str, Any]], results: dict[str, Any]) -> QFrame:
        panel = SmoothHeightFrame(self.body_scroll.viewport(), duration_ms=467)
        panel.setObjectName("InnerDrawer")
        panel.setProperty("kind", "cost")
        panel.setAttribute(Qt.WidgetAttribute.WA_StyledBackground, True)
        panel.heightProgressChanged.connect(self._restore_detail_scroll_anchor)
        layout = QVBoxLayout(panel)
        layout.setSizeConstraint(QLayout.SizeConstraint.SetNoConstraint)
        layout.setContentsMargins(14, 12, 14, 12)
        layout.setSpacing(8)
        if not applied_costs:
            total_cost = self._optional_float(results.get("total_costs")) or 0.0
            if total_cost > 0:
                applied_costs = [{"label": "Toplam maliyet", "gross": total_cost}]
            else:
                layout.addWidget(self._ledger_row("Maliyet kaydı bulunamadı", "-", tone="warn"))
        _current, page_count, start, end = self._detail_page_bounds("cost", len(applied_costs))
        for component in applied_costs[start:end]:
            label = _safe_str(component.get("label") or "Maliyet")
            label_key = label.casefold()
            if "stock" in label_key or "stok" in label_key:
                label = "Stok Maliyeti"
                note = "Eşleşen stok kartı maliyeti"
            elif "komisyon" in label_key or "commission" in label_key:
                label = "Komisyon bedeli"
                note = _safe_str(component.get("note")) or "Platform komisyonu"
            elif "kargo" in label_key or "cargo" in label_key:
                label = "Kargo"
                note = _safe_str(component.get("note")) or "Tahmini gönderim maliyeti"
            else:
                note = _safe_str(component.get("note")) or "Tanımlı maliyet kalemi"
            amount = self._optional_float(component.get("gross"))
            if amount is None:
                amount = self._optional_float(component.get("value"))
            if amount is None:
                continue
            layout.addWidget(self._ledger_row(label, self._amount_text(amount), note=note, tone="danger"))
        if page_count > 1:
            layout.addWidget(self._build_detail_pagination_bar("cost", len(applied_costs)))
        layout.addWidget(
            self._ledger_row(
                "Maliyet Toplamı",
                _safe_str(results.get("total_costs_text")) or self._amount_text(self._optional_float(results.get("total_costs"))),
                note="Stok ve tanımlı maliyet kalemlerinden toplandı.",
                tone="danger",
                summary=True,
            )
        )
        return panel

    def _render_side(self, summary: dict[str, Any], missing_reasons: list[Any], ready: bool) -> None:
        self._clear_layout(self.health_body_layout)
        linked = _safe_str(summary.get("linked_stock_label"))
        self.health_body_layout.addWidget(
            self._build_side_stat_row("Stok Eşleşmesi", "Var" if linked and linked != "-" else "Yok", tone="ok" if linked and linked != "-" else "warn")
        )
        self.health_body_layout.addWidget(
            self._build_side_stat_row("Satış Adedi", _safe_str(summary.get("sale_quantity_text")) or str(self.sale_qty.value()), tone="neutral")
        )
        self.health_body_layout.addWidget(
            self._build_side_stat_row("Eksik Kayıt", str(len(missing_reasons)), tone="danger" if missing_reasons else "ok")
        )
        self.health_body_layout.addWidget(
            self._build_side_stat_row("Kar Hesabı", "Hazır" if ready else "Bekliyor", tone="ok" if ready else "warn")
        )

        self._clear_layout(self.logic_body_layout)
        for note in (
            "İlan satış fiyatı satış toplamını oluşturur.",
            "Stok ve maliyet eşleşmeleri maliyeti oluşturur.",
            "Eksik eşleşme varsa kar hesabı beklemeye alınır.",
        ):
            self.logic_body_layout.addWidget(self._build_logic_note(note))
        self.logic_body_layout.addStretch(1)

    def _build_info_card(self, title: str, value: str, *, note: str = "") -> QFrame:
        card = QFrame(self)
        card.setObjectName("InfoCard")
        card.setMinimumHeight(132)
        layout = QVBoxLayout(card)
        layout.setContentsMargins(14, 12, 14, 12)
        layout.setSpacing(6)

        lbl_title = QLabel(title)
        lbl_title.setObjectName("CardTitle")
        lbl_value = QLabel(value)
        lbl_value.setObjectName("CardValue")
        lbl_value.setWordWrap(True)
        layout.addWidget(lbl_title)
        layout.addWidget(lbl_value)
        if note:
            lbl_note = QLabel(note)
            lbl_note.setObjectName("CardNote")
            lbl_note.setWordWrap(True)
            layout.addWidget(lbl_note)
        return card

    def _build_cost_card(self, component: dict) -> QFrame:
        card = QFrame(self)
        card.setObjectName("CostCard")
        card.setMinimumHeight(146)
        layout = QVBoxLayout(card)
        layout.setContentsMargins(14, 12, 14, 12)
        layout.setSpacing(8)

        title = QLabel(_safe_str(component.get("label") or "Maliyet"))
        title.setObjectName("CostTitle")
        layout.addWidget(title)

        note_text = _safe_str(component.get("note"))
        if note_text:
            note = QLabel(note_text)
            note.setObjectName("CardNote")
            note.setWordWrap(True)
            layout.addWidget(note)

        rows = [
            ("Net", _safe_str(component.get("net_text") or "-")),
            ("KDV", _safe_str(component.get("vat_text") or "-")),
            ("Toplam", _safe_str(component.get("gross_text") or "-")),
        ]
        for row_label, row_value in rows:
            row = QHBoxLayout()
            lbl = QLabel(row_label)
            lbl.setObjectName("CostRowLabel")
            val = QLabel(row_value)
            val.setObjectName("CostRowValue")
            val.setAlignment(Qt.AlignmentFlag.AlignRight | Qt.AlignmentFlag.AlignVCenter)
            row.addWidget(lbl, 1)
            row.addWidget(val, 0)
            layout.addLayout(row)
        return card

    def _build_result_card(self, title: str, value: str, *, note: str = "", tone: str = "primary") -> QFrame:
        palette = {
            "primary": ("#DBEAFE", "#BFDBFE", "#1D4ED8"),
            "warning": ("#FFEDD5", "#FED7AA", "#C2410C"),
            "danger": ("#FEE2E2", "#FECACA", "#B91C1C"),
            "success": ("#DCFCE7", "#BBF7D0", "#15803D"),
        }
        bg, border, value_color = palette.get(tone, palette["primary"])

        card = QFrame(self)
        card.setObjectName("InfoCard")
        card.setStyleSheet(
            f"""
            QFrame#InfoCard {{
                background: {bg};
                border: 1px solid {border};
                border-radius: 14px;
            }}
            QLabel#CardTitle {{ font-size: 11px; font-weight: 800; color: #0F172A; }}
            QLabel#CardValue {{ font-size: 26px; font-weight: 900; color: {value_color}; }}
            QLabel#CardNote {{ font-size: 11px; color: #334155; }}
            """
        )
        layout = QVBoxLayout(card)
        layout.setContentsMargins(14, 12, 14, 12)
        layout.setSpacing(6)

        lbl_title = QLabel(title)
        lbl_title.setObjectName("CardTitle")
        lbl_value = QLabel(value)
        lbl_value.setObjectName("CardValue")
        lbl_value.setWordWrap(True)
        layout.addWidget(lbl_title)
        layout.addWidget(lbl_value)
        if note:
            lbl_note = QLabel(note)
            lbl_note.setObjectName("CardNote")
            lbl_note.setWordWrap(True)
            layout.addWidget(lbl_note)
        return card

    def _format_desi_note(self, summary: dict[str, Any]) -> str:
        parts: list[str] = []
        source_text = _safe_str(summary.get("dimensional_weight_source_text"))
        if source_text:
            parts.append(f"Kaynak: {source_text}")
        stock_total_text = _safe_str(summary.get("stock_total_desi_text"))
        if stock_total_text and stock_total_text != "-":
            parts.append(f"Toplam stok desisi: {stock_total_text}")
        if bool(summary.get("stock_desi_is_partial")):
            parts.append("Bazı stok kartlarında ölçü/ağırlık eksik")
        tariff_text = _safe_str(summary.get("cargo_tariff_name"))
        if tariff_text:
            parts.append(f"Tarife: {tariff_text}")
        return "\n".join(parts)


class ListingListWidget(QListView):
    selection_changed = pyqtSignal()
    page_changed = pyqtSignal(int, int, int, int)
    active_row_changed = pyqtSignal()
    card_clicked = pyqtSignal(object, bool)
    selection_button_clicked = pyqtSignal(object, bool)
    blank_area_clicked = pyqtSignal()

    def __init__(
        self,
        parent=None,
        *,
        open_stock_picker_cb=None,
        open_cost_picker_cb=None,
        open_profit_detail_cb=None,
        open_image_cb=None,
        apply_price_cb=None,
    ):
        super().__init__(parent)
        self.setSelectionMode(QListView.SelectionMode.NoSelection)
        self.setMouseTracking(True)
        self.setSpacing(6)
        self.setUniformItemSizes(True)
        self.setResizeMode(QListView.ResizeMode.Adjust)
        self.setVerticalScrollBarPolicy(Qt.ScrollBarPolicy.ScrollBarAsNeeded)
        self.setHorizontalScrollBarPolicy(Qt.ScrollBarPolicy.ScrollBarAlwaysOff)

        self.items: list[dict] = []
        self.filtered_items: list[dict] = []
        self.page_size: int = 20
        self.current_page: int = 1
        self._open_stock_picker_cb = open_stock_picker_cb
        self._open_cost_picker_cb = open_cost_picker_cb
        self._open_profit_detail_cb = open_profit_detail_cb
        self._open_image_cb = open_image_cb
        self._open_price_cb = apply_price_cb

        self._suppress_click_once = False
        self._page_key_to_row: dict[str, int] = {}
        self._model = ListingListModel([], on_selection_changed=self.selection_changed.emit)
        self._delegate = ListingListDelegate(self, on_select_button=self._on_select_button_clicked)
        self.setModel(self._model)
        self.setItemDelegate(self._delegate)
        self._inline_controls: dict[int, InlinePriceControls] = {}
        self._inline_badge_buttons: dict[int, BadgeButton] = {}
        self._inline_cost_badge_buttons: dict[int, BadgeButton] = {}
        self._inline_profit_controls: dict[int, ProfitDetailControl] = {}
        self._inline_state: dict[str, tuple[str, float]] = {}
        self._active_listing_key: str = ""
        self._active_spin_allowed: bool = True
        self._active_border_animation_enabled: bool = False
        self._guide_select_animation_enabled: bool = False
        self._active_border_timer = QTimer(self)
        self._active_border_timer.setInterval(42)
        self._active_border_timer.timeout.connect(self._tick_active_border_animation)
        self.clicked.connect(self._on_row_clicked)
        self._startup_layout_synced = False
        try:
            self.verticalScrollBar().rangeChanged.connect(
                lambda *_: QTimer.singleShot(0, self._position_inline_controls)
            )
        except Exception:
            pass

    @staticmethod
    def _item_key_of(item: dict | None) -> str:
        if not isinstance(item, dict):
            return ""
        account_id = str(item.get("api_account_id") or "").strip()
        for key_name in (
            "id",
            "product_id",
            "productId",
            "platformListingId",
            "platform_listing_id",
            "productContentId",
            "product_content_id",
            "barcode",
            "stockCode",
            "stock_code",
        ):
            raw = str(item.get(key_name) or "").strip()
            if raw:
                if key_name != "id" and account_id:
                    return f"api_account_id:{account_id}|{key_name}:{raw}"
                return f"{key_name}:{raw}"
        return ""

    def _find_row_by_key(self, record_key: str) -> int:
        needle = str(record_key or "").strip()
        if not needle:
            return -1
        row = self._page_key_to_row.get(needle)
        if row is not None:
            return int(row)
        for row in range(self._model.rowCount()):
            current = self._item_key_of(self._model.item_at(row))
            if current == needle:
                return row
        return -1

    def _find_row_by_active_key(self) -> int:
        return self._find_row_by_key(self._active_listing_key)

    def _reapply_active_row(self) -> None:
        row = self._find_row_by_active_key()
        self._model.set_active_row(row)
        self._set_active_border_animation_enabled(row >= 0 and self._active_spin_allowed)

    def set_active_row(self, row: int | None) -> None:
        prev_key = self._active_listing_key
        if row is None or row < 0:
            self.clear_active_item()
            return
        row = int(row)
        item = self._model.item_at(row)
        self._active_listing_key = self._item_key_of(item)
        self._reapply_active_row()
        if self._active_listing_key != prev_key:
            self.active_row_changed.emit()

    def set_active_record_key(self, record_key: str) -> bool:
        row = self._find_row_by_key(record_key)
        if row < 0:
            self.clear_active_item()
            return False
        self.set_active_row(row)
        return True

    def set_active_record(self, record: dict | None) -> bool:
        return self.set_active_record_key(self._item_key_of(record))

    def clear_active_item(self) -> None:
        had_active = bool(self._active_listing_key)
        self._active_listing_key = ""
        self._model.clear_active_row()
        self._set_active_border_animation_enabled(False)
        if had_active:
            self.active_row_changed.emit()

    def has_active_item(self) -> bool:
        return self._find_row_by_active_key() >= 0

    def set_active_spin_allowed(self, allowed: bool) -> None:
        self._active_spin_allowed = bool(allowed)
        self._delegate.set_active_visual_enabled(self._active_spin_allowed)
        row = self._find_row_by_active_key()
        self._set_active_border_animation_enabled(row >= 0 and self._active_spin_allowed)

    def set_select_button_guide_enabled(self, enabled: bool) -> None:
        self._guide_select_animation_enabled = bool(enabled)
        self._delegate.set_guide_select_enabled(self._guide_select_animation_enabled)
        self._sync_border_animation_timer()

    def _set_active_border_animation_enabled(self, enabled: bool) -> None:
        self._active_border_animation_enabled = bool(enabled)
        self._sync_border_animation_timer()

    def _sync_border_animation_timer(self) -> None:
        should_tick = self._active_border_animation_enabled or self._guide_select_animation_enabled
        if should_tick:
            if not self._active_border_timer.isActive():
                self._active_border_timer.start()
        elif self._active_border_timer.isActive():
            self._active_border_timer.stop()

        if not self._active_border_animation_enabled:
            self._delegate.reset_spin_phase()
        if not self._guide_select_animation_enabled:
            self._delegate.reset_guide_phase()
        self._update_animation_rows()

    def _tick_active_border_animation(self) -> None:
        if self._active_border_animation_enabled:
            self._delegate.advance_spin_phase(1.35)
        if self._guide_select_animation_enabled:
            self._delegate.advance_guide_phase(1.35)
        self._update_animation_rows()

    def _update_animation_rows(self) -> None:
        if self._guide_select_animation_enabled:
            self.viewport().update()
            return
        active_row = self._find_row_by_active_key()
        if active_row < 0:
            return
        index = self._model.index(active_row)
        rect = self.visualRect(index)
        if rect.isValid() and not rect.isNull():
            self.viewport().update(rect.adjusted(-2, -2, 2, 2))

    def _on_row_clicked(self, index: QModelIndex) -> None:
        if self._suppress_click_once:
            self._suppress_click_once = False
            return
        if not index.isValid():
            return
        item = self._model.item_at(index.row()) or {}
        same_as_active = self._item_key_of(item) == self._active_listing_key
        self.card_clicked.emit(item, same_as_active)

    def _on_select_button_clicked(self, index: QModelIndex, selected: bool) -> bool:
        self._suppress_click_once = True
        if not index.isValid():
            return False
        self.set_active_row(index.row())
        item = self._model.item_at(index.row()) or {}
        self.selection_button_clicked.emit(item, bool(selected))
        return True

    def set_selected_for_record_key(self, record_key: str, selected: bool) -> bool:
        key = str(record_key or "").strip()
        if not key:
            return False
        matched = False
        for item in (self.filtered_items or []):
            if self._item_key_of(item) != key:
                continue
            item["_selected"] = bool(selected)
            matched = True
        if not matched:
            return False
        self._build_page()
        self.selection_changed.emit()
        return True

    def mouseMoveEvent(self, event):
        super().mouseMoveEvent(event)
        pos = event.position().toPoint()
        idx = self.indexAt(pos)
        if not idx.isValid():
            self.viewport().setCursor(Qt.CursorShape.ArrowCursor)
            return
        self.viewport().setCursor(Qt.CursorShape.PointingHandCursor)

    def mousePressEvent(self, event):
        if event.button() == Qt.MouseButton.LeftButton:
            idx = self.indexAt(event.position().toPoint())
            if not idx.isValid():
                self.clear_active_item()
                self.clearSelection()
                self.blank_area_clicked.emit()
        super().mousePressEvent(event)

    def leaveEvent(self, event):
        self.viewport().setCursor(Qt.CursorShape.ArrowCursor)
        super().leaveEvent(event)

    def set_items(self, items: list[dict]):
        self.items = items if isinstance(items, list) else list(items or [])
        self.filtered_items = self.items
        self.current_page = 1
        self._build_page()

    def refresh_record(self, record: dict | None) -> bool:
        if not isinstance(record, dict):
            return False
        payload = dict(record)
        record_key = self._item_key_of(record)
        if not record_key:
            return False

        matched = False
        collections = (self.items,) if self.filtered_items is self.items else (self.items, self.filtered_items)
        for collection in collections:
            for item in collection:
                if self._item_key_of(item) != record_key:
                    continue
                selected = bool(item.get("_selected"))
                item.clear()
                item.update(dict(payload))
                if selected:
                    item["_selected"] = True
                matched = True

        if not matched:
            return False

        self._build_page()
        return True

    def total_record_count(self) -> int:
        return len(self.filtered_items or [])

    def get_selected_records(self) -> list[dict]:
        selected: list[dict] = []
        for item in (self.filtered_items or []):
            if isinstance(item, dict) and bool(item.get("_selected")):
                selected.append(item)
        return selected

    def select_all_filtered_records(self) -> None:
        for item in (self.filtered_items or []):
            if isinstance(item, dict):
                item["_selected"] = True
        self._build_page()
        self.selection_changed.emit()

    def clear_selected_records(self) -> None:
        for item in (self.filtered_items or []):
            if isinstance(item, dict):
                item["_selected"] = False
        self._build_page()
        self.selection_changed.emit()

    def set_page_size(self, size: int) -> None:
        try:
            page_size = max(1, int(size))
        except Exception:
            page_size = 20
        if page_size == self.page_size:
            return
        self.page_size = page_size
        self.go_to_page(1)

    def apply_filter(self, text: str):
        q = (text or "").strip().casefold()
        if not q:
            self.filtered_items = self.items
        else:
            def _hit(it: dict) -> bool:
                return (
                    q in _safe_str(it.get("title") or it.get("name") or it.get("productName")).casefold()
                    or q in _safe_str(it.get("barcode")).casefold()
                    or q in _safe_str(it.get("stockCode") or it.get("stock_code")).casefold()
                    or q in _safe_str(it.get("brand") or it.get("brandName")).casefold()
                    or q in _safe_str(it.get("productId") or it.get("id")).casefold()
                    or q in _safe_str(it.get("store_name")).casefold()
                )
            self.filtered_items = [it for it in self.items if _hit(it)]

        self.current_page = 1
        self._build_page()

    def get_total_pages(self) -> int:
        total = len(self.filtered_items or [])
        if total <= 0:
            return 1
        return (total + self.page_size - 1) // self.page_size

    def next_page(self):
        self.go_to_page(self.current_page + 1)

    def prev_page(self):
        self.go_to_page(self.current_page - 1)

    def go_to_page(self, page: int):
        total = self.get_total_pages()
        page = max(1, min(page, total))
        self.current_page = page
        self._build_page()

    def get_current_page_items(self) -> list[dict]:
        total = len(self.filtered_items or [])
        if total <= 0:
            return []
        page_size = max(1, int(self.page_size))
        start = (self.current_page - 1) * page_size
        end = min(start + page_size, total)
        return self.filtered_items[start:end]

    def update_icon_for_url(self, url: str, path: str) -> None:
        # update current page only
        for i, it in enumerate(self._model._items):
            if it.get("image_url") == url:
                it["image_path"] = path
                if self._delegate:
                    self._delegate.invalidate_icon(path)
                idx = self._model.index(i)
                self._model.dataChanged.emit(idx, idx, [ListingListModel.IconRole])

    def _build_page(self):
        total = len(self.filtered_items or [])
        if total <= 0:
            self._model.set_items([])
            self._page_key_to_row = {}
            self.clear_active_item()
            self._rebuild_inline_controls()
            self.page_changed.emit(1, 1, 0, 0)
            return

        page_size = max(1, int(self.page_size))
        total_pages = (total + page_size - 1) // page_size
        self.current_page = max(1, min(self.current_page, total_pages))

        start = (self.current_page - 1) * page_size
        end = min(start + page_size, total)
        page_items = self.filtered_items[start:end]
        self._model.set_items(page_items)
        self._page_key_to_row = {
            key: row
            for row, item in enumerate(page_items)
            if (key := self._item_key_of(item))
        }
        self._reapply_active_row()
        self._rebuild_inline_controls()
        self.page_changed.emit(self.current_page, total_pages, total, end - start)
        QTimer.singleShot(120, self.sync_dynamic_row_width)

    def _rebuild_inline_controls(self) -> None:
        for w in self._inline_controls.values():
            w.setParent(None)
            w.deleteLater()
        self._inline_controls = {}
        for w in self._inline_badge_buttons.values():
            w.setParent(None)
            w.deleteLater()
        self._inline_badge_buttons = {}
        for w in self._inline_cost_badge_buttons.values():
            w.setParent(None)
            w.deleteLater()
        self._inline_cost_badge_buttons = {}
        for w in self._inline_profit_controls.values():
            w.setParent(None)
            w.deleteLater()
        self._inline_profit_controls = {}

        for row in range(self._model.rowCount()):
            product = self._model.item_at(row) or {}
            ctrl = InlinePriceControls(self.viewport())
            qty = product.get("quantity") or product.get("stockQuantity") or product.get("stock") or 0
            price = (
                product.get("sale_price")
                or product.get("salePrice")
                or product.get("list_price")
                or product.get("listPrice")
                or 0
            )
            ctrl.set_values(qty, price, currency_code=_listing_currency_code(product))
            ctrl.qty.submitted.connect(lambda _v, r=row: self._apply_inline_for_row(r))
            ctrl.price.submitted.connect(lambda _v, r=row: self._apply_inline_for_row(r))
            self._inline_controls[row] = ctrl
            self._apply_inline_state_for_row(product, ctrl)

            badge = BadgeButton(
                "Ürünler",
                icon_path=os.path.join(ASSETS_DIR, "stok-kartlari.png"),
                badge_count=0,
                badge_text=None,
                font_size=15,
                radius=10,
                base_w=LISTING_BADGE_BTN_W,
                base_h=LISTING_BADGE_BTN_H,
                min_w=LISTING_BADGE_MIN_W,
                min_h=LISTING_BADGE_MIN_H,
                icon_size=30,
                left_w=48,
                badge_r=17,
                badge_font=15,
                badge_offset_x=-2,
                badge_offset_y=-3,
                center_text=False,
                show_divider=False,
                size_scale=1.0,
                parent=self.viewport(),
            )
            badge.setActive(False)
            badge.clicked.connect(lambda _=False, r=row: self._open_products_for_row(r))
            self._sync_product_badge_for_row(product, badge)
            self._inline_badge_buttons[row] = badge

            cost_badge = BadgeButton(
                "Maliyetler",
                icon_path=os.path.join(ASSETS_DIR, "finans-icon.png"),
                badge_count=0,
                badge_text=None,
                font_size=15,
                radius=10,
                base_w=LISTING_BADGE_BTN_W,
                base_h=LISTING_BADGE_BTN_H,
                min_w=LISTING_BADGE_MIN_W,
                min_h=LISTING_BADGE_MIN_H,
                icon_size=30,
                left_w=48,
                badge_r=17,
                badge_font=15,
                badge_offset_x=-2,
                badge_offset_y=-3,
                center_text=False,
                show_divider=False,
                size_scale=1.0,
                parent=self.viewport(),
            )
            cost_badge.setActive(False)
            cost_badge.clicked.connect(lambda _=False, r=row: self._open_costs_for_row(r))
            self._sync_cost_badge_for_row(product, cost_badge)
            self._inline_cost_badge_buttons[row] = cost_badge

            profit_ctrl = ProfitDetailControl(self.viewport())
            profit_ctrl.set_profit(
                product.get("estimated_profit_text") or "-",
                bool(product.get("estimated_profit_ready")),
                product.get("estimated_profit"),
            )
            profit_ctrl.detailRequested.connect(lambda r=row: self._open_profit_detail_for_row(r))
            self._inline_profit_controls[row] = profit_ctrl

        self._position_inline_controls()
        QTimer.singleShot(0, self._position_inline_controls)

    def _apply_inline_state_for_row(self, product: dict, ctrl: InlinePriceControls) -> None:
        barcode = _safe_str(product.get("barcode"))
        if not barcode:
            return
        state = self._inline_state.get(barcode)
        if not state:
            return
        status, expires_at = state
        if time.monotonic() >= expires_at:
            self._inline_state.pop(barcode, None)
            return
        ctrl.set_state(status)

    def _sync_product_badge_for_row(self, product: dict, badge: BadgeButton) -> None:
        stock_items = product.get("stock_items") or []
        count = 0
        if isinstance(stock_items, list):
            for it in stock_items:
                if not isinstance(it, dict):
                    continue
                if it.get("stock_item_id") is None:
                    continue
                count += 1
        if count <= 0:
            badge.setBadgeColor("#EF4444")
            badge.setBadgeCount(0)
            badge.setBadgeText("!")
        else:
            badge.setBadgeColor("#22C55E")
            badge.setBadgeText(None)
            badge.setBadgeCount(count)

    def _sync_cost_badge_for_row(self, product: dict, badge: BadgeButton) -> None:
        cost_items = product.get("cost_items") or []
        count = 0
        if isinstance(cost_items, list):
            for it in cost_items:
                if not isinstance(it, dict):
                    continue
                if not _safe_str(it.get("cost_key") or it.get("costKey")):
                    continue
                count += 1
        if count <= 0:
            badge.setBadgeColor("#EF4444")
            badge.setBadgeCount(0)
            badge.setBadgeText("!")
        else:
            badge.setBadgeColor("#22C55E")
            badge.setBadgeText(None)
            badge.setBadgeCount(count)

    def _open_products_for_row(self, row: int) -> None:
        if not self._open_stock_picker_cb:
            return
        self.set_active_row(row)
        product = self._model.item_at(row) or {}
        if not product:
            return
        QTimer.singleShot(0, lambda p=product: self._open_stock_picker_cb(p))

    def _open_costs_for_row(self, row: int) -> None:
        if not self._open_cost_picker_cb:
            return
        self.set_active_row(row)
        product = self._model.item_at(row) or {}
        if not product:
            return
        QTimer.singleShot(0, lambda p=product: self._open_cost_picker_cb(p))

    def _open_profit_detail_for_row(self, row: int) -> None:
        if not self._open_profit_detail_cb:
            return
        self.set_active_row(row)
        product = self._model.item_at(row) or {}
        if not product:
            return
        QTimer.singleShot(0, lambda p=product: self._open_profit_detail_cb(p))

    def _set_inline_state(self, barcode: str, status: str, duration_ms: int) -> None:
        if not barcode:
            return
        status = (status or "pending").strip().lower()
        expires_at = time.monotonic() + (float(duration_ms) / 1000.0)
        self._inline_state[barcode] = (status, expires_at)

        for row, ctrl in self._inline_controls.items():
            product = self._model.item_at(row) or {}
            if _safe_str(product.get("barcode")) == barcode:
                ctrl.set_state(status)

        def _expire():
            current = self._inline_state.get(barcode)
            if not current or current[1] != expires_at:
                return
            self._inline_state.pop(barcode, None)
            for row, ctrl in self._inline_controls.items():
                product = self._model.item_at(row) or {}
                if _safe_str(product.get("barcode")) == barcode:
                    ctrl.set_state("idle")

        QTimer.singleShot(int(duration_ms), _expire)

    def _apply_inline_for_row(self, row: int) -> None:
        if not self._open_price_cb:
            return
        self.set_active_row(row)
        product = self._model.item_at(row) or {}
        ctrl = self._inline_controls.get(row)
        if not ctrl:
            return
        barcode = _safe_str(product.get("barcode"))
        if barcode:
            self._set_inline_state(barcode, "pending", PENDING_UI_MS)
        else:
            ctrl.set_state("pending")

        def _set_state(state: str) -> None:
            try:
                ctrl.set_state(state)
            except Exception:
                return

        def _reset_state() -> None:
            _set_state("idle")

        def _feedback(result) -> None:
            status = None
            if result is True:
                status = "success"
            elif result is False:
                status = "error"
            elif isinstance(result, str):
                status = result.strip().lower() or None

            if status == "pending":
                delay = PENDING_UI_MS
            else:
                delay = 30000 if status == "success" else 10000
            if barcode:
                self._set_inline_state(barcode, status or "pending", delay)
                return
            _set_state(status or "pending")
            QTimer.singleShot(delay, _reset_state)

        try:
            self._open_price_cb(product, ctrl.qty_value(), ctrl.price_value(), _feedback)
        except TypeError:
            self._open_price_cb(product, ctrl.qty_value(), ctrl.price_value())

    def _position_inline_controls(self) -> None:
        for row in range(self._model.rowCount()):
            idx = self._model.index(row)
            rect = self.visualRect(idx)
            ctrl = self._inline_controls.get(row)
            badge = self._inline_badge_buttons.get(row)
            cost_badge = self._inline_cost_badge_buttons.get(row)
            profit_ctrl = self._inline_profit_controls.get(row)
            if rect.isNull() or rect.height() <= 0:
                if ctrl:
                    ctrl.setVisible(False)
                    ctrl.set_in_view(False)
                if badge:
                    badge.setVisible(False)
                if cost_badge:
                    cost_badge.setVisible(False)
                if profit_ctrl:
                    profit_ctrl.setVisible(False)
                continue

            row_rect = rect.adjusted(6, 4, -6, -4)

            side_rects = _listing_right_side_rects(row_rect, self, hook_delegate=self._delegate)
            cost_btn_rect = side_rects["cost_badge"]
            product_btn_rect = side_rects["product_badge"]
            price_rect = side_rects["price"]
            qty_rect = side_rects["qty"]
            profit_rect = side_rects["profit"]

            if ctrl:
                ctrl.setGeometry(row_rect)
                ctrl.place_controls(
                    qty_rect.translated(-row_rect.left(), -row_rect.top()),
                    price_rect.translated(-row_rect.left(), -row_rect.top()),
                )
                ctrl.set_hook_context(self, row_rect)
                ctrl.setVisible(True)
                ctrl.set_in_view(True)
                ctrl.raise_()

            if badge:
                if product_btn_rect.isNull() or product_btn_rect.width() <= 0 or product_btn_rect.height() <= 0:
                    badge.setVisible(False)
                else:
                    badge.setSizeScale(
                        min(
                            product_btn_rect.width() / LISTING_BADGE_BTN_W,
                            product_btn_rect.height() / LISTING_BADGE_BTN_H,
                        )
                    )
                    badge.setGeometry(product_btn_rect)
                    badge.setVisible(True)
                    badge.raise_()
            if cost_badge:
                if cost_btn_rect.isNull() or cost_btn_rect.width() <= 0 or cost_btn_rect.height() <= 0:
                    cost_badge.setVisible(False)
                else:
                    cost_badge.setSizeScale(
                        min(
                            cost_btn_rect.width() / LISTING_BADGE_BTN_W,
                            cost_btn_rect.height() / LISTING_BADGE_BTN_H,
                        )
                    )
                    cost_badge.setGeometry(cost_btn_rect)
                    cost_badge.setVisible(True)
                    cost_badge.raise_()
            if profit_ctrl:
                profit_ctrl.setGeometry(profit_rect)
                profit_ctrl.setVisible(True)
                profit_ctrl.raise_()

    def scrollContentsBy(self, dx: int, dy: int) -> None:
        super().scrollContentsBy(dx, dy)
        self._position_inline_controls()

    def resizeEvent(self, event):
        super().resizeEvent(event)
        self._position_inline_controls()

    def showEvent(self, event):
        super().showEvent(event)
        if self._startup_layout_synced:
            return
        self._startup_layout_synced = True
        # First paint can occur before final viewport/layout metrics settle.
        # Re-sync shortly after show so badge and select button spacing stays stable.
        QTimer.singleShot(0, self.sync_dynamic_row_width)
        QTimer.singleShot(90, self.sync_dynamic_row_width)

    def sync_dynamic_row_width(self) -> None:
        try:
            self.doItemsLayout()
            self._position_inline_controls()
            self.viewport().update()
        except Exception:
            return


class StockItemRow(QWidget):
    def __init__(
        self,
        stock_items: list[dict],
        *,
        parent: Optional[QWidget] = None,
        row_index: int,
        selected_id: Optional[int] = None,
        qty: float = 1,
        on_change=None,
        on_remove=None,
    ):
        super().__init__(parent)
        self._row_index = row_index
        self._on_change = on_change
        self._on_remove = on_remove
        self._stock_items = list(stock_items or [])
        self._excluded_ids: set[int] = set()
        self._selected_id = selected_id

        outer = QHBoxLayout(self)
        outer.setContentsMargins(0, 0, 0, 0)
        outer.setSpacing(0)

        frame = QFrame(self)
        frame.setObjectName("RowCard")
        layout = QHBoxLayout(frame)
        layout.setContentsMargins(10, 8, 10, 8)
        layout.setSpacing(8)

        self.cb = SwitchSelectWidget(multi=False, all_text="Seçiniz")
        self.cb.setSizePolicy(QSizePolicy.Policy.Expanding, QSizePolicy.Policy.Fixed)
        control_h = max(34, int(self.cb.sizeHint().height()))
        row_h = control_h + 16
        self.cb.setFixedHeight(control_h)

        self._build_combo()

        self.qty = NumberStepper(minimum=1, maximum=10_000_000, step=1)
        self.qty.setValue(int(qty) if qty and qty > 0 else 1)
        self.qty.setFixedWidth(132)
        self.qty.setFixedHeight(min(control_h, 36))
        # tighten inner controls
        try:
            self.qty.input.setAlignment(Qt.AlignmentFlag.AlignCenter)
        except Exception:
            pass

        self.btn_remove = QPushButton("Sil")
        self.btn_remove.setObjectName("GhostBtn")
        self.btn_remove.setFixedHeight(control_h)

        layout.addWidget(self.cb, 1)
        layout.addWidget(self.qty, 0, Qt.AlignmentFlag.AlignVCenter)
        layout.addWidget(self.btn_remove, 0, Qt.AlignmentFlag.AlignVCenter)
        outer.addWidget(frame)

        frame.setFixedHeight(row_h)
        frame.setSizePolicy(QSizePolicy.Policy.Expanding, QSizePolicy.Policy.Fixed)
        self.setFixedHeight(row_h)
        self.setSizePolicy(QSizePolicy.Policy.Expanding, QSizePolicy.Policy.Fixed)

        if self._on_change:
            self.cb.currentIndexChanged.connect(lambda _=0: self._emit_change())
            self.qty.valueChanged.connect(lambda _v=0: self._emit_change())

        if self._on_remove:
            self.btn_remove.clicked.connect(lambda _=False: self._on_remove(self._row_index))

    def _emit_change(self):
        if self._on_change:
            self._on_change(self._row_index, self.value())

    def value(self) -> Optional[dict]:
        stock_id = self.cb.currentData()
        if stock_id is None:
            return None
        qty = int(self.qty.value())
        if qty <= 0:
            return None
        return {"stock_item_id": int(stock_id), "qty": qty}

    def set_excluded_ids(self, ids: set[int]):
        ids_set = {int(v) for v in (ids or []) if v is not None}
        if ids_set == self._excluded_ids:
            return
        self._excluded_ids = ids_set
        self._selected_id = self.cb.currentData()
        self._build_combo()

    def _build_combo(self):
        with QSignalBlocker(self.cb):
            self.cb.clear()
            for it in self._stock_items:
                sid = it.get("id")
                if sid is None:
                    continue
                try:
                    sid_int = int(sid)
                except Exception:
                    continue
                if sid_int in self._excluded_ids and sid_int != self._selected_id:
                    continue
                self.cb.addItem(it.get("label", ""), sid_int)
            if self._selected_id is not None:
                idx = self.cb.findData(self._selected_id)
                if idx >= 0:
                    self.cb.setCurrentIndex(idx)

    def replace_stock_items(self, stock_items: list[dict]) -> None:
        self._stock_items = list(stock_items or [])
        self._selected_id = self.cb.currentData()
        self._build_combo()


class StockMatchDialog(QDialog):
    def __init__(
        self,
        parent: QWidget,
        *,
        product_title: str,
        stock_items: list[dict],
        current_items: list[dict],
        open_stock_cards_cb=None,
    ):
        super().__init__(parent)
        self.setWindowTitle("Ürün Eşleştir")
        self.setFixedSize(640, 560)
        self.selected_items: list[dict] = []
        self._stock_items = stock_items
        self._open_stock_cards_cb = open_stock_cards_cb
        self._rows_data: list[dict] = []
        self._page_size = 6
        self._current_page = 1
        self._row_height: Optional[int] = None
        self._updating_page_size = False
        self._row_widgets: list[tuple[int, StockItemRow]] = []
        self.setStyleSheet("""
        QDialog {
            background: #F8FAFC;
        }
        QLabel#DialogTitle {
            font-size: 14px;
            font-weight: 800;
            color: #0F172A;
        }
        QLabel#DialogSub {
            font-size: 11px;
            color: #64748B;
        }
        QFrame#RowCard {
            background: #FFFFFF;
            border: 1px solid #E5E7EB;
            border-radius: 10px;
        }
        QFrame#RowCard QComboBox, QFrame#RowCard QLineEdit {
            padding: 6px 8px;
            border-radius: 8px;
            border: 1px solid #E5E7EB;
            background: #FFFFFF;
        }
        QFrame#RowCard QToolButton {
            border: 1px solid #E5E7EB;
            border-left: 0px;
            background: #F8FAFC;
        }
        QPushButton#PrimaryBtn {
            padding: 6px 12px;
            border-radius: 9px;
            border: 1px solid #0EA5E9;
            background: #0EA5E9;
            color: white;
            font-weight: 700;
        }
        QPushButton#PrimaryBtn:hover { background: #0284C7; border-color: #0284C7; }
        QPushButton#GhostBtn {
            padding: 6px 10px;
            border-radius: 9px;
            border: 1px solid #E5E7EB;
            background: #FFFFFF;
            font-weight: 600;
        }
        QPushButton#GhostBtn:hover { background: #F1F5F9; }
        """)

        layout = QVBoxLayout(self)
        layout.setContentsMargins(16, 16, 16, 16)
        layout.setSpacing(12)

        title_row = QVBoxLayout()
        lbl_title = QLabel("Ürün Eşleştir")
        lbl_title.setObjectName("DialogTitle")
        lbl_listing = QLabel(f"İlan: {product_title}")
        lbl_listing.setObjectName("DialogSub")
        lbl_listing.setWordWrap(True)
        lbl_info = QLabel(
            "Bu alanda ilanla stok kartlarını eşleştirir ve bu ilana hangi ürünlerin gönderileceğini belirlersin."
        )
        lbl_info.setObjectName("DialogSub")
        lbl_info.setWordWrap(True)
        title_row.addWidget(lbl_title)
        title_row.addWidget(lbl_listing)
        title_row.addWidget(lbl_info)
        layout.addLayout(title_row)

        self.rows_wrap = QWidget(self)
        self.rows_layout = QVBoxLayout(self.rows_wrap)
        self.rows_layout.setContentsMargins(0, 0, 0, 0)
        self.rows_layout.setSpacing(6)
        self.rows_layout.setAlignment(Qt.AlignmentFlag.AlignTop)
        layout.addWidget(self.rows_wrap, 1)

        tool_row = QHBoxLayout()
        self.btn_stock_cards = ActionPulseButton(text="", center_text=True)
        self.btn_stock_cards.setFixedSize(170, 48)
        self.btn_stock_cards.clicked.connect(self._on_primary_action)
        tool_row.addWidget(self.btn_stock_cards)
        tool_row.addStretch()

        self.prev_btn = QPushButton("<")
        self.prev_btn.setObjectName("GhostBtn")
        self.prev_btn.setFixedWidth(36)
        self.prev_btn.clicked.connect(self._prev_page)
        self.next_btn = QPushButton(">")
        self.next_btn.setObjectName("GhostBtn")
        self.next_btn.setFixedWidth(36)
        self.next_btn.clicked.connect(self._next_page)
        self.page_label = QLabel("Sayfa 0/1")
        self.page_label.setObjectName("DialogSub")

        tool_row.addWidget(self.prev_btn)
        tool_row.addWidget(self.page_label)
        tool_row.addWidget(self.next_btn)
        layout.addLayout(tool_row)

        if current_items:
            for it in current_items:
                self._rows_data.append(
                    {"stock_item_id": it.get("stock_item_id"), "qty": it.get("qty", 1)}
                )
        if not self._rows_data:
            self._rows_data.append({"stock_item_id": None, "qty": 1})

        self._refresh_stock_cards_button_text()
        self._render_page()

        action_row = QHBoxLayout()
        self.btn_save = QPushButton("Kaydet")
        self.btn_clear = QPushButton("Temizle")
        self.btn_cancel = QPushButton("İptal")
        self.btn_save.setObjectName("PrimaryBtn")
        self.btn_clear.setObjectName("GhostBtn")
        self.btn_cancel.setObjectName("GhostBtn")
        self.btn_save.clicked.connect(self._on_save)
        self.btn_clear.clicked.connect(self._on_clear)
        self.btn_cancel.clicked.connect(self.reject)
        action_row.addStretch()
        action_row.addWidget(self.btn_clear)
        action_row.addWidget(self.btn_cancel)
        action_row.addWidget(self.btn_save)
        layout.addLayout(action_row)

        QTimer.singleShot(0, self._refresh_deferred_state)

    def _refresh_deferred_state(self) -> None:
        self._refresh_page_size()
        self._refresh_action_state()

    def _refresh_stock_cards_button_text(self) -> None:
        if hasattr(self, "btn_stock_cards"):
            self.btn_stock_cards.setText("Ekle" if self._stock_items else "Stok Kartı Ekle")

    def _open_stock_cards(self):
        if self._open_stock_cards_cb:
            self._open_stock_cards_cb()

    def _add_row(self):
        self._rows_data.append({"stock_item_id": None, "qty": 1})
        self._current_page = self._get_total_pages()
        self._render_page()
        self._refresh_action_state()

    def _on_primary_action(self):
        if self._stock_items:
            if self._can_add_row():
                self._add_row()
            return
        self._open_stock_cards()

    def _remove_row(self, idx: int):
        if 0 <= idx < len(self._rows_data):
            self._rows_data = _remove_stock_match_row_at(self._rows_data, idx)
            self._current_page = max(1, min(self._current_page, self._get_total_pages()))
            self._render_page()
            self._refresh_action_state()

    def _on_row_change(self, idx: int, value: Optional[dict]):
        if 0 <= idx < len(self._rows_data):
            prev_id = self._rows_data[idx].get("stock_item_id")
            if value is None:
                self._rows_data[idx] = {"stock_item_id": None, "qty": 0}
            else:
                self._rows_data[idx] = value
            new_id = self._rows_data[idx].get("stock_item_id")
            if prev_id != new_id:
                self._refresh_exclusions()
        self._refresh_action_state()

    def _has_missing_selection(self) -> bool:
        for r in (self._rows_data or []):
            if not isinstance(r, dict):
                continue
            if r.get("stock_item_id") is None:
                return True
        return False

    def _can_add_row(self) -> bool:
        if not self._rows_data:
            return True
        return not self._has_missing_selection()

    def _refresh_action_state(self) -> None:
        if hasattr(self, "btn_stock_cards"):
            self.btn_stock_cards.setEnabled(True if not self._stock_items else self._can_add_row())
        if hasattr(self, "btn_save"):
            self.btn_save.setEnabled(not self._has_missing_selection())

    def _collect_items(self) -> list[dict]:
        items: list[dict] = []
        for r in self._rows_data:
            if not isinstance(r, dict):
                continue
            sid = r.get("stock_item_id")
            qty = r.get("qty")
            if sid is None:
                continue
            try:
                qty_val = float(qty)
            except Exception:
                qty_val = 0
            if qty_val <= 0:
                continue
            items.append({"stock_item_id": int(sid), "qty": qty_val})
        return items

    def _sync_rows_from_ui(self) -> None:
        for idx, row in self._row_widgets:
            value = row.value()
            if value is None:
                self._rows_data[idx] = {"stock_item_id": None, "qty": 0}
            else:
                self._rows_data[idx] = value

    def _get_total_pages(self) -> int:
        total = len(self._rows_data)
        if total <= 0:
            return 1
        return (total + self._page_size - 1) // self._page_size

    def _get_selected_ids(self, *, exclude_idx: Optional[int] = None) -> set[int]:
        ids: set[int] = set()
        for i, r in enumerate(self._rows_data):
            if exclude_idx is not None and i == exclude_idx:
                continue
            if not isinstance(r, dict):
                continue
            sid = r.get("stock_item_id")
            if sid is None:
                continue
            try:
                ids.add(int(sid))
            except Exception:
                continue
        return ids

    def _refresh_exclusions(self) -> None:
        for idx, row in self._row_widgets:
            row.set_excluded_ids(self._get_selected_ids(exclude_idx=idx))

    def _measure_row_height(self) -> int:
        if self._row_height:
            return self._row_height
        self._row_height = 50
        return self._row_height

    def _calc_page_size(self) -> int:
        row_h = self._measure_row_height()
        available = int(self.rows_wrap.height())
        if available <= 0:
            available = int(self.height() - 300)
        available = max(row_h, available)
        spacing = int(self.rows_layout.spacing())
        per = max(1, (available + spacing) // max(1, (row_h + spacing)))
        return max(1, int(per))

    def _refresh_page_size(self) -> None:
        if self._updating_page_size:
            return
        self._updating_page_size = True
        try:
            new_size = self._calc_page_size()
            if new_size != self._page_size:
                self._page_size = new_size
                self._current_page = max(1, min(self._current_page, self._get_total_pages()))
                self._render_page()
        finally:
            self._updating_page_size = False

    def _render_page(self):
        # clear existing rows safely (widgets + spacers)
        while self.rows_layout.count():
            item = self.rows_layout.takeAt(0)
            if item is None:
                continue
            w = item.widget()
            if w is not None:
                w.setParent(None)
                w.deleteLater()

        total_pages = self._get_total_pages()
        self._current_page = max(1, min(self._current_page, total_pages))

        start = (self._current_page - 1) * self._page_size
        end = min(start + self._page_size, len(self._rows_data))

        self._row_widgets = []
        for idx in range(start, end):
            row_data = self._rows_data[idx]
            row = StockItemRow(
                self._stock_items,
                parent=self.rows_wrap,
                row_index=idx,
                selected_id=row_data.get("stock_item_id"),
                qty=row_data.get("qty", 1),
                on_change=self._on_row_change,
                on_remove=self._remove_row,
            )
            self.rows_layout.addWidget(row)
            self._row_widgets.append((idx, row))
            row.set_excluded_ids(self._get_selected_ids(exclude_idx=idx))

        self.page_label.setText(f"Sayfa {self._current_page}/{total_pages}")
        self.prev_btn.setEnabled(self._current_page > 1)
        self.next_btn.setEnabled(self._current_page < total_pages)
        self._refresh_action_state()

    def replace_stock_items(self, stock_items: list[dict]) -> None:
        self._stock_items = list(stock_items or [])
        self._refresh_stock_cards_button_text()
        for idx, row in self._row_widgets:
            row.replace_stock_items(self._stock_items)
            row.set_excluded_ids(self._get_selected_ids(exclude_idx=idx))
        self._refresh_action_state()

    def resizeEvent(self, event):
        super().resizeEvent(event)
        self._refresh_page_size()

    def _prev_page(self):
        self._current_page = max(1, self._current_page - 1)
        self._render_page()

    def _next_page(self):
        self._current_page = min(self._get_total_pages(), self._current_page + 1)
        self._render_page()

    def _on_save(self):
        self._sync_rows_from_ui()
        if self._has_missing_selection():
            return
        self.selected_items = self._collect_items()
        self.accept()

    def _on_clear(self):
        self.selected_items = []
        self.accept()


class CostItemRow(QWidget):
    def __init__(
        self,
        cost_items: list[dict],
        *,
        cargo_providers: Optional[list[dict[str, Any]]] = None,
        parent: Optional[QWidget] = None,
        row_index: int,
        selected_key: Optional[str] = None,
        calculation_type: Optional[str] = None,
        source_type: Optional[str] = None,
        source_config: Optional[dict[str, Any]] = None,
        field_options: Optional[list[dict[str, str]]] = None,
        amount: Optional[float] = None,
        quantity: Optional[float] = None,
        on_change=None,
        on_remove=None,
    ):
        super().__init__(parent)
        self._row_index = row_index
        self._on_change = on_change
        self._on_remove = on_remove
        self._cost_items = list(cost_items or [])
        self._excluded_keys: set[str] = set()
        self._selected_key = _safe_str(selected_key)
        self._initial_match_mode = _safe_str(calculation_type) or "manual"
        self._initial_source_type = _safe_str(source_type)
        self._initial_source_config = dict(source_config or {}) or None
        self._source_type_override = self._initial_source_type
        self._source_config_override = dict(self._initial_source_config or {}) or None
        self._field_options = list(field_options or [])
        self._cargo_providers = list(cargo_providers or [])
        self._field_option_map: dict[str, dict[str, str]] = {
            _safe_str(opt.get("path")): dict(opt)
            for opt in self._field_options
            if isinstance(opt, dict) and _safe_str(opt.get("path"))
        }
        self._cargo_provider_map: dict[int, dict[str, Any]] = {}
        for option in self._cargo_providers:
            provider_pk = int(option.get("pk") or 0)
            if provider_pk > 0:
                self._cargo_provider_map[provider_pk] = dict(option)
        self._last_db_field_config = (
            dict(self._initial_source_config or {})
            if normalize_cost_source_type(self._initial_source_type, fallback="") == "db_field"
            else None
        )

        outer = QHBoxLayout(self)
        outer.setContentsMargins(0, 0, 0, 0)
        outer.setSpacing(0)

        frame = QFrame(self)
        frame.setObjectName("RowCard")
        layout = QGridLayout(frame)
        layout.setContentsMargins(12, 10, 12, 10)
        layout.setHorizontalSpacing(12)
        layout.setVerticalSpacing(0)

        self.cb = SwitchSelectWidget(frame, multi=False, all_text="Seçiniz")
        self.cb.setSizePolicy(QSizePolicy.Policy.Expanding, QSizePolicy.Policy.Fixed)
        control_h = max(34, int(self.cb.sizeHint().height()))
        row_h = max(COST_MATCH_ROW_HEIGHT + 30, control_h + 44)
        self.cb.setFixedHeight(control_h)

        self.mode_combo = SwitchSelectWidget(frame, multi=False, include_all=False, empty_text="Seçiniz")
        self.mode_combo.setPlaceholderText("Kaynak tipi")
        self.mode_combo.setSizePolicy(QSizePolicy.Policy.Expanding, QSizePolicy.Policy.Fixed)
        self.mode_combo.setMinimumWidth(168)
        self.mode_combo.setFixedHeight(control_h)
        self.mode_combo.setToolTip(COST_MATCH_SOURCE_TOOLTIP)

        self.amount = CurrencyAmountInput(
            frame,
            decimals=6,
            currency="TRY",
            placeholder="Tutar",
            badge_min_width=60,
            badge_max_width=72,
            spacing=8,
        )
        self.amount.setSizePolicy(QSizePolicy.Policy.Expanding, QSizePolicy.Policy.Fixed)
        self.amount.setMinimumWidth(212)
        self.amount.setFixedHeight(control_h + 12)

        self.percent_input = CircularPercentageInput(
            frame,
            decimals=4,
            diameter=COST_MATCH_PERCENT_DIAMETER,
        )
        self.percent_input.setPlaceholderText("0")
        self.percent_input.setVisible(False)

        self.source_field = SwitchSelectWidget(frame, multi=False, include_all=False, empty_text="Seçiniz", allow_empty=True)
        self.source_field.setPlaceholderText("Alan ara/seç")
        self.source_field.setSizePolicy(QSizePolicy.Policy.Expanding, QSizePolicy.Policy.Fixed)
        self.source_field.setMinimumWidth(204)
        self.source_field.setFixedHeight(control_h)
        self.source_field.setToolTip(COST_MATCH_DB_FIELD_TOOLTIP)

        self.cargo_provider_combo = SwitchSelectWidget(frame, multi=False, include_all=False, empty_text="Seçiniz", allow_empty=True)
        self.cargo_provider_combo.setPlaceholderText(COST_MATCH_CARGO_PROVIDER_PLACEHOLDER)
        self.cargo_provider_combo.setSizePolicy(QSizePolicy.Policy.Expanding, QSizePolicy.Policy.Fixed)
        self.cargo_provider_combo.setMinimumWidth(204)
        self.cargo_provider_combo.setFixedHeight(control_h)
        self.cargo_provider_combo.setToolTip(COST_MATCH_CARGO_PROVIDER_TOOLTIP)

        self._build_combo()
        self._refresh_mode_combo()
        self._refresh_source_field_combo()
        self._refresh_cargo_provider_combo()
        if amount not in (None, ""):
            try:
                self.amount.setValue(float(amount))
                self.percent_input.setValue(float(amount))
            except Exception:
                self.amount.clear()
                self.percent_input.clear()
        initial_field_path = _safe_str((self._source_config_override or {}).get("fieldPath"))
        if initial_field_path:
            self._set_source_field_value(initial_field_path)

        self.btn_remove = QPushButton("Sil")
        self.btn_remove.setObjectName("GhostBtn")
        self.btn_remove.setFixedSize(68, control_h + 8)
        self.btn_remove.setSizePolicy(QSizePolicy.Policy.Fixed, QSizePolicy.Policy.Fixed)

        self.cost_host = self._wrap_labeled_control(frame, "Maliyet Kartı", self.cb)
        self.mode_host = self._wrap_labeled_control(frame, "Ücret Hesaplama Türü", self.mode_combo)
        self.amount_host = self._wrap_labeled_control(frame, "Ücret", self.amount, self.percent_input)
        self.source_field_host = self._wrap_labeled_control(frame, "Veri Alanı", self.source_field)
        self.cargo_provider_host = self._wrap_labeled_control(frame, "Kargo Firması", self.cargo_provider_combo)
        for host in (
            self.cost_host,
            self.mode_host,
            self.amount_host,
            self.source_field_host,
            self.cargo_provider_host,
        ):
            host.setSizePolicy(QSizePolicy.Policy.Expanding, QSizePolicy.Policy.Fixed)
        self.cost_host.setMinimumWidth(220)
        self.mode_host.setMinimumWidth(168)
        self.amount_host.setMinimumWidth(212)
        self.source_field_host.setMinimumWidth(204)
        self.cargo_provider_host.setMinimumWidth(204)

        layout.addWidget(self.cost_host, 0, 0)
        layout.addWidget(self.mode_host, 0, 1)
        layout.addWidget(self.amount_host, 0, 2)
        layout.addWidget(self.source_field_host, 0, 3)
        layout.addWidget(self.cargo_provider_host, 0, 4)
        layout.addWidget(self.btn_remove, 0, 5, Qt.AlignmentFlag.AlignBottom)
        layout.setColumnStretch(0, 26)
        layout.setColumnStretch(1, 18)
        layout.setColumnStretch(2, 18)
        layout.setColumnStretch(3, 18)
        layout.setColumnStretch(4, 18)
        layout.setColumnStretch(5, 0)
        outer.addWidget(frame)

        frame.setFixedHeight(row_h)
        frame.setSizePolicy(QSizePolicy.Policy.Expanding, QSizePolicy.Policy.Fixed)
        self.setFixedHeight(row_h)
        self.setSizePolicy(QSizePolicy.Policy.Expanding, QSizePolicy.Policy.Fixed)

        self.cb.currentIndexChanged.connect(lambda _=0: self._on_select_changed())
        self.mode_combo.currentIndexChanged.connect(lambda _=0: self._on_mode_changed())
        if self._on_change:
            self.amount.textChanged.connect(lambda _v="": self._emit_change())
            self.percent_input.textChanged.connect(lambda _v="": self._emit_change())
            self.source_field.currentIndexChanged.connect(lambda _=0: self._on_source_field_changed())
            self.cargo_provider_combo.currentIndexChanged.connect(lambda _=0: self._on_cargo_provider_changed())
        if self._on_remove:
            self.btn_remove.clicked.connect(lambda _=False: self._on_remove(self._row_index))
        self._sync_dynamic_inputs(preserve_values=True)

    @staticmethod
    def _wrap_labeled_control(parent: QWidget, label_text: str, *widgets: QWidget) -> QWidget:
        host = QWidget(parent)
        host_layout = QVBoxLayout(host)
        host_layout.setContentsMargins(0, 0, 0, 0)
        host_layout.setSpacing(4)
        label = QLabel(label_text, host)
        label.setObjectName("DialogSub")
        host_layout.addWidget(label)
        for widget in widgets:
            host_layout.addWidget(widget)
        return host

    def _source_field_items(self) -> list[tuple[str, str]]:
        items: list[tuple[str, str]] = []
        for option in self._field_options:
            if not isinstance(option, dict):
                continue
            path = _safe_str(option.get("path"))
            label = _safe_str(option.get("label") or path)
            if not path or not label:
                continue
            items.append((label, path))
        current_path = _safe_str((self._source_config_override or {}).get("fieldPath"))
        if current_path and current_path not in {path for _, path in items}:
            current_note = _safe_str((self._source_config_override or {}).get("label")) or current_path
            items.append((f"{current_note} ({current_path})" if current_note != current_path else current_path, current_path))
        return items

    def _refresh_source_field_combo(self) -> None:
        selected_path = _safe_str((self._source_config_override or {}).get("fieldPath"))
        with QSignalBlocker(self.source_field):
            self.source_field.clear()
            for label, path in self._source_field_items():
                self.source_field.addItem(label, path)
            idx = self.source_field.findData(selected_path)
            self.source_field.setCurrentIndex(idx if idx >= 0 else -1)

    def _set_source_field_value(self, field_path: str) -> None:
        self._refresh_source_field_combo()
        idx = self.source_field.findData(field_path)
        with QSignalBlocker(self.source_field):
            self.source_field.setCurrentIndex(idx if idx >= 0 else -1)

    def _selected_requires_cargo_provider(self) -> bool:
        selected = self._lookup_selected_cost_item() or {}
        cost_key = _safe_str(selected.get("costKey")).casefold()
        if cost_key == "kargo":
            return True
        default_config = dict(selected.get("defaultSourceConfig") or {})
        resolver_key = _safe_str(
            default_config.get("ruleKey")
            or default_config.get("resolverKey")
            or default_config.get("resolver_key")
        ).casefold()
        return resolver_key == "cargo_amount_matcher"

    def _selected_cargo_provider_id(self) -> Optional[int]:
        payload = dict(self._source_config_override or {})
        try:
            parsed = int(payload.get("cargoProviderId") or payload.get("cargo_provider_id") or 0)
        except Exception:
            return None
        return parsed if parsed > 0 else None

    def _refresh_cargo_provider_combo(self) -> None:
        selected_provider_id = self._selected_cargo_provider_id()
        with QSignalBlocker(self.cargo_provider_combo):
            self.cargo_provider_combo.clear()
            for option in self._cargo_providers:
                provider_pk = int(option.get("pk") or 0)
                provider_name = _safe_str(option.get("cargoProviderName"))
                provider_code = _safe_str(option.get("cargoProviderCode")).upper()
                if provider_pk <= 0 or not provider_name:
                    continue
                label = provider_name if not provider_code else f"{provider_name} ({provider_code})"
                self.cargo_provider_combo.addItem(label, provider_pk)
            idx = self.cargo_provider_combo.findData(selected_provider_id)
            self.cargo_provider_combo.setCurrentIndex(idx if idx >= 0 else -1)

    def _lookup_selected_cost_item(self) -> Optional[dict]:
        selected_key = _safe_str(self.cb.currentData())
        if not selected_key:
            return None
        for it in self._cost_items:
            if _safe_str(it.get("costKey")) == selected_key:
                return it
        return None

    def _selected_value_type(self) -> str:
        selected = self._lookup_selected_cost_item()
        return _safe_str((selected or {}).get("valueType") or (selected or {}).get("calculationType")) or "numeric"

    def _selected_currency_code(self) -> str:
        return _cost_card_currency(self._lookup_selected_cost_item())

    def _refresh_amount_currency(self, *, visible: bool) -> None:
        currency_code = self._selected_currency_code() if visible else "TRY"
        self.amount.setCurrency(currency_code)

    def _selected_default_match_mode(self) -> str:
        selected = self._lookup_selected_cost_item()
        default_source_type = normalize_cost_source_type(
            (selected or {}).get("defaultSourceType") or (selected or {}).get("defaultMatchMode"),
            fallback="manual_value",
        )
        mode = "automatic" if default_source_type == "builtin_resolver" else "manual"
        return mode if mode in COST_MATCH_MODES else "manual"

    def _current_match_mode(self) -> str:
        mode = "automatic" if self._current_source_type() == "builtin_resolver" else "manual"
        return mode if mode in COST_MATCH_MODES else "manual"

    def _selected_default_source_type(self) -> str:
        selected = self._lookup_selected_cost_item()
        return normalize_cost_source_type(
            (selected or {}).get("defaultSourceType") or (selected or {}).get("defaultMatchMode"),
            fallback="manual_value",
        )

    def _selected_default_source_config(self) -> Optional[dict[str, Any]]:
        selected = self._lookup_selected_cost_item()
        payload = dict((selected or {}).get("defaultSourceConfig") or {})
        return payload or None

    def _current_source_type(self) -> str:
        explicit = _safe_str(self.mode_combo.currentData()) or _safe_str(self._source_type_override)
        if explicit:
            return normalize_cost_source_type(explicit, fallback="manual_value")
        return self._selected_default_source_type()

    def _refresh_mode_combo(self) -> None:
        selected = self._lookup_selected_cost_item() or {}
        supported = list(selected.get("supportedSourceTypes") or [])
        options: list[tuple[str, str]] = []
        for key in supported:
            source_type = normalize_cost_source_type(key, fallback="")
            if not source_type:
                continue
            source_meta = _finance_cost_source_definitions().get(source_type) or {}
            if not bool(source_meta.get("enabled", False)):
                continue
            options.append((cost_source_type_label(source_type), source_type))
        if not options:
            options = [(cost_source_type_label("manual_value"), "manual_value")]
        source_type = self._current_source_type()
        with QSignalBlocker(self.mode_combo):
            self.mode_combo.clear()
            for label, key in options:
                self.mode_combo.addItem(label, key)
            idx = self.mode_combo.findData(source_type)
            self.mode_combo.setCurrentIndex(idx if idx >= 0 else 0)
            self.mode_combo.setEnabled(len(options) > 1)

    def _on_select_changed(self) -> None:
        self._selected_key = _safe_str(self.cb.currentData())
        selected = self._lookup_selected_cost_item()
        default_amount = None if selected is None else selected.get("defaultAmount")
        self._source_type_override = self._selected_default_source_type()
        self._source_config_override = self._selected_default_source_config()
        if self._source_type_override == "db_field":
            self._last_db_field_config = dict(self._source_config_override or {}) or None
        self._refresh_mode_combo()
        field_path = _safe_str((self._source_config_override or {}).get("fieldPath"))
        self._set_source_field_value(field_path)
        self._refresh_cargo_provider_combo()
        if self._current_source_type() == "manual_value":
            if default_amount not in (None, ""):
                try:
                    self.amount.setValue(float(default_amount))
                    self.percent_input.setValue(float(default_amount))
                except Exception:
                    self.amount.clear()
                    self.percent_input.clear()
            else:
                self.amount.clear()
                self.percent_input.clear()
        self._sync_dynamic_inputs(preserve_values=False)
        self._emit_change()

    def _on_mode_changed(self) -> None:
        selected_source_type = _safe_str(self.mode_combo.currentData()) or self._selected_default_source_type()
        previous_source_type = normalize_cost_source_type(self._source_type_override, fallback="")
        previous_source_config = dict(self._source_config_override or {})
        self._source_type_override = selected_source_type
        if selected_source_type == "manual_value":
            self._source_config_override = None
        elif selected_source_type == "builtin_resolver":
            if selected_source_type == self._selected_default_source_type():
                self._source_config_override = self._selected_default_source_config()
            elif previous_source_type != "builtin_resolver":
                self._source_config_override = None
        elif selected_source_type == "db_field":
            existing = dict(self._last_db_field_config or {}) if self._last_db_field_config else {}
            if previous_source_type == "db_field":
                existing = dict(self._source_config_override or {}) or existing
            default_config = self._selected_default_source_config() if selected_source_type == self._selected_default_source_type() else None
            next_config = default_config or existing or {}
            next_config["fieldPath"] = _safe_str(next_config.get("fieldPath"))
            self._source_config_override = next_config
            self._last_db_field_config = dict(next_config) or None
        else:
            self._source_config_override = self._selected_default_source_config() if selected_source_type == self._selected_default_source_type() else None
        if self._selected_requires_cargo_provider():
            provider_payload = {
                key: previous_source_config.get(key)
                for key in ("cargoProviderId", "cargoProviderCode", "cargoProviderName")
                if previous_source_config.get(key) not in (None, "")
            }
            if provider_payload:
                merged_payload = dict(self._source_config_override or {})
                merged_payload.update(provider_payload)
                self._source_config_override = merged_payload or None
        field_path = _safe_str((self._source_config_override or {}).get("fieldPath"))
        self._set_source_field_value(field_path)
        self._refresh_cargo_provider_combo()
        self._sync_dynamic_inputs(preserve_values=False)
        self._emit_change()

    def _on_source_field_changed(self) -> None:
        if self._current_source_type() != "db_field":
            return
        field_path = _safe_str(self.source_field.currentData())
        payload = dict(self._source_config_override or {})
        if field_path:
            payload["fieldPath"] = field_path
            note_label = _safe_str((self._field_option_map.get(field_path) or {}).get("noteLabel"))
            if note_label:
                payload["label"] = note_label
        else:
            payload.pop("fieldPath", None)
            payload.pop("label", None)
        self._source_config_override = payload or None
        self._last_db_field_config = dict(self._source_config_override or {}) or None
        self._emit_change()

    def _on_cargo_provider_changed(self) -> None:
        payload = dict(self._source_config_override or {})
        provider_pk = int(self.cargo_provider_combo.currentData() or 0)
        provider_row = self._cargo_provider_map.get(provider_pk)
        if provider_pk > 0 and provider_row:
            payload["cargoProviderId"] = provider_pk
            payload["cargoProviderCode"] = _safe_str(provider_row.get("cargoProviderCode")).upper()
            payload["cargoProviderName"] = _safe_str(provider_row.get("cargoProviderName"))
        else:
            payload.pop("cargoProviderId", None)
            payload.pop("cargoProviderCode", None)
            payload.pop("cargoProviderName", None)
        self._source_config_override = payload or None
        self._emit_change()

    def _emit_change(self) -> None:
        if self._on_change:
            self._on_change(self._row_index, self.state())

    def _sync_dynamic_inputs(self, *, preserve_values: bool) -> None:
        value_type = self._selected_value_type()
        source_type = self._current_source_type()
        show_amount = source_type == "manual_value"
        show_percent = show_amount and value_type == "percentage"
        show_numeric_amount = show_amount and not show_percent
        show_source_field = source_type == "db_field"
        show_cargo_provider = self._selected_requires_cargo_provider()

        self.percent_input.setVisible(show_percent)
        self.amount.setVisible(show_numeric_amount)
        self.amount.setPlaceholderText(_cost_amount_placeholder(value_type))
        self._refresh_amount_currency(visible=show_numeric_amount)
        if hasattr(self, "amount_host"):
            self.amount_host.setVisible(show_amount)
        self.source_field.setVisible(show_source_field)
        self.source_field.setPlaceholderText(_cost_source_field_placeholder(value_type))
        if hasattr(self, "source_field_host"):
            self.source_field_host.setVisible(show_source_field)
        self.source_field.setEnabled(show_source_field)
        self.cargo_provider_combo.setVisible(show_cargo_provider)
        if hasattr(self, "cargo_provider_host"):
            self.cargo_provider_host.setVisible(show_cargo_provider)
        self.cargo_provider_combo.setEnabled(show_cargo_provider and bool(self._cargo_providers))

        if preserve_values:
            return

        if not show_amount:
            self.amount.clear()
            self.percent_input.clear()
        if not show_source_field:
            with QSignalBlocker(self.source_field):
                self.source_field.setCurrentIndex(-1)
        if not show_cargo_provider:
            payload = dict(self._source_config_override or {})
            payload.pop("cargoProviderId", None)
            payload.pop("cargoProviderCode", None)
            payload.pop("cargoProviderName", None)
            self._source_config_override = payload or None
            with QSignalBlocker(self.cargo_provider_combo):
                self.cargo_provider_combo.setCurrentIndex(-1)

    def state(self) -> dict:
        source_type = self._current_source_type() or None
        source_config = dict(self._source_config_override or {}) or None
        if source_type == "manual_value":
            payload = dict(source_config or {})
            if self._selected_value_type() == "percentage":
                amount = float(self.percent_input.value()) if self.percent_input.has_value() else None
            else:
                amount = float(self.amount.value()) if self.amount.has_value() else None
            if amount is None:
                payload.pop("amount", None)
            else:
                payload["amount"] = amount
            source_config = payload or None
        payload = {
            "cost_key": _safe_str(self.cb.currentData()),
            "calculation_type": self._current_match_mode() or None,
            "source_type": source_type,
            "source_config": source_config,
            "tariff_key": None,
        }
        return normalize_listing_cost_binding(payload, card=self._lookup_selected_cost_item())

    def value(self) -> Optional[dict]:
        payload = self.state()
        if not _cost_row_is_complete(payload):
            return None
        return _compact_cost_row(payload)

    def set_excluded_keys(self, keys: set[str]):
        clean_keys = {_safe_str(v) for v in (keys or set()) if _safe_str(v)}
        if clean_keys == self._excluded_keys:
            return
        self._excluded_keys = clean_keys
        self._selected_key = _safe_str(self.cb.currentData()) or self._selected_key
        self._build_combo()

    def _build_combo(self):
        with QSignalBlocker(self.cb):
            self.cb.clear()
            for it in self._cost_items:
                cost_key = _safe_str(it.get("costKey"))
                if not cost_key:
                    continue
                if cost_key in self._excluded_keys and cost_key != self._selected_key:
                    continue
                self.cb.addItem(_safe_str(it.get("label") or it.get("name") or cost_key), cost_key)
            if self._selected_key:
                idx = self.cb.findData(self._selected_key)
                if idx >= 0:
                    self.cb.setCurrentIndex(idx)
        self._sync_dynamic_inputs(preserve_values=True)


class CostMatchDialog(QDialog):
    def __init__(
        self,
        parent: QWidget,
        *,
        product_title: str,
        cost_items: list[dict],
        cargo_providers: Optional[list[dict[str, Any]]] = None,
        field_options: Optional[list[dict[str, str]]] = None,
        current_items: Optional[list[dict]] = None,
    ):
        super().__init__(parent)
        self.setWindowTitle("Maliyet Eşleştir")
        self.setFixedSize(1040, 620)
        self.selected_items: list[dict] = []
        self._cost_items = list(cost_items or [])
        self._cargo_providers = list(cargo_providers or [])
        self._field_options = list(field_options or [])
        self._rows_data: list[dict] = []
        self._page_size = 6
        self._current_page = 1
        self._row_height: Optional[int] = None
        self._updating_page_size = False
        self._row_widgets: list[tuple[int, CostItemRow]] = []
        self.setStyleSheet("""
        QDialog {
            background: #F8FAFC;
        }
        QLabel#DialogTitle {
            font-size: 14px;
            font-weight: 800;
            color: #0F172A;
        }
        QLabel#DialogSub {
            font-size: 11px;
            color: #64748B;
        }
        QFrame#RowCard {
            background: #FFFFFF;
            border: 1px solid #E5E7EB;
            border-radius: 10px;
        }
        QFrame#RowCard QComboBox, QFrame#RowCard QLineEdit {
            padding: 6px 8px;
            border-radius: 8px;
            border: 1px solid #E5E7EB;
            background: #FFFFFF;
        }
        QFrame#RowCard QToolButton {
            border: 1px solid #E5E7EB;
            border-left: 0px;
            background: #F8FAFC;
        }
        QPushButton#PrimaryBtn {
            padding: 6px 12px;
            border-radius: 9px;
            border: 1px solid #0EA5E9;
            background: #0EA5E9;
            color: white;
            font-weight: 700;
        }
        QPushButton#PrimaryBtn:hover { background: #0284C7; border-color: #0284C7; }
        QPushButton#GhostBtn {
            padding: 6px 10px;
            border-radius: 9px;
            border: 1px solid #E5E7EB;
            background: #FFFFFF;
            font-weight: 600;
        }
        QPushButton#GhostBtn:hover { background: #F1F5F9; }
        """)

        layout = QVBoxLayout(self)
        layout.setContentsMargins(16, 16, 16, 16)
        layout.setSpacing(12)

        title_row = QVBoxLayout()
        lbl_title = QLabel("Maliyet Eşleştir")
        lbl_title.setObjectName("DialogTitle")
        lbl_listing = QLabel(f"İlan: {product_title}")
        lbl_listing.setObjectName("DialogSub")
        lbl_listing.setWordWrap(True)
        lbl_info = QLabel(
            "Bu alanda ilanla maliyet kalemlerini eşleştirir ve bu ilana hangi ek maliyetlerin uygulanacağını belirlersin."
        )
        lbl_info.setObjectName("DialogSub")
        lbl_info.setWordWrap(True)
        title_row.addWidget(lbl_title)
        title_row.addWidget(lbl_listing)
        title_row.addWidget(lbl_info)
        layout.addLayout(title_row)

        self.rows_wrap = QWidget(self)
        self.rows_layout = QVBoxLayout(self.rows_wrap)
        self.rows_layout.setContentsMargins(0, 0, 0, 0)
        self.rows_layout.setSpacing(6)
        self.rows_layout.setAlignment(Qt.AlignmentFlag.AlignTop)
        layout.addWidget(self.rows_wrap, 1)

        tool_row = QHBoxLayout()
        self.btn_add = ActionPulseButton(text="Ekle", center_text=True)
        self.btn_add.setFixedSize(150, 48)
        self.btn_add.clicked.connect(self._add_row)
        tool_row.addWidget(self.btn_add)
        tool_row.addStretch()

        self.prev_btn = QPushButton("<")
        self.prev_btn.setObjectName("GhostBtn")
        self.prev_btn.setFixedWidth(36)
        self.prev_btn.clicked.connect(self._prev_page)
        self.next_btn = QPushButton(">")
        self.next_btn.setObjectName("GhostBtn")
        self.next_btn.setFixedWidth(36)
        self.next_btn.clicked.connect(self._next_page)
        self.page_label = QLabel("Sayfa 0/1")
        self.page_label.setObjectName("DialogSub")

        tool_row.addWidget(self.prev_btn)
        tool_row.addWidget(self.page_label)
        tool_row.addWidget(self.next_btn)
        layout.addLayout(tool_row)

        if current_items:
            for it in current_items:
                compact = _compact_cost_row(it)
                if compact.get("cost_key"):
                    self._rows_data.append(compact)

        self._render_page()

        action_row = QHBoxLayout()
        self.btn_save = QPushButton("Kaydet")
        self.btn_clear = QPushButton("Temizle")
        self.btn_cancel = QPushButton("İptal")
        self.btn_save.setObjectName("PrimaryBtn")
        self.btn_clear.setObjectName("GhostBtn")
        self.btn_cancel.setObjectName("GhostBtn")
        self.btn_save.clicked.connect(self._on_save)
        self.btn_clear.clicked.connect(self._on_clear)
        self.btn_cancel.clicked.connect(self.reject)
        action_row.addStretch()
        action_row.addWidget(self.btn_clear)
        action_row.addWidget(self.btn_cancel)
        action_row.addWidget(self.btn_save)
        layout.addLayout(action_row)

        QTimer.singleShot(0, self._refresh_deferred_state)

    def _refresh_deferred_state(self) -> None:
        self._refresh_page_size()
        self._refresh_action_state()

    def _add_row(self):
        self._rows_data.append(
            {
                "cost_key": "",
                "calculation_type": None,
                "amount": None,
                "quantity": None,
            }
        )
        self._current_page = self._get_total_pages()
        self._render_page()
        self._refresh_action_state()

    def _remove_row(self, idx: int):
        if 0 <= idx < len(self._rows_data):
            self._rows_data.pop(idx)
            self._current_page = max(1, min(self._current_page, self._get_total_pages()))
            self._render_page()
            self._refresh_action_state()

    def _on_row_change(self, idx: int, value: Optional[dict]):
        if 0 <= idx < len(self._rows_data):
            prev_key = _safe_str(self._rows_data[idx].get("cost_key"))
            self._rows_data[idx] = _compact_cost_row(value or {})
            new_key = _safe_str(self._rows_data[idx].get("cost_key"))
            if prev_key != new_key:
                self._refresh_exclusions()
        self._refresh_action_state()

    def _has_missing_selection(self) -> bool:
        for row in self._rows_data or []:
            if row and not _cost_row_is_complete(row):
                return True
        return False

    def _can_add_row(self) -> bool:
        if not self._rows_data:
            return True
        return not self._has_missing_selection()

    def _refresh_action_state(self) -> None:
        if hasattr(self, "btn_add"):
            self.btn_add.setEnabled(self._can_add_row())
        if hasattr(self, "btn_save"):
            self.btn_save.setEnabled(not self._has_missing_selection())

    def _collect_items(self) -> list[dict]:
        items: list[dict] = []
        for row in self._rows_data:
            compact = _compact_cost_row(row or {})
            if _cost_row_is_complete(compact):
                items.append(compact)
        return items

    def _sync_rows_from_ui(self) -> None:
        for idx, row in self._row_widgets:
            self._rows_data[idx] = _compact_cost_row(row.state())

    def _get_total_pages(self) -> int:
        total = len(self._rows_data)
        if total <= 0:
            return 1
        return (total + self._page_size - 1) // self._page_size

    def _get_selected_keys(self, *, exclude_idx: Optional[int] = None) -> set[str]:
        keys: set[str] = set()
        for i, row in enumerate(self._rows_data):
            if exclude_idx is not None and i == exclude_idx:
                continue
            cost_key = _safe_str((row or {}).get("cost_key"))
            if cost_key:
                keys.add(cost_key)
        return keys

    def _refresh_exclusions(self) -> None:
        for idx, row in self._row_widgets:
            row.set_excluded_keys(self._get_selected_keys(exclude_idx=idx))

    def _measure_row_height(self) -> int:
        if self._row_height:
            return self._row_height
        self._row_height = COST_MATCH_ROW_HEIGHT + 30
        return self._row_height

    def _calc_page_size(self) -> int:
        row_h = self._measure_row_height()
        available = int(self.rows_wrap.height())
        if available <= 0:
            available = int(self.height() - 300)
        available = max(row_h, available)
        spacing = int(self.rows_layout.spacing())
        per = max(1, (available + spacing) // max(1, (row_h + spacing)))
        return max(1, int(per))

    def _refresh_page_size(self) -> None:
        if self._updating_page_size:
            return
        self._updating_page_size = True
        try:
            new_size = self._calc_page_size()
            if new_size != self._page_size:
                self._page_size = new_size
                self._current_page = max(1, min(self._current_page, self._get_total_pages()))
                self._render_page()
        finally:
            self._updating_page_size = False

    def _render_page(self):
        while self.rows_layout.count():
            item = self.rows_layout.takeAt(0)
            if item is None:
                continue
            widget = item.widget()
            if widget is not None:
                widget.setParent(None)
                widget.deleteLater()

        total_pages = self._get_total_pages()
        self._current_page = max(1, min(self._current_page, total_pages))

        start = (self._current_page - 1) * self._page_size
        end = min(start + self._page_size, len(self._rows_data))

        self._row_widgets = []
        for idx in range(start, end):
            row_data = self._rows_data[idx]
            row = CostItemRow(
                self._cost_items,
                cargo_providers=self._cargo_providers,
                parent=self.rows_wrap,
                row_index=idx,
                selected_key=row_data.get("cost_key"),
                calculation_type=row_data.get("calculation_type"),
                source_type=row_data.get("source_type"),
                source_config=row_data.get("source_config"),
                field_options=self._field_options,
                amount=row_data.get("amount"),
                quantity=row_data.get("quantity"),
                on_change=self._on_row_change,
                on_remove=self._remove_row,
            )
            self.rows_layout.addWidget(row)
            self._row_widgets.append((idx, row))
            row.set_excluded_keys(self._get_selected_keys(exclude_idx=idx))

        self.page_label.setText(f"Sayfa {self._current_page}/{total_pages}")
        self.prev_btn.setEnabled(self._current_page > 1)
        self.next_btn.setEnabled(self._current_page < total_pages)
        self._refresh_action_state()

    def resizeEvent(self, event):
        super().resizeEvent(event)
        self._refresh_page_size()

    def _prev_page(self):
        self._current_page = max(1, self._current_page - 1)
        self._render_page()

    def _next_page(self):
        self._current_page = min(self._get_total_pages(), self._current_page + 1)
        self._render_page()

    def _on_save(self):
        self._sync_rows_from_ui()
        if self._has_missing_selection():
            return
        self.selected_items = self._collect_items()
        self.accept()

    def _on_clear(self):
        self.selected_items = []
        self.accept()


class PriceStockDialog(QDialog):
    def __init__(
        self,
        parent: QWidget,
        *,
        product_title: str,
        barcode: str,
        quantity: Optional[float],
        list_price: Optional[float],
        sale_price: Optional[float],
    ):
        super().__init__(parent)
        self.setWindowTitle("Stok & Fiyat Güncelle")
        self.setFixedSize(420, 260)

        self._barcode = barcode or ""

        self.setStyleSheet("""
        QDialog { background: #F8FAFC; }
        QLabel#DialogTitle { font-size: 14px; font-weight: 800; color: #0F172A; }
        QLabel#DialogSub { font-size: 11px; color: #64748B; }
        QPushButton#PrimaryBtn {
            padding: 6px 12px;
            border-radius: 9px;
            border: 1px solid #0EA5E9;
            background: #0EA5E9;
            color: white;
            font-weight: 700;
        }
        QPushButton#GhostBtn {
            padding: 6px 10px;
            border-radius: 9px;
            border: 1px solid #E5E7EB;
            background: #FFFFFF;
            font-weight: 600;
        }
        """)

        layout = QVBoxLayout(self)
        layout.setContentsMargins(16, 16, 16, 16)
        layout.setSpacing(10)

        title = QLabel("Stok & Fiyat Güncelle")
        title.setObjectName("DialogTitle")
        subtitle = QLabel(f"Seçim: {product_title}")
        subtitle.setObjectName("DialogSub")
        subtitle.setWordWrap(True)
        layout.addWidget(title)
        layout.addWidget(subtitle)

        if self._barcode:
            lbl_bar = QLabel(f"Barkod: {self._barcode}")
            lbl_bar.setObjectName("DialogSub")
            layout.addWidget(lbl_bar)

        form = QGridLayout()
        form.setHorizontalSpacing(10)
        form.setVerticalSpacing(8)

        self.qty = NumberStepper(minimum=0, maximum=10_000_000, step=1)
        qty_val = 0 if quantity is None else int(quantity)
        self.qty.setValue(qty_val)
        self.qty.setFixedHeight(34)

        self.list_price = NumericPulseLineEdit(decimals=2)
        self.list_price.setPlaceholderText("Liste Fiyatı")
        self.list_price.setFixedHeight(34)
        self.list_price.setValue(list_price)

        self.sale_price = NumericPulseLineEdit(decimals=2)
        self.sale_price.setPlaceholderText("Satış Fiyatı")
        self.sale_price.setFixedHeight(34)
        self.sale_price.setValue(sale_price)

        form.addWidget(QLabel("Stok", objectName="DialogSub"), 0, 0)
        form.addWidget(self.qty, 0, 1)
        form.addWidget(QLabel("Liste Fiyatı", objectName="DialogSub"), 1, 0)
        form.addWidget(self.list_price, 1, 1)
        form.addWidget(QLabel("Satış Fiyatı", objectName="DialogSub"), 2, 0)
        form.addWidget(self.sale_price, 2, 1)

        layout.addLayout(form)
        layout.addStretch(1)

        actions = QHBoxLayout()
        self.btn_cancel = QPushButton("İptal")
        self.btn_save = QPushButton("Güncelle")
        self.btn_cancel.setObjectName("GhostBtn")
        self.btn_save.setObjectName("PrimaryBtn")
        self.btn_cancel.clicked.connect(self.reject)
        self.btn_save.clicked.connect(self._on_save)
        actions.addStretch(1)
        actions.addWidget(self.btn_cancel)
        actions.addWidget(self.btn_save)
        layout.addLayout(actions)

        self._values: Optional[dict] = None

    def _on_save(self):
        if not self._barcode:
            AppMessageCenter.warning(self, "Eksik Barkod", "Barkod bulunamadı.")
            return
        if not self.list_price.has_value() or not self.sale_price.has_value():
            AppMessageCenter.warning(self, "Eksik Fiyat", "Liste ve satış fiyatı girilmelidir.")
            return
        list_price = float(self.list_price.value())
        sale_price = float(self.sale_price.value())
        if list_price <= 0 or sale_price <= 0:
            AppMessageCenter.warning(self, "Hatalı Fiyat", "Fiyatlar 0'dan büyük olmalıdır.")
            return
        if list_price < sale_price:
            AppMessageCenter.warning(self, "Hatalı Fiyat", "Liste fiyatı satış fiyatından küçük olamaz.")
            return

        self._values = {
            "barcode": self._barcode,
            "quantity": int(self.qty.value()),
            "list_price": list_price,
            "sale_price": sale_price,
        }
        self.accept()

    def values(self) -> Optional[dict]:
        return self._values

class ImageDownloadWorker(QThread):
    image_saved = pyqtSignal(str, str)
    MAX_CONCURRENT_DOWNLOADS = 5
    MAX_ATTEMPTS = 2

    def __init__(self, tasks: list[tuple[str, str]], parent=None):
        super().__init__(parent)
        self._tasks = tasks
        self._stop = False

    def _download_one(self, url: str, path: str) -> tuple[str, str] | None:
        if self._stop or not url or not path:
            return None
        target = Path(path)
        if target.exists():
            return url, path
        for _attempt in range(self.MAX_ATTEMPTS):
            if self._stop:
                return None
            try:
                import httpx

                with httpx.Client(
                    timeout=LISTING_IMAGE_DOWNLOAD_TIMEOUT_SECONDS,
                    follow_redirects=True,
                ) as client:
                    resp = client.get(url)
                if resp.status_code == 200 and resp.content:
                    target.parent.mkdir(parents=True, exist_ok=True)
                    target.write_bytes(resp.content)
                    return url, path
            except Exception:
                continue
        return None

    def run(self):
        tasks = [(url, path) for url, path in (self._tasks or []) if url and path]
        if not tasks:
            return
        max_workers = min(self.MAX_CONCURRENT_DOWNLOADS, len(tasks))
        with ThreadPoolExecutor(max_workers=max_workers, thread_name_prefix="listing-img") as pool:
            futures = [pool.submit(self._download_one, url, path) for url, path in tasks]
            for future in as_completed(futures):
                if self._stop:
                    break
                try:
                    result = future.result()
                except Exception:
                    continue
                if result is not None:
                    self.image_saved.emit(result[0], result[1])

    def stop(self):
        self._stop = True


class ImagePreviewDialog(QDialog):
    def __init__(self, parent: QWidget | None = None):
        super().__init__(parent)
        self.setWindowTitle("Görsel Önizleme")
        self.setMinimumSize(680, 520)
        self._pixmap: Optional[QPixmap] = None

        root = QVBoxLayout(self)
        root.setContentsMargins(12, 12, 12, 12)

        self.label = QLabel("Görsel yüklenemedi.")
        self.label.setAlignment(Qt.AlignmentFlag.AlignCenter)
        self.label.setStyleSheet("background:#0f172a;color:#e2e8f0;border-radius:8px;")
        root.addWidget(self.label, 1)

    def set_image_path(self, path: str | None) -> None:
        pix = QPixmap(path or "")
        if pix.isNull():
            self._pixmap = None
            self.label.setText("Görsel bulunamadı.")
            return
        self._pixmap = pix
        self._update_pixmap()

    def _update_pixmap(self) -> None:
        if not self._pixmap:
            return
        scaled = self._pixmap.scaled(
            self.label.size(),
            Qt.AspectRatioMode.KeepAspectRatio,
            Qt.TransformationMode.SmoothTransformation,
        )
        self.label.setPixmap(scaled)

    def resizeEvent(self, event):
        super().resizeEvent(event)
        self._update_pixmap()


class ListingManagerWindow(QWidget):
    def __init__(self, company_list_ref: CompanyListWidget | None = None, parent: QWidget | None = None):
        super().__init__(parent)
        self.company_list_ref = company_list_ref

        self._account_name_map: dict[int, str] = {}
        self._account_logo_map: dict[int, str] = {}
        self._account_platform_map: dict[int, str] = {}
        self._account_ids: list[int] = []
        self._stock_item_label_map: dict[int, str] = {}
        self._stock_cost_map: dict[int, float] = {}
        self._stock_cost_meta_map: dict[int, dict[str, Any]] = {}
        self._stock_vat_rate_map: dict[int, float | None] = {}
        self._stock_dimension_map: dict[int, dict[str, float | None]] = {}
        self._stock_items_cache: list[dict] = []
        self._cost_items_cache: list[dict] = []
        self._cargo_providers_cache: list[dict] = []
        self._cargo_tariffs_cache: list[dict] = []
        self._cache_loaded = False
        self._image_cache_dir = LISTING_IMAGE_CACHE_DIR
        self._image_cache_dir.mkdir(parents=True, exist_ok=True)
        self._image_cache_map: dict[str, str] = {}
        self._image_queue: list[tuple[str, str]] = []
        self._image_worker: Optional[ImageDownloadWorker] = None
        self._image_inflight: set[str] = set()
        self._pending_refresh = False
        self._filter_seq = 0
        self._last_filter_signature: tuple[object, ...] | None = None
        self._pending_verifications: dict[str, float] = {}
        self._pending_local_cost_refresh_keys: set[tuple[int, str]] = set()
        self._active_local_workers: list[SyncWorker] = []
        self._cache_worker: Optional[SyncWorker] = None
        self._filter_worker: Optional[SyncWorker] = None
        self._brand_worker: Optional[SyncWorker] = None
        self._cost_save_worker: Optional[SyncWorker] = None
        self._price_worker: Optional[AsyncWorker] = None
        self._cache_load_seq = 0
        self._image_preview = ImagePreviewDialog(self)
        listing_signals.products_saved.connect(self._on_listing_products_saved)
        listing_signals.stock_matches_changed.connect(self._on_listing_stock_matches_changed)
        listing_signals.cost_matches_changed.connect(self._on_listing_cost_matches_changed)
        listing_signals.listing_cache_updated.connect(self._on_listing_cache_updated)
        account_signals.company_changed.connect(self._request_catalog_refresh)
        marketplace_signals.marketplaces_changed.connect(self._request_catalog_refresh)
        logistics_signals.cargo_providers_changed.connect(self._request_catalog_refresh)
        logistics_signals.cargo_tariffs_changed.connect(self._request_catalog_refresh)

        self.setWindowTitle("İlanlar")
        self.setWindowIcon(QIcon(os.path.join(ASSETS_DIR, "ilanlar.png")))
        self.setObjectName("ListingRoot")
        self.setStyleSheet(compose_record_workspace_style(LISTING_STYLE))

        root = QVBoxLayout(self)
        root.setContentsMargins(10, 8, 10, 8)
        root.setSpacing(10)

        content_row = QHBoxLayout()
        content_row.setContentsMargins(0, 0, 0, 0)
        content_row.setSpacing(8)
        root.addLayout(content_row, 1)

        self._left_stick_size = 360
        self._left_collapse_trigger = 320
        self._right_min_size = 1120
        self._selection_mode = "detail"
        self._active_listing_record: dict | None = None

        self._workspace = RecordWorkspaceScaffold(
            self,
            left_title="Araçlar",
            left_subtitle="Filtreleri sıfırla ve toplu seçim akışını yönet.",
            filters_title="Filtreler",
            filters_subtitle="İlanları sağdan filtrele.",
            records_title="İlan Listesi",
            records_subtitle="",
            show_editor_actions_card=False,
            editor_canvas_color="#FFFFFF",
            editor_surface_color="#FFFFFF",
            shell_kwargs={
                "sidebar_width": 92,
                "left_stick_size": self._left_stick_size,
                "left_collapse_trigger": self._left_collapse_trigger,
                "min_left_width": 260,
                "min_right_width": self._right_min_size,
                "root_margins": (0, 0, 0, 0),
                "root_spacing": 8,
            },
            sidebar_buttons=(
                {
                    "panel_id": "new",
                    "text": "",
                    "tooltip": str(ui_text("listing.sidebar.new_tooltip", "Yeni ilan paneli")),
                    "icon_path": str(ASSETS_DIR / "yeni-kayit.png"),
                    "shape": "circular",
                    "icon_scale": 1.28,
                },
                {
                    "panel_id": "edit",
                    "text": "",
                    "tooltip": str(ui_text("listing.sidebar.edit_tooltip", "Düzenleme paneli")),
                    "icon_path": str(ASSETS_DIR / "duzenleme-icon.png"),
                    "shape": "circular",
                    "icon_scale": 1.28,
                },
                {
                    "panel_id": "bulk",
                    "text": "",
                    "tooltip": str(ui_text("listing.sidebar.bulk_tooltip", "Toplu işlem paneli")),
                    "icon_path": str(ASSETS_DIR / "toplu-secim (2).png"),
                    "shape": "circular",
                    "icon_scale": 1.28,
                },
            ),
        )
        self._workspace.setSizePolicy(QSizePolicy.Policy.Expanding, QSizePolicy.Policy.Expanding)
        content_row.addWidget(self._workspace, 1)
        self._workspace.left_card.setObjectName("Card")
        self._workspace.filters_card.setObjectName("Card")
        self._workspace.filters_card.setSizePolicy(QSizePolicy.Policy.Expanding, QSizePolicy.Policy.Maximum)
        self._workspace.records_card.setObjectName("Card")
        self._workspace.footer_bar.setObjectName("BottomBar")
        self._workspace.footer_bar.page_size_label.setText("Sayfa başına")

        self._workspace_shell = self._workspace.shell
        self._panel_sidebar = self._workspace.sidebar
        self._listing_splitter = self._workspace_shell.splitter
        self._listing_splitter.setCollapsible(0, True)
        self._listing_splitter.setCollapsible(1, False)
        self._tools_card = self._workspace.left_card
        self._list_card = self._workspace.right_card
        self._editor_scroll = self._workspace.editor_scroll
        self.bulk_panel = self._workspace.bulk_panel

        tools_content = QWidget(self._workspace.editor_surface)
        tools_layout = QVBoxLayout(tools_content)
        tools_layout.setContentsMargins(0, 0, 0, 0)
        tools_layout.setSpacing(8)
        tools_hint = QLabel("Filtreleri sıfırla ve liste durumunu soldan takip et.", objectName="Muted")
        tools_hint.setWordWrap(True)
        self.tools_summary_label = QLabel("Toplam 0 | Sayfada 0 | Seçili 0", objectName="Muted")
        self.tools_summary_label.setWordWrap(True)
        self.btn_tools_clear_filter = QPushButton("Filtreleri Sıfırla")
        self.btn_tools_clear_filter.setObjectName("GhostBtn")
        self.btn_tools_clear_filter.clicked.connect(self.clear_db_filters)
        tools_layout.addWidget(tools_hint)
        tools_layout.addWidget(self.tools_summary_label)
        tools_layout.addWidget(self.btn_tools_clear_filter)
        tools_layout.addStretch(1)
        self._workspace.set_editor_widget(tools_content)

        bulk_content = QWidget(self._workspace.bulk_panel)
        bulk_layout = QVBoxLayout(bulk_content)
        bulk_layout.setContentsMargins(0, 0, 0, 0)
        bulk_layout.setSpacing(8)
        bulk_title = QLabel("Toplu İşlem", objectName="SectionTitle")
        self.bulk_summary_label = QLabel("Seçili ilan: 0", objectName="Muted")
        self.bulk_summary_label.setWordWrap(True)
        bulk_hint = QLabel("SEÇ ile işaretlediğin ilanlar üzerinde toplu işlemleri yönet.", objectName="Muted")
        bulk_hint.setWordWrap(True)
        self.btn_bulk_clear_selection = QPushButton("Seçimi Temizle")
        self.btn_bulk_clear_selection.setObjectName("GhostBtn")
        self.btn_bulk_clear_selection.clicked.connect(self._clear_selection)
        bulk_layout.addWidget(bulk_title)
        bulk_layout.addWidget(self.bulk_summary_label)
        bulk_layout.addWidget(bulk_hint)
        bulk_layout.addWidget(self.btn_bulk_clear_selection)
        bulk_layout.addStretch(1)
        self._workspace.set_bulk_widget(bulk_content)
        self._workspace.bulk_panel.hide()

        filters_host = QWidget(self._workspace.filters_host)
        filter_layout = QVBoxLayout(filters_host)
        filter_layout.setContentsMargins(0, 0, 0, 0)
        filter_layout.setSpacing(10)

        search_row = QHBoxLayout()
        search_row.setSpacing(8)
        self.search_input = QLineEdit()
        self.search_input.setObjectName("Input")
        self.search_input.setPlaceholderText("Genel Ara (başlık / barkod / stok / marka / ürün id)")
        self.search_input.setMinimumWidth(260)
        self.search_input.setSizePolicy(QSizePolicy.Policy.Expanding, QSizePolicy.Policy.Fixed)
        search_row.addWidget(self.search_input, 1)
        self.btn_search_clear_filter = QPushButton("Filtreyi Sıfırla")
        self.btn_search_clear_filter.setObjectName("SoftDanger")
        self.btn_search_clear_filter.setMinimumHeight(42)
        self.btn_search_clear_filter.clicked.connect(self.clear_db_filters)
        search_row.addWidget(self.btn_search_clear_filter, 0)
        filter_layout.addLayout(search_row)

        filter_row = QHBoxLayout()
        self.status_filter = SwitchSelectWidget(title="Durum", all_text="Tümü", multi=False, default_all=False)
        self.status_filter.setMinimumWidth(140)
        self.status_filter.setSizePolicy(QSizePolicy.Policy.Expanding, QSizePolicy.Policy.Fixed)
        self.status_filter.addItem("Aktif", "active", checked=True)
        self.status_filter.addItem("Tükenen", "out_of_stock")
        self.status_filter.addItem("Arşiv", "archived")
        self.status_filter.selectionChanged.connect(lambda _v: self._schedule_filter())
        filter_row.addWidget(self.status_filter)

        self.platform_filter = SwitchSelectWidget(title="Platform", all_text="Tümü", multi=True)
        self.platform_filter.setMinimumWidth(130)
        self.platform_filter.setSizePolicy(QSizePolicy.Policy.Expanding, QSizePolicy.Policy.Fixed)
        self.platform_filter.selectionChanged.connect(lambda _v: self._schedule_filter())
        filter_row.addWidget(self.platform_filter)

        self.company_filter = SwitchSelectWidget(title="Şirket", all_text="Tümü", multi=True)
        self.company_filter.setMinimumWidth(130)
        self.company_filter.setSizePolicy(QSizePolicy.Policy.Expanding, QSizePolicy.Policy.Fixed)
        self.company_filter.selectionChanged.connect(lambda _v: self._schedule_filter())
        filter_row.addWidget(self.company_filter)

        self.brand_filter = SwitchSelectWidget(title="Marka", all_text="Tümü", multi=True)
        self.brand_filter.setMinimumWidth(130)
        self.brand_filter.setSizePolicy(QSizePolicy.Policy.Expanding, QSizePolicy.Policy.Fixed)
        self.brand_filter.selectionChanged.connect(lambda _v: self._schedule_filter())
        filter_row.addWidget(self.brand_filter)

        self.match_filter = SwitchSelectWidget(title="Eşleşme", all_text="Tümü", multi=False, default_all=True)
        self.match_filter.setMinimumWidth(145)
        self.match_filter.setSizePolicy(QSizePolicy.Policy.Expanding, QSizePolicy.Policy.Fixed)
        self.match_filter.addItem("Eşleştirilmiş", "matched")
        self.match_filter.addItem("Eşleştirilmemiş", "unmatched")
        self.match_filter.selectionChanged.connect(lambda _v: self._schedule_filter())
        filter_row.addWidget(self.match_filter)
        filter_layout.addLayout(filter_row)

        filter_row2 = QHBoxLayout()
        self.approval_filter = SwitchSelectWidget(title="Onay", all_text="Tümü", multi=False, default_all=True)
        self.approval_filter.setMinimumWidth(130)
        self.approval_filter.setSizePolicy(QSizePolicy.Policy.Expanding, QSizePolicy.Policy.Fixed)
        self.approval_filter.addItem("Onaylı", "approved")
        self.approval_filter.addItem("Onaysız", "unapproved")
        self.approval_filter.selectionChanged.connect(lambda _v: self._schedule_filter())
        filter_row2.addWidget(self.approval_filter)

        self.campaign_filter = SwitchSelectWidget(title="Kampanya", all_text="Tümü", multi=False, default_all=True)
        self.campaign_filter.setMinimumWidth(130)
        self.campaign_filter.setSizePolicy(QSizePolicy.Policy.Expanding, QSizePolicy.Policy.Fixed)
        self.campaign_filter.addItem("Var", "campaign_on")
        self.campaign_filter.addItem("Yok", "campaign_off")
        self.campaign_filter.selectionChanged.connect(lambda _v: self._schedule_filter())
        filter_row2.addWidget(self.campaign_filter)

        self.lock_filter = SwitchSelectWidget(title="Kilit", all_text="Tümü", multi=False, default_all=True)
        self.lock_filter.setMinimumWidth(125)
        self.lock_filter.setSizePolicy(QSizePolicy.Policy.Expanding, QSizePolicy.Policy.Fixed)
        self.lock_filter.addItem("Kilitli", "locked_on")
        self.lock_filter.addItem("Kilitsiz", "locked_off")
        self.lock_filter.selectionChanged.connect(lambda _v: self._schedule_filter())
        filter_row2.addWidget(self.lock_filter)

        self.stock_min_input = QLineEdit()
        self.stock_min_input.setObjectName("Input")
        self.stock_min_input.setPlaceholderText("Min Stok")
        self.stock_min_input.setValidator(QIntValidator(0, 999999999, self))
        self.stock_min_input.setMaximumWidth(108)
        filter_row2.addWidget(self.stock_min_input)

        self.stock_max_input = QLineEdit()
        self.stock_max_input.setObjectName("Input")
        self.stock_max_input.setPlaceholderText("Max Stok")
        self.stock_max_input.setValidator(QIntValidator(0, 999999999, self))
        self.stock_max_input.setMaximumWidth(108)
        filter_row2.addWidget(self.stock_max_input)

        decimal_validator = QRegularExpressionValidator(QRegularExpression(r"^\\d*(?:[.,]\\d{0,4})?$"), self)
        self.price_min_input = QLineEdit()
        self.price_min_input.setObjectName("Input")
        self.price_min_input.setPlaceholderText("Min Fiyat")
        self.price_min_input.setValidator(decimal_validator)
        self.price_min_input.setMaximumWidth(108)
        filter_row2.addWidget(self.price_min_input)

        self.price_max_input = QLineEdit()
        self.price_max_input.setObjectName("Input")
        self.price_max_input.setPlaceholderText("Max Fiyat")
        self.price_max_input.setValidator(decimal_validator)
        self.price_max_input.setMaximumWidth(108)
        filter_row2.addWidget(self.price_max_input)
        filter_layout.addLayout(filter_row2)
        self._workspace.set_filters_widget(filters_host)

        self.list_widget = ListingListWidget(
            self._workspace.records_card,
            open_stock_picker_cb=self._open_stock_picker,
            open_cost_picker_cb=self._open_cost_picker,
            open_profit_detail_cb=self._open_profit_detail,
            open_image_cb=self._open_image_preview,
            apply_price_cb=self._apply_inline_price_update,
        )
        self.list_widget.selection_changed.connect(self._update_tools_summary)
        self.list_widget.selection_changed.connect(self._sync_mode_guidance_state)
        self.list_widget.active_row_changed.connect(self._sync_mode_guidance_state)
        self.list_widget.page_changed.connect(self._on_list_page_changed)
        try:
            self.list_widget.verticalScrollBar().valueChanged.connect(self._sync_guidance_geometry)
            self.list_widget.horizontalScrollBar().valueChanged.connect(self._sync_guidance_geometry)
        except Exception:
            pass
        try:
            model = self.list_widget.model()
            if model is not None:
                model.modelReset.connect(self._sync_guidance_geometry)
                model.layoutChanged.connect(self._sync_guidance_geometry)
                model.rowsInserted.connect(lambda *_: self._sync_guidance_geometry())
                model.rowsRemoved.connect(lambda *_: self._sync_guidance_geometry())
        except Exception:
            pass
        self._workspace.set_list_widget(self.list_widget)
        self._workspace.set_records_summary_text("0 / 0 kayıt")
        self.btn_bulk_select_toggle = self._workspace.footer_bar.select_toggle
        self._listing_splitter.setSizes([0, 1380])
        self._workspace_adapter = ListingWorkspaceBehaviorAdapter(self)
        self._workspace_behavior = self._workspace.create_behavior_controller(
            adapter=self._workspace_adapter,
        )
        self._sync_sidebar_state_from_splitter()
        self._filter_timer = QTimer(self)
        self._filter_timer.setSingleShot(True)
        self._filter_timer.setInterval(250)
        self._filter_timer.timeout.connect(self.apply_db_filters)
        self._brand_source_key = None
        stock_signals.refreshed.connect(self._on_stock_domain_refreshed)

        self.search_input.textChanged.connect(self._schedule_filter)
        self.stock_min_input.textChanged.connect(self._schedule_filter)
        self.stock_max_input.textChanged.connect(self._schedule_filter)
        self.price_min_input.textChanged.connect(self._schedule_filter)
        self.price_max_input.textChanged.connect(self._schedule_filter)
        self._sync_guidance_geometry()
        self._sync_mode_guidance_state()
        self._update_tools_summary()

    def _on_stock_domain_refreshed(self) -> None:
        self._pending_refresh = True
        if not self._cache_loaded:
            return
        if self.isVisible():
            QTimer.singleShot(0, self._load_cached_products)

    def _on_list_page_changed(self, *_args) -> None:
        self._update_pagination_ui()
        self._cancel_image_worker()
        self._enqueue_images(self.list_widget.get_current_page_items())

    def _is_tools_panel_open(self) -> bool:
        return bool(self._workspace_shell.is_left_panel_open())

    def _sync_sidebar_state_from_splitter(self):
        self._workspace_behavior.sync_state()
        self._sync_guidance_geometry()

    def _set_list_loading(self, busy: bool, message: str = LISTING_LIST_LOADING_TEXT) -> None:
        self._workspace.set_list_loading(bool(busy), message=message)

    def _sync_guidance_geometry(self):
        self._workspace_shell.sync_guidance_geometry()
        self._workspace.sync_list_loading_geometry()

    def _sync_mode_guidance_state(self):
        mode = str(self._selection_mode or "detail").lower()
        if mode not in {"new", "detail", "bulk"}:
            mode = "detail"

        if mode == "new":
            self.list_widget.set_active_spin_allowed(False)
            self.list_widget.set_select_button_guide_enabled(False)
            self._workspace_behavior.sync_state()
            return

        panel_open = self._workspace_shell.is_left_panel_open()
        has_active = self.list_widget.has_active_item()
        self.list_widget.set_active_spin_allowed(panel_open and mode == "detail" and has_active)
        self._workspace_behavior.sync_state()

    def _selected_listing_count(self) -> int:
        count = 0
        for it in (self.list_widget.filtered_items or []):
            if isinstance(it, dict) and bool(it.get("_selected")):
                count += 1
        return count

    def _workspace_record_key(self, record: dict | None) -> str:
        return str(self.list_widget._item_key_of(record) or "")

    def _on_listing_behavior_cleared(self) -> None:
        self._selection_mode = "detail" if self._selection_mode != "bulk" else self._selection_mode

    def _sync_listing_workspace_state(self) -> None:
        mode = str(self._selection_mode or "detail").lower()
        if mode not in {"new", "detail", "bulk"}:
            mode = "detail"
        show_editor = mode in {"new", "detail"}
        self._editor_scroll.setVisible(show_editor)
        self.bulk_panel.setVisible(not show_editor)
        if mode == "bulk":
            self._workspace.set_left_title("Toplu İşlem")
            self._workspace.set_left_subtitle("SEÇ ile işaretlediğin ilanlar için toplu işlem akışını yönet.")
        else:
            self._workspace.set_left_title("Araçlar")
            self._workspace.set_left_subtitle("Filtreleri sıfırla ve toplu seçim akışını yönet.")
        panel_open = self._workspace_shell.is_left_panel_open()
        has_active = self.list_widget.has_active_item()
        self.list_widget.set_active_spin_allowed(panel_open and mode == "detail" and has_active)
        self._update_tools_summary()

    def showEvent(self, event):
        super().showEvent(event)
        self._sync_guidance_geometry()
        self._workspace_behavior.sync_state()
        if not self._cache_loaded:
            self._cache_loaded = True
            self._set_list_loading(True, LISTING_LIST_LOADING_TEXT)
            QTimer.singleShot(0, self._load_cached_products)
        elif self._pending_refresh:
            QTimer.singleShot(0, self._load_cached_products)

    def resizeEvent(self, event):
        super().resizeEvent(event)
        self.list_widget.sync_dynamic_row_width()
        self._sync_guidance_geometry()
        self._workspace_behavior.sync_state()

    @staticmethod
    def _shutdown_worker(worker, *, timeout_ms: int = 1500) -> None:
        if worker is None:
            return
        try:
            if hasattr(worker, "stop"):
                worker.stop()
        except Exception:
            pass
        try:
            if hasattr(worker, "requestInterruption"):
                worker.requestInterruption()
        except Exception:
            pass
        try:
            if worker.isRunning():
                worker.wait(timeout_ms)
        except Exception:
            pass

    def closeEvent(self, event):
        try:
            self._filter_timer.stop()
        except Exception:
            pass
        self._cancel_image_worker()

        workers = [
            getattr(self, "_image_worker", None),
            getattr(self, "_price_worker", None),
            getattr(self, "_cache_worker", None),
            getattr(self, "_filter_worker", None),
            getattr(self, "_brand_worker", None),
            getattr(self, "_cost_save_worker", None),
            *list(self._active_local_workers or []),
        ]
        seen: set[int] = set()
        for worker in workers:
            if worker is None:
                continue
            marker = id(worker)
            if marker in seen:
                continue
            seen.add(marker)
            self._shutdown_worker(worker)
        super().closeEvent(event)

    def reload_after_fetch(self):
        self._pending_refresh = True
        self._load_cached_products()

    def _update_counts(self):
        total = len(self.list_widget.filtered_items or [])
        shown = len(self.list_widget.get_current_page_items() or [])
        self._workspace.set_records_summary_text(f"{shown} / {total} kayıt")
        self._update_tools_summary()

    def _update_pagination_ui(self):
        self._workspace.footer_bar.refresh()
        total = len(self.list_widget.filtered_items or [])
        shown = len(self.list_widget.get_current_page_items() or [])
        self._workspace.set_records_summary_text(f"{shown} / {total} kayıt")
        self._update_tools_summary()

    def _update_tools_summary(self):
        total = len(self.list_widget.filtered_items or [])
        shown = len(self.list_widget.get_current_page_items() or [])
        selected = 0
        for it in (self.list_widget.filtered_items or []):
            if isinstance(it, dict) and bool(it.get("_selected")):
                selected += 1
        self.tools_summary_label.setText(f"Toplam {total} | Sayfa {shown} | Seçili {selected}")
        self._sync_selection_toggle_buttons(selected_count=selected, total_count=total)

    def _sync_selection_toggle_buttons(self, *, selected_count: int, total_count: int) -> None:
        self.btn_bulk_select_toggle.setSelectionCounts(int(selected_count or 0), int(total_count or 0))

    def _on_selection_toggle_action(self, action: str) -> None:
        mode = str(action or "").strip().lower()
        if mode == "clear":
            self._clear_selection()
            return
        self._select_all_rows()

    def _select_all_rows(self):
        for it in self.list_widget.filtered_items or []:
            if isinstance(it, dict):
                it["_selected"] = True
        self.list_widget._build_page()
        self.list_widget.selection_changed.emit()
        self._sync_mode_guidance_state()
        self._update_tools_summary()

    def _clear_selection(self):
        for it in self.list_widget.filtered_items or []:
            if isinstance(it, dict):
                it["_selected"] = False
        self.list_widget._build_page()
        self.list_widget.selection_changed.emit()
        self._sync_mode_guidance_state()
        self._update_tools_summary()

    def _load_account_maps(self):
        res = action_get_all_companies()
        if not res.success:
            return
        records = res.data.get("records", []) or []
        self._account_name_map = {
            getattr(r, "pk", None): (getattr(r, "comp_name", "") or "").strip()
            for r in records
            if getattr(r, "pk", None) is not None
        }
        self._account_logo_map = {
            getattr(r, "pk", None): (getattr(r, "logo_path", "") or "")
            for r in records
            if getattr(r, "pk", None) is not None
        }
        self._account_platform_map = {
            getattr(r, "pk", None): _normalized_account_platform(getattr(r, "platform", "") or "")
            for r in records
            if getattr(r, "pk", None) is not None
        }
        self._account_ids = [rid for rid in self._account_name_map.keys()]

        previous_selection = [
            _normalized_account_platform(value)
            for value in (self.platform_filter.selectedData() or [])
            if value not in (None, "all")
        ]
        # refresh filter sources
        with QSignalBlocker(self.platform_filter):
            self.platform_filter.clear()
            options = _platform_filter_options_from_account_platforms(self._account_platform_map)
            selected_texts: list[str] = []
            for label, code in options:
                self.platform_filter.addItem(label, code)
                if code in previous_selection:
                    selected_texts.append(label)
            if selected_texts:
                self.platform_filter.setSelectedTexts(selected_texts)

        with QSignalBlocker(self.company_filter):
            self.company_filter.clear()
            for acc_id, name in self._account_name_map.items():
                label = name or f"Şirket {acc_id}"
                self.company_filter.addItem(label, acc_id)

    def _load_stock_items_cache(self):
        res = action_list_items(only_active=True)
        if not res.success:
            return
        records = res.data.get("records", []) or []
        self._stock_items_cache = [
            {
                "id": _stock_item_identity(it),
                "label": f"{getattr(it, 'sku', '')} - {getattr(it, 'name', '')}",
            }
            for it in records
            if _stock_item_identity(it) is not None
        ]
        self._stock_item_label_map = {
            it["id"]: it["label"]
            for it in self._stock_items_cache
        }
        self._stock_cost_map = {}
        self._stock_cost_meta_map = {}
        self._stock_vat_rate_map = {}
        self._stock_dimension_map = {}
        for it in records:
            item_id = _stock_item_identity(it)
            if item_id is None:
                continue
            vat_rate = getattr(it, "vat_rate", None)
            try:
                self._stock_vat_rate_map[int(item_id)] = float(vat_rate) if vat_rate not in (None, "") else None
            except Exception:
                self._stock_vat_rate_map[int(item_id)] = None
            self._stock_dimension_map[int(item_id)] = {
                "weight_kg": getattr(it, "weight_kg", None),
                "length_cm": getattr(it, "length_cm", None),
                "width_cm": getattr(it, "width_cm", None),
                "height_cm": getattr(it, "height_cm", None),
            }
        balances_res = action_list_item_balances(only_active=True)
        if not balances_res.success:
            return
        for row in (balances_res.data or {}).get("records", []) or []:
            try:
                item = row.get("item")
                item_id = int(_stock_item_identity(item))
            except Exception:
                continue
            try:
                avg_cost = float(row.get("avg_unit_cost"))
            except Exception:
                continue
            self._stock_cost_map[item_id] = avg_cost
            self._stock_cost_meta_map[item_id] = {
                "input_mode": row.get("avg_cost_input_mode"),
            }

    def _load_cost_items_cache(self):
        res = action_load_finance_cost_items(active_state="Aktif")
        if not res.success:
            return
        records = (res.data or {}).get("records", []) or []
        self._cost_items_cache = [
            {
                "costKey": _safe_str(it.get("costKey")),
                "name": _safe_str(it.get("name")),
                "label": _safe_str(it.get("name")),
                "valueType": _safe_str(it.get("valueType") or it.get("calculationType")),
                "valueTypeLabel": _safe_str(it.get("valueTypeLabel") or it.get("calculationTypeLabel")),
                "defaultMatchMode": _safe_str(it.get("defaultMatchMode")) or "manual",
                "defaultMatchModeLabel": _safe_str(it.get("defaultMatchModeLabel")),
                "supportedSourceTypes": list(it.get("supportedSourceTypes") or []),
                "defaultSourceType": _safe_str(it.get("defaultSourceType")),
                "defaultSourceTypeLabel": _safe_str(it.get("defaultSourceTypeLabel")),
                "defaultSourceConfig": dict(it.get("defaultSourceConfig") or {}) or None,
                "defaultAmount": it.get("defaultAmount"),
                "currency": _safe_str(it.get("currency")) or "TRY",
                "isVatIncluded": bool(it.get("isVatIncluded", True)),
                "vatRate": it.get("vatRate"),
                "supportedPlatforms": list(it.get("supportedPlatforms") or []),
            }
            for it in records
            if _safe_str(it.get("costKey"))
            and (not bool(it.get("isBuiltin")) or bool(it.get("isBuiltinAccepted")))
        ]

    def _load_cargo_providers_cache(self):
        res = action_load_logistics_cargo_providers(source_filter="accepted")
        if not res.success:
            return
        records = (res.data or {}).get("records", []) or []
        self._cargo_providers_cache = [
            {
                "pk": int(it.get("pk") or 0),
                "platformCode": _safe_str(it.get("platformCode")),
                "cargoProviderCode": _safe_str(it.get("cargoProviderCode")).upper(),
                "cargoProviderName": _safe_str(it.get("cargoProviderName")),
            }
            for it in records
            if int(it.get("pk") or 0) > 0
            and _safe_str(it.get("cargoProviderCode"))
            and _safe_str(it.get("cargoProviderName"))
        ]

    def _load_cargo_tariffs_cache(self):
        res = action_load_logistics_cargo_tariffs(active_state="Aktif")
        if not res.success:
            return
        records = (res.data or {}).get("records", []) or []
        self._cargo_tariffs_cache = [
            {
                "tariffKey": _safe_str(it.get("tariffKey")),
                "name": _safe_str(it.get("name")),
                "label": _safe_str(
                    it.get("name")
                    or f"{_safe_str(it.get('cargoProviderName'))} / {_safe_str(it.get('platform'))}"
                ),
                "cargo_provider_id": int(it.get("cargo_provider_id") or 0) or None,
                "platformCode": _safe_str(it.get("platformCode") or it.get("platform")),
                "platform": _safe_str(it.get("platform")),
                "cargoProviderCode": _safe_str(it.get("cargoProviderCode")),
                "cargoProviderName": _safe_str(it.get("cargoProviderName")),
                "minDesi": it.get("minDesi"),
                "maxDesi": it.get("maxDesi"),
                "minSalePrice": it.get("minSalePrice"),
                "maxSalePrice": it.get("maxSalePrice"),
                "amount": it.get("amount"),
                "currency": _safe_str(it.get("currency")),
                "isVatIncluded": bool(it.get("isVatIncluded", True)),
                "vatRate": it.get("vatRate"),
                "isActive": bool(it.get("isActive", True)),
            }
            for it in records
            if _safe_str(it.get("tariffKey"))
            and (not bool(it.get("isBuiltin")) or bool(it.get("isBuiltinAccepted")))
        ]

    def _cargo_providers_for_platform(self, platform_code: Any) -> list[dict]:
        normalized_platform = _normalized_account_platform(platform_code)
        if not normalized_platform:
            return []
        return [
            dict(row)
            for row in self._cargo_providers_cache
            if _normalized_account_platform(row.get("platformCode")) == normalized_platform
        ]

    def _cost_items_for_platform(self, platform_code: Any) -> list[dict]:
        return _cost_items_for_platform_from_cache(self._cost_items_cache, platform_code)

    def _build_cost_card_map(self, *, platform_code: Any = None) -> dict[str, dict]:
        return _build_cost_card_map_from_cache(self._cost_items_cache, platform_code=platform_code)

    def _build_listing_enrichment_context(self) -> dict[str, Any]:
        return {
            "account_name_map": self._account_name_map,
            "account_logo_map": self._account_logo_map,
            "account_platform_map": self._account_platform_map,
            "stock_item_label_map": self._stock_item_label_map,
            "stock_cost_map": self._stock_cost_map,
            "stock_vat_rate_map": self._stock_vat_rate_map,
            "stock_cost_meta_map": self._stock_cost_meta_map,
            "stock_dimension_map": self._stock_dimension_map,
            "cost_items_cache": self._cost_items_cache,
            "cargo_tariffs": self._cargo_tariffs_cache,
        }

    def _register_local_worker(self, worker: SyncWorker) -> None:
        self._active_local_workers.append(worker)

        def _cleanup() -> None:
            try:
                self._active_local_workers.remove(worker)
            except ValueError:
                pass

        worker.finished.connect(_cleanup)

    def _refresh_product_profit_estimation(
        self,
        product: dict,
        *,
        platform_code: Any = None,
    ) -> None:
        if not isinstance(product, dict):
            return
        resolved_platform = _normalized_account_platform(
            platform_code or self._account_platform_map.get(product.get("api_account_id"))
        )
        product["linked_stock_label"] = _compose_linked_stock_label(
            product.get("stock_items"),
            stock_item_label_map=self._stock_item_label_map,
            fallback=_safe_str(product.get("linked_stock_label")),
        )
        profit_payload = estimate_listing_profit(
            item=product,
            stock_cost_map=self._stock_cost_map,
            stock_vat_rate_map=self._stock_vat_rate_map,
            stock_cost_meta_map=self._stock_cost_meta_map,
            stock_dimension_map=self._stock_dimension_map,
            cost_card_map=self._build_cost_card_map(platform_code=resolved_platform),
            cargo_tariffs=self._cargo_tariffs_cache,
            platform_code=resolved_platform,
        )
        product.update(profit_payload)

    def _decorate_product(
        self,
        product: dict,
        *,
        cost_card_map: Optional[dict[str, dict]] = None,
        sync_cost_snapshot: bool = True,
    ) -> None:
        if not isinstance(product, dict):
            return
        prepared = _prepare_listing_item_for_view(
            product,
            account_name_map=self._account_name_map,
            account_logo_map=self._account_logo_map,
            account_platform_map=self._account_platform_map,
            stock_item_label_map=self._stock_item_label_map,
            stock_cost_map=self._stock_cost_map,
            stock_vat_rate_map=self._stock_vat_rate_map,
            stock_cost_meta_map=self._stock_cost_meta_map,
            stock_dimension_map=self._stock_dimension_map,
            cost_items_cache=self._cost_items_cache,
            cargo_tariffs=self._cargo_tariffs_cache,
            cost_card_map=cost_card_map,
            sync_cost_snapshot=sync_cost_snapshot,
        )
        product.update(prepared)

    def _decorate_items(self, items: list[dict]):
        prepared_items = _prepare_listing_items_for_view(
            items,
            account_name_map=self._account_name_map,
            account_logo_map=self._account_logo_map,
            account_platform_map=self._account_platform_map,
            stock_item_label_map=self._stock_item_label_map,
            stock_cost_map=self._stock_cost_map,
            stock_vat_rate_map=self._stock_vat_rate_map,
            stock_cost_meta_map=self._stock_cost_meta_map,
            stock_dimension_map=self._stock_dimension_map,
            cost_items_cache=self._cost_items_cache,
            cargo_tariffs=self._cargo_tariffs_cache,
            sync_cost_snapshot=False,
        )
        for original, prepared in zip(items or [], prepared_items):
            if isinstance(original, dict):
                original.update(prepared)

    def _refresh_account_filter_widgets(self) -> None:
        previous_platforms = [
            _normalized_account_platform(value)
            for value in (self.platform_filter.selectedData() or [])
            if value not in (None, "all")
        ]
        previous_companies: set[int] = set()
        for value in (self.company_filter.selectedData() or []):
            if value in (None, "all"):
                continue
            try:
                previous_companies.add(int(value))
            except Exception:
                continue

        with QSignalBlocker(self.platform_filter):
            self.platform_filter.clear()
            selected_platform_texts: list[str] = []
            for label, code in _platform_filter_options_from_account_platforms(self._account_platform_map):
                self.platform_filter.addItem(label, code)
                if code in previous_platforms:
                    selected_platform_texts.append(label)
            if selected_platform_texts:
                self.platform_filter.setSelectedTexts(selected_platform_texts)

        with QSignalBlocker(self.company_filter):
            self.company_filter.clear()
            selected_company_texts: list[str] = []
            for acc_id, name in self._account_name_map.items():
                label = name or f"Şirket {acc_id}"
                self.company_filter.addItem(label, acc_id)
                if acc_id in previous_companies:
                    selected_company_texts.append(label)
            if selected_company_texts:
                self.company_filter.setSelectedTexts(selected_company_texts)

    def _apply_listing_cache_payload(self, payload: dict[str, Any]) -> None:
        self._account_name_map = dict(payload.get("account_name_map") or {})
        self._account_logo_map = dict(payload.get("account_logo_map") or {})
        self._account_platform_map = dict(payload.get("account_platform_map") or {})
        self._account_ids = list(payload.get("account_ids") or [])
        self._stock_items_cache = list(payload.get("stock_items_cache") or [])
        self._stock_item_label_map = dict(payload.get("stock_item_label_map") or {})
        self._stock_cost_map = dict(payload.get("stock_cost_map") or {})
        self._stock_cost_meta_map = dict(payload.get("stock_cost_meta_map") or {})
        self._stock_vat_rate_map = dict(payload.get("stock_vat_rate_map") or {})
        self._stock_dimension_map = dict(payload.get("stock_dimension_map") or {})
        self._cost_items_cache = list(payload.get("cost_items_cache") or [])
        self._cargo_providers_cache = list(payload.get("cargo_providers_cache") or [])
        self._cargo_tariffs_cache = list(payload.get("cargo_tariffs_cache") or [])
        self._refresh_account_filter_widgets()

    def _load_cached_products(self):
        self._cache_load_seq += 1
        seq = self._cache_load_seq
        self._set_list_loading(True, LISTING_LIST_LOADING_TEXT)
        self._cache_worker = SyncWorker(_load_listing_cache_payload, parent=self)
        self._register_local_worker(self._cache_worker)

        def _on_cache_loaded(res: Result):
            if seq != self._cache_load_seq:
                return
            if not res.success:
                MessageHandler.show(self, res, only_errors=True)
                self._set_list_loading(False)
                return
            self._apply_listing_cache_payload(res.data or {})
            self._pending_refresh = False
            self._last_filter_signature = None
            self._brand_source_key = None
            self.apply_db_filters()

        def _on_cache_finished():
            if seq == self._cache_load_seq:
                self._cache_worker = None

        self._cache_worker.result_ready.connect(_on_cache_loaded)
        self._cache_worker.finished.connect(_on_cache_finished)
        self._cache_worker.start()

    def _request_catalog_refresh(self) -> None:
        self._pending_refresh = True
        self._last_filter_signature = None
        self._brand_source_key = None
        if not self._cache_loaded:
            return
        if self.isVisible():
            QTimer.singleShot(0, self._load_cached_products)

    def _refresh_brand_filter(self, brands: list[str]) -> None:
        brands = sorted({b for b in (brands or []) if _safe_str(b)})
        current = set(self.brand_filter.selectedData() or [])
        with QSignalBlocker(self.brand_filter):
            self.brand_filter.clear()
            for brand in brands:
                self.brand_filter.addItem(brand, brand)
            if current and "all" not in current:
                keep = [b for b in brands if b in current]
                self.brand_filter.setSelectedTexts(keep)

    def apply_db_filters(self):
        force_refresh = bool(self._pending_refresh)
        self._pending_refresh = False

        def _parse_num(text: str) -> Optional[float]:
            raw = (text or "").strip().replace(",", ".")
            if not raw:
                return None
            try:
                return float(raw)
            except Exception:
                return None

        selected = []
        if self.company_list_ref is not None:
            try:
                selected = self.company_list_ref.get_selected_company_pks()
            except Exception:
                selected = []

        platform_sel = []
        company_sel: list[int] = []
        data = [d for d in (self.platform_filter.selectedData() or []) if d is not None]
        if "all" not in data:
            platform_sel = [
                code for code in (_normalized_account_platform(d) for d in data if d)
                if code
            ]
        data = [d for d in (self.company_filter.selectedData() or []) if d is not None]
        if "all" not in data:
            for v in data:
                try:
                    company_sel.append(int(v))
                except Exception:
                    continue

        brand_sel = []
        data = [d for d in (self.brand_filter.selectedData() or []) if d is not None]
        if "all" not in data:
            brand_sel = [str(d) for d in data if str(d)]

        acc_ids = list(selected or self._account_ids or [])
        if company_sel:
            acc_ids = [aid for aid in acc_ids if aid in company_sel]
        if platform_sel:
            acc_ids = [
                aid for aid in acc_ids
                if _normalized_account_platform(self._account_platform_map.get(aid)) in platform_sel
            ]
        if not acc_ids:
            acc_ids = []

        status_vals = self.status_filter.selectedData() or []
        status_val = status_vals[0] if status_vals else "all"

        match_val = "all"
        data = [d for d in (self.match_filter.selectedData() or []) if d is not None]
        match_val = data[0] if data else "all"

        approval_val = "all"
        data = [d for d in (self.approval_filter.selectedData() or []) if d is not None]
        approval_val = data[0] if data else "all"

        campaign_val = "all"
        data = [d for d in (self.campaign_filter.selectedData() or []) if d is not None]
        campaign_val = data[0] if data else "all"

        lock_val = "all"
        data = [d for d in (self.lock_filter.selectedData() or []) if d is not None]
        lock_val = data[0] if data else "all"

        stock_min = _parse_num(self.stock_min_input.text())
        stock_max = _parse_num(self.stock_max_input.text())
        price_min = _parse_num(self.price_min_input.text())
        price_max = _parse_num(self.price_max_input.text())

        search_text = (self.search_input.text() or "").strip()
        params = {
            "api_account_ids": acc_ids or None,
            "listing_status": status_val,
            "search": search_text or None,
            "brand": brand_sel or None,
        }
        filter_signature: tuple[object, ...] = (
            tuple(sorted(int(a) for a in (acc_ids or []))),
            str(status_val),
            str(search_text).casefold(),
            tuple(sorted(str(b).casefold() for b in (brand_sel or []))),
            str(match_val),
            str(approval_val),
            str(campaign_val),
            str(lock_val),
            stock_min,
            stock_max,
            price_min,
            price_max,
        )

        if (not force_refresh) and (filter_signature == self._last_filter_signature):
            self._update_counts()
            self._update_pagination_ui()
            self._sync_mode_guidance_state()
            return

        self._filter_seq += 1
        seq = self._filter_seq
        self._set_list_loading(True, LISTING_LIST_LOADING_TEXT)
        filter_params = {
            "match_val": str(match_val),
            "approval_val": str(approval_val),
            "campaign_val": str(campaign_val),
            "lock_val": str(lock_val),
            "stock_min": stock_min,
            "stock_max": stock_max,
            "price_min": price_min,
            "price_max": price_max,
        }
        enrichment_context = {
            "account_name_map": self._account_name_map,
            "account_logo_map": self._account_logo_map,
            "account_platform_map": self._account_platform_map,
            "stock_item_label_map": self._stock_item_label_map,
            "stock_cost_map": self._stock_cost_map,
            "stock_vat_rate_map": self._stock_vat_rate_map,
            "stock_cost_meta_map": self._stock_cost_meta_map,
            "stock_dimension_map": self._stock_dimension_map,
            "cost_items_cache": self._cost_items_cache,
            "cargo_tariffs": self._cargo_tariffs_cache,
        }
        self._filter_worker = SyncWorker(
            _prepare_filtered_listing_payload,
            parent=self,
            query_params=params,
            filter_params=filter_params,
            enrichment_context=enrichment_context,
        )

        brand_key = (tuple(sorted(acc_ids or [])), status_val)
        brand_params = {
            "api_account_ids": acc_ids or None,
            "listing_status": status_val,
        }

        def _on_filtered(res: Result):
            if seq != self._filter_seq:
                return
            try:
                if not res.success:
                    MessageHandler.show(self, res, only_errors=True)
                    return

                items = (res.data or {}).get("items", []) or []
                self.list_widget.set_items(items)
                self._last_filter_signature = filter_signature
                self._update_counts()
                self._update_pagination_ui()
                self._sync_mode_guidance_state()
                self._cancel_image_worker()
                self._enqueue_images(self.list_widget.get_current_page_items())
            finally:
                self._set_list_loading(False)

        def _on_filter_finished():
            if seq == self._filter_seq:
                self._set_list_loading(False)

        self._filter_worker.result_ready.connect(_on_filtered)
        self._filter_worker.finished.connect(_on_filter_finished)
        self._filter_worker.start()

        if self._brand_source_key == brand_key:
            return
        self._brand_source_key = brand_key
        self._brand_worker = SyncWorker(action_get_listing_brands_from_db, parent=self, **brand_params)

        def _on_brands(res: Result):
            if seq != self._filter_seq:
                return
            if not res.success:
                return
            brands = (res.data or {}).get("brands", []) or []
            self._refresh_brand_filter(brands)

        self._brand_worker.result_ready.connect(_on_brands)
        self._brand_worker.start()

    def _schedule_filter(self):
        self._filter_timer.start()

    def mark_stale(self):
        self._pending_refresh = True

    def _refresh_from_listing_signal(self):
        self.mark_stale()
        if self.isVisible():
            QTimer.singleShot(0, self.apply_db_filters)

    def _consume_pending_local_cost_refresh(self, payload: Any) -> bool:
        try:
            api_account_id = int((payload or {}).get("api_account_id"))
        except Exception:
            api_account_id = 0
        product_id = _safe_str((payload or {}).get("product_id"))
        if api_account_id <= 0 or not product_id:
            return False
        pending_key = (api_account_id, product_id)
        if pending_key not in self._pending_local_cost_refresh_keys:
            return False
        self._pending_local_cost_refresh_keys.discard(pending_key)
        return True

    def _on_listing_products_saved(self, _payload: Any):
        self._refresh_from_listing_signal()

    def _on_listing_stock_matches_changed(self, _payload: Any):
        self._refresh_from_listing_signal()

    def _on_listing_cost_matches_changed(self, _payload: Any):
        if self._consume_pending_local_cost_refresh(_payload):
            return
        self._refresh_from_listing_signal()

    def _on_listing_cache_updated(self, _payload: Any):
        self._refresh_from_listing_signal()

    def clear_db_filters(self):
        self.search_input.clear()
        self.approval_filter.setSelectedTexts([])
        self.campaign_filter.setSelectedTexts([])
        self.lock_filter.setSelectedTexts([])
        self.stock_min_input.clear()
        self.stock_max_input.clear()
        self.price_min_input.clear()
        self.price_max_input.clear()
        self.status_filter.setSelectedTexts(["Aktif"])
        self.platform_filter.setSelectedTexts([])
        self.company_filter.setSelectedTexts([])
        self.brand_filter.setSelectedTexts([])
        self.match_filter.setSelectedTexts([])
        self._load_cached_products()

    def _open_stock_picker(self, product: dict):
        try:
            if not self._stock_items_cache:
                self._load_stock_items_cache()

            title = _product_title(product) or "Ürün"
            current_items = product.get("stock_items") or []
            dialog_ref: dict[str, StockMatchDialog] = {}

            def _open_stock_cards_dialog() -> None:
                self._open_stock_items_manager(
                    on_finished=lambda: (
                        self._load_stock_items_cache(),
                        dialog_ref.get("dialog") and dialog_ref["dialog"].replace_stock_items(self._stock_items_cache),
                    )
                )

            dlg = StockMatchDialog(
                self,
                product_title=title,
                stock_items=self._stock_items_cache,
                current_items=current_items,
                open_stock_cards_cb=_open_stock_cards_dialog,
            )
            dialog_ref["dialog"] = dlg
            if dlg.exec() != QDialog.DialogCode.Accepted:
                return

            api_account_id = product.get("api_account_id")
            product_id = _safe_str(product.get("productId") or product.get("id"))
            if api_account_id is None or not product_id:
                MessageHandler.show(self, Result.fail("Ürün anahtarı bulunamadı."), only_errors=True)
                return

            res = action_set_listing_stock_items(
                api_account_id=int(api_account_id),
                product_id=str(product_id),
                stock_items=dlg.selected_items,
            )
            if not res.success:
                result = {"message": res.message}
                payload = {"sync_api_account_ids": []}
                MessageHandler.show(self, res, only_errors=True)
                if False:
                    listing_signals.products_save_failed.emit(
                    {
                        "message": result.get("message", "DB kayıt hatası."),
                        "api_account_ids": list(payload["sync_api_account_ids"]),
                    }
                )
                return

        except Exception as e:
            MessageHandler.show(self, Result.fail(map_error_to_message(e), error=e), only_errors=True)

    def _open_stock_items_manager(self, on_finished: Optional[Callable[[], None]] = None) -> None:
        try:
            from Stock.views.views import StockItemManagerDialog

            item_dialog = StockItemManagerDialog(self)
            if on_finished:
                item_dialog.finished.connect(lambda *_args: on_finished())
            item_dialog.open_full()
        except Exception as e:
            MessageHandler.show(self, Result.fail(map_error_to_message(e), error=e), only_errors=True)

    def _open_cost_picker(self, product: dict):
        try:
            if not self._cost_items_cache:
                self._load_cost_items_cache()
            if not self._cost_items_cache:
                MessageHandler.show(
                    self,
                    Result.fail("Önce finans tarafında en az bir maliyet kalemi tanımlayın."),
                    only_errors=True,
                )
                return

            title = _product_title(product) or "Ürün"
            current_items = product.get("cost_items") or []
            field_options = _listing_db_field_options(product)
            platform_code = _normalized_account_platform(self._account_platform_map.get(product.get("api_account_id")))
            cost_items = self._cost_items_for_platform(platform_code)
            if not self._cargo_providers_cache:
                self._load_cargo_providers_cache()
            cargo_providers = self._cargo_providers_for_platform(platform_code)
            if not cost_items:
                MessageHandler.show(
                    self,
                    Result.fail("Bu platform için aktif maliyet kartı bulunamadı."),
                    only_errors=True,
                )
                return

            dlg = CostMatchDialog(
                self,
                product_title=title,
                cost_items=cost_items,
                cargo_providers=cargo_providers,
                field_options=field_options,
                current_items=current_items,
            )
            if dlg.exec() != QDialog.DialogCode.Accepted:
                return

            api_account_id = product.get("api_account_id")
            product_id = _product_id_text(product)
            if api_account_id is None or not product_id:
                MessageHandler.show(self, Result.fail("Ürün anahtarı bulunamadı."), only_errors=True)
                return

            pending_key = (int(api_account_id), str(product_id))
            self._pending_local_cost_refresh_keys.add(pending_key)
            self._set_list_loading(True, "Maliyetler kaydediliyor...")
            self._cost_save_worker = SyncWorker(
                action_set_listing_cost_items,
                parent=self,
                api_account_id=int(api_account_id),
                product_id=str(product_id),
                cost_items=[dict(it or {}) for it in dlg.selected_items],
            )
            self._register_local_worker(self._cost_save_worker)

            def _on_saved(res: Result) -> None:
                self._pending_local_cost_refresh_keys.discard(pending_key)
                if not res.success:
                    MessageHandler.show(self, res, only_errors=True)
                    self._set_list_loading(False)
                    return
                updated_product = dict(product)
                updated_product["cost_items"] = [dict(it or {}) for it in dlg.selected_items]
                prepare_worker = SyncWorker(
                    _prepare_single_listing_item_payload,
                    parent=self,
                    product=updated_product,
                    enrichment_context=self._build_listing_enrichment_context(),
                    platform_code=platform_code,
                )
                self._register_local_worker(prepare_worker)

                def _on_prepared(prepare_res: Result) -> None:
                    if not prepare_res.success:
                        MessageHandler.show(self, prepare_res, only_errors=True)
                        return
                    prepared_product = dict(((prepare_res.data or {}).get("item") or {}))
                    if not prepared_product:
                        return
                    product.clear()
                    product.update(prepared_product)
                    self.list_widget.refresh_record(prepared_product)
                    self.list_widget.viewport().update()

                def _on_prepare_finished() -> None:
                    self._set_list_loading(False)

                prepare_worker.result_ready.connect(_on_prepared)
                prepare_worker.finished.connect(_on_prepare_finished)
                prepare_worker.start()

            def _on_finished() -> None:
                self._pending_local_cost_refresh_keys.discard(pending_key)

            self._cost_save_worker.result_ready.connect(_on_saved)
            self._cost_save_worker.finished.connect(_on_finished)
            self._cost_save_worker.start()

        except Exception as e:
            MessageHandler.show(self, Result.fail(map_error_to_message(e), error=e), only_errors=True)

    def _open_profit_detail(self, product: dict):
        try:
            if not product:
                return
            live_product = deepcopy(product)
            if not self._stock_items_cache or not self._stock_item_label_map:
                self._load_stock_items_cache()
            platform_code = _normalized_account_platform(self._account_platform_map.get(live_product.get("api_account_id")))
            has_cached_estimation = bool(
                live_product.get("estimated_profit_detail")
                or live_product.get("estimated_profit_breakdown")
            )
            estimation_context = {
                "stock_cost_map": self._stock_cost_map,
                "stock_vat_rate_map": self._stock_vat_rate_map,
                "stock_cost_meta_map": self._stock_cost_meta_map,
                "stock_dimension_map": self._stock_dimension_map,
                "stock_item_label_map": self._stock_item_label_map,
                "cost_card_map": self._build_cost_card_map(platform_code=platform_code),
                "cargo_tariffs": self._cargo_tariffs_cache,
                "platform_code": platform_code,
            }

            def _open_dialog(prepared_product: dict[str, Any]) -> None:
                dlg = ProfitBreakdownDialog(
                    self,
                    product_title=_product_title(prepared_product) or "İlan",
                    product=prepared_product,
                    estimation_context=estimation_context,
                )
                self._profit_detail_window = dlg
                dlg.showMaximized()
                dlg.raise_()
                dlg.activateWindow()

            if has_cached_estimation:
                _open_dialog(live_product)
                return

            self._set_list_loading(True, LISTING_PROFIT_DETAIL_LOADING_TEXT)
            prepare_worker = SyncWorker(
                _prepare_single_listing_item_payload,
                parent=self,
                product=live_product,
                enrichment_context=self._build_listing_enrichment_context(),
                platform_code=platform_code,
            )
            self._register_local_worker(prepare_worker)

            def _on_prepared(res: Result) -> None:
                if not res.success:
                    MessageHandler.show(self, res, only_errors=True)
                    return
                prepared_product = dict(((res.data or {}).get("item") or {}))
                if not prepared_product:
                    return
                _open_dialog(prepared_product)

            def _on_finished() -> None:
                self._set_list_loading(False)

            prepare_worker.result_ready.connect(_on_prepared)
            prepare_worker.finished.connect(_on_finished)
            prepare_worker.start()
        except Exception as e:
            MessageHandler.show(self, Result.fail(map_error_to_message(e), error=e), only_errors=True)

    def _open_price_editor(self, product: dict):
        try:
            if not product:
                return

            api_account_id = product.get("api_account_id")
            if api_account_id is None:
                MessageHandler.show(self, Result.fail("Şirket bilgisi bulunamadı."), only_errors=True)
                return

            platform = _normalized_account_platform(self._account_platform_map.get(api_account_id))
            if platform and platform != "TRENDYOL":
                MessageHandler.show(self, Result.fail("Bu işlem sadece Trendyol için destekleniyor."), only_errors=True)
                return

            barcode = _safe_str(product.get("barcode"))
            if not barcode:
                MessageHandler.show(self, Result.fail("Barkod bulunamadı."), only_errors=True)
                return

            qty = product.get("quantity") or product.get("stockQuantity") or product.get("stock")
            list_price = product.get("list_price") or product.get("listPrice")
            sale_price = product.get("sale_price") or product.get("salePrice")

            dlg = PriceStockDialog(
                self,
                product_title=_product_title(product),
                barcode=barcode,
                quantity=qty,
                list_price=list_price,
                sale_price=sale_price,
            )
            if dlg.exec() != QDialog.DialogCode.Accepted:
                return

            values = dlg.values()
            if not values:
                return
            self._apply_price_update(
                product,
                quantity=int(values["quantity"]),
                sale_price=float(values["sale_price"]),
                list_price=float(values["list_price"]),
            )

        except Exception as e:
            MessageHandler.show(self, Result.fail(map_error_to_message(e), error=e), only_errors=True)

    def _apply_inline_price_update(
        self,
        product: dict,
        quantity: int,
        price: float,
        feedback_cb: Optional[Callable[[object], None]] = None,
    ):
        if price <= 0:
            if feedback_cb:
                feedback_cb(False)
            MessageHandler.show(self, Result.fail("Fiyat 0'dan büyük olmalı."), only_errors=True)
            return
        list_price = None
        try:
            list_price = float(product.get("list_price") or product.get("listPrice") or 0)
        except Exception:
            list_price = 0
        final_list = max(list_price or 0, float(price))
        self._apply_price_update(
            product,
            quantity=quantity,
            sale_price=float(price),
            list_price=final_list,
            feedback_cb=feedback_cb,
        )

    def _apply_price_update(
        self,
        product: dict,
        *,
        quantity: int,
        sale_price: float,
        list_price: float,
        feedback_cb: Optional[Callable[[object], None]] = None,
    ):
        def _notify(status: Optional[str]) -> None:
            if feedback_cb:
                try:
                    feedback_cb(status)
                except Exception:
                    pass

        api_account_id = product.get("api_account_id")
        if api_account_id is None:
            _notify("error")
            MessageHandler.show(self, Result.fail("Şirket bilgisi bulunamadı."), only_errors=True)
            return

        platform = _normalized_account_platform(self._account_platform_map.get(api_account_id))
        if platform and platform != "TRENDYOL":
            _notify("error")
            MessageHandler.show(self, Result.fail("Bu işlem sadece Trendyol için destekleniyor."), only_errors=True)
            return

        barcode = _safe_str(product.get("barcode"))
        if not barcode:
            _notify("error")
            MessageHandler.show(self, Result.fail("Barkod bulunamadı."), only_errors=True)
            return

        res_creds = action_get_company_credentials([int(api_account_id)])
        if not res_creds.success:
            _notify("error")
            MessageHandler.show(self, res_creds, only_errors=True)
            return

        accounts = res_creds.data.get("accounts", []) or []
        if not accounts:
            _notify("error")
            MessageHandler.show(self, Result.fail("API bilgisi bulunamadı."), only_errors=True)
            return

        acc = accounts[0]
        api_key = acc[1]
        api_secret = acc[2]
        seller_id = acc[3]
        if not api_key or not api_secret or not seller_id:
            _notify("error")
            MessageHandler.show(self, Result.fail("API bilgisi eksik."), only_errors=True)
            return

        payload_items = [
            {
                "barcode": barcode,
                "quantity": int(max(0, quantity)),
                "salePrice": float(sale_price),
                "listPrice": float(list_price),
            }
        ]

        def _on_done(res: Result):
            if not res.success:
                _notify("error")
                MessageHandler.show(self, res, only_errors=True)
                return

            if (res.data or {}).get("pending"):
                _notify("pending")
                product_id = _safe_str(product.get("productId") or product.get("id"))
                if product_id:
                    action_update_listing_price_inventory_cache(
                        api_account_id=int(api_account_id),
                        product_id=product_id,
                        quantity=float(quantity),
                        list_price=float(list_price),
                        sale_price=float(sale_price),
                    )
                batch_request_id = (res.data or {}).get("batch_request_id")
                if batch_request_id:
                    self._schedule_batch_verify(
                        api_key=api_key,
                        api_secret=api_secret,
                        seller_id=seller_id,
                        barcode=barcode,
                        batch_request_id=str(batch_request_id),
                    )
                return

            _notify("success")
            product_id = _safe_str(product.get("productId") or product.get("id"))
            if product_id:
                action_update_listing_price_inventory_cache(
                    api_account_id=int(api_account_id),
                    product_id=product_id,
                    quantity=float(quantity),
                    list_price=float(list_price),
                    sale_price=float(sale_price),
                )

        self._price_worker = AsyncWorker(
            action_update_price_inventory_for_account,
            parent=self,
            kwargs={
                "api_key": api_key,
                "api_secret": api_secret,
                "seller_id": seller_id,
                "items": payload_items,
                "barcode": barcode,
                "wait_result": True,
                "poll_interval": 1.0,
                "max_wait": 300.0,
            },
        )
        self._price_worker.result_ready.connect(_on_done)
        self._price_worker.start()

    def _open_image_preview(self, product: dict):
        if not product:
            return
        path = _safe_str(product.get("image_path"))
        if not path:
            url = _safe_str(product.get("image_url"))
            if url:
                p = _image_cache_path(url)
                if p is not None:
                    path = str(p)

        if not path or not Path(path).exists():
            self._image_preview.set_image_path(None)
        else:
            self._image_preview.set_image_path(path)
        self._image_preview.show()
        self._image_preview.raise_()
        self._image_preview.activateWindow()

    def _schedule_batch_verify(
        self,
        *,
        api_key: str,
        api_secret: str,
        seller_id: str,
        barcode: str,
        batch_request_id: str,
    ) -> None:
        key = f"{barcode}:{batch_request_id}"
        now = time.monotonic()
        deadline = self._pending_verifications.get(key)
        if deadline and now < deadline:
            return
        if not deadline:
            deadline = now + (PENDING_UI_MS / 1000.0)
        self._pending_verifications[key] = deadline

        def _run_check():
            worker = AsyncWorker(
                action_check_batch_request_for_account,
                parent=self,
                kwargs={
                    "api_key": api_key,
                    "api_secret": api_secret,
                    "seller_id": seller_id,
                    "batch_request_id": batch_request_id,
                    "barcode": barcode,
                },
            )

            def _on_verify(res: Result):
                self._pending_verifications.pop(key, None)
                if not res.success:
                    if (res.data or {}).get("pending"):
                        remaining = max(5, int((deadline - time.monotonic()) * 1000))
                        self.list_widget._set_inline_state(barcode, "pending", min(PENDING_UI_MS, remaining))
                        if time.monotonic() < deadline:
                            self._schedule_batch_verify(
                                api_key=api_key,
                                api_secret=api_secret,
                                seller_id=seller_id,
                                barcode=barcode,
                                batch_request_id=batch_request_id,
                            )
                        return
                    self.list_widget._set_inline_state(barcode, "error", 10000)
                    MessageHandler.show(self, res, only_errors=True)
                    return

                if (res.data or {}).get("pending"):
                    remaining = max(5, int((deadline - time.monotonic()) * 1000))
                    self.list_widget._set_inline_state(barcode, "pending", min(PENDING_UI_MS, remaining))
                    if time.monotonic() < deadline:
                        self._schedule_batch_verify(
                            api_key=api_key,
                            api_secret=api_secret,
                            seller_id=seller_id,
                            barcode=barcode,
                            batch_request_id=batch_request_id,
                        )
                    return

                self.list_widget._set_inline_state(barcode, "success", 30000)
                self.apply_db_filters()

            worker.result_ready.connect(_on_verify)
            worker.start()

        QTimer.singleShot(VERIFY_RETRY_MS, _run_check)

    def _enqueue_images(self, items: list[dict]):
        # reset queue to avoid downloading images from previous pages
        self._image_queue.clear()
        tasks: list[tuple[str, str]] = []
        for it in items or []:
            url = it.get("image_url") or ""
            if not url:
                continue
            if url in self._image_inflight:
                continue
            path = it.get("image_path")
            if not path:
                p = _image_cache_path(url)
                if p is None:
                    continue
                path = str(p)
                it["image_path"] = path
            if Path(path).exists():
                self._image_cache_map[url] = path
                continue
            self._image_inflight.add(url)
            tasks.append((url, path))

        if tasks:
            self._image_queue.extend(tasks)
            self._start_image_worker()

    def _start_image_worker(self):
        if self._image_worker and self._image_worker.isRunning():
            return
        if not self._image_queue:
            return
        tasks = list(self._image_queue)
        self._image_queue.clear()
        self._image_worker = ImageDownloadWorker(tasks, parent=self)
        self._image_worker.image_saved.connect(self._on_image_saved)
        self._image_worker.finished.connect(self._start_image_worker)
        self._image_worker.start()

    def _cancel_image_worker(self):
        if self._image_worker and self._image_worker.isRunning():
            try:
                self._image_worker.stop()
            except Exception:
                pass
        self._image_inflight.clear()
        self._image_queue.clear()

    def _on_image_saved(self, url: str, path: str):
        self._image_cache_map[url] = path
        if url in self._image_inflight:
            self._image_inflight.discard(url)
        for it in self.list_widget.filtered_items or []:
            if it.get("image_url") == url:
                it["image_path"] = path
        self.list_widget.update_icon_for_url(url, path)


class ListingWorkspaceBehaviorAdapter(RecordWorkspaceBehaviorAdapter):
    def __init__(self, window: "ListingManagerWindow") -> None:
        self.tab = window
        self.new_panel_id = "new"
        self.edit_panel_id = "edit"
        self.bulk_panel_id = "bulk"
        self.clear_active_on_close = True

    @property
    def list_widget(self):
        return self.tab.list_widget

    def get_mode(self) -> str:
        mode = str(self.tab._selection_mode or "detail")
        return "edit" if mode == "detail" else mode

    def set_mode(self, mode: str) -> None:
        self.tab._selection_mode = "detail" if mode == "edit" else mode

    def get_active_key(self) -> str:
        return self.tab._workspace_record_key(getattr(self.tab, "_active_listing_record", None))

    def get_selected_count(self) -> int:
        return int(self.tab._selected_listing_count() or 0)

    def get_record_key(self, record) -> str:
        return self.tab._workspace_record_key(record)

    def prepare_new(self) -> None:
        return

    def activate_record(self, record) -> None:
        self.tab._active_listing_record = dict(record or {}) if isinstance(record, dict) else None
        self.tab.list_widget.set_active_record(record)

    def clear_active(self) -> None:
        self.tab._active_listing_record = None
        self.tab.list_widget.clear_active_item()
        self.tab._on_listing_behavior_cleared()

    def apply_selection(self, record, selected: bool) -> None:
        record_key = self.get_record_key(record)
        self.tab.list_widget.set_selected_for_record_key(record_key, bool(selected))

    def set_select_button_guide_enabled(self, enabled: bool) -> None:
        self.tab.list_widget.set_select_button_guide_enabled(bool(enabled))

    def on_state_changed(self) -> None:
        self.tab._sync_listing_workspace_state()

    def edit_guidance_message(self) -> str:
        return "Duzenlemek icin sagdaki listeden bir ilan sec."

    def bulk_guidance_message(self) -> str:
        return "Toplu işlem için sağdaki listeden en az bir ilanı SEÇ ile işaretle."




