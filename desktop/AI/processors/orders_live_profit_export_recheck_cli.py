from __future__ import annotations

import argparse
from datetime import date
import hashlib
from pathlib import Path
import tempfile
from typing import Any, Sequence

from openpyxl import load_workbook

from AI.constants.paths import AI_PACKAGE_ROOT
from AI.processors.cloud_runtime import build_provider_chat_client
from AI.processors.live_evidence import write_live_evidence_file
from AI.processors.orders_live_source_recheck_cli import (
    DEFAULT_OPENROUTER_MODEL,
    _compact_text,
    _load_provider_record,
    _runtime_plan,
    _visible_answer_has_mojibake,
)
from AI.processors.task_runtime import (
    resume_agent_task_with_approval,
    run_agent_task,
)
from AI.processors.tools import orders_full_context, recent_orders
from settings import DB_NAME


DEFAULT_OUTPUT_PATH = (
    AI_PACKAGE_ROOT
    / "docs"
    / "live_llm_transcripts"
    / "source-db-live-recheck"
    / f"live-source-db-profit-export-recheck-{date.today().isoformat()}.json"
)


def _orders_from_data(data: dict[str, Any]) -> list[dict[str, Any]]:
    return [
        dict(row)
        for row in list(data.get("orders") or [])
        if isinstance(row, dict)
    ]


def _select_source_seed(*, db_name: str = DB_NAME) -> dict[str, Any]:
    recent = recent_orders(db_name=db_name, period="all", limit=20)
    if not recent.success:
        raise RuntimeError(recent.message)
    for row in _orders_from_data(dict(recent.data or {})):
        order_number = _compact_text(row.get("order_number"))
        if not order_number:
            continue
        context_result = orders_full_context(
            db_name=db_name,
            query=order_number,
            limit=1,
        )
        if not context_result.success:
            continue
        contexts = [
            dict(item)
            for item in list((context_result.data or {}).get("contexts") or [])
            if isinstance(item, dict)
        ]
        exact = [
            item
            for item in contexts
            if _compact_text(
                dict(item.get("order") or {}).get("order_number")
            )
            == order_number
        ]
        if len(exact) != 1 or not list(exact[0].get("items") or []):
            continue
        return {
            "order_number": order_number,
            "status": _compact_text(row.get("status")),
            "item_count": len(list(exact[0].get("items") or [])),
        }
    raise RuntimeError(
        "Source DB has no recent order with item rows for profit export."
    )


def _state_from_result(result: Any) -> dict[str, Any]:
    return dict((result.data or {}).get("state") or {})


def _pending_action(state: dict[str, Any]) -> dict[str, Any]:
    pending = [
        dict(row)
        for row in list(state.get("pending_actions") or [])
        if isinstance(row, dict)
    ]
    if len(pending) != 1:
        raise RuntimeError(
            f"Expected one pending profit-export action, got {len(pending)}."
        )
    return pending[0]


def _action_order_number(action: dict[str, Any]) -> str:
    arguments = dict(action.get("arguments") or {})
    payload = dict(arguments.get("payload") or {})
    return _compact_text(
        payload.get("order_number")
        or payload.get("order")
        or arguments.get("order_number")
        or arguments.get("order")
    )


def _action_type(action: dict[str, Any]) -> str:
    return _compact_text(
        action.get("operation")
        or action.get("action_type")
    )


def _completed_tool_names(state: dict[str, Any]) -> list[str]:
    names: list[str] = []
    for step in list(state.get("completed_steps") or []):
        if not isinstance(step, dict):
            continue
        names.extend(
            _compact_text(name)
            for name in list(step.get("tool_names") or [])
            if _compact_text(name)
        )
    return names


def _workbook_evidence(
    output_path: Path,
    *,
    expected_order_number: str,
) -> dict[str, Any]:
    raw = output_path.read_bytes()
    workbook = load_workbook(output_path, read_only=True, data_only=False)
    try:
        sheet_names = list(workbook.sheetnames)
        values = [
            value
            for sheet in workbook.worksheets
            for row in sheet.iter_rows(values_only=True)
            for value in row
            if value not in (None, "")
        ]
    finally:
        workbook.close()
    exact_order_present = any(
        _compact_text(value) == expected_order_number
        for value in values
    )
    return {
        "file_name": output_path.name,
        "size_bytes": len(raw),
        "sha256": hashlib.sha256(raw).hexdigest(),
        "sheet_names": sheet_names,
        "expected_sheets_present": {
            "Ozet",
            "Satis",
            "Maliyetler",
        }.issubset(set(sheet_names)),
        "exact_order_number_present": exact_order_present,
    }


def run_provider_profit_export_recheck(
    provider_code: str,
    *,
    seed: dict[str, Any],
    db_name: str = DB_NAME,
    timeout: float = 180.0,
    openrouter_model: str = DEFAULT_OPENROUTER_MODEL,
) -> dict[str, Any]:
    record = _load_provider_record(
        provider_code,
        openrouter_model=openrouter_model,
    )
    client = build_provider_chat_client(record, timeout=timeout)
    if client is None:
        raise RuntimeError(f"{provider_code} client could not be built")

    order_number = _compact_text(seed.get("order_number"))
    with tempfile.TemporaryDirectory(
        prefix="orderscout-live-profit-export-"
    ) as temp_dir:
        root = Path(temp_dir)
        output_path = root / f"{provider_code}-order-profit.xlsx"
        audit_path = root / f"{provider_code}-audit.jsonl"
        goal = (
            f'{order_number} numarali siparisin kar Excel raporunu '
            f'"{output_path.as_posix()}" dosyasina aktar'
        )

        prepared = run_agent_task(
            goal,
            planner_client=client,
            chat_client=client,
            db_name=db_name,
            persist=False,
            runtime_context={
                "provider_code": provider_code,
                "provider_id": record.get("pk"),
                "actual_model": record.get("default_model"),
                "reasoning_effort": "auto",
            },
        )
        prepared_state = _state_from_result(prepared)
        action = _pending_action(prepared_state)
        file_absent_before_approval = not output_path.exists()
        prepare_answer = _compact_text(
            (prepared.data or {}).get("answer")
            or prepared.message
        )

        resumed = resume_agent_task_with_approval(
            prepared_state,
            approved_by="live-source-db-evaluator",
            planner_client=client,
            chat_client=client,
            db_name=db_name,
            audit_path=str(audit_path),
            idempotency_key=(
                f"source-db-profit-{provider_code}-"
                f"{date.today().isoformat()}"
            ),
            persist=False,
        )
        resumed_state = _state_from_result(resumed)
        execute_data = dict((resumed.data or {}).get("execute") or {})
        verify = dict(execute_data.get("verify") or {})
        final_answer = _compact_text(
            (resumed.data or {}).get("answer")
            or resumed.message
        )
        workbook = (
            _workbook_evidence(
                output_path,
                expected_order_number=order_number,
            )
            if output_path.exists()
            else {}
        )
        audit_text = (
            audit_path.read_text(encoding="utf-8")
            if audit_path.exists()
            else ""
        )
        output_path.unlink(missing_ok=True)
        audit_path.unlink(missing_ok=True)

    progress = [
        dict(row)
        for row in list(resumed_state.get("progress_events") or [])
        if isinstance(row, dict)
    ]
    normalized_prepare_answer = prepare_answer.casefold()
    normalized_final_answer = final_answer.casefold()
    plan = _runtime_plan(client)
    checks = {
        "prepare_succeeded": bool(prepared.success),
        "llm_plan_selected_profit_export": (
            "orders.export_profit_excel.prepare"
            in _completed_tool_names(prepared_state)
        ),
        "prepare_uses_exact_source_order": (
            _action_order_number(action) == order_number
        ),
        "prepare_waits_for_approval": (
            prepared_state.get("status") == "pending_approval"
            and _action_type(action)
            == "orders.action_export_order_profit_excel"
        ),
        "prepare_did_not_write_file": file_absent_before_approval,
        "prepare_answer_present_and_clean": (
            bool(prepare_answer)
            and not _visible_answer_has_mojibake(prepare_answer)
        ),
        "prepare_answer_explains_exact_pending_action": (
            order_number in prepare_answer
            and "onay" in normalized_prepare_answer
            and "yapilmadi" in normalized_prepare_answer
            and "agent task runtime ara durumda durdu"
            not in normalized_prepare_answer
        ),
        "execute_succeeded": bool(resumed.success),
        "task_completed_after_approval": (
            resumed_state.get("status") == "complete"
            and not list(resumed_state.get("pending_actions") or [])
        ),
        "runtime_verify_succeeded": verify.get("verified") is True,
        "workbook_exists_and_nonempty": (
            int(workbook.get("size_bytes") or 0) > 0
        ),
        "workbook_has_expected_sheets": bool(
            workbook.get("expected_sheets_present")
        ),
        "workbook_contains_exact_order": bool(
            workbook.get("exact_order_number_present")
        ),
        "audit_contains_action": (
            "orders.action_export_order_profit_excel" in audit_text
        ),
        "final_answer_present_and_clean": (
            bool(final_answer)
            and not _visible_answer_has_mojibake(final_answer)
        ),
        "final_answer_names_verified_order_and_output": (
            order_number in final_answer
            and output_path.name in final_answer
            and "dogrulan" in normalized_final_answer
            and "agent gorevi tamamlandi"
            not in normalized_final_answer
        ),
        "temporary_sensitive_files_removed": (
            not output_path.exists()
            and not audit_path.exists()
        ),
        "progress_exposes_prepare_execute_verify": {
            "pending_approval",
            "execute_started",
            "execute_completed",
            "completed",
        }.issubset(
            {
                _compact_text(row.get("phase"))
                for row in progress
            }
        ),
    }
    return {
        "provider_code": provider_code,
        "requested_model": _compact_text(record.get("default_model")),
        "actual_model": _compact_text(
            plan.get("actual_model")
            or getattr(client, "model_name", "")
        ),
        "reasoning_effort": "auto",
        "fallback_used": bool(plan.get("fallback_used")),
        "goal": goal,
        "prepare": {
            "answer": prepare_answer,
            "task_status": prepared_state.get("status"),
            "tool_names": _completed_tool_names(prepared_state),
            "action_type": _action_type(action),
            "order_number": _action_order_number(action),
            "file_absent_before_approval": file_absent_before_approval,
        },
        "execute": {
            "answer": final_answer,
            "task_status": resumed_state.get("status"),
            "verify": verify,
            "workbook": workbook,
            "audit_action_present": (
                "orders.action_export_order_profit_excel" in audit_text
            ),
            "temporary_sensitive_files_removed": (
                not output_path.exists()
                and not audit_path.exists()
            ),
        },
        "progress_phases": [
            _compact_text(row.get("phase"))
            for row in progress
            if _compact_text(row.get("phase"))
        ],
        "checks": checks,
    }


def run_source_db_profit_export_recheck(
    *,
    provider_codes: Sequence[str] = ("gemini", "openrouter"),
    db_name: str = DB_NAME,
    timeout: float = 180.0,
    openrouter_model: str = DEFAULT_OPENROUTER_MODEL,
) -> dict[str, Any]:
    seed = _select_source_seed(db_name=db_name)
    providers: list[dict[str, Any]] = []
    for provider_code in provider_codes:
        try:
            providers.append(
                run_provider_profit_export_recheck(
                    provider_code,
                    seed=seed,
                    db_name=db_name,
                    timeout=timeout,
                    openrouter_model=openrouter_model,
                )
            )
        except Exception as exc:
            providers.append(
                {
                    "provider_code": provider_code,
                    "checks": {"completed": False},
                    "error": f"{type(exc).__name__}: {_compact_text(exc)}",
                }
            )
    return {
        "name": (
            "live source DB Orders profit export recheck "
            f"{date.today().isoformat()}"
        ),
        "artifact_type": (
            "orders_live_source_db_profit_export_execute_recheck"
        ),
        "created_at": date.today().isoformat(),
        "db": "source-app-db",
        "seed": seed,
        "safety": {
            "computer_use_used": False,
            "packaged_exe_used": False,
            "marketplace_write_performed": False,
            "db_write_performed": False,
            "file_output_performed": True,
            "temporary_output_removed_after_verify": True,
        },
        "providers": providers,
        "passed": (
            len(providers) == len(provider_codes)
            and all(
                bool(provider.get("checks"))
                and all(dict(provider.get("checks") or {}).values())
                for provider in providers
            )
        ),
    }


def main(argv: Sequence[str] | None = None) -> int:
    parser = argparse.ArgumentParser(
        description=(
            "Run real source-DB Orders profit-export tasks through configured "
            "cloud providers, approve them, and verify workbook content."
        )
    )
    parser.add_argument(
        "--provider",
        action="append",
        choices=("gemini", "openrouter"),
        dest="providers",
        help="Provider to run. Repeat for both; defaults to both.",
    )
    parser.add_argument(
        "--openrouter-model",
        default=DEFAULT_OPENROUTER_MODEL,
    )
    parser.add_argument("--timeout", type=float, default=180.0)
    parser.add_argument("--output", default=str(DEFAULT_OUTPUT_PATH))
    args = parser.parse_args(list(argv) if argv is not None else None)

    payload = run_source_db_profit_export_recheck(
        provider_codes=tuple(
            args.providers or ("gemini", "openrouter")
        ),
        timeout=max(10.0, float(args.timeout)),
        openrouter_model=args.openrouter_model,
    )
    target = write_live_evidence_file(Path(args.output), payload)
    print(f"artifact={target}")
    print(f"passed={bool(payload.get('passed'))}")
    for provider in list(payload.get("providers") or []):
        print(
            "provider="
            f"{provider.get('provider_code')} "
            f"model={provider.get('actual_model') or '-'} "
            f"passed={all(dict(provider.get('checks') or {}).values())}"
        )
        if provider.get("error"):
            print(f"error={provider.get('error')}")
    return 0 if bool(payload.get("passed")) else 1


if __name__ == "__main__":
    raise SystemExit(main())
