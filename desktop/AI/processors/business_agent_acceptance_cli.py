"""Live cloud acceptance through the real Chat source UI on synthetic business data.

Run this file in a fresh process. Production is accessed only by a read-only query
of ai_provider_account; application settings and all business records are isolated.
"""
from __future__ import annotations

import argparse
from copy import deepcopy
from contextlib import ExitStack, closing
from datetime import datetime, timezone
from dataclasses import asdict, is_dataclass
from functools import wraps
import hashlib
import ipaddress
import json
import os
from pathlib import Path
import shutil
import sqlite3
import sys
import tempfile
import time
from typing import Any
from urllib.parse import urlsplit
from unittest.mock import patch
from uuid import UUID

if __package__:
    from AI.processors.business_acceptance_fixtures import WORKFLOW_ACCEPTANCE_SPEC, workflow_save_prompt
else:
    from business_acceptance_fixtures import WORKFLOW_ACCEPTANCE_SPEC, workflow_save_prompt


SCENARIOS = {
    "child_stock_report_prepare": {
        "prompt": "Stok incelemesini bağımsız bir alt göreve devret: alt görev seçili modelle kendi planını kurup güncel stok kaynağını okusun; dosya veya işlem hazırlamasın. Kayıtlı gözlem kaynaklarını ana görevde ayrıca açıp incele ve ürün kodu, miktarı, birimi ile dövizlerine ayrılmış stok değerini çok sayfalı Excel raporuna hazırla. TRY ile USD'yi toplama; verinin alt görevin hangi çalışmasında gözlendiğini belirt ve ana görevin tekrar güncel stok okuduğunu söyleme. Dosyayı henüz oluşturma, tam raporu yeni onayıma sun.",
        "phase": "prepare", "child_stock_report": True,
        "required_operations": ["agent.delegate", "artifacts.read", "reports.build_snapshot", "reports.xlsx.prepare"],
    },
    "child_stock_report_approve": {"prompt": "Onaylıyorum.", "phase": "approve", "child_stock_export": True},
    "child_stock_report_followup": {
        "prompt": "Alt araştırmadan gelen rapor gerçekten üretildi mi? Hangi alt çalışma ve gözlemden geldiğini, dövizlerin ve kaynak miktarlarının korunmasını ve dosyanın nasıl doğrulandığını açıkla. Kaynak yeniden okunmadıysa yeni güncel stok okuması yapıldığını söyleme.",
        "phase": "read", "required_goal_complete": True,
    },
    "workflow_save_prepare": {
        "prompt": workflow_save_prompt(), "phase": "prepare", "effect_kind": "workflow_save",
        "required_operations": ["workflows.save.prepare"], "workflow_spec_fixture": True,
    },
    "workflow_save_pending": {
        "prompt": "Onay bekleyen yönergenin tam talimatı, amaç şablonu ve parametre sınırları neler? Henüz kaydedildi mi ve çalıştırıldı mı?",
        "phase": "pending",
    },
    "workflow_save_approve": {
        "prompt": "Onaylıyorum.", "phase": "approve", "effect_kind": "workflow_save",
    },
    "workflow_run": {
        "prompt": "Kayıtlı Pilot Stok İncelemesi yönergesini Pilot Tekrar Stok Raporu başlığıyla yeniden çalıştır.",
        "phase": "prepare", "workflow_invocation": {"report_title": "Pilot Tekrar Stok Raporu"},
        "required_operations": ["stock.inventory_overview", "reports.build_snapshot", "reports.xlsx.prepare"],
        "workflow_fresh_sources": True,
    },
    "workflow_schedule_review": {
        "prompt": "Pilot Stok İncelemesi yönergesini Pilot Planlı Stok Raporu başlığıyla, ilk çalışma şimdi olacak şekilde her saat çalışacak plana hazırla; planı henüz kaydetme.",
        "phase": "schedule_review", "schedule_stage": "review",
        "schedule_parameters": {"report_title": "Pilot Planlı Stok Raporu"},
    },
    "workflow_schedule_approve": {
        "prompt": "Gözden geçirdiğim bu tam zamanlama planını onaylıyorum.",
        "phase": "schedule_approve", "schedule_stage": "approve",
    },
    "workflow_schedule_run": {
        "prompt": "Onaylanan planın zamanı geldi; seçili modelle güncel stok raporunu hazırla ve dosya oluşturmadan yeniden onayımı bekle.",
        "phase": "prepare", "schedule_stage": "run", "workflow_fresh_sources": True,
        "required_operations": ["stock.inventory_overview", "reports.build_snapshot", "reports.xlsx.prepare"],
    },
    "workflow_report_approve": {"prompt": "Onaylıyorum.", "phase": "approve"},
    "workflow_report_followup": {
        "prompt": "Kayıtlı yönergeyle yeni rapor gerçekten üretildi mi? Kaynaklar bu çalışmada yeniden okundu mu, dövizler ve miktarlar nasıl doğrulandı?",
        "phase": "read", "required_goal_complete": True,
    },
    "business_research": {
        "prompt": "Pilot Ana mağazamda son 7 günün satışlarını önceki 7 günle karşılaştır; stokta bağlı parayı ve kayıtlı finans hareketlerini de araştır. Sonuçları dövizleri karıştırmadan açıkla. Ortak işletme/ürün bağlantısı veya gerçekleşmiş banka bakiyesi kanıtı yoksa uydurma; eksikleri ve sonraki araştırma adımlarını belirt.",
        "required_domains": ["orders", "stock", "finance"], "phase": "read", "fixture_oracle": True,
        "required_goal_complete": True,
    },
    "general_report_prepare": {
        "prompt": "Pilot Ana mağazasının son 7 gün ve önceki 7 gün sipariş sayılarını ve tahmini cirolarını, stokta bağlı değeri ve kayıtlı finans hareketlerini araştırıp benim için çok sayfalı bir Excel raporu hazırla. Satış dönemlerini ayrı satırlarda, stok ve finansı ayrı sayfalarda göster; her dövizi ayrı tut ve kaynak kapsamını belirt. Raporun düzenini bu ihtiyaca göre seç. İşletme/mağaza bağlantısını veya gerçekleşmiş banka bakiyesini kanıt olmadan çıkarma; varsayımları ve önerileri kaynak verisinden ayrı işaretle. Dosyayı henüz kaydetme, oluşturulacak raporu onayıma sun.",
        "required_operations": ["reports.build_snapshot", "reports.xlsx.prepare"],
        "phase": "prepare", "fixture_oracle": True, "general_report_sources": True,
    },
    "general_report_pending": {
        "prompt": "Onay bekleyen Excel raporunun kaynakları, mağaza kapsamı, dönemleri ve dövizleri neler? Dosya henüz oluştu mu?",
        "phase": "pending",
    },
    "general_report_approve": {
        "prompt": "Onaylıyorum.", "phase": "approve",
    },
    "general_report_followup": {
        "prompt": "Bu rapor gerçekten oluşturuldu mu, dosyadaki satış dönemleri, stok ve finans kaynakları neler ve sonucu nasıl doğruladın? Gerçekleşmiş banka bakiyesi veya mağazaya bağlanmış stok/finans kanıtı olduğunu söyleme.",
        "phase": "read", "required_goal_complete": True,
    },
    "attachment_json_prepare": {
        "prompt": "Sohbete eklediğim pilot-sales.json dosyasını incele ve benim için bir Excel raporu hazırla. Her kaydın sku, quantity, unit, revenue, currency ve discount alanlarını kaynak değerleriyle tek tabloda koru; sayıları metne, boş discount değerlerini sıfıra çevirme. TRY ve USD'yi toplamadan kısa bir değerlendirme yap; varsa yorumlarını kaynak tablosundan ayrı tut. Rapor düzenini ihtiyaca göre seç. Dosyayı henüz kaydetme, raporu onayıma sun.",
        "phase": "prepare", "attachment_formats": ["json"], "attachment_report": True,
        "required_operations": ["files.read", "reports.build_snapshot", "reports.xlsx.prepare"],
    },
    "attachment_xlsx_prepare": {
        "prompt": "Sohbete eklediğim pilot-sales.xlsx dosyasını incele ve benim için yeni bir Excel raporu hazırla. Her kaydın sku, quantity, unit, revenue, currency ve discount alanlarını kaynak değerleriyle tek tabloda koru; sayıları metne, boş discount değerlerini sıfıra çevirme. TRY ve USD'yi toplamadan kısa bir değerlendirme yap; varsa yorumlarını kaynak tablosundan ayrı tut. Rapor düzenini ihtiyaca göre seç. Dosyayı henüz kaydetme, raporu onayıma sun.",
        "phase": "prepare", "attachment_formats": ["xlsx"], "attachment_report": True,
        "required_operations": ["files.read", "reports.build_snapshot", "reports.xlsx.prepare"],
    },
    "attachment_pending": {
        "prompt": "Onay bekleyen dosya raporunda hangi ek dosya, sütunlar, kayıtlar ve dövizler bulunuyor? Kaynakta boş discount değerleri nasıl korunacak? Dosya henüz oluştu mu?",
        "phase": "pending",
    },
    "attachment_approve": {"prompt": "Onaylıyorum.", "phase": "approve", "attachment_export": True},
    "attachment_followup": {
        "prompt": "Ek dosyadan istediğim rapor gerçekten oluşturuldu mu? Hangi kaynak dosyadan geldiğini, kayıtların birim ve dövizlerinin korunduğunu ve boş discount alanlarının sonucunu nasıl doğruladığını söyle.",
        "phase": "read", "required_goal_complete": True,
    },
    "scope_correction": {
        "prompt": "Yan mağazayı dahil etme, yalnız Pilot Ana mağazasını kastettim. Satış karşılaştırmasını bu kapsamda düzelt; stok ve finansın bu mağazaya bağlanamadığı noktaları ayrı söyle.",
        "phase": "read", "primary_order_scope": True,
    },
    "stock_followup": {
        "prompt": "Kupa ve filtre stokunun değerini kendi dövizinde ayrı göster; TRY ile USD'yi toplama. Hangi ürünün azaldığını mevcut bakiyeyle söyle, hareket okumadıysan azalma nedeni çıkarma.",
        "required_domains": ["stock"], "phase": "read",
    },
    "report_prepare": {
        "prompt": "Pilot Ana mağazasının son 7 gündeki ürün satışlarını Excel raporuna dönüştürmek için hazırla. Dosyayı henüz üretme; tam kapsamı ve kaydetme işlemini benim onayıma sun.",
        "phase": "prepare", "required_domains": ["orders"],
    },
    "pending_scope": {
        "prompt": "Onay bekleyen raporda hangi mağaza, dönem ve ürünler var; dosya henüz oluştu mu?",
        "phase": "pending",
    },
    "cancel_report": {"prompt": "Vazgeç, bu raporu kaydetme.", "phase": "cancel"},
    "public_research": {
        "prompt": "Trendyol'un kamuya açık satıcı yardım kaynaklarında kargo etiketindeki barkodun ne için kullanıldığını araştır. Sayfaları açarak doğrula ve yalnız kaynakta geçen bilgiyle kısa bir kaynaklı özet ver. Özel sipariş veya müşteri bilgilerini internete gönderme.",
        "phase": "read", "required_domains": ["web"], "requires_public_web": True, "required_goal_complete": True,
        "required_operations": ["web.search", "web.open"],
    },
    "research_report_prepare": {
        "prompt": "Az önce gerçekten açtığın kaynakların başlıklarını, bağlantılarını ve kanıtladıkları bilgileri Excel araştırma raporuna hazırla. Kaynak veya olgu ekleme; kaydetmeden önce onayımı bekle.",
        "phase": "prepare", "requires_public_web": True,
    },
    "approve_report": {
        "prompt": "Onaylıyorum; yalnız şu anda onay bekleyen araştırma raporunu üret.",
        "phase": "approve", "requires_public_web": True,
    },
    "report_followup": {
        "prompt": "Üretilen dosyada hangi kaynakların bulunduğunu ve sonucu nasıl doğruladığını söyle. Üretilmediyse açıkça belirt.",
        "phase": "read", "requires_public_web": True, "required_goal_complete": True,
    },
}

LOCAL_EXPORT_ACTIONS = frozenset({
    "orders.action_export_report_snapshot_xlsx", "ai.action_export_report_snapshot_xlsx",
})


class ObservedCloudClient:
    """Record model responses, never prompts; preserve the actual live client."""

    def __init__(self, client: Any, sink: Any, *, before_child_callback: Any = None):
        object.__setattr__(self, "_client", client)
        object.__setattr__(self, "_sink", sink)
        object.__setattr__(self, "_ordinal", 0)
        object.__setattr__(self, "_before_child_callback", before_child_callback)
        object.__setattr__(self, "_retained_child_ids", set())

    def __setattr__(self, name: str, value: Any) -> None:
        if name.startswith("_"):
            object.__setattr__(self, name, value)
        else:
            setattr(self._client, name, value)

    def __getattr__(self, name: str) -> Any:
        value = getattr(self._client, name)
        if name in {"invoke_structured", "invoke_tools", "invoke_json_text", "invoke_final_text"} and callable(value):
            @wraps(value)
            def observed(*args, **kwargs):
                return self._invoke(name, value, *args, **kwargs)
            return observed
        return value

    def __call__(self, *args, **kwargs):
        return self._invoke("text", self._client, *args, **kwargs)

    def _invoke(self, method: str, callback: Any, *args, **kwargs) -> Any:
        self._ordinal += 1
        row: dict[str, Any] = {"invocation": self._ordinal, "method": method,
                               "started_at": datetime.now(timezone.utc).isoformat()}
        from AI.processors.child_task_runtime import current_child_policy, current_child_execution_id
        policy = current_child_policy()
        if policy is not None:
            row["delegated_child"] = {"task_id": current_child_execution_id(), "parent_task_id": policy["parent_task_id"],
                                      "depth": policy["depth"], "read_only": True, "effect_authority": False}
        started = time.monotonic()
        try:
            if policy is not None and self._before_child_callback is not None:
                child_id = current_child_execution_id()
                if child_id not in self._retained_child_ids:
                    row["durable_checkpoint_retention"] = self._before_child_callback(child_id, policy["parent_task_id"])
                    self._retained_child_ids.add(child_id)
            result = callback(*args, **kwargs)
            if isinstance(result, str):
                try:
                    row["response"] = json.loads(result)
                except (ValueError, TypeError):
                    row["response"] = result
            elif is_dataclass(result):
                row["response"] = asdict(result)
            elif callable(getattr(result, "model_dump", None)):
                row["response"] = result.model_dump()
            elif isinstance(result, (dict, list, int, float, bool)) or result is None:
                row["response"] = result
            else:
                row["response"] = str(result)
            row["response_received"] = True
            return result
        except Exception as exc:
            row["response_received"] = False
            row["exception_type"] = type(exc).__name__
            raise
        finally:
            plan = dict(getattr(self._client, "resource_plan", {}) or {})
            row["contract"] = plan.get("last_invocation_contract") or method
            row["requested_model"] = plan.get("requested_model") or getattr(self._client, "model_name", "")
            row["actual_model"] = plan.get("actual_model") or plan.get("last_actual_model") or getattr(self._client, "model_name", "")
            row["fallback_used"] = bool(plan.get("fallback_used") or plan.get("any_model_fallback_used"))
            row["duration_seconds"] = round(time.monotonic() - started, 3)
            self._sink(row)


def load_provider_readonly(source_db: str | Path, provider_id: int, model: str) -> dict[str, Any]:
    """Resolve credentials in memory without the provider migration/upsert path."""
    database = Path(source_db).resolve()
    with closing(sqlite3.connect(database.as_uri() + "?mode=ro", uri=True)) as connection:
        connection.row_factory = sqlite3.Row
        row = connection.execute(
            "SELECT pk, provider_code, label, endpoint_url, api_key, default_model, is_enabled "
            "FROM ai_provider_account WHERE pk = ?", (int(provider_id),),
        ).fetchone()
    if row is None or not row["is_enabled"]:
        raise ValueError("Enabled provider account not found.")
    from Core.processors.credential_store import is_credential_reference, resolve_secret
    secret = str(row["api_key"] or "")
    if is_credential_reference(secret):
        secret = resolve_secret(secret)
    if not secret:
        raise ValueError("Selected provider has no usable credential.")
    return {"id": int(row["pk"]), "provider_code": row["provider_code"],
            "label": row["label"], "endpoint_url": row["endpoint_url"],
            "api_key": secret, "default_model": model, "is_enabled": True,
            "allow_model_fallback": False}


def acceptance_provider_endpoint(record: dict[str, Any]) -> str:
    """Use the same configured endpoint/default owner as the provider adapter."""
    from AI.processors.cloud_runtime import DEFAULT_BASE_URLS
    endpoint = str(record.get("endpoint_url") or DEFAULT_BASE_URLS.get(str(record.get("provider_code") or "").casefold(), ""))
    if not endpoint:
        raise ValueError("The selected provider has no configured or canonical default endpoint.")
    return endpoint


def secret_safe(value: Any, secrets: tuple[str, ...]) -> Any:
    if isinstance(value, dict):
        output = {str(key): secret_safe(item, secrets) for key, item in value.items()
                  if str(key).casefold() not in {"api_key", "api_secret", "authorization", "token", "credentials", "capability_token"}}
        if any(str(key).casefold() == "capability_token" for key in value):
            output["private_capability_redacted"] = True
        return output
    if isinstance(value, list):
        return [secret_safe(item, secrets) for item in value]
    if isinstance(value, tuple):
        return [secret_safe(item, secrets) for item in value]
    if isinstance(value, str):
        for secret in secrets:
            if secret:
                value = value.replace(secret, "[REDACTED_PROVIDER_CREDENTIAL]")
        return value
    return value


def local_export_approval_problems(action: dict[str, Any], isolation_root: str | Path) -> list[str]:
    problems = []
    if action.get("action_type") not in LOCAL_EXPORT_ACTIONS:
        problems.append("approval_action_not_allowlisted_local_export")
    if action.get("external_write") or action.get("destructive"):
        problems.append("approval_has_external_or_destructive_effect")
    arguments = action.get("arguments") or {}
    payload = arguments.get("payload") or action.get("safe_payload") or {}
    output = payload.get("output_path") or arguments.get("output_path")
    if not output:
        problems.append("approval_missing_output_path")
    else:
        try:
            Path(output).resolve().relative_to(Path(isolation_root).resolve())
        except (ValueError, TypeError):
            problems.append("approval_output_outside_isolation")
    return problems


def local_acceptance_approval_problems(action: dict[str, Any], isolation_root: str | Path) -> list[str]:
    """Allow only exact local XLSX exports or immutable synthetic workflow saves."""
    if action.get("action_type") != "ai.action_save_workflow_spec":
        return local_export_approval_problems(action, isolation_root)
    problems = []
    if action.get("external_write") or action.get("destructive"):
        problems.append("approval_has_external_or_destructive_effect")
    arguments = action.get("arguments") or {}
    payload = arguments.get("payload") or action.get("safe_payload") or {}
    try:
        from AI.processors.workflow_store import _validated_action, workflow_store_path
        if set(payload) != {"proposal"}:
            raise ValueError("Unexpected workflow effect payload.")
        _validated_action(payload["proposal"])
        workflow_store_path().resolve().relative_to(Path(isolation_root).resolve())
    except Exception as exc:
        problems.append("workflow_approval_boundary_rejected:" + type(exc).__name__)
    return problems


def workflow_proposal_problems(action: dict[str, Any]) -> list[str]:
    try:
        from AI.processors.workflow_store import _validated_action
        if action.get("action_type") != "ai.action_save_workflow_spec":
            return ["workflow_save_action_missing"]
        proposal = _validated_action(((action.get("arguments") or {}).get("payload") or {}).get("proposal") or {})
        expected = WORKFLOW_ACCEPTANCE_SPEC
        actual = proposal["spec"]
        # Descriptive parameter labels may be authored by the model; the user's
        # exact instruction text and every constraint must remain unchanged.
        problems = []
        for key in ("name", "instructions", "objective_template", "source_policy", "effect_policy"):
            if actual[key] != expected[key]:
                problems.append("workflow_user_contract_changed:" + key)
        schema = actual["parameter_schema"]
        if (set(schema["properties"]) != {"report_title"} or schema["required"] != ["report_title"]
                or schema["additionalProperties"] is not False):
            problems.append("workflow_parameter_contract_changed")
        parameter = schema["properties"].get("report_title") or {}
        if parameter.get("type") != "string" or parameter.get("maxLength") != 80:
            problems.append("workflow_parameter_constraints_changed")
        if set(parameter) - {"type", "maxLength", "description"}:
            problems.append("workflow_unrequested_parameter_constraints")
        return problems
    except Exception as exc:
        return ["workflow_proposal_invalid:" + type(exc).__name__]


class NetworkBoundary:
    """Provider inference/catalog plus optional public GET; no marketplace API."""

    def __init__(self, provider_origin: str, *, allow_public_web: bool = False):
        self.provider_origin = urlsplit(provider_origin).netloc.casefold()
        self.allow_public_web = allow_public_web
        self.events: list[dict[str, Any]] = []

    def check(self, method: str, url: str, headers: Any = None) -> None:
        parsed = urlsplit(str(url))
        origin = parsed.netloc.casefold()
        host = (parsed.hostname or "").casefold()
        method = str(method).upper()
        provider = parsed.scheme == "https" and origin == self.provider_origin
        private_api = host.startswith(("api.", "apigw.", "stageapi.", "stageapigw."))
        private_host = host in {"localhost", "127.0.0.1", "::1"} or host.endswith(".local")
        try:
            private_host = private_host or not ipaddress.ip_address(host).is_global
        except ValueError:
            pass
        auth = any(str(key).casefold() in {"authorization", "x-api-key"} for key in (headers or {}))
        public = (self.allow_public_web and method == "GET" and parsed.scheme == "https"
                  and not private_api and not private_host and not auth)
        if not (provider or public):
            self.events.append({"method": method, "host": host, "allowed": False})
            raise RuntimeError("Acceptance network boundary rejected a non-authorized request.")
        self.events.append({"method": method, "host": host, "allowed": True,
                            "kind": "provider" if provider else "public_read"})

    def install(self, stack: ExitStack) -> None:
        import requests
        import urllib.request
        original_send = requests.Session.send
        original_open = urllib.request.OpenerDirector.open
        boundary = self

        def send(session, request, *args, **kwargs):
            boundary.check(request.method, request.url, request.headers)
            return original_send(session, request, *args, **kwargs)

        def open_url(opener, request, *args, **kwargs):
            url = request.full_url if hasattr(request, "full_url") else str(request)
            body = kwargs.get("data") if "data" in kwargs else (args[0] if args else None)
            method = request.get_method() if hasattr(request, "get_method") else ("POST" if body is not None else "GET")
            headers = dict(request.header_items()) if hasattr(request, "header_items") else {}
            boundary.check(method, url, headers)
            return original_open(opener, request, *args, **kwargs)

        stack.enter_context(patch.object(requests.Session, "send", send))
        stack.enter_context(patch.object(urllib.request.OpenerDirector, "open", open_url))


def operation_names(turn: dict[str, Any]) -> list[str]:
    from AI.processors.business_acceptance_evidence import completed_observation_receipts
    names = []
    for row in list(turn.get("tool_calls") or []) + [turn.get("tool_call") or {}]:
        if isinstance(row, dict):
            name = str(row.get("tool") or row.get("name") or "")
            if name and name not in names:
                names.append(name)
    for receipt in completed_observation_receipts(turn):
        name = receipt["operation_id"]
        if name not in names:
            names.append(name)
    return names


def canonical_execution_evidence(turn: dict[str, Any]) -> tuple[bool, dict[str, Any]]:
    """Resolve captured Result and shared resume envelopes, never arbitrary data."""
    value = turn.get("execute")
    if not isinstance(value, dict):
        return False, {}
    success = value.get("success") is True
    # A captured Result has only these public fields. Business payloads containing
    # a `data` field are not wrappers and cannot grant execution authority.
    if set(value).issubset({"success", "message", "data"}) and isinstance(value.get("data"), dict):
        value = value["data"]
    if (isinstance(value.get("execute"), dict)
            and (isinstance(value.get("state"), dict) or isinstance(value.get("session_state"), dict))):
        value = value["execute"]
        if value.get("success") is False:
            success = False
    return success, value


def pending_control_reference(turn: dict[str, Any]) -> dict[str, Any]:
    """Freeze the actual current proposal and its task owner before inquiry."""
    raw = turn.get("raw_turn_data") if isinstance(turn.get("raw_turn_data"), dict) else {}
    state = raw.get("state") or raw.get("session_state") or {}
    state = state if isinstance(state, dict) else {}
    action = turn.get("pending_action") or turn.get("action_proposal") or {}
    return deepcopy({"action": action if isinstance(action, dict) else {},
                     "task_id": state.get("session_id"),
                     "status": state.get("status"),
                     "active_pending_action_id": state.get("active_pending_action_id"),
                     "pending_actions": state.get("pending_actions"),
                     "action_history": state.get("action_history"),
                     "completed_steps": state.get("completed_steps")})


def _immutable_pending_identity(action: dict[str, Any]) -> dict[str, Any]:
    # Provider/runtime continuation hints are not the immutable approval packet.
    return {key: action.get(key) for key in (
        "action_id", "action_type", "domain", "tool", "arguments", "safe_payload",
        "prepare_fingerprint", "selection_id", "record_refs",
        "title", "summary", "operation_metadata", "rows", "warnings", "risk_level",
        "external_write", "destructive", "execute_enabled", "selected_record_count",
        "preview_item_count", "execute_input_count",
    )}


def _same_pending_json(left: Any, right: Any) -> bool:
    try:
        return json.dumps(left, sort_keys=True, ensure_ascii=False, separators=(",", ":"), allow_nan=False) == json.dumps(
            right, sort_keys=True, ensure_ascii=False, separators=(",", ":"), allow_nan=False)
    except (ValueError, TypeError):
        return False


def pending_control_problems(turn: dict[str, Any], reference: dict[str, Any] | None) -> list[str]:
    """A read cannot erase, replace, approve or move the active control packet."""
    problems = []
    reference = reference if isinstance(reference, dict) else {}
    expected = reference.get("action")
    expected = expected if isinstance(expected, dict) else {}
    expected_task_id = reference.get("task_id")
    expected_actions = reference.get("pending_actions")
    expected_identity = _immutable_pending_identity(expected)
    if (not all(isinstance(expected.get(key), str) and expected[key].strip() for key in (
            "action_id", "action_type", "domain", "tool", "prepare_fingerprint"))
            or not isinstance(expected.get("arguments"), dict)
            or not isinstance(expected_task_id, str) or not expected_task_id.strip()
            or expected.get("_agent_task_session_id") != expected_task_id
            or expected.get("status") != "pending_approval" or expected.get("approval_required") is not True
            or reference.get("status") != "pending_approval"
            or reference.get("active_pending_action_id") != expected.get("action_id")
            or not isinstance(expected_actions, list) or len(expected_actions) != 1
            or not isinstance(expected_actions[0], dict)
            or expected_actions[0].get("_agent_task_session_id") != expected_task_id
            or expected_actions[0].get("status") != "pending_approval"
            or expected_actions[0].get("approval_required") is not True
            or not _same_pending_json(_immutable_pending_identity(expected_actions[0]), expected_identity)
            or not isinstance(reference.get("action_history"), list)
            or not isinstance(reference.get("completed_steps"), list)):
        return ["pending_current_proposal_reference_missing"]
    raw = turn.get("raw_turn_data") if isinstance(turn.get("raw_turn_data"), dict) else {}
    state = raw.get("state") or raw.get("session_state") or {}
    state = state if isinstance(state, dict) else {}
    current = turn.get("pending_action") or turn.get("action_proposal") or {}
    current = current if isinstance(current, dict) else {}
    proposal_alias = turn.get("action_proposal")
    if (not _same_pending_json(_immutable_pending_identity(current), expected_identity)
            or (proposal_alias and (not isinstance(proposal_alias, dict)
                or not _same_pending_json(_immutable_pending_identity(proposal_alias), expected_identity)))):
        problems.append("pending_current_immutable_proposal_lost_or_changed")
    if (state.get("session_id") != expected_task_id
            or current.get("_agent_task_session_id") != expected_task_id
            or (proposal_alias and isinstance(proposal_alias, dict)
                and proposal_alias.get("_agent_task_session_id") != expected_task_id)):
        problems.append("pending_current_task_owner_lost_or_changed")
    pending_actions = state.get("pending_actions")
    canonical = [action for action in pending_actions if isinstance(action, dict)
                 and action.get("action_id") == expected["action_id"]] if isinstance(pending_actions, list) else []
    if (state.get("status") != "pending_approval"
            or state.get("active_pending_action_id") != expected["action_id"]
            or len(canonical) != 1
            or not _same_pending_json(pending_actions, expected_actions)
            or any(action.get("status") != "pending_approval" or action.get("approval_required") is not True
                   or action.get("_agent_task_session_id") != expected_task_id
                   or not _same_pending_json(_immutable_pending_identity(action), expected_identity) for action in canonical)
            or current.get("status") != "pending_approval" or current.get("approval_required") is not True):
        problems.append("pending_canonical_active_approval_state_missing_or_changed")
    if not _same_pending_json(state.get("action_history"), reference["action_history"]):
        problems.append("pending_read_changed_action_history")
    if not _same_pending_json(state.get("completed_steps"), reference["completed_steps"]):
        problems.append("pending_read_changed_completed_observations")
    executed, _ = canonical_execution_evidence(turn)
    effect_sources = [turn, raw, current, *canonical]
    execution_attempted = False
    for container in (turn, raw):
        execution = container.get("execute")
        if not isinstance(execution, dict):
            continue
        execution_attempted = execution_attempted or isinstance(execution.get("success"), bool)
        effect_sources.append(execution)
        if (set(execution).issubset({"success", "message", "data"})
                and isinstance(execution.get("data"), dict)):
            effect_sources.append(execution["data"])
            nested = execution["data"].get("execute")
            if isinstance(nested, dict):
                effect_sources.append(nested)
    effect_claim = any(source.get(key) is True for source in effect_sources for key in (
        "api_call_performed", "db_write_performed", "file_written", "executed", "approved",
    ))
    if executed or execution_attempted or effect_claim:
        problems.append("pending_read_claimed_or_performed_effect")
    return problems


def validate_business_turn(turn: dict[str, Any], scenario: dict[str, Any], *, model: str,
                           pending_reference: dict[str, Any] | None = None) -> list[str]:
    problems = []
    if not turn.get("success"):
        problems.append("application_turn_failed")
    if not str(turn.get("answer") or "").strip():
        problems.append("visible_answer_missing")
    if scenario.get("required_goal_complete"):
        raw = turn.get("raw_turn_data") or {}
        state = raw.get("state") or raw.get("session_state") or {}
        if not isinstance(state, dict) or state.get("status") != "complete":
            problems.append("required_goal_not_complete:" + str(state.get("status") if isinstance(state, dict) else "missing"))
        elif state.get("remaining_steps") or state.get("pending_actions") or state.get("active_pending_action_id"):
            problems.append("completed_goal_has_unfinished_work")
    operations = operation_names(turn)
    domains = {name.split(".", 1)[0] for name in operations}
    for domain in scenario.get("required_domains") or []:
        if domain not in domains:
            problems.append("missing_observation_domain:" + domain)
    for operation in scenario.get("required_operations") or []:
        if operation not in operations:
            problems.append("missing_operation:" + operation)
    runtime = turn.get("runtime") or {}
    phase = scenario.get("phase")
    if phase not in {"pending", "cancel", "approve", "schedule_review", "schedule_approve"}:
        if not int(runtime.get("model_response_count") or 0):
            problems.append("live_model_response_missing")
        actual = str(runtime.get("actual_model") or runtime.get("final_actual_model") or "")
        if actual.casefold() != model.casefold():
            problems.append("exact_model_identity_not_proven")
    if runtime.get("fallback_used") or runtime.get("any_model_fallback_used"):
        problems.append("model_fallback_used")
    executed, data = canonical_execution_evidence(turn)
    if phase in {"read", "prepare", "pending", "cancel", "schedule_review", "schedule_approve"} and executed:
        problems.append("unexpected_execute_before_approval")
    if phase == "pending":
        problems.extend(pending_control_problems(turn, pending_reference))
    if phase == "prepare":
        pending = turn.get("pending_action") or turn.get("action_proposal") or {}
        if not pending or not pending.get("action_id"):
            problems.append("pending_immutable_action_missing")
        if not pending.get("approval_required"):
            problems.append("approval_requirement_missing")
        if pending.get("external_write"):
            problems.append("unexpected_external_proposal")
    if phase == "approve":
        if not executed:
            problems.append("approved_workflow_not_executed" if scenario.get("effect_kind") == "workflow_save" else "approved_export_not_executed")
        verify = data.get("verify") or data.get("verification") or {}
        if not (data.get("verified") is True or verify.get("success") is True or verify.get("verified") is True):
            problems.append("independent_export_readback_missing")
        if scenario.get("effect_kind") == "workflow_save" and (
                verify.get("verified_by") != "workflow_exact_revision_sqlite_readback"
                or verify.get("current_head_matches_approved_revision") is not True
                or verify.get("saved_execution_authority") is not False):
            problems.append("workflow_exact_current_revision_readback_missing")
        execution = turn.get("execute") or {}
        execution_data = execution.get("data") or {}
        executed_state = execution_data.get("state") or execution_data.get("session_state")
        if executed and isinstance(executed_state, dict) and executed_state:
            from AI.processors.response_verifier import _executed_action_state_problems
            from AI.processors.session_state import AgentSessionState
            problems.extend(_executed_action_state_problems(
                str(turn.get("answer") or ""), AgentSessionState.from_dict(executed_state)))
    if scenario.get("primary_order_scope"):
        combined = json.dumps(turn.get("direct_tool_results") or [], ensure_ascii=False)
        if "PILOT-B-NEW-1" in str(turn.get("answer") or ""):
            problems.append("other_store_visible_record")
        if '"api_account_id": 102' in combined or '"api_account_ids": [102]' in combined:
            problems.append("other_store_observed_after_correction")
    return problems


def _source_snapshot() -> dict[str, Any]:
    from AI.processors.orders_source_ui_human_recheck_cli import _runtime_source_evidence
    evidence = _runtime_source_evidence()
    root = Path(__file__).resolve().parents[2]
    rows = list(evidence["files"])
    extra_sources = (
        "requirements.txt", "tests/domain_ai/test_agent_runtime_architecture.py",
        "tests/domain_ai/test_agentic_user_followup.py", "tests/domain_stock/test_stock_database_scope.py",
        "tests/domain_ai/test_turn_control_service.py", "tests/domain_ai/test_turn_entrypoint_parity.py",
        "tests/domain_ai/test_verified_export_followup.py",
        "tests/domain_ai/test_model_loop.py", "Testing/processors/order_snapshot_fixture.py",
        "AI/processors/business_agent_acceptance_cli.py", "AI/processors/business_acceptance_fixtures.py",
        "AI/processors/business_acceptance_evidence.py", "AI/processors/analysis_operations.py",
        "AI/processors/report_builder.py", "AI/processors/observation_projection.py",
        "AI/processors/observation_analysis_evidence.py", "AI/processors/file_workspace.py",
        "AI/processors/business_attachment_fixtures.py", "AI/processors/tabular_analysis.py",
        "AI/processors/observation_context.py",
        "AI/processors/task_control.py", "AI/processors/observation_artifacts.py",
        "AI/processors/child_task_runtime.py", "AI/processors/business_child_acceptance.py",
        "AI/processors/workflow_store.py",
        "AI/processors/workflow_scheduler.py",
        "AI/processors/scheduled_task_runtime.py",
        "AI/views/workflow_library.py",
        "AI/views/workflow_schedule_dialog.py",
        "tests/domain_ai/test_workflow_scheduler.py", "tests/domain_ai/test_workflow_schedule_dialog.py",
        "tests/domain_ai/test_workflow_library.py", "tests/domain_ai/test_scheduled_task_runtime.py",
        "tests/domain_ai/test_chat_workflow_scheduler.py", "tests/domain_ai/test_business_schedule_acceptance.py",
        "tests/domain_ai/test_child_task_runtime.py", "tests/domain_ai/test_business_child_acceptance.py",
        "tests/domain_ai/test_calculation_goal_context.py", "tests/domain_ai/test_order_item_quantity_units.py",
        "tests/domain_ai/test_business_period_oracle.py", "tests/support/run_isolated_ai_checks.py",
        "tests/domain_ai/test_business_tool_result_schemas.py",
        "tests/domain_ai/test_business_attachment_acceptance.py",
        "tests/domain_ai/test_business_child_checkpoint_retention.py",
        "Core/processors/credential_store.py",
        "tests/domain_ai/test_general_action_interpretation.py",
        "tests/domain_ai/test_artifact_task_runtime.py",
        "tests/domain_ai/test_existing_sources_new_report.py",
        "tests/domain_ai/fixtures/schedule_glm_first_turn_interpretation.json",
        "tests/domain_ai/fixtures/schedule_glm_approval_control_first_turn_interpretation.json",
        "tests/domain_ai/test_verified_export_followup.py",
        "tests/domain_ai/test_report_projection_diagnostics.py",
        "tests/domain_ai/test_failed_step_recovery_context.py",
        "tests/domain_ai/test_failed_plan_contract_repair.py",
        "tests/domain_ai/test_failed_web_replan_budget.py",
        "tests/domain_ai/test_configuration_schema_action_verification.py",
        "tests/domain_ai/fixtures/schedule_glm_schema_constraint_drift.json",
        "tests/domain_ai/test_workflow_task_runtime.py",
        "tests/domain_ai/test_pending_action_existing_observation.py",
        "tests/domain_ai/test_ai_session_state.py",
        "tests/domain_ai/fixtures/schedule_glm_pending_interpretation.json",
        "tests/domain_ai/fixtures/attachment_glm_pending_interpretation.json",
        "tests/domain_ai/fixtures/schedule_glm_pending_authority_loss.json",
        "tests/domain_ai/test_pending_control_acceptance.py",
        "tests/domain_ai/test_general_business_summary_projection.py",
        "tests/domain_ai/fixtures/business_glm_broad_summary_evidence.json",
        "tests/domain_ai/test_business_campaign_evidence.py",
        "tests/domain_ai/test_business_agent_acceptance.py",
        "tests/domain_core/test_report_artifacts.py",
        "tests/domain_ai/test_business_numeric_comparison_obligation.py",
        "tests/domain_ai/test_live_projection_replays.py",
        "tests/domain_ai/test_provider_adapters.py",
        "tests/domain_ai/fixtures/business_glm_local_research_projection.json",
        "tests/domain_ai/fixtures/business_glm_formatted_comparison_projection.json",
        "tests/domain_ai/fixtures/child_deepseek_stock_amount_projection.json",
        "tests/domain_ai/fixtures/child_deepseek_unprepared_delivery.json",
        "tests/domain_ai/fixtures/child_deepseek_stock_records_projection.json",
        "tests/domain_ai/fixtures/schedule_glm_saved_workflow_contradiction.json",
        "tests/domain_ai/fixtures/attachment_deepseek_pending_scope_projection.json",
        "tests/domain_ai/fixtures/child_deepseek_pending_bindings_projection.json",
        "tests/domain_ai/fixtures/business_glm_quantity_per_order_labels.json",
        "tests/domain_ai/fixtures/attachment_deepseek_delivery_projection.json",
        "tests/domain_ai/test_prepare_source_lineage.py",
        "tests/domain_ai/fixtures/child_deepseek_report_lineage_denial.json",
        "tests/domain_ai/test_observation_replan_recovery.py",
        "tests/domain_ai/fixtures/agentic_observation_recovery_evidence.json",
        "tests/domain_ai/test_web_model_plan_preservation.py",
        "tests/domain_ai/fixtures/web_skipped_plan_recovery.json",
        "tests/domain_ai/fixtures/attachment_unrequested_calculation_projection.json",
        "tests/domain_ai/test_live_source_binding_and_public_read.py",
        "tests/domain_ai/fixtures/live_source_binding_and_public_read.json",
        "tests/domain_ai/test_public_scope_and_finance_periods.py",
        "tests/domain_ai/fixtures/public_scope_and_finance_periods.json",
        "tests/domain_ai/test_historical_report_scope_and_web_reuse.py",
        "tests/domain_ai/fixtures/historical_report_scope_and_web_reuse.json",
        "tests/domain_ai/test_web_native_delivery_followup.py",
        "tests/domain_ai/fixtures/web_native_delivery_followup.json",
        "tests/domain_ai/fixtures/schedule_glm_local_delivery_followup.json",
        "tests/domain_ai/fixtures/attachment_glm_currency_projection.json",
        "tests/domain_ai/fixtures/business_glm_numeric_comparison_evidence.json",
        "tests/domain_ai/fixtures/business_glm_comparison_refresh_evidence.json",
        "requirements-dev.txt",
        "Testing/processors/business_agent_fixtures.py", "Testing/constants/trendyol_stage_constants.py",
        "tests/domain_testing/test_business_agent_fixtures.py",
        "Marketplaces/processors/catalog.py", "Marketplaces/constants/paths.py",
        "Marketplaces/constants/files/trendyol/platform_api_config.json",
    )
    seen = {row["path"] for row in rows}
    missing = []
    for relative in extra_sources:
        path = root / relative
        if not path.is_file():
            missing.append(relative)
            continue
        if relative in seen:
            continue
        rows.append({"path": relative, "sha256": hashlib.sha256(path.read_bytes()).hexdigest()})
    aggregate = hashlib.sha256(json.dumps([(row["path"], row["sha256"]) for row in rows],
                                         separators=(",", ":")).encode()).hexdigest()
    return {"fingerprint": aggregate, "files": rows, "missing_requested_sources": missing,
            "scope": "Explicit source UI runtime dependencies and this driver, not the entire repository."}


def finalize_campaign_source_evidence(artifact: dict[str, Any]) -> bool:
    """The final/only phase cannot escape the campaign's frozen source check."""
    current = _source_snapshot()
    original = artifact.get("source_snapshot") or {}
    unchanged = bool(not current["missing_requested_sources"]
                     and current["fingerprint"] == original.get("fingerprint"))
    artifact["end_source_snapshot"] = current
    artifact["source_snapshot_unchanged"] = unchanged
    if not unchanged:
        artifact["status_before_source_recheck"] = artifact.get("status")
        artifact["status"] = "source_changed_restart_required"
    return unchanged


def _read_isolated_task_checkpoint(task_id: str, *, isolation_root: str | Path,
                                   task_store_root: str | Path | None = None) -> tuple[Path, bytes, dict[str, Any]]:
    from AI.processors.task_store import TASK_STORE_DIR

    if not isinstance(task_id, str) or UUID(task_id).hex != task_id:
        raise ValueError("Checkpoint retention requires the actual task UUID.")
    isolated = Path(isolation_root).resolve()
    root = Path(task_store_root or TASK_STORE_DIR).resolve()
    root.relative_to(isolated)
    path = (root / (task_id + ".json")).resolve()
    relative = path.relative_to(isolated)
    raw = path.read_bytes()
    payload = json.loads(raw.decode("utf-8"))
    if not isinstance(payload, dict) or payload.get("session_id") != task_id:
        raise ValueError("Durable checkpoint identity differs from its actual task.")
    return relative, raw, payload


def _retain_checkpoint_bytes(task_id: str, *, role: str, stage: str, relative: Path, raw: bytes,
                              payload: dict[str, Any], isolation_root: str | Path,
                              artifact_path: str | Path) -> dict[str, Any]:
    from Core.processors.credential_store import _protect, _unprotect

    artifact = Path(artifact_path).resolve()
    try:
        artifact.parent.relative_to(Path(isolation_root).resolve())
    except ValueError:
        pass
    else:
        raise ValueError("Retained encrypted checkpoints must survive isolation cleanup.")
    directory = (artifact.parent / (artifact.stem + "-checkpoints")).resolve()
    directory.relative_to(artifact.parent)
    ciphertext = _protect(raw.decode("utf-8"))
    if _unprotect(ciphertext).encode("utf-8") != raw:
        raise ValueError("DPAPI did not preserve the original checkpoint bytes.")
    directory.mkdir(parents=True, exist_ok=True)
    digest = hashlib.sha256(raw).hexdigest()
    stage_code = {"before_first_child_model": "cb", "after_child_oracle": "cf",
                  "scheduled_task_attached": "si", "scheduled_task_observed": "sf"}[stage]
    destination = directory / (stage_code + "-" + role + "-" + task_id + "-" + digest[:16] + ".dpapi")
    if destination.exists():
        stored = destination.read_bytes()
    else:
        destination.write_bytes(ciphertext)
        stored = destination.read_bytes()
        if stored != ciphertext:
            raise ValueError("Retained encrypted checkpoint verification failed.")
    if _unprotect(stored).encode("utf-8") != raw:
        raise ValueError("Retained encrypted checkpoint verification failed.")
    return {"role": role, "task_id": task_id, "source_relative_path": relative.as_posix(),
            "artifact_relative_path": destination.relative_to(artifact.parent).as_posix(),
            "checkpoint_phase": (payload.get("active_entities", {}).get("task_checkpoint") or {}).get("phase"),
            "status": payload.get("status"), "original_sha256": digest,
            "ciphertext_sha256": hashlib.sha256(stored).hexdigest(),
            "original_byte_count": len(raw), "ciphertext_byte_count": len(stored),
            "decrypt_roundtrip_verified": True}


def retain_child_task_checkpoints(child_task_id: str, parent_task_id: str, *, stage: str,
                                  isolation_root: str | Path, artifact_path: str | Path,
                                  task_store_root: str | Path | None = None) -> dict[str, Any]:
    """Retain original durable bytes with current-user DPAPI, never public grants."""
    if stage not in {"before_first_child_model", "after_child_oracle"}:
        raise ValueError("Unsupported child checkpoint retention stage.")
    for task_id in (child_task_id, parent_task_id):
        if not isinstance(task_id, str) or UUID(task_id).hex != task_id:
            raise ValueError("Checkpoint retention requires the actual task UUID.")
    if child_task_id == parent_task_id:
        raise ValueError("Parent and child execution identities must differ.")
    originals = [(role, task_id, *_read_isolated_task_checkpoint(task_id, isolation_root=isolation_root,
                    task_store_root=task_store_root))
                 for role, task_id in (("parent", parent_task_id), ("child", child_task_id))]
    parent, child = originals[0][4], originals[1][4]
    parent_active, child_active = parent.get("active_entities") or {}, child.get("active_entities") or {}
    policy = child_active.get("child_delegation_policy") or {}
    reservations = [row for row in (parent_active.get("child_task_reservations") or {}).values()
                    if isinstance(row, dict) and row.get("child_task_id") == child_task_id]
    if (len(reservations) != 1 or policy.get("parent_task_id") != parent_task_id
            or policy.get("depth") != 1 or policy.get("effects_allowed") is not False
            or reservations[0].get("scope_policy") != policy):
        raise ValueError("Durable child and parent reservation do not have the same authority.")
    child_budget = child_active.get("execution_budget") or {}
    if (child_budget.get("max_steps") != reservations[0].get("max_steps")
            or child_budget.get("max_replans") != reservations[0].get("max_replans")):
        raise ValueError("Durable child budget differs from its parent reservation.")
    if stage == "before_first_child_model" and (
            child.get("status") != "running" or child_active.get("planning_not_started") is not True
            or child.get("completed_steps") or child.get("pending_actions")
            or (child_active.get("task_checkpoint") or {}).get("phase") != "initialized"
            or reservations[0].get("status") != "running"):
        raise ValueError("Original pre-model child checkpoint is unavailable.")
    files = [_retain_checkpoint_bytes(task_id, role=role, stage=stage, relative=relative, raw=raw,
                payload=payload, isolation_root=isolation_root, artifact_path=artifact_path)
             for role, task_id, relative, raw, payload in originals]
    return {"protection": "windows_dpapi_current_user", "stage": stage,
            "parent_task_id": parent_task_id, "child_task_id": child_task_id,
            "captured_at": datetime.now(timezone.utc).isoformat(), "files": files}


def retain_scheduled_task_checkpoint(scheduler: Any, run_id: str, session_id: str, *, stage: str,
                                     isolation_root: str | Path, artifact_path: str | Path,
                                     task_store_root: str | Path | None = None) -> dict[str, Any]:
    """Preserve the original scheduled execution and actual safe lease metadata."""
    if stage not in {"scheduled_task_attached", "scheduled_task_observed"}:
        raise ValueError("Unsupported scheduled checkpoint retention stage.")
    assert_isolated_scheduler(scheduler, isolation_root)
    runs = [row for row in scheduler.list_runs() if row["run_id"] == run_id]
    if (len(runs) != 1 or runs[0].get("session_id") != session_id
            or not (runs[0].get("receipt") or {}).get("submission_session_id")
            or runs[0]["receipt"]["submission_session_id"] == session_id):
        raise ValueError("Actual scheduled task differs from its durable run binding.")
    relative, raw, payload = _read_isolated_task_checkpoint(session_id, isolation_root=isolation_root,
                                                         task_store_root=task_store_root)
    active = payload.get("active_entities") or {}
    if stage == "scheduled_task_attached" and (
            active.get("planning_not_started") is not True or payload.get("completed_steps")
            or payload.get("pending_actions") or payload.get("action_history")
            or (active.get("task_checkpoint") or {}).get("phase") != "initialized"):
        raise ValueError("The original scheduled checkpoint preceded neither planning nor effects.")
    checkpoint = _retain_checkpoint_bytes(session_id, role="scheduled", stage=stage, relative=relative,
        raw=raw, payload=payload, isolation_root=isolation_root, artifact_path=artifact_path)
    run = runs[0]
    metadata = {key: run[key] for key in ("run_id", "schedule_id", "schedule_hash", "due_at", "status", "owner_id",
                "attempts", "lease_expires_at", "execution_started_at", "session_id", "created_at", "updated_at")}
    metadata["submission_session_id"] = run["receipt"]["submission_session_id"]
    return {"protection": "windows_dpapi_current_user", "stage": stage,
            "captured_at": datetime.now(timezone.utc).isoformat(), "checkpoint": checkpoint,
            "run_lease_evidence": metadata}


def retain_approved_workflow(turn: dict[str, Any], *, isolation_root: str | Path, artifact_path: str | Path) -> dict[str, Any]:
    """Reopen only the approved instruction revision; never retain its catalog."""
    from AI.processors.workflow_store import _validated_action, read_workflow
    action = turn.get("approved_action") or {}
    if local_acceptance_approval_problems(action, isolation_root):
        raise ValueError("Workflow save is outside the isolated approval boundary.")
    executed, data = canonical_execution_evidence(turn)
    verify = data.get("verify") or {}
    proposal = _validated_action(((action.get("arguments") or {}).get("payload") or {}).get("proposal") or {})
    if (not executed or verify.get("verified") is not True
            or verify.get("verified_by") != "workflow_exact_revision_sqlite_readback"
            or verify.get("workflow_id") != proposal["workflow_id"]
            or verify.get("revision_id") != proposal["revision_id"]
            or verify.get("spec_hash") != proposal["spec_hash"]):
        raise ValueError("Workflow save lacks exact canonical readback identity.")
    current = read_workflow(proposal["workflow_id"])
    exact = read_workflow(proposal["workflow_id"], revision_id=proposal["revision_id"])
    if (current != exact or exact.spec.to_dict() != proposal["spec"]
            or exact.spec.content_hash != proposal["spec_hash"]):
        raise ValueError("Saved workflow differs from the current approved revision.")
    destination_dir = Path(artifact_path).resolve().parent / (Path(artifact_path).stem + "-files")
    destination_dir.mkdir(parents=True, exist_ok=True)
    destination = destination_dir / (str(turn["scenario_id"]) + ".workflow.json")
    destination.write_text(json.dumps(exact.to_dict(), ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return {"artifact_path": str(destination), "sha256": hashlib.sha256(destination.read_bytes()).hexdigest(),
            "approved_action_id": action.get("action_id"), "workflow_id": exact.workflow_id,
            "revision_id": exact.revision_id, "spec_hash": exact.spec.content_hash,
            "independent_reader": "workflow_store.read_workflow immutable revision and current head", "saved_execution_authority": False}


def workflow_fresh_source_problems(turn: dict[str, Any]) -> list[str]:
    from AI.processors.business_acceptance_evidence import completed_observation_receipts
    raw = turn.get("raw_turn_data") or {}
    state = raw.get("state") or raw.get("session_state") or {}
    invocation = (state.get("active_entities") or {}).get("workflow_invocation") or {}
    problems = []
    if (invocation.get("fresh_model_plan_required") is not True
            or invocation.get("requires_current_sources") is not True
            or invocation.get("requires_current_approval") is not True):
        problems.append("workflow_fresh_invocation_authority_boundary_missing")
    reads = [row for row in completed_observation_receipts(turn) if row["operation_id"] == "stock.inventory_overview"]
    if not reads:
        problems.append("workflow_current_stock_read_missing")
    return problems


def assert_isolated_scheduler(scheduler: Any, isolation_root: str | Path) -> None:
    """A UI schedule may write only the synthetic campaign's local store."""
    try:
        scheduler.path.resolve().relative_to(Path(isolation_root).resolve())
    except (ValueError, TypeError):
        raise ValueError("Acceptance scheduler storage is outside isolation.") from None


def prepare_source_ui_schedule(tab: Any, *, parameters: dict[str, Any], isolation_root: str | Path) -> tuple[Any, dict[str, Any]]:
    """Simulate the user's real library and schedule-review clicks, without saving."""
    from PyQt6.QtCore import QDateTime, Qt
    from AI.processors.workflow_scheduler import WorkflowScheduler
    from AI.views.workflow_library import WorkflowLibraryDialog
    from AI.views.workflow_schedule_dialog import WorkflowScheduleDialog
    bindings = tab._current_workflow_scope_bindings()
    library = WorkflowLibraryDialog(tab, allowed_scope_bindings=bindings)
    try:
        selected = next((index for index in range(library.workflow_list.count())
                         if library.workflow_list.item(index).data(Qt.ItemDataRole.UserRole)["name"] == WORKFLOW_ACCEPTANCE_SPEC["name"]), None)
        if selected is None:
            raise ValueError("Approved campaign workflow is absent from the real library.")
        library.workflow_list.setCurrentRow(selected)
        for name, value in parameters.items():
            library._fields[name][1].setText(str(value))
        library.schedule_button.click()
        request = library.run_request
        if request is None or library.requested_mode != "schedule":
            raise ValueError("The real typed parameter dialog rejected scheduling.")
        scheduler = WorkflowScheduler()
        assert_isolated_scheduler(scheduler, isolation_root)
        dialog = WorkflowScheduleDialog(tab, scheduler=scheduler, request=request, allowed_scope_bindings=bindings)
        dialog.interval.setCurrentIndex(dialog.interval.findData(3600))
        due = datetime.now(timezone.utc).replace(second=0, microsecond=0)
        dialog.first_due.setDateTime(QDateTime.fromSecsSinceEpoch(int(due.timestamp())).toLocalTime())
        dialog.prepare_button.click()
        if dialog.proposal is None or not dialog.confirm.isEnabled() or dialog.saved:
            raise ValueError("The real schedule review dialog rejected the proposal.")
        proposal = json.loads(json.dumps(dialog.proposal, ensure_ascii=False, allow_nan=False))
        if (proposal["schedule"]["first_due_at"] != due.isoformat()
                or proposal["schedule"]["interval_seconds"] != 3600
                or proposal["schedule"]["parameters"] != parameters):
            raise ValueError("Reviewed schedule differs from the requested typed time/parameters.")
        if any(row["schedule_id"] == proposal["schedule_id"] for row in scheduler.list_schedules()):
            raise ValueError("Schedule was written before its exact confirmation.")
        return dialog, {"proposal": proposal, "reviewed_text": dialog.review.toPlainText(),
                        "library_request": {"workflow_id": request.workflow_id, "revision_id": request.revision_id,
                                            "spec_hash": request.spec_hash, "parameters": request.parameters},
                        "source_path": "WorkflowLibraryDialog.schedule_button -> WorkflowScheduleDialog.prepare_button",
                        "schedule_saved": False, "workflow_executed": False, "model_invoked": False}
    finally:
        library.deleteLater()


def confirm_source_ui_schedule(dialog: Any, *, reviewed: dict[str, Any], isolation_root: str | Path) -> dict[str, Any]:
    """Click the exact reviewed confirmation and independently reopen its row."""
    assert_isolated_scheduler(dialog.scheduler, isolation_root)
    proposal = reviewed["proposal"]
    if dialog.proposal != proposal or not dialog.confirm.isEnabled():
        raise ValueError("The UI schedule no longer matches the reviewed proposal.")
    dialog.confirm.click()
    if not dialog.saved:
        raise ValueError("The real schedule confirmation did not commit.")
    row = next((row for row in dialog.scheduler.list_schedules() if row["schedule_id"] == proposal["schedule_id"]), None)
    if (row is None or row["status"] != "active" or row["schedule_hash"] != proposal["action_hash"]
            or row["schedule"] != proposal["schedule"] or row["active_run_id"] is not None):
        raise ValueError("Independent scheduler read differs from the exact approved proposal.")
    return {"schedule": row, "confirmed_action_id": proposal["action_id"],
            "independent_reader": "WorkflowScheduler.list_schedules exact SQLite approval chain",
            "source_path": "WorkflowScheduleDialog.confirm.click", "effect_approval_granted": False}


def schedule_task_binding_evidence(scheduler: Any, lease: Any, session_id: str, *, isolation_root: str | Path) -> dict[str, Any]:
    """Observe the real runtime callback after canonical binding, before planning."""
    from AI.processors.task_store import TASK_STORE_DIR, load_agent_task_state
    assert_isolated_scheduler(scheduler, isolation_root)
    TASK_STORE_DIR.resolve().relative_to(Path(isolation_root).resolve())
    run = next(row for row in scheduler.list_runs() if row["run_id"] == lease.run_id)
    checkpoint = load_agent_task_state(session_id)
    if (run["session_id"] != session_id or checkpoint.session_id != session_id
            or run["receipt"]["submission_session_id"] == session_id
            or checkpoint.active_entities.get("planning_not_started") is not True
            or checkpoint.completed_steps or checkpoint.pending_actions or checkpoint.action_history
            or (checkpoint.active_entities.get("task_checkpoint") or {}).get("phase") != "initialized"):
        raise ValueError("Actual schedule execution identity is not fresh/persisted before planning.")
    return {"run_id": run["run_id"], "schedule_id": run["schedule_id"], "schedule_hash": run["schedule_hash"],
            "actual_task_session_id": session_id, "submission_session_id": run["receipt"]["submission_session_id"],
            "observed_at": datetime.now(timezone.utc).isoformat(), "checkpoint_phase": "initialized",
            "planning_not_started": True, "completed_step_count": 0, "pending_action_count": 0,
            "prior_approval_count": 0, "persisted_before_goal_plan_and_tools": True}


def schedule_run_problems(turn: dict[str, Any], *, schedule_receipt: dict[str, Any], binding_receipts: list[dict[str, Any]]) -> list[str]:
    from AI.processors.business_acceptance_evidence import completed_observation_receipts, _currency_rows, _equal_number
    raw = turn.get("raw_turn_data") or {}
    state = raw.get("state") or raw.get("session_state") or {}
    run = raw.get("scheduled_run") or {}
    approved = schedule_receipt["schedule"]
    configuration = approved["schedule"]
    problems = []
    if (run.get("schedule_id") != approved["schedule_id"] or run.get("schedule_hash") != approved["schedule_hash"]
            or run.get("status") != "awaiting_approval" or run.get("session_id") != state.get("session_id")):
        problems.append("scheduled_run_not_bound_to_current_approved_schedule_and_task")
    matching = [row for row in binding_receipts if row.get("run_id") == run.get("run_id")
                and row.get("schedule_id") == approved["schedule_id"]
                and row.get("actual_task_session_id") == state.get("session_id")
                and row.get("submission_session_id") == (run.get("receipt") or {}).get("submission_session_id")
                and row.get("actual_task_session_id") != row.get("submission_session_id")
                and row.get("schedule_hash") == approved["schedule_hash"]
                and row.get("checkpoint_phase") == "initialized" and row.get("planning_not_started") is True
                and all(type(row.get(key)) is int and row[key] == 0 for key in
                        ("completed_step_count", "pending_action_count", "prior_approval_count"))
                and row.get("persisted_before_goal_plan_and_tools") is True]
    if len(matching) != 1:
        problems.append("scheduled_actual_task_checkpoint_before_planning_missing")
    invocation = (state.get("active_entities") or {}).get("workflow_invocation") or {}
    if (invocation.get("revision_id") != configuration["revision_id"]
            or invocation.get("spec_hash") != configuration["spec_hash"]):
        problems.append("scheduled_run_pinned_workflow_revision_mismatch")
    reads = [row["result"] for row in completed_observation_receipts(turn) if row["operation_id"] == "stock.inventory_overview"]
    fresh = False
    for data in reads:
        currency_values = _currency_rows(data, "stock_value_by_currency", "amount")
        stock_rows = data.get("top_balances") or []
        records = {row.get("sku"): row for row in stock_rows}
        if len(records) != len(stock_rows):
            continue
        kupa, filtre = records.get("PILOT-KUPA", {}), records.get("PILOT-FILTRE", {})
        fresh = fresh or (set(currency_values) == {"TRY", "USD"}
                         and _equal_number(currency_values["TRY"], 400) and _equal_number(currency_values["USD"], 30)
                         and _equal_number(kupa.get("balance"), 4) and _equal_number(kupa.get("stock_value"), 400)
                         and kupa.get("unit") == "adet" and kupa.get("stock_value_currency") == "TRY"
                         and _equal_number(filtre.get("balance"), 3) and _equal_number(filtre.get("stock_value"), 30)
                         and filtre.get("unit") == "adet" and filtre.get("stock_value_currency") == "USD")
    if not fresh:
        problems.append("scheduled_fresh_stock_fixture_values_missing")
    if state.get("status") != "pending_approval" or not state.get("active_pending_action_id"):
        problems.append("scheduled_current_report_approval_missing")
    return problems


def retain_approved_export(turn: dict[str, Any], *, isolation_root: str | Path, artifact_path: str | Path) -> dict[str, Any]:
    """Keep only the explicitly approved synthetic XLSX, never database files."""
    action = turn.get("approved_action") or {}
    executed, data = canonical_execution_evidence(turn)
    verify = data.get("verify") or data.get("verification") or {}
    if not executed or verify.get("verified") is not True:
        raise ValueError("Approved export lacks successful exact readback evidence.")
    output = verify.get("output_path")
    if not output:
        raise ValueError("Approved export has no actual output path.")
    source = Path(output).resolve()
    source.relative_to(Path(isolation_root).resolve())
    if source.suffix.casefold() != ".xlsx" or not source.is_file():
        raise ValueError("Approved export is not an existing isolated XLSX file.")
    if local_export_approval_problems(action, isolation_root):
        raise ValueError("Approved action fails the local export boundary.")
    payload = (action.get("arguments") or {}).get("payload") or action.get("safe_payload") or {}
    if Path(payload.get("output_path") or "").resolve() != source:
        raise ValueError("Actual output path differs from the approved proposal.")
    from Core.processors.report_artifacts import ReportSnapshot
    snapshot = ReportSnapshot.from_dict(payload.get("report_snapshot") or {})
    if (verify.get("verified_by") != "xlsx_exact_snapshot_readback"
            or verify.get("snapshot_id") != snapshot.snapshot_id
            or verify.get("content_hash") != snapshot.content_hash):
        raise ValueError("Executed report identity differs from the approved snapshot.")
    file_hash = hashlib.sha256(source.read_bytes()).hexdigest()
    if verify.get("file_sha256") != file_hash:
        raise ValueError("Export file bytes differ from the recorded exact readback.")
    from openpyxl import load_workbook
    workbook = load_workbook(source, read_only=True, data_only=False)
    try:
        sheets = [{"name": sheet.title, "rows": [list(row) for row in sheet.iter_rows(values_only=True)]}
                  for sheet in workbook.worksheets]
        actual_names = [sheet["name"] for sheet in sheets if sheet["name"] != "_OrderScout"]
        if actual_names != [sheet.name for sheet in snapshot.sheets]:
            raise ValueError("Saved report sheets differ from the approved snapshot.")
        manifest = {row[0]: row[1] for sheet in sheets if sheet["name"] == "_OrderScout" for row in sheet["rows"]}
        if manifest.get("snapshot_id") != snapshot.snapshot_id or manifest.get("content_hash") != snapshot.content_hash:
            raise ValueError("Saved workbook manifest differs from the approved snapshot.")
        for expected in snapshot.sheets:
            actual = next(sheet for sheet in sheets if sheet["name"] == expected.name)
            wanted = [[column.label for column in expected.columns]]
            for row in expected.rows:
                # The report writer escapes Excel formula-like text with a single
                # apostrophe; numeric, boolean and null cells retain their types.
                wanted.append([(None if value == "" else "'" + value if value.startswith(("=", "+", "-", "@")) else value)
                               if isinstance(value, str) else value for value in row])
            if actual["rows"] != wanted:
                raise ValueError("Saved report cells differ from the approved snapshot: " + expected.name)
    finally:
        workbook.close()
    destination_dir = Path(artifact_path).resolve().parent / (Path(artifact_path).stem + "-files")
    destination_dir.mkdir(parents=True, exist_ok=True)
    destination = destination_dir / (str(turn["scenario_id"]) + ".xlsx")
    shutil.copy2(source, destination)
    return {"artifact_path": str(destination), "sha256": hashlib.sha256(destination.read_bytes()).hexdigest(),
            "byte_count": destination.stat().st_size, "approved_action_id": action.get("action_id"),
            "approved_snapshot_id": snapshot.snapshot_id, "approved_content_hash": snapshot.content_hash,
            "independent_reader": "openpyxl.load_workbook(read_only=True,data_only=False)", "sheets": sheets}


def run_source_ui_campaign(
    *, provider_record: dict[str, Any], artifact_path: str | Path,
    isolation_root: str | Path, fixture: dict[str, Any], scenario_ids: list[str],
    allow_public_web: bool = False, timeout: float = 180, stop_on_failure: bool = True,
) -> dict[str, Any]:
    from PyQt6.QtWidgets import QApplication
    from AI.views import chat_tab as chat_module
    from AI.processors.cloud_runtime import build_provider_chat_client
    from AI.processors.pipeline import list_provider_models
    from AI.processors.orders_source_ui_human_recheck_cli import (
        _capture_turn, _ensure_provider_selection, _wait_for_workers,
    )
    from Feedback.processors.pipeline import Result
    from main import build_source_app_ai_acceptance_surface
    from AI.processors.business_acceptance_evidence import (
        completed_observation_receipts, validate_attachment_export_cells, validate_attachment_observations,
        validate_fixture_observations, validate_general_report_sources,
    )
    from AI.processors.business_attachment_fixtures import seed_attachment_fixtures
    from AI.processors.business_child_acceptance import validate_child_stock_report, validate_child_stock_cells

    model = str(provider_record["default_model"])
    provider_code = str(provider_record["provider_code"])
    endpoint = acceptance_provider_endpoint(provider_record)
    boundary = NetworkBoundary(endpoint, allow_public_web=allow_public_web)
    source = _source_snapshot()
    secrets = (str(provider_record.get("api_key") or ""),)
    artifact_path = Path(artifact_path).resolve()
    artifact_path.parent.mkdir(parents=True, exist_ok=True)
    artifact: dict[str, Any] = {
        "artifact_type": "orderscout_live_business_agent_source_ui_acceptance",
        "schema_version": 2, "created_at": datetime.now(timezone.utc).isoformat(),
        "requested_model": model, "provider_code": provider_code, "fallback_allowed": False,
        "source_entrypoint": "desktop/main.py", "source_send_path": "AITab._send_message",
        "render_mode": "offscreen", "planner_kind": "real_cloud_model",
        "fixture": fixture, "source_snapshot": source, "turns": [], "status": "running",
        "scenario_ids": scenario_ids,
        "safety": {"business_data": "synthetic_temporary_sqlite", "customer_data_included": False,
                   "provider_credential_included": False, "source_provider_read_only": True,
                   "source_business_tables_read": False, "marketplace_api_allowed": False,
                   "external_mutation_allowed": False, "local_export_requires_explicit_approval": True},
        "acceptance_validation": {"accepted": False, "direct_review_status": "pending",
                                  "review_requirements": ["Visible answer matches actual observations and scope",
                                                          "No inferred bank balance, settlement realization or unsupported account/company ownership",
                                                          "Research usefulness and report layout meet the natural user request",
                                                          "Follow-up and approval apply to the current immutable proposal"]},
    }
    attachment_fixtures = (seed_attachment_fixtures(isolation_root)
                           if any(SCENARIOS[key].get("attachment_formats") for key in scenario_ids) else {})
    if attachment_fixtures:
        input_directory = artifact_path.parent / (artifact_path.stem + "-inputs")
        input_directory.mkdir(parents=True, exist_ok=True)
        selected_formats = {kind for key in scenario_ids for kind in SCENARIOS[key].get("attachment_formats") or []}
        for kind in selected_formats:
            source_path = Path(attachment_fixtures[kind]["path"]).resolve()
            source_path.relative_to(Path(isolation_root).resolve())
            retained_source = input_directory / attachment_fixtures[kind]["name"]
            shutil.copy2(source_path, retained_source)
            attachment_fixtures[kind]["retained_source_path"] = str(retained_source)
        artifact["attachment_fixtures"] = attachment_fixtures
        artifact["source_attachment_import_path"] = "AITab._import_chat_attachments"

    def save():
        artifact["network_events"] = boundary.events
        artifact_path.write_text(json.dumps(secret_safe(artifact, secrets), ensure_ascii=False,
                                            indent=2, default=str) + "\n", encoding="utf-8")

    save()
    app = QApplication.instance() or QApplication([])
    window = None
    with ExitStack() as stack:
        boundary.install(stack)
        try:
            catalog = list_provider_models(provider_code=provider_code, endpoint_url=endpoint,
                                           api_key=provider_record["api_key"], free_only=False, limit=500)
            if not catalog.success:
                raise RuntimeError("Live model catalog lookup failed.")
            metadata = next((dict(row) for row in (catalog.data or {}).get("records") or []
                             if str(row.get("id") or row.get("model_id") or "").casefold() == model.casefold()), {})
            if not metadata:
                raise RuntimeError("Exact model is absent from the live provider catalog.")
        except Exception as exc:
            artifact.update({"status": "failed", "exception_type": type(exc).__name__,
                             "exception_message": str(exc), "completed_at": datetime.now(timezone.utc).isoformat()})
            artifact["acceptance_validation"]["automated_checks_passed"] = False
            save()
            return secret_safe(artifact, secrets)
        artifact["model_metadata"] = metadata
        capture: dict[str, Any] = {}
        original_turn = chat_module._run_chat_turn_with_runtime
        original_scheduled_turn = chat_module._run_scheduled_chat_with_runtime
        original_execute = chat_module._run_approved_chat_action

        def provider_getter(_provider_id):
            return Result.ok("Live acceptance provider resolved in memory.", data={"record": dict(provider_record)})

        def client_builder(selected_record, **kwargs):
            record = {**provider_record, **dict(selected_record or {})}
            record["api_key"] = provider_record["api_key"]
            record["allow_model_fallback"] = False
            record["model_metadata"] = metadata
            options = dict(kwargs)
            options.setdefault("timeout", min(float(timeout), 120))
            client = build_provider_chat_client(record, **options)
            def before_child(child_id, parent_id):
                retained = retain_child_task_checkpoints(child_id, parent_id, stage="before_first_child_model",
                    isolation_root=isolation_root, artifact_path=artifact_path)
                capture.setdefault("retained_child_checkpoints", []).append(retained)
                return retained
            return (ObservedCloudClient(client, lambda row: capture.setdefault("model_responses", []).append(row),
                                        before_child_callback=before_child)
                    if client is not None else None)

        def run_turn(*args, **kwargs):
            result = original_turn(*args, **kwargs)
            capture["turn_result"] = result
            capture["raw_turn_data"] = dict(result.data or {})
            return result

        def run_scheduled_turn(*args, **kwargs):
            result = original_scheduled_turn(*args, **kwargs)
            capture["turn_result"] = result
            capture["raw_turn_data"] = dict(result.data or {})
            return result

        from AI.processors.workflow_scheduler import WorkflowScheduler
        original_attach = WorkflowScheduler.attach_task_session
        def attach_schedule_task(scheduler, lease, session_id):
            value = original_attach(scheduler, lease, session_id)
            receipt = schedule_task_binding_evidence(scheduler, lease, session_id, isolation_root=isolation_root)
            retained = retain_scheduled_task_checkpoint(scheduler, lease.run_id, session_id,
                stage="scheduled_task_attached", isolation_root=isolation_root, artifact_path=artifact_path)
            capture.setdefault("retained_schedule_checkpoints", []).append(retained)
            capture.setdefault("schedule_bindings", []).append(receipt)
            return value

        def execute(action, *args, **kwargs):
            capture["approved_action"] = dict(action or {})
            problems = local_acceptance_approval_problems(action, isolation_root)
            result = (Result.fail("Acceptance approval boundary rejected the action.", data={"guard_problems": problems})
                      if problems else original_execute(action, *args, **kwargs))
            capture["execute_result"] = result
            return result

        stack.enter_context(patch.object(chat_module.ai_actions, "action_get_provider_account", provider_getter))
        stack.enter_context(patch.object(chat_module, "build_provider_chat_client", client_builder))
        stack.enter_context(patch.object(chat_module, "_run_chat_turn_with_runtime", run_turn))
        stack.enter_context(patch.object(chat_module, "_run_scheduled_chat_with_runtime", run_scheduled_turn))
        stack.enter_context(patch.object(chat_module, "_run_approved_chat_action", execute))
        stack.enter_context(patch.object(WorkflowScheduler, "attach_task_session", attach_schedule_task))
        schedule_dialog = None
        schedule_review_receipt = None
        schedule_confirmation = None
        try:
            window, tab = build_source_app_ai_acceptance_surface(app)
            artifact["source_window_runtime_class"] = f"{type(window).__module__}.{type(window).__name__}"
            artifact["source_widget_runtime_class"] = f"{type(tab).__module__}.{type(tab).__name__}"
            if any(SCENARIOS[key].get("schedule_stage") for key in scenario_ids):
                # Preserve the actual Qt timeout handler; this staged campaign
                # emits it explicitly after the separate human review/confirm
                # steps, avoiding an unrelated wall-clock timer race.
                tab._schedule_timer.stop()
                artifact["source_schedule_run_path"] = "AITab._schedule_timer.timeout -> AITab._poll_scheduled_workflows -> selected provider worker"
            last_pending_action_id = None
            last_pending_reference = None
            last_attachment_formats: tuple[str, ...] = ()
            last_attachment_source_receipts: list[dict[str, Any]] = []
            for scenario_id in scenario_ids:
                scenario = SCENARIOS[scenario_id]
                if scenario.get("requires_public_web") and not allow_public_web:
                    artifact["status"] = "configuration_blocked"
                    artifact["configuration_problem"] = "public_research_requires_explicit_public_web_option"
                    break
                if _source_snapshot()["fingerprint"] != source["fingerprint"]:
                    artifact["status"] = "source_changed_restart_required"
                    break
                capture.clear()
                attachment_formats = tuple(scenario.get("attachment_formats") or [])
                selected_metadata: list[dict[str, Any]] = []
                if attachment_formats:
                    tab._import_chat_attachments([attachment_fixtures[kind]["path"] for kind in attachment_formats])
                    workspace = ((tab._conversation_state.get("active_entities") or {}).get("file_workspace") or {})
                    selected_metadata = [dict(row) for row in workspace.get("attachments") or []]
                    artifact.setdefault("source_ui_attachment_imports", []).append({
                        "scenario_id": scenario_id, "selected_formats": list(attachment_formats),
                        "registered_metadata": selected_metadata,
                    })
                    last_attachment_formats = attachment_formats
                _ensure_provider_selection(tab, provider_id=int(provider_record["id"]),
                                           provider_label=str(provider_record.get("label") or provider_code),
                                           provider_code=provider_code, model=model, model_metadata=metadata)
                started_at = datetime.now(timezone.utc).isoformat()
                started = time.monotonic()
                actual_user_prompt = scenario["prompt"]
                invocation_receipt = None
                schedule_stage = scenario.get("schedule_stage")
                if schedule_stage == "review":
                    schedule_dialog, schedule_review_receipt = prepare_source_ui_schedule(
                        tab, parameters=scenario["schedule_parameters"], isolation_root=isolation_root)
                    data = {"answer": schedule_dialog.review.toPlainText(), "mode": "schedule_review",
                            "schedule_review": schedule_review_receipt, "state": dict(tab._conversation_state)}
                    capture["turn_result"] = Result.ok(data["answer"], close_dialog=False, data=data)
                    capture["raw_turn_data"] = data
                elif schedule_stage == "approve":
                    if schedule_dialog is None or schedule_review_receipt is None:
                        raise ValueError("No current exact schedule review is available for confirmation.")
                    schedule_confirmation = confirm_source_ui_schedule(
                        schedule_dialog, reviewed=schedule_review_receipt, isolation_root=isolation_root)
                    data = {"answer": "Gözden geçirilen tam plan kaydedildi; çalışma kendi güncel işlem onayını bekleyecek.",
                            "mode": "schedule_approved", "schedule_confirmation": schedule_confirmation,
                            "state": dict(tab._conversation_state)}
                    capture["turn_result"] = Result.ok(data["answer"], close_dialog=False, data=data)
                    capture["raw_turn_data"] = data
                elif schedule_stage == "run":
                    if schedule_dialog is None or schedule_confirmation is None:
                        raise ValueError("No current human-confirmed schedule is available to run.")
                    assert_isolated_scheduler(schedule_dialog.scheduler, isolation_root)
                    tab._schedule_timer.timeout.emit()
                    if tab._chat_worker is None and capture.get("turn_result") is None:
                        raise ValueError("The real due-schedule Qt handler did not start a worker.")
                elif scenario.get("workflow_invocation"):
                    from PyQt6.QtCore import Qt
                    from AI.views.workflow_library import WorkflowLibraryDialog
                    bindings = tab._current_workflow_scope_bindings()
                    dialog = WorkflowLibraryDialog(tab, allowed_scope_bindings=bindings)
                    selected = next((index for index in range(dialog.workflow_list.count())
                                     if dialog.workflow_list.item(index).data(Qt.ItemDataRole.UserRole)["name"] == WORKFLOW_ACCEPTANCE_SPEC["name"]), None)
                    if selected is None:
                        raise ValueError("The approved campaign workflow is absent from the real library.")
                    dialog.workflow_list.setCurrentRow(selected)
                    for name, value in scenario["workflow_invocation"].items():
                        dialog._fields[name][1].setText(str(value))
                    dialog.run_button.click()
                    request = dialog.run_request
                    if request is None:
                        raise ValueError("The real workflow parameter dialog rejected the invocation.")
                    actual_user_prompt = request.objective
                    invocation_receipt = {"workflow_id": request.workflow_id, "revision_id": request.revision_id,
                                          "spec_hash": request.spec_hash, "parameters": request.parameters,
                                          "source_path": "WorkflowLibraryDialog.run_button -> AITab._start_saved_workflow -> shared chat"}
                    tab._start_saved_workflow(request, allowed_scope_bindings=bindings)
                    dialog.deleteLater()
                else:
                    tab.composer.setPlainText(actual_user_prompt)
                    tab._send_message()
                try:
                    _wait_for_workers(tab, app, timeout=timeout)
                except TimeoutError:
                    artifact["status"] = "timeout_restart_required"
                    artifact["turns"].append({"scenario_id": scenario_id, "success": False,
                                              "problems": ["source_ui_worker_timeout"], "prompt": scenario["prompt"]})
                    save()
                    os._exit(124)  # No destructor can safely join the still-running QThread.
                raw_data = capture.pop("raw_turn_data", {})
                model_responses = capture.pop("model_responses", [])
                turn = _capture_turn(tab, prompt=actual_user_prompt, review_previous={}, capture=capture,
                                     started_at=started_at, completed_at=datetime.now(timezone.utc).isoformat(),
                                     duration_seconds=time.monotonic() - started)
                turn["scenario_id"] = scenario_id
                turn["scenario_contract"] = scenario
                turn["raw_turn_data"] = raw_data
                turn["model_responses"] = model_responses
                turn["retained_child_checkpoints"] = capture.pop("retained_child_checkpoints", [])
                turn["retained_schedule_checkpoints"] = capture.pop("retained_schedule_checkpoints", [])
                turn["approved_action"] = capture.pop("approved_action", {})
                if invocation_receipt:
                    turn["workflow_invocation_receipt"] = invocation_receipt
                turn["problems"] = validate_business_turn(turn, scenario, model=model,
                                                         pending_reference=last_pending_reference)
                if schedule_stage == "review":
                    turn["schedule_review_receipt"] = schedule_review_receipt
                elif schedule_stage == "approve":
                    turn["schedule_confirmation_receipt"] = schedule_confirmation
                elif schedule_stage == "run":
                    turn["schedule_confirmation_receipt"] = schedule_confirmation
                    turn["schedule_task_binding_receipts"] = capture.pop("schedule_bindings", [])
                    turn["problems"].extend(schedule_run_problems(
                        turn, schedule_receipt=schedule_confirmation, binding_receipts=turn["schedule_task_binding_receipts"]))
                    try:
                        run = raw_data.get("scheduled_run") or {}
                        retained = retain_scheduled_task_checkpoint(schedule_dialog.scheduler, run["run_id"], run["session_id"],
                            stage="scheduled_task_observed", isolation_root=isolation_root, artifact_path=artifact_path)
                        turn["retained_schedule_checkpoints"].append(retained)
                    except Exception as exc:
                        turn["problems"].append("scheduled_checkpoint_retention_failed:" + type(exc).__name__)
                if scenario.get("fixture_oracle"):
                    turn["fixture_observation_evidence"] = validate_fixture_observations(turn, fixture)
                    turn["problems"].extend(turn["fixture_observation_evidence"]["problems"])
                if scenario.get("child_stock_report"):
                    turn["child_stock_evidence"] = validate_child_stock_report(turn, fixture, isolation_root=isolation_root, expected_model=model)
                    turn["problems"].extend(turn["child_stock_evidence"]["problems"])
                    parent = (raw_data.get("state") or raw_data.get("session_state") or {})
                    reservations = (parent.get("active_entities") or {}).get("child_task_reservations") or {}
                    for reservation in reservations.values():
                        if not isinstance(reservation, dict) or not reservation.get("child_task_id"):
                            continue
                        try:
                            retained = retain_child_task_checkpoints(reservation["child_task_id"], parent["session_id"],
                                stage="after_child_oracle", isolation_root=isolation_root, artifact_path=artifact_path)
                            turn["retained_child_checkpoints"].append(retained)
                        except Exception as exc:
                            turn["problems"].append("child_checkpoint_retention_failed:" + type(exc).__name__)
                if scenario.get("general_report_sources"):
                    turn["report_source_evidence"] = validate_general_report_sources(turn)
                    turn["problems"].extend(turn["report_source_evidence"]["problems"])
                if attachment_formats:
                    turn["selected_attachment_metadata"] = selected_metadata
                    turn["attachment_observation_evidence"] = validate_attachment_observations(
                        turn, attachment_fixtures, selected_metadata, required_formats=attachment_formats,
                    )
                    turn["problems"].extend(turn["attachment_observation_evidence"]["problems"])
                if scenario.get("attachment_report"):
                    last_attachment_source_receipts = deepcopy(completed_observation_receipts(turn))
                    turn["report_source_evidence"] = validate_general_report_sources(
                        turn, required_domains=("files",), required_source_operations=("files.read",),
                    )
                    turn["problems"].extend(turn["report_source_evidence"]["problems"])
                    pending = turn.get("pending_action") or turn.get("action_proposal") or {}
                    snapshot = ((pending.get("arguments") or {}).get("payload") or pending.get("safe_payload") or {}).get("report_snapshot") or {}
                    turn["attachment_prepared_cell_evidence"] = validate_attachment_export_cells(
                        {"sheets": snapshot.get("sheets") or []}, attachment_fixtures, required_formats=attachment_formats,
                        report_snapshot=snapshot, source_receipts=completed_observation_receipts(turn),
                    )
                    turn["problems"].extend(turn["attachment_prepared_cell_evidence"]["problems"])
                if scenario.get("phase") == "prepare":
                    pending = turn.get("pending_action") or turn.get("action_proposal") or {}
                    last_pending_action_id = pending.get("action_id")
                    last_pending_reference = pending_control_reference(turn)
                    turn["problems"].extend(local_acceptance_approval_problems(pending, isolation_root))
                    if scenario.get("workflow_spec_fixture"):
                        turn["problems"].extend(workflow_proposal_problems(pending))
                    if scenario.get("workflow_fresh_sources"):
                        turn["problems"].extend(workflow_fresh_source_problems(turn))
                    output = ((pending.get("arguments") or {}).get("payload") or pending.get("safe_payload") or {}).get("output_path")
                    if output and Path(output).is_file():
                        turn["problems"].append("export_file_exists_before_explicit_approval")
                if scenario.get("phase") == "approve" and (turn.get("execute") or {}).get("success"):
                    if not last_pending_action_id or turn["approved_action"].get("action_id") != last_pending_action_id:
                        turn["problems"].append("executed_action_not_current_campaign_proposal")
                    try:
                        if scenario.get("effect_kind") == "workflow_save":
                            turn["retained_workflow"] = retain_approved_workflow(turn, isolation_root=isolation_root, artifact_path=artifact_path)
                        else:
                            turn["retained_export"] = retain_approved_export(turn, isolation_root=isolation_root, artifact_path=artifact_path)
                        if scenario.get("child_stock_export"):
                            turn["child_saved_cell_evidence"] = validate_child_stock_cells(turn["retained_export"], fixture)
                            turn["problems"].extend(turn["child_saved_cell_evidence"]["problems"])
                        if scenario.get("attachment_export"):
                            if not last_attachment_formats:
                                turn["problems"].append("attachment_export_without_current_attachment_prepare")
                            else:
                                turn["attachment_saved_cell_evidence"] = validate_attachment_export_cells(
                                    turn["retained_export"], attachment_fixtures, required_formats=last_attachment_formats,
                                    report_snapshot=((turn["approved_action"].get("arguments") or {}).get("payload")
                                                     or turn["approved_action"].get("safe_payload") or {}).get("report_snapshot"),
                                    source_receipts=last_attachment_source_receipts,
                                )
                                turn["problems"].extend(turn["attachment_saved_cell_evidence"]["problems"])
                    except Exception as exc:
                        turn["problems"].append("independent_export_artifact_read_failed:" + type(exc).__name__)
                turn["problems"] = list(dict.fromkeys(turn["problems"]))
                artifact["turns"].append(turn)
                save()
                print(json.dumps(secret_safe({"scenario_id": scenario_id, "answer": turn.get("answer"),
                                              "operations": operation_names(turn), "problems": turn["problems"],
                                              "runtime": turn.get("runtime")}, secrets), ensure_ascii=False), flush=True)
                if turn["problems"] and stop_on_failure:
                    artifact["status"] = "failed"
                    break
            else:
                artifact["status"] = "automated_checks_passed" if all(not row["problems"] for row in artifact["turns"]) else "failed"
        except Exception as exc:
            artifact.update({"status": "failed", "exception_type": type(exc).__name__, "exception_message": str(exc)})
        finally:
            if schedule_dialog is not None:
                schedule_dialog.close()
                schedule_dialog.deleteLater()
            if window is not None:
                window.close()
                app.processEvents()
    finalize_campaign_source_evidence(artifact)
    artifact["acceptance_validation"]["automated_checks_passed"] = artifact["status"] == "automated_checks_passed"
    artifact["acceptance_validation"]["scope_limit"] = "Scenario-specific live synthetic acceptance; visible answers still require direct review. No whole-product or production acceptance."
    artifact["completed_at"] = datetime.now(timezone.utc).isoformat()
    save()
    return secret_safe(artifact, secrets)


def main(argv: list[str] | None = None) -> int:
    for stream in (sys.stdout, sys.stderr):
        if callable(getattr(stream, "reconfigure", None)):
            stream.reconfigure(encoding="utf-8", errors="backslashreplace")
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source-provider-db", required=True)
    parser.add_argument("--provider-id", type=int, required=True)
    parser.add_argument("--model", required=True)
    parser.add_argument("--artifact-path", required=True)
    parser.add_argument("--scenario", action="append", choices=tuple(SCENARIOS))
    parser.add_argument("--allow-public-web", action="store_true")
    parser.add_argument("--continue-after-failure", action="store_true")
    parser.add_argument("--timeout", type=float, default=180)
    args = parser.parse_args(argv)
    source_db = Path(args.source_provider_db).resolve()
    artifact_path = Path(args.artifact_path).resolve()
    if "settings" in sys.modules:
        raise RuntimeError("Live acceptance CLI requires a fresh process before settings import.")
    os.environ["QT_QPA_PLATFORM"] = "offscreen"
    os.environ["PYTHONDONTWRITEBYTECODE"] = "1"
    sys.dont_write_bytecode = True
    desktop = Path(__file__).resolve().parents[2]
    sys.path.insert(0, str(desktop))
    with tempfile.TemporaryDirectory(prefix="orderscout-live-business-") as isolated:
        os.environ["LOCALAPPDATA"] = isolated
        try:
            from settings import DB_NAME
            from AI.processors.business_acceptance_fixtures import seed_business_fixture
            provider = load_provider_readonly(source_db, args.provider_id, args.model)
            fixture = seed_business_fixture(DB_NAME, isolation_root=isolated)
            artifact = run_source_ui_campaign(provider_record=provider, artifact_path=artifact_path,
                                              isolation_root=isolated, fixture=fixture,
                                              scenario_ids=args.scenario or ["business_research", "general_report_prepare",
                                                                            "general_report_pending", "general_report_approve",
                                                                            "general_report_followup"],
                                              allow_public_web=args.allow_public_web,
                                              timeout=args.timeout,
                                              stop_on_failure=not args.continue_after_failure)
        finally:
            from Core.utils.model_utils import clear_engine_cache
            import logging
            clear_engine_cache()
            logging.shutdown()
    return 0 if artifact["status"] == "automated_checks_passed" else 1


if __name__ == "__main__":
    raise SystemExit(main())
