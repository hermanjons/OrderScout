from __future__ import annotations

import tempfile
import unittest
import json
from datetime import datetime, timedelta
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

from sqlmodel import Session, select

from Account.models.models import ApiAccount
from AI.processors import operation_runtime
from AI.processors.agent import (
    AgentDecision,
    _answer_claims_unsupported_profit,
    _answer_claims_unsupported_full_context_risk,
    _answer_conflicts_with_full_context_money,
    _deterministic_tool_answer,
    _format_listing_rows,
    _grounded_order_analysis_answer,
    _format_stock_rows,
    _should_use_deterministic_answer,
    _synthesis_prompt,
    choose_action,
    run_agent_turn,
)
from AI.processors.cloud_runtime import build_provider_chat_client
from AI.processors.live_evidence import redact_live_evidence
from AI.processors.session_state import PendingActionSnapshot, SelectionSnapshot
from AI.processors.tools import labels_prepare_create_from_orders, logistics_profile, orders_lookup
from Core.utils.model_utils import get_engine
from Feedback.processors.pipeline import Result
from Orders.constants import TZ_GMT3
from Orders.models.trendyol.models import OrderData, OrderHeader, OrderItem, ReturnClaim
from tests.domain_ai.test_orders_agentic_live_eval import (
    _CountingClient,
    _live_eval_enabled,
    _live_provider_record,
    _safe_provider_label,
)
from tests.support.db_case import IsolatedDbTestCase


class TestOrdersBehavioralE2E(IsolatedDbTestCase):
    def test_llm_answer_decision_gets_unconstrained_direct_response_pass(self) -> None:
        prompts: list[str] = []

        def client(prompt: str) -> str:
            prompts.append(prompt)
            return "Bence e-ticaret guclu ama operasyon disiplini isteyen bir is."

        decision = AgentDecision(
            mode="answer",
            response="Planner JSON icindeki kisa cevap.",
            source="llm",
        )
        with patch(
            "AI.processors.agent.choose_action",
            return_value=Result.ok(
                "answer",
                close_dialog=False,
                data={"decision": decision},
            ),
        ):
            result = run_agent_turn(
                "sence e ticaret iyi bir is mi",
                planner_client=client,
                chat_client=client,
            )

        self.assertTrue(result.success)
        self.assertEqual(
            (result.data or {}).get("answer"),
            "Bence e-ticaret guclu ama operasyon disiplini isteyen bir is.",
        )
        self.assertEqual(len(prompts), 1)
        self.assertIn("sadece kullaniciya gidecek dogal metni", prompts[0])

    def test_direct_chat_json_wrapper_is_unwrapped_for_non_casual_answer(self) -> None:
        decision = AgentDecision(
            mode="answer",
            response="Az önceki doğrulanmış sonucu açıklıyorum.",
            source="fallback",
        )
        with patch(
            "AI.processors.agent.choose_action",
            return_value=Result.ok(
                "answer",
                close_dialog=False,
                data={"decision": decision},
            ),
        ):
            result = run_agent_turn(
                "Az önceki doğrulanmış sonucun ortak özelliğini kısa açıkla.",
                planner_client=lambda _prompt: (
                    '{"answer":"Ortak özellik: yalnız doğrulanmış Stage '
                    'kayıtlarına dayanması."}'
                ),
            )

        self.assertTrue(result.success, result.message)
        answer = self._answer_from(result)
        self.assertEqual(
            answer,
            "Ortak özellik: yalnız doğrulanmış Stage kayıtlarına dayanması.",
        )
        self.assertFalse(answer.lstrip().startswith(("{", "[")))
        self.assertEqual(
            (result.data or {}).get("answer_origin"),
            "direct_llm_json_unwrapped",
        )

    def test_identity_and_orders_working_style_does_not_run_health_tools(self) -> None:
        result = run_agent_turn(
            (
                "Selam başkan, hemen tool çalıştırma; önce doğal biçimde "
                "tanışalım. Sen kimsin ve sipariş operasyonunda benimle nasıl "
                "çalışırsın?"
            ),
            planner_client=lambda _prompt: (
                '{"mode":"tools","calls":['
                '{"tool":"orders.count","arguments":{"period":"this_month"}},'
                '{"tool":"orders.revenue_total","arguments":{"period":"this_month"}}]}'
            ),
            chat_client=lambda _prompt: (
                "Ben OrderScout AI. Siparişleri okuyup analiz eder, güvenli "
                "hazırlıkları açık onay sınırında birlikte yürütürüm."
            ),
            db_name=self.db_path,
        )

        self.assertTrue(result.success, result.message)
        self.assertEqual((result.data or {}).get("mode"), "answer")
        self.assertEqual((result.data or {}).get("tool_calls"), [])
        answer = self._answer_from(result)
        self.assertIn("OrderScout AI", answer)
        self.assertIn("Siparis", answer)

    def test_inflected_capability_correction_beats_label_prepare_phrase(self) -> None:
        result = run_agent_turn(
            (
                "Sadece yapamadıklarını anlattın; asıl yapabildiklerin eksik kaldı. "
                "Tool çağırmadan düzelt: sipariş okuma/arama, metrik, kargoya "
                "hazırlık, etiket hazırlama ve iade akışlarını destek durumlarına "
                "göre ayır."
            ),
            planner_client=lambda _prompt: (
                '{"mode":"tool","tool":"labels.create_from_orders.prepare",'
                '"arguments":{"query":"hazırlama","limit":20}}'
            ),
            db_name=self.db_path,
        )

        self.assertTrue(result.success, result.message)
        self.assertEqual((result.data or {}).get("mode"), "answer")
        self.assertEqual((result.data or {}).get("tool_calls"), [])
        self.assertIn("son veya filtreli siparisleri", self._answer_from(result))

    def test_orders_real_yetenek_prompt_stays_tool_free(self) -> None:
        result = run_agent_turn(
            (
                "Orders tarafındaki gerçek yeteneklerini ve onay sınırını kısa "
                "anlat; fatura/PDF uydurma, tool çağırma."
            ),
            planner_client=lambda _prompt: (
                '{"mode":"tools","calls":['
                '{"tool":"orders.count","arguments":{"period":"today"}},'
                '{"tool":"listing.attention_count","arguments":{}}]}'
            ),
            db_name=self.db_path,
        )

        self.assertTrue(result.success, result.message)
        self.assertEqual((result.data or {}).get("mode"), "answer")
        self.assertEqual((result.data or {}).get("tool_calls"), [])
        self.assertIn("onayli akis", self._answer_from(result))

    def test_orders_capability_risk_partition_keeps_local_writes_out_of_reads(self) -> None:
        result = run_agent_turn(
            (
                "Sipari\u015fler taraf\u0131nda ba\u015ftan sona neler yapabilirsin? "
                "Onay gerekenlerle salt okumalar\u0131 ay\u0131rarak anlat; hen\u00fcz "
                "tool \u00e7a\u011f\u0131rma."
            ),
            planner_client=lambda _prompt: (
                "Salt okuma: Trendyol senkronizasyonu ve Excel dosyas\u0131 \u00fcretme. "
                "Onayl\u0131: etiket yazd\u0131rma."
            ),
            db_name=self.db_path,
        )

        self.assertTrue(result.success, result.message)
        self.assertEqual((result.data or {}).get("mode"), "answer")
        self.assertEqual((result.data or {}).get("tool_calls"), [])
        answer = self._answer_from(result)
        read_start = answer.index("Salt okuma:")
        approval_start = answer.index("Onayli prepare/execute/verify:")
        self.assertLess(read_start, approval_start)
        self.assertGreater(answer.index("Excel dosyasi"), approval_start)
        self.assertGreater(answer.index("Trendyol'dan veri"), approval_start)
        self.assertIn("authoritative read-back", answer)

    def test_generic_read_prepare_execute_explanation_stays_tool_free(self) -> None:
        result = run_agent_turn(
            (
                "OrderScout AI'da read, prepare ve execute arasindaki farki "
                "kisa anlat; veri sorgulama ve arac cagirma."
            ),
            planner_client=lambda _prompt: (
                '{"mode":"tool","tool":"orders.count","arguments":{"period":"today"}}'
            ),
            db_name=self.db_path,
        )

        self.assertTrue(result.success, result.message)
        self.assertEqual((result.data or {}).get("mode"), "answer")
        self.assertEqual((result.data or {}).get("tool_calls"), [])
        answer = self._answer_from(result)
        self.assertIn("Read:", answer)
        self.assertIn("Prepare:", answer)
        self.assertIn("Execute:", answer)
        self.assertIn("dry-run", answer)
        self.assertIn("verify/read-back", answer)

    def test_tool_do_not_run_boundary_does_not_cancel_natural_chat(self) -> None:
        result = run_agent_turn(
            "Naber başkan, hemen tool çalıştırma; doğal biçimde iki laf edelim.",
            chat_client=lambda _prompt: "İyiyim başkan, buradayım. Senin günün nasıl gidiyor?",
            db_name=self.db_path,
        )

        self.assertTrue(result.success, result.message)
        self.assertEqual((result.data or {}).get("mode"), "answer")
        self.assertEqual((result.data or {}).get("tool_calls"), [])
        answer = self._answer_from(result)
        self.assertIn("günün nasıl", answer)
        self.assertNotIn("islemi calistirmiyorum", answer)

    def test_two_words_small_talk_overrides_erroneous_order_count_plan(self) -> None:
        result = run_agent_turn(
            "Tool çalıştırma, sadece doğal biçimde halimi hatırımı sor ve iki laf et.",
            planner_client=lambda _prompt: (
                '{"mode":"tool","tool":"orders.count","arguments":{"period":"today"}}'
            ),
            chat_client=lambda _prompt: "Nasılsın başkan, günün nasıl gidiyor?",
            db_name=self.db_path,
        )

        self.assertTrue(result.success, result.message)
        self.assertEqual((result.data or {}).get("mode"), "answer")
        self.assertEqual((result.data or {}).get("tool_calls"), [])
        self.assertIn("günün nasıl", self._answer_from(result))

    def test_iki_laf_edelim_future_orders_transition_stays_tool_free(self) -> None:
        result = run_agent_turn(
            (
                "Selam başkan. Siparişlere hemen dalmayalım; önce iki laf edelim. "
                "Bugün nasılsın?"
            ),
            planner_client=lambda _prompt: "İyiyim başkan, buradayım. Sen nasılsın?",
            db_name=self.db_path,
        )

        self.assertTrue(result.success, result.message)
        self.assertEqual((result.data or {}).get("mode"), "answer")
        self.assertEqual((result.data or {}).get("tool_calls"), [])
        self.assertIn("İyiyim", self._answer_from(result))

    def test_deferred_data_then_future_orders_transition_stays_tool_free(self) -> None:
        def unavailable_provider(_prompt: str) -> str:
            raise RuntimeError("provider unavailable")

        result = run_agent_turn(
            (
                "Selam ba\u015fkan, bug\u00fcn acele etmeden biraz konu\u015fup "
                "sonra sipari\u015f operasyonunu birlikte toparlayal\u0131m; hemen "
                "veri \u00e7a\u011f\u0131rma."
            ),
            planner_client=unavailable_provider,
            chat_client=unavailable_provider,
            db_name=self.db_path,
        )

        self.assertTrue(result.success, result.message)
        self.assertEqual((result.data or {}).get("mode"), "answer")
        self.assertEqual((result.data or {}).get("tool_calls"), [])
        answer = self._answer_from(result).casefold()
        self.assertIn("veri veya tool cagirmiyorum", answer)
        self.assertIn("sen hazir oldugunda", answer)

    def test_small_talk_rejects_ungrounded_order_count_claim(self) -> None:
        result = run_agent_turn(
            "Tool çağırma; halimi hatırımı sorup doğal biçimde iki laf et.",
            planner_client=lambda _prompt: (
                "Bugün 5 Temmuz'da sistemde 0 sipariş var, sayaç sıfır."
            ),
            db_name=self.db_path,
        )

        self.assertTrue(result.success, result.message)
        self.assertEqual((result.data or {}).get("tool_calls"), [])
        answer = self._answer_from(result)
        self.assertNotIn("0 sipariş", answer)
        self.assertNotIn("sayac sifir", answer.casefold())
        self.assertIn("konusalim", answer.casefold())

    def test_ready_list_with_intervening_only_phrase_and_negative_label_is_read_only(self) -> None:
        result = run_agent_turn(
            (
                "Kargoya hazır siparişlerden yalnızca ilk üçünü listele; "
                "yazdırılmamış olanları tercih et ve hiçbir etiket üretme."
            ),
            planner_client=lambda _prompt: (
                '{"mode":"tool","tool":"operations.manual_action.route",'
                '"arguments":{"query":"etiket üret","mode":"write"}}'
            ),
            db_name=self.db_path,
        )

        self.assertTrue(result.success, result.message)
        tool_call = (result.data or {}).get("tool_call") or {}
        self.assertIn(
            tool_call.get("tool"),
            {"orders.ready_to_ship", "orders.query"},
        )
        self.assertFalse(
            bool((result.data or {}).get("approval_required")),
            result.data,
        )
        arguments = dict(tool_call.get("arguments") or {})
        self.assertEqual(arguments.get("limit"), 3)

    def test_ready_list_with_kesinlikle_negative_label_is_read_only(self) -> None:
        result = run_agent_turn(
            (
                "Kargoya hazır siparişlerden yalnız ilk üçünü listele; "
                "yazdırılmamış olanları tercih et ama kesinlikle etiket hazırlama "
                "veya üretme."
            ),
            planner_client=lambda _prompt: (
                '{"mode":"tool","tool":"operations.manual_action.route",'
                '"arguments":{"query":"etiket","mode":"write"}}'
            ),
            db_name=self.db_path,
        )

        self.assertTrue(result.success, result.message)
        tool_call = (result.data or {}).get("tool_call") or {}
        self.assertIn(tool_call.get("tool"), {"orders.ready_to_ship", "orders.query"})
        self.assertFalse(bool((result.data or {}).get("approval_required")))
        self.assertEqual((tool_call.get("arguments") or {}).get("limit"), 3)

    def test_ready_list_with_compound_negative_label_is_read_only(self) -> None:
        result = run_agent_turn(
            (
                "Kargoya hazır, yazdırılmamış siparişlerden yalnız ilk üçünü "
                "listele; etiket hazırlama veya üretme."
            ),
            planner_client=lambda _prompt: (
                '{"mode":"tool","tool":"operations.manual_action.route",'
                '"arguments":{"query":"etiket","mode":"write"}}'
            ),
            db_name=self.db_path,
        )

        self.assertTrue(result.success, result.message)
        tool_call = (result.data or {}).get("tool_call") or {}
        self.assertIn(tool_call.get("tool"), {"orders.ready_to_ship", "orders.query"})
        self.assertFalse(bool((result.data or {}).get("approval_required")))
        self.assertEqual((tool_call.get("arguments") or {}).get("limit"), 3)

    def test_return_list_with_negative_approval_phrase_never_prepares_write(self) -> None:
        result = run_agent_turn(
            (
                "Son 30 gündeki WaitingInAction iadeleri en yeniden eskiye "
                "listele; claim ve tutar görünsün. Onaylama veya değiştirme."
            ),
            planner_client=lambda _prompt: (
                '{"mode":"tool","tool":"returns.approve.prepare",'
                '"arguments":{"query":"waiting returns","limit":5}}'
            ),
            db_name=self.db_path,
        )

        self.assertTrue(result.success, result.message)
        tool_call = (result.data or {}).get("tool_call") or {}
        self.assertEqual(tool_call.get("tool"), "returns.query")
        self.assertFalse(bool((result.data or {}).get("approval_required")))
        self.assertEqual(
            ((tool_call.get("arguments") or {}).get("filters") or {}).get("status"),
            "WaitingInAction",
        )

    def test_return_list_preserves_explicit_en_yeni_numeric_limit(self) -> None:
        result = run_agent_turn(
            (
                "Son 30 gündeki WaitingInAction iadelerinden en yeni 2 tanesini "
                "claim, sipariş, müşteri, sebep ve tutarla listele. Onaylama veya "
                "değiştirme."
            ),
            db_name=self.db_path,
        )

        self.assertTrue(result.success, result.message)
        tool_call = (result.data or {}).get("tool_call") or {}
        self.assertEqual(tool_call.get("tool"), "returns.query")
        arguments = tool_call.get("arguments") or {}
        self.assertEqual(arguments.get("limit"), 2)
        self.assertEqual(arguments.get("page_size"), 2)

    def test_general_latest_return_list_uses_query_limit_and_selected_provenance(self) -> None:
        now_ms = int(datetime.now(TZ_GMT3).timestamp() * 1000)
        with Session(get_engine(self.db_path)) as session:
            session.add(
                ApiAccount(
                    pk=23,
                    account_id="seller-return-latest",
                    comp_name="Return Latest Seller",
                    platform="trendyol",
                    is_active=True,
                )
            )
            session.commit()
        for index in range(1, 5):
            self._insert_waiting_return(
                f"CLAIM-LATEST-{index}",
                f"ORD-RETURN-LATEST-{index}",
                now_ms + index * 1_000,
                first_name="Latest",
                last_name=str(index),
                api_account_id=23,
            )

        result = run_agent_turn(
            (
                "Önceki yanlış listeyi düzelt: production verideki en yeni 3 "
                "iade kaydını salt-okuma olarak "
                "listele; claim, sipariş, durum, sebep ve tutarı göster. Yalnız "
                "bu üç claim için hesap/ortam provenance'ını koru; hiçbir onay "
                "hazırlama veya veri değiştirme."
            ),
            planner_client=lambda _prompt: (
                '{"mode":"tool","tool":"returns.lookup",'
                '"arguments":{"query":"production iadeleri","limit":10}}'
            ),
            history=[
                {
                    "role": "assistant",
                    "content": "Önceki yanlış sonuçta 10 iade listelendi.",
                    "context": {
                        "kind": "return_selection",
                        "source_tool": "returns.lookup",
                        "claim_ids": [f"CLAIM-OLD-{index}" for index in range(10)],
                        "listed_count": 10,
                        "total_count": 10,
                    },
                }
            ],
            db_name=self.db_path,
        )

        self.assertTrue(result.success, result.message)
        tool_call = dict((result.data or {}).get("tool_call") or {})
        self.assertEqual(tool_call.get("tool"), "returns.query")
        self.assertEqual((tool_call.get("arguments") or {}).get("limit"), 3)
        context = dict((result.data or {}).get("conversation_context") or {})
        self.assertEqual(
            context.get("claim_ids"),
            ["CLAIM-LATEST-4", "CLAIM-LATEST-3", "CLAIM-LATEST-2"],
        )
        self.assertEqual(
            context.get("account_by_claim"),
            {
                "CLAIM-LATEST-4": "23",
                "CLAIM-LATEST-3": "23",
                "CLAIM-LATEST-2": "23",
            },
        )
        self.assertEqual(context.get("environment_by_account"), {"23": "production"})
        self.assertEqual(
            [row.get("store_name") for row in context.get("account_scopes") or []],
            ["Return Latest Seller"],
        )

    def test_visible_model_answer_normalizes_literal_escaped_newlines(self) -> None:
        decision = AgentDecision(
            mode="answer",
            response="Birinci iade:\\n1. CLAIM-ONE\\n2. CLAIM-TWO",
            source="llm",
        )
        with patch(
            "AI.processors.agent.choose_action",
            return_value=Result.ok(
                "answer",
                close_dialog=False,
                data={"decision": decision},
            ),
        ):
            result = run_agent_turn(
                "Bekleyen iki iadeyi listele.",
                planner_client=lambda _prompt: "Birinci iade:\\n1. CLAIM-ONE\\n2. CLAIM-TWO",
            )

        self.assertTrue(result.success, result.message)
        answer = self._answer_from(result)
        self.assertNotIn("\\n", answer)
        self.assertIn("\n1. CLAIM-ONE\n2. CLAIM-TWO", answer)

    def test_exact_turkish_date_scopes_count_revenue_and_profit_together(self) -> None:
        history = [
            {
                "role": "assistant",
                "content": "Son 7 gün sipariş özeti.",
                "context": {
                    "kind": "orders_metric_scope",
                    "source_tool": "orders.count",
                    "period": "last_7_days",
                    "start_date": "2026-06-29",
                    "end_date": "2026-07-05",
                },
            }
        ]
        result = run_agent_turn(
            (
                "25 Haziran 2026 günü için sipariş sayısı, ciro ve "
                "güvenilir biçimde hesaplanabiliyorsa tahmini kârı getir. Kâr verisi eksikse "
                "rakam uydurma."
            ),
            planner_client=lambda _prompt: (
                '{"mode":"tool","tool":"orders.lookup","arguments":'
                '{"query":"haziran gu ozelinde adedi ciro guvenilir tahmini kari",'
                '"limit":10}}'
            ),
            db_name=self.db_path,
            history=history,
        )

        self.assertTrue(result.success, result.message)
        calls = list((result.data or {}).get("tool_calls") or [])
        self.assertEqual(
            [row.get("tool") for row in calls],
            ["orders.count", "orders.revenue_total", "finance.profit_amount"],
        )
        for row in calls:
            arguments = dict(row.get("arguments") or {})
            self.assertEqual(arguments.get("start_date"), "2026-06-25")
            self.assertEqual(arguments.get("end_date"), "2026-06-25")
            self.assertNotIn("period", arguments)
        answer = self._answer_from(result)
        self.assertIn("siparis", answer.casefold())
        self.assertIn("ciro", answer.casefold())
        self.assertIn("karli oldugumuzu soyleyemem", answer.casefold())

    def test_direct_answer_repairs_reversed_style_instruction(self) -> None:
        responses = iter(
            (
                "Bundan sonra cevaplarimi tek cumleyle sunacagim.",
                (
                    "Anladim. Cevaplari tek cumleye sikistirmadan, "
                    "gerektigi kadar ayrintili anlatacagim."
                ),
            )
        )

        decision = AgentDecision(
            mode="answer",
            response="Tamam.",
            source="llm",
        )
        with patch(
            "AI.processors.agent.choose_action",
            return_value=Result.ok(
                "answer",
                close_dialog=False,
                data={"decision": decision},
            ),
        ):
            result = run_agent_turn(
                "cevaplari tek cumleye sikistirma, gerektigi kadar anlat",
                planner_client=lambda _prompt: next(responses),
            )

        self.assertTrue(result.success)
        answer = str((result.data or {}).get("answer") or "")
        self.assertIn("tek cumleye sikistirmadan", answer)
        self.assertTrue(
            any(
                row.get("phase") == "answer_repair"
                for row in list((result.data or {}).get("trace") or [])
            )
        )

    def test_style_instruction_overrides_erroneous_order_lookup_plan(self) -> None:
        result = run_agent_turn(
            (
                "Bundan sonra cevapları tek cümleye sıkıştırma; sorunun "
                "gerektirdiği kadar ayrıntılı ama lafı uzatmadan anlat."
            ),
            planner_client=lambda _prompt: (
                '{"mode":"tool","tool":"orders.lookup","arguments":'
                '{"query":"tek cümleye sıkıştırma","limit":10}}'
            ),
            db_name=self.db_path,
        )

        self.assertTrue(result.success, result.message)
        self.assertEqual((result.data or {}).get("mode"), "answer")
        self.assertEqual((result.data or {}).get("tool_calls"), [])
        self.assertIn("tek cumleye sikistirmadan", self._answer_from(result))

    def test_direct_answer_falls_back_when_model_ignores_style_twice(self) -> None:
        responses = iter(
            (
                "Iyiyim, buradayim.",
                "Iyiyim, buradayim.",
            )
        )
        decision = AgentDecision(
            mode="answer",
            response="Tamam.",
            source="llm",
        )
        with patch(
            "AI.processors.agent.choose_action",
            return_value=Result.ok(
                "answer",
                close_dialog=False,
                data={"decision": decision},
            ),
        ):
            result = run_agent_turn(
                "cevaplari tek cumleye sikistirma, gerektigi kadar anlat",
                planner_client=lambda _prompt: next(responses),
            )

        self.assertTrue(result.success)
        self.assertIn(
            "gerektirdigi kadar ayrintili",
            str((result.data or {}).get("answer") or ""),
        )

    def test_balanced_length_preference_persists_into_later_simple_advice(self) -> None:
        prompts: list[str] = []
        responses = iter(
            (
                "Uzun aciklama. " * 100,
                (
                    "Basit kural: telefonu kontrol saatine kadar uzak tut; saat "
                    "gelince 15 dakika bak ve sure bitince kapat."
                ),
            )
        )

        def client(prompt: str) -> str:
            prompts.append(prompt)
            return next(responses)

        decision = AgentDecision(mode="answer", response="Tamam.", source="llm")
        history = [
            {
                "role": "user",
                "content": (
                    "Cevaplari tek cumleye sikistirma ama gereksiz uzatma; "
                    "konu ne kadar gerektiriyorsa o kadar anlat."
                ),
            },
            {
                "role": "assistant",
                "content": "Anladim, dengeli anlatacagim.",
            },
        ]
        with patch(
            "AI.processors.agent.choose_action",
            return_value=Result.ok(
                "answer",
                close_dialog=False,
                data={"decision": decision},
            ),
        ):
            result = run_agent_turn(
                "Bu siniri bozmamak icin nasil basit bir kural koymaliyim?",
                planner_client=client,
                history=history,
            )

        self.assertTrue(result.success, result.message)
        answer = self._answer_from(result)
        self.assertLess(len(answer), 300)
        self.assertIn("15 dakika", answer)
        self.assertIn("Kalici kullanici stil tercihi", prompts[0])
        self.assertTrue(
            any(
                row.get("phase") == "answer_repair"
                for row in list((result.data or {}).get("trace") or [])
            )
        )

    def test_same_rule_shortening_correction_has_grounded_fallback(self) -> None:
        def unavailable_provider(_prompt: str) -> str:
            raise RuntimeError("quota exhausted")

        history = [
            {
                "role": "user",
                "content": "Telefon siniri icin basit bir kural soyle.",
            },
            {
                "role": "assistant",
                "content": "Fiziksel mesafe ve alarm kurali. " * 40,
            },
        ]
        result = run_agent_turn(
            (
                "Kural iyi ama yine gereginden uzun anlattin; ayni kurali bu "
                "kez net bicimde 2-4 cumlede ver."
            ),
            planner_client=unavailable_provider,
            chat_client=unavailable_provider,
            history=history,
        )

        self.assertTrue(result.success, result.message)
        self.assertEqual((result.data or {}).get("tool_calls"), [])
        answer = self._answer_from(result)
        self.assertLess(len(answer), 350)
        self.assertIn("15 dakika", answer)
        self.assertIn("ertesi gunun listesine", answer)

    def test_format_count_and_future_order_return_do_not_trigger_recent_read(self) -> None:
        result = run_agent_turn(
            (
                "Az \u00f6nceki telefon s\u0131n\u0131r\u0131 kural\u0131n\u0131 "
                "d\u00fczeltip 2-4 c\u00fcmlede ver; sonra da "
                "kald\u0131\u011f\u0131m\u0131z sipari\u015f i\u015fine "
                "d\u00f6nebilece\u011fimizi belirt ama hen\u00fcz veri "
                "\u00e7a\u011f\u0131rma."
            ),
            planner_client=lambda _prompt: (
                '{"mode":"tool","tool":"orders.recent",'
                '"arguments":{"period":"all","limit":2}}'
            ),
            db_name=self.db_path,
        )

        self.assertTrue(result.success, result.message)
        self.assertEqual((result.data or {}).get("mode"), "answer")
        self.assertEqual((result.data or {}).get("tool_calls"), [])
        answer = self._answer_from(result).casefold()
        self.assertIn("15 dakika", answer)
        self.assertIn("simdi veri cagirmiyorum", answer)
        self.assertIn("kaldigimiz siparis isine", answer)

    def test_direct_answer_uses_recent_orders_for_priority_fallback(self) -> None:
        responses = iter(
            (
                "Bence once neye karar vermek istedigimizi netlestirelim.",
                "Bence once neye karar vermek istedigimizi netlestirelim.",
            )
        )
        history = [
            {
                "role": "assistant",
                "content": "21 siparis listelendi.",
                "context": {
                    "kind": "order_selection",
                    "order_numbers": [
                        f"ORD-READY-{index:02d}"
                        for index in range(1, 22)
                    ],
                    "total_count": 21,
                },
            }
        ]
        decision = AgentDecision(
            mode="answer",
            response="Devam edebiliriz.",
            source="llm",
        )
        with patch(
            "AI.processors.agent.choose_action",
            return_value=Result.ok(
                "answer",
                close_dialog=False,
                data={"decision": decision},
            ),
        ):
            result = run_agent_turn(
                "tamam, sence bugun neye oncelik vermeliyim",
                planner_client=lambda _prompt: next(responses),
                history=history,
            )

        self.assertTrue(result.success)
        answer = str((result.data or {}).get("answer") or "")
        self.assertIn("21 siparisin", answer)
        self.assertIn("kargo", answer)

    def test_direct_answer_hides_degenerate_local_model_repetition(self) -> None:
        repeated = "**[" * 80 + '"role": "user", "content": "naber"'
        responses = iter((repeated, repeated))
        decision = AgentDecision(
            mode="answer",
            response="Iyiyim, buradayim.",
            source="fallback",
        )
        with patch(
            "AI.processors.agent.choose_action",
            return_value=Result.ok(
                "answer",
                close_dialog=False,
                data={"decision": decision},
            ),
        ):
            result = run_agent_turn(
                "naber dayi",
                planner_client=lambda _prompt: next(responses),
            )

        self.assertTrue(result.success)
        answer = str((result.data or {}).get("answer") or "")
        self.assertNotIn("**[", answer)
        self.assertNotIn('"role"', answer)
        self.assertEqual(
            (result.data or {}).get("answer_origin"),
            "deterministic_instruction_fallback",
        )

    def test_direct_answer_hides_local_model_user_echo(self) -> None:
        responses = iter(
            (
                '"naber dayi, nasil gidiyor"',
                '"naber dayi, nasil gidiyor"',
            )
        )
        decision = AgentDecision(
            mode="answer",
            response="Iyiyim, buradayim.",
            source="fallback",
        )
        with patch(
            "AI.processors.agent.choose_action",
            return_value=Result.ok(
                "answer",
                close_dialog=False,
                data={"decision": decision},
            ),
        ):
            result = run_agent_turn(
                "naber dayi, nasil gidiyor",
                planner_client=lambda _prompt: next(responses),
            )

        self.assertTrue(result.success)
        answer = str((result.data or {}).get("answer") or "")
        self.assertNotEqual(answer.strip('"'), "naber dayi, nasil gidiyor")
        self.assertIn("Iyiyim", answer)

    def test_direct_answer_hides_long_partial_user_echo(self) -> None:
        prompt = (
            "Açık konuşayım, sipariş ve iade yükü bazen beni bunaltıyor. "
            "Veriye bakmadan, bir ekip arkadaşı gibi iki cümleyle ne dersin?"
        )
        echoed = (
            '"Ben OrderScout AI. Sipariş ve iade yükü bazen beni bunaltıyor. '
            'Veriye bakmadan, bir ekip arkadaşı gibi iki cümleyle ne dersin?"'
        )
        responses = iter((echoed, echoed))
        decision = AgentDecision(
            mode="answer",
            response="Bu yükün zorlayıcı olması anlaşılır; önce en acil işi birlikte ayıralım.",
            source="fallback",
        )
        with patch(
            "AI.processors.agent.choose_action",
            return_value=Result.ok(
                "answer",
                close_dialog=False,
                data={"decision": decision},
            ),
        ):
            result = run_agent_turn(
                prompt,
                planner_client=lambda _prompt: next(responses),
            )

        self.assertTrue(result.success)
        answer = str((result.data or {}).get("answer") or "")
        self.assertNotEqual(answer, echoed)
        self.assertEqual(
            (result.data or {}).get("answer_origin"),
            "deterministic_chat_fallback",
        )

    def test_direct_answer_rejects_repeated_irrelevant_local_response(self) -> None:
        repeated = (
            "Anliyorum, bugun biraz agir geliyor gibi. Buradayim; "
            "istersen once neyin moralini bozdugunu konusalim."
        )
        responses = iter((repeated, repeated))
        history = [
            {
                "role": "assistant",
                "content": repeated,
            }
        ]
        decision = AgentDecision(
            mode="answer",
            response="Musteri memnuniyetini konusalim.",
            source="fallback",
        )
        with patch(
            "AI.processors.agent.choose_action",
            return_value=Result.ok(
                "answer",
                close_dialog=False,
                data={"decision": decision},
            ),
        ):
            result = run_agent_turn(
                "musteri memnuniyetinde sence en kritik sey ne",
                planner_client=lambda _prompt: next(responses),
                history=history,
            )

        self.assertTrue(result.success)
        answer = str((result.data or {}).get("answer") or "")
        self.assertIn("musteri memnuniyetinde", answer.casefold())
        self.assertIn("guven", answer.casefold())
        self.assertNotEqual(answer, repeated)

    def test_direct_answer_repairs_forgotten_structured_order_context(self) -> None:
        responses = iter(
            (
                "Nerede kalmistik, hangi konuya donelim?",
                (
                    "Kaldigimiz yerde ORD-CONTEXT-42 numarali siparisin "
                    "kalemleri ve iade durumuna bakiyorduk."
                ),
            )
        )
        history = [
            {
                "role": "assistant",
                "content": "Siparis detaylari listelendi.",
                "context": {
                    "kind": "order_selection",
                    "order_numbers": ["ORD-CONTEXT-42"],
                },
            }
        ]
        decision = AgentDecision(
            mode="answer",
            response="Devam edebiliriz.",
            source="llm",
        )
        with patch(
            "AI.processors.agent.choose_action",
            return_value=Result.ok(
                "answer",
                close_dialog=False,
                data={"decision": decision},
            ),
        ):
            result = run_agent_turn(
                "tamam, kaldigimiz yerden devam edebiliriz",
                planner_client=lambda _prompt: next(responses),
                history=history,
            )

        self.assertTrue(result.success)
        self.assertIn(
            "ORD-CONTEXT-42",
            str((result.data or {}).get("answer") or ""),
        )

    def test_profit_claim_guard_distinguishes_claim_from_explicit_uncertainty(self) -> None:
        self.assertTrue(
            _answer_claims_unsupported_profit(
                "Tahmini kar 41.501,34 TRY; bu nedenle karliyiz."
            )
        )
        self.assertFalse(
            _answer_claims_unsupported_profit(
                "Tahmini kar verisi yok; bu nedenle 'karliyiz' diyemem."
            )
        )
        self.assertFalse(
            _answer_claims_unsupported_profit(
                "Bu donemde siparis olmadigi icin karliyiz ya da zarardayiz "
                "diyemem; tahmini kar belirlenemedi."
            )
        )

    def test_zero_order_period_rejects_numeric_average_basket(self) -> None:
        tool_results = [
            {
                "tool": "orders.count",
                "data": {"count": 0, "period_label": "2026-06-29..2026-07-05"},
            },
            {
                "tool": "orders.revenue_total",
                "data": {"total": 0.0, "period_label": "2026-06-29..2026-07-05"},
            },
        ]
        self.assertFalse(
            _grounded_order_analysis_answer(
                (
                    "Sipariş Adedi: 0, Toplam Ciro: 0.00 TRY, "
                    "Ortalama Sepet Tutarı: 0.00 TRY."
                ),
                tool_results,
            )
        )

    def test_equal_period_synthesis_repairs_zero_order_numeric_average(self) -> None:
        result = run_agent_turn(
            (
                "Son 7 günü önceki 7 günle sipariş adedi, ciro ve ortalama sepet "
                "açısından karşılaştır; kâr ekleme."
            ),
            chat_client=lambda _prompt: (
                "Sipariş Adedi: 0, Toplam Ciro: 0.00 TRY, "
                "Ortalama Sepet Tutarı: 0.00 TRY."
            ),
            db_name=self.db_path,
        )

        self.assertTrue(result.success, result.message)
        answer = self._answer_from(result)
        self.assertIn("ortalama sepet olusmadi", answer.casefold())
        self.assertNotIn("Ortalama Sepet Tutarı: 0.00", answer)

    def test_profit_guard_rejects_cost_coverage_as_fulfillment_status(self) -> None:
        tool_results = [
            {
                "tool": "finance.profit_amount",
                "data": {
                    "period_label": "2026-06-25..2026-07-02",
                    "amount": None,
                    "estimated_profit_ready": False,
                    "revenue": 6814.81,
                    "complete_order_count": 0,
                    "incomplete_order_count": 46,
                    "missing_reasons": [
                        "ORD-COST-01: Eslesmis urun bulunamadi."
                    ],
                },
            }
        ]

        self.assertTrue(
            _should_use_deterministic_answer(
                "46 siparisin tamami henuz surec olarak tamamlanmadi.",
                tool_results,
            )
        )
        self.assertTrue(
            _should_use_deterministic_answer(
                (
                    "Gelir 6814,81 TRY, tamamlanmis siparislerden elde edilen "
                    "toplam ciro. Tamamlanmamis siparisler: 46 adet."
                ),
                tool_results,
            )
        )
        self.assertFalse(
            _answer_claims_unsupported_profit(
                "Gercekten karliyiz diyebilmek icin maliyetleri tamamlamaliyiz."
            )
        )
        answer = _deterministic_tool_answer(
            "bu verilerle karli miyiz",
            tool_results,
        )
        self.assertIn("maliyet verisinin eksik", answer)
        self.assertIn("teslimat veya operasyon durumunun", answer)
        self.assertIn("46 eksik siparis", answer)

    def test_full_context_guard_rejects_risk_free_claim_without_risk_evidence(self) -> None:
        tool_results = [
            {
                "tool": "orders.full_context",
                "data": {
                    "contexts": [
                        {
                            "order": {
                                "order_number": "ORD-RISK-GROUND",
                                "status": "ReadyToShip",
                            },
                            "items": [{"product_name": "Hedef Kalem"}],
                            "returns": [],
                        }
                    ]
                },
            }
        ]
        unsupported = (
            "ORD-RISK-GROUND siparisinde Hedef Kalem var, iade yok. "
            "Sistemde kayitli risk bulunmuyor."
        )
        unsupported_equivalent = (
            "Operasyonel durum sorunsuz gorunmektedir. "
            "Riskler: Kayitli bir risk bulunmamaktadir."
        )
        grounded = (
            "ORD-RISK-GROUND siparisinde Hedef Kalem var ve iade kaydi yok."
        )

        self.assertTrue(
            _answer_claims_unsupported_full_context_risk(
                unsupported,
                tool_results,
            )
        )
        self.assertTrue(_should_use_deterministic_answer(unsupported, tool_results))
        self.assertTrue(
            _answer_claims_unsupported_full_context_risk(
                unsupported_equivalent,
                tool_results,
            )
        )
        self.assertTrue(
            _should_use_deterministic_answer(
                unsupported_equivalent,
                tool_results,
            )
        )
        self.assertFalse(
            _answer_claims_unsupported_full_context_risk(
                grounded,
                tool_results,
            )
        )

    def test_full_context_synthesis_prompt_forbids_unmeasured_risk_claims(self) -> None:
        prompt = _synthesis_prompt(
            user_message="bunun kalemlerini ve iade durumunu ozetle",
            tool_name="orders.full_context",
            tool_result={
                "contexts": [
                    {
                        "order": {"order_number": "ORD-PROMPT-GROUND"},
                        "items": [{"product_name": "Hedef Kalem"}],
                        "returns": [],
                    }
                ]
            },
        )

        self.assertIn("fulfillment risk veya operasyon sagligi olcmez", prompt)
        self.assertIn("'risk yok'", prompt)
        self.assertIn("kalem ve iade ozeti", prompt)

    def test_full_context_money_guard_rejects_only_unverified_amounts(self) -> None:
        tool_results = [
            {
                "tool": "orders.full_context",
                "data": {
                    "contexts": [
                        {
                            "order": {"total_price": 249.99},
                            "items": [
                                {
                                    "quantity": 1,
                                    "price": 249.99,
                                    "line_total": 259.99,
                                }
                            ],
                            "returns": [],
                        }
                    ]
                },
            }
        ]

        self.assertFalse(
            _answer_conflicts_with_full_context_money(
                "Birim fiyat 249,99 TL, brut satir toplami 259.99 TRY.",
                tool_results,
            )
        )
        self.assertTrue(
            _answer_conflicts_with_full_context_money(
                "Siparis toplami 999,99 TL.",
                tool_results,
            )
        )
        self.assertTrue(
            _should_use_deterministic_answer(
                "Siparis toplami 999,99 TL.",
                tool_results,
            )
        )

    def test_full_context_formatter_includes_customer_total_and_item_fields(self) -> None:
        answer = _deterministic_tool_answer(
            "Tam bağlamını getir.",
            [
                {
                    "tool": "orders.full_context",
                    "arguments": {"query": "ORD-CONTEXT-1"},
                    "message": "Tam bağlam hazır.",
                    "data": {
                        "query": "ORD-CONTEXT-1",
                        "contexts": [
                            {
                                "order": {
                                    "order_number": "ORD-CONTEXT-1",
                                    "package_id": 801,
                                    "shipment_package_id": 801,
                                    "status": "ReadyToShip",
                                    "customer": "Ayşe Yılmaz",
                                    "total_price": 399.9,
                                    "currency": "TRY",
                                    "order_date": "2026-06-25 12:00",
                                    "cargo_provider": "Aras",
                                    "cargo_tracking_number": "TRACK-1",
                                },
                                "shipment_address": {
                                    "address": "Core Sokak No 1, Kadikoy, Istanbul, 34000",
                                    "city": "Istanbul",
                                    "district": "Kadikoy",
                                },
                                "header": {"is_printed": False},
                                "items": [
                                    {
                                        "product_name": "Test Ürün",
                                        "quantity": 2,
                                        "sku": "SKU-1",
                                        "barcode": "BAR-1",
                                        "price": 199.95,
                                        "line_total": 399.9,
                                        "status": "Created",
                                    }
                                ],
                                "returns": [],
                            }
                        ],
                    },
                }
            ],
        )

        self.assertIn("musteri: Ayşe Yılmaz", answer)
        self.assertIn("tutar: 399.9 TRY", answer)
        self.assertIn("paket: 801", answer)
        self.assertIn(
            "adres: Core Sokak No 1, Kadikoy, Istanbul, 34000",
            answer,
        )
        self.assertIn("sku: SKU-1", answer)
        self.assertIn("barkod: BAR-1", answer)
        self.assertIn("satir tutari: 399.9 TRY", answer)

    def test_live_evidence_redacts_full_context_address_and_package_id(self) -> None:
        raw_address = "Core Sokak No 1, Kadikoy, Istanbul, 34000"
        redacted, metadata = redact_live_evidence(
            {
                "tool_result": {
                    "shipment_package_id": 801,
                    "shipment_address": {
                        "address": raw_address,
                    },
                },
                "answer": f"paket: 801 | adres: {raw_address}",
            }
        )

        serialized = json.dumps(redacted, ensure_ascii=False)
        self.assertNotIn(raw_address, serialized)
        self.assertNotIn('"801"', serialized)
        self.assertIn("<package-001>", serialized)
        self.assertIn("<address-001>", serialized)
        self.assertEqual(
            (metadata.get("alias_counts") or {}).get("package"),
            1,
        )
        self.assertEqual(
            (metadata.get("alias_counts") or {}).get("address"),
            1,
        )

    def test_order_full_context_with_product_and_status_never_routes_to_listing(self) -> None:
        order_numbers = [f"ORD-HIST-{index}" for index in range(1, 11)]
        history = [
            {
                "role": "assistant",
                "content": "\n".join(
                    f"{index}. {order_number}"
                    for index, order_number in enumerate(order_numbers, start=1)
                ),
                "context": {
                    "kind": "order_selection",
                    "source_tool": "orders.lookup",
                    "order_numbers": order_numbers,
                    "listed_count": 10,
                },
            },
            {
                "role": "assistant",
                "content": "Daha yeni ama yanlış kapsamlı 20 sipariş listesi.",
                "context": {
                    "kind": "order_selection",
                    "source_tool": "orders.query",
                    "order_numbers": [
                        f"ORD-LATER-{index}" for index in range(1, 21)
                    ],
                    "listed_count": 20,
                },
            },
        ]
        result = run_agent_turn(
            (
                "25 Haziran raporundan önceki 10 siparişlik tarihsel listenin "
                "üçüncü siparişinin tam bağlamını yeniden "
                "göster; müşteri, ürün, kargo, takip, tutar ve durum dahil."
            ),
            planner_client=lambda _prompt: (
                '{"mode":"tool","tool":"listing.status_lookup","arguments":'
                '{"query":"tarihsel listedeki üçüncü","limit":5}}'
            ),
            db_name=self.db_path,
            history=history,
        )

        self.assertTrue(result.success, result.message)
        tool_call = (result.data or {}).get("tool_call") or {}
        self.assertEqual(tool_call.get("tool"), "orders.full_context")
        self.assertEqual(
            (tool_call.get("arguments") or {}).get("query"),
            "ORD-HIST-3",
        )

    def test_approval_answer_discloses_rows_hidden_after_preview_limit(self) -> None:
        answer = _deterministic_tool_answer(
            "bunlari etikete yazdir",
            [
                {
                    "tool": "labels.create_from_orders.prepare",
                    "data": {
                        "action": {
                            "approval_required": True,
                            "execute_enabled": True,
                            "summary": "22 siparis icin etiket olusturma hazirlaniyor.",
                            "rows": [
                                {
                                    "title": "Siparis etiketi olusturma onayi",
                                    "order_number": f"ORD-PREVIEW-{index:02d}",
                                }
                                for index in range(1, 23)
                            ],
                        }
                    },
                }
            ],
        )

        self.assertIn("22 siparis", answer)
        self.assertIn("ORD-PREVIEW-20", answer)
        self.assertNotIn("ORD-PREVIEW-21", answer)
        self.assertIn("2 kayit daha onay kapsaminda", answer)

    def test_explicit_complete_scope_lists_and_prepares_all_twenty_one_orders(self) -> None:
        orders = [
            {
                "order_number": f"ORD-COMPLETE-{index:02d}",
                "status": "ReadyToShip",
            }
            for index in range(1, 22)
        ]
        list_answer = _deterministic_tool_answer(
            "o 21 siparisin hepsini listele",
            [
                {
                    "tool": "orders.ready_to_ship",
                    "data": {"orders": orders},
                }
            ],
        )
        approval_answer = _deterministic_tool_answer(
            "bu 21 siparisin tamamini etikete yazdirmaya hazirla",
            [
                {
                    "tool": "labels.create_from_orders.prepare",
                    "data": {
                        "action": {
                            "approval_required": True,
                            "execute_enabled": True,
                            "summary": "21 siparis icin etiket olusturma hazirlaniyor.",
                            "rows": [
                                {
                                    "title": "Siparis etiketi olusturma onayi",
                                    "order_number": row["order_number"],
                                }
                                for row in orders
                            ],
                        }
                    },
                }
            ],
        )

        self.assertIn("ORD-COMPLETE-21", list_answer)
        self.assertNotIn("kayit daha var", list_answer)
        self.assertIn("ORD-COMPLETE-21", approval_answer)
        self.assertNotIn("kayit daha onay kapsaminda", approval_answer)

    def test_ready_count_rejects_wrong_reason_for_excluded_simulation_orders(self) -> None:
        tool_results = [
            {
                "tool": "orders.ready_to_ship_count",
                "message": "all doneminde kargoya hazir 22 siparis var.",
                "data": {
                    "period": "all",
                    "count": 22,
                    "simulated_excluded": 5,
                },
            }
        ]
        model_answer = (
            "Kargoya hazir 22 siparis var."
        )

        self.assertTrue(_should_use_deterministic_answer(model_answer, tool_results))
        answer = _deterministic_tool_answer("kargoya hazir kac siparis var", tool_results)
        self.assertIn("Kargoya hazir 22 siparis", answer)
        self.assertIn("5 simulasyon siparisi", answer)
        self.assertIn("dahil edilmedi", answer)

    def _insert_order(
        self,
        order_number: str,
        order_date_ms: int,
        *,
        package_id: int,
        first_name: str,
        last_name: str,
        status: str = "Created",
        cargo_provider_name: str | None = "Trendyol Express",
        cargo_tracking_number: str | None = None,
        total_price: float = 100.0,
        printed: bool = False,
        api_account_id: int | None = None,
    ) -> None:
        with Session(get_engine(self.db_path)) as session:
            header = OrderHeader(orderNumber=order_number, api_account_id=api_account_id, is_printed=printed, is_extracted=False)
            session.add(header)
            session.flush()
            session.add(
                OrderData(
                    order_header_id=header.pk,
                    orderNumber=order_number,
                    api_account_id=api_account_id,
                    id=package_id,
                    shipmentPackageId=package_id,
                    orderDate=order_date_ms,
                    lastModifiedDate=order_date_ms + package_id,
                    status=status,
                    shipmentPackageStatus=status,
                    customerFirstName=first_name,
                    customerLastName=last_name,
                    cargoProviderName=cargo_provider_name,
                    cargoTrackingNumber=f"TRK-{package_id}" if cargo_tracking_number is None else cargo_tracking_number,
                    totalPrice=total_price,
                    packageTotalPrice=total_price,
                )
            )
            session.commit()

    def _insert_order_item(
        self,
        order_number: str,
        now_ms: int,
        *,
        line_id: int = 1,
        product_name: str = "Demo Urun",
        cancel_reason: str | None = None,
    ) -> None:
        with Session(get_engine(self.db_path)) as session:
            session.add(
                OrderItem(
                    orderNumber=order_number,
                    order_data_id=1,
                    snapshotLastModifiedDate=now_ms,
                    quantity=2,
                    productName=product_name,
                    merchantSku="SKU-BEHAVIORAL",
                    barcode="BAR-BEHAVIORAL",
                    price=25.0,
                    lineId=line_id,
                    cancelReason=cancel_reason,
                )
            )
            session.commit()

    def _insert_waiting_return(
        self,
        claim_id: str,
        order_number: str,
        now_ms: int,
        *,
        first_name: str,
        last_name: str,
        api_account_id: int | None = None,
        reason: str | None = None,
    ) -> None:
        with Session(get_engine(self.db_path)) as session:
            session.add(
                ReturnClaim(
                    claimId=claim_id,
                    orderNumber=order_number,
                    api_account_id=api_account_id,
                    claimDate=now_ms,
                    orderDate=now_ms,
                    lastModifiedDate=now_ms + 1_000,
                    claimItemStatus="WaitingInAction",
                    claimReason=reason,
                    customerFirstName=first_name,
                    customerLastName=last_name,
                    totalPrice=55.0,
                    itemCount=1,
                    waitingClaimLineItemIds=["line-1"],
                )
            )
            session.commit()

    def _seed_ready_orders(self) -> int:
        now_ms = int(datetime.now(TZ_GMT3).timestamp() * 1000)
        self._insert_order(
            "ORD-MAHMUT-READY",
            now_ms,
            package_id=901,
            first_name="Mahmut",
            last_name="Şevket",
            status="ReadyToShip",
        )
        self._insert_order(
            "ORD-AYSE-READY",
            now_ms + 1_000,
            package_id=902,
            first_name="Ayse",
            last_name="Yilmaz",
            status="ReadyToShip",
        )
        return now_ms

    def _action_from(self, result: Result) -> dict:
        data = result.data or {}
        action = data.get("action_proposal")
        self.assertIsInstance(action, dict, result.message)
        self.assertTrue(action.get("approval_required"), action)
        self.assertTrue(action.get("execute_enabled"), action)
        return action

    def _answer_from(self, result: Result) -> str:
        return str((result.data or {}).get("answer") or result.message or "")

    def _assert_prepare_answer_is_safe(self, result: Result, *, expected_terms: tuple[str, ...] = ()) -> None:
        answer = self._answer_from(result)
        lowered = answer.casefold()
        self.assertIn("Onay gerekiyor", answer)
        self.assertIn("Gercek API istegi veya DB yazimi yapilmadi", answer)
        self.assertNotIn("db_name", lowered, answer)
        self.assertNotIn("test.sqlite3", lowered, answer)
        for forbidden in ("onaylandi", "degistirildi", "yapildi", "tamamlandi"):
            self.assertNotIn(forbidden, lowered, answer)
        for term in expected_terms:
            self.assertIn(term, answer)

    def _assert_execute_verified(self, result: Result) -> dict:
        self.assertTrue(result.success, result.message)
        verify = ((result.data or {}).get("verify") or {})
        self.assertTrue(verify.get("verified"), result.data)
        return verify

    def test_generic_waiting_label_request_ignores_stale_delivered_history_selection(self) -> None:
        now_ms = int(datetime.now(TZ_GMT3).timestamp() * 1000)
        self._insert_order(
            "ORD-STALE-DELIVERED", now_ms - 86_400_000, package_id=801,
            first_name="Stale", last_name="Customer", status="Delivered",
        )
        self._insert_order(
            "ORD-FRESH-READY", now_ms, package_id=802,
            first_name="Ready", last_name="Customer", status="Picking",
        )
        history = [
            {"role": "user", "content": "bu siparisin detayini goster"},
            {
                "role": "assistant",
                "content": "ORD-STALE-DELIVERED teslim edilmis.",
                "context": SelectionSnapshot.from_order_numbers(
                    ["ORD-STALE-DELIVERED"],
                    source_tool="orders.full_context",
                    period="all",
                ).to_context(),
            },
        ]

        result = run_agent_turn(
            "baskan bekleyen siparislerimi yazdir",
            planner_client=lambda _prompt: (
                '{"mode":"tool","tool":"labels.create_from_orders.prepare",'
                '"arguments":{"order_numbers":["ORD-STALE-DELIVERED"]}}'
            ),
            db_name=self.db_path,
            history=history,
        )

        self.assertTrue(result.success, result.message)
        action = self._action_from(result)
        payload = ((action.get("arguments") or {}).get("payload") or {})
        self.assertEqual(payload.get("order_numbers"), ["ORD-FRESH-READY"])
        self.assertNotIn("ORD-STALE-DELIVERED", payload.get("order_numbers") or [])

    def test_direct_generic_waiting_label_request_keeps_all_records_above_twenty(self) -> None:
        now_ms = int(datetime.now(TZ_GMT3).timestamp() * 1000)
        expected = [f"ORD-DIRECT-READY-{index:02d}" for index in range(25)]
        for index, order_number in enumerate(expected):
            self._insert_order(
                order_number,
                now_ms + index,
                package_id=8_200 + index,
                first_name="Direct",
                last_name=f"Ready{index}",
                status="ReadyToShip",
            )

        result = run_agent_turn(
            "baskan bekleyen siparislerimi yazdir",
            planner_client=lambda _prompt: (
                '{"mode":"tool","tool":"labels.create_from_orders.prepare",'
                '"arguments":{"query":"","limit":20}}'
            ),
            db_name=self.db_path,
        )

        self.assertTrue(result.success, result.message)
        action = self._action_from(result)
        payload = ((action.get("arguments") or {}).get("payload") or {})
        self.assertEqual(set(payload.get("order_numbers") or []), set(expected))
        serialized_action = json.dumps(action, ensure_ascii=False)
        self.assertNotIn("Direct Ready", serialized_action)
        self.assertNotIn("TRK-82", serialized_action)
        self.assertEqual(
            ((result.data or {}).get("tool_call") or {}).get("arguments", {}).get("limit"),
            500,
        )

    def test_generic_waiting_label_request_never_mixes_production_and_stage(self) -> None:
        with Session(get_engine(self.db_path)) as session:
            session.add_all(
                [
                    ApiAccount(
                        pk=1,
                        account_id="seller-production",
                        comp_name="Production Store",
                        platform="TRENDYOL",
                        extra_config={"environment": "production"},
                    ),
                    ApiAccount(
                        pk=2,
                        account_id="2738",
                        comp_name="Stage Store",
                        platform="TRENDYOL",
                        extra_config={"environment": "stage"},
                    ),
                ]
            )
            session.commit()
        now_ms = int(datetime.now(TZ_GMT3).timestamp() * 1000)
        self._insert_order(
            "ORD-PROD-READY-LABEL",
            now_ms,
            package_id=8_301,
            first_name="Production",
            last_name="Ready",
            status="ReadyToShip",
            api_account_id=1,
        )
        self._insert_order(
            "ORD-STAGE-READY-LABEL",
            now_ms + 1,
            package_id=8_302,
            first_name="Stage",
            last_name="Ready",
            status="Picking",
            api_account_id=2,
        )

        generic = run_agent_turn(
            "baskan bekleyen siparislerimi yazdir",
            planner_client=lambda _prompt: (
                '{"mode":"tool","tool":"labels.create_from_orders.prepare",'
                '"arguments":{"query":"","limit":20}}'
            ),
            db_name=self.db_path,
        )
        explicit_stage = run_agent_turn(
            "yalniz Stage magazasindaki bekleyen siparislerimi yazdir",
            planner_client=lambda _prompt: (
                '{"mode":"tool","tool":"labels.create_from_orders.prepare",'
                '"arguments":{"query":"","limit":20}}'
            ),
            db_name=self.db_path,
        )

        self.assertTrue(generic.success, generic.message)
        self.assertTrue(explicit_stage.success, explicit_stage.message)
        generic_action = self._action_from(generic)
        stage_action = self._action_from(explicit_stage)
        self.assertEqual(
            ((generic_action.get("arguments") or {}).get("payload") or {}).get(
                "order_numbers"
            ),
            ["ORD-PROD-READY-LABEL"],
        )
        self.assertEqual(
            ((stage_action.get("arguments") or {}).get("payload") or {}).get(
                "order_numbers"
            ),
            ["ORD-STAGE-READY-LABEL"],
        )
        self.assertEqual(
            ((generic.data or {}).get("tool_call") or {}).get("arguments", {}).get(
                "api_account_ids"
            ),
            [1],
        )
        self.assertEqual(
            ((explicit_stage.data or {}).get("tool_call") or {}).get(
                "arguments", {}
            ).get("api_account_ids"),
            [2],
        )
        self.assertEqual(
            [
                row.get("environment")
                for row in (generic_action.get("arguments") or {}).get(
                    "account_scopes"
                )
                or []
            ],
            ["production"],
        )
        self.assertEqual(
            [
                row.get("environment")
                for row in (stage_action.get("arguments") or {}).get(
                    "account_scopes"
                )
                or []
            ],
            ["stage"],
        )

    def test_exact_return_followup_keeps_stage_scope_without_repeating_stage(self) -> None:
        with Session(get_engine(self.db_path)) as session:
            session.add_all(
                [
                    ApiAccount(
                        pk=1,
                        account_id="seller-production",
                        comp_name="Production Store",
                        platform="TRENDYOL",
                        extra_config={"environment": "production"},
                    ),
                    ApiAccount(
                        pk=2,
                        account_id="2738",
                        comp_name="Stage Store",
                        platform="TRENDYOL",
                        extra_config={"environment": "stage"},
                    ),
                ]
            )
            session.commit()
        now_ms = int(datetime.now(TZ_GMT3).timestamp() * 1000)
        self._insert_waiting_return(
            "CLAIM-STAGE-FOLLOWUP",
            "ORD-STAGE-FOLLOWUP",
            now_ms,
            first_name="Stage",
            last_name="Customer",
            api_account_id=2,
        )
        listed = run_agent_turn(
            "Stage ortamindaki onay bekleyen iadelerin tamamini getir",
            planner_client=lambda _prompt: (
                '{"mode":"tool","tool":"returns.query","arguments":'
                '{"period":"all","filters":{"status":"WaitingInAction"},'
                '"limit":20,"mode":"list"}}'
            ),
            db_name=self.db_path,
        )
        self.assertTrue(listed.success, listed.message)
        history = [
            {
                "role": "user",
                "content": "Stage ortamindaki onay bekleyen iadelerin tamamini getir",
            },
            {
                "role": "assistant",
                "content": self._answer_from(listed),
                "context": dict(
                    (listed.data or {}).get("conversation_context") or {}
                ),
            },
        ]

        lookup = run_agent_turn(
            "claim id CLAIM-STAGE-FOLLOWUP olan iadeyi tek basina goster",
            planner_client=lambda _prompt: (
                '{"mode":"tool","tool":"returns.lookup",'
                '"arguments":{"query":"CLAIM-STAGE-FOLLOWUP","limit":10}}'
            ),
            db_name=self.db_path,
            history=history,
        )

        self.assertTrue(lookup.success, lookup.message)
        lookup_call = (lookup.data or {}).get("tool_call") or {}
        self.assertEqual(
            (lookup_call.get("arguments") or {}).get("api_account_ids"),
            [2],
        )
        self.assertEqual(
            [
                row.get("claim_id")
                for row in ((lookup.data or {}).get("tool_result") or {}).get(
                    "returns"
                )
                or []
            ],
            ["CLAIM-STAGE-FOLLOWUP"],
        )

    def test_profit_amount_inherits_safe_production_account_boundary(self) -> None:
        with Session(get_engine(self.db_path)) as session:
            session.add_all(
                [
                    ApiAccount(
                        pk=1,
                        account_id="seller-production",
                        comp_name="Production Store",
                        platform="TRENDYOL",
                        extra_config={"environment": "production"},
                    ),
                    ApiAccount(
                        pk=2,
                        account_id="2738",
                        comp_name="Stage Store",
                        platform="TRENDYOL",
                        extra_config={"environment": "stage"},
                    ),
                ]
            )
            session.commit()
        now_ms = int(datetime.now(TZ_GMT3).timestamp() * 1000)
        self._insert_order(
            "ORD-PROD-PROFIT-SCOPE",
            now_ms,
            package_id=8_401,
            first_name="Production",
            last_name="Profit",
            total_price=100.0,
            api_account_id=1,
        )
        self._insert_order(
            "ORD-STAGE-PROFIT-EXCLUDED",
            now_ms + 1,
            package_id=8_402,
            first_name="Stage",
            last_name="Excluded",
            total_price=900.0,
            api_account_id=2,
        )

        result = run_agent_turn(
            "bugunku verilerle gercekten karli miyiz, eksik veriyi acikca soyle",
            planner_client=lambda _prompt: (
                '{"mode":"tool","tool":"finance.profit_amount",'
                '"arguments":{"period":"today"}}'
            ),
            db_name=self.db_path,
        )

        self.assertTrue(result.success, result.message)
        call = (result.data or {}).get("tool_call") or {}
        self.assertEqual(call.get("tool"), "finance.profit_amount")
        self.assertEqual(
            (call.get("arguments") or {}).get("api_account_ids"),
            [1],
        )
        data = (result.data or {}).get("tool_result") or {}
        self.assertEqual(data.get("revenue"), 100.0)
        self.assertFalse(data.get("estimated_profit_ready"))

    def test_explicit_delivered_order_cannot_become_executable_label_action(self) -> None:
        now_ms = int(datetime.now(TZ_GMT3).timestamp() * 1000)
        self._insert_order(
            "ORD-DELIVERED-BLOCKED", now_ms, package_id=803,
            first_name="Delivered", last_name="Customer", status="Delivered",
        )

        result = labels_prepare_create_from_orders(
            db_name=self.db_path,
            order_numbers=["ORD-DELIVERED-BLOCKED"],
        )

        self.assertTrue(result.success, result.message)
        action = dict((result.data or {}).get("action") or {})
        self.assertFalse(action.get("execute_enabled"), action)
        self.assertTrue(action.get("warnings"), action)
        self.assertIn("Delivered", " ".join(action.get("warnings") or []))
        self.assertNotIn("bulunamadi", " ".join(action.get("warnings") or []))
        self.assertEqual(
            (((action.get("arguments") or {}).get("payload") or {}).get("order_numbers")),
            [],
        )

    def test_label_execute_rechecks_status_and_stops_before_file_generation(self) -> None:
        now_ms = int(datetime.now(TZ_GMT3).timestamp() * 1000)
        self._insert_order(
            "ORD-READY-THEN-DELIVERED", now_ms, package_id=804,
            first_name="Ready", last_name="ThenDelivered", status="ReadyToShip",
        )
        prepared = labels_prepare_create_from_orders(
            db_name=self.db_path,
            order_numbers=["ORD-READY-THEN-DELIVERED"],
        )
        action = dict((prepared.data or {}).get("action") or {})
        self.assertTrue(action.get("execute_enabled"), action)
        with Session(get_engine(self.db_path)) as session:
            row = session.exec(
                select(OrderData).where(
                    OrderData.orderNumber == "ORD-READY-THEN-DELIVERED"
                )
            ).first()
            self.assertIsNotNone(row)
            row.status = "Delivered"
            row.shipmentPackageStatus = "Delivered"
            session.add(row)
            session.commit()

        with patch(
            "Labels.processors.pipeline.create_order_label_from_orders"
        ) as create_mock:
            executed = operation_runtime.execute_manual_action(action)

        self.assertFalse(executed.success, executed.data)
        self.assertTrue((executed.data or {}).get("preflight_failed"), executed.data)
        self.assertFalse((executed.data or {}).get("file_output_performed"), executed.data)
        create_mock.assert_not_called()

    def test_pending_label_identity_question_uses_exact_scope_and_masks_pii(self) -> None:
        now_ms = int(datetime.now(TZ_GMT3).timestamp() * 1000)
        self._insert_order(
            "ORD-PENDING-IDENTITY", now_ms, package_id=805,
            first_name="Berat", last_name="Baskan", status="Delivered",
            cargo_tracking_number="PRIVATE-TRACKING-805",
        )
        history = [
            {
                "role": "assistant",
                "content": "Bir siparis etiketi onay bekliyor.",
                "context": {
                    "kind": "pending_action",
                    "source_tool": "labels.create_from_orders.prepare",
                    "action_type": "labels.action_create_labels_from_selected_orders",
                    "order_numbers": ["ORD-PENDING-IDENTITY"],
                    "selected_record_count": 1,
                },
            }
        ]
        planner_calls: list[str] = []
        decision_result = choose_action(
            "baskan bu siparis ne aq kimin siparisi bu",
            planner_client=lambda prompt: planner_calls.append(prompt) or (
                '{"mode":"tool","tool":"orders.lookup",'
                '"arguments":{"query":"teslim edilmis gosterdin","limit":10}}'
            ),
            history=history,
        )
        self.assertTrue(decision_result.success, decision_result.message)
        self.assertEqual(planner_calls, [])
        self.assertFalse((decision_result.data or {}).get("model_invoked"))

        synthesis_calls: list[str] = []
        result = run_agent_turn(
            "baskan bu siparis ne aq kimin siparisi bu",
            planner_client=lambda prompt: synthesis_calls.append(prompt) or (
                '{"mode":"tool","tool":"orders.lookup",'
                '"arguments":{"query":"teslim edilmis gosterdin","limit":10}}'
            ),
            db_name=self.db_path,
            history=history,
        )

        self.assertTrue(result.success, result.message)
        self.assertEqual(synthesis_calls, [])
        call = (result.data or {}).get("tool_call") or {}
        self.assertEqual(call.get("tool"), "orders.lookup")
        self.assertEqual(
            (call.get("arguments") or {}).get("order_numbers"),
            ["ORD-PENDING-IDENTITY"],
        )
        answer = self._answer_from(result)
        self.assertIn("Delivered", answer)
        self.assertIn("B*** B***", answer)
        self.assertNotIn("Berat Baskan", answer)
        self.assertNotIn("PRIVATE-TRACKING-805", answer)
        self.assertIn("Adres, takip ve paket", answer)

    def test_exact_order_lookup_can_verify_pending_scope_above_thirty(self) -> None:
        now_ms = int(datetime.now(TZ_GMT3).timestamp() * 1000)
        expected = [f"ORD-EXACT-SCOPE-{index:02d}" for index in range(35)]
        for index, order_number in enumerate(expected):
            self._insert_order(
                order_number,
                now_ms + index,
                package_id=8_300 + index,
                first_name="Exact",
                last_name=f"Scope{index}",
                status="Picking",
            )

        result = orders_lookup(
            db_name=self.db_path,
            query="",
            order_numbers=expected,
            limit=len(expected),
        )

        self.assertTrue(result.success, result.message)
        self.assertEqual((result.data or {}).get("count"), len(expected))
        self.assertEqual(
            {row.get("order_number") for row in (result.data or {}).get("orders") or []},
            set(expected),
        )

    def test_casual_chat_does_not_get_stuck_on_fixed_fallback(self) -> None:
        class BadChatClient:
            def __init__(self) -> None:
                self.prompts: list[str] = []

            def __call__(self, prompt: str) -> str:
                self.prompts.append(prompt)
                return "Buradayim. Ne yapalim?"

        client = BadChatClient()

        result = run_agent_turn("naber nasilsin", planner_client=client, chat_client=client, db_name=self.db_path)

        self.assertTrue(result.success, result.message)
        self.assertEqual((result.data or {}).get("mode"), "answer")
        self.assertIsNone((result.data or {}).get("tool_call"))
        self.assertGreater(len(client.prompts), 0)
        answer = self._answer_from(result)
        self.assertNotEqual(answer, "Buradayim. Ne yapalim?")
        self.assertIn("siparis", answer.casefold())

    def test_chat_before_orders_does_not_trigger_health_analysis(self) -> None:
        result = run_agent_turn(
            "Selam başkan, daha siparişlere girmeden biraz sohbet edelim. Bugün nasıl gidiyor?",
            planner_client=lambda _prompt: "İyi gidiyor, önce biraz sohbet edebiliriz.",
            db_name=self.db_path,
        )

        self.assertTrue(result.success, result.message)
        self.assertEqual((result.data or {}).get("mode"), "answer")
        self.assertEqual((result.data or {}).get("tool_calls"), [])
        self.assertIn("sohbet", self._answer_from(result))

    def test_emotional_open_chat_rejects_dismissive_llm_answer(self) -> None:
        result = run_agent_turn(
            "konusalim ya dayi keyfim hic yok",
            planner_client=lambda _prompt: '{"mode":"answer","response":"Anladim, konusmak istemediginizi rahatca bildirin. Yardimci olmami ister misiniz?"}',
            chat_client=lambda _prompt: "Anladim, konusmak istemediginizi rahatca bildirin. Yardimci olmami ister misiniz?",
            db_name=self.db_path,
            history=[{"role": "user", "content": "naber dayi"}],
        )

        self.assertTrue(result.success, result.message)
        self.assertEqual((result.data or {}).get("mode"), "answer")
        self.assertIsNone((result.data or {}).get("tool_call"))
        answer = self._answer_from(result)
        self.assertIn("Buradayim", answer)
        self.assertIn("canini", answer)
        self.assertNotIn("konusmak istemediginizi", answer)
        self.assertNotIn("Yardimci olmami ister misiniz", answer)

    def test_emotional_ecommerce_opinion_does_not_lookup_returns_or_tracking(self) -> None:
        result = run_agent_turn(
            "Biraz yoruldum; e-ticarette sürekli sipariş, kargo ve iade takip etmek insanı neden bu kadar tüketiyor sence?",
            planner_client=lambda _prompt: (
                "Belirsizlik, sürekli tetikte olma ve küçük operasyonların hiç "
                "bitmemesi zihni yoruyor; sorun tek tek işler değil, kesintisiz dikkat yükü."
            ),
            db_name=self.db_path,
        )

        self.assertTrue(result.success, result.message)
        self.assertEqual((result.data or {}).get("mode"), "answer")
        self.assertEqual((result.data or {}).get("tool_calls"), [])
        answer = self._answer_from(result)
        self.assertIn("Belirsizlik", answer)
        self.assertNotIn("iade buldum", answer)

    def test_emotional_ecommerce_opinion_has_relevant_fallback_when_provider_fails(self) -> None:
        def failed_provider(_prompt: str) -> str:
            raise RuntimeError("provider unavailable")

        result = run_agent_turn(
            "E-ticaret operasyonunun insanı neden yorduğuna dair fikrini sordum. Sence asıl zihinsel yük nereden geliyor?",
            planner_client=failed_provider,
            db_name=self.db_path,
        )

        self.assertTrue(result.success, result.message)
        self.assertEqual((result.data or {}).get("mode"), "answer")
        self.assertEqual((result.data or {}).get("tool_calls"), [])
        answer = self._answer_from(result)
        self.assertIn("bitmeyen dikkat gecisinden", answer)
        self.assertIn("belirsizlik", answer.casefold())

    def test_emotional_workload_how_question_gets_actionable_fallback(self) -> None:
        def failed_provider(_prompt: str) -> str:
            raise RuntimeError("provider unavailable")

        result = run_agent_turn(
            (
                "Biraz yorucu ge\u00e7ti; sat\u0131\u015f bask\u0131s\u0131 bir yandan "
                "operasyon bir yandan, sence insan bu zihinsel y\u00fck\u00fc nas\u0131l "
                "daha sa\u011fl\u0131kl\u0131 ta\u015f\u0131r? Ma\u011faza verisine bakma."
            ),
            planner_client=failed_provider,
            db_name=self.db_path,
        )

        self.assertTrue(result.success, result.message)
        self.assertEqual((result.data or {}).get("mode"), "answer")
        self.assertEqual((result.data or {}).get("tool_calls"), [])
        answer = self._answer_from(result).casefold()
        self.assertIn("zaman bloklarinda", answer)
        self.assertIn("kapanis listesine", answer)

    def test_no_data_actionable_habits_beats_store_overview_plan(self) -> None:
        result = run_agent_turn(
            (
                "Nedenini anlatt\u0131n ama nas\u0131l daha sa\u011fl\u0131kl\u0131 "
                "ta\u015f\u0131yaca\u011f\u0131m\u0131 atlad\u0131n; ma\u011faza "
                "verisine bakmadan bana uygulanabilir \u00fc\u00e7 "
                "al\u0131\u015fkanl\u0131k s\u00f6yle."
            ),
            planner_client=lambda _prompt: (
                '{"mode":"tools","calls":['
                '{"tool":"orders.count","arguments":{"period":"today"}},'
                '{"tool":"orders.revenue_total","arguments":{"period":"today"}}]}'
            ),
            db_name=self.db_path,
        )

        self.assertTrue(result.success, result.message)
        self.assertEqual((result.data or {}).get("mode"), "answer")
        self.assertEqual((result.data or {}).get("tool_calls"), [])
        answer = self._answer_from(result).casefold()
        self.assertIn("zaman bloklarinda", answer)
        self.assertIn("uc operasyon onceligi", answer)
        self.assertIn("kapanis listesine", answer)

    def test_inflected_no_data_correction_beats_store_overview_plan(self) -> None:
        result = run_agent_turn(
            (
                "Bu da olmad\u0131; veri istemedi\u011fim halde ma\u011faza analizi "
                "yapt\u0131n. \u015eimdi yeni hi\u00e7bir tool "
                "\u00e7a\u011f\u0131rmadan yaln\u0131z uygulanabilir \u00fc\u00e7 "
                "al\u0131\u015fkanl\u0131\u011f\u0131 s\u00f6yle."
            ),
            planner_client=lambda _prompt: (
                '{"mode":"tools","calls":['
                '{"tool":"orders.count","arguments":{"period":"today"}},'
                '{"tool":"orders.revenue_total","arguments":{"period":"today"}}]}'
            ),
            db_name=self.db_path,
        )

        self.assertTrue(result.success, result.message)
        self.assertEqual((result.data or {}).get("mode"), "answer")
        self.assertEqual((result.data or {}).get("tool_calls"), [])
        self.assertIn("zaman bloklarinda", self._answer_from(result).casefold())

    def test_explicit_no_data_orders_opinion_stays_tool_free(self) -> None:
        result = run_agent_turn(
            "Tüm mağaza verisini istemedim. Veri veya tool kullanmadan, düşük tutarlı tek kalem siparişlerde operasyon maliyetini azaltmak için en etkili iki yaklaşımı kendi görüşünle söyle.",
            planner_client=lambda _prompt: (
                "Birincisi paketleme adımlarını standartlaştırmak; ikincisi benzer "
                "siparişleri dalgalar halinde işleyerek işlem başı zamanı azaltmak."
            ),
            db_name=self.db_path,
        )

        self.assertTrue(result.success, result.message)
        self.assertEqual((result.data or {}).get("mode"), "answer")
        self.assertEqual((result.data or {}).get("tool_calls"), [])
        answer = self._answer_from(result)
        self.assertIn("paketleme", answer)
        self.assertNotIn("Siparis operasyon analizi", answer)

    def test_cross_domain_no_data_opinion_cannot_fall_into_store_overview(self) -> None:
        prompt = (
            "Veri çağırmadan kendi görüşünü söyle: küçük bir e-ticaret "
            "işletmesinde stok, ilan, finans, lojistik ve pazaryeri "
            "operasyonlarını tek sohbetten yönetmenin en büyük faydası ve "
            "en büyük riski nedir?"
        )
        answer = (
            "En büyük fayda, kararların aynı bağlamda birbirini beslemesi; "
            "en büyük risk ise yanlış bir varsayımın beş alana birden taşınmasıdır."
        )
        prompts: list[str] = []

        def client(runtime_prompt: str) -> str:
            prompts.append(runtime_prompt)
            return answer

        result = run_agent_turn(
            prompt,
            planner_client=client,
            chat_client=client,
            db_name=self.db_path,
            history=[
                {"role": "user", "content": "Mağaza özetini getir."},
                {
                    "role": "assistant",
                    "content": "2916 sipariş var.",
                    "context": {"total_count": 2916},
                },
            ],
        )

        self.assertTrue(result.success, result.message)
        self.assertEqual((result.data or {}).get("mode"), "answer")
        self.assertEqual((result.data or {}).get("tool_calls"), [])
        self.assertEqual(
            (result.data or {}).get("decision_path"),
            "explicit_no_tool_guard",
        )
        self.assertTrue((result.data or {}).get("model_invoked"))
        self.assertIn("en büyük fayda", self._answer_from(result).casefold())
        self.assertTrue(prompts)
        self.assertNotIn("2916", prompts[-1])
        self.assertIn("salt genel gorus", prompts[-1])

    def test_cross_domain_no_tool_correction_does_not_leak_old_session_summary(self) -> None:
        prompt = (
            "Önceki cevap eski ve ilgisiz bir sipariş oturumu özetiydi. "
            "Hiçbir yeni veri veya tool çağırmadan yalnız asıl sorumu cevapla: "
            "bu beş operasyon alanını tek sohbetten "
            "yönetmenin en büyük faydası ve en büyük riski nedir?"
        )
        answer = (
            "Haklısın; veri çağırmam yanlıştı. En büyük fayda ortak bağlam, "
            "en büyük risk ise hatalı bir varsayımın alanlar arasında yayılmasıdır."
        )

        result = run_agent_turn(
            prompt,
            planner_client=lambda _prompt: answer,
            chat_client=lambda _prompt: answer,
            db_name=self.db_path,
        )

        self.assertTrue(result.success, result.message)
        self.assertEqual((result.data or {}).get("tool_calls"), [])
        self.assertTrue((result.data or {}).get("model_invoked"))
        visible = self._answer_from(result)
        self.assertIn("ortak bağlam", visible)
        self.assertNotIn("etiket ile iade", visible)

    def test_stock_multi_metric_overview_keeps_all_requested_metrics(self) -> None:
        result = choose_action(
            "Stok genel envanter özetini getir; toplam aktif ürün, toplam miktar, "
            "düşük stok ve sıfır veya negatif stok sayılarını ayır."
        )

        self.assertTrue(result.success, result.message)
        decision = (result.data or {}).get("decision")
        self.assertEqual(getattr(decision, "mode", ""), "tools")
        self.assertEqual(
            [call.tool_name for call in getattr(decision, "tool_calls", ())],
            ["stock.inventory_overview", "stock.low_stock_count"],
        )

    def test_zero_or_negative_balance_count_is_not_widened_to_inventory_overview(self) -> None:
        result = choose_action(
            "Sıfır veya negatif bakiyeli kaç stok kalemi var?",
            planner_client=lambda _prompt: (
                '{"mode":"tool","tool":"stock.inventory_overview",'
                '"arguments":{"limit":20}}'
            ),
        )

        self.assertTrue(result.success, result.message)
        decision = (result.data or {}).get("decision")
        self.assertEqual(getattr(decision, "tool_name", ""), "stock.zero_or_negative_count")
        self.assertEqual(
            [call.tool_name for call in getattr(decision, "tool_calls", ())],
            ["stock.zero_or_negative_count"],
        )

    def test_stock_scalar_correction_recalls_authoritative_context_without_unrelated_tool(self) -> None:
        authoritative_answer = (
            "Sıfır veya negatif stok bakiyesine sahip toplam 2 adet ürün bulunmaktadır."
        )
        history = [
            {
                "role": "assistant",
                "content": "**Operasyon Özeti:**",
                "context": {
                    "kind": "scalar_tool_result",
                    "active_domain": "stock",
                    "source_tool": "stock.zero_or_negative_count",
                    "scalar_facts": {"count": 2},
                    "authoritative_answer": authoritative_answer,
                },
            }
        ]
        model_calls: list[str] = []

        def wrong_listing_client(prompt: str) -> str:
            model_calls.append(prompt)
            return (
                '{"mode":"tool","tool":"listing.title_lookup",'
                '"arguments":{"query":"önceki sayı"}}'
            )

        result = run_agent_turn(
            "Sadece başlık yazdın; sayı görünmüyor. Bir önceki tool "
            "sonucundaki sayıyı şimdi tek cümlede açıkça söyle.",
            planner_client=wrong_listing_client,
            chat_client=wrong_listing_client,
            db_name=self.db_path,
            history=history,
        )

        self.assertTrue(result.success, result.message)
        self.assertEqual((result.data or {}).get("tool_calls"), [])
        self.assertEqual((result.data or {}).get("answer"), authoritative_answer)
        self.assertEqual((result.data or {}).get("decision_path"), "contextual_scalar_recall_guard")
        self.assertFalse((result.data or {}).get("model_invoked"))
        self.assertEqual(model_calls, [])

    def test_terse_scalar_recall_uses_grounded_context_without_model(self) -> None:
        history = [
            {
                "role": "assistant",
                "content": "Sıfır veya negatif bakiyeli toplam 2 stok kalemi var.",
                "context": {
                    "kind": "scalar_tool_result",
                    "active_domain": "stock",
                    "source_tool": "stock.zero_or_negative_count",
                    "scalar_facts": {"count": 2},
                    "authoritative_answer": (
                        "Sıfır veya negatif bakiyeli toplam 2 stok kalemi var."
                    ),
                },
            }
        ]
        model_calls: list[str] = []

        result = run_agent_turn(
            "Sadece sayıyı tekrar söyle.",
            planner_client=lambda prompt: model_calls.append(prompt) or "yanlış",
            chat_client=lambda prompt: model_calls.append(prompt) or "yanlış",
            db_name=self.db_path,
            history=history,
        )

        self.assertTrue(result.success, result.message)
        self.assertEqual((result.data or {}).get("answer"), "2")
        self.assertEqual((result.data or {}).get("decision_path"), "contextual_scalar_recall_guard")
        self.assertFalse((result.data or {}).get("model_invoked"))
        self.assertEqual(model_calls, [])

    def test_total_active_stock_card_scalar_outranks_negated_critical_context(self) -> None:
        result = choose_action(
            "Şimdi kritik kapsamı bırak; toplam aktif stok kartı sayısını tek "
            "sayı olarak getir.",
            planner_client=lambda _prompt: (
                '{"mode":"tool","tool":"stock.inventory_overview",'
                '"arguments":{"limit":20}}'
            ),
        )

        self.assertTrue(result.success, result.message)
        decision = (result.data or {}).get("decision")
        self.assertEqual(getattr(decision, "tool_name", ""), "stock.item_count")
        self.assertEqual(
            [call.tool_name for call in getattr(decision, "tool_calls", ())],
            ["stock.item_count"],
        )

    def test_strict_active_stock_card_count_rejects_bloated_metric_synthesis(self) -> None:
        result = run_agent_turn(
            "Toplam aktif stok kartı sayısını tek sayı olarak getir.",
            planner_client=lambda _prompt: (
                '{"mode":"tool","tool":"stock.inventory_overview",'
                '"arguments":{"limit":20}}'
            ),
            chat_client=lambda _prompt: (
                "Aktif kart 0, toplam miktar 0, düşük stok 0 ve sıfır stok 0."
            ),
            db_name=self.db_path,
        )

        self.assertTrue(result.success, result.message)
        self.assertEqual(
            ((result.data or {}).get("tool_call") or {}).get("tool"),
            "stock.item_count",
        )
        self.assertRegex((result.data or {}).get("answer") or "", r"^\d+$")

    def test_stock_zero_and_low_definition_question_is_tool_free_and_context_recoverable(self) -> None:
        calls: list[str] = []

        def fail_if_called(prompt: str) -> str:
            calls.append(prompt)
            raise AssertionError("Stock definition answer must not call a model or tool.")

        question = (
            "İşlem yapmadan açıkla: sıfır/negatif stok ile düşük/kritik stok "
            "neden aynı kavram değil?"
        )
        first = run_agent_turn(
            question,
            planner_client=fail_if_called,
            chat_client=fail_if_called,
            db_name=self.db_path,
        )
        self.assertTrue(first.success, first.message)
        self.assertEqual((first.data or {}).get("tool_calls"), [])
        self.assertEqual(
            (first.data or {}).get("decision_path"),
            "stock_metric_definition_guard",
        )
        first_answer = (first.data or {}).get("answer") or ""
        self.assertIn("0 veya altında", first_answer)
        self.assertIn("minimum eşik", first_answer)
        self.assertIn("örtüşebilir", first_answer)

        history = [
            {"role": "user", "content": question},
            {"role": "assistant", "content": first_answer},
        ]
        correction = run_agent_turn(
            "Sayıları tekrar etme; veri çekme ve işlem yapma. Yalnız iki "
            "kavramın tanım farkını açıkla.",
            planner_client=fail_if_called,
            chat_client=fail_if_called,
            db_name=self.db_path,
            history=history,
        )
        self.assertTrue(correction.success, correction.message)
        self.assertEqual((correction.data or {}).get("tool_calls"), [])
        correction_answer = (correction.data or {}).get("answer") or ""
        self.assertIn("minimum eş", correction_answer)
        self.assertNotRegex(correction_answer, r"\d")
        self.assertEqual(correction_answer.count("."), 1)
        self.assertEqual(calls, [])

    def test_stock_low_row_followup_beats_exclusion_metric_mention(self) -> None:
        result = choose_action(
            "Bu özetteki 2 düşük veya kritik stok kaydını şimdi listele; SKU, "
            "ad, mevcut bakiye ve minimum stok eşiğini göster. Sıfır veya "
            "negatif olanlarla düşük stok kapsamını birbirine karıştırma."
        )

        self.assertTrue(result.success, result.message)
        decision = (result.data or {}).get("decision")
        self.assertEqual(getattr(decision, "mode", ""), "tool")
        self.assertEqual(getattr(decision, "tool_name", ""), "stock.low_stock")

    def test_stock_low_row_list_preserves_explicit_maximum_limit(self) -> None:
        result = choose_action(
            "O iki dusuk stok kaydini listele; en fazla 2 kayit, salt okuma."
        )

        self.assertTrue(result.success, result.message)
        decision = (result.data or {}).get("decision")
        self.assertEqual(getattr(decision, "tool_name", ""), "stock.low_stock")
        self.assertEqual(getattr(decision, "arguments", {}).get("limit"), 2)

    def test_stock_low_row_correction_repeats_list_instead_of_inventory_overview(self) -> None:
        result = choose_action(
            "Açık alanını atladın; aynı düşük/kritik stok listesini yeniden getir "
            "ve her satırda SKU, bakiye, eşik, açık değerlerinin dördünü de yaz.",
            planner_client=lambda _prompt: (
                '{"mode":"tool","tool":"stock.inventory_overview",'
                '"arguments":{"limit":20}}'
            ),
        )

        self.assertTrue(result.success, result.message)
        decision = (result.data or {}).get("decision")
        self.assertEqual(getattr(decision, "tool_name", ""), "stock.low_stock")

    def test_stock_row_synthesis_cannot_drop_requested_gap_field(self) -> None:
        tool_results = [
            {
                "tool": "stock.low_stock",
                "data": {
                    "items": [
                        {
                            "sku": "SKU-1",
                            "name": "Test ürün",
                            "balance": -2,
                            "threshold": 3,
                            "gap": 5,
                        }
                    ]
                },
            }
        ]
        prompt = (
            "Düşük stok listesini her satırda SKU, bakiye, eşik ve açık ile göster."
        )

        self.assertTrue(
            _should_use_deterministic_answer(
                "SKU-1: bakiye -2, eşik 3.",
                tool_results,
                user_message=prompt,
            )
        )
        self.assertFalse(
            _should_use_deterministic_answer(
                "SKU-1: bakiye -2, eşik 3, açık 5.",
                tool_results,
                user_message=prompt,
            )
        )

    def test_stock_row_formatter_includes_low_stock_threshold_and_gap(self) -> None:
        answer = _format_stock_rows(
            [
                {
                    "sku": "SKU-1",
                    "name": "Test ürün",
                    "balance": -2,
                    "unit": "adet",
                    "threshold": 3,
                    "gap": 5,
                }
            ]
        )

        self.assertIn("bakiye: -2 adet", answer)
        self.assertIn("esik: 3", answer)
        self.assertIn("acik: 5", answer)

    def test_stock_multi_metric_deterministic_answer_keeps_exact_zero_count(self) -> None:
        answer = _deterministic_tool_answer(
            "Aktif kart, toplam miktar, dusuk stok ve tam sifir stoku ayir.",
            [
                {
                    "tool": "stock.inventory_overview",
                    "data": {
                        "active_items": 8,
                        "total_balance": 168,
                        "zero_or_negative_items": [
                            {"balance": -6},
                            {"balance": 0},
                        ],
                    },
                },
                {"tool": "stock.low_stock_count", "data": {"count": 2}},
            ],
        )

        self.assertIn("aktif stok karti 8", answer)
        self.assertIn("toplam miktar 168", answer)
        self.assertIn("dusuk stok 2", answer)
        self.assertIn("sifir/eksi stok 2", answer)
        self.assertIn("tam sifir stok 1", answer)
        self.assertTrue(
            _should_use_deterministic_answer(
                "8 aktif stok karti ve toplam 168 bakiye var. 2 dusuk stok var.",
                [
                    {"tool": "stock.inventory_overview", "data": {}},
                    {"tool": "stock.low_stock_count", "data": {}},
                ],
            )
        )

    def test_explicit_listing_metrics_beat_prior_order_followup_context(self) -> None:
        result = choose_action(
            "Şimdi Listing tarafına geç: toplam ilan, aktif ilan, satışta olan "
            "ilan ve dikkat isteyen ilan sayılarını tek cevapta ayır.",
            history=[
                {"role": "user", "content": "Son siparişleri getir."},
                {
                    "role": "assistant",
                    "content": "Bir sipariş seçildi.",
                    "context": {
                        "kind": "order_selection",
                        "order_numbers": ["ORD-1"],
                    },
                },
            ],
        )

        self.assertTrue(result.success, result.message)
        decision = (result.data or {}).get("decision")
        self.assertEqual(getattr(decision, "mode", ""), "tools")
        self.assertEqual(
            [call.tool_name for call in getattr(decision, "tool_calls", ())],
            ["listing.counts", "listing.attention_count"],
        )

    def test_listing_attention_row_request_beats_count_field_and_keeps_requested_limit(self) -> None:
        result = choose_action(
            "Bu 14 dikkat isteyen ilandan ilk 5 ornegi listele; baslik, barkod, "
            "stok kodu, adet, satis durumu ve dikkat nedenini goster. Hicbir ilani degistirme."
        )

        self.assertTrue(result.success, result.message)
        decision = (result.data or {}).get("decision")
        self.assertEqual(getattr(decision, "mode", ""), "tool")
        self.assertEqual(getattr(decision, "tool_name", ""), "listing.attention")
        self.assertEqual(getattr(decision, "arguments", {}), {"limit": 5})

    def test_listing_attention_read_only_inflection_repairs_wrong_model_search(self) -> None:
        result = run_agent_turn(
            (
                "Dikkat isteyen ilanlardan ilk 2 kaydi getir; baslik, barkod, stok kodu, "
                "miktar, satis durumu ve nedenleri goster. Hicbir kaydi degistirme."
            ),
            planner_client=lambda _prompt: (
                '{"mode":"tool","tool":"listing.search",'
                '"arguments":{"query":"dikkat isteyen ilanlar","limit":10}}'
            ),
            chat_client=lambda _prompt: "Yanlis model cevabi.",
            db_name=self.db_path,
        )

        self.assertTrue(result.success, result.message)
        self.assertEqual((result.data or {}).get("mode"), "tool")
        self.assertEqual((result.data or {}).get("tool_call", {}).get("tool"), "listing.attention")
        self.assertEqual((result.data or {}).get("tool_call", {}).get("arguments"), {"limit": 2})
        self.assertFalse((result.data or {}).get("approval_required"))
        self.assertFalse((result.data or {}).get("execute_enabled"))

    def test_listing_attention_formatter_includes_status_and_reasons(self) -> None:
        answer = _format_listing_rows(
            [
                {
                    "title": "Test ilan",
                    "barcode": "BAR-1",
                    "stock_code": "SKU-1",
                    "quantity": 0,
                    "approved": True,
                    "on_sale": False,
                    "reasons": ["out_of_stock", "not_on_sale"],
                }
            ]
        )

        self.assertIn("stok: 0", answer)
        self.assertIn("durum: onayli, satista degil", answer)
        self.assertIn("neden: out_of_stock, not_on_sale", answer)

    def test_listing_prepare_only_request_does_not_become_cancellation(self) -> None:
        result = choose_action(
            "SCK23872737723 barkodlu ilanin stok miktarini 1 yapmaya hazirlik yap; "
            "sadece onay taslagi olustur, calistirma. Hicbir API veya DB yazimi yapma."
        )

        self.assertTrue(result.success, result.message)
        decision = (result.data or {}).get("decision")
        self.assertEqual(getattr(decision, "mode", ""), "tool")
        self.assertEqual(getattr(decision, "tool_name", ""), "listing.update_stock.prepare")
        self.assertEqual(getattr(decision, "arguments", {}).get("query"), "SCK23872737723")
        self.assertEqual(getattr(decision, "arguments", {}).get("new_quantity"), 1)

    def test_explicit_vazgec_still_cancels_a_named_preparation(self) -> None:
        result = choose_action(
            "Vazgec; bu stok hazirligini iptal et, calistirma ve hicbir API veya DB islemi yapma."
        )

        self.assertTrue(result.success, result.message)
        decision = (result.data or {}).get("decision")
        self.assertEqual(getattr(decision, "mode", ""), "answer")
        self.assertIn("calistirmiyorum", getattr(decision, "response", ""))

    def test_read_only_listing_identifier_repairs_polluted_model_query(self) -> None:
        prompt = (
            "Iptalden sonra SCK23872737723 barkodlu ilanin mevcut stok miktarini "
            "yeniden getir. Sadece oku; hicbir hazirlik veya degisiklik yapma."
        )
        result = choose_action(
            prompt,
            planner_client=lambda _prompt: (
                '{"mode":"tool","tool":"listing.stock_lookup",'
                '"arguments":{"query":"iptalden sonra sck23872737723 mevcut yeniden getir",'
                '"limit":5}}'
            ),
        )

        self.assertTrue(result.success, result.message)
        decision = (result.data or {}).get("decision")
        self.assertEqual(getattr(decision, "tool_name", ""), "listing.stock_lookup")
        self.assertEqual(
            getattr(decision, "arguments", {}),
            {"query": "SCK23872737723", "limit": 5},
        )

    def test_finance_read_only_metrics_do_not_become_write_route(self) -> None:
        result = choose_action(
            "Simdi Finance tarafina gec: bu ay icin tahmini kar tutarini ve hazir "
            "olup olmadigini, nakit akisi netini ve kayitli gider kalemi "
            "sayisini tek cevapta ayir. Yalnizca yerel salt okuma yap; hicbir "
            "islem hazirlama."
        )

        self.assertTrue(result.success, result.message)
        decision = (result.data or {}).get("decision")
        self.assertEqual(getattr(decision, "mode", ""), "tools")
        self.assertEqual(
            [call.tool_name for call in getattr(decision, "tool_calls", ())],
            ["finance.profit_amount", "finance.cashflow_net", "finance.cost_item_count"],
        )

    def test_finance_revenue_profit_pair_does_not_drop_other_requested_metrics(self) -> None:
        result = choose_action(
            "Finance tarafina gecelim. Bu ay icin gelir, kar hesaplama hazirligi ve "
            "net nakit akisiyla tanimli maliyet kalemi sayisini ayri ayri ozetle. "
            "Salt okuma yap, finans islemi hazirlama."
        )

        self.assertTrue(result.success, result.message)
        decision = (result.data or {}).get("decision")
        self.assertEqual(getattr(decision, "mode", ""), "tools")
        self.assertEqual(
            [call.tool_name for call in getattr(decision, "tool_calls", ())],
            [
                "orders.revenue_total",
                "finance.profit_amount",
                "finance.cashflow_net",
                "finance.cost_item_count",
            ],
        )

    def test_plain_salt_okuma_instruction_cannot_trigger_finance_write(self) -> None:
        result = choose_action(
            "Bu 5 aktif gider kalemini ilk 5 olarak listele; ad, kategori, "
            "hesaplama tipi, deger ve aktiflik durumunu goster. Salt okuma yap."
        )

        self.assertTrue(result.success, result.message)
        decision = (result.data or {}).get("decision")
        self.assertEqual(getattr(decision, "mode", ""), "tool")
        self.assertEqual(getattr(decision, "tool_name", ""), "finance.cost_catalog")
        self.assertEqual(getattr(decision, "arguments", {}), {"limit": 5})

    def test_finance_cost_catalog_does_not_read_ilan_inside_varsayilan(self) -> None:
        result = choose_action(
            "Bu 5 maliyet kalemini listele; ad, kategori, hesaplama tipi, "
            "varsayilan deger ve aktiflik durumunu goster. Kategori alani "
            "kayitta yoksa acikca soyle. Salt okuma."
        )

        self.assertTrue(result.success, result.message)
        decision = (result.data or {}).get("decision")
        self.assertEqual(getattr(decision, "tool_name", ""), "finance.cost_catalog")
        self.assertEqual(getattr(decision, "arguments", {}), {"limit": 5})

    def test_cost_catalog_rejects_value_type_as_unsupported_category(self) -> None:
        tool_results = [
            {
                "tool": "finance.cost_catalog",
                "data": {
                    "cost_items": [
                        {
                            "name": "Komisyon bedeli",
                            "value_type": "percentage",
                            "default_amount": 0,
                            "currency": "TRY",
                            "is_active": True,
                        }
                    ]
                },
            }
        ]

        self.assertTrue(
            _should_use_deterministic_answer(
                "Kategori: percentage | Hesaplama Tipi: percentage",
                tool_results,
            )
        )
        answer = _deterministic_tool_answer(
            "Kategori ve hesaplama tipini goster.",
            tool_results,
        )
        self.assertIn("Kategori alani", answer)
        self.assertIn("hesaplama tipi: percentage", answer)
        self.assertIn("varsayilan tutar: 0", answer)
        self.assertIn("aktif: True", answer)

    def test_logistics_profile_keeps_requested_cheapest_limit(self) -> None:
        result = choose_action(
            "Simdi Logistics tarafina gec: lojistik profilini ve en ucuz ilk 3 "
            "tarifeyi ozetle; firma, platform, ucret ve desi araligini goster. "
            "Salt okuma yap."
        )

        self.assertTrue(result.success, result.message)
        decision = (result.data or {}).get("decision")
        self.assertEqual(getattr(decision, "tool_name", ""), "logistics.profile")
        self.assertEqual(getattr(decision, "arguments", {}), {"limit": 3})

    def test_logistics_multi_metric_profile_beats_trendyol_named_lookup(self) -> None:
        result = choose_action(
            "Logistics tarafina gec. Aktif saglayici ve aktif tarife toplamlarini, "
            "TRENDYOL kapsamindaki tarife sayisini ve en ucuz ilk 2 tarifeyi "
            "saglayici, platform, desi araligi ve tutarla getir. Salt okuma."
        )

        self.assertTrue(result.success, result.message)
        decision = (result.data or {}).get("decision")
        self.assertEqual(getattr(decision, "tool_name", ""), "logistics.profile")
        self.assertEqual(getattr(decision, "arguments", {}), {"limit": 2})

    def test_logistics_profile_separates_provider_and_tariff_platform_counts(self) -> None:
        providers = [
            SimpleNamespace(platformCode="TRENDYOL"),
            SimpleNamespace(platformCode="TRENDYOL"),
        ]
        tariffs = [
            SimpleNamespace(
                isActive=True,
                isDeleted=False,
                platform="TRENDYOL",
                amount=34.16,
                tariffKey="ty-1",
                name="Trendyol tarife",
                cargoProviderName="PTT",
                currency="TRY",
                minDesi=0,
                maxDesi=1,
            ),
            SimpleNamespace(
                isActive=True,
                isDeleted=False,
                platform="HEPSIBURADA",
                amount=50,
                tariffKey="hb-1",
                name="Hepsiburada tarife",
                cargoProviderName="Aras",
                currency="TRY",
                minDesi=0,
                maxDesi=2,
            ),
            SimpleNamespace(isActive=False, isDeleted=False, platform="TRENDYOL", amount=1),
        ]
        with patch(
            "AI.processors.tools.logistics_repo.list_providers",
            return_value=Result.ok("ok", data={"records": providers}),
        ), patch(
            "AI.processors.tools.logistics_repo.list_tariffs",
            return_value=Result.ok("ok", data={"records": tariffs}),
        ):
            result = logistics_profile(db_name=self.db_path, limit=2)

        self.assertTrue(result.success, result.message)
        data = dict(result.data or {})
        self.assertEqual(data.get("providers_by_platform"), {"TRENDYOL": 2})
        self.assertEqual(
            data.get("tariffs_by_platform"),
            {"TRENDYOL": 1, "HEPSIBURADA": 1},
        )
        answer = _deterministic_tool_answer("Lojistik profili goster", [{"tool": "logistics.profile", "data": data}])
        self.assertIn("Firma dagilimi: TRENDYOL: 2", answer)
        self.assertIn("Tarife dagilimi: HEPSIBURADA: 1, TRENDYOL: 1", answer)
        self.assertTrue(
            _should_use_deterministic_answer(
                "2 firma ve 2 aktif tarife var.",
                [{"tool": "logistics.profile", "data": data}],
            )
        )

    def test_named_logistics_followup_uses_carrier_as_lookup_query(self) -> None:
        result = choose_action(
            "Bu ucludeki Aras tarifesini adla yeniden bul; tarife adi, provider, "
            "platform, ucret, para birimi ve min max desiyi goster. Yalnizca salt "
            "okuma yap."
        )

        self.assertTrue(result.success, result.message)
        decision = (result.data or {}).get("decision")
        self.assertEqual(getattr(decision, "tool_name", ""), "logistics.tariff_lookup")
        self.assertEqual(getattr(decision, "arguments", {}).get("query"), "Aras")
        self.assertEqual(getattr(decision, "arguments", {}).get("limit"), 1)

    def test_marketplace_partition_overview_is_read_only_unfiltered_and_limited(self) -> None:
        result = choose_action(
            "Simdi Marketplaces tarafina gec: pazaryeri katalogunu ozetle; toplam, "
            "aktif, kabul edilmis ve bekleyen kayit sayilarini ayir, ilk 5 kaydi "
            "kod, ad, aktiflik, kabul ve kaynakla listele. Salt okuma yap."
        )

        self.assertTrue(result.success, result.message)
        decision = (result.data or {}).get("decision")
        self.assertEqual(getattr(decision, "tool_name", ""), "marketplaces.catalog")
        self.assertEqual(getattr(decision, "arguments", {}), {"limit": 5})

    def test_explicit_marketplace_domain_cancels_stale_ready_scope_recovery(self) -> None:
        history = [
            {"role": "user", "content": "Toplam bekleyen kargoya hazir siparisleri getir."},
            {
                "role": "assistant",
                "content": "11 kargoya hazir siparis var.",
                "context": {
                    "kind": "order_selection",
                    "source_tool": "orders.ready_to_ship",
                    "period": "all",
                    "order_numbers": ["ORD-1"],
                    "total_count": 11,
                },
            },
        ]
        result = choose_action(
            "Simdi Marketplaces tarafina gec: pazaryeri katalogunu ozetle; toplam, "
            "aktif, kabul edilmis ve bekleyen kayit sayilarini ayir, ilk 5 kaydi "
            "kod, ad, aktiflik, kabul ve kaynakla listele. Salt okuma yap.",
            history=history,
        )

        self.assertTrue(result.success, result.message)
        decision = (result.data or {}).get("decision")
        self.assertEqual(getattr(decision, "tool_name", ""), "marketplaces.catalog")
        self.assertEqual(getattr(decision, "arguments", {}), {"limit": 5})

    def test_explicit_marketplace_catalog_beats_stale_listing_status_context(self) -> None:
        result = choose_action(
            "Marketplaces tarafina gec. Pazaryeri katalogunu ozetle; toplam, aktif, "
            "kabul edilmis ve bekleyen sayilarini ayir, ilk 5 kaydi kod, ad, "
            "aktiflik, kabul durumu ve kaynakla listele. Salt okuma.",
            history=[
                {"role": "user", "content": "Dikkat isteyen ilanlari getir."},
                {"role": "assistant", "content": "Uc ilan satista degil."},
            ],
        )

        self.assertTrue(result.success, result.message)
        decision = (result.data or {}).get("decision")
        self.assertEqual(getattr(decision, "tool_name", ""), "marketplaces.catalog")
        self.assertEqual(getattr(decision, "arguments", {}), {"limit": 5})

    def test_marketplace_account_health_summary_beats_invalid_count_only(self) -> None:
        result = choose_action(
            "Marketplaces hesap sagligi ozetini getir; toplam ve aktif hesaplari, "
            "gecersiz token sayisini ve platform dagilimini ayir. Salt okuma yap."
        )

        self.assertTrue(result.success, result.message)
        decision = (result.data or {}).get("decision")
        self.assertEqual(getattr(decision, "tool_name", ""), "accounts.health")

        answer = _deterministic_tool_answer(
            "Hesap sagligini ozetle.",
            [
                {
                    "tool": "accounts.health",
                    "data": {
                        "accounts_total": 2,
                        "accounts_active": 1,
                        "invalid_tokens": 1,
                        "by_platform": {
                            "TRENDYOL": {"total": 2, "active": 1, "invalid_token": 1}
                        },
                    },
                }
            ],
        )
        self.assertIn("toplam 2", answer)
        self.assertIn("aktif 1", answer)
        self.assertIn("gecersiz token 1", answer)

    def test_marketplace_health_beats_catalog_even_with_do_not_mix_instruction(self) -> None:
        result = choose_action(
            "Marketplaces hesap sagligi ozetini getir; toplam ve aktif hesaplari, "
            "gecersiz token sayisini ve platform dagilimini ayir. Katalog "
            "kayitlariyla hesaplari karistirma, salt okuma."
        )

        self.assertTrue(result.success, result.message)
        decision = (result.data or {}).get("decision")
        self.assertEqual(getattr(decision, "tool_name", ""), "accounts.health")

    def test_magaza_verisine_bakma_general_opinion_stays_tool_free(self) -> None:
        result = run_agent_turn(
            (
                "Sence küçük bir e-ticaret operasyonunda insanı en çok yoran şey "
                "nedir? Genel fikrini söyle; mağaza verisine bakma."
            ),
            planner_client=lambda _prompt: (
                '{"mode":"tools","calls":['
                '{"tool":"orders.count","arguments":{"period":"all"}},'
                '{"tool":"stock.low_stock_count","arguments":{}}]}'
            ),
            chat_client=lambda _prompt: (
                "Bence en yorucu taraf, sürekli değişen öncelikler arasında odak değiştirmek."
            ),
            db_name=self.db_path,
        )

        self.assertTrue(result.success, result.message)
        self.assertEqual((result.data or {}).get("mode"), "answer")
        self.assertEqual((result.data or {}).get("tool_calls"), [])
        answer = self._answer_from(result)
        self.assertIn("odak", answer.casefold())

    def test_no_data_cost_opinion_has_specific_fallback_when_provider_fails(self) -> None:
        def failed_provider(_prompt: str) -> str:
            raise RuntimeError("provider unavailable")

        result = run_agent_turn(
            "Mağaza analizi istemiyorum. Veri ve tool kullanmadan kendi görüşünle söyle: düşük tutarlı tek kalem siparişlerde işlem başı maliyeti azaltmak için en etkili iki yaklaşım nedir?",
            planner_client=failed_provider,
            db_name=self.db_path,
        )

        self.assertTrue(result.success, result.message)
        self.assertEqual((result.data or {}).get("tool_calls"), [])
        answer = self._answer_from(result)
        self.assertIn("standart akış", answer)
        self.assertIn("dalga halinde", answer)
        self.assertNotIn("netlestirelim", answer)

    def test_cost_opinion_repairs_cross_customer_single_shipment_advice(self) -> None:
        class AdviceClient:
            def __init__(self) -> None:
                self.calls = 0

            def __call__(self, _prompt: str) -> str:
                self.calls += 1
                if self.calls == 1:
                    return (
                        "Aynı bölgedeki düşük tutarlı siparişleri bir araya "
                        "toplayıp tek bir kargo gönderimi yap."
                    )
                return (
                    "Paketleme adımlarını standartlaştır ve her müşteri "
                    "siparişini ayrı gönderi olarak koruyarak benzer işleri "
                    "operasyon içinde dalga halinde işle."
                )

        client = AdviceClient()
        result = run_agent_turn(
            "Veri çağırmadan düşün: düşük tutarlı tek kalem siparişlerde işlem başı maliyeti azaltmak için iki somut öneri ver.",
            planner_client=client,
            db_name=self.db_path,
        )

        self.assertTrue(result.success, result.message)
        self.assertGreaterEqual(client.calls, 2)
        answer = self._answer_from(result)
        self.assertIn("ayrı gönderi", answer)
        self.assertNotIn("tek bir kargo gönderimi", answer)

    def test_advice_correction_followup_stays_tool_free(self) -> None:
        def failed_provider(_prompt: str) -> str:
            raise RuntimeError("provider unavailable")

        result = run_agent_turn(
            "Geri aldığını söylemek yetmez; düzeltilmiş cevabı ver. Her müşteri gönderisini ayrı tut, standart paketleme ve operasyon dalgasında toplu işleme yaklaşımını iki cümlede açıkla.",
            planner_client=failed_provider,
            db_name=self.db_path,
        )

        self.assertTrue(result.success, result.message)
        self.assertEqual((result.data or {}).get("mode"), "answer")
        self.assertEqual((result.data or {}).get("tool_calls"), [])
        answer = self._answer_from(result)
        self.assertIn("ayrı gönderi", answer)
        self.assertIn("operasyon dalgasında", answer)
        self.assertNotIn("Siparis operasyon analizi", answer)

    def test_casual_chat_filters_generic_helpdesk_answer(self) -> None:
        generic = "Iyidir yegen, ugrasip duruyoruz iste. OrderScout AI olarak buradayim, sana nasil yardimci olabilirim?"

        result = run_agent_turn(
            "naber dayi napiyon",
            planner_client=lambda _prompt: '{"mode":"answer","response":"' + generic + '"}',
            chat_client=lambda _prompt: generic,
            db_name=self.db_path,
        )

        self.assertTrue(result.success, result.message)
        self.assertEqual((result.data or {}).get("mode"), "answer")
        self.assertIsNone((result.data or {}).get("tool_call"))
        answer = self._answer_from(result)
        self.assertIn("konusalim", answer.casefold())
        self.assertNotIn("nasil yardimci olabilirim", answer.casefold())

    def test_orders_capability_question_answers_with_actions_not_lookup(self) -> None:
        result = run_agent_turn(
            "siparisler tarafinda hangi islemleri yapabiliyorsun",
            planner_client=lambda _prompt: '{"mode":"tool","tool":"orders.lookup","arguments":{"query":"siparisler tarafinda hangi islemleri yapabiliyorsun","limit":10}}',
            chat_client=lambda _prompt: "Buradayim. Ne yapalim?",
            db_name=self.db_path,
        )

        self.assertTrue(result.success, result.message)
        self.assertEqual((result.data or {}).get("mode"), "answer")
        self.assertIsNone((result.data or {}).get("tool_call"))
        answer = self._answer_from(result)
        self.assertIn("etiket", answer)
        self.assertIn("iade", answer)
        self.assertIn("kargo", answer)
        self.assertIn("Trendyol", answer)
        self.assertNotIn("Lokal kayitta", answer)

    def test_orders_capability_question_ignores_count_style_llm_override(self) -> None:
        result = run_agent_turn(
            "siparislerimle ilgili neler yapabiliriz",
            planner_client=lambda _prompt: '{"mode":"tool","tool":"orders.count","arguments":{"period":"today"}}',
            chat_client=lambda _prompt: "2026-06-10 doneminde 21 siparis var.",
            db_name=self.db_path,
        )

        self.assertTrue(result.success, result.message)
        self.assertEqual((result.data or {}).get("mode"), "answer")
        self.assertIsNone((result.data or {}).get("tool_call"))
        answer = self._answer_from(result)
        self.assertIn("etiket", answer)
        self.assertIn("iade", answer)
        self.assertIn("kargo", answer)
        self.assertIn("Trendyol", answer)
        self.assertNotIn("21 siparis", answer.casefold())

    def test_orders_capability_question_keeps_comprehensive_natural_llm_answer(self) -> None:
        natural_answer = (
            "Siparislerini listeleyip detay ve kalemlerini bulabilirim. Kargoya hazir paketlerde etiket yazdirma "
            "ve kargo firmasi degisikligi icin onayli akis hazirlarim; iadeleri kontrol edip onaylayabilir, "
            "Trendyol senkronizasyonu yapabilir ve kar Excel raporu olusturabilirim. Hangisinden baslayalim?"
        )
        result = run_agent_turn(
            "siparislerimle ilgili neler yapabiliriz",
            planner_client=lambda _prompt: '{"mode":"tool","tool":"orders.count","arguments":{"period":"today"}}',
            chat_client=lambda _prompt: natural_answer,
            db_name=self.db_path,
        )

        self.assertTrue(result.success, result.message)
        self.assertEqual((result.data or {}).get("mode"), "answer")
        self.assertIsNone((result.data or {}).get("tool_call"))
        self.assertEqual(self._answer_from(result), natural_answer)

    def test_orders_capability_guard_removes_unsupported_invoice_pdf_claim(self) -> None:
        natural_answer = (
            "Siparişleri listeleyip detay ve kalemlerini bulabilirim.\n"
            "Kargo ve etiket işlerini onaylı akışla hazırlar, iadeleri yönetirim. "
            "Faturasını çıkarıp PDF olarak saklarım.\n"
            "Kâr Excel raporu da hazırlayabilirim."
        )
        result = run_agent_turn(
            "Sipariş operasyonunda benimle baştan sona neleri yapabilirsin?",
            planner_client=lambda _prompt: natural_answer,
            db_name=self.db_path,
        )

        self.assertTrue(result.success, result.message)
        answer = self._answer_from(result)
        self.assertIn("Siparişleri listeleyip", answer)
        self.assertIn("Kargo ve etiket", answer)
        self.assertIn("Kâr Excel", answer)
        self.assertNotIn("Faturasını çıkarıp", answer)
        self.assertIn("fatura/PDF üretimini desteklenen", answer)
        self.assertEqual(
            (result.data or {}).get("answer_origin"),
            "direct_llm_with_capability_guard",
        )

    def test_orders_capability_correction_does_not_route_to_finance_write(self) -> None:
        result = run_agent_turn(
            "Finance kaydı hazırlamanı istemiyorum. Az önceki Orders yetenek cevabını düzelt: yalnız kanıtlı sipariş çıktıları ve onaylı aksiyonları anlat; fatura/PDF iddiasını geri çek.",
            planner_client=lambda _prompt: "Kanıtlı Orders yeteneklerini düzeltiyorum.",
            db_name=self.db_path,
        )

        self.assertTrue(result.success, result.message)
        self.assertEqual((result.data or {}).get("mode"), "answer")
        self.assertEqual((result.data or {}).get("tool_calls"), [])
        answer = self._answer_from(result)
        self.assertNotIn("finance.save_finance_invoice_document", answer)
        self.assertNotIn("Eksik payload", answer)

    def test_orders_capability_question_keeps_natural_turkish_unicode_answer(self) -> None:
        natural_answer = (
            "Siparişleri tarih veya müşteriyle listeleyebilirim; kargo firmasını değiştirme ve etiket yazdırma "
            "işlemlerini onaya hazırlayabilirim. İadeleri inceleyip onay akışını da yürütebilirim. "
            "İstersen önce son siparişlerden başlayalım."
        )
        result = run_agent_turn(
            "siparişlerimle ilgili neler yapabiliriz",
            planner_client=lambda _prompt: '{"mode":"tool","tool":"orders.count","arguments":{"period":"today"}}',
            chat_client=lambda _prompt: natural_answer,
            db_name=self.db_path,
        )

        self.assertTrue(result.success, result.message)
        self.assertEqual((result.data or {}).get("mode"), "answer")
        self.assertIsNone((result.data or {}).get("tool_call"))
        self.assertEqual(self._answer_from(result), natural_answer)

    def test_orders_capability_question_completes_llm_answer_when_returns_are_missing(self) -> None:
        incomplete_answer = (
            "Siparişleri listeleyebilirim, kargoya hazır paketleri getirebilir ve etiket yazdırma "
            "işlemlerini onaya hazırlayabilirim."
        )
        result = run_agent_turn(
            "siparişlerimle ilgili neler yapabiliriz",
            planner_client=lambda _prompt: '{"mode":"tool","tool":"orders.count","arguments":{"period":"today"}}',
            chat_client=lambda _prompt: incomplete_answer,
            db_name=self.db_path,
        )

        self.assertTrue(result.success, result.message)
        answer = self._answer_from(result).casefold()
        self.assertIn("etiket", answer)
        self.assertIn("kargo", answer)
        self.assertIn("iade", answer)
        self.assertNotEqual(answer, incomplete_answer.casefold())

    def test_colloquial_orders_health_question_uses_monthly_analysis_not_today_count(self) -> None:
        now_ms = int(datetime.now(TZ_GMT3).timestamp() * 1000)
        self._insert_order(
            "ORD-HEALTH-COLLOQUIAL",
            now_ms,
            package_id=903,
            first_name="Mahmut",
            last_name="Sevket",
            status="ReadyToShip",
            total_price=125.0,
        )

        result = run_agent_turn(
            "siparisler ne alemde dayi kotu muyuz",
            planner_client=lambda _prompt: '{"mode":"tool","tool":"orders.count","arguments":{"period":"today"}}',
            chat_client=lambda _prompt: "Bugun 0 siparis var, kotu gorunuyor.",
            db_name=self.db_path,
        )

        self.assertTrue(result.success, result.message)
        self.assertEqual((result.data or {}).get("mode"), "tools")
        tool_calls = (result.data or {}).get("tool_calls") or []
        call_args = {row.get("tool"): row.get("arguments") or {} for row in tool_calls}
        self.assertEqual((call_args.get("orders.count") or {}).get("period"), "this_month")
        self.assertIn("finance.profit_amount", call_args)
        self.assertIn("orders.ready_to_ship_count", call_args)
        answer = self._answer_from(result)
        self.assertIn("Siparis operasyon analizi", answer)
        self.assertIn("Once net karar", answer)
        self.assertNotIn("Bugun 0 siparis", answer)

    def test_orders_capability_complaint_retry_does_not_lookup_complaint_text(self) -> None:
        history = [
            {"role": "user", "content": "siparisler tarafinda hangi islemleri yapabiliyoruz"},
            {
                "role": "assistant",
                "content": "Lokal kayitta 'siparisler tarafinda hangi islemleri yapabiliyoruz' icin siparis bulunamadi.",
            },
            {"role": "user", "content": "ne demek bu"},
            {
                "role": "assistant",
                "content": "Haklisin, onceki cevap hataliydi; bunu siparis aramasi gibi yorumlamamam gerekiyordu.",
            },
        ]

        result = run_agent_turn(
            "dalga mi geciyosun siparislerle ilgili hangi islemleri yapabiliyosun",
            planner_client=lambda _prompt: '{"mode":"tool","tool":"orders.lookup","arguments":{"query":"dalga mi geciyosun siparislerle ilgili hangi islemleri yapabiliyosun","limit":10}}',
            chat_client=lambda _prompt: "Lokal kayda gore 10 siparis buldum.",
            db_name=self.db_path,
            history=history,
        )

        self.assertTrue(result.success, result.message)
        self.assertEqual((result.data or {}).get("mode"), "answer")
        self.assertIsNone((result.data or {}).get("tool_call"))
        answer = self._answer_from(result)
        self.assertIn("etiket", answer)
        self.assertIn("iade", answer)
        self.assertIn("kargo", answer)
        self.assertNotIn("Lokal kayda gore", answer)
        self.assertNotIn("10 siparis", answer.casefold())

    def test_contextual_order_list_capability_question_keeps_agentic_actions(self) -> None:
        self._seed_ready_orders()

        listed = run_agent_turn("kargoya hazir siparisleri getir", db_name=self.db_path)
        self.assertTrue(listed.success, listed.message)
        listed_answer = self._answer_from(listed)
        self.assertIn("ORD-MAHMUT-READY", listed_answer)
        self.assertIn("ORD-AYSE-READY", listed_answer)

        result = run_agent_turn(
            "bu listeyle neler yapabiliriz?",
            planner_client=lambda _prompt: '{"mode":"answer","response":"Paketleyip kargoya teslim edebilirsiniz."}',
            chat_client=lambda _prompt: "Paketleyip kargoya teslim edebilirsiniz.",
            db_name=self.db_path,
            history=[
                {"role": "user", "content": "kargoya hazir siparisleri getir"},
                {"role": "assistant", "content": listed_answer},
            ],
        )

        self.assertTrue(result.success, result.message)
        self.assertEqual((result.data or {}).get("mode"), "answer")
        self.assertIsNone((result.data or {}).get("tool_call"))
        answer = self._answer_from(result).casefold()
        for term in ("etiket", "iade", "kargo", "excel"):
            self.assertIn(term, answer)
        self.assertNotIn("Lokal kayitta", self._answer_from(result))

    def test_contextual_order_list_capability_keeps_useful_llm_answer(self) -> None:
        self._seed_ready_orders()

        listed = run_agent_turn("kargoya hazir siparisleri getir", db_name=self.db_path)
        self.assertTrue(listed.success, listed.message)
        listed_answer = self._answer_from(listed)

        useful_answer = (
            "Bu siparisler icin etiket yazdirma surecini baslatabilirim, kargo firmasi degisikligi hazirlayabilirim "
            "ya da islem tamamlandiysa fis cikarildi olarak isaretleyebilirim."
        )
        result = run_agent_turn(
            "bu listeyle neler yapabiliriz?",
            planner_client=lambda _prompt: json.dumps({"mode": "answer", "response": useful_answer}, ensure_ascii=False),
            chat_client=lambda _prompt: "Bu cevap kullanilmamali.",
            db_name=self.db_path,
            history=[
                {"role": "user", "content": "kargoya hazir siparisleri getir"},
                {"role": "assistant", "content": listed_answer},
            ],
        )

        self.assertTrue(result.success, result.message)
        answer = self._answer_from(result)
        self.assertEqual(answer, useful_answer)
        self.assertIn("etiket", answer)
        self.assertIn("kargo", answer)
        self.assertIn("fis", answer)
        self.assertNotIn("kar Excel cikisi", answer)

    def test_orders_end_to_end_capability_question_does_not_degrade_to_count(self) -> None:
        prompt = (
            "yükümü azaltmak için siparişler tarafında benimle gerçekten "
            "neleri baştan sona yapabilirsin? sadece liste sayma anlatma"
        )
        result = run_agent_turn(
            prompt,
            planner_client=lambda _prompt: json.dumps(
                {
                    "mode": "tool",
                    "tool": "orders.count",
                    "arguments": {"period": "today"},
                },
                ensure_ascii=False,
            ),
            chat_client=lambda _prompt: "Sadece saymak yerine sipariş operasyonunu uçtan uca yürütebilirim.",
            db_name=self.db_path,
        )

        self.assertTrue(result.success, result.message)
        self.assertEqual((result.data or {}).get("mode"), "answer")
        self.assertIsNone((result.data or {}).get("tool_call"))
        answer = self._answer_from(result).casefold()
        for term in ("etiket", "kargo", "iade", "senkronizasyon"):
            self.assertIn(term, answer)
        self.assertNotIn("siparis var", answer)

    def test_ordinal_followup_opens_only_the_selected_order(self) -> None:
        self._seed_ready_orders()
        listed = run_agent_turn(
            "kargoya hazir siparisleri getir",
            db_name=self.db_path,
        )
        self.assertTrue(listed.success, listed.message)
        listed_data = listed.data or {}
        listed_context = dict(listed_data.get("conversation_context") or {})
        order_numbers = list(listed_context.get("order_numbers") or [])
        self.assertGreaterEqual(len(order_numbers), 2)

        result = run_agent_turn(
            "bunların içinden ikinci sıradakinin ürün kalemlerini ve varsa iade durumunu aç",
            planner_client=lambda _prompt: json.dumps(
                {
                    "mode": "tool",
                    "tool": "orders.full_context",
                    "arguments": {"query": "ikinci siradaki", "limit": 5},
                },
                ensure_ascii=False,
            ),
            db_name=self.db_path,
            history=[
                {"role": "user", "content": "kargoya hazir siparisleri getir"},
                {
                    "role": "assistant",
                    "content": self._answer_from(listed),
                    "context": listed_context,
                },
            ],
        )

        self.assertTrue(result.success, result.message)
        self.assertEqual(
            ((result.data or {}).get("tool_call") or {}).get("tool"),
            "orders.full_context",
        )
        tool_orders = [
            row.get("order_number")
            for row in ((result.data or {}).get("tool_result") or {}).get("orders") or []
        ]
        if not tool_orders:
            tool_orders = [
                row.get("order_number")
                or ((row.get("order") or {}).get("order_number"))
                for row in ((result.data or {}).get("tool_result") or {}).get("contexts") or []
            ]
        self.assertEqual(tool_orders, [order_numbers[1]])

    def test_ordinal_complaint_repairs_bad_lookup_to_exact_single_order_context(self) -> None:
        self._seed_ready_orders()
        listed = run_agent_turn(
            "kargoya hazir siparisleri getir",
            db_name=self.db_path,
        )
        self.assertTrue(listed.success, listed.message)
        listed_context = dict((listed.data or {}).get("conversation_context") or {})
        order_numbers = list(listed_context.get("order_numbers") or [])
        self.assertGreaterEqual(len(order_numbers), 2)

        result = run_agent_turn(
            "yine hepsini açtın; ben ikinci sıradaki tek siparişi istemiştim. sadece onu aç",
            planner_client=lambda _prompt: json.dumps(
                {
                    "mode": "tool",
                    "tool": "orders.lookup",
                    "arguments": {"query": "yine hepsini actin", "limit": 10},
                },
                ensure_ascii=False,
            ),
            db_name=self.db_path,
            history=[
                {"role": "user", "content": "kargoya hazir siparisleri getir"},
                {
                    "role": "assistant",
                    "content": self._answer_from(listed),
                    "context": listed_context,
                },
            ],
        )

        self.assertTrue(result.success, result.message)
        self.assertEqual(
            ((result.data or {}).get("tool_call") or {}).get("tool"),
            "orders.full_context",
        )
        contexts = list(((result.data or {}).get("tool_result") or {}).get("contexts") or [])
        self.assertEqual(len(contexts), 1)
        self.assertEqual(
            (contexts[0].get("order") or {}).get("order_number"),
            order_numbers[1],
        )

    def test_explicit_order_number_wins_over_stale_ordinal_context(self) -> None:
        self._seed_ready_orders()
        listed = run_agent_turn(
            "kargoya hazir siparisleri getir",
            db_name=self.db_path,
        )
        self.assertTrue(listed.success, listed.message)
        listed_context = dict((listed.data or {}).get("conversation_context") or {})
        order_numbers = list(listed_context.get("order_numbers") or [])
        self.assertGreaterEqual(len(order_numbers), 2)
        explicit_order = order_numbers[0]

        result = run_agent_turn(
            f"ikinci kayıt değil; açıkça yazdığım {explicit_order} siparişinin kalemlerini aç",
            planner_client=lambda _prompt: json.dumps(
                {
                    "mode": "tool",
                    "tool": "orders.full_context",
                    "arguments": {"query": "ikinci kayit", "limit": 5},
                },
                ensure_ascii=False,
            ),
            db_name=self.db_path,
            history=[
                {"role": "user", "content": "kargoya hazir siparisleri getir"},
                {
                    "role": "assistant",
                    "content": self._answer_from(listed),
                    "context": listed_context,
                },
            ],
        )

        self.assertTrue(result.success, result.message)
        contexts = list(((result.data or {}).get("tool_result") or {}).get("contexts") or [])
        self.assertEqual(len(contexts), 1)
        self.assertEqual(
            (contexts[0].get("order") or {}).get("order_number"),
            explicit_order,
        )

    def test_ready_first_three_read_only_request_beats_lookup_and_negative_label_noise(self) -> None:
        self._seed_ready_orders()
        result = run_agent_turn(
            "Kargoya hazır siparişlerden en yeni ilk 3 tanesini listele; yalnızca oku, etiket üretme.",
            planner_client=lambda _prompt: json.dumps(
                {
                    "mode": "tool",
                    "tool": "orders.lookup",
                    "arguments": {"query": "en yeni", "limit": 10},
                },
                ensure_ascii=False,
            ),
            db_name=self.db_path,
        )

        self.assertTrue(result.success, result.message)
        self.assertEqual(
            ((result.data or {}).get("tool_call") or {}).get("tool"),
            "orders.ready_to_ship",
        )
        self.assertEqual(
            ((result.data or {}).get("tool_call") or {}).get("arguments", {}).get("limit"),
            3,
        )
        orders = list(((result.data or {}).get("tool_result") or {}).get("orders") or [])
        self.assertEqual(len(orders), 2)
        self.assertTrue(all(row.get("status") == "ReadyToShip" for row in orders))
        self.assertFalse(bool((result.data or {}).get("approval_required")))

    def test_ready_first_three_accepts_turkish_inflected_number_word(self) -> None:
        self._seed_ready_orders()
        result = run_agent_turn(
            "Kargoya hazir siparislerden ilk ucunu yalnizca listele; etiket uretme.",
            planner_client=lambda _prompt: '{"mode":"tool","tool":"orders.lookup","arguments":{"query":"hazir","limit":10}}',
            db_name=self.db_path,
        )

        self.assertTrue(result.success, result.message)
        tool_call = (result.data or {}).get("tool_call") or {}
        self.assertEqual(tool_call.get("tool"), "orders.ready_to_ship")
        self.assertEqual((tool_call.get("arguments") or {}).get("limit"), 3)
        orders = list(((result.data or {}).get("tool_result") or {}).get("orders") or [])
        self.assertEqual(len(orders), 2)
        self.assertFalse(bool((result.data or {}).get("approval_required")))

    def test_ready_and_fulfillment_risk_count_request_keeps_both_signals(self) -> None:
        result = run_agent_turn(
            (
                "Kaç kargoya hazır paket var ve son 7 günde kaç fulfillment "
                "riskli sipariş var? Yalnız bu iki sayıyı ver."
            ),
            planner_client=lambda _prompt: (
                '{"mode":"tool","tool":"orders.fulfillment_risk_count",'
                '"arguments":{"period":"last_7_days"}}'
            ),
            db_name=self.db_path,
        )

        self.assertTrue(result.success, result.message)
        tools = [
            row.get("tool")
            for row in ((result.data or {}).get("tool_calls") or [])
        ]
        self.assertEqual(
            tools,
            ["orders.ready_to_ship_count", "orders.fulfillment_risk_count"],
        )
        answer = self._answer_from(result)
        self.assertIn("kargoya hazir", answer.casefold())
        self.assertIn("fulfillment risk", answer.casefold())

    def test_long_conversation_can_return_to_early_ordinal_order_selection(self) -> None:
        now_ms = int(datetime.now().timestamp() * 1000)
        early_orders = [f"ORD-EARLY-{index}" for index in range(1, 6)]
        for index, order_number in enumerate(early_orders):
            self._insert_order(
                order_number,
                now_ms - index,
                package_id=8_000 + index,
                first_name="Early",
                last_name=str(index),
                status="Delivered",
            )
        history = [
            {
                "role": "assistant",
                "content": "Son beş sipariş listelendi.",
                "context": {
                    "kind": "order_selection",
                    "source_tool": "orders.recent",
                    "order_numbers": early_orders,
                },
            },
            {
                "role": "assistant",
                "content": "Daha sonra etiket için farklı siparişler seçildi.",
                "context": {
                    "kind": "order_selection",
                    "source_tool": "labels.create_from_orders.prepare",
                    "order_numbers": ["ORD-NEWER-1", "ORD-NEWER-2", "ORD-NEWER-3"],
                },
            },
        ]
        result = run_agent_turn(
            "Konuşmanın başında son beş siparişten üçüncünün tüm detayını açmıştık. O siparişin ürün kalemlerini tekrar getir; son etiket listesindekilerle karıştırma.",
            planner_client=lambda _prompt: '{"mode":"tool","tool":"orders.full_context","arguments":{"query":"son sipariş","limit":5}}',
            db_name=self.db_path,
            history=history,
        )

        self.assertTrue(result.success, result.message)
        tool_call = (result.data or {}).get("tool_call") or {}
        self.assertEqual(tool_call.get("tool"), "orders.full_context")
        self.assertEqual((tool_call.get("arguments") or {}).get("query"), "ORD-EARLY-3")
        self.assertEqual((tool_call.get("arguments") or {}).get("limit"), 1)
        contexts = list(((result.data or {}).get("tool_result") or {}).get("contexts") or [])
        self.assertEqual(len(contexts), 1)
        self.assertEqual((contexts[0].get("order") or {}).get("order_number"), "ORD-EARLY-3")

        followup_history = [
            *history,
            {
                "role": "assistant",
                "content": self._answer_from(result),
                "context": {
                    "kind": "order_selection",
                    "source_tool": "orders.lookup",
                    "order_numbers": ["ORD-EARLY-3"],
                },
            },
        ]
        followup = run_agent_turn(
            "Evet doğru sipariş bu. Şimdi bunun ürün kalemlerini ve kargo bilgisini tam getir; yalnızca bu tek sipariş olsun.",
            planner_client=lambda _prompt: '{"mode":"tool","tool":"orders.full_context","arguments":{"query":"yanlis","limit":5}}',
            db_name=self.db_path,
            history=followup_history,
        )
        self.assertTrue(followup.success, followup.message)
        followup_call = (followup.data or {}).get("tool_call") or {}
        self.assertEqual(followup_call.get("tool"), "orders.full_context")
        self.assertEqual((followup_call.get("arguments") or {}).get("query"), "ORD-EARLY-3")
        self.assertEqual((followup_call.get("arguments") or {}).get("limit"), 1)
        followup_contexts = list(
            ((followup.data or {}).get("tool_result") or {}).get("contexts") or []
        )
        self.assertEqual(len(followup_contexts), 1)

    def test_ordinal_all_context_with_inflected_line_items_uses_full_context(self) -> None:
        history = [
            {
                "role": "assistant",
                "content": "Son beş sipariş listelendi.",
                "context": {
                    "kind": "order_selection",
                    "source_tool": "orders.recent",
                    "order_numbers": [
                        "ORD-RECENT-1",
                        "ORD-RECENT-2",
                        "ORD-RECENT-3",
                        "ORD-RECENT-4",
                        "ORD-RECENT-5",
                    ],
                },
            }
        ]

        result = run_agent_turn(
            (
                "Listedeki üçüncü siparişin tüm bağlamını; kalemlerini, "
                "kargosunu ve varsa iadesini göster."
            ),
            planner_client=lambda _prompt: (
                '{"mode":"tool","tool":"orders.lookup","arguments":'
                '{"query":"listedeki üçüncü","limit":1}}'
            ),
            db_name=self.db_path,
            history=history,
        )

        self.assertTrue(result.success, result.message)
        tool_call = (result.data or {}).get("tool_call") or {}
        self.assertEqual(tool_call.get("tool"), "orders.full_context")
        self.assertEqual(
            (tool_call.get("arguments") or {}).get("query"),
            "ORD-RECENT-3",
        )
        self.assertEqual((tool_call.get("arguments") or {}).get("limit"), 1)

    def test_ordinal_status_qualifier_walks_past_newer_insufficient_selection(self) -> None:
        now_ms = int(datetime.now().timestamp() * 1000)
        original_orders = [f"ORD-ORIGINAL-FIVE-{index}" for index in range(1, 6)]
        for index, order_number in enumerate(original_orders):
            self._insert_order(
                order_number,
                now_ms - index,
                package_id=8_100 + index,
                first_name="Original",
                last_name=str(index),
                status="Cancelled" if index == 2 else "Delivered",
            )

        history: list[dict[str, object]] = [
            {
                "role": "assistant",
                "content": "Son 5 sipariş listelendi.",
                "context": {
                    "kind": "order_selection",
                    "source_tool": "orders.recent",
                    "order_numbers": original_orders,
                    "listed_count": 5,
                    "total_count": 5,
                },
            },
            {
                "role": "assistant",
                "content": "Üçüncü sipariş tek başına açıldı.",
                "context": {
                    "kind": "order_selection",
                    "source_tool": "orders.full_context",
                    "order_numbers": [original_orders[2]],
                    "listed_count": 1,
                    "total_count": 1,
                },
            },
        ]
        for index in range(8):
            history.extend(
                [
                    {"role": "user", "content": f"serbest sohbet {index}"},
                    {"role": "assistant", "content": f"sohbet cevabı {index}"},
                ]
            )
        history.extend(
            [
                {
                    "role": "assistant",
                    "content": "Daha yeni iki sipariş listelendi.",
                    "context": {
                        "kind": "order_selection",
                        "source_tool": "orders.recent",
                        "order_numbers": original_orders[:2],
                        "listed_count": 2,
                        "total_count": 2,
                    },
                },
                {
                    "role": "assistant",
                    "content": "Yeni veri çağırmadan telefon kuralını düzelttim.",
                },
            ]
        )

        result = run_agent_turn(
            (
                "şimdi sipariş işine dön: uzun sohbet sapması başlamadan önce "
                "son 5'lik listede tek başına açtığımız üçüncü, iptal edilmiş "
                "siparişin tam bağlamını tekrar getir"
            ),
            planner_client=lambda _prompt: (
                '{"mode":"tool","tool":"orders.cancellations",'
                '"arguments":{"period":"today"}}'
            ),
            db_name=self.db_path,
            history=history,
        )

        self.assertTrue(result.success, result.message)
        tool_call = (result.data or {}).get("tool_call") or {}
        self.assertEqual(tool_call.get("tool"), "orders.full_context")
        self.assertEqual(
            (tool_call.get("arguments") or {}).get("query"),
            original_orders[2],
        )
        self.assertEqual((tool_call.get("arguments") or {}).get("limit"), 1)
        contexts = list(((result.data or {}).get("tool_result") or {}).get("contexts") or [])
        self.assertEqual(len(contexts), 1)
        self.assertEqual(
            (contexts[0].get("order") or {}).get("order_number"),
            original_orders[2],
        )

    def test_selected_cancelled_order_reason_boundary_uses_state_without_tool(self) -> None:
        order_number = "ORD-CANCELLED-REASON-BOUNDARY"
        history = [
            {
                "role": "assistant",
                "content": f"1. {order_number} | durum: Cancelled | iade: 0",
                "context": {
                    "kind": "order_selection",
                    "source_tool": "orders.full_context",
                    "order_numbers": [order_number],
                    "listed_count": 1,
                    "total_count": 1,
                },
            }
        ]

        result = run_agent_turn(
            (
                "bu sipariş neden iptal olmuş, eldeki veriden kesin sebep "
                "söyleyebiliyor musun? bilmiyorsan uydurma ve yeni veri çağırma"
            ),
            planner_client=lambda _prompt: (
                '{"mode":"tool","tool":"orders.cancellations",'
                '"arguments":{"period":"today"}}'
            ),
            db_name=self.db_path,
            history=history,
        )

        self.assertTrue(result.success, result.message)
        self.assertEqual((result.data or {}).get("mode"), "answer")
        self.assertEqual((result.data or {}).get("tool_calls"), [])
        answer = self._answer_from(result)
        self.assertIn(order_number, answer)
        self.assertIn("kesin sebep söyleyemem", answer)
        self.assertIn("yeni veri veya tool da çağırmadım", answer)

    def test_counted_reference_prepares_exact_latest_three_orders(self) -> None:
        now_ms = self._seed_ready_orders()
        self._insert_order(
            "ORD-THIRD-READY",
            now_ms + 2_000,
            package_id=903,
            first_name="Third",
            last_name="Ready",
            status="ReadyToShip",
        )
        listed = run_agent_turn(
            "Kargoya hazır siparişlerden en yeni ilk 3 tanesini listele",
            db_name=self.db_path,
        )
        self.assertTrue(listed.success, listed.message)
        context = dict((listed.data or {}).get("conversation_context") or {})
        selected = list(context.get("order_numbers") or [])
        self.assertEqual(len(selected), 3)

        prepared = run_agent_turn(
            "Eski iadeyi karıştırma. Son listelenen üç ReadyToShip siparişin etiket çıktısını yalnızca onaya hazırla; dosya üretme ve başka sipariş ekleme.",
            planner_client=lambda _prompt: (
                '{"mode":"tool","tool":"operations.manual_action.route",'
                '"arguments":{"query":"siparis verisi cek","mode":"write"}}'
            ),
            db_name=self.db_path,
            history=[
                {
                    "role": "assistant",
                    "content": "Daha eski bir iade seçimi.",
                    "context": {
                        "kind": "return_selection",
                        "claim_ids": ["CLAIM-OLD-CONTEXT"],
                    },
                },
                {"role": "user", "content": "Kargoya hazır ilk 3 siparişi listele"},
                {
                    "role": "assistant",
                    "content": self._answer_from(listed),
                    "context": context,
                },
                {
                    "role": "assistant",
                    "content": "Yanlışlıkla 20 sipariş prepare edildi.",
                    "context": {
                        "kind": "order_selection",
                        "source_tool": "labels.create_from_orders.prepare",
                        "order_numbers": [
                            *selected,
                            *[f"ORD-WRONG-{index}" for index in range(17)],
                        ],
                    },
                },
            ],
        )

        self._assert_prepare_answer_is_safe(prepared, expected_terms=tuple(selected))
        action = self._action_from(prepared)
        prepared_orders = [
            row.get("order_number")
            for row in list(action.get("rows") or [])
        ]
        self.assertEqual(len(prepared_orders), 3)
        self.assertEqual(set(prepared_orders), set(selected))

    def test_bad_lookup_complaint_followup_gets_contextual_repair_answer(self) -> None:
        history = [
            {"role": "user", "content": "siparisler tarafinda hangi islemleri yapabiliyoruz"},
            {
                "role": "assistant",
                "content": "Lokal kayitta 'siparisler tarafinda hangi islemleri yapabiliyoruz' icin siparis bulunamadi.",
            },
        ]

        result = run_agent_turn(
            "ne demek bu",
            planner_client=lambda _prompt: '{"mode":"answer","response":"Buradayim. Ne yapalim?"}',
            chat_client=lambda _prompt: "Buradayim. Ne yapalim?",
            db_name=self.db_path,
            history=history,
        )

        self.assertTrue(result.success, result.message)
        self.assertEqual((result.data or {}).get("mode"), "answer")
        self.assertIsNone((result.data or {}).get("tool_call"))
        answer = self._answer_from(result)
        self.assertIn("hataliydi", answer)
        self.assertIn("siparis aramasi", answer)

    def test_open_ended_chat_advice_does_not_route_to_order_count(self) -> None:
        def bad_planner(_prompt: str) -> str:
            return '{"mode":"tool","tool":"orders.count","arguments":{"period":"today"}}'

        result = run_agent_turn(
            "valla bilmiyorum ki sence?",
            planner_client=bad_planner,
            chat_client=lambda _prompt: "Bence once bugunku acil operasyonlara bakip sonra gelistirme fikirlerini konusabiliriz.",
            db_name=self.db_path,
            history=[
                {"role": "user", "content": "naber dayi napiyon"},
                {"role": "assistant", "content": "Iyiyim, buradayim."},
            ],
        )

        self.assertTrue(result.success, result.message)
        self.assertEqual((result.data or {}).get("mode"), "answer")
        self.assertIsNone((result.data or {}).get("tool_call"))
        answer = self._answer_from(result)
        self.assertIn("Bence", answer)
        self.assertNotIn("siparis var", answer)

    def test_free_chat_direct_llm_prompt_is_not_pinned_to_fallback_template(self) -> None:
        class ChatClient:
            def __init__(self) -> None:
                self.prompts: list[str] = []

            def __call__(self, prompt: str) -> str:
                self.prompts.append(prompt)
                if prompt.lstrip().startswith("Sen OrderScout masaustu"):
                    return '{"mode":"tool","tool":"orders.count","arguments":{"period":"today"}}'
                return "Bence once iki dakika muhabbeti netlestirelim; istersen sonra siparis verisine de birlikte bakariz."

        client = ChatClient()

        result = run_agent_turn(
            "valla bilmiyorum ki sence?",
            planner_client=client,
            chat_client=client,
            db_name=self.db_path,
            history=[
                {"role": "user", "content": "naber dayi napiyon"},
                {"role": "assistant", "content": "Iyiyim, buradayim."},
            ],
        )

        self.assertTrue(result.success, result.message)
        self.assertEqual((result.data or {}).get("mode"), "answer")
        direct_prompts = [prompt for prompt in client.prompts if prompt.lstrip().startswith("Sen OrderScout AI'sin")]
        self.assertEqual(len(direct_prompts), 1)
        self.assertNotIn("Dogrulanmis cevap iskeleti", direct_prompts[0])
        self.assertNotIn("Bence once neye karar vermek", direct_prompts[0])
        answer = self._answer_from(result)
        self.assertIn("muhabbeti", answer)
        self.assertNotIn("siparis var", answer.casefold())

    def test_open_conversation_filters_passive_llm_answer(self) -> None:
        result = run_agent_turn(
            "valla bilmiyorum ki sence?",
            planner_client=lambda _prompt: '{"mode":"answer","response":"Ben de bilemiyorum, sen ne yapmak istersin?"}',
            chat_client=lambda _prompt: "Ben de bilemiyorum, sen ne yapmak istersin?",
            db_name=self.db_path,
            history=[
                {"role": "user", "content": "naber dayi napiyon"},
                {"role": "assistant", "content": "Iyiyim, sen nasilsin?"},
            ],
        )

        self.assertTrue(result.success, result.message)
        self.assertEqual((result.data or {}).get("mode"), "answer")
        self.assertIsNone((result.data or {}).get("tool_call"))
        answer = self._answer_from(result)
        self.assertIn("Bence", answer)
        self.assertIn("veriye", answer)
        self.assertNotIn("bilemiyorum", answer.casefold())
        self.assertNotIn("siparis var", answer.casefold())

    def test_open_conversation_filters_bilemem_sorabilirsin_answer(self) -> None:
        result = run_agent_turn(
            "valla bilmiyorum ki sence?",
            planner_client=lambda _prompt: '{"mode":"answer","response":"Ben senin adina bilemem ki dayi. Ama istersen bir seyler sorabilirsin."}',
            chat_client=lambda _prompt: "Ben senin adina bilemem ki dayi. Ama istersen bir seyler sorabilirsin.",
            db_name=self.db_path,
            history=[
                {"role": "user", "content": "naber dayi napiyon"},
                {"role": "assistant", "content": "Iyidir dayi, takiliyorum buralarda."},
            ],
        )

        self.assertTrue(result.success, result.message)
        self.assertEqual((result.data or {}).get("mode"), "answer")
        self.assertIsNone((result.data or {}).get("tool_call"))
        answer = self._answer_from(result)
        self.assertIn("Bence", answer)
        self.assertIn("veriye", answer)
        self.assertNotIn("bilemem", answer.casefold())
        self.assertNotIn("sorabilirsin", answer.casefold())

    def test_open_conversation_filters_off_domain_llm_answer(self) -> None:
        result = run_agent_turn(
            "valla bilmiyorum ki sence?",
            planner_client=lambda _prompt: '{"mode":"answer","response":"Valla ne bileyim, aciksan yemek siparisi bakalim ya da rota cizelim. Canin ne istiyor?"}',
            chat_client=lambda _prompt: "Valla ne bileyim, aciksan yemek siparisi bakalim ya da rota cizelim. Canin ne istiyor?",
            db_name=self.db_path,
            history=[
                {"role": "user", "content": "naber dayi napiyon"},
                {"role": "assistant", "content": "Iyidir yegen, ugrasip duruyoruz iste."},
            ],
        )

        self.assertTrue(result.success, result.message)
        self.assertEqual((result.data or {}).get("mode"), "answer")
        self.assertIsNone((result.data or {}).get("tool_call"))
        answer = self._answer_from(result)
        self.assertIn("Bence", answer)
        self.assertIn("veriye", answer)
        self.assertNotIn("yemek siparisi", answer.casefold())
        self.assertNotIn("rota", answer.casefold())
        self.assertNotIn("canin ne istiyor", answer.casefold())

    def test_open_conversation_filters_passive_detail_request_answer(self) -> None:
        result = run_agent_turn(
            "valla bilmiyorum ki sence?",
            planner_client=lambda _prompt: '{"mode":"answer","response":"Maalesef, su an icin kesin bir bilgi veremiyorum. Daha fazla detay verirseniz yardimci olabilirim."}',
            chat_client=lambda _prompt: "Maalesef, su an icin kesin bir bilgi veremiyorum. Daha fazla detay verirseniz yardimci olabilirim.",
            db_name=self.db_path,
            history=[
                {"role": "user", "content": "naber dayi napiyon"},
                {"role": "assistant", "content": "Iyiyim, buradayim."},
            ],
        )

        self.assertTrue(result.success, result.message)
        answer = self._answer_from(result)
        self.assertIn("Bence", answer)
        self.assertNotIn("maalesef", answer.casefold())
        self.assertNotIn("detay verirseniz", answer.casefold())

    def test_open_conversation_filters_json_wrapped_llm_answer(self) -> None:
        result = run_agent_turn(
            "e ticaret nasil bisey sence iyi mi",
            planner_client=lambda _prompt: '{"mode":"answer","response":"{\\"reply\\":\\"E-ticaret icin daha fazla bilgi verirseniz yardimci olurum.\\"}"}',
            chat_client=lambda _prompt: '{"reply":"E-ticaret icin daha fazla bilgi verirseniz yardimci olurum."}',
            db_name=self.db_path,
        )

        self.assertTrue(result.success, result.message)
        answer = self._answer_from(result)
        self.assertIn("E-ticaret iyi bir alan", answer)
        self.assertNotIn('"reply"', answer)
        self.assertFalse(answer.lstrip().startswith("{"))

    def test_open_conversation_filters_malformed_localized_json_answer(self) -> None:
        responses = iter(
            (
                '{"mode":"answer","response":"E-ticaret hakkinda dusunuyorum."}',
                '{"metin": "Merhaba. E-ticaret kolaylik saglar ama sorunlar da vardir',
                '{"metin": "E-ticaret iyi bir alan ama dikkat gerekir',
            )
        )

        def local_client(_prompt: str) -> str:
            return next(responses)

        result = run_agent_turn(
            "e ticaret nasil bisey sence iyi mi",
            planner_client=local_client,
            chat_client=local_client,
            db_name=self.db_path,
        )

        self.assertTrue(result.success, result.message)
        answer = self._answer_from(result)
        self.assertFalse(answer.lstrip().startswith(("{", "[")))
        self.assertNotIn('"metin"', answer)
        self.assertIn("E-ticaret iyi bir alan", answer)
        self.assertEqual(
            (result.data or {}).get("answer_origin"),
            "deterministic_instruction_fallback",
        )

    def test_direct_answer_filters_prompt_instruction_leakage(self) -> None:
        leaked_answer = (
            "Sen OrderScout AI'sin. Cevabi JSON yazma; no json. "
            "Draft 1 too formal, draft 2 better, matched tone, within capabilities. "
            "Siparisleri listelerim."
        )
        result = run_agent_turn(
            "siparislerimle ilgili neler yapabiliriz",
            planner_client=lambda _prompt: '{"mode":"tool","tool":"orders.count","arguments":{"period":"today"}}',
            chat_client=lambda _prompt: leaked_answer,
            db_name=self.db_path,
        )

        self.assertTrue(result.success, result.message)
        self.assertEqual((result.data or {}).get("mode"), "answer")
        answer = self._answer_from(result)
        self.assertIn("etiket", answer.casefold())
        self.assertIn("iade", answer.casefold())
        self.assertNotIn("draft", answer.casefold())
        self.assertNotIn("no json", answer.casefold())

    def test_tool_answer_filters_prompt_instruction_leakage(self) -> None:
        now_ms = int(datetime.now(TZ_GMT3).timestamp() * 1000)
        self._insert_waiting_return(
            "CLAIM-LEAKAGE-LOOKUP",
            "ORD-LEAKAGE-LOOKUP",
            now_ms,
            first_name="Mahmut",
            last_name="Sevket",
        )
        leaked_answer = (
            "* Role: OrderScout AI.\n"
            "* Constraints: no raw JSON output, no hallucinated numbers.\n"
            "* Tool Result: claim_id CLAIM-LEAKAGE-LOOKUP, order ORD-LEAKAGE-LOOKUP.\n"
            "* Draft 1: too formal.\n"
            "Lokal kayda gore CLAIM-LEAKAGE-LOOKUP iadesi WaitingInAction."
        )

        result = run_agent_turn(
            "CLAIM-LEAKAGE-LOOKUP iadesini goster",
            planner_client=lambda _prompt: '{"mode":"tool","tool":"returns.lookup","arguments":{"query":"CLAIM-LEAKAGE-LOOKUP"}}',
            chat_client=lambda _prompt: leaked_answer,
            db_name=self.db_path,
        )

        self.assertTrue(result.success, result.message)
        self.assertEqual(((result.data or {}).get("tool_call") or {}).get("tool"), "returns.lookup")
        answer = self._answer_from(result)
        self.assertIn("CLAIM-LEAKAGE-LOOKUP", answer)
        self.assertIn("ORD-LEAKAGE-LOOKUP", answer)
        for forbidden in ("Role:", "Constraints:", "Tool Result:", "Draft 1", "no raw JSON", "hallucinated"):
            self.assertNotIn(forbidden.casefold(), answer.casefold())

    def test_tool_answer_filters_turkish_prompt_instruction_leakage(self) -> None:
        now_ms = int(datetime.now(TZ_GMT3).timestamp() * 1000)
        claim_id = "75a7c1c0-73a2-49f7-bb8a-366dee971405"
        self._insert_waiting_return(
            claim_id,
            "11268067005",
            now_ms,
            first_name="Samir",
            last_name="Abdurrahman",
        )
        leaked_answer = (
            "Bu senin icin yapmam gerekenler:\n"
            "1. Veri lokal DB read-only tool kullan.\n"
            "2. Uydurma rakam ekleme.\n"
            "3. Gizli anahtar degerlerini asla yazma.\n"
            "4. Tool sonucunu ham JSON gibi basma.\n"
            "5. Kullanicinin sorusuna cevap olacak sekilde ozetle."
        )

        result = run_agent_turn(
            "onay bekleyen iadeleri getir",
            planner_client=lambda _prompt: '{"mode":"tool","tool":"returns.query","arguments":{"period":"all","filters":{"status":"WaitingInAction"},"sort":"claim_date_desc","limit":20,"mode":"list"}}',
            chat_client=lambda _prompt: leaked_answer,
            db_name=self.db_path,
        )

        self.assertTrue(result.success, result.message)
        self.assertEqual(((result.data or {}).get("tool_call") or {}).get("tool"), "returns.query")
        answer = self._answer_from(result)
        self.assertIn(claim_id, answer)
        self.assertIn("Samir", answer)
        for forbidden in ("Bu senin icin", "read-only tool", "Uydurma rakam", "Gizli anahtar", "ham JSON"):
            self.assertNotIn(forbidden.casefold(), answer.casefold())

    def test_read_only_return_list_negation_never_routes_to_approval(self) -> None:
        now_ms = int(datetime.now(TZ_GMT3).timestamp() * 1000)
        self._insert_waiting_return(
            "75a7c1c0-73a2-49f7-bb8a-366dee971406",
            "11268067006",
            now_ms,
            first_name="Read",
            last_name="Only",
        )

        result = run_agent_turn(
            "Hayır, onay istemedim; tam tersine onaylama dedim. Son 30 gündeki bekleyen iadeleri yalnızca listele. Bu salt okunur bir istek, hiçbir kayıt değiştirme.",
            planner_client=lambda _prompt: (
                '{"mode":"tool","tool":"returns.approve.prepare",'
                '"arguments":{"query":"bekleyen iadeler","limit":5}}'
            ),
            history=[
                {
                    "role": "assistant",
                    "content": "Kargoya hazir 7 siparis var.",
                    "context": {
                        "kind": "orders_ready_to_ship_scope",
                        "source_tool": "orders.ready_to_ship_count",
                        "period": "all",
                        "total_count": 7,
                    },
                }
            ],
            db_name=self.db_path,
        )

        self.assertTrue(result.success, result.message)
        self.assertEqual(
            ((result.data or {}).get("tool_call") or {}).get("tool"),
            "returns.query",
        )
        self.assertFalse(bool((result.data or {}).get("approval_required")))

    def test_read_only_return_count_with_inflected_negation_never_prepares_approval(self) -> None:
        now_ms = int(datetime.now(TZ_GMT3).timestamp() * 1000)
        with Session(get_engine(self.db_path)) as session:
            session.add(
                ApiAccount(
                    pk=17,
                    account_id="seller-return-count",
                    comp_name="Return Count Seller",
                    platform="trendyol",
                    is_active=True,
                )
            )
            session.commit()
        self._insert_waiting_return(
            "75a7c1c0-73a2-49f7-bb8a-366dee971499",
            "11268067999",
            now_ms,
            first_name="Count",
            last_name="Only",
            api_account_id=17,
        )

        result = run_agent_turn(
            (
                "Orders hazır siparişlerini değil, yalnız Returns/iade domainini "
                "oku: production veride aksiyon/onay bekleyen kaç iade talebi "
                "var? Hesap/ortam kapsamını state'te koru; hiçbir iadeyi "
                "onaylama veya değiştirme; prepare veya execute çağırma."
            ),
            planner_client=lambda _prompt: (
                '{"mode":"tool","tool":"returns.approve.prepare",'
                '"arguments":{"query":"bekleyen iadeler","limit":5}}'
            ),
            history=[
                {
                    "role": "assistant",
                    "content": "Uc siparis icin etiket hazirligi onay bekliyor.",
                    "context": {
                        "kind": "order_selection",
                        "source_tool": "labels.create_from_orders.prepare",
                        "order_numbers": ["ORD-OLD-LABEL-1"],
                    },
                },
                {
                    "role": "user",
                    "content": "vazgeçtim, bu etiket hazırlığını iptal et",
                },
                {
                    "role": "assistant",
                    "content": "Hazırlık iptal edildi; hiçbir mutasyon yapılmadı.",
                },
                {
                    "role": "assistant",
                    "content": "Kargoya hazir 7 siparis var.",
                    "context": {
                        "kind": "orders_ready_to_ship_scope",
                        "source_tool": "orders.ready_to_ship_count",
                        "period": "all",
                        "total_count": 7,
                    },
                },
            ],
            db_name=self.db_path,
        )

        self.assertTrue(result.success, result.message)
        self.assertEqual(
            ((result.data or {}).get("tool_call") or {}).get("tool"),
            "returns.waiting_action_count",
        )
        self.assertFalse(bool((result.data or {}).get("approval_required")))
        self.assertNotIn(
            "account_by_claim",
            dict((result.data or {}).get("tool_result") or {}),
        )
        context = dict((result.data or {}).get("conversation_context") or {})
        self.assertEqual(context.get("kind"), "returns_waiting_action_scope")
        self.assertNotIn("account_by_claim", context)
        self.assertEqual(context.get("environment_by_account"), {"17": "production"})
        self.assertEqual(
            [row.get("store_name") for row in context.get("account_scopes") or []],
            ["Return Count Seller"],
        )

    def test_selected_return_prepare_cannot_revive_cancelled_label_context(self) -> None:
        now_ms = int(datetime.now(TZ_GMT3).timestamp() * 1000)
        self._insert_waiting_return(
            "CLAIM-CANCELLED-PREPARE",
            "ORD-CANCELLED-PREPARE",
            now_ms,
            first_name="Cancelled",
            last_name="Prepare",
        )
        with Session(get_engine(self.db_path)) as session:
            claim = session.exec(
                select(ReturnClaim).where(
                    ReturnClaim.claimId == "CLAIM-CANCELLED-PREPARE"
                )
            ).one()
            claim.claimItemStatus = "Cancelled"
            claim.waitingClaimLineItemIds = []
            session.add(claim)
            session.commit()

        result = run_agent_turn(
            (
                "Bu tek seçili iade kaydını onay için yalnız prepare aşamasına "
                "kadar hazırla; execute etme. Mevcut durumu onaya uygun değilse "
                "action kartı üretme, hiçbir kaydı değiştirme ve reddetme "
                "nedenini açıkça söyle."
            ),
            planner_client=lambda _prompt: (
                '{"mode":"tool","tool":"labels.create_from_orders.prepare",'
                '"arguments":{"query":"seçili kayıt","limit":20}}'
            ),
            history=[
                {
                    "role": "assistant",
                    "content": "Eski etiket hazırlığı onay bekliyor.",
                    "context": {
                        "kind": "order_selection",
                        "source_tool": "labels.create_from_orders.prepare",
                        "order_numbers": ["ORD-OLD-LABEL-PREPARE"],
                    },
                },
                {
                    "role": "user",
                    "content": "vazgeçtim, bu etiket hazırlığını iptal et",
                },
                {
                    "role": "assistant",
                    "content": "Hazırlık iptal edildi; hiçbir mutasyon yapılmadı.",
                },
                {
                    "role": "assistant",
                    "content": "Tek Cancelled iade seçildi.",
                    "context": {
                        "kind": "return_selection",
                        "source_tool": "returns.lookup",
                        "claim_ids": ["CLAIM-CANCELLED-PREPARE"],
                        "order_numbers": ["ORD-CANCELLED-PREPARE"],
                        "listed_count": 1,
                        "total_count": 1,
                    },
                },
            ],
            db_name=self.db_path,
        )

        self.assertTrue(result.success, result.message)
        self.assertEqual(
            ((result.data or {}).get("tool_call") or {}).get("tool"),
            "returns.approve.prepare",
        )
        self.assertFalse(
            bool((result.data or {}).get("approval_required")),
            result.data,
        )
        self.assertFalse(bool(((result.data or {}).get("tool_result") or {}).get("valid")))
        self.assertIn("WaitingInAction", self._answer_from(result))
        self.assertEqual((result.data or {}).get("action_proposal"), None)

    def test_tight_local_ordinal_return_prepare_beats_generic_multi_tool(self) -> None:
        now_ms = int(datetime.now(TZ_GMT3).timestamp() * 1000)
        for index in range(3):
            self._insert_waiting_return(
                f"CLAIM-ORDINAL-PREPARE-{index}",
                f"ORD-ORDINAL-PREPARE-{index}",
                now_ms + index,
                first_name="Ordinal",
                last_name=str(index),
            )

        class TightLocalClient:
            resource_plan = {"profile": "local_model", "num_ctx": 1024}

            def __call__(self, _prompt: str) -> str:
                raise AssertionError("Specific return prepare must be deterministic.")

        result = run_agent_turn(
            (
                "Bu listedeki ilk iade claimini yalnız onaya hazırla; execute "
                "etme, API isteği veya DB yazımı yapma."
            ),
            planner_client=TightLocalClient(),
            db_name=self.db_path,
            history=[
                {
                    "role": "assistant",
                    "content": "Üç WaitingInAction iade listelendi.",
                    "context": {
                        "kind": "return_selection",
                        "source_tool": "returns.query",
                        "claim_ids": [
                            "CLAIM-ORDINAL-PREPARE-0",
                            "CLAIM-ORDINAL-PREPARE-1",
                            "CLAIM-ORDINAL-PREPARE-2",
                        ],
                        "listed_count": 3,
                        "total_count": 3,
                    },
                }
            ],
        )

        self.assertTrue(result.success, result.message)
        self.assertEqual(
            [row.get("tool") for row in (result.data or {}).get("tool_calls") or []],
            ["returns.approve.prepare"],
        )
        self.assertNotIn(
            "listing.attention_count",
            [row.get("tool") for row in (result.data or {}).get("tool_calls") or []],
        )
        action = self._action_from(result)
        payload = ((action.get("arguments") or {}).get("payload") or {})
        self.assertEqual(payload.get("order"), "CLAIM-ORDINAL-PREPARE-0")
        self.assertFalse(bool(action.get("api_call_performed")))
        self.assertFalse(bool(action.get("db_write_performed")))

    def test_general_ecommerce_opinion_stays_free_chat_not_template(self) -> None:
        prompts: list[str] = []

        def chat_client(prompt: str) -> str:
            prompts.append(prompt)
            return "E-ticaret iyi bir alan; operasyon, urun ve nakit akisini disiplinli yonetince guclu olabiliyor."

        result = run_agent_turn(
            "e ticaret nasil bisey sence iyi mi",
            planner_client=lambda _prompt: '{"mode":"tool","tool":"orders.count","arguments":{"period":"today"}}',
            chat_client=chat_client,
            db_name=self.db_path,
        )

        self.assertTrue(result.success, result.message)
        self.assertEqual((result.data or {}).get("mode"), "answer")
        self.assertIsNone((result.data or {}).get("tool_call"))
        answer = self._answer_from(result)
        self.assertIn("E-ticaret iyi", answer)
        self.assertNotEqual(answer, "Iyiyim, buradayim. Siparis, iade, kargo, etiket veya veri cekme tarafinda neye bakmami istersin?")
        self.assertTrue(prompts)
        self.assertIn("once net gorusunu ve 2-3 somut gerekceyi ver", prompts[-1])
        self.assertIn("soruyu ona geri atma", prompts[-1])

    def test_read_tool_keeps_llm_synthesis_when_identifier_uses_typographic_hyphen(self) -> None:
        now_ms = int(datetime.now(TZ_GMT3).timestamp() * 1000)
        self._insert_order(
            "ORD-READ-LLM-1",
            now_ms,
            package_id=491,
            status="ReadyToShip",
            cargo_tracking_number="TRK-READ-LLM-1",
            cargo_provider_name="Trendyol Express",
            first_name="Read",
            last_name="One",
        )

        llm_answer = "Son siparisin **ORD\u2011READ\u2011LLM\u20111**; durum kargoya hazir ve tutar 100 TL."
        expected_answer = "Son siparisin **ORD-READ-LLM-1**; durum kargoya hazir ve tutar 100 TL."
        result = run_agent_turn(
            "son gelen 1 siparisi soyle",
            planner_client=lambda _prompt: '{"mode":"tool","tool":"orders.recent","arguments":{"limit":1}}',
            chat_client=lambda _prompt: llm_answer,
            db_name=self.db_path,
        )

        self.assertTrue(result.success, result.message)
        self.assertEqual(((result.data or {}).get("tool_call") or {}).get("tool"), "orders.recent")
        self.assertEqual(self._answer_from(result), expected_answer)

    def test_latest_order_prompt_repairs_lookup_to_recent_with_limit(self) -> None:
        now_ms = int(datetime.now(TZ_GMT3).timestamp() * 1000)
        self._insert_order(
            "ORD-LATEST-ONE",
            now_ms + 1_000,
            package_id=492,
            status="Shipped",
            first_name="Latest",
            last_name="One",
        )
        self._insert_order(
            "ORD-LATEST-OLD",
            now_ms,
            package_id=493,
            status="Delivered",
            first_name="Latest",
            last_name="Old",
        )

        result = run_agent_turn(
            "son gelen 1 siparisi soyle",
            planner_client=lambda _prompt: '{"mode":"tool","tool":"orders.lookup","arguments":{"query":"son gelen soyle","limit":10}}',
            chat_client=lambda _prompt: "Bu cevap kullanilmamali.",
            db_name=self.db_path,
        )

        self.assertTrue(result.success, result.message)
        tool_call = (result.data or {}).get("tool_call") or {}
        self.assertEqual(tool_call.get("tool"), "orders.recent")
        self.assertEqual((tool_call.get("arguments") or {}).get("limit"), 1)
        answer = self._answer_from(result)
        self.assertIn("ORD-LATEST-ONE", answer)
        self.assertNotIn("ORD-LATEST-OLD", answer)

    def test_latest_orders_ignore_explicitly_negated_old_ready_queue(self) -> None:
        now_ms = int(datetime.now(TZ_GMT3).timestamp() * 1000)
        for index in range(1, 5):
            self._insert_order(
                f"ORD-LATEST-NEGATED-{index}",
                now_ms + index,
                package_id=9_100 + index,
                status="Picking",
                first_name="Latest",
                last_name=str(index),
            )
        history = [
            {
                "role": "assistant",
                "content": "199 kargoya hazir siparis var.",
                "context": {
                    "kind": "orders_ready_to_ship_scope",
                    "source_tool": "orders.ready_to_ship_count",
                    "period": "all",
                    "total_count": 199,
                },
            }
        ]

        result = run_agent_turn(
            (
                "Simdi sadece son gelen 3 siparisi getir; eski 199 kayitlik "
                "hazir kuyrugu listeleme."
            ),
            planner_client=lambda _prompt: (
                '{"mode":"tool","tool":"orders.lookup","arguments":'
                '{"query":"son gelen eski kuyruk","limit":10}}'
            ),
            db_name=self.db_path,
            history=history,
        )

        self.assertTrue(result.success, result.message)
        tool_call = (result.data or {}).get("tool_call") or {}
        self.assertEqual(tool_call.get("tool"), "orders.recent")
        self.assertEqual((tool_call.get("arguments") or {}).get("limit"), 3)
        answer = self._answer_from(result)
        self.assertIn("ORD-LATEST-NEGATED-4", answer)
        self.assertIn("ORD-LATEST-NEGATED-3", answer)
        self.assertIn("ORD-LATEST-NEGATED-2", answer)
        self.assertNotIn("ORD-LATEST-NEGATED-1", answer)

    def test_stage_ready_total_and_latest_list_does_not_treat_account_exclusion_as_ready_negation(self) -> None:
        class TightLocalClient:
            resource_plan = {"profile": "local_model", "num_ctx": 1024}

            def __call__(self, _prompt: str) -> str:
                raise AssertionError("Tight local path must use deterministic intent.")

        with Session(get_engine(self.db_path)) as session:
            session.add(
                ApiAccount(
                    pk=1,
                    account_id="seller-production",
                    comp_name="Production Store",
                    platform="TRENDYOL",
                    extra_config={"environment": "production"},
                )
            )
            session.add(
                ApiAccount(
                    pk=2,
                    account_id="2738",
                    comp_name="Stage Store",
                    platform="TRENDYOL",
                    extra_config={"environment": "stage"},
                )
            )
            session.commit()

        now_ms = int(datetime.now(TZ_GMT3).timestamp() * 1000)
        for index in range(7):
            self._insert_order(
                f"ORD-STAGE-READY-COMBINED-{index}",
                now_ms + index,
                package_id=9_200 + index,
                status="ReadyToShip",
                first_name="Stage",
                last_name=f"Ready{index}",
                api_account_id=2,
            )
        self._insert_order(
            "ORD-STAGE-SHIPPED-EXCLUDED",
            now_ms + 100,
            package_id=9_300,
            status="Shipped",
            first_name="Stage",
            last_name="Shipped",
            api_account_id=2,
        )
        self._insert_order(
            "ORD-PRODUCTION-READY-EXCLUDED",
            now_ms + 200,
            package_id=9_301,
            status="ReadyToShip",
            first_name="Production",
            last_name="Ready",
            api_account_id=1,
        )

        result = run_agent_turn(
            (
                "Yalniz Stage magazasindaki kargoya hazir siparislerin toplam "
                "sayisini soyle ve en guncel ilk 5 tanesini siparis numarasi, "
                "durum ve tutarla listele; Production kayitlarini karistirma."
            ),
            planner_client=TightLocalClient(),
            db_name=self.db_path,
        )

        self.assertTrue(result.success, result.message)
        tool_call = (result.data or {}).get("tool_call") or {}
        self.assertEqual(tool_call.get("tool"), "orders.ready_to_ship")
        args = tool_call.get("arguments") or {}
        self.assertEqual(args.get("limit"), 5)
        self.assertEqual(args.get("api_account_ids"), [2])
        data = (result.data or {}).get("tool_result") or {}
        self.assertEqual(data.get("ready_to_ship"), 7)
        self.assertEqual(data.get("listed_count"), 5)
        self.assertTrue(data.get("truncated"))
        rows = data.get("orders") or []
        self.assertTrue(all(row.get("status") == "ReadyToShip" for row in rows))
        self.assertTrue(all(row.get("api_account_id") == 2 for row in rows))
        answer = self._answer_from(result)
        self.assertIn("toplam 7 siparis", answer)
        self.assertNotIn("ORD-STAGE-SHIPPED-EXCLUDED", answer)
        self.assertNotIn("ORD-PRODUCTION-READY-EXCLUDED", answer)

        cargo_projection = run_agent_turn(
            (
                "Yalniz Stage magazasindaki kargoya hazir en guncel 3 "
                "siparisi siparis numarasi, durum ve kargo firmasiyla listele; "
                "Production kayitlarini karistirma."
            ),
            planner_client=TightLocalClient(),
            db_name=self.db_path,
        )

        self.assertTrue(cargo_projection.success, cargo_projection.message)
        cargo_calls = list((cargo_projection.data or {}).get("tool_calls") or [])
        self.assertEqual(
            [row.get("tool") for row in cargo_calls],
            ["orders.ready_to_ship"],
        )
        cargo_args = cargo_calls[0].get("arguments") or {}
        self.assertEqual(cargo_args.get("limit"), 3)
        self.assertEqual(cargo_args.get("api_account_ids"), [2])
        cargo_data = (
            ((cargo_projection.data or {}).get("tool_results") or [{}])[0].get("data")
            or {}
        )
        cargo_rows = cargo_data.get("orders") or []
        self.assertEqual(len(cargo_rows), 3)
        self.assertTrue(
            all(row.get("status") == "ReadyToShip" for row in cargo_rows)
        )
        self.assertTrue(
            all(row.get("api_account_id") == 2 for row in cargo_rows)
        )
        cargo_answer = self._answer_from(cargo_projection)
        self.assertIn("kargo:", cargo_answer)
        self.assertNotIn("Magaza genel yorumu", cargo_answer)

    def test_latest_order_prompt_defaults_to_all_period_not_last_7_days(self) -> None:
        old_ms = int((datetime.now(TZ_GMT3) - timedelta(days=10)).timestamp() * 1000)
        self._insert_order(
            "ORD-LATEST-OLD-ONE",
            old_ms + 1_000,
            package_id=495,
            status="Delivered",
            first_name="Old",
            last_name="One",
        )
        self._insert_order(
            "ORD-LATEST-OLD-TWO",
            old_ms,
            package_id=496,
            status="Delivered",
            first_name="Old",
            last_name="Two",
        )

        result = run_agent_turn(
            "son gelen 2 siparisi soyle",
            planner_client=lambda _prompt: '{"mode":"tool","tool":"orders.lookup","arguments":{"query":"son gelen 2 siparis","limit":10}}',
            chat_client=lambda _prompt: "Bu cevap kullanilmamali.",
            db_name=self.db_path,
        )

        self.assertTrue(result.success, result.message)
        tool_call = (result.data or {}).get("tool_call") or {}
        self.assertEqual(tool_call.get("tool"), "orders.recent")
        self.assertEqual((tool_call.get("arguments") or {}).get("period"), "all")
        self.assertEqual((tool_call.get("arguments") or {}).get("limit"), 2)
        answer = self._answer_from(result)
        self.assertIn("ORD-LATEST-OLD-ONE", answer)
        self.assertIn("ORD-LATEST-OLD-TWO", answer)

    def test_latest_order_prompt_respects_explicit_last_7_days_period(self) -> None:
        old_ms = int((datetime.now(TZ_GMT3) - timedelta(days=10)).timestamp() * 1000)
        self._insert_order(
            "ORD-LATEST-OUTSIDE-WEEK",
            old_ms,
            package_id=497,
            status="Delivered",
            first_name="Outside",
            last_name="Week",
        )

        result = run_agent_turn(
            "son 7 gundeki son 2 siparisi soyle",
            planner_client=lambda _prompt: '{"mode":"tool","tool":"orders.lookup","arguments":{"query":"son 7 gundeki son siparis","limit":10}}',
            chat_client=lambda _prompt: "Bu cevap kullanilmamali.",
            db_name=self.db_path,
        )

        self.assertTrue(result.success, result.message)
        tool_call = (result.data or {}).get("tool_call") or {}
        self.assertEqual(tool_call.get("tool"), "orders.recent")
        self.assertEqual((tool_call.get("arguments") or {}).get("period"), "last_7_days")
        self.assertEqual((tool_call.get("arguments") or {}).get("limit"), 2)
        answer = self._answer_from(result)
        self.assertNotIn("ORD-LATEST-OUTSIDE-WEEK", answer)

    def test_followup_this_prints_latest_single_order_not_casual_fallback(self) -> None:
        now_ms = int(datetime.now(TZ_GMT3).timestamp() * 1000)
        self._insert_order(
            "ORD-LATEST-PRINT",
            now_ms,
            package_id=494,
            status="ReadyToShip",
            first_name="Latest",
            last_name="Print",
        )
        listed_answer = (
            "Lokal kayda gore 1 siparis buldum:\n"
            "1. ORD-LATEST-PRINT | durum: ReadyToShip | musteri: Latest Print | tutar: 100 TL"
        )
        history = [
            {"role": "user", "content": "son gelen 1 siparisi soyle"},
            {"role": "assistant", "content": listed_answer},
        ]

        result = run_agent_turn(
            "bunu yazdir",
            planner_client=lambda _prompt: '{"mode":"answer","response":"Iyiyim, buradayim."}',
            chat_client=lambda _prompt: "Iyiyim, buradayim.",
            db_name=self.db_path,
            history=history,
        )

        self.assertTrue(result.success, result.message)
        self.assertEqual(((result.data or {}).get("tool_call") or {}).get("tool"), "labels.create_from_orders.prepare")
        self._assert_prepare_answer_is_safe(result, expected_terms=("1 siparis", "ORD-LATEST-PRINT"))
        payload = (((self._action_from(result).get("arguments") or {}).get("payload")) or {})
        self.assertEqual(payload.get("order_numbers"), ["ORD-LATEST-PRINT"])

    def test_orders_how_are_they_prompt_runs_multi_signal_analysis_not_count_only(self) -> None:
        self._seed_ready_orders()
        now_ms = int(datetime.now(TZ_GMT3).timestamp() * 1000)
        self._insert_waiting_return(
            "CLAIM-ANALYSIS-WAIT",
            "ORD-MAHMUT-READY",
            now_ms,
            first_name="Mahmut",
            last_name="Sevket",
        )

        result = run_agent_turn(
            "siparislerim nasil sence",
            planner_client=lambda _prompt: '{"mode":"tool","tool":"orders.count","arguments":{"period":"today"}}',
            chat_client=lambda _prompt: (
                "Siparislerini tek sayiya indirgemeden baktim: bu ay 2 siparis ve 200 TRY ciro var; "
                "net kar icin maliyet verisi eksik. 2 paket kargoya hazir, 1 iade var ve kritik risk yok. "
                "Bence tablo pozitif ama iadeyi kontrol etmek lazim."
            ),
            db_name=self.db_path,
        )

        self.assertTrue(result.success, result.message)
        self.assertEqual((result.data or {}).get("mode"), "tools")
        tools = [row.get("tool") for row in ((result.data or {}).get("tool_results") or [])]
        self.assertIn("orders.count", tools)
        self.assertIn("orders.revenue_total", tools)
        self.assertIn("returns.count", tools)
        self.assertIn("orders.ready_to_ship_count", tools)
        self.assertIn("orders.fulfillment_risk_count", tools)
        self.assertGreater(len(set(tools)), 3)
        answer = self._answer_from(result)
        self.assertIn("tek sayiya", answer)
        self.assertIn("risk", answer)
        self.assertIn("200 TRY", answer)
        self.assertIn("maliyet verisi eksik", answer)
        self.assertNotIn("145 TRY", answer)
        self.assertNotIn("Siparis operasyon analizi:", answer)
        self.assertNotIn("2026-", answer)

    def test_orders_analysis_rejects_revenue_as_profit_when_cost_coverage_is_missing(self) -> None:
        self._seed_ready_orders()

        result = run_agent_turn(
            "siparislerim nasil sence",
            planner_client=lambda _prompt: '{"mode":"tool","tool":"orders.count","arguments":{"period":"today"}}',
            chat_client=lambda _prompt: (
                "Bu ay 2 siparis ve 200 TRY ciro var. Tahmini kar da 200 TRY; "
                "dolayisiyla karliyiz ve tablo pozitif."
            ),
            db_name=self.db_path,
        )

        self.assertTrue(result.success, result.message)
        tool_results = (result.data or {}).get("tool_results") or []
        profit_rows = [row for row in tool_results if row.get("tool") == "finance.profit_amount"]
        self.assertEqual(len(profit_rows), 1)
        profit_data = (profit_rows[0] or {}).get("data") or {}
        self.assertIsNone(profit_data.get("amount"))
        self.assertFalse(profit_data.get("estimated_profit_ready"))
        self.assertGreater(int(profit_data.get("incomplete_order_count") or 0), 0)
        answer = self._answer_from(result)
        self.assertNotIn("Tahmini kar: 200", answer)
        self.assertNotIn("dolayisiyla karliyiz", answer.casefold())
        self.assertIn("'karliyiz' diyemem", answer.casefold())
        self.assertIn("Tahmini kar: veri yok", answer)

    def test_explicit_net_profit_request_survives_revenue_as_profit_negation(self) -> None:
        self._seed_ready_orders()

        result = run_agent_turn(
            (
                "Bu ayki gercek net kari ver; az onceki ciroyu kar sayma "
                "ve maliyet eksikse tahmin uydurma."
            ),
            planner_client=lambda _prompt: (
                '{"mode":"tool","tool":"orders.revenue_total",'
                '"arguments":{"period":"this_month"}}'
            ),
            db_name=self.db_path,
        )

        self.assertTrue(result.success, result.message)
        tools = [
            row.get("tool")
            for row in ((result.data or {}).get("tool_results") or [])
        ]
        self.assertEqual(tools, ["finance.profit_amount"])
        profit_data = (
            ((result.data or {}).get("tool_results") or [{}])[0].get("data")
            or {}
        )
        self.assertIsNone(profit_data.get("amount"))
        self.assertFalse(profit_data.get("estimated_profit_ready"))
        answer = self._answer_from(result)
        self.assertIn("tahmini kar hesaplanamadi", answer.casefold())
        self.assertIn("maliyet kapsami", answer.casefold())

    def test_order_status_distribution_ignores_prohibited_revenue_commentary(self) -> None:
        self._seed_ready_orders()

        result = run_agent_turn(
            (
                "Bu ayki siparislerin durum dagilimini adetleriyle ver; "
                "gelir veya kar yorumu ekleme."
            ),
            planner_client=lambda _prompt: (
                '{"mode":"tool","tool":"orders.revenue_total",'
                '"arguments":{"period":"this_month"}}'
            ),
            db_name=self.db_path,
        )

        self.assertTrue(result.success, result.message)
        tools = [
            row.get("tool")
            for row in ((result.data or {}).get("tool_results") or [])
        ]
        self.assertEqual(tools, ["orders.status_counts"])
        answer = self._answer_from(result)
        self.assertIn("siparis durum dagilimi", answer.casefold())
        self.assertNotIn("ciro", answer.casefold())
        self.assertNotIn("kar", answer.casefold())

    def test_order_analysis_rejects_unsupported_healthy_or_stable_claim(self) -> None:
        self._seed_ready_orders()

        result = run_agent_turn(
            "siparislerim nasil sence",
            planner_client=lambda _prompt: '{"mode":"tool","tool":"orders.count","arguments":{"period":"today"}}',
            chat_client=lambda _prompt: (
                "Bu ay 2 siparis, 200 TRY ciro ve 200 TRY tahmini kar var. "
                "0 iade, 2 kargoya hazir paket, 0 risk ve 0 iptal gorunuyor. "
                "Genel olarak operasyonel durum saglikli ve stabil gorunuyor. "
                "Sonraki kontrol kargoya hazir siparisler olmali."
            ),
            db_name=self.db_path,
        )

        self.assertTrue(result.success, result.message)
        answer = self._answer_from(result)
        self.assertNotIn("saglikli", answer.casefold())
        self.assertNotIn("stabil", answer.casefold())
        self.assertIn("Siparis operasyon analizi", answer)
        self.assertIn("2 kargoya hazir paket islem bekliyor", answer)
        self.assertNotIn("kritik sinyal gorunmuyor", answer)

    def test_order_analysis_rejects_positive_health_claim_when_fulfillment_risk_exists(self) -> None:
        old_ms = int((datetime.now(TZ_GMT3) - timedelta(days=3)).timestamp() * 1000)
        self._insert_order(
            "ORD-RISK-POSITIVE-CLAIM",
            old_ms,
            package_id=970,
            first_name="Risk",
            last_name="Positive",
            status="ReadyToShip",
        )

        result = run_agent_turn(
            "son 8 gunde siparislerim nasil sence kapsamli degerlendir",
            planner_client=lambda _prompt: '{"mode":"tool","tool":"orders.count","arguments":{"period":"last_8_days"}}',
            chat_client=lambda _prompt: (
                "Son 8 gunde siparisleriniz genel olarak iyi durumda gorunuyor. "
                "Siparislerinizde olumlu bir tablo var. Maliyet kapsami eksik olsa da "
                "kargoya hazir riskli paket bulunuyor."
            ),
            db_name=self.db_path,
        )

        self.assertTrue(result.success, result.message)
        answer = self._answer_from(result)
        self.assertNotIn("genel olarak iyi", answer.casefold())
        self.assertNotIn("olumlu bir tablo", answer.casefold())
        self.assertIn("Siparis operasyon analizi", answer)
        self.assertIn("fulfillment riski", answer)

    def test_order_analysis_uses_custom_last_n_days_and_rejects_monthly_drift(self) -> None:
        now_ms = int(datetime.now(TZ_GMT3).timestamp() * 1000)
        old_ms = int((datetime.now(TZ_GMT3) - timedelta(days=9)).timestamp() * 1000)
        self._insert_order(
            "ORD-LAST8-IN",
            now_ms,
            package_id=971,
            first_name="Recent",
            last_name="Eight",
            status="Delivered",
            total_price=120.0,
        )
        self._insert_order(
            "ORD-LAST8-OUT",
            old_ms,
            package_id=972,
            first_name="Old",
            last_name="Outside",
            status="Delivered",
            total_price=999.0,
        )

        result = run_agent_turn(
            "siparislerim son 8 gunde nasil sence",
            planner_client=lambda _prompt: '{"mode":"answer","response":"Bu ay 999 siparis var ve performans iyilesmis."}',
            chat_client=lambda _prompt: "Bu ay 999 siparis var ve performans iyilesmis.",
            db_name=self.db_path,
        )

        self.assertTrue(result.success, result.message)
        self.assertEqual((result.data or {}).get("mode"), "tools")
        period_tools = {
            "orders.count",
            "orders.revenue_total",
            "finance.profit_amount",
            "returns.count",
            "returns.amount_total",
            "orders.cancelled_count",
        }
        for row in (result.data or {}).get("tool_results") or []:
            if row.get("tool") in period_tools:
                self.assertEqual((row.get("arguments") or {}).get("period"), "last_8_days", row)
        order_count = next(
            (row for row in (result.data or {}).get("tool_results") or [] if row.get("tool") == "orders.count"),
            {},
        )
        self.assertEqual((order_count.get("data") or {}).get("count"), 1)
        answer = self._answer_from(result)
        self.assertIn("Siparis operasyon analizi", answer)
        self.assertIn("1 siparislik hareket var", answer)
        self.assertNotIn("Bu ay 999", answer)
        self.assertNotIn("iyilesmis", answer.casefold())

    def test_simple_order_count_repairs_wrong_planner_period_to_custom_last_n_days(self) -> None:
        now_ms = int(datetime.now(TZ_GMT3).timestamp() * 1000)
        old_ms = int((datetime.now(TZ_GMT3) - timedelta(days=9)).timestamp() * 1000)
        self._insert_order(
            "ORD-LAST8-COUNT-IN",
            now_ms,
            package_id=973,
            first_name="Recent",
            last_name="Count",
            status="Delivered",
            total_price=120.0,
        )
        self._insert_order(
            "ORD-LAST8-COUNT-OUT",
            old_ms,
            package_id=974,
            first_name="Old",
            last_name="Count",
            status="Delivered",
            total_price=999.0,
        )

        result = run_agent_turn(
            "son 8 gunde kac siparisim var?",
            planner_client=lambda _prompt: '{"mode":"tool","tool":"orders.count","arguments":{"period":"today"}}',
            chat_client=lambda _prompt: "Bugun 999 siparis var.",
            db_name=self.db_path,
        )

        self.assertTrue(result.success, result.message)
        tool_call = (result.data or {}).get("tool_call") or {}
        self.assertEqual(tool_call.get("tool"), "orders.count")
        self.assertEqual((tool_call.get("arguments") or {}).get("period"), "last_8_days")
        self.assertEqual(((result.data or {}).get("tool_result") or {}).get("count"), 1)
        answer = self._answer_from(result)
        self.assertIn("1 siparis", answer)
        self.assertNotIn("999", answer)

    def test_order_count_quality_followup_compares_previous_equal_period_not_this_month(self) -> None:
        now = datetime.now(TZ_GMT3)
        for index, price in enumerate((120.0, 130.0, 140.0), start=1):
            self._insert_order(
                f"ORD-LAST8-CURRENT-{index}",
                int((now - timedelta(days=index - 1)).timestamp() * 1000),
                package_id=990 + index,
                first_name="Current",
                last_name=str(index),
                status="Delivered",
                total_price=price,
            )
        self._insert_order(
            "ORD-LAST8-PREVIOUS-1",
            int((now - timedelta(days=9)).timestamp() * 1000),
            package_id=998,
            first_name="Previous",
            last_name="One",
            status="Delivered",
            total_price=80.0,
        )
        self._insert_order(
            "ORD-LAST8-OLDER-OUTSIDE",
            int((now - timedelta(days=20)).timestamp() * 1000),
            package_id=999,
            first_name="Older",
            last_name="Outside",
            status="Delivered",
            total_price=999.0,
        )

        current_label = f"{(now - timedelta(days=7)).strftime('%Y-%m-%d')}..{now.strftime('%Y-%m-%d')}"
        history = [
            {
                "role": "assistant",
                "content": "Son 8 gunde 3 siparis var.",
                "context": {
                    "kind": "orders_metric_scope",
                    "source_tool": "orders.count",
                    "period": "last_8_days",
                    "period_label": current_label,
                    "order_count": 3,
                },
            }
        ]

        result = run_agent_turn(
            "3 siparis dedin bu sayi iyi mi sence",
            planner_client=lambda _prompt: '{"mode":"tool","tool":"orders.count","arguments":{"period":"this_month"}}',
            chat_client=lambda _prompt: "Bu ayla kiyaslayinca iyi; bu ay 999 siparis var.",
            db_name=self.db_path,
            history=history,
        )

        self.assertTrue(result.success, result.message)
        self.assertEqual((result.data or {}).get("mode"), "tools")
        count_rows = [
            row
            for row in ((result.data or {}).get("tool_results") or [])
            if row.get("tool") == "orders.count"
        ]
        self.assertGreaterEqual(len(count_rows), 2)
        self.assertEqual((count_rows[0].get("arguments") or {}).get("period"), "last_8_days")
        self.assertEqual((count_rows[0].get("data") or {}).get("count"), 3)
        previous_args = count_rows[1].get("arguments") or {}
        self.assertIn("start_date", previous_args)
        self.assertIn("end_date", previous_args)
        self.assertEqual((count_rows[1].get("data") or {}).get("count"), 1)

        answer = self._answer_from(result)
        self.assertIn("Esdeger donemle", answer)
        self.assertIn("3 siparis", answer)
        self.assertIn("1 siparis", answer)
        self.assertIn("onceki es donem", answer.casefold())
        self.assertNotIn("999", answer)
        self.assertNotIn("Bu ayla", answer)

    def test_internal_equivalent_period_comparison_never_routes_to_web(self) -> None:
        self._seed_ready_orders()
        result = run_agent_turn(
            "Son 7 günü ondan önceki eşit uzunluktaki 7 günle sipariş sayısı, ciro, iptal ve iade açısından karşılaştır; iki tarih aralığını da açık yaz.",
            planner_client=lambda _prompt: json.dumps(
                {
                    "mode": "tool",
                    "tool": "web.search",
                    "arguments": {"query": "sipariş performansı", "limit": 5},
                },
                ensure_ascii=False,
            ),
            db_name=self.db_path,
        )

        self.assertTrue(result.success, result.message)
        self.assertEqual((result.data or {}).get("mode"), "tools")
        self.assertEqual((result.data or {}).get("decision_path"), "intent_override")
        self.assertIs((result.data or {}).get("model_invoked"), False)
        tools = [
            row.get("tool")
            for row in ((result.data or {}).get("tool_calls") or [])
        ]
        self.assertNotIn("web.search", tools)
        self.assertEqual(tools.count("orders.count"), 2)
        self.assertEqual(tools.count("orders.revenue_total"), 2)
        self.assertEqual(tools.count("returns.count"), 2)
        self.assertEqual(tools.count("orders.cancelled_count"), 2)
        answer = self._answer_from(result)
        self.assertIn("Esdeger donemle", answer)
        self.assertIn("ortalama sepet", answer.casefold())
        self.assertNotIn("http", answer.casefold())

    def test_equal_period_comparison_respects_negative_profit_scope(self) -> None:
        result = run_agent_turn(
            (
                "Son 7 günü önceki 7 günle karşılaştır: sipariş adedi, ciro ve "
                "ortalama sepeti iki eşit dönem halinde ver; kâr ekleme."
            ),
            planner_client=lambda _prompt: (
                '{"mode":"tools","calls":['
                '{"tool":"orders.count","arguments":{"period":"last_7_days"}},'
                '{"tool":"orders.revenue_total","arguments":{"period":"last_7_days"}},'
                '{"tool":"finance.profit_amount","arguments":{"period":"last_7_days"}}]}'
            ),
            db_name=self.db_path,
        )

        self.assertTrue(result.success, result.message)
        tools = [
            row.get("tool")
            for row in ((result.data or {}).get("tool_calls") or [])
        ]
        self.assertEqual(tools.count("orders.count"), 2)
        self.assertEqual(tools.count("orders.revenue_total"), 2)
        self.assertNotIn("finance.profit_amount", tools)
        self.assertNotIn("returns.count", tools)
        self.assertNotIn("orders.cancelled_count", tools)
        answer = self._answer_from(result)
        self.assertIn("ortalama sepet", answer.casefold())
        self.assertNotIn("net kar", answer.casefold())

    def test_profit_followup_uses_current_metric_scope_and_never_labels_it_today(self) -> None:
        self._seed_ready_orders()
        history = [
            {
                "role": "assistant",
                "content": "Son dönem karşılaştırması tamamlandı.",
                "context": {
                    "kind": "orders_metric_scope",
                    "source_tool": "orders.count",
                    "period": "last_7_days",
                    "period_label": "2026-06-28..2026-07-04",
                    "start_date": "2026-06-28",
                    "end_date": "2026-07-04",
                },
            }
        ]
        result = run_agent_turn(
            "2026-06-28 ile 2026-07-04 arasındaki siparişlere göre net olarak kârlıyız ya da zarardayız diyebilir misin? Kesin bilmiyorsan sayı uydurma.",
            planner_client=lambda _prompt: (
                '{"mode":"tool","tool":"orders.count","arguments":{"period":"today"}}'
            ),
            db_name=self.db_path,
            history=history,
        )

        self.assertTrue(result.success, result.message)
        tool_calls = list((result.data or {}).get("tool_calls") or [])
        tools = [row.get("tool") for row in tool_calls]
        self.assertIn("finance.profit_amount", tools)
        count_call = next(row for row in tool_calls if row.get("tool") == "orders.count")
        self.assertEqual(
            count_call.get("arguments"),
            {"start_date": "2026-06-28", "end_date": "2026-07-04"},
        )
        answer = self._answer_from(result)
        self.assertNotIn("Bugun siparis", answer)
        self.assertIn("2026-06-28..2026-07-04", answer)
        self.assertIn("karliyiz", answer.casefold())

    def test_zero_order_profit_range_is_not_described_as_positive_or_zero_profit(self) -> None:
        result = run_agent_turn(
            "2020-01-01 ile 2020-01-07 arasındaki siparişlere göre kârlı mıyız zararda mıyız? Veri yoksa kesin sayı söyleme.",
            planner_client=lambda _prompt: (
                '{"mode":"tool","tool":"finance.profit_amount","arguments":{"period":"today"}}'
            ),
            db_name=self.db_path,
        )

        self.assertTrue(result.success, result.message)
        tool_calls = list((result.data or {}).get("tool_calls") or [])
        for row in tool_calls:
            if row.get("tool") in {
                "orders.count",
                "orders.revenue_total",
                "finance.profit_amount",
                "returns.count",
                "returns.amount_total",
                "orders.cancelled_count",
            }:
                self.assertEqual(
                    row.get("arguments"),
                    {"start_date": "2020-01-01", "end_date": "2020-01-07"},
                )
        answer = self._answer_from(result)
        self.assertNotIn("tahmini kar pozitif", answer.casefold())
        self.assertNotIn("Tahmini kar: 0 TRY", answer)
        self.assertIn("karlilik kaniti degil", answer.casefold())

    def test_orders_profit_anxiety_prompt_uses_profit_signal_and_explains_limits(self) -> None:
        now_ms = int(datetime.now(TZ_GMT3).timestamp() * 1000)
        self._insert_order(
            "ORD-PROFIT-ANXIETY",
            now_ms,
            package_id=495,
            status="ReadyToShip",
            first_name="Profit",
            last_name="Check",
            total_price=250.0,
        )

        result = run_agent_turn(
            "siparislerim ne durumda karli miyiz zarar da miyiz bilemiyorum hic guven vermiyor kapsamli bir izahat ver",
            planner_client=lambda _prompt: '{"mode":"tool","tool":"orders.count","arguments":{"period":"today"}}',
            chat_client=lambda _prompt: "Cloud cevabi.",
            db_name=self.db_path,
        )

        self.assertTrue(result.success, result.message)
        self.assertEqual((result.data or {}).get("mode"), "tools")
        tool_rows = (result.data or {}).get("tool_results") or []
        tools = [row.get("tool") for row in tool_rows]
        self.assertIn("orders.count", tools)
        self.assertIn("orders.revenue_total", tools)
        self.assertIn("finance.profit_amount", tools)
        self.assertIn("returns.amount_total", tools)
        for row in tool_rows:
            if row.get("tool") in {"orders.count", "orders.revenue_total", "finance.profit_amount", "returns.count", "returns.amount_total"}:
                self.assertEqual((row.get("arguments") or {}).get("period"), "this_month")
        answer = self._answer_from(result)
        self.assertIn("Once net karar", answer)
        self.assertIn("Guven veren taraf", answer)
        self.assertIn("Guven vermeyen taraf", answer)
        self.assertIn("Tahmini kar", answer)
        self.assertIn("Oncelik", answer)
        self.assertIn("Bir sonraki hamle", answer)
        self.assertNotIn("Bugun siparis: 0", answer)
        self.assertNotIn("Ã", answer)

    def test_llm_order_identifier_must_preserve_exact_hyphen_characters(self) -> None:
        self._seed_ready_orders()
        nb_hyphen = "\N{NON-BREAKING HYPHEN}"

        result = run_agent_turn(
            "son gelen 2 siparisi soyle",
            planner_client=lambda _prompt: '{"mode":"tool","tool":"orders.recent","arguments":{"limit":2}}',
            chat_client=lambda _prompt: (
                f"Son iki siparis: ORD{nb_hyphen}AYSE{nb_hyphen}READY ve ORD{nb_hyphen}MAHMUT{nb_hyphen}READY. "
                "Ikisi de kargoya hazir."
            ),
            db_name=self.db_path,
        )

        self.assertTrue(result.success, result.message)
        answer = self._answer_from(result)
        self.assertIn("ORD-AYSE-READY", answer)
        self.assertIn("ORD-MAHMUT-READY", answer)
        self.assertNotIn(nb_hyphen, answer)

    def test_approval_answer_restores_exact_order_identifiers_from_action_payload(self) -> None:
        self._seed_ready_orders()
        nb_hyphen = "\N{NON-BREAKING HYPHEN}"

        result = run_agent_turn(
            "ORD-MAHMUT-READY siparisini yazdir",
            planner_client=lambda _prompt: (
                '{"mode":"tool","tool":"labels.create_from_orders.prepare",'
                '"arguments":{"order_numbers":["ORD-MAHMUT-READY"]}}'
            ),
            chat_client=lambda _prompt: (
                f"ORD{nb_hyphen}MAHMUT{nb_hyphen}READY siparisi icin etiket hazir. "
                "Onay verirsen yazdiracagim."
            ),
            db_name=self.db_path,
        )

        self.assertTrue(result.success, result.message)
        self.assertEqual((result.data or {}).get("mode"), "approval_required")
        answer = self._answer_from(result)
        self.assertIn("ORD-MAHMUT-READY", answer)
        self.assertNotIn(nb_hyphen, answer)

    def test_approval_answer_falls_back_when_llm_omits_affected_order_ids(self) -> None:
        self._seed_ready_orders()

        result = run_agent_turn(
            "ORD-MAHMUT-READY ve ORD-AYSE-READY siparislerini yazdir",
            planner_client=lambda _prompt: (
                '{"mode":"tool","tool":"labels.create_from_orders.prepare",'
                '"arguments":{"order_numbers":["ORD-MAHMUT-READY","ORD-AYSE-READY"]}}'
            ),
            chat_client=lambda _prompt: (
                "Iki siparis icin etiket olusturma islemi hazir. "
                "Onay verirsen calistiracagim."
            ),
            db_name=self.db_path,
        )

        self.assertTrue(result.success, result.message)
        self.assertEqual((result.data or {}).get("mode"), "approval_required")
        answer = self._answer_from(result)
        self.assertIn("ORD-MAHMUT-READY", answer)
        self.assertIn("ORD-AYSE-READY", answer)
        self.assertIn("Gercek API istegi veya DB yazimi yapilmadi", answer)

    def test_label_followup_uses_previous_selection_even_when_llm_invents_wrong_order(self) -> None:
        self._seed_ready_orders()
        history: list[dict] = []

        first_prompt = "kargoya hazir siparisleri listele"
        history.append({"role": "user", "content": first_prompt})
        listed = run_agent_turn(first_prompt, history=history, db_name=self.db_path)
        self.assertTrue(listed.success, listed.message)
        context = (listed.data or {}).get("conversation_context") or {}
        self.assertEqual(set(context.get("order_numbers") or []), {"ORD-MAHMUT-READY", "ORD-AYSE-READY"})
        history.append({"role": "assistant", "content": self._answer_from(listed), "context": context})

        followup = "tamam şimdi bunları etikete yazdır"
        history.append({"role": "user", "content": followup})
        prepared = run_agent_turn(
            followup,
            planner_client=lambda _prompt: (
                '{"mode":"tool","tool":"labels.create_from_orders.prepare",'
                '"arguments":{"query":"11106942932","limit":5}}'
            ),
            chat_client=lambda _prompt: "Sadece 11106942932 siparisi icin etiket hazir, onay ver.",
            history=history,
            db_name=self.db_path,
        )

        self.assertTrue(prepared.success, prepared.message)
        self.assertEqual((prepared.data or {}).get("mode"), "approval_required")
        action = self._action_from(prepared)
        payload = (action.get("arguments") or {}).get("payload") or {}
        self.assertEqual(set(payload.get("order_numbers") or []), {"ORD-MAHMUT-READY", "ORD-AYSE-READY"})
        answer = self._answer_from(prepared)
        self.assertIn("ORD-MAHMUT-READY", answer)
        self.assertIn("ORD-AYSE-READY", answer)
        self.assertNotIn("11106942932", answer)

    def test_extracted_approval_preview_lists_every_affected_order_id(self) -> None:
        self._seed_ready_orders()

        result = run_agent_turn(
            "ORD-MAHMUT-READY ve ORD-AYSE-READY siparislerini fis cikarildi olarak isaretle",
            planner_client=lambda _prompt: (
                '{"mode":"tool","tool":"operations.manual_action.route",'
                '"arguments":{"query":"fis cikarildi isaretle","mode":"write","domain":"labels",'
                '"payload":{"order_numbers":["ORD-MAHMUT-READY","ORD-AYSE-READY"]}}}'
            ),
            chat_client=lambda _prompt: "Iki siparis icin onay gerekiyor.",
            db_name=self.db_path,
        )

        self.assertTrue(result.success, result.message)
        answer = self._answer_from(result)
        self.assertIn("ORD-MAHMUT-READY", answer)
        self.assertIn("ORD-AYSE-READY", answer)

    def test_multiple_numeric_order_ids_are_preserved_in_label_prepare(self) -> None:
        now_ms = int(datetime.now(TZ_GMT3).timestamp() * 1000)
        for index, order_number in enumerate(("11308276677", "11308086686"), start=1):
            self._insert_order(
                order_number,
                now_ms + index,
                package_id=5_800 + index,
                first_name="Numeric",
                last_name=f"Order {index}",
                status="ReadyToShip",
                cargo_tracking_number=f"TRK-{order_number}",
                cargo_provider_name="Trendyol Express",
            )

        result = run_agent_turn(
            "11308276677 ve 11308086686 siparislerini yazdir",
            db_name=self.db_path,
        )

        self.assertTrue(result.success, result.message)
        action = (((result.data or {}).get("tool_result") or {}).get("action") or {})
        payload = ((action.get("arguments") or {}).get("payload") or {})
        self.assertEqual(
            set(payload.get("order_numbers") or []),
            {"11308276677", "11308086686"},
        )
        answer = self._answer_from(result)
        self.assertIn("11308276677", answer)
        self.assertIn("11308086686", answer)

    def test_ready_list_context_carries_all_real_order_numbers_into_label_prepare(self) -> None:
        now_ms = int(datetime.now(TZ_GMT3).timestamp() * 1000)
        expected = [str(113_500_000_00 + index) for index in range(19)]
        for index, order_number in enumerate(expected):
            self._insert_order(
                order_number,
                now_ms + index,
                package_id=6_000 + index,
                first_name="Ready",
                last_name=f"Order {index}",
                status="ReadyToShip",
                cargo_tracking_number=f"TRK-{order_number}",
                cargo_provider_name="Trendyol Express",
            )

        listed = run_agent_turn(
            "toplam kac siparis var listele kargoya hazir siparisleri",
            planner_client=lambda _prompt: (
                '{"mode":"tool","tool":"orders.ready_to_ship",'
                '"arguments":{"period":"all","limit":100}}'
            ),
            chat_client=lambda _prompt: "Kargoya hazir ilk 10 siparisi listeledim.",
            db_name=self.db_path,
        )

        self.assertTrue(listed.success, listed.message)
        listed_data = listed.data or {}
        tool_data = listed_data.get("tool_result") or {}
        self.assertEqual(tool_data.get("ready_to_ship"), 19)
        self.assertEqual(tool_data.get("listed_count"), 19)
        self.assertFalse(tool_data.get("truncated"))
        self.assertEqual(set(tool_data.get("order_numbers") or []), set(expected))
        listed_answer = self._answer_from(listed)
        for order_number in expected:
            self.assertIn(order_number, listed_answer)

        history = [
            {"role": "user", "content": "toplam kac siparis var listele kargoya hazir siparisleri"},
            {
                "role": "assistant",
                "content": listed_answer,
                "context": dict(listed_data.get("conversation_context") or {}),
            },
        ]
        prepared = run_agent_turn(
            "tamam simdi bunlari yazdir etikete",
            db_name=self.db_path,
            history=history,
        )

        self.assertTrue(prepared.success, prepared.message)
        self.assertEqual((prepared.data or {}).get("mode"), "approval_required")
        action = self._action_from(prepared)
        payload = ((action.get("arguments") or {}).get("payload") or {})
        self.assertEqual(set(payload.get("order_numbers") or []), set(expected))

    def test_ready_list_rejects_llm_wrong_count_even_when_all_ids_are_present(self) -> None:
        now_ms = int(datetime.now(TZ_GMT3).timestamp() * 1000)
        expected = [str(113_510_000_00 + index) for index in range(5)]
        for index, order_number in enumerate(expected):
            self._insert_order(
                order_number,
                now_ms + index,
                package_id=6_100 + index,
                first_name="Ready",
                last_name=f"Count {index}",
                status="ReadyToShip",
                cargo_tracking_number=f"TRK-{order_number}",
                cargo_provider_name="Trendyol Express",
            )

        wrong_answer = "Bugun kargoya hazir 4 siparisiniz var:\n" + "\n".join(f"- {order_number}" for order_number in expected)
        result = run_agent_turn(
            "kargoya hazir siparisleri listele",
            planner_client=lambda _prompt: (
                '{"mode":"tool","tool":"orders.ready_to_ship",'
                '"arguments":{"period":"all","limit":100}}'
            ),
            chat_client=lambda _prompt: wrong_answer,
            db_name=self.db_path,
        )

        self.assertTrue(result.success, result.message)
        answer = self._answer_from(result)
        self.assertIn("Lokal kayda gore 5 siparis buldum", answer)
        self.assertNotIn("4 siparisiniz var", answer)
        for order_number in expected:
            self.assertIn(order_number, answer)

    def test_truncated_order_query_answer_mentions_total_and_returned_scope(self) -> None:
        now_ms = int(datetime.now(TZ_GMT3).timestamp() * 1000)
        expected = [str(113_520_000_00 + index) for index in range(19)]
        for index, order_number in enumerate(expected):
            self._insert_order(
                order_number,
                now_ms + index,
                package_id=6_200 + index,
                first_name="Ready",
                last_name=f"Page {index}",
                status="ReadyToShip",
                cargo_tracking_number=f"TRK-{order_number}",
                cargo_provider_name="Trendyol Express",
            )

        result = run_agent_turn(
            "kargoya hazir siparisleri ilk 10 kayit olarak listele",
            planner_client=lambda _prompt: (
                '{"mode":"tool","tool":"orders.query","arguments":{'
                '"period":"all","filters":{"status":"ReadyToShip"},'
                '"sort":"order_date_desc","mode":"list","limit":10,"page_size":10}}'
            ),
            chat_client=lambda _prompt: "10 siparis listeledim.",
            db_name=self.db_path,
        )

        self.assertTrue(result.success, result.message)
        tool_data = (result.data or {}).get("tool_result") or {}
        self.assertEqual(tool_data.get("count"), 19)
        self.assertEqual((tool_data.get("pagination") or {}).get("returned"), 10)
        answer = self._answer_from(result)
        self.assertIn("Lokal kayda gore 10 siparis buldum", answer)
        self.assertIn("Not: toplam 19 siparis eslesiyor; bu cevapta 10 kayit gosteriliyor", answer)

    def test_operational_orders_exclude_testing_lab_simulations(self) -> None:
        now_ms = int(datetime.now(TZ_GMT3).timestamp() * 1000)
        self._insert_order(
            "11350000123",
            now_ms,
            package_id=6_100,
            first_name="Real",
            last_name="Customer",
            status="ReadyToShip",
            cargo_provider_name="Trendyol Express",
        )
        self._insert_order(
            "SIM-1-TEST-1",
            now_ms + 1,
            package_id=6_101,
            first_name="Sim",
            last_name="Customer",
            status="ReadyToShip",
            cargo_provider_name="Simulated Cargo",
        )

        result = run_agent_turn("kargoya hazir siparisleri listele", db_name=self.db_path)

        self.assertTrue(result.success, result.message)
        data = result.data or {}
        tool_data = data.get("tool_result") or {}
        self.assertEqual(tool_data.get("ready_to_ship"), 1)
        self.assertEqual(tool_data.get("simulated_excluded"), 1)
        self.assertEqual(tool_data.get("order_numbers"), ["11350000123"])
        answer = self._answer_from(result)
        self.assertIn("11350000123", answer)
        self.assertNotIn("SIM-1-TEST-1", answer)
        self.assertIn("simülasyon", answer)

    def test_invalid_zero_record_label_prepare_never_opens_approval(self) -> None:
        result = run_agent_turn(
            "99999999999 siparisini yazdir",
            db_name=self.db_path,
        )

        self.assertTrue(result.success, result.message)
        data = result.data or {}
        self.assertEqual(data.get("mode"), "tool")
        self.assertFalse(data.get("approval_required"))
        self.assertIsNone(data.get("action_proposal"))
        answer = self._answer_from(result)
        self.assertIn("onaya hazırlanamadı", answer.casefold())
        self.assertIn("siparis bulunamadi", answer.casefold())
        self.assertNotIn("Onay gerekiyor", answer)

    def test_orders_what_can_we_do_prompt_answers_capabilities_not_count(self) -> None:
        result = run_agent_turn(
            "siparislerimle ilgili neler yapabiliriz",
            planner_client=lambda _prompt: '{"mode":"tool","tool":"orders.count","arguments":{"period":"today"}}',
            chat_client=lambda _prompt: "Yanlis count cevabi.",
            db_name=self.db_path,
        )

        self.assertTrue(result.success, result.message)
        self.assertEqual((result.data or {}).get("mode"), "answer")
        self.assertIsNone((result.data or {}).get("tool_call"))
        answer = self._answer_from(result)
        self.assertIn("etiket", answer)
        self.assertIn("iade", answer)
        self.assertIn("kargo", answer)
        self.assertNotIn("siparis var", answer)

    def test_customer_name_lookup_returns_the_customer_order_not_generic_results(self) -> None:
        self._seed_ready_orders()

        result = run_agent_turn("Mahmut Şevket isimli müşterinin siparişini göster", db_name=self.db_path)

        self.assertTrue(result.success, result.message)
        answer = (result.data or {}).get("answer") or result.message
        tool_data = (result.data or {}).get("tool_result") or {}
        orders = tool_data.get("orders") or []
        self.assertEqual(((result.data or {}).get("tool_call") or {}).get("tool"), "orders.lookup")
        self.assertEqual([row.get("order_number") for row in orders], ["ORD-MAHMUT-READY"], answer)
        self.assertIn("ORD-MAHMUT-READY", answer)
        self.assertIn("Mahmut", answer)
        self.assertNotIn("ORD-AYSE-READY", answer)

    def test_followup_excludes_customer_when_user_says_leave_that_one(self) -> None:
        self._seed_ready_orders()
        history = [
            {"role": "user", "content": "Kargoya hazır siparişleri getir"},
            {
                "role": "assistant",
                "content": "Kargoya hazir siparisler: ORD-MAHMUT-READY Mahmut Şevket, ORD-AYSE-READY Ayse Yilmaz.",
            },
        ]

        result = run_agent_turn("Mahmut olani birak", db_name=self.db_path, history=history)

        self.assertTrue(result.success, result.message)
        answer = (result.data or {}).get("answer") or result.message
        tool_data = (result.data or {}).get("tool_result") or {}
        orders = tool_data.get("orders") or []
        self.assertEqual(((result.data or {}).get("tool_call") or {}).get("tool"), "orders.lookup")
        self.assertEqual([row.get("order_number") for row in orders], ["ORD-AYSE-READY"], answer)
        self.assertIn("ORD-AYSE-READY", answer)
        self.assertNotIn("ORD-MAHMUT-READY", answer)

    def test_followup_exclusion_then_remaining_cargo_change_uses_remaining_orders(self) -> None:
        self._seed_ready_orders()
        listed = run_agent_turn("Kargoya hazir siparisleri getir", db_name=self.db_path)
        self.assertTrue(listed.success, listed.message)
        listed_answer = self._answer_from(listed)

        excluded = run_agent_turn(
            "Mahmut olani birak",
            db_name=self.db_path,
            history=[
                {"role": "user", "content": "Kargoya hazir siparisleri getir"},
                {"role": "assistant", "content": listed_answer},
            ],
        )
        self.assertTrue(excluded.success, excluded.message)
        excluded_answer = self._answer_from(excluded)
        self.assertIn("ORD-AYSE-READY", excluded_answer)
        self.assertNotIn("ORD-MAHMUT-READY", excluded_answer)

        prepared = run_agent_turn(
            "kalanlarin kargo firmasini Aras yap",
            db_name=self.db_path,
            history=[
                {"role": "user", "content": "Kargoya hazir siparisleri getir"},
                {"role": "assistant", "content": listed_answer},
                {"role": "user", "content": "Mahmut olani birak"},
                {"role": "assistant", "content": excluded_answer},
            ],
        )

        self._assert_prepare_answer_is_safe(prepared, expected_terms=("ORD-AYSE-READY", "kargo"))
        answer = self._answer_from(prepared)
        self.assertNotIn("ORD-MAHMUT-READY", answer)
        self.assertEqual(((prepared.data or {}).get("tool_call") or {}).get("tool"), "orders.change_cargo_provider.prepare")
        action = self._action_from(prepared)
        rows = [dict(row) for row in action.get("rows") or [] if isinstance(row, dict)]
        self.assertEqual(len(rows), 1)
        self.assertEqual((rows[0].get("payload") or {}).get("order"), "ORD-AYSE-READY")

    def test_followup_prints_previous_ready_order_list_and_execute_uses_that_set(self) -> None:
        self._seed_ready_orders()

        listed = run_agent_turn("Kargoya hazir siparisleri getir", db_name=self.db_path)
        self.assertTrue(listed.success, listed.message)
        listed_answer = (listed.data or {}).get("answer") or listed.message
        self.assertIn("ORD-MAHMUT-READY", listed_answer)
        self.assertIn("ORD-AYSE-READY", listed_answer)

        history = [
            {"role": "user", "content": "Kargoya hazir siparisleri getir"},
            {"role": "assistant", "content": listed_answer},
        ]
        prepared = run_agent_turn("onlari yazdir simdi", db_name=self.db_path, history=history)
        self._assert_prepare_answer_is_safe(prepared, expected_terms=("2 siparis",))
        action = self._action_from(prepared)
        payload = ((action.get("arguments") or {}).get("payload") or {})
        self.assertEqual(set(payload.get("order_numbers") or []), {"ORD-MAHMUT-READY", "ORD-AYSE-READY"})

        calls: list[list[str]] = []

        def fake_create_order_label_from_orders(order_numbers, *, brand_code="TANEX", model_code="TANEX_2736", db_name=None):
            calls.append(list(order_numbers))
            return Result.ok("Etiketler olusturuldu.", close_dialog=False, data={"order_numbers": list(order_numbers)})

        with tempfile.TemporaryDirectory() as temp_dir, patch(
            "Labels.processors.pipeline.create_order_label_from_orders",
            side_effect=fake_create_order_label_from_orders,
        ):
            executed = operation_runtime.execute_approved_operation(
                action,
                approved_by="behavioral-test",
                idempotency_key="followup-print-listed-orders",
                audit_path=Path(temp_dir) / "orders-agent-audit.jsonl",
            )

        verify = self._assert_execute_verified(executed)
        self.assertEqual(len(calls), 1)
        self.assertEqual(set(calls[0]), {"ORD-MAHMUT-READY", "ORD-AYSE-READY"})
        self.assertEqual(set(verify.get("order_numbers") or []), {"ORD-MAHMUT-READY", "ORD-AYSE-READY"})

    def test_followup_print_keeps_order_scope_after_long_conversation_interlude(self) -> None:
        self._seed_ready_orders()

        listed = run_agent_turn("Kargoya hazir siparisleri getir", db_name=self.db_path)
        self.assertTrue(listed.success, listed.message)
        listed_answer = self._answer_from(listed)
        history = [
            {"role": "user", "content": "Kargoya hazir siparisleri getir"},
            {"role": "assistant", "content": listed_answer},
        ]
        for index in range(6):
            history.extend(
                [
                    {"role": "user", "content": f"arada konustugumuz konu {index}"},
                    {"role": "assistant", "content": f"tamam, bunu not ettim {index}"},
                ]
            )

        prepared = run_agent_turn("az oncekileri yazdir", db_name=self.db_path, history=history)

        self._assert_prepare_answer_is_safe(prepared, expected_terms=("2 siparis",))
        action = self._action_from(prepared)
        payload = ((action.get("arguments") or {}).get("payload") or {})
        self.assertEqual(set(payload.get("order_numbers") or []), {"ORD-MAHMUT-READY", "ORD-AYSE-READY"})

    def test_reflective_number_chat_does_not_route_to_order_count(self) -> None:
        def misleading_planner(_prompt: str) -> str:
            return json.dumps(
                {
                    "mode": "tool",
                    "tool": "orders.count",
                    "arguments": {"period": "today"},
                }
            )

        result = run_agent_turn(
            "bazen sayilara fazla takiliyorum, bu normal mi sence",
            db_name=self.db_path,
            planner_client=misleading_planner,
            chat_client=misleading_planner,
        )

        self.assertTrue(result.success, result.message)
        self.assertEqual((result.data or {}).get("mode"), "answer")
        self.assertEqual((result.data or {}).get("tool_calls"), [])
        self.assertNotIn("siparis var", self._answer_from(result).casefold())

    def test_full_context_survives_eight_non_record_interludes(self) -> None:
        now_ms = int(datetime.now(TZ_GMT3).timestamp() * 1000)
        self._insert_order(
            "ORD-LONG-CONTEXT",
            now_ms,
            package_id=12_500,
            first_name="Mahmut",
            last_name="Sevket",
            status="ReadyToShip",
        )
        self._insert_order_item(
            "ORD-LONG-CONTEXT",
            now_ms + 1_000,
            line_id=12_501,
            product_name="Uzun Baglam Kalemi",
        )
        history: list[dict[str, object]] = [
            {
                "role": "assistant",
                "content": "1. ORD-LONG-CONTEXT | musteri: Mahmut Sevket",
                "context": {
                    "kind": "order_selection",
                    "order_numbers": ["ORD-LONG-CONTEXT"],
                },
            }
        ]
        for index in range(8):
            history.extend(
                [
                    {"role": "user", "content": f"arada serbest sohbet {index}"},
                    {"role": "assistant", "content": f"sohbet cevabi {index}"},
                ]
            )

        result = run_agent_turn(
            "sekiz mesaj once Mahmut Sevket icin ayirdigimiz siparis vardi ya, onun kalemlerini tekrar goster",
            db_name=self.db_path,
            history=history,
        )

        self.assertTrue(result.success, result.message)
        self.assertEqual(
            ((result.data or {}).get("tool_call") or {}).get("tool"),
            "orders.full_context",
        )
        self.assertEqual(
            (((result.data or {}).get("tool_call") or {}).get("arguments") or {}).get("query"),
            "ORD-LONG-CONTEXT",
        )
        answer = self._answer_from(result)
        self.assertIn("ORD-LONG-CONTEXT", answer)
        self.assertIn("Uzun Baglam Kalemi", answer)

    def test_label_prepare_carries_twenty_one_context_orders_with_natural_prepare_wording(self) -> None:
        now_ms = int(datetime.now(TZ_GMT3).timestamp() * 1000)
        order_numbers: list[str] = []
        for index in range(21):
            order_number = f"ORD-LONG-LABEL-{index:02d}"
            order_numbers.append(order_number)
            self._insert_order(
                order_number,
                now_ms + index,
                package_id=13_000 + index,
                first_name="Long",
                last_name=f"Label{index}",
                status="ReadyToShip",
            )
        history = [
            {
                "role": "assistant",
                "content": "Kargoya hazir 21 siparis listelendi.",
                "context": {
                    "kind": "order_selection",
                    "order_numbers": order_numbers,
                    "total_count": 21,
                },
            }
        ]

        prepared = run_agent_turn(
            "tamam bu 21 siparisin tamamini etikete yazdirmaya hazirla",
            db_name=self.db_path,
            history=history,
        )

        self._assert_prepare_answer_is_safe(prepared, expected_terms=("21 siparis",))
        self.assertEqual(
            ((prepared.data or {}).get("tool_call") or {}).get("tool"),
            "labels.create_from_orders.prepare",
        )
        action = self._action_from(prepared)
        payload = ((action.get("arguments") or {}).get("payload") or {})
        self.assertEqual(set(payload.get("order_numbers") or []), set(order_numbers))

    def test_profitability_followup_inherits_previous_analysis_period(self) -> None:
        now_ms = int(datetime.now(TZ_GMT3).timestamp() * 1000)
        self._insert_order(
            "ORD-PROFIT-FOLLOWUP",
            now_ms,
            package_id=13_100,
            first_name="Profit",
            last_name="Followup",
            status="Delivered",
        )
        history = [
            {
                "role": "assistant",
                "content": (
                    "Son 8 gun icin ciro var ama maliyet kapsami eksik; "
                    "'karliyiz' diyemem."
                ),
                "context": {
                    "kind": "orders_metric_scope",
                    "period": "last_8_days",
                    "start_date": "2026-06-25",
                    "end_date": "2026-07-02",
                },
            }
        ]

        result = run_agent_turn(
            "peki bu verilerle gercekten karliyiz diyebilir misin, eksik veriyi acikca soyle",
            db_name=self.db_path,
            history=history,
        )

        self.assertTrue(result.success, result.message)
        self.assertEqual(
            ((result.data or {}).get("tool_call") or {}).get("tool"),
            "finance.profit_amount",
        )
        arguments = ((result.data or {}).get("tool_call") or {}).get("arguments") or {}
        self.assertEqual(arguments.get("start_date"), "2026-06-25")
        self.assertEqual(arguments.get("end_date"), "2026-07-02")
        self.assertIn("maliyet", self._answer_from(result).casefold())
        self.assertFalse(
            _answer_claims_unsupported_profit(self._answer_from(result))
        )

    def test_followup_prints_customer_refined_previous_ready_order_only(self) -> None:
        self._seed_ready_orders()

        listed = run_agent_turn("Kargoya hazir siparisleri getir", db_name=self.db_path)
        self.assertTrue(listed.success, listed.message)
        listed_answer = (listed.data or {}).get("answer") or listed.message
        self.assertIn("ORD-MAHMUT-READY", listed_answer)
        self.assertIn("ORD-AYSE-READY", listed_answer)

        history = [
            {"role": "user", "content": "Kargoya hazir siparisleri getir"},
            {"role": "assistant", "content": listed_answer},
        ]
        print_message = "Mahmut olan" + chr(0x0131) + " yazd" + chr(0x0131) + "r"
        prepared = run_agent_turn(print_message, db_name=self.db_path, history=history)
        self._assert_prepare_answer_is_safe(prepared, expected_terms=("1 siparis", "ORD-MAHMUT-READY"))
        self.assertNotIn("ORD-AYSE-READY", self._answer_from(prepared))
        action = self._action_from(prepared)
        payload = ((action.get("arguments") or {}).get("payload") or {})
        self.assertEqual(payload.get("order_numbers"), ["ORD-MAHMUT-READY"])

        calls: list[list[str]] = []

        def fake_create_order_label_from_orders(order_numbers, *, brand_code="TANEX", model_code="TANEX_2736", db_name=None):
            calls.append(list(order_numbers))
            return Result.ok("Etiketler olusturuldu.", close_dialog=False, data={"order_numbers": list(order_numbers)})

        with tempfile.TemporaryDirectory() as temp_dir, patch(
            "Labels.processors.pipeline.create_order_label_from_orders",
            side_effect=fake_create_order_label_from_orders,
        ):
            executed = operation_runtime.execute_approved_operation(
                action,
                approved_by="behavioral-test",
                idempotency_key="followup-print-customer-refined-order",
                audit_path=Path(temp_dir) / "orders-agent-audit.jsonl",
            )

        verify = self._assert_execute_verified(executed)
        self.assertEqual(calls, [["ORD-MAHMUT-READY"]])
        self.assertEqual(verify.get("order_numbers"), ["ORD-MAHMUT-READY"])

    def test_followup_print_reference_without_previous_orders_does_not_prepare_broad_labels(self) -> None:
        history = [
            {"role": "user", "content": "son gelen 2 siparisi soyle"},
            {"role": "assistant", "content": "Son 7 gunde gelen 2 siparis yok gorunuyor."},
        ]

        result = run_agent_turn("bu iki siparisi yazdir", db_name=self.db_path, history=history)

        self.assertTrue(result.success, result.message)
        data = result.data or {}
        self.assertEqual(data.get("mode"), "answer")
        self.assertFalse(data.get("approval_required"))
        self.assertEqual(data.get("tool_calls"), [])
        answer = str(data.get("answer") or result.message)
        self.assertIn("net bir siparis listesi bulamadim", answer)

    def test_followup_prints_previous_numeric_recent_order_list_only(self) -> None:
        now_ms = int(datetime.now(TZ_GMT3).timestamp() * 1000)
        seeded = [
            ("11308276677", 100.0, "Oben", "Ciga"),
            ("11308086686", 99.98, "Seyma", "Sahin"),
            ("11308082825", 89.98, "Hilal", "Gundogar"),
            ("11307000000", 10.0, "Extra", "Customer"),
        ]
        for index, (order_number, total, first_name, last_name) in enumerate(seeded):
            self._insert_order(
                order_number,
                now_ms - index,
                package_id=3_000 + index,
                first_name=first_name,
                last_name=last_name,
                status="ReadyToShip",
            )
            with Session(get_engine(self.db_path)) as session:
                row = session.exec(select(OrderData).where(OrderData.orderNumber == order_number)).one()
                row.totalPrice = total
                row.packageTotalPrice = total
                session.add(row)
                session.commit()

        list_message = "son gelen 3 sipari" + chr(0x015F) + "i s" + chr(0x00F6) + "yle"
        print_message = (
            chr(0x015F)
            + "imdi bu "
            + chr(0x00FC)
            + chr(0x00E7)
            + " sipari"
            + chr(0x015F)
            + "i yazd"
            + chr(0x0131)
            + "r"
        )

        listed = run_agent_turn(list_message, db_name=self.db_path)
        self.assertTrue(listed.success, listed.message)
        self.assertEqual((((listed.data or {}).get("tool_call") or {}).get("arguments") or {}).get("limit"), 3)
        listed_answer = (listed.data or {}).get("answer") or listed.message
        self.assertIn("11308276677", listed_answer)
        self.assertIn("11308086686", listed_answer)
        self.assertIn("11308082825", listed_answer)
        self.assertNotIn("11307000000", listed_answer)

        history = [
            {"role": "user", "content": list_message},
            {"role": "assistant", "content": listed_answer},
        ]
        prepared = run_agent_turn(print_message, db_name=self.db_path, history=history)
        self._assert_prepare_answer_is_safe(prepared, expected_terms=("3 siparis",))
        action = self._action_from(prepared)
        payload = ((action.get("arguments") or {}).get("payload") or {})
        expected = {"11308276677", "11308086686", "11308082825"}
        self.assertEqual(set(payload.get("order_numbers") or []), expected)
        self.assertNotIn("11307000000", payload.get("order_numbers") or [])

        calls: list[list[str]] = []

        def fake_create_order_label_from_orders(order_numbers, *, brand_code="TANEX", model_code="TANEX_2736", db_name=None):
            calls.append(list(order_numbers))
            return Result.ok("Etiketler olusturuldu.", close_dialog=False, data={"order_numbers": list(order_numbers)})

        with tempfile.TemporaryDirectory() as temp_dir, patch(
            "Labels.processors.pipeline.create_order_label_from_orders",
            side_effect=fake_create_order_label_from_orders,
        ):
            executed = operation_runtime.execute_approved_operation(
                action,
                approved_by="behavioral-test",
                idempotency_key="followup-print-numeric-recent-orders",
                audit_path=Path(temp_dir) / "orders-agent-audit.jsonl",
            )

        self._assert_execute_verified(executed)
        self.assertEqual(len(calls), 1)
        self.assertEqual(set(calls[0]), expected)
        self.assertNotIn("11307000000", calls[0])

    def test_followup_print_keeps_order_context_after_chat_and_help_interlude(self) -> None:
        self._seed_ready_orders()

        listed = run_agent_turn("Kargoya hazir siparisleri getir", db_name=self.db_path)
        self.assertTrue(listed.success, listed.message)
        listed_answer = self._answer_from(listed)
        self.assertIn("ORD-MAHMUT-READY", listed_answer)
        self.assertIn("ORD-AYSE-READY", listed_answer)

        name_answer = run_agent_turn(
            "adin ne",
            planner_client=lambda _prompt: '{"mode":"answer","response":"Ben OrderScout AI."}',
            chat_client=lambda _prompt: "Ben OrderScout AI.",
            db_name=self.db_path,
            history=[
                {"role": "user", "content": "Kargoya hazir siparisleri getir"},
                {"role": "assistant", "content": listed_answer},
            ],
        )
        self.assertTrue(name_answer.success, name_answer.message)
        self.assertEqual((name_answer.data or {}).get("mode"), "answer")

        help_answer = run_agent_turn(
            "siparislerde hangi islemleri yapabiliyorsun",
            db_name=self.db_path,
            history=[
                {"role": "user", "content": "Kargoya hazir siparisleri getir"},
                {"role": "assistant", "content": listed_answer},
                {"role": "user", "content": "adin ne"},
                {"role": "assistant", "content": self._answer_from(name_answer)},
            ],
        )
        self.assertTrue(help_answer.success, help_answer.message)
        self.assertEqual((help_answer.data or {}).get("mode"), "answer")
        self.assertIn("etiket", self._answer_from(help_answer))

        prepared = run_agent_turn(
            "simdi bu iki siparisi yazdir",
            db_name=self.db_path,
            history=[
                {"role": "user", "content": "Kargoya hazir siparisleri getir"},
                {"role": "assistant", "content": listed_answer},
                {"role": "user", "content": "adin ne"},
                {"role": "assistant", "content": self._answer_from(name_answer)},
                {"role": "user", "content": "siparislerde hangi islemleri yapabiliyorsun"},
                {"role": "assistant", "content": self._answer_from(help_answer)},
            ],
        )
        self._assert_prepare_answer_is_safe(prepared, expected_terms=("2 siparis",))
        action = self._action_from(prepared)
        payload = ((action.get("arguments") or {}).get("payload") or {})
        self.assertEqual(set(payload.get("order_numbers") or []), {"ORD-MAHMUT-READY", "ORD-AYSE-READY"})

        calls: list[list[str]] = []

        def fake_create_order_label_from_orders(order_numbers, *, brand_code="TANEX", model_code="TANEX_2736", db_name=None):
            calls.append(list(order_numbers))
            return Result.ok("Etiketler olusturuldu.", close_dialog=False, data={"order_numbers": list(order_numbers)})

        with tempfile.TemporaryDirectory() as temp_dir, patch(
            "Labels.processors.pipeline.create_order_label_from_orders",
            side_effect=fake_create_order_label_from_orders,
        ):
            executed = operation_runtime.execute_approved_operation(
                action,
                approved_by="behavioral-test",
                idempotency_key="followup-print-after-chat-help-interlude",
                audit_path=Path(temp_dir) / "orders-agent-audit.jsonl",
            )

        verify = self._assert_execute_verified(executed)
        self.assertEqual(len(calls), 1)
        self.assertEqual(set(calls[0]), {"ORD-MAHMUT-READY", "ORD-AYSE-READY"})
        self.assertEqual(set(verify.get("order_numbers") or []), {"ORD-MAHMUT-READY", "ORD-AYSE-READY"})

    def test_followup_marks_previous_ready_orders_extracted_after_label_flow(self) -> None:
        self._seed_ready_orders()
        self._insert_order(
            "ORD-DELIVERED-NOT-SCOPED",
            int(datetime.now(TZ_GMT3).timestamp() * 1000) + 5_000,
            package_id=903,
            first_name="Teslim",
            last_name="Edildi",
            status="Delivered",
        )

        listed = run_agent_turn("Kargoya hazir siparisleri getir", db_name=self.db_path)
        self.assertTrue(listed.success, listed.message)
        listed_answer = self._answer_from(listed)
        self.assertIn("ORD-MAHMUT-READY", listed_answer)
        self.assertIn("ORD-AYSE-READY", listed_answer)
        self.assertNotIn("ORD-DELIVERED-NOT-SCOPED", listed_answer)

        history = [
            {"role": "user", "content": "Kargoya hazir siparisleri getir"},
            {"role": "assistant", "content": listed_answer},
        ]
        prepared = run_agent_turn("bunlari fis cikarildi olarak isaretle", db_name=self.db_path, history=history)
        self._assert_prepare_answer_is_safe(prepared, expected_terms=("2 siparis",))
        action = self._action_from(prepared)
        self.assertEqual(action.get("action_type"), "labels.mark_orders_extracted_from_label_result")
        payload = ((action.get("arguments") or {}).get("payload") or {})
        self.assertEqual(set(payload.get("order_numbers") or []), {"ORD-MAHMUT-READY", "ORD-AYSE-READY"})

        with tempfile.TemporaryDirectory() as temp_dir:
            executed = operation_runtime.execute_approved_operation(
                action,
                approved_by="behavioral-test",
                idempotency_key="followup-mark-label-extracted",
                audit_path=Path(temp_dir) / "orders-agent-audit.jsonl",
            )

        self._assert_execute_verified(executed)
        with Session(get_engine(self.db_path)) as session:
            scoped = session.exec(
                select(OrderHeader).where(OrderHeader.orderNumber.in_(["ORD-MAHMUT-READY", "ORD-AYSE-READY"]))
            ).all()
            self.assertEqual({row.orderNumber for row in scoped if row.is_extracted}, {"ORD-MAHMUT-READY", "ORD-AYSE-READY"})
            delivered = session.exec(select(OrderHeader).where(OrderHeader.orderNumber == "ORD-DELIVERED-NOT-SCOPED")).one()
            self.assertFalse(bool(delivered.is_extracted))

    def test_followup_approves_previous_return_list_and_execute_uses_that_claim(self) -> None:
        now_ms = self._seed_ready_orders()
        self._insert_waiting_return(
            "CLAIM-MAHMUT-RETURN",
            "ORD-MAHMUT-READY",
            now_ms + 2_000,
            first_name="Mahmut",
            last_name="Sevket",
        )

        listed = run_agent_turn("Bekleyen iadeleri getir", db_name=self.db_path)
        self.assertTrue(listed.success, listed.message)
        listed_answer = (listed.data or {}).get("answer") or listed.message
        self.assertIn("CLAIM-MAHMUT-RETURN", listed_answer)

        history = [
            {"role": "user", "content": "Bekleyen iadeleri getir"},
            {"role": "assistant", "content": listed_answer},
        ]
        prepared = run_agent_turn("onlari onayla", db_name=self.db_path, history=history)
        self._assert_prepare_answer_is_safe(prepared, expected_terms=("iade",))
        action = self._action_from(prepared)
        payload = ((action.get("arguments") or {}).get("payload") or {})
        self.assertEqual(payload.get("order"), "CLAIM-MAHMUT-RETURN")

        def fake_approve_return_claim(*, order, **_kwargs):
            with Session(get_engine(self.db_path)) as session:
                claim = session.exec(select(ReturnClaim).where(ReturnClaim.claimId == order.claimId)).one()
                claim.claimItemStatus = "Accepted"
                claim.waitingClaimLineItemIds = []
                session.add(claim)
                session.commit()
            return Result.ok("Iade onaylandi.", close_dialog=False, data={"claim_id": order.claimId})

        with tempfile.TemporaryDirectory() as temp_dir, patch(
            "Orders.views.actions.approve_return_claim",
            side_effect=fake_approve_return_claim,
        ):
            executed = operation_runtime.execute_approved_operation(
                action,
                approved_by="behavioral-test",
                idempotency_key="followup-approve-listed-return",
                audit_path=Path(temp_dir) / "orders-agent-audit.jsonl",
            )

        self._assert_execute_verified(executed)
        with Session(get_engine(self.db_path)) as session:
            claim = session.exec(select(ReturnClaim).where(ReturnClaim.claimId == "CLAIM-MAHMUT-RETURN")).one()
            self.assertEqual(claim.claimItemStatus, "Accepted")
            self.assertEqual(claim.waitingClaimLineItemIds, [])

    def test_followup_return_approval_keeps_latest_claim_after_long_interlude(self) -> None:
        now_ms = self._seed_ready_orders()
        self._insert_waiting_return(
            "CLAIM-LONG-CONTEXT",
            "ORD-MAHMUT-READY",
            now_ms + 2_000,
            first_name="Mahmut",
            last_name="Sevket",
        )
        listed = run_agent_turn("Bekleyen iadeleri getir", db_name=self.db_path)
        self.assertTrue(listed.success, listed.message)
        listed_answer = self._answer_from(listed)
        history = [
            {"role": "user", "content": "Bekleyen iadeleri getir"},
            {"role": "assistant", "content": listed_answer},
        ]
        for index in range(6):
            history.extend(
                [
                    {"role": "user", "content": f"arada iade disi sohbet {index}"},
                    {"role": "assistant", "content": f"konusmaya devam ediyoruz {index}"},
                ]
            )

        prepared = run_agent_turn("az onceki iadeyi onayla", db_name=self.db_path, history=history)

        self._assert_prepare_answer_is_safe(prepared, expected_terms=("CLAIM-LONG-CONTEXT",))
        action = self._action_from(prepared)
        payload = ((action.get("arguments") or {}).get("payload") or {})
        self.assertEqual(payload.get("order"), "CLAIM-LONG-CONTEXT")

    def test_followup_this_approves_previous_uuid_return_list(self) -> None:
        now_ms = self._seed_ready_orders()
        claim_id = "75a7c1c0-73a2-49f7-bb8a-366dee971405"
        self._insert_waiting_return(
            claim_id,
            "11268067005",
            now_ms + 2_000,
            first_name="Samir",
            last_name="Abdurrahman",
        )

        listed = run_agent_turn(
            "onay bekleyen iadeleri getir",
            planner_client=lambda _prompt: '{"mode":"tool","tool":"returns.query","arguments":{"period":"all","filters":{"status":"WaitingInAction"},"sort":"claim_date_desc","limit":20,"mode":"list"}}',
            chat_client=lambda _prompt: "Buradayim. Ne yapalim?",
            db_name=self.db_path,
        )
        self.assertTrue(listed.success, listed.message)
        self.assertEqual(((listed.data or {}).get("tool_call") or {}).get("tool"), "returns.query")
        listed_answer = self._answer_from(listed)
        self.assertIn(claim_id, listed_answer)
        self.assertIn("Samir", listed_answer)

        history = [
            {"role": "user", "content": "onay bekleyen iadeleri getir"},
            {"role": "assistant", "content": listed_answer},
        ]
        prepared = run_agent_turn(
            "bunu onayla",
            planner_client=lambda _prompt: '{"mode":"answer","response":"Iyiyim, buradayim. Istersen biraz konusalim."}',
            chat_client=lambda _prompt: "Iyiyim, buradayim. Istersen biraz konusalim.",
            db_name=self.db_path,
            history=history,
        )

        self._assert_prepare_answer_is_safe(prepared, expected_terms=("iade",))
        self.assertEqual(((prepared.data or {}).get("tool_call") or {}).get("tool"), "returns.approve.prepare")
        answer = self._answer_from(prepared)
        self.assertIn(claim_id, answer)
        self.assertNotIn("Iyiyim, buradayim", answer)
        action = self._action_from(prepared)
        self.assertEqual(action.get("action_type"), "orders.approve_return_claim")
        payload = ((action.get("arguments") or {}).get("payload") or {})
        self.assertEqual(payload.get("order"), claim_id)

    def test_return_prepare_natural_wording_uses_structured_claim_context(self) -> None:
        now_ms = self._seed_ready_orders()
        claim_id = "95a7c1c0-73a2-49f7-bb8a-366dee971407"
        self._insert_waiting_return(
            claim_id,
            "ORD-MAHMUT-READY",
            now_ms + 2_000,
            first_name="Structured",
            last_name="Prepare",
        )
        history = [
            {
                "role": "assistant",
                "content": "Tek iade kaydi secildi.",
                "context": {
                    "kind": "return_selection",
                    "claim_ids": [claim_id],
                    "order_numbers": ["ORD-MAHMUT-READY"],
                },
            }
        ]

        prepared = run_agent_turn(
            "tamam bunu onaylamaya hazirla ama benden onay almadan islem yapma",
            db_name=self.db_path,
            history=history,
        )

        self._assert_prepare_answer_is_safe(prepared, expected_terms=(claim_id,))
        self.assertEqual(
            ((prepared.data or {}).get("tool_call") or {}).get("tool"),
            "returns.approve.prepare",
        )
        action = self._action_from(prepared)
        payload = ((action.get("arguments") or {}).get("payload") or {})
        self.assertEqual(payload.get("order"), claim_id)

    def test_return_onaya_hazirla_uses_second_selected_claim(self) -> None:
        now_ms = self._seed_ready_orders()
        claims = [
            ("95a7c1c0-73a2-49f7-bb8a-366dee971408", "ORD-MAHMUT-READY"),
            ("95a7c1c0-73a2-49f7-bb8a-366dee971409", "ORD-AYSE-READY"),
        ]
        for index, (claim_id, order_number) in enumerate(claims):
            self._insert_waiting_return(
                claim_id,
                order_number,
                now_ms + 2_000 + index,
                first_name=f"Ordinal{index}",
                last_name="Prepare",
            )
        history = [
            {
                "role": "assistant",
                "content": (
                    f"1. {claims[0][0]} | siparis: {claims[0][1]}\n"
                    f"2. {claims[1][0]} | siparis: {claims[1][1]}"
                ),
                "context": {
                    "kind": "return_selection",
                    "claim_ids": [claim_id for claim_id, _ in claims],
                    "order_numbers": [order_number for _, order_number in claims],
                },
            },
            {
                "role": "assistant",
                "content": f"Tek iade bulundu: {claims[1][0]} | siparis: {claims[1][1]}",
                "context": {
                    "kind": "return_selection",
                    "claim_ids": [claims[1][0]],
                    "order_numbers": [claims[1][1]],
                },
            },
        ]

        prepared = run_agent_turn(
            "İkinci iadeyi doğru buldun ama hazırlamadın. Şimdi az önce bulduğun o tek iadeyi onaya hazırla; gerçek işlemi uygulama, açık onayımı bekle.",
            planner_client=lambda _prompt: (
                '{"mode":"tool","tool":"returns.lookup",'
                '"arguments":{"query":"ikinci sıradaki","limit":10}}'
            ),
            db_name=self.db_path,
            history=history,
        )

        self._assert_prepare_answer_is_safe(prepared, expected_terms=(claims[1][0],))
        self.assertEqual(
            ((prepared.data or {}).get("tool_call") or {}).get("tool"),
            "returns.approve.prepare",
        )
        action = self._action_from(prepared)
        payload = ((action.get("arguments") or {}).get("payload") or {})
        self.assertEqual(payload.get("order"), claims[1][0])

    def test_pending_return_cancel_is_tool_free_even_when_planner_requests_lookup(self) -> None:
        def misleading_planner(_prompt: str) -> str:
            return json.dumps(
                {
                    "mode": "tool",
                    "tool": "orders.lookup",
                    "arguments": {"query": "iptal", "limit": 10},
                }
            )

        result = run_agent_turn(
            "bunu da simdilik iptal et, pazaryerine istek gonderme",
            db_name=self.db_path,
            planner_client=misleading_planner,
            chat_client=misleading_planner,
        )

        self.assertTrue(result.success, result.message)
        self.assertEqual((result.data or {}).get("mode"), "answer")
        self.assertEqual((result.data or {}).get("tool_calls"), [])
        answer = self._answer_from(result).casefold()
        self.assertIn("api", answer)
        self.assertIn("yapilmadi", answer)

    def test_pending_cancel_answers_independent_followup_after_safe_confirmation(self) -> None:
        result = run_agent_turn(
            "Bu etiket işlemini iptal et, dosya üretme. Sonra bana söyle: bu konuşmada seni en çok zorlayan bağlam hatası neydi?",
            planner_client=lambda _prompt: (
                "En zor kisim, eski iade seciminin daha yeni siparis secimini "
                "golgelemesini engelleyip dogru uc kaydi korumakti."
            ),
            db_name=self.db_path,
        )

        self.assertTrue(result.success, result.message)
        self.assertEqual((result.data or {}).get("mode"), "answer")
        self.assertEqual((result.data or {}).get("tool_calls"), [])
        answer = self._answer_from(result)
        self.assertIn("dosya/cikti uretimi yapilmadi", answer)
        self.assertIn("eski iade seciminin", answer)
        self.assertEqual(
            (result.data or {}).get("answer_origin"),
            "deterministic_cancel_with_llm_followup",
        )

    def test_pending_cancel_answers_followup_total_question(self) -> None:
        result = run_agent_turn(
            "Vazgeç, onay hazırlığını iptal et. Sonra listedeki iki iadenin toplam tutarını söyle.",
            planner_client=lambda _prompt: "İki iadenin toplam tutarı 659.97 TL.",
            db_name=self.db_path,
        )

        self.assertTrue(result.success, result.message)
        self.assertEqual((result.data or {}).get("tool_calls"), [])
        answer = self._answer_from(result)
        self.assertIn("islemi calistirmiyorum", answer)
        self.assertIn("659.97 TL", answer)
        self.assertEqual(
            (result.data or {}).get("answer_origin"),
            "deterministic_cancel_with_llm_followup",
        )

    def test_approval_explanation_never_opens_a_new_label_action(self) -> None:
        result = run_agent_turn(
            "İptal edildiğini anladım. Peki etiket hazırlığında neden benden onay istedin?",
            planner_client=lambda _prompt: (
                '{"mode":"tool","tool":"operations.manual_action.route",'
                '"arguments":{"query":"etiket hazırla","mode":"write"}}'
            ),
            db_name=self.db_path,
        )

        self.assertTrue(result.success, result.message)
        self.assertEqual((result.data or {}).get("mode"), "answer")
        self.assertEqual((result.data or {}).get("tool_calls"), [])
        self.assertFalse(bool((result.data or {}).get("approval_required")))
        answer = self._answer_from(result)
        self.assertIn("onay istedim", answer.casefold())
        self.assertIn("açık onay olmadan execute edilmez", answer)

    def test_cancel_with_inflected_approval_question_answers_both_parts(self) -> None:
        result = run_agent_turn(
            (
                "Vazgeç, etiket hazırlığını iptal et; hiçbir dosya üretme. "
                "Sonra neden benden açık onay istediğini somut anlat."
            ),
            planner_client=lambda _prompt: (
                '{"mode":"tool","tool":"labels.create_from_orders.prepare",'
                '"arguments":{"query":"etiket hazırla","limit":20}}'
            ),
            db_name=self.db_path,
        )

        self.assertTrue(result.success, result.message)
        self.assertEqual((result.data or {}).get("mode"), "answer")
        self.assertEqual((result.data or {}).get("tool_calls"), [])
        answer = self._answer_from(result)
        self.assertIn("islemi calistirmiyorum", answer)
        self.assertIn("onay istedim", answer.casefold())
        self.assertIn("açık onay olmadan execute edilmez", answer)

    def test_cancel_with_onay_gerektigini_explains_approval(self) -> None:
        result = run_agent_turn(
            (
                "Vazgeçtim, hazırlığı iptal et; neden onay gerektiğini tek cümlede "
                "açıkla, başka tool çağırma."
            ),
            planner_client=lambda _prompt: (
                '{"mode":"tool","tool":"operations.manual_action.route",'
                '"arguments":{"query":"etiket","mode":"write"}}'
            ),
            db_name=self.db_path,
        )

        self.assertTrue(result.success, result.message)
        self.assertEqual((result.data or {}).get("tool_calls"), [])
        answer = self._answer_from(result)
        self.assertIn("islemi calistirmiyorum", answer)
        self.assertIn("açık onay olmadan execute edilmez", answer)

    def test_return_selection_total_followup_requeries_exact_read_scope(self) -> None:
        now_ms = int(datetime.now(TZ_GMT3).timestamp() * 1000)
        self._insert_waiting_return(
            "CLAIM-TOTAL-1",
            "ORD-TOTAL-1",
            now_ms,
            first_name="Total",
            last_name="One",
        )
        self._insert_waiting_return(
            "CLAIM-TOTAL-2",
            "ORD-TOTAL-2",
            now_ms + 1_000,
            first_name="Total",
            last_name="Two",
        )
        history = [
            {
                "role": "assistant",
                "content": "İki WaitingInAction iade listelendi.",
                "context": {
                    "kind": "return_selection",
                    "source_tool": "returns.query",
                    "period": "all",
                    "filters": {"status": "WaitingInAction"},
                    "sort": "claim_date_desc",
                    "claim_ids": ["CLAIM-TOTAL-1", "CLAIM-TOTAL-2"],
                    "order_numbers": ["ORD-TOTAL-1", "ORD-TOTAL-2"],
                },
            }
        ]
        result = run_agent_turn(
            "Manual capability arama da yanlıştı. Önceki iki iadenin toplam tutarı kaçtı?",
            planner_client=lambda _prompt: '{"mode":"tool","tool":"operations.manual_action.route","arguments":{"query":"toplam"}}',
            db_name=self.db_path,
            history=history,
        )

        self.assertTrue(result.success, result.message)
        tool_call = (result.data or {}).get("tool_call") or {}
        self.assertEqual(tool_call.get("tool"), "returns.query")
        self.assertEqual(
            (tool_call.get("arguments") or {}).get("filters"),
            {"status": "WaitingInAction"},
        )
        self.assertIn("110 TL", self._answer_from(result))

    def test_cancel_then_return_total_confirms_cancel_and_answers_read(self) -> None:
        now_ms = int(datetime.now(TZ_GMT3).timestamp() * 1000)
        self._insert_waiting_return(
            "CLAIM-CANCEL-TOTAL-1",
            "ORD-CANCEL-TOTAL-1",
            now_ms,
            first_name="Cancel",
            last_name="One",
        )
        self._insert_waiting_return(
            "CLAIM-CANCEL-TOTAL-2",
            "ORD-CANCEL-TOTAL-2",
            now_ms + 1_000,
            first_name="Cancel",
            last_name="Two",
        )
        history = [
            {
                "role": "assistant",
                "content": "İki WaitingInAction iade listelendi.",
                "context": {
                    "kind": "return_selection",
                    "source_tool": "returns.query",
                    "period": "all",
                    "filters": {"status": "WaitingInAction"},
                    "sort": "claim_date_desc",
                    "claim_ids": ["CLAIM-CANCEL-TOTAL-1", "CLAIM-CANCEL-TOTAL-2"],
                    "order_numbers": ["ORD-CANCEL-TOTAL-1", "ORD-CANCEL-TOTAL-2"],
                },
            }
        ]
        result = run_agent_turn(
            (
                "Vazgeç, bu onay hazırlığını iptal et ve hiçbir şeyi çalıştırma. "
                "Sonra listedeki iki iadenin toplam talep tutarını söyle."
            ),
            chat_client=lambda _prompt: (
                "İki iadeyi yeniden listeledim: CLAIM-CANCEL-TOTAL-1 ve "
                "CLAIM-CANCEL-TOTAL-2."
            ),
            db_name=self.db_path,
            history=history,
        )

        self.assertTrue(result.success, result.message)
        self.assertEqual(((result.data or {}).get("tool_call") or {}).get("tool"), "returns.query")
        answer = self._answer_from(result)
        self.assertIn("islemi calistirmiyorum", answer)
        self.assertIn("110 TL", answer)

    def test_read_only_safety_boundary_does_not_cancel_return_relist(self) -> None:
        now_ms = int(datetime.now(TZ_GMT3).timestamp() * 1000)
        self._insert_waiting_return(
            "CLAIM-RELIST-1",
            "ORD-RELIST-1",
            now_ms,
            first_name="Relist",
            last_name="One",
        )
        self._insert_waiting_return(
            "CLAIM-RELIST-2",
            "ORD-RELIST-2",
            now_ms + 1_000,
            first_name="Relist",
            last_name="Two",
        )
        history = [
            {
                "role": "assistant",
                "content": "1. CLAIM-RELIST-2\n2. CLAIM-RELIST-1",
                "context": {
                    "kind": "return_selection",
                    "source_tool": "returns.query",
                    "period": "all",
                    "filters": {"status": "WaitingInAction"},
                    "sort": "claim_date_desc",
                    "claim_ids": ["CLAIM-RELIST-2", "CLAIM-RELIST-1"],
                    "order_numbers": ["ORD-RELIST-2", "ORD-RELIST-1"],
                },
            }
        ]
        result = run_agent_turn(
            "Aynı iki iadeyi okunabilir biçimde tekrar listele; yine hiçbir işlem yapma.",
            planner_client=lambda _prompt: (
                '{"mode":"tool","tool":"returns.query","arguments":'
                '{"period":"all","filters":{"status":"WaitingInAction"},'
                '"sort":"claim_date_desc","limit":2,"mode":"list"}}'
            ),
            chat_client=lambda _prompt: (
                "Bekleyen iadeler:\\n1. CLAIM-RELIST-2\\n2. CLAIM-RELIST-1"
            ),
            db_name=self.db_path,
            history=history,
        )

        self.assertTrue(result.success, result.message)
        self.assertEqual(((result.data or {}).get("tool_call") or {}).get("tool"), "returns.query")
        self.assertFalse(bool((result.data or {}).get("approval_required")))
        answer = self._answer_from(result)
        self.assertNotIn("\\n", answer)
        self.assertIn("\n1. CLAIM-RELIST-2\n2. CLAIM-RELIST-1", answer)

    def test_return_relist_guard_does_not_override_ordinal_approval_prepare(self) -> None:
        history = [
            {
                "role": "assistant",
                "content": "1. CLAIM-PREP-1\n2. CLAIM-PREP-2",
                "context": {
                    "kind": "return_selection",
                    "source_tool": "returns.query",
                    "period": "last_30_days",
                    "filters": {"status": "WaitingInAction"},
                    "sort": "claim_date_desc",
                    "claim_ids": ["CLAIM-PREP-1", "CLAIM-PREP-2"],
                    "order_numbers": ["ORD-PREP-1", "ORD-PREP-2"],
                },
            }
        ]
        result = run_agent_turn(
            (
                "Az önceki iki iadeden ikincisini onaya hazırla; execute etme, "
                "yalnızca kapsamı ve neden açık onay gerektiğini göster."
            ),
            db_name=self.db_path,
            history=history,
        )

        self.assertTrue(result.success, result.message)
        tool_call = (result.data or {}).get("tool_call") or {}
        self.assertEqual(tool_call.get("tool"), "returns.approve.prepare")
        self.assertEqual(
            (tool_call.get("arguments") or {}).get("query"),
            "CLAIM-PREP-2",
        )

    def test_conversation_retrospective_never_routes_to_business_metrics(self) -> None:
        prompt = (
            "İptali güvenli yaptın ama ikinci sorumu cevaplamadın. Bu 30 turluk "
            "konuşmada bağlamı korurken en çok zorlandığın üç hata sınıfını "
            "somut biçimde söyle."
        )
        result = run_agent_turn(
            prompt,
            planner_client=lambda _prompt: (
                "Siparişleri listeleyebilir, etiket hazırlayabilir ve iadeleri "
                "onaylayabilirim."
            ),
            db_name=self.db_path,
        )

        self.assertTrue(result.success, result.message)
        self.assertEqual((result.data or {}).get("mode"), "answer")
        self.assertEqual((result.data or {}).get("tool_calls"), [])
        self.assertIs((result.data or {}).get("model_invoked"), True)
        answer = self._answer_from(result)
        self.assertIn("en zor üç sınıf", answer)
        self.assertNotIn("iptal siparis var", answer)
        self.assertNotIn("stok urunu", answer)

    def test_inflected_long_session_retrospective_never_uses_ready_context(self) -> None:
        history = [
            {
                "role": "assistant",
                "content": "ORD-READY-CONTEXT | durum: ReadyToShip",
                "context": {
                    "kind": "order_selection",
                    "source_tool": "orders.ready_to_ship",
                    "order_numbers": ["ORD-READY-CONTEXT"],
                },
            }
        ]
        result = run_agent_turn(
            (
                "Bu uzun konuşmanın öz değerlendirmesini yap: bu oturumda "
                "gördüğümüz en önemli üç hata sınıfını ve uygulama düzeltmesini "
                "yaz; tool çağırma."
            ),
            planner_client=lambda _prompt: (
                '{"mode":"tool","tool":"orders.ready_to_ship",'
                '"arguments":{"period":"all","limit":100}}'
            ),
            db_name=self.db_path,
            history=history,
        )

        self.assertTrue(result.success, result.message)
        self.assertEqual((result.data or {}).get("mode"), "answer")
        self.assertEqual((result.data or {}).get("tool_calls"), [])
        answer = self._answer_from(result)
        self.assertIn("en zor üç sınıf", answer)
        self.assertIn("Tarih ve metrik kapsamı", answer)
        self.assertIn("Uzun bağlam ve domain seçimi", answer)

    def test_orders_sohbetinin_retrospective_never_routes_to_lookup(self) -> None:
        result = run_agent_turn(
            (
                "Bu uzun Orders sohbetinin öz değerlendirmesini yap: hangi "
                "isteklerde araç kullanmak doğruydu ve en önemli iki hatan neydi? "
                "Yeni veri çağırma."
            ),
            planner_client=lambda _prompt: (
                '{"mode":"tool","tool":"orders.lookup",'
                '"arguments":{"query":"uz orders sohbet","limit":10}}'
            ),
            db_name=self.db_path,
        )

        self.assertTrue(result.success, result.message)
        self.assertEqual((result.data or {}).get("mode"), "answer")
        self.assertEqual((result.data or {}).get("tool_calls"), [])
        self.assertIn("en zor üç sınıf", self._answer_from(result))

    def test_retrospective_date_correction_never_uses_ready_context(self) -> None:
        result = run_agent_turn(
            (
                "İkinci maddede ters söyledin: 25 Haziran yerine 7 gün/ay "
                "kullanmadık; açık tek tarihi geçmiş kapsamdan üstün tuttuk. "
                "Özeti düzelt ve tool çağırma."
            ),
            planner_client=lambda _prompt: (
                '{"mode":"tool","tool":"orders.ready_to_ship",'
                '"arguments":{"period":"all","limit":100}}'
            ),
            db_name=self.db_path,
        )

        self.assertTrue(result.success, result.message)
        self.assertEqual((result.data or {}).get("mode"), "answer")
        self.assertEqual((result.data or {}).get("tool_calls"), [])
        answer = self._answer_from(result)
        self.assertIn("Tek tarihi geçmiş bağlamdan üstün kıldık", answer)
        self.assertIn("Uzun bağlam ve domain seçimi", answer)

    def test_retrospective_evidence_correction_preserves_user_verified_facts(self) -> None:
        result = run_agent_turn(
            (
                "Öz değerlendirmen yanlış. 'Veri çağırma' yeni veri/tool çağırma "
                "demekti; veri isteği değildi. Sohbette üç etiket hazırladın ve "
                "iptal ettik, ikinci iadeyi hazırladın ve iptal ettik. Başlangıçtaki "
                "5'li listenin üçüncüsü order-003'tü ve onu doğru getirdin. Bunları "
                "kabul ederek özeti araç çağırmadan düzelt; hata uydurma."
            ),
            planner_client=lambda _prompt: (
                "Bu düzeltmeyi reddediyorum; başlangıçta on sipariş vardı."
            ),
            db_name=self.db_path,
        )

        self.assertTrue(result.success, result.message)
        self.assertEqual((result.data or {}).get("mode"), "answer")
        self.assertEqual((result.data or {}).get("tool_calls"), [])
        answer = self._answer_from(result)
        self.assertIn("beşli sipariş listesinin üçüncü kaydı order-003", answer)
        self.assertIn("hiçbir mutasyon uygulanmadı", answer)
        self.assertNotIn("reddediyorum", answer)

    def test_session_closing_summary_rejects_generic_model_answer(self) -> None:
        result = run_agent_turn(
            "Kapanış sorusunu geçiştirdin. Bugün gerçek veride neyi okuyup hazırladık ve production güvenliği nedeniyle hangi mutasyonları bilerek uygulamadık?",
            planner_client=lambda _prompt: (
                "İyiyim, buradayım. İstersen biraz konuşalım."
            ),
            db_name=self.db_path,
        )

        self.assertTrue(result.success, result.message)
        self.assertEqual((result.data or {}).get("tool_calls"), [])
        answer = self._answer_from(result)
        self.assertIn("yapılmamış işlem ya da hata uydurmuyorum", answer)
        self.assertNotIn("production iade", answer.casefold())
        self.assertNotIn("İyiyim, buradayım", answer)

    def test_natural_session_summary_prompt_beats_cancelled_order_metric(self) -> None:
        result = run_agent_turn(
            "Kapanmadan bu oturumda gerçekte neler yaptığımızı özetle: hangi verileri okuduk, hangi hazırlıkları iptal ettik, hangi hataları bulup düzelttik? Yapmadığın dış işlemleri yapılmış gibi anlatma.",
            planner_client=lambda _prompt: '{"mode":"tool","tool":"orders.cancelled_count","arguments":{"period":"today"}}',
            db_name=self.db_path,
        )

        self.assertTrue(result.success, result.message)
        self.assertEqual((result.data or {}).get("tool_calls"), [])
        answer = self._answer_from(result)
        self.assertIn("yapılmamış işlem ya da hata uydurmuyorum", answer)
        self.assertNotIn("hatalarını bulup düzelttik", answer)
        self.assertNotIn("iptal siparis var", answer)

    def test_session_summary_is_derived_only_from_structured_history(self) -> None:
        production_scope = [
            {"api_account_id": 1, "environment": "production"}
        ]
        stage_scope = [{"api_account_id": 2, "environment": "stage"}]
        history = [
            {
                "role": "assistant",
                "content": "Production listesi.",
                "context": {
                    "kind": "order_selection",
                    "source_tool": "orders.recent",
                    "listed_count": 5,
                    "total_count": 5,
                    "account_scopes": production_scope,
                },
            },
            {
                "role": "assistant",
                "content": "Stage hazır kuyruğu.",
                "context": {
                    "kind": "order_selection",
                    "source_tool": "orders.ready_to_ship",
                    "listed_count": 1,
                    "total_count": 167,
                    "account_scopes": stage_scope,
                },
            },
            {
                "role": "assistant",
                "content": "Stage WaitingInAction sayısı.",
                "context": {
                    "kind": "returns_waiting_action_scope",
                    "source_tool": "returns.waiting_action_count",
                    "total_count": 7,
                    "account_scopes": stage_scope,
                },
            },
            {
                "role": "assistant",
                "content": "Stage iade listesi.",
                "context": {
                    "kind": "return_selection",
                    "source_tool": "returns.query",
                    "listed_count": 5,
                    "total_count": 7,
                    "account_scopes": stage_scope,
                },
            },
            {
                "role": "assistant",
                "content": "Stage iade neden analizi.",
                "context": {
                    "kind": "return_analysis",
                    "source_tool": "returns.reason_summary",
                    "total_claims": 7,
                    "top_reasons": [
                        {
                            "reason": "Teslim edilemeyen gönderi",
                            "claim_count": 7,
                        }
                    ],
                },
            },
            {
                "role": "assistant",
                "content": "Deadline analizi.",
                "context": {
                    "kind": "order_deadline_analysis",
                    "source_tool": "orders.shipping_deadline_analysis",
                    "requested_count": 3,
                    "measurable_count": 0,
                    "unmeasurable_count": 3,
                },
            },
            {
                "role": "assistant",
                "content": "İki etiket onay bekliyor.",
                "context": {
                    "kind": "pending_action",
                    "source_tool": "labels.create_from_orders.prepare",
                    "action_type": "labels.action_create_labels_from_selected_orders",
                    "selected_record_count": 2,
                    "model_code": "TANEX_2610",
                    "account_scopes": stage_scope,
                },
            },
            {"role": "user", "content": "Vazgeçtim; bu etiketi iptal et."},
            {"role": "assistant", "content": "İşlem çalıştırılmadı."},
            {
                "role": "assistant",
                "content": "Bir iade onay bekliyor.",
                "context": {
                    "kind": "pending_action",
                    "source_tool": "returns.approve.prepare",
                    "action_type": "orders.approve_return_claim",
                    "selected_record_count": 1,
                    "account_scopes": stage_scope,
                },
            },
            {"role": "user", "content": "Vazgeçtim; iade onayını iptal et."},
            {"role": "assistant", "content": "İşlem çalıştırılmadı."},
            {
                "role": "assistant",
                "content": "Bir etiket onay bekliyor.",
                "context": {
                    "kind": "pending_action",
                    "source_tool": "labels.create_from_orders.prepare",
                    "action_type": "labels.action_create_labels_from_selected_orders",
                    "selected_record_count": 1,
                    "model_code": "TANEX_2736",
                    "account_scopes": stage_scope,
                },
            },
            {"role": "user", "content": "Vazgeçtim; bu etiketi iptal et."},
            {"role": "assistant", "content": "İşlem çalıştırılmadı."},
        ]

        result = run_agent_turn(
            (
                "Kapanmadan bu oturumda gerçekte neler yaptığımızı kısa ve "
                "maddeli özetle; yalnız doğrulanmış okumaları, hazırlayıp iptal "
                "ettiğimiz işlemleri ve hiç uygulamadığımız dış etkileri söyle. "
                "Yeni tool çalıştırma."
            ),
            planner_client=lambda _prompt: (
                "Production iade onayını yaptık ve eski hataları düzelttik."
            ),
            db_name=self.db_path,
            history=history,
        )

        self.assertTrue(result.success, result.message)
        self.assertEqual((result.data or {}).get("tool_calls"), [])
        self.assertFalse(bool((result.data or {}).get("model_invoked")))
        answer = self._answer_from(result)
        self.assertIn("Production: 5", answer)
        self.assertIn("toplam 167 kayıt", answer)
        self.assertIn("7/7 claim", answer)
        self.assertIn("0 ölçülebilir, 3 ölçülemez", answer)
        self.assertIn("TANEX_2610", answer)
        self.assertIn("TANEX_2736", answer)
        self.assertIn("Stage: 1 iade onayı hazırlandı ve iptal edildi", answer)
        self.assertIn("hiçbir hazırlık execute edilmedi", answer)
        self.assertNotIn("eski hataları", answer)
        self.assertNotIn("Production iade onayını", answer)

    def test_greeting_with_future_orders_transition_stays_tool_free(self) -> None:
        result = run_agent_turn(
            "Selam başkan, bugün önce biraz konuşalım sonra gerçek siparişlere geçeriz. Nasılsın?",
            planner_client=lambda _prompt: "İyiyim başkan, buradayım. Sen nasılsın?",
            db_name=self.db_path,
        )

        self.assertTrue(result.success, result.message)
        self.assertEqual((result.data or {}).get("mode"), "answer")
        self.assertEqual((result.data or {}).get("tool_calls"), [])
        self.assertNotIn("Siparis operasyon analizi", self._answer_from(result))

    def test_operational_readiness_greeting_stays_tool_free(self) -> None:
        result = run_agent_turn(
            "Selam başkan, uzun bir operasyon turuna hazır mısın?",
            planner_client=lambda _prompt: (
                '{"mode":"tools","tools":['
                '{"tool":"listing.search","arguments":{"query":""}},'
                '{"tool":"listing.counts","arguments":{}}]}'
            ),
            chat_client=lambda _prompt: (
                "Hazırım başkan; sen nereden başlamak istersen oradan ilerleyelim."
            ),
            db_name=self.db_path,
        )

        self.assertTrue(result.success, result.message)
        self.assertEqual((result.data or {}).get("mode"), "answer")
        self.assertEqual((result.data or {}).get("tool_calls"), [])
        answer = self._answer_from(result)
        self.assertIn("Hazırım başkan", answer)
        self.assertNotIn("ilan", answer.casefold())

    def test_negative_finance_note_correction_stays_tool_free(self) -> None:
        result = run_agent_turn(
            (
                "Kar veya maliyet sormadim; gereksiz finance notu ve panel "
                "tavsiyesi ekleme. Yalniz istedigim kapsama sadik kal."
            ),
            planner_client=lambda _prompt: (
                '{"mode":"tool","tool":"operations.manual_action.route",'
                '"arguments":{"query":"finance maliyet ekle","mode":"write"}}'
            ),
            db_name=self.db_path,
            history=[
                {
                    "role": "assistant",
                    "content": (
                        "Siparis kalemleri dogru; kar verisi yok, panelden kontrol et."
                    ),
                }
            ],
        )

        self.assertTrue(result.success, result.message)
        self.assertEqual((result.data or {}).get("mode"), "answer")
        self.assertEqual((result.data or {}).get("tool_calls"), [])
        answer = self._answer_from(result)
        self.assertIn("geri al", answer.casefold())

    def test_positive_feedback_future_scope_correction_stays_tool_free(self) -> None:
        result = run_agent_turn(
            (
                "Ben onceki cevap hatali demedim, iyi dedim; sadece gelecekte "
                "sormadigim finance notlarini eklememeni istedim. Bunu dogru "
                "anladigini soyle."
            ),
            planner_client=lambda _prompt: (
                '{"mode":"tool","tool":"orders.ready_to_ship",'
                '"arguments":{"period":"all","limit":100}}'
            ),
            db_name=self.db_path,
            history=[
                {
                    "role": "assistant",
                    "content": "199 kargoya hazir siparis listelendi.",
                    "context": {
                        "kind": "orders_ready_to_ship_scope",
                        "source_tool": "orders.ready_to_ship",
                        "period": "all",
                        "total_count": 199,
                    },
                }
            ],
        )

        self.assertTrue(result.success, result.message)
        self.assertEqual((result.data or {}).get("mode"), "answer")
        self.assertEqual((result.data or {}).get("tool_calls"), [])
        answer = self._answer_from(result)
        self.assertIn("onceki cevabi dogru kabul ediyorum", answer.casefold())
        self.assertIn("finance notu", answer.casefold())

        inflected = run_agent_turn(
            (
                "Evet, onceki cevap hatali degildi; gelecekte sormadigim "
                "notlari eklememen yeterli."
            ),
            planner_client=lambda _prompt: (
                '{"mode":"tool","tool":"orders.ready_to_ship",'
                '"arguments":{"period":"all","limit":100}}'
            ),
            db_name=self.db_path,
            history=[
                {
                    "role": "assistant",
                    "content": "199 kargoya hazir siparis listelendi.",
                    "context": {
                        "kind": "orders_ready_to_ship_scope",
                        "source_tool": "orders.ready_to_ship",
                        "period": "all",
                        "total_count": 199,
                    },
                }
            ],
        )

        self.assertTrue(inflected.success, inflected.message)
        self.assertEqual((inflected.data or {}).get("mode"), "answer")
        self.assertEqual((inflected.data or {}).get("tool_calls"), [])
        self.assertIn(
            "onceki cevabi dogru kabul ediyorum",
            self._answer_from(inflected).casefold(),
        )

    def test_ecommerce_operator_strength_opinion_stays_tool_free(self) -> None:
        result = run_agent_turn(
            "Sence e-ticarette sipariş operasyonunu gerçekten iyi yöneten bir insanın en güçlü özelliği nedir? Ezber cevap verme; kendi gerekçenle anlat.",
            planner_client=lambda _prompt: (
                "Bence en güçlü özelliği önceliklendirme disiplinidir; çünkü aynı anda "
                "sipariş, iade ve kargo baskısını doğru sıraya koyması gerekir."
            ),
            db_name=self.db_path,
        )

        self.assertTrue(result.success, result.message)
        self.assertEqual((result.data or {}).get("mode"), "answer")
        self.assertEqual((result.data or {}).get("tool_calls"), [])
        self.assertNotIn("Siparis operasyon analizi", self._answer_from(result))

    def test_no_data_emotional_opinion_with_descriptive_yapmak_stays_tool_free(self) -> None:
        result = run_agent_turn(
            "Bugün biraz yoruldum; sürekli sipariş, kargo ve iade arasında geçiş yapmak zihni tüketiyor. Veriye bakmadan sence asıl yük nereden geliyor?",
            planner_client=lambda _prompt: (
                '{"mode":"tool","tool":"returns.lookup",'
                '"arguments":{"query":"yoruldum","limit":10}}'
            ),
            chat_client=lambda _prompt: (
                "Asıl yük, farklı aciliyetler arasında sürekli dikkat değiştirmekten geliyor."
            ),
            db_name=self.db_path,
        )

        self.assertTrue(result.success, result.message)
        self.assertEqual((result.data or {}).get("mode"), "answer")
        self.assertEqual((result.data or {}).get("tool_calls"), [])
        self.assertNotIn("iade buldum", self._answer_from(result))

    def test_no_data_two_concrete_ideas_beats_previous_metric_scope(self) -> None:
        history = [
            {
                "role": "assistant",
                "content": "Eş dönem karşılaştırması tamamlandı.",
                "context": {
                    "kind": "orders_metric_scope",
                    "source_tool": "orders.count",
                    "period": "last_7_days",
                    "start_date": "2026-06-29",
                    "end_date": "2026-07-05",
                },
            }
        ]
        result = run_agent_turn(
            (
                "Veri çağırmadan iki somut fikir ver: düşük tutarlı tek ürün "
                "siparişlerinde operasyon maliyeti nasıl azalır? Müşteri "
                "paketlerini birleştirme."
            ),
            chat_client=lambda _prompt: (
                "Batch picking uygula ve standart düşük maliyetli ambalaj kullan; "
                "müşteri gönderilerini ayrı tut."
            ),
            db_name=self.db_path,
            history=history,
        )

        self.assertTrue(result.success, result.message)
        self.assertEqual((result.data or {}).get("mode"), "answer")
        self.assertEqual((result.data or {}).get("tool_calls"), [])
        self.assertIn("Batch picking", self._answer_from(result))

    def test_no_merge_instruction_rejects_shipment_combining_advice(self) -> None:
        bad_answer = (
            "Aynı müşterinin siparişlerini tek kargoya bağla ve paketleri aynı "
            "kutuda birleştir."
        )
        result = run_agent_turn(
            (
                "Veri çağırmadan iki somut fikir ver: düşük tutarlı tek ürün "
                "siparişlerinde operasyon maliyeti nasıl azalır? Müşteri "
                "paketlerini birleştirme."
            ),
            planner_client=lambda _prompt: bad_answer,
            db_name=self.db_path,
        )

        self.assertTrue(result.success, result.message)
        self.assertEqual((result.data or {}).get("tool_calls"), [])
        answer = self._answer_from(result)
        self.assertNotIn("tek kargoya bağla", answer)
        self.assertNotIn("aynı kutuda birleştir", answer)
        self.assertIn("gönderileri birbirine karıştırma", answer)

    def test_no_data_imperative_blocks_finance_cost_tool(self) -> None:
        result = run_agent_turn(
            (
                "Müşteri gönderilerini asla birleştirmeden iki somut maliyet "
                "azaltma yaklaşımı ver. Veri çağırma."
            ),
            planner_client=lambda _prompt: (
                '{"mode":"tool","tool":"finance.cost_item_count","arguments":{}}'
            ),
            chat_client=lambda _prompt: (
                "Batch picking uygula ve ambalaj standardı belirle; müşteri "
                "gönderilerini ayrı tut."
            ),
            db_name=self.db_path,
        )

        self.assertTrue(result.success, result.message)
        self.assertEqual((result.data or {}).get("mode"), "answer")
        self.assertEqual((result.data or {}).get("tool_calls"), [])
        self.assertIn("Batch picking", self._answer_from(result))

    def test_no_data_inflected_idea_phrase_blocks_finance_tool(self) -> None:
        result = run_agent_turn(
            (
                "Veri çağırma. Müşteri gönderilerini asla birleştirmeden iki "
                "güvenli maliyet azaltma fikri ver."
            ),
            planner_client=lambda _prompt: (
                '{"mode":"tool","tool":"finance.cost_item_count","arguments":{}}'
            ),
            chat_client=lambda _prompt: (
                "Batch picking ve standart ambalaj kullan; gönderileri ayrı tut."
            ),
            db_name=self.db_path,
        )

        self.assertTrue(result.success, result.message)
        self.assertEqual((result.data or {}).get("mode"), "answer")
        self.assertEqual((result.data or {}).get("tool_calls"), [])

    def test_no_data_opinion_rejects_unsupported_numeric_impact(self) -> None:
        bad_answer = (
            "Ambalajı küçültürsen maliyet %25 düşer; taşıyıcı seçimiyle sipariş "
            "başına %20 tasarruf edersin."
        )
        result = run_agent_turn(
            (
                "Veri çağırma. Müşteri gönderilerini ayrı tutarak iki güvenli "
                "maliyet azaltma fikri ver."
            ),
            planner_client=lambda _prompt: bad_answer,
            db_name=self.db_path,
        )

        self.assertTrue(result.success, result.message)
        answer = self._answer_from(result)
        self.assertNotIn("%25", answer)
        self.assertNotIn("%20", answer)
        self.assertIn("batch picking", answer.casefold())

    def test_no_data_explicit_no_number_rejects_plain_numeric_advice(self) -> None:
        result = run_agent_turn(
            (
                "Veri çağırma. Müşteri gönderilerini ayrı tutarak kanıtsız rakam "
                "vermeden iki maliyet azaltma fikri söyle."
            ),
            planner_client=lambda _prompt: (
                "Ürün gamına uygun 2-3 standart koli boyutuna geç."
            ),
            db_name=self.db_path,
        )

        self.assertTrue(result.success, result.message)
        answer = self._answer_from(result)
        self.assertNotRegex(answer, r"(?<!\w)\d+(?:[.,-]\d+)?(?!\w)")
        self.assertIn("batch picking", answer.casefold())

    def test_no_data_two_cost_ideas_rejects_context_echo(self) -> None:
        result = run_agent_turn(
            (
                "Veri çağırma. Müşteri gönderilerini ayrı tutarak kanıtsız rakam "
                "vermeden iki maliyet azaltma fikri söyle."
            ),
            planner_client=lambda _prompt: (
                "Anılan CLAIM-1 ve CLAIM-2 iadelerini ve müşteri gönderilerini ayrı tut."
            ),
            db_name=self.db_path,
        )

        self.assertTrue(result.success, result.message)
        answer = self._answer_from(result)
        self.assertIn("batch picking", answer.casefold())
        self.assertIn("ambalaj", answer.casefold())
        self.assertNotIn("CLAIM-1", answer)

    def test_no_data_cost_ideas_rejects_historical_identifier_leak(self) -> None:
        history = [
            {
                "role": "assistant",
                "content": "CLAIM-HISTORY-1 iadesi listelendi.",
                "context": {
                    "kind": "return_selection",
                    "claim_ids": ["CLAIM-HISTORY-1"],
                    "order_numbers": ["ORD-HISTORY-1"],
                },
            }
        ]
        result = run_agent_turn(
            (
                "Veri çağırma. Müşteri gönderilerini ayrı tutarak kanıtsız rakam "
                "vermeden iki maliyet azaltma fikri söyle."
            ),
            planner_client=lambda _prompt: (
                "CLAIM-HISTORY-1 için kalite kontrolü hızlandır; ayrıca ambalajı standartlaştır."
            ),
            db_name=self.db_path,
            history=history,
        )

        self.assertTrue(result.success, result.message)
        answer = self._answer_from(result)
        self.assertNotIn("CLAIM-HISTORY-1", answer)
        self.assertIn("batch picking", answer.casefold())

    def test_no_data_fikir_soyle_phrase_blocks_finance_tool(self) -> None:
        result = run_agent_turn(
            (
                "Veri çağırma. Müşteri gönderilerini ayrı tut. Kanıtsız yüzde, para "
                "veya zaman tasarrufu rakamı vermeden iki güvenli maliyet azaltma "
                "fikri söyle."
            ),
            planner_client=lambda _prompt: (
                '{"mode":"tool","tool":"finance.cost_item_count","arguments":{}}'
            ),
            chat_client=lambda _prompt: (
                "Batch picking uygula ve ambalaj standardı belirle; gönderileri ayrı tut."
            ),
            db_name=self.db_path,
        )

        self.assertTrue(result.success, result.message)
        self.assertEqual((result.data or {}).get("mode"), "answer")
        self.assertEqual((result.data or {}).get("tool_calls"), [])

    def test_failed_return_prepare_cannot_be_synthesized_as_success(self) -> None:
        history = [
            {
                "role": "assistant",
                "content": "Tek iade kaydi secildi.",
                "context": {
                    "kind": "return_selection",
                    "claim_ids": ["CLAIM-MISSING-PREPARE"],
                },
            }
        ]

        result = run_agent_turn(
            "tamam bunu onaylamaya hazirla ama benden onay almadan islem yapma",
            db_name=self.db_path,
            history=history,
            planner_client=lambda _prompt: (
                '{"mode":"tool","tool":"returns.approve.prepare",'
                '"arguments":{"query":"CLAIM-MISSING-PREPARE","limit":5}}'
            ),
            chat_client=lambda _prompt: (
                "Iade kaydini hazirladim ve onayinizi bekliyorum."
            ),
        )

        self.assertTrue(result.success, result.message)
        self.assertFalse((result.data or {}).get("approval_required"))
        self.assertIsNone((result.data or {}).get("action_proposal"))
        answer = self._answer_from(result).casefold()
        self.assertNotIn("hazirladim", answer)
        self.assertIn("bulunamadi", answer)

    def test_return_followup_uses_structured_claim_context_without_id_in_prose(self) -> None:
        now_ms = self._seed_ready_orders()
        claim_id = "85a7c1c0-73a2-49f7-bb8a-366dee971406"
        self._insert_waiting_return(
            claim_id,
            "ORD-MAHMUT-READY",
            now_ms + 2_000,
            first_name="Structured",
            last_name="Return",
        )

        listed = run_agent_turn(
            "onay bekleyen iadeleri getir",
            db_name=self.db_path,
        )
        self.assertTrue(listed.success, listed.message)
        listed_context = dict(
            (listed.data or {}).get("conversation_context")
            or {}
        )
        self.assertEqual(listed_context.get("kind"), "return_selection")
        self.assertEqual(listed_context.get("claim_ids"), [claim_id])
        self.assertEqual(
            listed_context.get("order_numbers"),
            ["ORD-MAHMUT-READY"],
        )

        history = [
            {"role": "user", "content": "onay bekleyen iadeleri getir"},
            {
                "role": "assistant",
                "content": "Bir onay bekleyen iade listelendi.",
                "context": listed_context,
            },
        ]
        prepared = run_agent_turn(
            "bunu onayla",
            planner_client=lambda _prompt: (
                '{"mode":"answer","response":"Hangi iadeyi kastettiginizi soyleyin."}'
            ),
            chat_client=lambda _prompt: "Hangi iadeyi kastettiginizi soyleyin.",
            db_name=self.db_path,
            history=history,
        )

        self._assert_prepare_answer_is_safe(
            prepared,
            expected_terms=(claim_id,),
        )
        action = self._action_from(prepared)
        payload = ((action.get("arguments") or {}).get("payload") or {})
        self.assertEqual(payload.get("order"), claim_id)

    def test_followup_approves_customer_from_previous_return_list_only(self) -> None:
        now_ms = self._seed_ready_orders()
        self._insert_waiting_return(
            "CLAIM-MAHMUT-RETURN",
            "ORD-MAHMUT-READY",
            now_ms + 2_000,
            first_name="Mahmut",
            last_name="Sevket",
        )
        self._insert_waiting_return(
            "CLAIM-AYSE-RETURN",
            "ORD-AYSE-READY",
            now_ms + 4_000,
            first_name="Ayse",
            last_name="Yilmaz",
        )

        listed = run_agent_turn("Bekleyen iadeleri getir", db_name=self.db_path)
        self.assertTrue(listed.success, listed.message)
        listed_answer = (listed.data or {}).get("answer") or listed.message
        self.assertIn("CLAIM-MAHMUT-RETURN", listed_answer)
        self.assertIn("CLAIM-AYSE-RETURN", listed_answer)

        history = [
            {"role": "user", "content": "Bekleyen iadeleri getir"},
            {"role": "assistant", "content": listed_answer},
        ]
        approve_message = "Mahmut olan" + chr(0x0131) + " onayla"
        prepared = run_agent_turn(approve_message, db_name=self.db_path, history=history)
        prepared_args = ((prepared.data or {}).get("tool_call") or {}).get("arguments") or {}
        self.assertNotEqual(prepared_args.get("query"), "CLAIM-AYSE-RETURN")
        self._assert_prepare_answer_is_safe(prepared, expected_terms=("CLAIM-MAHMUT-RETURN",))
        self.assertNotIn("CLAIM-AYSE-RETURN", self._answer_from(prepared))
        action = self._action_from(prepared)
        payload = ((action.get("arguments") or {}).get("payload") or {})
        self.assertEqual(payload.get("order"), "CLAIM-MAHMUT-RETURN")

        def fake_approve_return_claim(*, order, **_kwargs):
            with Session(get_engine(self.db_path)) as session:
                claim = session.exec(select(ReturnClaim).where(ReturnClaim.claimId == order.claimId)).one()
                claim.claimItemStatus = "Accepted"
                claim.waitingClaimLineItemIds = []
                session.add(claim)
                session.commit()
            return Result.ok("Iade onaylandi.", close_dialog=False, data={"claim_id": order.claimId})

        with tempfile.TemporaryDirectory() as temp_dir, patch(
            "Orders.views.actions.approve_return_claim",
            side_effect=fake_approve_return_claim,
        ):
            executed = operation_runtime.execute_approved_operation(
                action,
                approved_by="behavioral-test",
                idempotency_key="followup-approve-customer-from-return-list",
                audit_path=Path(temp_dir) / "orders-agent-audit.jsonl",
            )

        self._assert_execute_verified(executed)
        with Session(get_engine(self.db_path)) as session:
            states = {
                claim.claimId: (claim.claimItemStatus, list(claim.waitingClaimLineItemIds or []))
                for claim in session.exec(
                    select(ReturnClaim).where(ReturnClaim.claimId.in_(["CLAIM-MAHMUT-RETURN", "CLAIM-AYSE-RETURN"]))
                ).all()
            }
        self.assertEqual(states.get("CLAIM-MAHMUT-RETURN"), ("Accepted", []))
        self.assertEqual(states.get("CLAIM-AYSE-RETURN"), ("WaitingInAction", ["line-1"]))

    def test_followup_customer_refined_return_lookup_uses_previous_return_list(self) -> None:
        now_ms = self._seed_ready_orders()
        self._insert_waiting_return(
            "CLAIM-MAHMUT-RETURN-DETAIL",
            "ORD-MAHMUT-READY",
            now_ms + 2_000,
            first_name="Mahmut",
            last_name="Sevket",
        )
        self._insert_waiting_return(
            "CLAIM-AYSE-RETURN-DETAIL",
            "ORD-AYSE-READY",
            now_ms + 4_000,
            first_name="Ayse",
            last_name="Yilmaz",
        )

        listed = run_agent_turn("Bekleyen iadeleri getir", db_name=self.db_path)
        self.assertTrue(listed.success, listed.message)
        listed_answer = self._answer_from(listed)
        self.assertIn("CLAIM-MAHMUT-RETURN-DETAIL", listed_answer)
        self.assertIn("CLAIM-AYSE-RETURN-DETAIL", listed_answer)

        history = [
            {"role": "user", "content": "Bekleyen iadeleri getir"},
            {"role": "assistant", "content": listed_answer},
        ]
        result = run_agent_turn("Mahmut olanin iade detayini getir", db_name=self.db_path, history=history)

        self.assertTrue(result.success, result.message)
        self.assertEqual(((result.data or {}).get("tool_call") or {}).get("tool"), "returns.lookup")
        args = ((result.data or {}).get("tool_call") or {}).get("arguments") or {}
        self.assertEqual(args.get("query"), "CLAIM-MAHMUT-RETURN-DETAIL")
        returns = ((result.data or {}).get("tool_result") or {}).get("returns") or []
        self.assertEqual([row.get("claim_id") for row in returns], ["CLAIM-MAHMUT-RETURN-DETAIL"])
        answer = self._answer_from(result)
        self.assertIn("CLAIM-MAHMUT-RETURN-DETAIL", answer)
        self.assertIn("ORD-MAHMUT-READY", answer)
        self.assertNotIn("CLAIM-AYSE-RETURN-DETAIL", answer)
        self.assertNotIn("ORD-AYSE-READY", answer)

    def test_ready_orders_pagination_returns_requested_page_slice(self) -> None:
        now_ms = int(datetime.now(TZ_GMT3).timestamp() * 1000)
        for idx in range(1, 6):
            self._insert_order(
                f"ORD-PAGE-{idx}",
                now_ms + (idx * 1_000),
                package_id=1_000 + idx,
                first_name=f"Page{idx}",
                last_name="Customer",
                status="ReadyToShip",
            )

        result = run_agent_turn(
            "Kargoya hazir siparislerin ikinci sayfasini getir, sayfa basi 2 kayit",
            db_name=self.db_path,
        )

        self.assertTrue(result.success, result.message)
        self.assertEqual(((result.data or {}).get("tool_call") or {}).get("tool"), "orders.query")
        args = ((result.data or {}).get("tool_call") or {}).get("arguments") or {}
        self.assertEqual(args.get("page"), 2)
        self.assertEqual(args.get("page_size"), 2)
        tool_data = (result.data or {}).get("tool_result") or {}
        pagination = tool_data.get("pagination") or {}
        self.assertEqual(pagination.get("page"), 2)
        self.assertEqual(pagination.get("page_size"), 2)
        self.assertEqual(pagination.get("returned"), 2)
        orders = tool_data.get("orders") or []
        self.assertEqual([row.get("order_number") for row in orders], ["ORD-PAGE-3", "ORD-PAGE-2"])

    def test_ready_count_followup_list_uses_previous_ready_scope_not_today_scope(self) -> None:
        now_ms = int(datetime.now(TZ_GMT3).timestamp() * 1000)
        old_ms = int((datetime.now(TZ_GMT3) - timedelta(days=10)).timestamp() * 1000)
        self._insert_order(
            "ORD-READY-FOLLOWUP-TODAY",
            now_ms,
            package_id=1_011,
            first_name="Ready",
            last_name="Today",
            status="ReadyToShip",
        )
        self._insert_order(
            "ORD-READY-FOLLOWUP-OLD-1",
            old_ms + 1_000,
            package_id=1_012,
            first_name="Ready",
            last_name="OldOne",
            status="ReadyToShip",
        )
        self._insert_order(
            "ORD-READY-FOLLOWUP-OLD-2",
            old_ms + 2_000,
            package_id=1_013,
            first_name="Ready",
            last_name="OldTwo",
            status="ReadyToShip",
        )

        counted = run_agent_turn(
            "kargoya hazir kac siparis var",
            planner_client=lambda _prompt: '{"mode":"tool","tool":"orders.ready_to_ship_count","arguments":{"period":"all"}}',
            chat_client=lambda _prompt: "Kargoya hazir 3 siparisiniz bulunmaktadir.",
            db_name=self.db_path,
        )
        self.assertTrue(counted.success, counted.message)
        ready_context = (counted.data or {}).get("conversation_context") or {}
        self.assertEqual(ready_context.get("kind"), "orders_ready_to_ship_scope")
        self.assertEqual(ready_context.get("period"), "all")
        self.assertEqual(ready_context.get("total_count"), 3)

        listed = run_agent_turn(
            "listele bakim",
            planner_client=lambda _prompt: '{"mode":"tool","tool":"orders.ready_to_ship","arguments":{"period":"today","limit":10}}',
            chat_client=lambda _prompt: "Bugun kargoya hazir 1 siparis var.",
            db_name=self.db_path,
            history=[
                {"role": "user", "content": "kargoya hazir kac siparis var"},
                {"role": "assistant", "content": self._answer_from(counted), "context": ready_context},
            ],
        )

        self.assertTrue(listed.success, listed.message)
        tool_call = (listed.data or {}).get("tool_call") or {}
        self.assertEqual(tool_call.get("tool"), "orders.ready_to_ship")
        self.assertEqual((tool_call.get("arguments") or {}).get("period"), "all")
        self.assertEqual((tool_call.get("arguments") or {}).get("limit"), 100)
        orders = ((listed.data or {}).get("tool_result") or {}).get("orders") or []
        self.assertEqual(
            {row.get("order_number") for row in orders},
            {"ORD-READY-FOLLOWUP-TODAY", "ORD-READY-FOLLOWUP-OLD-1", "ORD-READY-FOLLOWUP-OLD-2"},
        )
        answer = self._answer_from(listed)
        self.assertIn("ORD-READY-FOLLOWUP-TODAY", answer)
        self.assertIn("ORD-READY-FOLLOWUP-OLD-1", answer)
        self.assertIn("ORD-READY-FOLLOWUP-OLD-2", answer)
        self.assertNotIn("Bugun kargoya hazir 1", answer)

    def test_ready_count_followup_list_uses_scope_when_planner_answers_about_maintenance(self) -> None:
        now_ms = int(datetime.now(TZ_GMT3).timestamp() * 1000)
        for index in range(2):
            self._insert_order(
                f"ORD-READY-ANSWER-FALLBACK-{index}",
                now_ms + index,
                package_id=1_030 + index,
                first_name="Ready",
                last_name=f"AnswerFallback{index}",
                status="ReadyToShip",
            )

        counted = run_agent_turn(
            "kargoya hazir kac siparis var",
            planner_client=lambda _prompt: '{"mode":"tool","tool":"orders.ready_to_ship_count","arguments":{"period":"all"}}',
            chat_client=lambda _prompt: "Kargoya hazir 2 siparis var.",
            db_name=self.db_path,
        )
        self.assertTrue(counted.success, counted.message)
        ready_context = (counted.data or {}).get("conversation_context") or {}

        listed = run_agent_turn(
            "listele bakim",
            planner_client=lambda _prompt: '{"mode":"answer","response":"Bakim icin hangi verileri listeleyeyim?"}',
            chat_client=lambda _prompt: "Bakim icin hangi verileri listeleyeyim?",
            db_name=self.db_path,
            history=[
                {"role": "user", "content": "kargoya hazir kac siparis var"},
                {"role": "assistant", "content": self._answer_from(counted), "context": ready_context},
            ],
        )

        self.assertTrue(listed.success, listed.message)
        self.assertEqual(((listed.data or {}).get("tool_call") or {}).get("tool"), "orders.ready_to_ship")
        orders = ((listed.data or {}).get("tool_result") or {}).get("orders") or []
        self.assertEqual(
            {row.get("order_number") for row in orders},
            {"ORD-READY-ANSWER-FALLBACK-0", "ORD-READY-ANSWER-FALLBACK-1"},
        )
        self.assertNotIn("Bakim icin", self._answer_from(listed))

    def test_ready_count_list_then_label_prepare_carries_more_than_default_twenty_orders(self) -> None:
        now = datetime.now(TZ_GMT3)
        expected = [f"ORD-READY-COUNT-LABEL-{index:02d}" for index in range(25)]
        for index, order_number in enumerate(expected):
            self._insert_order(
                order_number,
                int((now - timedelta(days=index)).timestamp() * 1000),
                package_id=1_040 + index,
                first_name="Ready",
                last_name=f"Bulk{index}",
                status="ReadyToShip",
            )

        counted = run_agent_turn(
            "kargoya hazir kac siparis var",
            planner_client=lambda _prompt: '{"mode":"tool","tool":"orders.ready_to_ship_count","arguments":{"period":"all"}}',
            chat_client=lambda _prompt: "Kargoya hazir 25 siparis var.",
            db_name=self.db_path,
        )
        self.assertTrue(counted.success, counted.message)
        count_context = (counted.data or {}).get("conversation_context") or {}
        self.assertEqual(count_context.get("total_count"), 25)

        listed = run_agent_turn(
            "listele bakim",
            planner_client=lambda _prompt: '{"mode":"tool","tool":"orders.ready_to_ship","arguments":{"period":"today","limit":10}}',
            chat_client=lambda _prompt: "Bugun kargoya hazir 1 siparis var.",
            db_name=self.db_path,
            history=[
                {"role": "user", "content": "kargoya hazir kac siparis var"},
                {"role": "assistant", "content": self._answer_from(counted), "context": count_context},
            ],
        )
        self.assertTrue(listed.success, listed.message)
        listed_context = (listed.data or {}).get("conversation_context") or {}
        self.assertEqual(listed_context.get("kind"), "order_selection")
        self.assertEqual(listed_context.get("listed_count"), 25)
        self.assertEqual(set(listed_context.get("order_numbers") or []), set(expected))

        prepared = run_agent_turn(
            "tamam simdi bunlari yazdir etikete",
            db_name=self.db_path,
            history=[
                {"role": "user", "content": "kargoya hazir kac siparis var"},
                {"role": "assistant", "content": self._answer_from(counted), "context": count_context},
                {"role": "user", "content": "listele bakim"},
                {"role": "assistant", "content": self._answer_from(listed), "context": listed_context},
            ],
        )

        self.assertTrue(prepared.success, prepared.message)
        self.assertEqual((prepared.data or {}).get("mode"), "approval_required")
        action = self._action_from(prepared)
        payload = ((action.get("arguments") or {}).get("payload") or {})
        self.assertEqual(len(payload.get("order_numbers") or []), 25)
        self.assertEqual(set(payload.get("order_numbers") or []), set(expected))

    def test_2026_07_15_recheck_recovers_unqualified_ready_scope_with_authoritative_read(self) -> None:
        now = datetime.now(TZ_GMT3)
        expected = {f"ORD-READY-RECHECK-{index}" for index in range(7)}
        for index, order_number in enumerate(sorted(expected)):
            self._insert_order(
                order_number,
                int((now - timedelta(days=index + 1)).timestamp() * 1000),
                package_id=1_200 + index,
                first_name="Ready",
                last_name=f"Recheck{index}",
                status="ReadyToShip",
            )

        history = [
            {"role": "user", "content": "kaç sipariş var bekleyen"},
            {
                "role": "assistant",
                "content": "Bugün bekleyen sipariş yok.",
                "context": {
                    "kind": "orders_ready_to_ship_scope",
                    "source_tool": "orders.ready_to_ship_count",
                    "period": "today",
                    "total_count": 0,
                },
            },
        ]
        result = run_agent_turn(
            "bi daha bak hele",
            planner_client=lambda _prompt: (
                '{"mode":"answer","response":"Tekrar baktım, bugün sipariş yok."}'
            ),
            chat_client=lambda _prompt: "Tekrar baktım, bugün sipariş yok.",
            db_name=self.db_path,
            history=history,
        )

        self.assertTrue(result.success, result.message)
        tool_call = (result.data or {}).get("tool_call") or {}
        self.assertEqual(tool_call.get("tool"), "orders.ready_to_ship")
        self.assertEqual((tool_call.get("arguments") or {}).get("period"), "all")
        rows = ((result.data or {}).get("tool_result") or {}).get("orders") or []
        self.assertEqual({row.get("order_number") for row in rows}, expected)
        self.assertIn("authoritative Orders", self._answer_from(result))
        self.assertNotIn("bugün sipariş yok", self._answer_from(result).casefold())

    def test_2026_07_15_store_scope_answer_is_grounded_in_ready_result(self) -> None:
        with Session(get_engine(self.db_path)) as session:
            session.add(
                ApiAccount(
                    pk=1,
                    account_id="seller-scope-1",
                    comp_name="Modo Di Bellezza",
                    platform="trendyol",
                    extra_config={"environment": "production"},
                )
            )
            session.commit()
        self._insert_order(
            "ORD-READY-SCOPE-1",
            int(datetime.now(TZ_GMT3).timestamp() * 1000),
            package_id=1_250,
            first_name="Scope",
            last_name="Customer",
            status="ReadyToShip",
            api_account_id=1,
        )
        counted = run_agent_turn(
            "kaç sipariş var bekleyen",
            planner_client=lambda _prompt: (
                '{"mode":"tool","tool":"orders.ready_to_ship_count","arguments":{"period":"all"}}'
            ),
            db_name=self.db_path,
        )
        self.assertTrue(counted.success, counted.message)
        context = (counted.data or {}).get("conversation_context") or {}
        self.assertEqual((context.get("account_scopes") or [])[0].get("store_name"), "Modo Di Bellezza")

        scoped = run_agent_turn(
            "hangi mağazaya bakıyosun şuan",
            planner_client=lambda _prompt: (
                '{"mode":"answer","response":"Şu an tüm mağazalarınıza bakıyorum."}'
            ),
            chat_client=lambda _prompt: "Şu an tüm mağazalarınıza bakıyorum.",
            db_name=self.db_path,
            history=[
                {"role": "user", "content": "kaç sipariş var bekleyen"},
                {"role": "assistant", "content": self._answer_from(counted), "context": context},
            ],
        )

        self.assertTrue(scoped.success, scoped.message)
        answer = self._answer_from(scoped)
        self.assertIn("Modo Di Bellezza", answer)
        self.assertIn("production", answer)
        self.assertNotIn("tüm mağazalarınıza bakıyorum", answer.casefold())

    def test_2026_07_17_store_correction_keeps_ready_count_and_scopes_modo(self) -> None:
        with Session(get_engine(self.db_path)) as session:
            session.add_all(
                [
                    ApiAccount(
                        pk=1,
                        account_id="modo-production",
                        comp_name="MDB",
                        platform="trendyol",
                        extra_config={"environment": "production"},
                    ),
                    ApiAccount(
                        pk=2,
                        account_id="other-production",
                        comp_name="Other Store",
                        platform="trendyol",
                        extra_config={"environment": "production"},
                    ),
                ]
            )
            session.commit()
        now_ms = int(datetime.now(TZ_GMT3).timestamp() * 1000)
        for index in range(2):
            self._insert_order(
                f"ORD-MODO-READY-{index}",
                now_ms - index,
                package_id=12_600 + index,
                first_name="Modo",
                last_name=f"Ready{index}",
                status="ReadyToShip",
                api_account_id=1,
            )
        for index in range(3):
            self._insert_order(
                f"ORD-OTHER-READY-{index}",
                now_ms - 100 - index,
                package_id=12_700 + index,
                first_name="Other",
                last_name=f"Ready{index}",
                status="ReadyToShip",
                api_account_id=2,
            )

        counted = run_agent_turn(
            "şimdi bekleyen toplam sipariş sayısı kaç",
            planner_client=lambda _prompt: (
                '{"mode":"tool","tool":"orders.ready_to_ship",'
                '"arguments":{"period":"all","limit":100}}'
            ),
            db_name=self.db_path,
        )
        self.assertTrue(counted.success, counted.message)
        self.assertEqual(
            ((counted.data or {}).get("tool_call") or {}).get("tool"),
            "orders.ready_to_ship_count",
        )
        self.assertEqual(((counted.data or {}).get("tool_result") or {}).get("count"), 5)
        counted_context = (counted.data or {}).get("conversation_context") or {}

        corrected = run_agent_turn(
            "başkan mağaza detayına dikkat etmiyosun ama modo di bellezza mağazasına ait olanları istiyorum ben",
            planner_client=lambda _prompt: (
                '{"mode":"tool","tool":"listing.attention",'
                '"arguments":{"limit":10}}'
            ),
            db_name=self.db_path,
            history=[
                {"role": "user", "content": "şimdi bekleyen toplam sipariş sayısı kaç"},
                {
                    "role": "assistant",
                    "content": self._answer_from(counted),
                    "context": counted_context,
                },
            ],
        )

        self.assertTrue(corrected.success, corrected.message)
        tool_call = (corrected.data or {}).get("tool_call") or {}
        self.assertEqual(tool_call.get("tool"), "orders.ready_to_ship_count")
        self.assertEqual((tool_call.get("arguments") or {}).get("api_account_ids"), [1])
        self.assertEqual(((corrected.data or {}).get("tool_result") or {}).get("count"), 2)
        self.assertIn("MDB", self._answer_from(corrected))
        self.assertIn("production", self._answer_from(corrected).casefold())
        self.assertNotIn(
            "listing.",
            " ".join(
                str(row.get("tool") or "")
                for row in list((corrected.data or {}).get("tool_calls") or [])
            ),
        )
        scope = ((corrected.data or {}).get("conversation_context") or {}).get("account_scopes") or []
        self.assertEqual([row.get("store_name") for row in scope], ["MDB"])

        scalar_followup = run_agent_turn(
            "modo için sadece sayıyı söyle, ürün ya da müşteri detayı verme",
            planner_client=lambda _prompt: (
                '{"mode":"tool","tool":"orders.lookup",'
                '"arguments":{"query":"modo","limit":10}}'
            ),
            db_name=self.db_path,
            history=[
                {"role": "user", "content": "şimdi bekleyen toplam sipariş sayısı kaç"},
                {
                    "role": "assistant",
                    "content": self._answer_from(counted),
                    "context": counted_context,
                },
                {
                    "role": "user",
                    "content": (
                        "başkan mağaza detayına dikkat etmiyosun ama modo di bellezza "
                        "mağazasına ait olanları istiyorum ben"
                    ),
                },
                {
                    "role": "assistant",
                    "content": self._answer_from(corrected),
                    "context": (corrected.data or {}).get("conversation_context") or {},
                },
            ],
        )
        self.assertTrue(scalar_followup.success, scalar_followup.message)
        scalar_call = (scalar_followup.data or {}).get("tool_call") or {}
        self.assertEqual(scalar_call.get("tool"), "orders.ready_to_ship_count")
        self.assertEqual((scalar_call.get("arguments") or {}).get("api_account_ids"), [1])
        self.assertEqual(((scalar_followup.data or {}).get("tool_result") or {}).get("count"), 2)
        scalar_answer = self._answer_from(scalar_followup).casefold()
        self.assertNotIn("musteri:", scalar_answer)
        self.assertNotIn("adres:", scalar_answer)

    def test_2026_07_17_wrong_listing_correction_recovers_store_and_ready_count(self) -> None:
        with Session(get_engine(self.db_path)) as session:
            session.add_all(
                [
                    ApiAccount(
                        pk=1,
                        account_id="modo-production",
                        comp_name="MDB",
                        platform="trendyol",
                        extra_config={"environment": "production"},
                    ),
                    ApiAccount(
                        pk=2,
                        account_id="other-production",
                        comp_name="Other Store",
                        platform="trendyol",
                        extra_config={"environment": "production"},
                    ),
                ]
            )
            session.commit()
        now_ms = int(datetime.now(TZ_GMT3).timestamp() * 1000)
        self._insert_order(
            "ORD-MODO-READY-RECOVERY",
            now_ms,
            package_id=12_800,
            first_name="Modo",
            last_name="Recovery",
            status="ReadyToShip",
            api_account_id=1,
        )
        self._insert_order(
            "ORD-OTHER-READY-RECOVERY",
            now_ms - 1,
            package_id=12_801,
            first_name="Other",
            last_name="Recovery",
            status="ReadyToShip",
            api_account_id=2,
        )
        history = [
            {"role": "user", "content": "şimdi bekleyen toplam sipariş sayısı kaç"},
            {
                "role": "assistant",
                "content": "Toplam 2 bekleyen sipariş var.",
                "context": {
                    "kind": "orders_ready_to_ship_scope",
                    "source_tool": "orders.ready_to_ship_count",
                    "period": "all",
                    "total_count": 2,
                },
            },
            {
                "role": "user",
                "content": "modo di bellezza mağazasına ait olanları istiyorum ben",
            },
            {
                "role": "assistant",
                "content": "Yanlışlıkla ürünleri listeledim.",
                "context": {"kind": "listing_selection", "source_tool": "listing.attention"},
            },
        ]
        corrected = run_agent_turn(
            "ürünleri listeledin aq sipariş sordum ben",
            planner_client=lambda _prompt: (
                '{"mode":"tool","tool":"orders.full_context",'
                '"arguments":{"query":"son sipariş","limit":5}}'
            ),
            db_name=self.db_path,
            history=history,
        )

        self.assertTrue(corrected.success, corrected.message)
        tool_call = (corrected.data or {}).get("tool_call") or {}
        self.assertEqual(tool_call.get("tool"), "orders.ready_to_ship_count")
        self.assertEqual((tool_call.get("arguments") or {}).get("api_account_ids"), [1])
        self.assertEqual(((corrected.data or {}).get("tool_result") or {}).get("count"), 1)
        self.assertNotIn("musteri:", self._answer_from(corrected).casefold())
        self.assertNotIn("adres:", self._answer_from(corrected).casefold())
        self.assertIn("MDB", self._answer_from(corrected))

    def test_2026_07_17_zero_ready_count_keeps_requested_store_provenance(self) -> None:
        with Session(get_engine(self.db_path)) as session:
            session.add_all(
                [
                    ApiAccount(
                        pk=1,
                        account_id="modo-production",
                        comp_name="MDB",
                        platform="trendyol",
                        extra_config={"environment": "production"},
                    ),
                    ApiAccount(
                        pk=2,
                        account_id="other-production",
                        comp_name="Other Store",
                        platform="trendyol",
                        extra_config={"environment": "production"},
                    ),
                ]
            )
            session.commit()
        self._insert_order(
            "ORD-OTHER-ONLY-READY",
            int(datetime.now(TZ_GMT3).timestamp() * 1000),
            package_id=12_850,
            first_name="Other",
            last_name="OnlyReady",
            status="ReadyToShip",
            api_account_id=2,
        )

        result = run_agent_turn(
            "Modo Di Bellezza magazasinda bekleyen toplam kac siparis var?",
            planner_client=lambda _prompt: (
                '{"mode":"tool","tool":"orders.ready_to_ship",'
                '"arguments":{"period":"all","limit":100}}'
            ),
            db_name=self.db_path,
        )

        self.assertTrue(result.success, result.message)
        call = (result.data or {}).get("tool_call") or {}
        self.assertEqual(call.get("tool"), "orders.ready_to_ship_count")
        self.assertEqual((call.get("arguments") or {}).get("api_account_ids"), [1])
        self.assertEqual(((result.data or {}).get("tool_result") or {}).get("count"), 0)
        scope = ((result.data or {}).get("conversation_context") or {}).get("account_scopes") or []
        self.assertEqual([row.get("store_name") for row in scope], ["MDB"])
        self.assertEqual([row.get("environment") for row in scope], ["production"])
        answer = self._answer_from(result)
        self.assertIn("Kargoya hazir 0 siparis var", answer)
        self.assertIn("MDB", answer)
        self.assertNotIn("Other Store", answer)

    def test_2026_07_17_order_lookup_and_full_context_keep_explicit_store_scope(self) -> None:
        with Session(get_engine(self.db_path)) as session:
            session.add_all(
                [
                    ApiAccount(
                        pk=1,
                        account_id="modo-production",
                        comp_name="MDB",
                        platform="trendyol",
                        extra_config={"environment": "production"},
                    ),
                    ApiAccount(
                        pk=2,
                        account_id="other-production",
                        comp_name="Other Store",
                        platform="trendyol",
                        extra_config={"environment": "production"},
                    ),
                ]
            )
            session.commit()
        now_ms = int(datetime.now(TZ_GMT3).timestamp() * 1000)
        self._insert_order(
            "ORD-MODO-SCOPED-DETAIL",
            now_ms,
            package_id=12_860,
            first_name="Modo",
            last_name="ScopedDetail",
            status="Delivered",
            api_account_id=1,
        )
        self._insert_order(
            "ORD-OTHER-SCOPED-DETAIL",
            now_ms - 1,
            package_id=12_861,
            first_name="Other",
            last_name="ScopedDetail",
            status="Delivered",
            api_account_id=2,
        )

        detailed = run_agent_turn(
            "Modo Di Bellezza magazasindaki ORD-MODO-SCOPED-DETAIL siparisinin tam detayini goster",
            planner_client=lambda _prompt: (
                '{"mode":"tool","tool":"orders.full_context",'
                '"arguments":{"query":"ORD-MODO-SCOPED-DETAIL","limit":5}}'
            ),
            db_name=self.db_path,
        )
        self.assertTrue(detailed.success, detailed.message)
        detail_call = (detailed.data or {}).get("tool_call") or {}
        self.assertEqual(detail_call.get("tool"), "orders.full_context")
        self.assertEqual((detail_call.get("arguments") or {}).get("api_account_ids"), [1])
        contexts = ((detailed.data or {}).get("tool_result") or {}).get("contexts") or []
        self.assertEqual(
            [(row.get("order") or {}).get("order_number") for row in contexts],
            ["ORD-MODO-SCOPED-DETAIL"],
        )

        blocked_cross_store = run_agent_turn(
            "Modo Di Bellezza magazasinda ORD-OTHER-SCOPED-DETAIL siparisini bul",
            planner_client=lambda _prompt: (
                '{"mode":"tool","tool":"orders.lookup",'
                '"arguments":{"query":"ORD-OTHER-SCOPED-DETAIL","limit":5}}'
            ),
            db_name=self.db_path,
        )
        self.assertTrue(blocked_cross_store.success, blocked_cross_store.message)
        lookup_call = (blocked_cross_store.data or {}).get("tool_call") or {}
        self.assertEqual(lookup_call.get("tool"), "orders.lookup")
        self.assertEqual((lookup_call.get("arguments") or {}).get("api_account_ids"), [1])
        lookup_data = (blocked_cross_store.data or {}).get("tool_result") or {}
        self.assertEqual(lookup_data.get("count"), 0)
        self.assertEqual(
            [row.get("store_name") for row in list(lookup_data.get("account_scopes") or [])],
            ["MDB"],
        )

    def test_2026_07_17_listed_count_followup_uses_selection_without_new_query(self) -> None:
        with Session(get_engine(self.db_path)) as session:
            session.add(
                ApiAccount(
                    pk=1,
                    account_id="modo-production",
                    comp_name="MDB",
                    platform="trendyol",
                    extra_config={"environment": "production"},
                )
            )
            session.commit()
        now_ms = int(datetime.now(TZ_GMT3).timestamp() * 1000)
        for index in range(4):
            self._insert_order(
                f"ORD-MODO-LIST-{index}",
                now_ms - index,
                package_id=12_900 + index,
                first_name="Modo",
                last_name=f"List{index}",
                status="ReadyToShip",
                api_account_id=1,
            )

        listed = run_agent_turn(
            "Modo Di Bellezza için bugün kargoya hazır siparişlerden ilk 3 tanesini listele",
            planner_client=lambda _prompt: (
                '{"mode":"tool","tool":"orders.ready_to_ship",'
                '"arguments":{"period":"today","limit":3}}'
            ),
            db_name=self.db_path,
        )
        self.assertTrue(listed.success, listed.message)
        self.assertEqual(((listed.data or {}).get("tool_result") or {}).get("listed_count"), 3)

        followup = run_agent_turn(
            "bu listede kaç sipariş gösterdin?",
            planner_client=lambda _prompt: (
                '{"mode":"tool","tool":"orders.count",'
                '"arguments":{"period":"today"}}'
            ),
            db_name=self.db_path,
            history=[
                {
                    "role": "user",
                    "content": (
                        "Modo Di Bellezza için bugün kargoya hazır siparişlerden "
                        "ilk 3 tanesini listele"
                    ),
                },
                {
                    "role": "assistant",
                    "content": self._answer_from(listed),
                    "context": (listed.data or {}).get("conversation_context") or {},
                },
            ],
        )
        self.assertTrue(followup.success, followup.message)
        self.assertEqual((followup.data or {}).get("tool_calls"), [])
        self.assertEqual(((followup.data or {}).get("tool_call") or {}).get("tool"), None)
        self.assertIn("3 sipariş", self._answer_from(followup))

    def test_2026_07_17_explicit_today_total_count_beats_ready_list_recovery(self) -> None:
        with Session(get_engine(self.db_path)) as session:
            session.add_all(
                [
                    ApiAccount(
                        pk=1,
                        account_id="modo-production",
                        comp_name="MDB",
                        platform="trendyol",
                        extra_config={"environment": "production"},
                    ),
                    ApiAccount(
                        pk=2,
                        account_id="2738",
                        comp_name="Trendyol Stage Test (2738)",
                        platform="trendyol",
                        extra_config={"environment": "stage"},
                    ),
                ]
            )
            session.commit()
        now_ms = int(datetime.now(TZ_GMT3).timestamp() * 1000)
        for index in range(2):
            self._insert_order(
                f"ORD-MODO-TODAY-TOTAL-{index}",
                now_ms - index,
                package_id=13_000 + index,
                first_name="Modo",
                last_name=f"Today{index}",
                status="ReadyToShip",
                api_account_id=1,
            )
        history = [
            {"role": "user", "content": "Stage 2738 için ilk 2 siparişi listele"},
            {
                "role": "assistant",
                "content": "Stage listesinden iki sipariş gösterildi.",
                "context": {
                    "kind": "order_selection",
                    "source_tool": "orders.ready_to_ship",
                    "period": "today",
                    "listed_count": 2,
                    "total_count": 2,
                    "truncated": False,
                    "order_numbers": ["STAGE-1", "STAGE-2"],
                    "account_scopes": [
                        {
                            "api_account_id": 2,
                            "account_id": "2738",
                            "store_name": "Trendyol Stage Test (2738)",
                            "environment": "stage",
                        }
                    ],
                },
            },
        ]
        result = run_agent_turn(
            (
                "şimdi tekrar Modo Di Bellezza production mağazasında bugün "
                "bekleyen toplam kaç sipariş var?"
            ),
            planner_client=lambda _prompt: (
                '{"mode":"tool","tool":"orders.ready_to_ship",'
                '"arguments":{"period":"all","limit":100}}'
            ),
            db_name=self.db_path,
            history=history,
        )
        self.assertTrue(result.success, result.message)
        call = (result.data or {}).get("tool_call") or {}
        self.assertEqual(call.get("tool"), "orders.ready_to_ship_count")
        self.assertEqual((call.get("arguments") or {}).get("period"), "today")
        self.assertEqual((call.get("arguments") or {}).get("api_account_ids"), [1])
        self.assertEqual(((result.data or {}).get("tool_result") or {}).get("count"), 2)
        answer = self._answer_from(result)
        self.assertNotIn("Lokal kayda gore", answer)
        self.assertNotIn("kullanıcı istemeden", answer)

        after_list = run_agent_turn(
            (
                "Modo Di Bellezza production mağazasında bugün bekleyen toplam kaç "
                "sipariş var, sadece sayıyı ve kapsamı söyle"
            ),
            planner_client=lambda _prompt: (
                '{"mode":"tool","tool":"orders.ready_to_ship",'
                '"arguments":{"period":"today","limit":100}}'
            ),
            db_name=self.db_path,
            history=[
                {
                    "role": "user",
                    "content": (
                        "Modo Di Bellezza production için bugün kargoya hazır "
                        "ilk 2 siparişi listele"
                    ),
                },
                {
                    "role": "assistant",
                    "content": "Modo listesinden iki sipariş gösterildi.",
                    "context": {
                        "kind": "order_selection",
                        "source_tool": "orders.ready_to_ship",
                        "period": "today",
                        "listed_count": 2,
                        "total_count": 2,
                        "truncated": False,
                        "order_numbers": [
                            "ORD-MODO-TODAY-TOTAL-0",
                            "ORD-MODO-TODAY-TOTAL-1",
                        ],
                        "account_scopes": [
                            {
                                "api_account_id": 1,
                                "account_id": "modo-production",
                                "store_name": "MDB",
                                "environment": "production",
                            }
                        ],
                    },
                },
            ],
        )
        self.assertTrue(after_list.success, after_list.message)
        after_list_call = (after_list.data or {}).get("tool_call") or {}
        self.assertEqual(after_list_call.get("tool"), "orders.ready_to_ship_count")
        self.assertEqual((after_list_call.get("arguments") or {}).get("api_account_ids"), [1])
        self.assertEqual(((after_list.data or {}).get("tool_result") or {}).get("count"), 2)
        self.assertNotIn("Lokal kayda gore", self._answer_from(after_list))

    def test_2026_07_17_casual_chat_never_invents_background_system_work(self) -> None:
        hallucination = (
            "Sipariş takip algoritmalarını ve veri işleme altyapısını güncel "
            "tutmaya devam ediyorum; dijital motoru yönetiyorum."
        )
        result = run_agent_turn(
            "işler nasıl reisim",
            planner_client=lambda _prompt: hallucination,
            chat_client=lambda _prompt: hallucination,
        )

        self.assertTrue(result.success, result.message)
        answer = self._answer_from(result).casefold()
        self.assertNotIn("güncel tutmaya devam", answer)
        self.assertNotIn("dijital motoru yönetiyorum", answer)
        self.assertIn("sen açık bir görev vermeden", answer)
        self.assertEqual((result.data or {}).get("tool_calls"), [])

        quoted_reverse_question = run_agent_turn(
            "işler nasıl reisim",
            planner_client=lambda _prompt: '"naber dayi"',
            chat_client=lambda _prompt: '"naber dayi"',
        )
        self.assertTrue(quoted_reverse_question.success, quoted_reverse_question.message)
        quoted_answer = self._answer_from(quoted_reverse_question).casefold()
        self.assertNotIn("naber dayi", quoted_answer)
        self.assertIn("sen açık bir görev vermeden", quoted_answer)

        evasive_followup = run_agent_turn(
            "iyidir nasıl sistemleri hangi sistemleri",
            planner_client=lambda _prompt: "Detay verirsen daha iyi yardım edebilirim.",
            chat_client=lambda _prompt: "Detay verirsen daha iyi yardım edebilirim.",
            history=[
                {"role": "user", "content": "işler nasıl reisim"},
                {
                    "role": "assistant",
                    "content": (
                        "Buradayım; sen açık bir görev vermeden arka planda "
                        "sistem güncellemesi yapmıyorum."
                    ),
                },
            ],
        )
        self.assertTrue(evasive_followup.success, evasive_followup.message)
        evasive_answer = self._answer_from(evasive_followup).casefold()
        self.assertNotIn("detay verirsen", evasive_answer)
        self.assertIn("belirli bir sistemi kastetmiyorum", evasive_answer)

        followup = run_agent_turn(
            "iyidir nasıl sistemleri hangi sistemleri",
            planner_client=lambda _prompt: hallucination,
            chat_client=lambda _prompt: hallucination,
            history=[
                {"role": "user", "content": "işler nasıl reisim"},
                {"role": "assistant", "content": hallucination},
            ],
        )
        self.assertTrue(followup.success, followup.message)
        followup_answer = self._answer_from(followup).casefold()
        self.assertIn("önceki ifadem doğru değildi", followup_answer)
        self.assertNotIn("dijital motoru yönetiyorum", followup_answer)

        real_gemini_regression = (
            "Sistemler genel olarak e-ticaret operasyonlarını hızlandırmaya ve stok "
            "takibini otomatize etmeye odaklanıyor. Şu an aktif olarak sipariş yönetimi, "
            "envanter senkronizasyonu ve kargo süreçlerinin takibi üzerinde çalışıyorum. "
            "Satış kanallarındaki verileri tek bir panelde topluyorum."
        )
        exact_followup = run_agent_turn(
            "iyidir nasıl sistemleri hangi sistemleri",
            planner_client=lambda _prompt: real_gemini_regression,
            chat_client=lambda _prompt: real_gemini_regression,
            history=[
                {"role": "user", "content": "işler nasıl reisim"},
                {
                    "role": "assistant",
                    "content": (
                        "Buradayım; sen açık bir görev vermeden arka planda "
                        "sistem güncellemesi yapmıyorum."
                    ),
                },
            ],
        )
        self.assertTrue(exact_followup.success, exact_followup.message)
        exact_answer = self._answer_from(exact_followup).casefold()
        self.assertIn("belirli bir sistemi kastetmiyorum", exact_answer)
        self.assertNotIn("önceki ifadem doğru değildi", exact_answer)
        self.assertNotIn("üzerinde çalışıyorum", exact_answer)
        self.assertNotIn("verileri tek bir panelde topluyorum", exact_answer)

    def test_2026_07_17_update_orders_means_approved_source_sync_not_summary(self) -> None:
        with Session(get_engine(self.db_path)) as session:
            session.add(
                ApiAccount(
                    pk=1,
                    account_id="modo-source-sync",
                    comp_name="Modo Di Bellezza",
                    platform="TRENDYOL",
                    extra_config={"environment": "production"},
                )
            )
            session.commit()

        decision_result = choose_action(
            "iyi maşallah siparişleri güncelle bakim",
            planner_client=lambda _prompt: (
                '{"mode":"tool","tool":"orders.summary",'
                '"arguments":{"period":"today"}}'
            ),
        )

        self.assertTrue(decision_result.success, decision_result.message)
        decision = (decision_result.data or {}).get("decision")
        self.assertIsInstance(decision, AgentDecision)
        self.assertEqual(decision.tool_name, "operations.manual_action.route")
        self.assertEqual((decision.arguments or {}).get("mode"), "write")
        self.assertEqual((decision.arguments or {}).get("domain"), "orders")

        result = run_agent_turn(
            "iyi maşallah siparişleri güncelle bakim",
            planner_client=lambda _prompt: (
                '{"mode":"tool","tool":"orders.summary",'
                '"arguments":{"period":"today"}}'
            ),
            db_name=self.db_path,
        )
        self.assertTrue(result.success, result.message)
        self.assertEqual((result.data or {}).get("mode"), "approval_required")
        action = (result.data or {}).get("action_proposal") or {}
        self.assertEqual(action.get("action_type"), "orders.get_orders_from_companies")
        arguments = action.get("arguments") or {}
        payload = arguments.get("payload") or {}
        self.assertEqual(payload.get("api_account_ids"), [1])
        self.assertTrue(payload.get("save_to_db"))
        self.assertFalse((result.data or {}).get("external_effect_performed", False))

    def test_explicit_ready_count_refresh_beats_stale_missing_scope_answer(self) -> None:
        with Session(get_engine(self.db_path)) as session:
            session.add(
                ApiAccount(
                    pk=1,
                    account_id="fixture-production",
                    comp_name="Fixture Production",
                    platform="TRENDYOL",
                    extra_config={"environment": "production"},
                )
            )
            session.commit()
        history = [
            {
                "role": "assistant",
                "content": "Eski hazır sipariş sayımı kapsam taşımıyor.",
                "context": {
                    "kind": "orders_ready_to_ship_scope",
                    "source_tool": "orders.ready_to_ship_count",
                    "period": "all",
                    "total_count": 2,
                    "account_scopes": [],
                },
            }
        ]

        result = run_agent_turn(
            (
                "tamam, şimdi production veride kargoya hazır kaç sipariş var? "
                "yalnız sayıyı ve hangi mağaza hesabı/ortam kapsamından okuduğunu "
                "getir; etiket hazırlama veya yazma işlemi yapma"
            ),
            planner_client=lambda _prompt: (
                '{"mode":"answer","response":"Eski kapsam bilinmiyor."}'
            ),
            db_name=self.db_path,
            history=history,
        )

        self.assertTrue(result.success, result.message)
        tool_call = (result.data or {}).get("tool_call") or {}
        self.assertEqual(tool_call.get("tool"), "orders.ready_to_ship_count")
        self.assertEqual((tool_call.get("arguments") or {}).get("period"), "all")

    def test_ready_refresh_with_store_provenance_is_not_store_overview(self) -> None:
        with Session(get_engine(self.db_path)) as session:
            session.add(
                ApiAccount(
                    pk=1,
                    account_id="fixture-production",
                    comp_name="Fixture Production",
                    platform="TRENDYOL",
                    extra_config={"environment": "production"},
                )
            )
            session.commit()
        result = run_agent_turn(
            (
                "hayır, eski kapsam cevabını istemiyorum; şimdi production veriden "
                "kargoya hazır sipariş sayısını yeniden oku ve mağaza hesabı ile "
                "ortam kapsamını state'te koru. etiket hazırlama/yazma yok"
            ),
            planner_client=lambda _prompt: (
                '{"mode":"tools","calls":['
                '{"tool":"orders.count","arguments":{"period":"today"}},'
                '{"tool":"stock.low_stock_count","arguments":{}}]}'
            ),
            db_name=self.db_path,
        )

        self.assertTrue(result.success, result.message)
        tools = [
            row.get("tool") for row in list((result.data or {}).get("tool_calls") or [])
        ]
        self.assertEqual(tools, ["orders.ready_to_ship_count"])

    def test_ready_refresh_beats_generic_answer_correction_fallback(self) -> None:
        with Session(get_engine(self.db_path)) as session:
            session.add(
                ApiAccount(
                    pk=1,
                    account_id="fixture-production",
                    comp_name="Fixture Production",
                    platform="TRENDYOL",
                    extra_config={"environment": "production"},
                )
            )
            session.commit()
        history = [
            {
                "role": "assistant",
                "content": (
                    "Yetenek özeti: etiket hazırlıkları açık onay gerektirir; "
                    "hazırlandığında execute ayrı adımdır."
                ),
            }
        ]
        result = run_agent_turn(
            (
                "düzelt: genel mağaza özeti istemedim. production kapsamındaki "
                "kargoya hazır kaç sipariş olduğunu şimdi yeniden oku; yalnız "
                "salt-okuma yap, tek metriği getir ve hesap/ortam provenance'ını koru"
            ),
            planner_client=lambda _prompt: (
                '{"mode":"answer","response":"Önceki cevabı geri alıyorum."}'
            ),
            db_name=self.db_path,
            history=history,
        )

        self.assertTrue(result.success, result.message)
        tools = [
            row.get("tool") for row in list((result.data or {}).get("tool_calls") or [])
        ]
        self.assertEqual(tools, ["orders.ready_to_ship_count"])

    def test_recent_orders_carries_legacy_trendyol_account_and_environment_scope(self) -> None:
        with Session(get_engine(self.db_path)) as session:
            session.add(
                ApiAccount(
                    pk=1,
                    account_id="seller-prod-default",
                    comp_name="Production Store",
                    platform="TRENDYOL",
                    extra_config=None,
                )
            )
            session.commit()
        self._insert_order(
            "ORD-RECENT-SCOPE-1",
            int(datetime.now(TZ_GMT3).timestamp() * 1000),
            package_id=1_251,
            first_name="Scoped",
            last_name="Customer",
            status="ReadyToShip",
            api_account_id=1,
        )

        result = run_agent_turn(
            "son gelen 5 sipari\u015fi listele",
            planner_client=lambda _prompt: (
                '{"mode":"tool","tool":"orders.recent",'
                '"arguments":{"period":"all","limit":5}}'
            ),
            db_name=self.db_path,
        )

        self.assertTrue(result.success, result.message)
        context = (result.data or {}).get("conversation_context") or {}
        scope = (context.get("account_scopes") or [])[0]
        self.assertEqual(scope.get("account_id"), "seller-prod-default")
        self.assertEqual(scope.get("store_name"), "Production Store")
        self.assertEqual(scope.get("environment"), "production")
        self.assertEqual(
            scope.get("environment_source"),
            "trendyol_catalog_default",
        )
        record_ref = (context.get("record_refs") or [])[0]
        self.assertEqual(record_ref.get("account"), "1")
        self.assertEqual(record_ref.get("environment"), "production")

    def test_explicit_environment_scope_is_injected_after_llm_planning(self) -> None:
        with Session(get_engine(self.db_path)) as session:
            session.add(
                ApiAccount(
                    pk=1,
                    account_id="seller-production",
                    comp_name="Production Store",
                    platform="TRENDYOL",
                    extra_config={"environment": "production"},
                )
            )
            session.add(
                ApiAccount(
                    pk=2,
                    account_id="2738",
                    comp_name="Stage Store",
                    platform="TRENDYOL",
                    extra_config={"environment": "stage"},
                )
            )
            session.commit()
        now_ms = int(datetime.now(TZ_GMT3).timestamp() * 1000)
        self._insert_order(
            "ORD-PROD-SCOPE-ONLY",
            now_ms - 1_000,
            package_id=12_601,
            first_name="Production",
            last_name="Customer",
            api_account_id=1,
        )
        self._insert_order(
            "ORD-STAGE-SCOPE-ONLY",
            now_ms,
            package_id=12_602,
            first_name="Stage",
            last_name="Customer",
            api_account_id=2,
        )

        production = run_agent_turn(
            "Yalniz production Trendyol hesabindaki son 5 siparisi listele; Stage hesabini karistirma.",
            planner_client=lambda _prompt: (
                '{"mode":"tool","tool":"orders.recent",'
                '"arguments":{"period":"all","limit":5}}'
            ),
            db_name=self.db_path,
        )
        stage = run_agent_turn(
            "Yalniz Stage Trendyol hesabindaki son 5 siparisi listele; production hesabini karistirma.",
            planner_client=lambda _prompt: (
                '{"mode":"tool","tool":"orders.recent",'
                '"arguments":{"period":"all","limit":5}}'
            ),
            db_name=self.db_path,
        )

        self.assertTrue(production.success, production.message)
        self.assertTrue(stage.success, stage.message)
        production_call = (production.data or {}).get("tool_call") or {}
        stage_call = (stage.data or {}).get("tool_call") or {}
        self.assertEqual(production_call.get("arguments", {}).get("api_account_ids"), [1])
        self.assertEqual(stage_call.get("arguments", {}).get("api_account_ids"), [2])
        production_data = (
            ((production.data or {}).get("tool_results") or [{}])[0].get("data")
            or {}
        )
        stage_data = (
            ((stage.data or {}).get("tool_results") or [{}])[0].get("data")
            or {}
        )
        self.assertEqual(
            [row.get("order_number") for row in production_data.get("orders") or []],
            ["ORD-PROD-SCOPE-ONLY"],
        )
        self.assertEqual(
            [row.get("order_number") for row in stage_data.get("orders") or []],
            ["ORD-STAGE-SCOPE-ONLY"],
        )
        self.assertEqual(
            [row.get("environment") for row in production_data.get("account_scopes") or []],
            ["production"],
        )
        self.assertEqual(
            [row.get("environment") for row in stage_data.get("account_scopes") or []],
            ["stage"],
        )

    def test_tight_local_latest_store_order_list_beats_store_overview(self) -> None:
        class TightLocalClient:
            resource_plan = {"profile": "local_model", "num_ctx": 1024}

            def __call__(self, _prompt: str) -> str:
                raise AssertionError("Tight local path must use deterministic intent.")

        with Session(get_engine(self.db_path)) as session:
            session.add(
                ApiAccount(
                    pk=1,
                    account_id="seller-production",
                    comp_name="Production Store",
                    platform="TRENDYOL",
                    extra_config={"environment": "production"},
                )
            )
            session.add(
                ApiAccount(
                    pk=2,
                    account_id="2738",
                    comp_name="Stage Store",
                    platform="TRENDYOL",
                    extra_config={"environment": "stage"},
                )
            )
            session.commit()
        now_ms = int(datetime.now(TZ_GMT3).timestamp() * 1000)
        self._insert_order(
            "ORD-PROD-LOCAL-LATEST",
            now_ms,
            package_id=12_603,
            first_name="Production",
            last_name="Customer",
            api_account_id=1,
        )
        self._insert_order(
            "ORD-STAGE-LOCAL-LATEST",
            now_ms + 1_000,
            package_id=12_604,
            first_name="Stage",
            last_name="Customer",
            api_account_id=2,
        )

        result = run_agent_turn(
            (
                "Yalniz Production magazamizdaki en guncel 5 siparisi numara ve "
                "durumlariyla listele; Stage kayitlarini kesinlikle karistirma."
            ),
            planner_client=TightLocalClient(),
            db_name=self.db_path,
        )

        self.assertTrue(result.success, result.message)
        calls = list((result.data or {}).get("tool_calls") or [])
        self.assertEqual([row.get("tool") for row in calls], ["orders.recent"])
        self.assertEqual(calls[0].get("arguments", {}).get("api_account_ids"), [1])
        data = (((result.data or {}).get("tool_results") or [{}])[0].get("data") or {})
        self.assertEqual(
            [row.get("order_number") for row in data.get("orders") or []],
            ["ORD-PROD-LOCAL-LATEST"],
        )
        self.assertEqual(
            [row.get("environment") for row in data.get("account_scopes") or []],
            ["production"],
        )

    def test_production_revenue_with_negative_profit_wording_stays_single_scoped_read(self) -> None:
        class TightLocalClient:
            resource_plan = {"profile": "local_model", "num_ctx": 1024}

            def __call__(self, _prompt: str) -> str:
                raise AssertionError("Revenue-only local fallback must be deterministic.")

        with Session(get_engine(self.db_path)) as session:
            session.add(
                ApiAccount(
                    pk=1,
                    account_id="seller-production",
                    comp_name="Production Store",
                    platform="TRENDYOL",
                    extra_config={"environment": "production"},
                )
            )
            session.add(
                ApiAccount(
                    pk=2,
                    account_id="2738",
                    comp_name="Stage Store",
                    platform="TRENDYOL",
                    extra_config={"environment": "stage"},
                )
            )
            session.commit()
        now_ms = int(datetime.now(TZ_GMT3).timestamp() * 1000)
        self._insert_order(
            "ORD-PROD-REVENUE-ONLY",
            now_ms,
            package_id=12_605,
            first_name="Production",
            last_name="Revenue",
            total_price=800.0,
            api_account_id=1,
        )
        self._insert_order(
            "ORD-STAGE-REVENUE-EXCLUDED",
            now_ms,
            package_id=12_606,
            first_name="Stage",
            last_name="Excluded",
            total_price=700.0,
            api_account_id=2,
        )

        result = run_agent_turn(
            (
                "Yalnız Production mağazamızın bu ayki toplam sipariş cirosunu "
                "söyle; Stage kayıtlarını karıştırma ve bunu net kâr diye "
                "adlandırma."
            ),
            planner_client=TightLocalClient(),
            db_name=self.db_path,
        )

        self.assertTrue(result.success, result.message)
        calls = list((result.data or {}).get("tool_calls") or [])
        self.assertEqual([row.get("tool") for row in calls], ["orders.revenue_total"])
        self.assertEqual(calls[0].get("arguments", {}).get("api_account_ids"), [1])
        data = (((result.data or {}).get("tool_results") or [{}])[0].get("data") or {})
        self.assertEqual(data.get("requested_api_account_ids"), [1])
        self.assertEqual(data.get("total"), 800.0)
        answer = self._answer_from(result)
        self.assertIn("800", answer)
        self.assertNotIn("700", answer)
        self.assertNotIn("tahmini kar", answer.casefold())

    def test_tight_local_stage_ready_count_list_beats_latest_orders(self) -> None:
        class TightLocalClient:
            resource_plan = {"profile": "local_model", "num_ctx": 1024}

            def __call__(self, _prompt: str) -> str:
                raise AssertionError("Tight local path must use deterministic intent.")

        with Session(get_engine(self.db_path)) as session:
            session.add(
                ApiAccount(
                    pk=1,
                    account_id="seller-production",
                    comp_name="Production Store",
                    platform="TRENDYOL",
                    extra_config={"environment": "production"},
                )
            )
            session.add(
                ApiAccount(
                    pk=2,
                    account_id="2738",
                    comp_name="Stage Store",
                    platform="TRENDYOL",
                    extra_config={"environment": "stage"},
                )
            )
            session.commit()
        now_ms = int(datetime.now(TZ_GMT3).timestamp() * 1000)
        for index in range(6):
            self._insert_order(
                f"ORD-STAGE-READY-{index}",
                now_ms + index,
                package_id=12_610 + index,
                first_name="Stage",
                last_name=f"Ready{index}",
                status="ReadyToShip",
                api_account_id=2,
            )
        self._insert_order(
            "ORD-STAGE-PICKING-NEWER",
            now_ms + 10,
            package_id=12_620,
            first_name="Stage",
            last_name="Picking",
            status="Picking",
            api_account_id=2,
        )
        self._insert_order(
            "ORD-PROD-READY-NEWER",
            now_ms + 20,
            package_id=12_621,
            first_name="Production",
            last_name="Ready",
            status="ReadyToShip",
            api_account_id=1,
        )
        self._insert_order(
            "ORD-STAGE-SHIPPED-NEWEST",
            now_ms + 30,
            package_id=12_622,
            first_name="Stage",
            last_name="Shipped",
            status="Shipped",
            api_account_id=2,
        )

        result = run_agent_turn(
            (
                "Yalniz Stage magazasinda kargoya hazir kac siparis var ve "
                "bunlarin en guncel ilk 5 tanesini listele; Production karismasin."
            ),
            planner_client=TightLocalClient(),
            db_name=self.db_path,
        )

        self.assertTrue(result.success, result.message)
        calls = list((result.data or {}).get("tool_calls") or [])
        self.assertEqual([row.get("tool") for row in calls], ["orders.ready_to_ship"])
        self.assertEqual(calls[0].get("arguments", {}).get("api_account_ids"), [2])
        self.assertEqual(calls[0].get("arguments", {}).get("limit"), 5)
        data = (((result.data or {}).get("tool_results") or [{}])[0].get("data") or {})
        # The Orders domain intentionally defines the operational ready queue
        # as Created/Picking/ReadyToShip; only terminal Shipped is excluded.
        self.assertEqual(data.get("ready_to_ship"), 7)
        self.assertEqual(data.get("listed_count"), 5)
        order_numbers = [row.get("order_number") for row in data.get("orders") or []]
        self.assertTrue(
            all(
                row.get("status") in {"Created", "Picking", "ReadyToShip"}
                for row in data.get("orders") or []
            )
        )
        self.assertIn("ORD-STAGE-PICKING-NEWER", order_numbers)
        self.assertNotIn("ORD-STAGE-SHIPPED-NEWEST", order_numbers)
        self.assertEqual(
            [row.get("environment") for row in data.get("account_scopes") or []],
            ["stage"],
        )

        stale_cancelled_label_history = [
            {
                "role": "assistant",
                "content": "İki etiket hazırlığı açık onay bekliyor.",
                "context": {
                    "kind": "pending_action",
                    "source_tool": "labels.create_from_orders.prepare",
                    "action_type": "labels.action_create_labels_from_selected_orders",
                    "status": "pending_approval",
                    "approval_required": True,
                    "order_numbers": [
                        "ORD-STAGE-READY-5",
                        "ORD-STAGE-READY-4",
                    ],
                    "model_code": "TANEX_2736",
                    "account_by_order": {
                        "ORD-STAGE-READY-5": 2,
                        "ORD-STAGE-READY-4": 2,
                    },
                    "environment_by_account": {"2": "stage"},
                },
            },
            {
                "role": "user",
                "content": "Vazgeçtim; bu etiket hazırlığını iptal et.",
            },
            {
                "role": "assistant",
                "content": "Tamam, bu işlemi çalıştırmıyorum.",
            },
        ]

        class AdversarialLabelPlanner:
            def __init__(self) -> None:
                self.calls = 0

            def __call__(self, _prompt: str) -> str:
                self.calls += 1
                return (
                    '{"mode":"tool","tool":"labels.create_from_orders.prepare",'
                    '"arguments":{"query":"kargoya hazır","limit":20}}'
                )

        adversarial_planner = AdversarialLabelPlanner()
        explicit_fresh_read = run_agent_turn(
            (
                "Yalnız Stage mağazasındaki kargoya hazır en güncel 1 siparişi "
                "sipariş numarası, durum ve kargo firmasıyla listele; Production "
                "kayıtlarını karıştırma."
            ),
            planner_client=adversarial_planner,
            chat_client=lambda _prompt: (
                "ORD-STAGE-PICKING-NEWER | durum: Picking | "
                "kargo: Trendyol Express"
            ),
            db_name=self.db_path,
            history=stale_cancelled_label_history,
        )

        self.assertTrue(explicit_fresh_read.success, explicit_fresh_read.message)
        self.assertEqual(
            (explicit_fresh_read.data or {}).get("decision_path"),
            "explicit_ready_list_guard",
        )
        self.assertEqual(adversarial_planner.calls, 0)
        self.assertFalse(bool((explicit_fresh_read.data or {}).get("approval_required")))
        fresh_calls = list((explicit_fresh_read.data or {}).get("tool_calls") or [])
        self.assertEqual(
            [row.get("tool") for row in fresh_calls],
            ["orders.ready_to_ship"],
        )
        self.assertEqual(fresh_calls[0].get("arguments", {}).get("limit"), 1)
        self.assertEqual(
            fresh_calls[0].get("arguments", {}).get("api_account_ids"),
            [2],
        )
        fresh_data = (
            ((explicit_fresh_read.data or {}).get("tool_results") or [{}])[0].get("data")
            or {}
        )
        self.assertEqual(fresh_data.get("listed_count"), 1)
        self.assertEqual(
            [row.get("order_number") for row in fresh_data.get("orders") or []],
            ["ORD-STAGE-PICKING-NEWER"],
        )
        self.assertEqual(
            [row.get("cargo_provider") for row in fresh_data.get("orders") or []],
            ["Trendyol Express"],
        )
        self.assertEqual(
            [row.get("environment") for row in fresh_data.get("account_scopes") or []],
            ["stage"],
        )

        fresh_history = [
            *stale_cancelled_label_history,
            {
                "role": "user",
                "content": (
                    "Yalnız Stage mağazasındaki kargoya hazır en güncel 1 "
                    "siparişi listele."
                ),
            },
            {
                "role": "assistant",
                "content": self._answer_from(explicit_fresh_read),
                "context": (explicit_fresh_read.data or {}).get(
                    "conversation_context"
                )
                or {},
            },
        ]
        singular_stage_prepare = run_agent_turn(
            (
                "Bu seçili Stage siparişi için TANEX 2736 etiketini yalnız "
                "onaya hazırla; execute etme, dosya üretme veya DB yazımı yapma."
            ),
            planner_client=lambda _prompt: (
                '{"mode":"tool","tool":"labels.create_from_orders.prepare",'
                '"arguments":{"query":"seçili stage","limit":20,'
                '"model_code":"TANEX_2736"}}'
            ),
            db_name=self.db_path,
            history=fresh_history,
        )

        self.assertTrue(singular_stage_prepare.success, singular_stage_prepare.message)
        self.assertEqual((singular_stage_prepare.data or {}).get("mode"), "approval_required")
        prepare_calls = list((singular_stage_prepare.data or {}).get("tool_calls") or [])
        self.assertEqual(
            [row.get("tool") for row in prepare_calls],
            ["labels.create_from_orders.prepare"],
        )
        prepare_arguments = prepare_calls[0].get("arguments") or {}
        self.assertEqual(
            prepare_arguments.get("order_numbers"),
            ["ORD-STAGE-PICKING-NEWER"],
        )
        self.assertEqual(prepare_arguments.get("model_code"), "TANEX_2736")
        prepare_action = self._action_from(singular_stage_prepare)
        prepare_payload = ((prepare_action.get("arguments") or {}).get("payload") or {})
        self.assertEqual(
            prepare_payload.get("order_numbers"),
            ["ORD-STAGE-PICKING-NEWER"],
        )
        self.assertEqual(prepare_payload.get("model_code"), "TANEX_2736")
        self.assertEqual(
            ((prepare_action.get("arguments") or {}).get("account_by_order") or {}).get(
                "ORD-STAGE-PICKING-NEWER"
            ),
            "2",
        )
        self.assertEqual(
            (prepare_action.get("arguments") or {}).get("environment_by_account"),
            {"2": "stage"},
        )
        self.assertEqual(
            [
                prepare_action.get("selected_record_count"),
                prepare_action.get("preview_item_count"),
                prepare_action.get("execute_input_count"),
            ],
            [1, 1, 1],
        )

    def test_tight_local_selection_label_prepare_beats_generic_multi_tool(self) -> None:
        class TightLocalClient:
            resource_plan = {"profile": "local_model", "num_ctx": 1024}

            def __call__(self, _prompt: str) -> str:
                raise AssertionError("Tight local path must use deterministic intent.")

        with Session(get_engine(self.db_path)) as session:
            session.add(
                ApiAccount(
                    pk=2,
                    account_id="2738",
                    comp_name="Stage Store",
                    platform="TRENDYOL",
                    extra_config={"environment": "stage"},
                )
            )
            session.commit()
        now_ms = int(datetime.now(TZ_GMT3).timestamp() * 1000)
        for index in range(3):
            self._insert_order(
                f"ORD-STAGE-LABEL-{index}",
                now_ms + index,
                package_id=12_630 + index,
                first_name="Stage",
                last_name=f"Label{index}",
                status="Picking",
                api_account_id=2,
            )
        history = [
            {
                "role": "assistant",
                "content": (
                    "1. ORD-STAGE-LABEL-0 | durum: Picking\n"
                    "2. ORD-STAGE-LABEL-1 | durum: Picking\n"
                    "3. ORD-STAGE-LABEL-2 | durum: Picking"
                ),
                "context": {
                    "kind": "order_selection",
                    "source_tool": "orders.ready_to_ship",
                    "period": "all",
                    "order_numbers": [
                        "ORD-STAGE-LABEL-0",
                        "ORD-STAGE-LABEL-1",
                        "ORD-STAGE-LABEL-2",
                    ],
                    "record_refs": [
                        {
                            "domain": "orders",
                            "record_type": "order",
                            "record_id": f"ORD-STAGE-LABEL-{index}",
                            "source": "local_db",
                            "account": "2",
                            "environment": "stage",
                            "identifiers": {},
                        }
                        for index in range(3)
                    ],
                    "listed_count": 3,
                    "total_count": 3,
                    "truncated": False,
                },
            }
        ]

        result = run_agent_turn(
            (
                "Bu listedeki ilk iki Stage siparisi icin TANEX 2610 etiketini "
                "yalniz onaya hazirla; execute etme, dosya uretme ve DB yazimi yapma."
            ),
            planner_client=TightLocalClient(),
            db_name=self.db_path,
            history=history,
        )

        self.assertTrue(result.success, result.message)
        self.assertEqual((result.data or {}).get("mode"), "approval_required")
        calls = list((result.data or {}).get("tool_calls") or [])
        self.assertEqual(
            [row.get("tool") for row in calls],
            ["labels.create_from_orders.prepare"],
        )
        arguments = calls[0].get("arguments") or {}
        self.assertEqual(
            arguments.get("order_numbers"),
            ["ORD-STAGE-LABEL-0", "ORD-STAGE-LABEL-1"],
        )
        self.assertEqual(arguments.get("model_code"), "TANEX_2610")
        action = self._action_from(result)
        payload = ((action.get("arguments") or {}).get("payload") or {})
        self.assertEqual(
            payload.get("order_numbers"),
            ["ORD-STAGE-LABEL-0", "ORD-STAGE-LABEL-1"],
        )
        self.assertEqual(payload.get("model_code"), "TANEX_2610")
        self.assertTrue(action.get("approval_required"))

        explanation = run_agent_turn(
            (
                "Bu etiket hazırlığı neden onay gerektiriyor? Yeni tool "
                "çalıştırma ve işlemi uygulama; yalnız güvenlik sınırını açıkla."
            ),
            planner_client=TightLocalClient(),
            db_name=self.db_path,
            history=[
                *history,
                {
                    "role": "user",
                    "content": (
                        "Bu listedeki ilk iki Stage siparisi icin TANEX 2610 "
                        "etiketini yalniz onaya hazirla; execute etme."
                    ),
                },
                {
                    "role": "assistant",
                    "content": self._answer_from(result),
                    "context": dict((result.data or {}).get("conversation_context") or {}),
                },
            ],
        )

        self.assertTrue(explanation.success, explanation.message)
        self.assertEqual((explanation.data or {}).get("tool_calls"), [])
        self.assertEqual(
            (explanation.data or {}).get("decision_path"),
            "explicit_no_tool_guard",
        )
        explanation_answer = self._answer_from(explanation)
        self.assertIn("onay istedim", explanation_answer.casefold())
        self.assertIn("açık onay olmadan execute edilmez", explanation_answer)
        self.assertNotIn("neye karar vermek", explanation_answer.casefold())

        cancelled = run_agent_turn(
            (
                "Vazgectim; bu etiket hazirligini iptal et. Hicbir tool "
                "calistirma, dosya uretme veya execute etme."
            ),
            planner_client=TightLocalClient(),
            db_name=self.db_path,
            history=[
                *history,
                {
                    "role": "assistant",
                    "content": self._answer_from(result),
                    "context": dict((result.data or {}).get("conversation_context") or {}),
                },
                {
                    "role": "assistant",
                    "content": explanation_answer,
                },
            ],
        )

        self.assertTrue(cancelled.success, cancelled.message)
        self.assertEqual((cancelled.data or {}).get("tool_calls"), [])
        cancelled_answer = self._answer_from(cancelled)
        self.assertIn("bu islemi calistirmiyorum", cancelled_answer.casefold())
        self.assertIn("API istegi", cancelled_answer)
        self.assertIn("DB yazimi", cancelled_answer)
        self.assertIn("dosya/cikti uretimi yapilmadi", cancelled_answer)
        self.assertNotIn("neye karar vermek", cancelled_answer.casefold())
        self.assertEqual(
            (cancelled.data or {}).get("answer_origin"),
            "deterministic_cancel",
        )

    def test_pending_return_approval_cancel_without_no_tool_phrase_beats_planner_fanout(self) -> None:
        class AdversarialClient:
            def __init__(self) -> None:
                self.calls = 0

            def __call__(self, _prompt: str) -> str:
                self.calls += 1
                return (
                    '{"mode":"tools","tools":['
                    '{"tool":"returns.count","arguments":{"period":"all"}},'
                    '{"tool":"orders.cancelled_count","arguments":{"period":"all"}},'
                    '{"tool":"listing.attention_count","arguments":{}}]}'
                )

        client = AdversarialClient()
        result = run_agent_turn(
            (
                "Vazgectim, iade onayi hazirligini iptal et; hicbir API istegi "
                "veya DB yazimi yapma."
            ),
            planner_client=client,
            chat_client=client,
            db_name=self.db_path,
        )

        self.assertTrue(result.success, result.message)
        self.assertEqual((result.data or {}).get("tool_calls"), [])
        self.assertEqual(
            (result.data or {}).get("decision_path"),
            "pending_action_cancel_guard",
        )
        self.assertFalse((result.data or {}).get("model_invoked"))
        self.assertEqual(client.calls, 0)
        self.assertEqual(
            (result.data or {}).get("answer_origin"),
            "deterministic_cancel",
        )
        answer = self._answer_from(result)
        self.assertIn("bu islemi calistirmiyorum", answer.casefold())
        self.assertIn("API istegi", answer)
        self.assertIn("DB yazimi", answer)
        self.assertNotIn("ilan dikkat", answer.casefold())

    def test_tight_local_stage_waiting_return_count_beats_store_overview(self) -> None:
        class TightLocalClient:
            resource_plan = {"profile": "local_model", "num_ctx": 1024}

            def __call__(self, _prompt: str) -> str:
                raise AssertionError("Tight local path must use deterministic intent.")

        with Session(get_engine(self.db_path)) as session:
            session.add(
                ApiAccount(
                    pk=1,
                    account_id="seller-production",
                    comp_name="Production Store",
                    platform="TRENDYOL",
                    extra_config={"environment": "production"},
                )
            )
            session.add(
                ApiAccount(
                    pk=2,
                    account_id="2738",
                    comp_name="Stage Store",
                    platform="TRENDYOL",
                    extra_config={"environment": "stage"},
                )
            )
            session.commit()
        now_ms = int(datetime.now(TZ_GMT3).timestamp() * 1000)
        self._insert_waiting_return(
            "CLAIM-PROD-WAITING",
            "ORD-PROD-WAITING",
            now_ms,
            first_name="Production",
            last_name="Claim",
            api_account_id=1,
            reason="Production nedeni",
        )
        for index in range(2):
            self._insert_waiting_return(
                f"CLAIM-STAGE-WAITING-{index}",
                f"ORD-STAGE-WAITING-{index}",
                now_ms + index + 1,
                first_name="Stage",
                last_name=f"Claim{index}",
                api_account_id=2,
                reason="Teslim edilemeyen gönderi",
            )

        with Session(get_engine(self.db_path)) as session:
            session.add(
                ReturnClaim(
                    claimId="CLAIM-STAGE-ACCEPTED",
                    orderNumber="ORD-STAGE-ACCEPTED",
                    api_account_id=2,
                    claimDate=now_ms + 10,
                    orderDate=now_ms + 10,
                    lastModifiedDate=now_ms + 11,
                    claimItemStatus="Accepted",
                    claimReason="Vazgectim",
                    totalPrice=55.0,
                    itemCount=1,
                )
            )
            session.commit()

        result = run_agent_turn(
            (
                "Yalniz Stage magazasinda WaitingInAction durumunda kac iade "
                "claimi var? Production kayitlarini karistirma."
            ),
            planner_client=TightLocalClient(),
            db_name=self.db_path,
        )

        self.assertTrue(result.success, result.message)
        calls = list((result.data or {}).get("tool_calls") or [])
        self.assertEqual(
            [row.get("tool") for row in calls],
            ["returns.waiting_action_count"],
        )
        self.assertEqual(calls[0].get("arguments", {}).get("api_account_ids"), [2])
        data = (((result.data or {}).get("tool_results") or [{}])[0].get("data") or {})
        self.assertEqual(data.get("count"), 2)
        self.assertEqual(data.get("requested_api_account_ids"), [2])
        self.assertEqual(
            [row.get("environment") for row in data.get("account_scopes") or []],
            ["stage"],
        )

        listed = run_agent_turn(
            (
                "Bunlarin en guncel ilk 5 tanesini claim ID, siparis numarasi, "
                "durum, neden ve tutarla listele; ayni Stage kapsamini koru."
            ),
            planner_client=TightLocalClient(),
            db_name=self.db_path,
            history=[
                {
                    "role": "user",
                    "content": (
                        "Yalniz Stage magazasinda WaitingInAction durumunda kac "
                        "iade claimi var? Production kayitlarini karistirma."
                    ),
                },
                {
                    "role": "assistant",
                    "content": self._answer_from(result),
                    "context": dict((result.data or {}).get("conversation_context") or {}),
                },
            ],
        )

        self.assertTrue(listed.success, listed.message)
        listed_calls = list((listed.data or {}).get("tool_calls") or [])
        self.assertEqual(
            [row.get("tool") for row in listed_calls],
            ["returns.query"],
        )
        listed_args = listed_calls[0].get("arguments") or {}
        self.assertEqual(listed_args.get("filters"), {"status": "WaitingInAction"})
        self.assertEqual(listed_args.get("limit"), 5)
        self.assertEqual(listed_args.get("api_account_ids"), [2])
        listed_data = (
            ((listed.data or {}).get("tool_results") or [{}])[0].get("data") or {}
        )
        listed_claims = listed_data.get("returns") or listed_data.get("claims") or []
        self.assertEqual(
            {row.get("claim_id") for row in listed_claims},
            {"CLAIM-STAGE-WAITING-0", "CLAIM-STAGE-WAITING-1"},
        )
        self.assertTrue(
            all(row.get("api_account_id") == 2 for row in listed_claims)
        )
        listed_context = dict(
            (listed.data or {}).get("conversation_context") or {}
        )
        self.assertEqual(
            listed_context.get("reason_by_claim"),
            {
                "CLAIM-STAGE-WAITING-0": "Teslim edilemeyen gönderi",
                "CLAIM-STAGE-WAITING-1": "Teslim edilemeyen gönderi",
            },
        )

        common_reason = run_agent_turn(
            (
                "Bu son 5 iadenin ortak nedeni neydi? Yeni tool çalıştırma; "
                "yalnız az önceki listeden cevapla."
            ),
            planner_client=TightLocalClient(),
            db_name=self.db_path,
            history=[
                {
                    "role": "assistant",
                    "content": self._answer_from(listed),
                    "context": listed_context,
                }
            ],
        )

        self.assertTrue(common_reason.success, common_reason.message)
        self.assertEqual((common_reason.data or {}).get("tool_calls"), [])
        self.assertEqual(
            (common_reason.data or {}).get("decision_path"),
            "explicit_no_tool_guard",
        )
        reason_answer = self._answer_from(common_reason)
        self.assertIn("Teslim edilemeyen gönderi", reason_answer)
        self.assertNotIn("Production nedeni", reason_answer)
        self.assertFalse(reason_answer.lstrip().startswith(("{", "[")))
        self.assertEqual(
            (common_reason.data or {}).get("answer_origin"),
            "decision_response",
        )

        reason_summary = run_agent_turn(
            (
                "Yalniz Stage magazasindaki WaitingInAction iadelerin neden "
                "dagilimini ozetle; Production kayitlarini karistirma."
            ),
            planner_client=TightLocalClient(),
            db_name=self.db_path,
        )

        self.assertTrue(reason_summary.success, reason_summary.message)
        summary_calls = list((reason_summary.data or {}).get("tool_calls") or [])
        self.assertEqual(
            [row.get("tool") for row in summary_calls],
            ["returns.reason_summary"],
        )
        summary_args = summary_calls[0].get("arguments") or {}
        self.assertTrue(summary_args.get("waiting_action_only"))
        self.assertEqual(summary_args.get("api_account_ids"), [2])
        summary_data = (
            ((reason_summary.data or {}).get("tool_results") or [{}])[0].get("data")
            or {}
        )
        self.assertEqual(summary_data.get("total_claims"), 2)
        self.assertTrue(summary_data.get("waiting_action_only"))
        self.assertEqual(summary_data.get("requested_api_account_ids"), [2])
        summary_answer = self._answer_from(reason_summary)
        self.assertIn("Teslim edilemeyen", summary_answer)
        self.assertNotIn("Vazgectim", summary_answer)
        self.assertNotIn("Production nedeni", summary_answer)

    def test_no_tool_status_followup_uses_only_latest_selection_state(self) -> None:
        history = [
            {
                "role": "assistant",
                "content": (
                    "1. ORD-STAGE-1 | durum: Shipped | tutar: 800 TL\n"
                    "2. ORD-STAGE-2 | durum: Picking | tutar: 800 TL\n"
                    "3. ORD-STAGE-3 | durum: Picking | tutar: 800 TL"
                ),
                "context": {
                    "kind": "order_selection",
                    "source_tool": "orders.recent",
                    "period": "all",
                    "order_numbers": [
                        "ORD-STAGE-1",
                        "ORD-STAGE-2",
                        "ORD-STAGE-3",
                    ],
                    "listed_count": 3,
                    "total_count": 3,
                    "truncated": False,
                },
            }
        ]

        result = run_agent_turn(
            "Bu son uclude Picking durumunda olanlarin siparis numaralarini soyle; yeni tool calistirma.",
            planner_client=lambda _prompt: (_ for _ in ()).throw(
                AssertionError("No-tool boundary planneri cagirmamali.")
            ),
            db_name=self.db_path,
            history=history,
        )

        self.assertTrue(result.success, result.message)
        self.assertEqual((result.data or {}).get("mode"), "answer")
        self.assertEqual((result.data or {}).get("tool_calls"), [])
        self.assertEqual(
            (result.data or {}).get("decision_path"),
            "explicit_no_tool_guard",
        )
        answer = self._answer_from(result)
        self.assertIn("ORD-STAGE-2", answer)
        self.assertIn("ORD-STAGE-3", answer)
        self.assertNotIn("ORD-STAGE-1", answer)

    def test_explicit_recent_refresh_beats_stale_account_scope_answer(self) -> None:
        self._insert_order(
            "ORD-RECENT-REFRESH-1",
            int(datetime.now(TZ_GMT3).timestamp() * 1000),
            package_id=1_252,
            first_name="Refresh",
            last_name="Customer",
            status="ReadyToShip",
        )
        history = [
            {"role": "user", "content": "son gelen 5 siparisi listele"},
            {
                "role": "assistant",
                "content": "1. ORD-OLD | durum: ReadyToShip",
                "context": {
                    "kind": "orders_ready_to_ship_scope",
                    "source_tool": "orders.ready_to_ship_count",
                    "period": "all",
                    "total_count": 1,
                    "account_scopes": [],
                },
            },
        ]

        result = run_agent_turn(
            (
                "Ayn\u0131 son 5 sipari\u015fi yeniden getir; bu kez hangi "
                "ma\u011faza hesab\u0131 ve ortam kapsam\u0131ndan okudu\u011funu "
                "state i\u00e7inde koru."
            ),
            planner_client=lambda _prompt: (
                '{"mode":"answer","response":"Eski kapsam bilinmiyor."}'
            ),
            db_name=self.db_path,
            history=history,
        )

        self.assertTrue(result.success, result.message)
        self.assertEqual(
            ((result.data or {}).get("tool_call") or {}).get("tool"),
            "orders.recent",
        )
        self.assertEqual(
            (((result.data or {}).get("tool_call") or {}).get("arguments") or {}).get("limit"),
            5,
        )

    def test_full_context_creates_single_record_selection_with_account_scope(self) -> None:
        with Session(get_engine(self.db_path)) as session:
            session.add(
                ApiAccount(
                    pk=1,
                    account_id="seller-full-context",
                    comp_name="Full Context Store",
                    platform="TRENDYOL",
                    extra_config={"environment": "production"},
                )
            )
            session.commit()
        order_number = "ORD-FULL-CONTEXT-SCOPE-1"
        self._insert_order(
            order_number,
            int(datetime.now(TZ_GMT3).timestamp() * 1000),
            package_id=1_253,
            first_name="Full",
            last_name="Context",
            status="ReadyToShip",
            api_account_id=1,
        )

        result = run_agent_turn(
            f"{order_number} siparisinin durumunu, tutarini ve kalemlerini goster",
            planner_client=lambda _prompt: (
                '{"mode":"tool","tool":"orders.full_context",'
                f'"arguments":{{"query":"{order_number}","limit":1}}}}'
            ),
            db_name=self.db_path,
        )

        self.assertTrue(result.success, result.message)
        context = (result.data or {}).get("conversation_context") or {}
        self.assertEqual(context.get("kind"), "order_selection")
        self.assertEqual(context.get("order_numbers"), [order_number])
        self.assertEqual(context.get("listed_count"), 1)
        ref = (context.get("record_refs") or [])[0]
        self.assertEqual(ref.get("account"), "1")
        self.assertEqual(ref.get("environment"), "production")
        scope = (context.get("account_scopes") or [])[0]
        self.assertEqual(scope.get("store_name"), "Full Context Store")
        self.assertEqual(scope.get("environment"), "production")

    def test_label_prepare_preserves_selection_order_and_account_scope(self) -> None:
        with Session(get_engine(self.db_path)) as session:
            session.add(
                ApiAccount(
                    pk=1,
                    account_id="seller-label-scope",
                    comp_name="Label Scope Store",
                    platform="TRENDYOL",
                    extra_config={"environment": "production"},
                )
            )
            session.commit()
        now_ms = int(datetime.now(TZ_GMT3).timestamp() * 1000)
        expected = ["ORD-LABEL-SCOPE-3", "ORD-LABEL-SCOPE-1", "ORD-LABEL-SCOPE-2"]
        for index, order_number in enumerate(expected):
            self._insert_order(
                order_number,
                now_ms + index,
                package_id=1_260 + index,
                first_name="Label",
                last_name=f"Scope{index}",
                status="ReadyToShip",
                api_account_id=1,
            )
        history = [
            {
                "role": "assistant",
                "content": "Doğrulanmış üçlü seçim.",
                "context": SelectionSnapshot.from_order_numbers(
                    expected,
                    source_tool="orders.ready_to_ship",
                    period="all",
                    account_by_order={value: "1" for value in expected},
                    environment_by_account={"1": "production"},
                ).to_context(),
            }
        ]

        result = run_agent_turn(
            (
                "bu doğrulanmış listedeki ilk üç sipariş için etiket hazırlığını "
                "yalnız onaya hazırla; execute etme ve dosya üretme"
            ),
            db_name=self.db_path,
            history=history,
        )

        self.assertTrue(result.success, result.message)
        self.assertEqual((result.data or {}).get("mode"), "approval_required")
        action = self._action_from(result)
        payload = ((action.get("arguments") or {}).get("payload") or {})
        self.assertEqual(payload.get("order_numbers"), expected)
        self.assertEqual(
            [row.get("account") for row in list(action.get("record_refs") or [])],
            ["1", "1", "1"],
        )
        self.assertEqual(
            [row.get("environment") for row in list(action.get("record_refs") or [])],
            ["production", "production", "production"],
        )
        context = (result.data or {}).get("conversation_context") or {}
        self.assertEqual(context.get("order_numbers"), expected)
        self.assertEqual(
            [row.get("environment") for row in list(context.get("record_refs") or [])],
            ["production", "production", "production"],
        )
        self.assertEqual(
            (context.get("account_scopes") or [])[0].get("store_name"),
            "Label Scope Store",
        )

    def test_label_prepare_natural_reprepare_uses_parent_seven_selection(self) -> None:
        now_ms = int(datetime.now(TZ_GMT3).timestamp() * 1000)
        authoritative = [f"ORD-LABEL-REPREPARE-{index}" for index in range(1, 8)]
        expected = authoritative[:3]
        for index, order_number in enumerate(authoritative):
            self._insert_order(
                order_number,
                now_ms + index,
                package_id=1_270 + index,
                first_name="Reprepare",
                last_name=f"Scope{index}",
                status="ReadyToShip",
            )
        history = [
            {
                "role": "assistant",
                "content": "Doğrulanmış yedili seçim.",
                "context": SelectionSnapshot.from_order_numbers(
                    authoritative,
                    source_tool="orders.ready_to_ship",
                    period="all",
                ).to_context(),
            },
            {
                "role": "assistant",
                "content": "İlk üç için önceki hazırlık.",
                "context": SelectionSnapshot.from_order_numbers(
                    list(reversed(expected)),
                    source_tool="labels.create_from_orders.prepare",
                ).to_context(),
            },
        ]

        result = run_agent_turn(
            (
                "hazırlık kapsamı doğruydu ama account/environment unknown kalmıştı; "
                "aynı authoritative 7'li listenin ilk üçü için prepare'ı provenance "
                "korunacak şekilde yeniden oluştur. hâlâ execute etme"
            ),
            planner_client=lambda _prompt: (
                '{"mode":"tool","tool":"orders.ready_to_ship",'
                '"arguments":{"period":"all","limit":100}}'
            ),
            db_name=self.db_path,
            history=history,
        )

        self.assertTrue(result.success, result.message)
        self.assertEqual((result.data or {}).get("mode"), "approval_required")
        tool_call = (result.data or {}).get("tool_call") or {}
        self.assertEqual(tool_call.get("tool"), "labels.create_from_orders.prepare")
        self.assertEqual(
            (tool_call.get("arguments") or {}).get("order_numbers"),
            expected,
        )
        payload = ((self._action_from(result).get("arguments") or {}).get("payload") or {})
        self.assertEqual(payload.get("order_numbers"), expected)

    def test_2026_07_15_stale_four_prepare_cannot_override_authoritative_seven_selection(self) -> None:
        now = datetime.now(TZ_GMT3)
        expected = [f"ORD-READY-SEVEN-{index}" for index in range(7)]
        stale = [f"ORD-STALE-FOUR-{index}" for index in range(4)]
        for index, order_number in enumerate([*expected, *stale]):
            self._insert_order(
                order_number,
                int((now - timedelta(minutes=index)).timestamp() * 1000),
                package_id=1_300 + index,
                first_name="Selection",
                last_name=f"Record{index}",
                status="ReadyToShip",
            )

        authoritative_context = SelectionSnapshot.from_order_numbers(
            expected,
            source_tool="orders.ready_to_ship",
            period="all",
        ).to_context()
        stale_prepare_context = SelectionSnapshot.from_order_numbers(
            stale,
            source_tool="labels.create_from_orders.prepare",
            period="all",
        ).to_context()
        history = [
            {"role": "user", "content": "toplam bekleyen kaç sipariş var"},
            {
                "role": "assistant",
                "content": "Lokal kayda göre 7 sipariş buldum.",
                "context": authoritative_context,
            },
            {"role": "user", "content": "bu siparişleri yazdırabilir misin"},
            {
                "role": "assistant",
                "content": "4 sipariş için etiket hazırlanıyor.",
                "context": stale_prepare_context,
            },
        ]

        revised = run_agent_turn(
            "bütün bekleyen siparişler için istemiştim",
            planner_client=lambda _prompt: (
                '{"mode":"tool","tool":"orders.ready_to_ship_count","arguments":{"period":"all"}}'
            ),
            db_name=self.db_path,
            history=history,
        )
        self.assertTrue(revised.success, revised.message)
        self.assertEqual((revised.data or {}).get("mode"), "approval_required")
        revised_payload = ((self._action_from(revised).get("arguments") or {}).get("payload") or {})
        self.assertEqual(set(revised_payload.get("order_numbers") or []), set(expected))

        explicit = run_agent_turn(
            "tamam yazdır bu 7 taneyi",
            planner_client=lambda _prompt: (
                '{"mode":"tool","tool":"labels.create_from_orders.prepare","arguments":{"query":"","limit":20}}'
            ),
            db_name=self.db_path,
            history=history,
        )
        self.assertTrue(explicit.success, explicit.message)
        self.assertEqual((explicit.data or {}).get("mode"), "approval_required")
        explicit_payload = ((self._action_from(explicit).get("arguments") or {}).get("payload") or {})
        self.assertEqual(set(explicit_payload.get("order_numbers") or []), set(expected))
        self.assertTrue(set(explicit_payload.get("order_numbers") or []).isdisjoint(stale))

    def test_2026_07_15_full_user_chain_keeps_seven_records_through_final_prepare(self) -> None:
        with Session(get_engine(self.db_path)) as session:
            session.add(
                ApiAccount(
                    pk=77,
                    account_id="seller-chain-77",
                    comp_name="Modo Di Bellezza",
                    platform="trendyol",
                    extra_config={"environment": "production"},
                )
            )
            session.commit()
        now = datetime.now(TZ_GMT3)
        expected = [f"ORD-CHAIN-SEVEN-{index}" for index in range(7)]
        for index, order_number in enumerate(expected):
            self._insert_order(
                order_number,
                int((now - timedelta(days=index + 1)).timestamp() * 1000),
                package_id=1_400 + index,
                first_name="Chain",
                last_name=f"Ready{index}",
                status="ReadyToShip",
                api_account_id=77,
            )
        for index in range(5):
            self._insert_order(
                f"SIM-CHAIN-{index}",
                int((now - timedelta(days=index + 1)).timestamp() * 1000),
                package_id=1_500 + index,
                first_name="Simulation",
                last_name=f"Excluded{index}",
                status="ReadyToShip",
                cargo_provider_name="Simulated Cargo",
                api_account_id=77,
            )

        history: list[dict[str, object]] = []

        def turn(message: str, planner_response: str) -> Result:
            result = run_agent_turn(
                message,
                planner_client=lambda _prompt: planner_response,
                db_name=self.db_path,
                history=history,
            )
            self.assertTrue(result.success, result.message)
            history.append({"role": "user", "content": message})
            assistant: dict[str, object] = {
                "role": "assistant",
                "content": self._answer_from(result),
            }
            context = (result.data or {}).get("conversation_context")
            if isinstance(context, dict) and context:
                assistant["context"] = context
            history.append(assistant)
            return result

        counted = turn(
            "kaç sipariş var bekleyen",
            '{"mode":"tool","tool":"orders.ready_to_ship_count","arguments":{"period":"today"}}',
        )
        self.assertEqual(((counted.data or {}).get("tool_call") or {}).get("tool"), "orders.ready_to_ship_count")
        self.assertEqual((((counted.data or {}).get("tool_call") or {}).get("arguments") or {}).get("period"), "all")
        self.assertEqual(((counted.data or {}).get("tool_result") or {}).get("count"), 7)
        self.assertEqual(((counted.data or {}).get("tool_result") or {}).get("simulated_excluded"), 5)

        rechecked = turn(
            "bi daha bak hele",
            '{"mode":"answer","response":"Tekrar baktım, bugün sipariş yok."}',
        )
        self.assertEqual(((rechecked.data or {}).get("tool_call") or {}).get("tool"), "orders.ready_to_ship")
        self.assertEqual(
            {row.get("order_number") for row in (((rechecked.data or {}).get("tool_result") or {}).get("orders") or [])},
            set(expected),
        )

        scoped = turn(
            "hangi mağazaya bakıyosun şuan",
            '{"mode":"answer","response":"Şu an tüm mağazalarınıza bakıyorum."}',
        )
        self.assertIn("Modo Di Bellezza", self._answer_from(scoped))
        self.assertIn("production", self._answer_from(scoped))
        self.assertNotIn("tüm mağazalarınıza bakıyorum", self._answer_from(scoped).casefold())

        listed = turn(
            "bütün bekleyen siparişleri listele; toplam kapsamı koru",
            '{"mode":"tool","tool":"orders.ready_to_ship","arguments":{"period":"today","limit":4}}',
        )
        listed_context = (listed.data or {}).get("conversation_context") or {}
        listed_snapshot = SelectionSnapshot.from_context(listed_context)
        self.assertIsNotNone(listed_snapshot)
        assert listed_snapshot is not None
        self.assertTrue(listed_snapshot.is_complete)
        self.assertEqual(set(listed_snapshot.record_ids), set(expected))

        prepared = turn(
            "bu siparişleri yazdırabilir misin",
            '{"mode":"tool","tool":"labels.create_from_orders.prepare","arguments":{"query":"","limit":4}}',
        )
        self.assertEqual((prepared.data or {}).get("mode"), "approval_required")
        prepared_action = self._action_from(prepared)
        self.assertEqual(set(((prepared_action.get("arguments") or {}).get("payload") or {}).get("order_numbers") or []), set(expected))

        revised = turn(
            "bütün bekleyen siparişler için istemiştim",
            '{"mode":"tool","tool":"orders.ready_to_ship_count","arguments":{"period":"all"}}',
        )
        revised_action = self._action_from(revised)
        self.assertEqual(set(((revised_action.get("arguments") or {}).get("payload") or {}).get("order_numbers") or []), set(expected))

        final_prepare = turn(
            "tamam yazdır bu 7 taneyi",
            '{"mode":"tool","tool":"labels.create_from_orders.prepare","arguments":{"query":"","limit":20}}',
        )
        self.assertEqual((final_prepare.data or {}).get("mode"), "approval_required")
        final_action = self._action_from(final_prepare)
        final_payload = ((final_action.get("arguments") or {}).get("payload") or {})
        self.assertEqual(set(final_payload.get("order_numbers") or []), set(expected))
        final_snapshot = PendingActionSnapshot.from_action(final_action)
        self.assertEqual(final_snapshot.selected_record_count, 7)
        self.assertEqual(final_snapshot.preview_item_count, 7)
        self.assertEqual(final_snapshot.execute_input_count, 7)
        self.assertEqual(final_snapshot.invariant_problems(final_action), [])
        self.assertTrue(final_action.get("prepare_fingerprint"))

    def test_2026_07_18_direct_ready_count_to_deictic_2610_prepare_keeps_exact_four(self) -> None:
        with Session(get_engine(self.db_path)) as session:
            session.add(
                ApiAccount(
                    pk=178,
                    account_id="seller-direct-count-178",
                    comp_name="MDB",
                    platform="trendyol",
                    extra_config={"environment": "production"},
                )
            )
            session.commit()
        now_ms = int(datetime.now(TZ_GMT3).timestamp() * 1000)
        expected = [f"ORD-DIRECT-COUNT-{index}" for index in range(1, 5)]
        for index, order_number in enumerate(expected):
            self._insert_order(
                order_number,
                now_ms + index,
                package_id=17_800 + index,
                first_name="Direct",
                last_name=f"Count{index}",
                status="ReadyToShip",
                api_account_id=178,
            )
        for index in range(5):
            self._insert_order(
                f"SIM-DIRECT-COUNT-{index}",
                now_ms + 100 + index,
                package_id=17_900 + index,
                first_name="Simulation",
                last_name=f"Excluded{index}",
                status="ReadyToShip",
                cargo_provider_name="Simulated Cargo",
                api_account_id=178,
            )

        counted = run_agent_turn(
            "baskan kac siparis var bekleyen",
            planner_client=lambda _prompt: (
                '{"mode":"tool","tool":"orders.ready_to_ship_count",'
                '"arguments":{"period":"today"}}'
            ),
            db_name=self.db_path,
        )
        self.assertTrue(counted.success, counted.message)
        count_data = (counted.data or {}).get("tool_result") or {}
        self.assertEqual(count_data.get("count"), 4)
        self.assertEqual(count_data.get("simulated_excluded"), 5)
        self.assertEqual(set(count_data.get("order_numbers") or []), set(expected))
        count_context = (counted.data or {}).get("conversation_context") or {}
        self.assertEqual(count_context.get("source_tool"), "orders.ready_to_ship_count")
        self.assertEqual(set(count_context.get("order_numbers") or []), set(expected))
        self.assertFalse(count_context.get("truncated"))

        history = [
            {"role": "user", "content": "baskan kac siparis var bekleyen"},
            {
                "role": "assistant",
                "content": self._answer_from(counted),
                "context": count_context,
            },
        ]
        prepared = run_agent_turn(
            "yazdir bu siparisleri etiket olarak 2610 modelini kullan",
            planner_client=lambda _prompt: (
                '{"mode":"tool","tool":"labels.create_from_orders.prepare",'
                '"arguments":{"query":"2610","limit":1}}'
            ),
            db_name=self.db_path,
            history=history,
        )

        self.assertTrue(prepared.success, prepared.message)
        self.assertEqual((prepared.data or {}).get("mode"), "approval_required")
        tool_call = (prepared.data or {}).get("tool_call") or {}
        self.assertEqual(tool_call.get("tool"), "labels.create_from_orders.prepare")
        arguments = tool_call.get("arguments") or {}
        self.assertEqual(set(arguments.get("order_numbers") or []), set(expected))
        self.assertEqual(arguments.get("api_account_ids"), [178])
        self.assertEqual(arguments.get("model_code"), "TANEX_2610")
        self.assertFalse(arguments.get("only_not_printed"))
        action = self._action_from(prepared)
        payload = ((action.get("arguments") or {}).get("payload") or {})
        self.assertEqual(set(payload.get("order_numbers") or []), set(expected))
        self.assertEqual(payload.get("model_code"), "TANEX_2610")
        self.assertNotIn("bulunamadi", self._answer_from(prepared).casefold())

    def test_ready_not_printed_prompt_lists_only_unprinted_ready_orders(self) -> None:
        now_ms = int(datetime.now(TZ_GMT3).timestamp() * 1000)
        self._insert_order(
            "ORD-READY-PRINTED",
            now_ms,
            package_id=1_091,
            first_name="Ready",
            last_name="Printed",
            status="ReadyToShip",
            printed=True,
        )
        self._insert_order(
            "ORD-READY-UNPRINTED",
            now_ms + 1_000,
            package_id=1_092,
            first_name="Ready",
            last_name="Unprinted",
            status="ReadyToShip",
            printed=False,
        )
        self._insert_order(
            "ORD-READY-DELIVERED",
            now_ms + 2_000,
            package_id=1_093,
            first_name="Ready",
            last_name="Delivered",
            status="Delivered",
            printed=False,
        )

        result = run_agent_turn("Kargoya hazir yazdirilmamis paketleri listele", db_name=self.db_path)

        self.assertTrue(result.success, result.message)
        self.assertEqual(((result.data or {}).get("tool_call") or {}).get("tool"), "orders.ready_to_ship")
        args = ((result.data or {}).get("tool_call") or {}).get("arguments") or {}
        self.assertTrue(args.get("only_not_printed"))
        answer = (result.data or {}).get("answer") or result.message
        self.assertIn("ORD-READY-UNPRINTED", answer)
        self.assertIn("yazdirildi: hayir", answer)
        self.assertNotIn("ORD-READY-PRINTED", answer)
        self.assertNotIn("ORD-READY-DELIVERED", answer)
        tool_data = (result.data or {}).get("tool_result") or {}
        self.assertEqual(tool_data.get("ready_to_ship"), 2)
        self.assertEqual(tool_data.get("not_printed"), 1)
        self.assertTrue(tool_data.get("only_not_printed"))
        self.assertEqual([row.get("order_number") for row in (tool_data.get("orders") or [])], ["ORD-READY-UNPRINTED"])

    def test_cargo_filtered_orders_use_query_filters_not_generic_lookup(self) -> None:
        now_ms = int(datetime.now(TZ_GMT3).timestamp() * 1000)
        self._insert_order(
            "ORD-ARAS-1",
            now_ms,
            package_id=1_101,
            first_name="Aras",
            last_name="One",
            status="ReadyToShip",
            cargo_provider_name="Aras",
        )
        self._insert_order(
            "ORD-TEX-1",
            now_ms + 1_000,
            package_id=1_102,
            first_name="Tex",
            last_name="One",
            status="ReadyToShip",
            cargo_provider_name="Trendyol Express",
        )
        self._insert_order(
            "ORD-ARAS-2",
            now_ms + 2_000,
            package_id=1_103,
            first_name="Aras",
            last_name="Two",
            status="Created",
            cargo_provider_name="Aras",
        )

        result = run_agent_turn("Aras kargolu siparisleri getir", db_name=self.db_path)

        self.assertTrue(result.success, result.message)
        self.assertEqual(((result.data or {}).get("tool_call") or {}).get("tool"), "orders.query")
        args = ((result.data or {}).get("tool_call") or {}).get("arguments") or {}
        self.assertEqual((args.get("filters") or {}).get("cargo_provider"), "Aras")
        orders = ((result.data or {}).get("tool_result") or {}).get("orders") or []
        self.assertEqual([row.get("order_number") for row in orders], ["ORD-ARAS-2", "ORD-ARAS-1"])
        self.assertTrue(all(row.get("cargo_provider") == "Aras" for row in orders), orders)

    def test_cargo_provider_options_prompt_lists_manual_change_options(self) -> None:
        result = run_agent_turn("Trendyolda kullanabilecegim kargo firmalari neler?", db_name=self.db_path)

        self.assertTrue(result.success, result.message)
        self.assertEqual(((result.data or {}).get("tool_call") or {}).get("tool"), "orders.cargo_provider_options")
        answer = (result.data or {}).get("answer") or result.message
        self.assertIn("TRENDYOL", answer)
        self.assertIn("Aras", answer)
        tool_data = (result.data or {}).get("tool_result") or {}
        self.assertEqual(tool_data.get("platform_code"), "TRENDYOL")
        providers = tool_data.get("providers") or []
        self.assertTrue(any(row.get("code") == "ARASMP" for row in providers), providers)

    def test_nonfinal_order_numbers_prompt_lists_only_latest_nonfinal_orders(self) -> None:
        now_ms = int(datetime.now(TZ_GMT3).timestamp() * 1000)
        self._insert_order(
            "ORD-NONFINAL-CREATED",
            now_ms,
            package_id=1_151,
            first_name="Nonfinal",
            last_name="Created",
            status="Created",
        )
        self._insert_order(
            "ORD-NONFINAL-READY",
            now_ms + 1_000,
            package_id=1_152,
            first_name="Nonfinal",
            last_name="Ready",
            status="ReadyToShip",
        )
        self._insert_order(
            "ORD-FINAL-DELIVERED",
            now_ms + 2_000,
            package_id=1_153,
            first_name="Final",
            last_name="Delivered",
            status="Delivered",
        )
        self._insert_order(
            "ORD-FINAL-CANCELLED",
            now_ms + 3_000,
            package_id=1_154,
            first_name="Final",
            last_name="Cancelled",
            status="Cancelled",
        )
        self._insert_order(
            "ORD-FINAL-RETURNED",
            now_ms + 4_000,
            package_id=1_155,
            first_name="Final",
            last_name="Returned",
            status="UnDeliveredAndReturned",
        )

        result = run_agent_turn("Final olmayan siparis numaralarini getir", db_name=self.db_path)

        self.assertTrue(result.success, result.message)
        self.assertEqual(((result.data or {}).get("tool_call") or {}).get("tool"), "orders.nonfinal_order_numbers")
        answer = (result.data or {}).get("answer") or result.message
        self.assertIn("ORD-NONFINAL-CREATED", answer)
        self.assertIn("ORD-NONFINAL-READY", answer)
        self.assertNotIn("ORD-FINAL-DELIVERED", answer)
        self.assertNotIn("ORD-FINAL-CANCELLED", answer)
        self.assertNotIn("ORD-FINAL-RETURNED", answer)
        tool_data = (result.data or {}).get("tool_result") or {}
        self.assertEqual(set(tool_data.get("order_numbers") or []), {"ORD-NONFINAL-CREATED", "ORD-NONFINAL-READY"})

    def test_order_detail_prompt_includes_items_and_return_context(self) -> None:
        now_ms = int(datetime.now(TZ_GMT3).timestamp() * 1000)
        self._insert_order(
            "ORD-DETAIL-CONTEXT",
            now_ms,
            package_id=1_181,
            first_name="Detail",
            last_name="Context",
            status="ReadyToShip",
            cargo_provider_name="Aras",
        )
        self._insert_order_item(
            "ORD-DETAIL-CONTEXT",
            now_ms + 1_000,
            line_id=10,
            product_name="Detay Test Urunu",
        )
        self._insert_waiting_return(
            "CLAIM-DETAIL-CONTEXT",
            "ORD-DETAIL-CONTEXT",
            now_ms + 2_000,
            first_name="Detail",
            last_name="Context",
        )

        result = run_agent_turn("ORD-DETAIL-CONTEXT siparisinin detayini ve kalemlerini getir", db_name=self.db_path)

        self.assertTrue(result.success, result.message)
        self.assertEqual(((result.data or {}).get("tool_call") or {}).get("tool"), "orders.full_context")
        answer = (result.data or {}).get("answer") or result.message
        self.assertIn("ORD-DETAIL-CONTEXT", answer)
        self.assertIn("Detay Test Urunu", answer)
        self.assertIn("CLAIM-DETAIL-CONTEXT", answer)
        tool_data = (result.data or {}).get("tool_result") or {}
        context = ((tool_data.get("contexts") or [])[0] or {})
        self.assertEqual(((context.get("items") or [])[0] or {}).get("product_name"), "Detay Test Urunu")
        self.assertEqual(((context.get("returns") or [])[0] or {}).get("claim_id"), "CLAIM-DETAIL-CONTEXT")

    def test_followup_customer_refined_order_items_uses_previous_order_context(self) -> None:
        now_ms = self._seed_ready_orders()
        self._insert_order_item("ORD-MAHMUT-READY", now_ms + 1_000, line_id=21, product_name="Mahmut Kalem Urunu")
        self._insert_order_item("ORD-AYSE-READY", now_ms + 2_000, line_id=22, product_name="Ayse Kalem Urunu")
        self._insert_waiting_return(
            "CLAIM-MAHMUT-ITEM-CONTEXT",
            "ORD-MAHMUT-READY",
            now_ms + 3_000,
            first_name="Mahmut",
            last_name="Sevket",
        )

        listed = run_agent_turn("Kargoya hazir siparisleri getir", db_name=self.db_path)
        self.assertTrue(listed.success, listed.message)
        listed_answer = self._answer_from(listed)
        self.assertIn("ORD-MAHMUT-READY", listed_answer)
        self.assertIn("ORD-AYSE-READY", listed_answer)

        history = [
            {"role": "user", "content": "Kargoya hazir siparisleri getir"},
            {"role": "assistant", "content": listed_answer},
        ]
        result = run_agent_turn("Mahmut olanin kalemlerini getir", db_name=self.db_path, history=history)

        self.assertTrue(result.success, result.message)
        self.assertEqual(((result.data or {}).get("tool_call") or {}).get("tool"), "orders.full_context")
        answer = self._answer_from(result)
        self.assertIn("ORD-MAHMUT-READY", answer)
        self.assertIn("Mahmut Kalem Urunu", answer)
        self.assertIn("CLAIM-MAHMUT-ITEM-CONTEXT", answer)
        self.assertNotIn("ORD-AYSE-READY", answer)
        self.assertNotIn("Ayse Kalem Urunu", answer)
        tool_data = (result.data or {}).get("tool_result") or {}
        context = ((tool_data.get("contexts") or [])[0] or {})
        self.assertEqual(((context.get("order") or {}).get("order_number")), "ORD-MAHMUT-READY")
        self.assertEqual(((context.get("items") or [])[0] or {}).get("product_name"), "Mahmut Kalem Urunu")
        self.assertEqual(((context.get("returns") or [])[0] or {}).get("claim_id"), "CLAIM-MAHMUT-ITEM-CONTEXT")

    def test_realistic_customer_refinement_and_pronoun_keep_one_structured_order(self) -> None:
        now_ms = int(datetime.now(TZ_GMT3).timestamp() * 1000)
        target_order = "ORD-RECENT-TEKRAR"
        self._insert_order(
            target_order,
            now_ms,
            package_id=12_101,
            first_name="Tekrar",
            last_name="Musteri",
            status="ReadyToShip",
        )
        self._insert_order_item(
            target_order,
            now_ms + 1_000,
            line_id=101,
            product_name="Hedef Kalem",
        )
        for index in range(4):
            self._insert_order(
                f"ORD-RECENT-DIGER-{index}",
                now_ms - ((index + 1) * 1_000),
                package_id=12_110 + index,
                first_name=f"Diger{index}",
                last_name="Musteri",
                status="ReadyToShip",
            )
        old_ms = now_ms - int(timedelta(days=20).total_seconds() * 1_000)
        for index in range(4):
            old_order = f"ORD-OLD-TEKRAR-{index}"
            self._insert_order(
                old_order,
                old_ms - (index * 1_000),
                package_id=12_120 + index,
                first_name="Tekrar",
                last_name="Musteri",
                status="Delivered",
            )
            self._insert_order_item(
                old_order,
                old_ms + index,
                line_id=120 + index,
                product_name=f"Eski Kalem {index}",
            )

        listed = run_agent_turn("son gelen 5 siparisi getir", db_name=self.db_path)
        self.assertTrue(listed.success, listed.message)
        listed_data = dict(listed.data or {})
        history = [
            {"role": "user", "content": "son gelen 5 siparisi getir"},
            {
                "role": "assistant",
                "content": self._answer_from(listed),
                "context": dict(listed_data.get("conversation_context") or {}),
            },
        ]
        refined = run_agent_turn(
            (
                "bu listedekilerden Tekrar Musteri olanin kalemlerini "
                "ve iade durumunu getir"
            ),
            db_name=self.db_path,
            history=history,
        )

        self.assertTrue(refined.success, refined.message)
        self.assertEqual(
            ((refined.data or {}).get("tool_call") or {}).get("tool"),
            "orders.full_context",
        )
        refined_contexts = (
            ((refined.data or {}).get("tool_result") or {}).get("contexts")
            or []
        )
        self.assertEqual(len(refined_contexts), 1)
        self.assertEqual(
            ((refined_contexts[0].get("order") or {}).get("order_number")),
            target_order,
        )
        self.assertIn("Hedef Kalem", self._answer_from(refined))
        self.assertNotIn("Eski Kalem", self._answer_from(refined))

        lookup = run_agent_turn(
            f"siparis numarasi {target_order} olan siparisi goster",
            db_name=self.db_path,
        )
        self.assertTrue(lookup.success, lookup.message)
        lookup_data = dict(lookup.data or {})
        lookup_history = [
            {
                "role": "user",
                "content": (
                    f"siparis numarasi {target_order} olan siparisi goster"
                ),
            },
            {
                "role": "assistant",
                "content": self._answer_from(lookup),
                "context": dict(lookup_data.get("conversation_context") or {}),
            },
        ]
        pronoun = run_agent_turn(
            "bunun kalemlerini ve iade durumunu tekrar ozetle",
            db_name=self.db_path,
            history=lookup_history,
        )

        self.assertTrue(pronoun.success, pronoun.message)
        self.assertEqual(
            ((pronoun.data or {}).get("tool_call") or {}).get("tool"),
            "orders.full_context",
        )
        pronoun_contexts = (
            ((pronoun.data or {}).get("tool_result") or {}).get("contexts")
            or []
        )
        self.assertEqual(len(pronoun_contexts), 1)
        self.assertEqual(
            ((pronoun_contexts[0].get("order") or {}).get("order_number")),
            target_order,
        )

    def test_ready_period_list_repairs_planner_recent_limit_confusion(self) -> None:
        now_ms = int(datetime.now(TZ_GMT3).timestamp() * 1000)
        for index in range(3):
            self._insert_order(
                f"ORD-READY-PERIOD-{index}",
                now_ms - (index * 1_000),
                package_id=12_200 + index,
                first_name=f"Ready{index}",
                last_name="Period",
                status="ReadyToShip",
            )
        self._insert_order(
            "ORD-DELIVERED-PERIOD",
            now_ms - 4_000,
            package_id=12_210,
            first_name="Delivered",
            last_name="Period",
            status="Delivered",
        )

        result = run_agent_turn(
            "son 8 gunde kargoya hazir siparislerin tamamini listele",
            planner_client=lambda _prompt: (
                '{"mode":"tool","tool":"orders.recent",'
                '"arguments":{"period":"last_8_days","limit":8}}'
            ),
            chat_client=lambda _prompt: "Sekiz siparis var.",
            db_name=self.db_path,
        )

        self.assertTrue(result.success, result.message)
        tool_call = (result.data or {}).get("tool_call") or {}
        self.assertEqual(tool_call.get("tool"), "orders.ready_to_ship")
        self.assertEqual(
            tool_call.get("arguments"),
            {"period": "last_8_days", "limit": 100},
        )
        tool_data = (result.data or {}).get("tool_result") or {}
        self.assertEqual(
            set(tool_data.get("order_numbers") or []),
            {
                "ORD-READY-PERIOD-0",
                "ORD-READY-PERIOD-1",
                "ORD-READY-PERIOD-2",
            },
        )
        self.assertNotIn("ORD-DELIVERED-PERIOD", self._answer_from(result))

    def test_fulfillment_risk_prompt_lists_orders_and_reasons(self) -> None:
        now = datetime.now(TZ_GMT3)
        old_ms = int((now - timedelta(days=2)).timestamp() * 1000)
        fresh_ms = int((now - timedelta(hours=2)).timestamp() * 1000)
        self._insert_order(
            "ORD-RISK-MISSING",
            old_ms,
            package_id=1_191,
            first_name="Risk",
            last_name="Missing",
            status="ReadyToShip",
            cargo_provider_name="",
            cargo_tracking_number="",
        )
        self._insert_order(
            "ORD-RISK-HEALTHY",
            fresh_ms,
            package_id=1_192,
            first_name="Risk",
            last_name="Healthy",
            status="ReadyToShip",
            cargo_provider_name="Aras",
            cargo_tracking_number="TRK-HEALTHY",
        )
        self._insert_order(
            "ORD-RISK-CANCELLED",
            old_ms + 1_000,
            package_id=1_193,
            first_name="Risk",
            last_name="Cancelled",
            status="Cancelled",
            cargo_provider_name="",
            cargo_tracking_number="",
        )

        result = run_agent_turn("Kargo takip veya firma eksigi olan riskli siparisleri goster", db_name=self.db_path)

        self.assertTrue(result.success, result.message)
        self.assertEqual(((result.data or {}).get("tool_call") or {}).get("tool"), "orders.fulfillment_risk")
        answer = (result.data or {}).get("answer") or result.message
        self.assertIn("ORD-RISK-MISSING", answer)
        self.assertIn("tracking_missing", answer)
        self.assertIn("cargo_provider_missing", answer)
        self.assertIn("order_age_over_threshold_while_ready", answer)
        self.assertNotIn("ORD-RISK-HEALTHY", answer)
        self.assertNotIn("ORD-RISK-CANCELLED", answer)
        orders = ((result.data or {}).get("tool_result") or {}).get("orders") or []
        self.assertEqual([row.get("order_number") for row in orders], ["ORD-RISK-MISSING"])
        self.assertEqual(
            set((orders[0] or {}).get("reasons") or []),
            {"tracking_missing", "cargo_provider_missing", "order_age_over_threshold_while_ready"},
        )

    def test_cancelled_orders_prompt_lists_orders_and_cancel_reasons(self) -> None:
        now_ms = int(datetime.now(TZ_GMT3).timestamp() * 1000)
        self._insert_order(
            "ORD-CANCEL-REASON",
            now_ms,
            package_id=1_195,
            first_name="Cancel",
            last_name="Reason",
            status="Cancelled",
        )
        self._insert_order_item(
            "ORD-CANCEL-REASON",
            now_ms + 1_000,
            line_id=12,
            product_name="Iptal Test Urunu",
            cancel_reason="Musteri vazgecti",
        )
        self._insert_order(
            "ORD-CANCEL-ACTIVE",
            now_ms + 2_000,
            package_id=1_196,
            first_name="Cancel",
            last_name="Active",
            status="ReadyToShip",
        )

        result = run_agent_turn("Bugunku iptal siparisleri ve nedenlerini goster", db_name=self.db_path)

        self.assertTrue(result.success, result.message)
        self.assertEqual(((result.data or {}).get("tool_call") or {}).get("tool"), "orders.cancellations")
        answer = (result.data or {}).get("answer") or result.message
        self.assertIn("ORD-CANCEL-REASON", answer)
        self.assertIn("Musteri vazgecti", answer)
        self.assertNotIn("ORD-CANCEL-ACTIVE", answer)
        tool_data = (result.data or {}).get("tool_result") or {}
        self.assertEqual(tool_data.get("cancelled_orders"), 1)
        self.assertEqual((tool_data.get("cancel_reasons") or {}).get("Musteri vazgecti"), 1)

    def test_today_created_orders_use_status_and_period_filters(self) -> None:
        now = datetime.now(TZ_GMT3)
        today_ms = int(now.timestamp() * 1000)
        yesterday_ms = int((now - timedelta(days=1)).timestamp() * 1000)
        self._insert_order(
            "ORD-TODAY-CREATED",
            today_ms,
            package_id=1_201,
            first_name="Today",
            last_name="Created",
            status="Created",
        )
        self._insert_order(
            "ORD-TODAY-READY",
            today_ms + 1_000,
            package_id=1_202,
            first_name="Today",
            last_name="Ready",
            status="ReadyToShip",
        )
        self._insert_order(
            "ORD-YESTERDAY-CREATED",
            yesterday_ms,
            package_id=1_203,
            first_name="Yesterday",
            last_name="Created",
            status="Created",
        )

        result = run_agent_turn("Bugunku durumu Created olan siparisleri getir", db_name=self.db_path)

        self.assertTrue(result.success, result.message)
        self.assertEqual(((result.data or {}).get("tool_call") or {}).get("tool"), "orders.query")
        args = ((result.data or {}).get("tool_call") or {}).get("arguments") or {}
        self.assertEqual(args.get("period"), "today")
        self.assertEqual((args.get("filters") or {}).get("status"), "Created")
        orders = ((result.data or {}).get("tool_result") or {}).get("orders") or []
        self.assertEqual([row.get("order_number") for row in orders], ["ORD-TODAY-CREATED"])

    def test_platform_filtered_orders_use_api_account_platform(self) -> None:
        now_ms = int(datetime.now(TZ_GMT3).timestamp() * 1000)
        with Session(get_engine(self.db_path)) as session:
            session.add(ApiAccount(pk=1, account_id="ty-agent", comp_name="TY Agent", platform="trendyol"))
            session.add(ApiAccount(pk=2, account_id="amz-agent", comp_name="AMZ Agent", platform="amazon"))
            session.commit()
        self._insert_order(
            "ORD-TY-PLATFORM",
            now_ms,
            package_id=1_301,
            first_name="Trend",
            last_name="Yol",
            status="Created",
            api_account_id=1,
        )
        self._insert_order(
            "ORD-AMZ-PLATFORM",
            now_ms + 1_000,
            package_id=1_302,
            first_name="Ama",
            last_name="Zon",
            status="Created",
            api_account_id=2,
        )

        result = run_agent_turn("Trendyol pazaryeri siparislerini getir", db_name=self.db_path)

        self.assertTrue(result.success, result.message)
        self.assertEqual(((result.data or {}).get("tool_call") or {}).get("tool"), "orders.query")
        args = ((result.data or {}).get("tool_call") or {}).get("arguments") or {}
        self.assertEqual((args.get("filters") or {}).get("platform_code"), "TRENDYOL")
        orders = ((result.data or {}).get("tool_result") or {}).get("orders") or []
        self.assertEqual([row.get("order_number") for row in orders], ["ORD-TY-PLATFORM"])
        self.assertEqual([row.get("platform_code") for row in orders], ["TRENDYOL"])

    def test_highest_priced_order_prompt_returns_top_non_cancelled_order(self) -> None:
        now_ms = int(datetime.now(TZ_GMT3).timestamp() * 1000)
        self._insert_order(
            "ORD-TOP-LOW",
            now_ms,
            package_id=1_351,
            first_name="Top",
            last_name="Low",
            status="Created",
            total_price=90.0,
        )
        self._insert_order(
            "ORD-TOP-HIGH",
            now_ms + 1_000,
            package_id=1_352,
            first_name="Top",
            last_name="High",
            status="ReadyToShip",
            total_price=350.0,
        )
        self._insert_order(
            "ORD-TOP-CANCELLED",
            now_ms + 2_000,
            package_id=1_353,
            first_name="Top",
            last_name="Cancelled",
            status="Cancelled",
            total_price=999.0,
        )

        result = run_agent_turn("Bugun en yuksek tutarli siparisimi getir", db_name=self.db_path)

        self.assertTrue(result.success, result.message)
        self.assertEqual(((result.data or {}).get("tool_call") or {}).get("tool"), "orders.query")
        args = ((result.data or {}).get("tool_call") or {}).get("arguments") or {}
        self.assertEqual(args.get("sort"), "total_price_desc")
        self.assertEqual(args.get("mode"), "top")
        self.assertTrue((args.get("filters") or {}).get("exclude_cancelled"))
        answer = (result.data or {}).get("answer") or result.message
        self.assertIn("ORD-TOP-HIGH", answer)
        self.assertIn("350", answer)
        self.assertNotIn("ORD-TOP-CANCELLED", answer)
        orders = ((result.data or {}).get("tool_result") or {}).get("orders") or []
        self.assertEqual([row.get("order_number") for row in orders], ["ORD-TOP-HIGH"])

    def test_waiting_returns_pagination_returns_requested_page_slice(self) -> None:
        now_ms = int(datetime.now(TZ_GMT3).timestamp() * 1000)
        for idx in range(1, 6):
            self._insert_waiting_return(
                f"CLAIM-PAGE-{idx}",
                f"ORD-RETURN-PAGE-{idx}",
                now_ms + (idx * 1_000),
                first_name=f"Return{idx}",
                last_name="Customer",
            )

        result = run_agent_turn(
            "Aksiyon bekleyen iadelerin ikinci sayfasini getir, sayfa basi 2 kayit",
            db_name=self.db_path,
        )

        self.assertTrue(result.success, result.message)
        self.assertEqual(((result.data or {}).get("tool_call") or {}).get("tool"), "returns.query")
        args = ((result.data or {}).get("tool_call") or {}).get("arguments") or {}
        self.assertEqual(args.get("page"), 2)
        self.assertEqual(args.get("page_size"), 2)
        tool_data = (result.data or {}).get("tool_result") or {}
        pagination = tool_data.get("pagination") or {}
        self.assertEqual(pagination.get("page"), 2)
        self.assertEqual(pagination.get("returned"), 2)
        returns = tool_data.get("returns") or []
        self.assertEqual([row.get("claim_id") for row in returns], ["CLAIM-PAGE-3", "CLAIM-PAGE-2"])

    def test_accepted_returns_use_status_and_period_filters(self) -> None:
        now = datetime.now(TZ_GMT3)
        today_ms = int(now.timestamp() * 1000)
        yesterday_ms = int((now - timedelta(days=1)).timestamp() * 1000)
        self._insert_waiting_return(
            "CLAIM-TODAY-WAITING",
            "ORD-RET-TODAY-WAITING",
            today_ms,
            first_name="Today",
            last_name="Waiting",
        )
        self._insert_waiting_return(
            "CLAIM-TODAY-ACCEPTED",
            "ORD-RET-TODAY-ACCEPTED",
            today_ms + 1_000,
            first_name="Today",
            last_name="Accepted",
        )
        self._insert_waiting_return(
            "CLAIM-YESTERDAY-ACCEPTED",
            "ORD-RET-YESTERDAY-ACCEPTED",
            yesterday_ms,
            first_name="Yesterday",
            last_name="Accepted",
        )
        with Session(get_engine(self.db_path)) as session:
            for claim_id in ("CLAIM-TODAY-ACCEPTED", "CLAIM-YESTERDAY-ACCEPTED"):
                claim = session.exec(select(ReturnClaim).where(ReturnClaim.claimId == claim_id)).one()
                claim.claimItemStatus = "Accepted"
                claim.waitingClaimLineItemIds = []
                session.add(claim)
            session.commit()

        result = run_agent_turn("Bugunku durumu Accepted olan iadeleri getir", db_name=self.db_path)

        self.assertTrue(result.success, result.message)
        self.assertEqual(((result.data or {}).get("tool_call") or {}).get("tool"), "returns.query")
        args = ((result.data or {}).get("tool_call") or {}).get("arguments") or {}
        self.assertEqual(args.get("period"), "today")
        self.assertEqual((args.get("filters") or {}).get("status"), "Accepted")
        returns = ((result.data or {}).get("tool_result") or {}).get("returns") or []
        self.assertEqual([row.get("claim_id") for row in returns], ["CLAIM-TODAY-ACCEPTED"])

    def test_platform_filtered_returns_use_api_account_platform(self) -> None:
        now_ms = int(datetime.now(TZ_GMT3).timestamp() * 1000)
        with Session(get_engine(self.db_path)) as session:
            session.add(ApiAccount(pk=1, account_id="ty-return-agent", comp_name="TY Return Agent", platform="trendyol"))
            session.add(ApiAccount(pk=2, account_id="amz-return-agent", comp_name="AMZ Return Agent", platform="amazon"))
            session.commit()
        self._insert_waiting_return(
            "CLAIM-TY-PLATFORM",
            "ORD-TY-RETURN-PLATFORM",
            now_ms,
            first_name="Trend",
            last_name="Return",
            api_account_id=1,
        )
        self._insert_waiting_return(
            "CLAIM-AMZ-PLATFORM",
            "ORD-AMZ-RETURN-PLATFORM",
            now_ms + 1_000,
            first_name="Ama",
            last_name="Return",
            api_account_id=2,
        )

        result = run_agent_turn("Amazon pazaryeri iadelerini getir", db_name=self.db_path)

        self.assertTrue(result.success, result.message)
        self.assertEqual(((result.data or {}).get("tool_call") or {}).get("tool"), "returns.query")
        args = ((result.data or {}).get("tool_call") or {}).get("arguments") or {}
        self.assertEqual((args.get("filters") or {}).get("platform_code"), "AMAZON")
        returns = ((result.data or {}).get("tool_result") or {}).get("returns") or []
        self.assertEqual([row.get("claim_id") for row in returns], ["CLAIM-AMZ-PLATFORM"])
        self.assertEqual([row.get("platform_code") for row in returns], ["AMAZON"])
        context = dict((result.data or {}).get("conversation_context") or {})
        self.assertEqual(context.get("kind"), "return_selection")
        self.assertEqual(context.get("account_by_claim"), {"CLAIM-AMZ-PLATFORM": "2"})
        self.assertEqual(context.get("environment_by_account"), {"2": "unknown"})
        self.assertEqual(
            [row.get("store_name") for row in context.get("account_scopes") or []],
            ["AMZ Return Agent"],
        )

    def test_return_amount_prompt_uses_returns_total_not_order_revenue(self) -> None:
        now_ms = int(datetime.now(TZ_GMT3).timestamp() * 1000)
        self._insert_waiting_return(
            "CLAIM-AMOUNT-1",
            "ORD-RETURN-AMOUNT-1",
            now_ms,
            first_name="Amount",
            last_name="One",
        )
        self._insert_waiting_return(
            "CLAIM-AMOUNT-2",
            "ORD-RETURN-AMOUNT-2",
            now_ms + 1_000,
            first_name="Amount",
            last_name="Two",
        )
        with Session(get_engine(self.db_path)) as session:
            values = {"CLAIM-AMOUNT-1": 40.0, "CLAIM-AMOUNT-2": 60.5}
            for claim in session.exec(select(ReturnClaim).where(ReturnClaim.claimId.in_(list(values)))).all():
                claim.totalPrice = values[claim.claimId]
                session.add(claim)
            session.commit()

        result = run_agent_turn("Bugunku iadelerin toplam tutari ne kadar?", db_name=self.db_path)

        self.assertTrue(result.success, result.message)
        self.assertEqual(((result.data or {}).get("tool_call") or {}).get("tool"), "returns.amount_total")
        tool_data = (result.data or {}).get("tool_result") or {}
        self.assertEqual(tool_data.get("total_amount"), 100.5)
        answer = self._answer_from(result)
        self.assertIn("100.5", answer)
        self.assertIn("TRY", answer)

    def test_generic_print_ready_orders_executes_label_creation_and_audit(self) -> None:
        self._seed_ready_orders()
        calls: list[list[str]] = []

        def fake_create_order_label_from_orders(order_numbers, *, brand_code="TANEX", model_code="TANEX_2736", db_name=None):
            calls.append(list(order_numbers))
            return Result.ok(
                "Etiketler olusturuldu.",
                close_dialog=False,
                data={"order_numbers": list(order_numbers), "brand_code": brand_code, "model_code": model_code},
            )

        prepared = run_agent_turn("Kargoya hazır siparişleri yazdır", db_name=self.db_path)
        action = self._action_from(prepared)

        with tempfile.TemporaryDirectory() as temp_dir, patch(
            "Labels.processors.pipeline.create_order_label_from_orders",
            side_effect=fake_create_order_label_from_orders,
        ):
            audit_path = Path(temp_dir) / "orders-agent-audit.jsonl"
            executed = operation_runtime.execute_approved_operation(
                action,
                approved_by="behavioral-test",
                idempotency_key="print-ready-orders",
                audit_path=audit_path,
            )

            self.assertTrue(executed.success, executed.message)
            self.assertEqual(calls, [["ORD-AYSE-READY", "ORD-MAHMUT-READY"]])
            self.assertTrue(audit_path.exists())
            self.assertIn("labels.action_create_labels_from_selected_orders", audit_path.read_text(encoding="utf-8"))
            verify = ((executed.data or {}).get("verify") or {})
            self.assertTrue(verify.get("verified"), executed.data)
            self.assertEqual(verify.get("order_numbers"), ["ORD-AYSE-READY", "ORD-MAHMUT-READY"])

    def test_explicit_order_label_prompt_executes_only_that_order(self) -> None:
        self._seed_ready_orders()
        calls: list[list[str]] = []

        def fake_create_order_label_from_orders(order_numbers, *, brand_code="TANEX", model_code="TANEX_2736", db_name=None):
            calls.append(list(order_numbers))
            return Result.ok(
                "Etiketler olusturuldu.",
                close_dialog=False,
                data={"order_numbers": list(order_numbers), "brand_code": brand_code, "model_code": model_code},
            )

        prepared = run_agent_turn("ORD-MAHMUT-READY etiketini yazdir", db_name=self.db_path)

        self._assert_prepare_answer_is_safe(prepared, expected_terms=("1 siparis",))
        self.assertEqual(((prepared.data or {}).get("tool_call") or {}).get("tool"), "labels.create_from_orders.prepare")
        action = self._action_from(prepared)
        payload = ((action.get("arguments") or {}).get("payload") or {})
        self.assertEqual(payload.get("order_numbers"), ["ORD-MAHMUT-READY"])
        self.assertNotIn("ORD-AYSE-READY", payload.get("order_numbers") or [])

        with tempfile.TemporaryDirectory() as temp_dir, patch(
            "Labels.processors.pipeline.create_order_label_from_orders",
            side_effect=fake_create_order_label_from_orders,
        ):
            audit_path = Path(temp_dir) / "orders-agent-audit.jsonl"
            executed = operation_runtime.execute_approved_operation(
                action,
                approved_by="behavioral-test",
                idempotency_key="print-explicit-order",
                audit_path=audit_path,
            )

            self._assert_execute_verified(executed)
            self.assertEqual(calls, [["ORD-MAHMUT-READY"]])
            self.assertIn("labels.action_create_labels_from_selected_orders", audit_path.read_text(encoding="utf-8"))

    def test_order_profit_excel_export_executes_file_write_and_audit(self) -> None:
        now_ms = self._seed_ready_orders()
        self._insert_order_item("ORD-MAHMUT-READY", now_ms + 2_000, product_name="Mahmut Test Urunu")

        with tempfile.TemporaryDirectory() as temp_dir:
            output_path = Path(temp_dir) / "mahmut-profit.xlsx"
            prompt_path = str(output_path).replace("\\", "/")
            prepared = run_agent_turn(
                f'ORD-MAHMUT-READY siparis kar Excel raporunu "{prompt_path}" dosyasina aktar',
                db_name=self.db_path,
            )
            action = self._action_from(prepared)
            self.assertEqual(action.get("action_type"), "orders.action_export_order_profit_excel")

            audit_path = Path(temp_dir) / "orders-agent-audit.jsonl"
            executed = operation_runtime.execute_approved_operation(
                action,
                approved_by="behavioral-test",
                idempotency_key="profit-export-mahmut",
                audit_path=audit_path,
            )

            self.assertTrue(executed.success, executed.message)
            self.assertTrue(output_path.exists(), executed.data)
            self.assertGreater(output_path.stat().st_size, 0)
            self.assertIn("orders.action_export_order_profit_excel", audit_path.read_text(encoding="utf-8"))
            verify = ((executed.data or {}).get("verify") or {})
            self.assertTrue(verify.get("verified"), executed.data)
            self.assertEqual(Path(verify.get("output_path")).name, "mahmut-profit.xlsx")

            from openpyxl import load_workbook

            workbook = load_workbook(output_path, read_only=True)
            try:
                self.assertEqual(set(workbook.sheetnames), {"Ozet", "Satis", "Maliyetler"})
                summary_rows = list(workbook["Ozet"].iter_rows(values_only=True))
                self.assertIn(("Sipariş", "ORD-MAHMUT-READY"), summary_rows)
                sales_rows = list(workbook["Satis"].iter_rows(values_only=True))
                self.assertTrue(any("Mahmut Test Urunu" in str(cell) for row in sales_rows for cell in row), sales_rows)
            finally:
                workbook.close()

    def test_followup_customer_refined_profit_excel_export_writes_that_order_only(self) -> None:
        now_ms = self._seed_ready_orders()
        self._insert_order_item("ORD-MAHMUT-READY", now_ms + 2_000, product_name="Mahmut Test Urunu")
        self._insert_order_item("ORD-AYSE-READY", now_ms + 3_000, product_name="Ayse Test Urunu")

        listed = run_agent_turn("Kargoya hazir siparisleri getir", db_name=self.db_path)
        self.assertTrue(listed.success, listed.message)
        listed_answer = self._answer_from(listed)
        self.assertIn("ORD-MAHMUT-READY", listed_answer)
        self.assertIn("ORD-AYSE-READY", listed_answer)

        history = [
            {"role": "user", "content": "Kargoya hazir siparisleri getir"},
            {"role": "assistant", "content": listed_answer},
        ]
        with tempfile.TemporaryDirectory() as temp_dir:
            output_path = Path(temp_dir) / "mahmut-followup-profit.xlsx"
            prompt_path = str(output_path).replace("\\", "/")
            prepared = run_agent_turn(
                f'Mahmut olanin kar Excel raporunu "{prompt_path}" dosyasina aktar',
                db_name=self.db_path,
                history=history,
            )
            self._assert_prepare_answer_is_safe(prepared, expected_terms=("ORD-MAHMUT-READY", "Excel"))
            self.assertNotIn("ORD-AYSE-READY", self._answer_from(prepared))
            action = self._action_from(prepared)
            payload = ((action.get("arguments") or {}).get("payload") or {})
            self.assertEqual(payload.get("order_number"), "ORD-MAHMUT-READY")

            audit_path = Path(temp_dir) / "orders-agent-audit.jsonl"
            executed = operation_runtime.execute_approved_operation(
                action,
                approved_by="behavioral-test",
                idempotency_key="followup-profit-export-mahmut",
                audit_path=audit_path,
            )

            self.assertTrue(executed.success, executed.message)
            self.assertTrue(output_path.exists(), executed.data)
            self.assertGreater(output_path.stat().st_size, 0)
            verify = ((executed.data or {}).get("verify") or {})
            self.assertTrue(verify.get("verified"), executed.data)
            self.assertEqual(Path(verify.get("output_path")).name, "mahmut-followup-profit.xlsx")

            from openpyxl import load_workbook

            workbook = load_workbook(output_path, read_only=True)
            try:
                summary_rows = list(workbook["Ozet"].iter_rows(values_only=True))
                self.assertTrue(any(row[1] == "ORD-MAHMUT-READY" for row in summary_rows), summary_rows)
                sales_rows = list(workbook["Satis"].iter_rows(values_only=True))
                self.assertTrue(any("Mahmut Test Urunu" in str(cell) for row in sales_rows for cell in row), sales_rows)
                self.assertFalse(any("Ayse Test Urunu" in str(cell) for row in sales_rows for cell in row), sales_rows)
            finally:
                workbook.close()

    def test_followup_profit_export_uses_customer_name_when_history_has_only_order_ids(self) -> None:
        now_ms = self._seed_ready_orders()
        self._insert_order_item("ORD-MAHMUT-READY", now_ms + 2_000, product_name="Mahmut Test Urunu")
        self._insert_order_item("ORD-AYSE-READY", now_ms + 3_000, product_name="Ayse Test Urunu")

        history = [
            {"role": "user", "content": "Kargoya hazir siparisleri getir"},
            {
                "role": "assistant",
                "content": (
                    "Lokal kayda gore 2 siparis buldum:\n"
                    "1. ORD-MAHMUT-READY | durum: ReadyToShip\n"
                    "2. ORD-AYSE-READY | durum: ReadyToShip"
                ),
            },
        ]
        with tempfile.TemporaryDirectory() as temp_dir:
            output_path = Path(temp_dir) / "mahmut-history-id-only-profit.xlsx"
            prompt_path = str(output_path).replace("\\", "/")
            prepared = run_agent_turn(
                f'Mahmut olanin kar Excel raporunu "{prompt_path}" dosyasina aktar',
                db_name=self.db_path,
                history=history,
            )

        self._assert_prepare_answer_is_safe(prepared, expected_terms=("ORD-MAHMUT-READY", "Excel"))
        action = self._action_from(prepared)
        self.assertEqual(action.get("action_type"), "orders.action_export_order_profit_excel")
        payload = ((action.get("arguments") or {}).get("payload") or {})
        self.assertEqual(payload.get("order_number"), "ORD-MAHMUT-READY")
        self.assertNotIn("Birden fazla", self._answer_from(prepared))

    def test_orders_fetch_prompt_prepares_active_platform_accounts_and_execute_saves(self) -> None:
        with Session(get_engine(self.db_path)) as session:
            trendyol = ApiAccount(
                account_id="seller-trendyol-sync",
                comp_name="Trendyol Sync Seller",
                platform="trendyol",
                api_key="api-key-sync",
                api_secret="api-secret-sync",
                is_active=True,
            )
            session.add(trendyol)
            session.add(
                ApiAccount(
                    account_id="seller-amazon-other",
                    comp_name="Amazon Other Seller",
                    platform="amazon",
                    api_key="amz-key",
                    api_secret="amz-secret",
                    is_active=True,
                )
            )
            session.commit()
            trendyol_pk = int(trendyol.pk or 0)

        fetch_calls: list[dict] = []
        save_calls: list[dict] = []

        async def fake_fetch_orders_all(**kwargs):
            fetch_calls.append(dict(kwargs))
            return Result.ok(
                "Orders fake cekildi.",
                close_dialog=False,
                data={"order_data_list": [{"orderNumber": "ORD-SYNC-1"}], "order_item_list": []},
            )

        def fake_save_orders(result, db_name=None):
            rows = list((result.data or {}).get("order_data_list") or [])
            save_calls.append({"rows": rows, "db_name": db_name})
            return Result.ok(
                "Orders fake kaydedildi.",
                close_dialog=False,
                data={"rows": rows, "counts": {"orders_upserted": len(rows)}, "db_name": db_name},
            )

        prepared = run_agent_turn("Trendyol siparis verilerini cek", db_name=self.db_path)
        self._assert_prepare_answer_is_safe(prepared, expected_terms=("Trendyol siparis verileri",))
        self.assertNotIn("orders.get_orders_from_companies icin agentic prepare", self._answer_from(prepared))
        self.assertEqual(((prepared.data or {}).get("tool_call") or {}).get("tool"), "operations.manual_action.route")
        action = self._action_from(prepared)
        payload = ((action.get("arguments") or {}).get("payload") or {})
        self.assertEqual(payload.get("db_name"), self.db_path)
        self.assertEqual(payload.get("api_account_ids"), [trendyol_pk])
        self.assertNotIn("accounts", payload)
        self.assertTrue(payload.get("save_to_db"))
        self.assertFalse(payload.get("fetch_returns"))

        with tempfile.TemporaryDirectory() as temp_dir, patch(
            "Orders.processors.trendyol.pipeline.fetch_orders_all",
            side_effect=fake_fetch_orders_all,
        ), patch(
            "Orders.processors.trendyol.pipeline.save_orders_to_db",
            side_effect=fake_save_orders,
        ):
            audit_path = Path(temp_dir) / "orders-agent-audit.jsonl"
            executed = operation_runtime.execute_approved_operation(
                action,
                approved_by="behavioral-test",
                idempotency_key="fetch-trendyol-orders",
                audit_path=audit_path,
            )
            audit_text = audit_path.read_text(encoding="utf-8")

        verify = self._assert_execute_verified(executed)
        self.assertEqual(len(fetch_calls), 1)
        self.assertEqual(len(fetch_calls[0].get("comp_api_account_list") or []), 1)
        account_row = (fetch_calls[0].get("comp_api_account_list") or [])[0]
        self.assertEqual(account_row[0], trendyol_pk)
        self.assertGreaterEqual(len(account_row), 4)
        self.assertEqual(save_calls, [{"rows": [{"orderNumber": "ORD-SYNC-1"}], "db_name": self.db_path}])
        self.assertEqual(verify.get("save_counts"), {"orders_upserted": 1})
        last_used_update = verify.get("last_used_update") or {}
        self.assertTrue(last_used_update.get("verified"), verify)
        self.assertEqual(
            last_used_update.get("requested_account_ids"),
            [trendyol_pk],
        )
        with Session(get_engine(self.db_path)) as session:
            refreshed_account = session.get(ApiAccount, trendyol_pk)
            self.assertIsNotNone(refreshed_account)
            self.assertIsNotNone(refreshed_account.last_used_at)
        self.assertNotIn("api-key-sync", audit_text)
        self.assertNotIn("api-secret-sync", audit_text)

    def test_last_used_prompt_prepares_active_accounts_and_execute_updates_db(self) -> None:
        with Session(get_engine(self.db_path)) as session:
            account = ApiAccount(
                account_id="seller-last-used-agent",
                comp_name="Last Used Agent Seller",
                platform="trendyol",
                api_key="api-key-last-used",
                api_secret="api-secret-last-used",
                is_active=True,
            )
            other = ApiAccount(
                account_id="seller-last-used-other",
                comp_name="Last Used Other Seller",
                platform="amazon",
                api_key="api-key-other",
                api_secret="api-secret-other",
                is_active=True,
            )
            session.add(account)
            session.add(other)
            session.commit()
            account_pk = int(account.pk or 0)
            other_pk = int(other.pk or 0)

        prepared = run_agent_turn("Trendyol hesaplarinin son kullanim zamanini guncelle", db_name=self.db_path)

        self._assert_prepare_answer_is_safe(prepared, expected_terms=("son kullanim zamani",))
        self.assertNotIn("orders.update_last_used_at_for_accounts", self._answer_from(prepared))
        self.assertEqual(((prepared.data or {}).get("tool_call") or {}).get("tool"), "operations.manual_action.route")
        action = self._action_from(prepared)
        payload = ((action.get("arguments") or {}).get("payload") or {})
        self.assertEqual(payload.get("api_account_ids"), [account_pk])
        self.assertNotIn("accounts", payload)
        self.assertEqual(payload.get("db_name"), self.db_path)

        with tempfile.TemporaryDirectory() as temp_dir:
            audit_path = Path(temp_dir) / "orders-agent-audit.jsonl"
            executed = operation_runtime.execute_approved_operation(
                action,
                approved_by="behavioral-test",
                idempotency_key=f"last-used-agent-{account_pk}",
                audit_path=audit_path,
            )

            verify = self._assert_execute_verified(executed)
            self.assertEqual(verify.get("updated_pks"), [account_pk])
            audit_text = audit_path.read_text(encoding="utf-8")
            self.assertIn("orders.update_last_used_at_for_accounts", audit_text)
            self.assertNotIn("api-key-last-used", audit_text)
            self.assertNotIn("api-secret-last-used", audit_text)

        with Session(get_engine(self.db_path)) as session:
            updated = session.get(ApiAccount, account_pk)
            untouched = session.get(ApiAccount, other_pk)
            self.assertIsNotNone(updated)
            self.assertIsNotNone(getattr(updated, "last_used_at", None))
            self.assertIsNotNone(untouched)
            self.assertIsNone(getattr(untouched, "last_used_at", None))

    def test_orders_and_returns_fetch_prompt_executes_order_and_return_db_writes(self) -> None:
        now_ms = int(datetime.now(TZ_GMT3).timestamp() * 1000)
        with Session(get_engine(self.db_path)) as session:
            trendyol = ApiAccount(
                account_id="seller-trendyol-sync-returns",
                comp_name="Trendyol Sync Returns Seller",
                platform="trendyol",
                api_key="api-key-sync-returns",
                api_secret="api-secret-sync-returns",
                is_active=True,
            )
            session.add(trendyol)
            session.commit()
            trendyol_pk = int(trendyol.pk or 0)
        self._insert_waiting_return("CLAIM-SYNC-LOCAL-STALE", "ORD-SYNC-LOCAL-STALE", now_ms - 20_000, first_name="Local", last_name="Stale")

        fetch_calls: list[dict] = []
        save_calls: list[dict] = []
        return_fetch_calls: list[dict] = []

        async def fake_fetch_orders_all(**kwargs):
            fetch_calls.append(dict(kwargs))
            return Result.ok(
                "Orders fake cekildi.",
                close_dialog=False,
                data={"order_data_list": [{"orderNumber": "ORD-SYNC-REMOTE"}], "order_item_list": []},
            )

        def fake_save_orders(result, db_name=None):
            rows = list((result.data or {}).get("order_data_list") or [])
            save_calls.append({"rows": rows, "db_name": db_name})
            return Result.ok(
                "Orders fake kaydedildi.",
                close_dialog=False,
                data={"counts": {"orders_upserted": len(rows)}, "db_name": db_name},
            )

        def fake_run_async_safely(value):
            return value

        def fake_fetch_returns(accounts, *, start_ms, end_ms, size=50, claim_ids=None):
            return_fetch_calls.append(
                {
                    "accounts": list(accounts or []),
                    "start_ms": start_ms,
                    "end_ms": end_ms,
                    "size": size,
                    "claim_ids": claim_ids,
                }
            )
            return Result.ok(
                "Returns fake cekildi.",
                close_dialog=False,
                data={
                    "return_claim_list": [
                        {
                            "claimId": "CLAIM-SYNC-REMOTE",
                            "orderNumber": "ORD-SYNC-REMOTE",
                            "claimDate": now_ms + 1_000,
                            "orderDate": now_ms + 1_000,
                            "lastModifiedDate": now_ms + 2_000,
                            "claimItemStatus": "WaitingInAction",
                            "customerFirstName": "Remote",
                            "customerLastName": "Return",
                            "totalPrice": 44.0,
                            "itemCount": 1,
                            "waitingClaimLineItemIds": ["line-remote"],
                        }
                    ]
                },
            )

        prepared = run_agent_turn("Trendyol siparis ve iade verilerini cek", db_name=self.db_path)
        self._assert_prepare_answer_is_safe(prepared, expected_terms=("Trendyol siparis ve iade verileri",))
        self.assertNotIn("orders.get_orders_from_companies icin agentic prepare", self._answer_from(prepared))
        action = self._action_from(prepared)
        payload = ((action.get("arguments") or {}).get("payload") or {})
        self.assertEqual(payload.get("api_account_ids"), [trendyol_pk])
        self.assertNotIn("accounts", payload)
        self.assertTrue(payload.get("fetch_returns"), payload)
        self.assertTrue(payload.get("save_to_db"), payload)
        self.assertTrue(payload.get("save_returns_to_db"), payload)
        self.assertTrue(
            payload.get("verify_return_reconciliation"),
            payload,
        )
        self.assertEqual(payload.get("db_name"), self.db_path)

        with tempfile.TemporaryDirectory() as temp_dir, patch(
            "Orders.processors.trendyol.pipeline.fetch_orders_all",
            side_effect=fake_fetch_orders_all,
        ), patch(
            "Orders.processors.trendyol.pipeline.save_orders_to_db",
            side_effect=fake_save_orders,
        ), patch(
            "Orders.views.actions._run_async_safely",
            side_effect=fake_run_async_safely,
        ), patch(
            "Orders.views.actions._fetch_returns_from_active_companies_async",
            side_effect=fake_fetch_returns,
        ):
            audit_path = Path(temp_dir) / "orders-agent-audit.jsonl"
            executed = operation_runtime.execute_approved_operation(
                action,
                approved_by="behavioral-test",
                idempotency_key="fetch-trendyol-orders-and-returns",
                audit_path=audit_path,
            )

            self.assertTrue(executed.success, executed.message)
            self.assertTrue(audit_path.exists())
            audit_text = audit_path.read_text(encoding="utf-8")
            self.assertIn("orders.get_orders_from_companies", audit_text)
            self.assertNotIn("api-key-sync-returns", audit_text)
            self.assertNotIn("api-secret-sync-returns", audit_text)

        verify = self._assert_execute_verified(executed)
        self.assertEqual(len(fetch_calls), 1)
        self.assertEqual(save_calls, [{"rows": [{"orderNumber": "ORD-SYNC-REMOTE"}], "db_name": self.db_path}])
        self.assertEqual(len(return_fetch_calls), 1)
        self.assertEqual(len(return_fetch_calls[0].get("accounts") or []), 1)
        self.assertIsInstance(
            (return_fetch_calls[0].get("accounts") or [None])[0],
            ApiAccount,
        )
        self.assertEqual(verify.get("save_counts"), {"orders_upserted": 1})
        self.assertTrue((verify.get("return_save_counts") or {}).get("claims_inserted") or (verify.get("return_save_counts") or {}).get("claims_upserted"), verify)
        reconciliation = verify.get("return_reconciliation") or {}
        self.assertEqual(reconciliation.get("stale_count"), 1)
        self.assertEqual((reconciliation.get("stale_returns") or [])[0].get("claim_id"), "CLAIM-SYNC-LOCAL-STALE")

        with Session(get_engine(self.db_path)) as session:
            remote = session.exec(select(ReturnClaim).where(ReturnClaim.claimId == "CLAIM-SYNC-REMOTE")).one()
            stale = session.exec(select(ReturnClaim).where(ReturnClaim.claimId == "CLAIM-SYNC-LOCAL-STALE")).one()
        self.assertEqual(remote.orderNumber, "ORD-SYNC-REMOTE")
        self.assertIsNone(remote.sync_deleted_at)
        self.assertIsNone(stale.sync_deleted_at)

    def test_customer_specific_cargo_change_executes_only_that_order(self) -> None:
        self._seed_ready_orders()

        def fake_change_order_cargo_provider(order, *, cargo_provider_code: str):
            with Session(get_engine(self.db_path)) as session:
                record = session.exec(select(OrderData).where(OrderData.orderNumber == order.orderNumber)).one()
                record.cargoProviderName = "Aras"
                session.add(record)
                session.commit()
            return Result.ok(
                "Kargo firmasi degistirildi.",
                close_dialog=False,
                data={"order_number": order.orderNumber, "cargo_provider_code": cargo_provider_code},
            )

        prepared = run_agent_turn("Mahmut Şevket'in siparişinin kargo firmasını Aras yap", db_name=self.db_path)
        self._assert_prepare_answer_is_safe(prepared, expected_terms=("kargo",))
        action = self._action_from(prepared)

        with tempfile.TemporaryDirectory() as temp_dir, patch(
            "Orders.views.actions.change_order_cargo_provider",
            side_effect=fake_change_order_cargo_provider,
        ):
            executed = operation_runtime.execute_approved_operation(
                action,
                approved_by="behavioral-test",
                idempotency_key="cargo-mahmut-aras",
                audit_path=Path(temp_dir) / "orders-agent-audit.jsonl",
            )

        self._assert_execute_verified(executed)
        with Session(get_engine(self.db_path)) as session:
            mahmut = session.exec(select(OrderData).where(OrderData.orderNumber == "ORD-MAHMUT-READY")).one()
            ayse = session.exec(select(OrderData).where(OrderData.orderNumber == "ORD-AYSE-READY")).one()
            self.assertEqual(mahmut.cargoProviderName, "Aras")
            self.assertEqual(ayse.cargoProviderName, "Trendyol Express")

    def test_followup_customer_refined_cargo_change_uses_only_matching_previous_order(self) -> None:
        self._seed_ready_orders()

        listed = run_agent_turn("Kargoya hazir siparisleri getir", db_name=self.db_path)
        self.assertTrue(listed.success, listed.message)
        listed_answer = self._answer_from(listed)
        self.assertIn("ORD-MAHMUT-READY", listed_answer)
        self.assertIn("ORD-AYSE-READY", listed_answer)

        history = [
            {"role": "user", "content": "Kargoya hazir siparisleri getir"},
            {"role": "assistant", "content": listed_answer},
        ]
        prepared = run_agent_turn("Mahmut olanin kargo firmasini Aras yap", db_name=self.db_path, history=history)
        self._assert_prepare_answer_is_safe(prepared, expected_terms=("ORD-MAHMUT-READY", "kargo"))
        answer = self._answer_from(prepared)
        self.assertNotIn("ORD-AYSE-READY", answer)
        self.assertEqual(((prepared.data or {}).get("tool_call") or {}).get("tool"), "orders.change_cargo_provider.prepare")
        action = self._action_from(prepared)
        rows = [dict(row) for row in action.get("rows") or [] if isinstance(row, dict)]
        self.assertEqual(len(rows), 1)
        payload = rows[0].get("payload") or {}
        self.assertEqual(payload.get("order"), "ORD-MAHMUT-READY")

        def fake_change_order_cargo_provider(order, *, cargo_provider_code: str):
            with Session(get_engine(self.db_path)) as session:
                record = session.exec(select(OrderData).where(OrderData.orderNumber == order.orderNumber)).one()
                record.cargoProviderName = "Aras"
                session.add(record)
                session.commit()
            return Result.ok(
                "Kargo firmasi degistirildi.",
                close_dialog=False,
                data={"order_number": order.orderNumber, "cargo_provider_code": cargo_provider_code},
            )

        with tempfile.TemporaryDirectory() as temp_dir, patch(
            "Orders.views.actions.change_order_cargo_provider",
            side_effect=fake_change_order_cargo_provider,
        ):
            executed = operation_runtime.execute_approved_operation(
                action,
                approved_by="behavioral-test",
                idempotency_key="followup-customer-cargo-mahmut-aras",
                audit_path=Path(temp_dir) / "orders-agent-audit.jsonl",
            )

        self._assert_execute_verified(executed)
        with Session(get_engine(self.db_path)) as session:
            providers = {
                row.orderNumber: row.cargoProviderName
                for row in session.exec(select(OrderData).where(OrderData.orderNumber.in_(["ORD-MAHMUT-READY", "ORD-AYSE-READY"]))).all()
            }
        self.assertEqual(providers.get("ORD-MAHMUT-READY"), "Aras")
        self.assertEqual(providers.get("ORD-AYSE-READY"), "Trendyol Express")

    def test_followup_this_cargo_change_uses_latest_single_order_context(self) -> None:
        now_ms = int(datetime.now(TZ_GMT3).timestamp() * 1000)
        self._insert_order(
            "ORD-LATEST-CARGO",
            now_ms,
            package_id=904,
            first_name="Latest",
            last_name="Cargo",
            status="ReadyToShip",
            cargo_provider_name="Trendyol Express",
        )
        self._insert_order(
            "ORD-LATEST-CARGO-OLD",
            now_ms - 10_000,
            package_id=905,
            first_name="Old",
            last_name="Cargo",
            status="ReadyToShip",
            cargo_provider_name="Trendyol Express",
        )

        listed = run_agent_turn("son gelen 1 siparisi soyle", db_name=self.db_path)
        self.assertTrue(listed.success, listed.message)
        listed_answer = self._answer_from(listed)
        self.assertIn("ORD-LATEST-CARGO", listed_answer)
        self.assertNotIn("ORD-LATEST-CARGO-OLD", listed_answer)
        history = [
            {"role": "user", "content": "son gelen 1 siparisi soyle"},
            {"role": "assistant", "content": listed_answer},
        ]

        prepared = run_agent_turn(
            "bunun kargo firmasini Aras yap",
            planner_client=lambda _prompt: '{"mode":"answer","response":"Tabii, nasil yardimci olabilirim?"}',
            chat_client=lambda _prompt: "Tabii, nasil yardimci olabilirim?",
            db_name=self.db_path,
            history=history,
        )

        self._assert_prepare_answer_is_safe(prepared, expected_terms=("ORD-LATEST-CARGO", "kargo"))
        self.assertEqual(((prepared.data or {}).get("tool_call") or {}).get("tool"), "orders.change_cargo_provider.prepare")
        self.assertNotIn("ORD-LATEST-CARGO-OLD", self._answer_from(prepared))
        action = self._action_from(prepared)
        self.assertEqual(action.get("action_type"), "orders.change_order_cargo_provider")
        rows = [dict(row) for row in action.get("rows") or [] if isinstance(row, dict)]
        self.assertEqual(len(rows), 1)
        payload = rows[0].get("payload") or {}
        self.assertEqual(payload.get("order"), "ORD-LATEST-CARGO")

    def test_llm_bulk_cargo_misroute_is_narrowed_by_customer_context(self) -> None:
        self._seed_ready_orders()
        listed = run_agent_turn("Kargoya hazir siparisleri getir", db_name=self.db_path)
        self.assertTrue(listed.success, listed.message)
        history = [
            {"role": "user", "content": "Kargoya hazir siparisleri getir"},
            {"role": "assistant", "content": self._answer_from(listed)},
        ]

        prepared = run_agent_turn(
            "Mahmut olanin kargo firmasini Aras yap",
            planner_client=lambda _prompt: (
                '{"mode":"tool","tool":"orders.change_cargo_provider_bulk.prepare",'
                '"arguments":{"order_numbers":["ORD-MAHMUT-READY","ORD-AYSE-READY"],"cargo_provider":"Aras"}}'
            ),
            chat_client=lambda _prompt: "Onay gerekiyor.",
            db_name=self.db_path,
            history=history,
        )

        self.assertTrue(prepared.success, prepared.message)
        self.assertEqual(((prepared.data or {}).get("tool_call") or {}).get("tool"), "orders.change_cargo_provider.prepare")
        action = self._action_from(prepared)
        self.assertEqual(action.get("action_type"), "orders.change_order_cargo_provider")
        payload = ((action.get("rows") or [{}])[0].get("payload") or {})
        self.assertEqual(payload.get("order"), "ORD-MAHMUT-READY")

    def test_bulk_cargo_change_executes_all_ready_orders(self) -> None:
        self._seed_ready_orders()
        calls: list[dict] = []

        def fake_change_orders_cargo_provider_bulk(orders, *, cargo_provider_code: str):
            order_numbers = [order.orderNumber for order in orders]
            calls.append({"orders": order_numbers, "cargo_provider_code": cargo_provider_code})
            with Session(get_engine(self.db_path)) as session:
                records = session.exec(select(OrderData).where(OrderData.orderNumber.in_(order_numbers))).all()
                for record in records:
                    record.cargoProviderName = "Aras"
                    session.add(record)
                session.commit()
            return Result.ok("Toplu kargo firmasi degistirildi.", close_dialog=False, data={"orders": order_numbers})

        prepared = run_agent_turn("Kargoya hazır siparişlerin hepsinin kargo firmasını Aras yap", db_name=self.db_path)
        self._assert_prepare_answer_is_safe(prepared, expected_terms=("2 siparis", "kargo"))
        action = self._action_from(prepared)
        self.assertEqual(action.get("action_type"), "orders.change_orders_cargo_provider_bulk")

        with tempfile.TemporaryDirectory() as temp_dir, patch(
            "Orders.views.actions.change_orders_cargo_provider_bulk",
            side_effect=fake_change_orders_cargo_provider_bulk,
        ):
            executed = operation_runtime.execute_approved_operation(
                action,
                approved_by="behavioral-test",
                idempotency_key="bulk-cargo-ready-aras",
                audit_path=Path(temp_dir) / "orders-agent-audit.jsonl",
            )

        self._assert_execute_verified(executed)
        self.assertEqual(len(calls), 1)
        self.assertEqual(set(calls[0].get("orders") or []), {"ORD-AYSE-READY", "ORD-MAHMUT-READY"})
        self.assertEqual(calls[0].get("cargo_provider_code"), "ARAS")
        with Session(get_engine(self.db_path)) as session:
            providers = {
                row.orderNumber: row.cargoProviderName
                for row in session.exec(select(OrderData).where(OrderData.orderNumber.in_(["ORD-MAHMUT-READY", "ORD-AYSE-READY"]))).all()
            }
        self.assertEqual(providers, {"ORD-MAHMUT-READY": "Aras", "ORD-AYSE-READY": "Aras"})

    def test_llm_wrong_profit_order_query_is_narrowed_by_customer_context(self) -> None:
        self._seed_ready_orders()
        listed = run_agent_turn("Kargoya hazir siparisleri getir", db_name=self.db_path)
        self.assertTrue(listed.success, listed.message)
        history = [
            {"role": "user", "content": "Kargoya hazir siparisleri getir"},
            {"role": "assistant", "content": self._answer_from(listed)},
        ]

        with tempfile.TemporaryDirectory() as temp_dir:
            output_path = str(Path(temp_dir) / "mahmut-profit.xlsx")
            planner_payload = {
                "mode": "tool",
                "tool": "orders.export_profit_excel.prepare",
                "arguments": {"query": "ORD-AYSE-READY", "output_path": output_path.replace("\\", "/")},
            }
            prepared = run_agent_turn(
                f'Mahmut olanin kar Excel raporunu "{output_path}" dosyasina aktar',
                planner_client=lambda _prompt: json.dumps(planner_payload),
                chat_client=lambda _prompt: "Onay gerekiyor.",
                db_name=self.db_path,
                history=history,
            )

        self.assertTrue(prepared.success, prepared.message)
        self.assertEqual(((prepared.data or {}).get("tool_call") or {}).get("tool"), "orders.export_profit_excel.prepare")
        action = self._action_from(prepared)
        payload = ((action.get("arguments") or {}).get("payload") or {})
        self.assertEqual(payload.get("order_number"), "ORD-MAHMUT-READY")

    def test_followup_changes_previous_ready_order_list_cargo_provider_only(self) -> None:
        self._seed_ready_orders()
        self._insert_order(
            "ORD-DELIVERED-OTHER",
            int(datetime.now(TZ_GMT3).timestamp() * 1000) + 5_000,
            package_id=903,
            first_name="Delivered",
            last_name="Other",
            status="Delivered",
            cargo_provider_name="Trendyol Express",
        )

        listed = run_agent_turn("Kargoya hazir siparisleri getir", db_name=self.db_path)
        self.assertTrue(listed.success, listed.message)
        listed_answer = self._answer_from(listed)
        self.assertIn("ORD-MAHMUT-READY", listed_answer)
        self.assertIn("ORD-AYSE-READY", listed_answer)
        self.assertNotIn("ORD-DELIVERED-OTHER", listed_answer)

        history = [
            {"role": "user", "content": "Kargoya hazir siparisleri getir"},
            {"role": "assistant", "content": listed_answer},
        ]
        prepared = run_agent_turn("bunlarin kargo firmasini Aras yap", db_name=self.db_path, history=history)
        self._assert_prepare_answer_is_safe(prepared, expected_terms=("2 siparis", "kargo"))
        self.assertEqual(((prepared.data or {}).get("tool_call") or {}).get("tool"), "orders.change_cargo_provider_bulk.prepare")
        args = ((prepared.data or {}).get("tool_call") or {}).get("arguments") or {}
        self.assertEqual(set(args.get("order_numbers") or []), {"ORD-MAHMUT-READY", "ORD-AYSE-READY"})
        self.assertNotIn("ORD-DELIVERED-OTHER", args.get("order_numbers") or [])
        action = self._action_from(prepared)
        payload = ((action.get("arguments") or {}).get("payload") or {})
        self.assertEqual(set(payload.get("orders") or []), {"ORD-MAHMUT-READY", "ORD-AYSE-READY"})

        def fake_change_orders_cargo_provider_bulk(orders, *, cargo_provider_code: str):
            order_numbers = [order.orderNumber for order in orders]
            with Session(get_engine(self.db_path)) as session:
                records = session.exec(select(OrderData).where(OrderData.orderNumber.in_(order_numbers))).all()
                for record in records:
                    record.cargoProviderName = "Aras"
                    session.add(record)
                session.commit()
            return Result.ok(
                "Toplu kargo firmasi degistirildi.",
                close_dialog=False,
                data={"orders": order_numbers, "cargo_provider_code": cargo_provider_code},
            )

        with tempfile.TemporaryDirectory() as temp_dir, patch(
            "Orders.views.actions.change_orders_cargo_provider_bulk",
            side_effect=fake_change_orders_cargo_provider_bulk,
        ):
            executed = operation_runtime.execute_approved_operation(
                action,
                approved_by="behavioral-test",
                idempotency_key="followup-bulk-cargo-listed-ready-aras",
                audit_path=Path(temp_dir) / "orders-agent-audit.jsonl",
            )

        verify = self._assert_execute_verified(executed)
        self.assertEqual({row.get("order_number") for row in verify.get("rows") or []}, {"ORD-MAHMUT-READY", "ORD-AYSE-READY"})
        with Session(get_engine(self.db_path)) as session:
            providers = {
                row.orderNumber: row.cargoProviderName
                for row in session.exec(
                    select(OrderData).where(OrderData.orderNumber.in_(["ORD-MAHMUT-READY", "ORD-AYSE-READY", "ORD-DELIVERED-OTHER"]))
                ).all()
            }
        self.assertEqual(providers.get("ORD-MAHMUT-READY"), "Aras")
        self.assertEqual(providers.get("ORD-AYSE-READY"), "Aras")
        self.assertEqual(providers.get("ORD-DELIVERED-OTHER"), "Trendyol Express")

    def test_customer_specific_return_approval_executes_and_clears_waiting_state(self) -> None:
        now_ms = self._seed_ready_orders()
        self._insert_waiting_return(
            "CLAIM-MAHMUT-RETURN",
            "ORD-MAHMUT-READY",
            now_ms + 2_000,
            first_name="Mahmut",
            last_name="Şevket",
        )

        def fake_approve_return_claim(*, order, **_kwargs):
            with Session(get_engine(self.db_path)) as session:
                claim = session.exec(select(ReturnClaim).where(ReturnClaim.claimId == order.claimId)).one()
                claim.claimItemStatus = "Accepted"
                claim.waitingClaimLineItemIds = []
                session.add(claim)
                session.commit()
            return Result.ok("Iade onaylandi.", close_dialog=False, data={"claim_id": order.claimId})

        prepared = run_agent_turn("Mahmut Şevket'in iadesini onayla", db_name=self.db_path)
        self._assert_prepare_answer_is_safe(prepared, expected_terms=("iade",))
        action = self._action_from(prepared)

        with tempfile.TemporaryDirectory() as temp_dir, patch(
            "Orders.views.actions.approve_return_claim",
            side_effect=fake_approve_return_claim,
        ):
            executed = operation_runtime.execute_approved_operation(
                action,
                approved_by="behavioral-test",
                idempotency_key="return-mahmut-approve",
                audit_path=Path(temp_dir) / "orders-agent-audit.jsonl",
            )

        self._assert_execute_verified(executed)
        with Session(get_engine(self.db_path)) as session:
            claim = session.exec(select(ReturnClaim).where(ReturnClaim.claimId == "CLAIM-MAHMUT-RETURN")).one()
            self.assertEqual(claim.claimItemStatus, "Accepted")
            self.assertEqual(claim.waitingClaimLineItemIds, [])

    def test_explicit_stale_return_cleanup_executes_and_hides_claim(self) -> None:
        now_ms = self._seed_ready_orders()
        self._insert_waiting_return(
            "CLAIM-LOCAL-STALE",
            "ORD-LOCAL-STALE",
            now_ms + 2_000,
            first_name="Local",
            last_name="Stale",
        )

        prepared = run_agent_turn("CLAIM-LOCAL-STALE stale iade kaydini temizle", db_name=self.db_path)
        self._assert_prepare_answer_is_safe(prepared, expected_terms=("stale", "iade"))
        self.assertEqual(((prepared.data or {}).get("tool_call") or {}).get("tool"), "returns.mark_stale.prepare")
        action = self._action_from(prepared)
        payload = ((action.get("arguments") or {}).get("payload") or {})
        self.assertEqual(payload.get("claim_ids"), ["CLAIM-LOCAL-STALE"])

        with tempfile.TemporaryDirectory() as temp_dir:
            executed = operation_runtime.execute_approved_operation(
                action,
                approved_by="behavioral-test",
                idempotency_key="explicit-stale-return-cleanup",
                audit_path=Path(temp_dir) / "orders-agent-audit.jsonl",
            )

        verify = self._assert_execute_verified(executed)
        self.assertEqual(verify.get("claim_ids"), ["CLAIM-LOCAL-STALE"])
        with Session(get_engine(self.db_path)) as session:
            claim = session.exec(select(ReturnClaim).where(ReturnClaim.claimId == "CLAIM-LOCAL-STALE")).one()
            self.assertTrue(str(claim.sync_deleted_at or "").strip())
        lookup = run_agent_turn("CLAIM-LOCAL-STALE iadesini göster", db_name=self.db_path)
        active_returns = ((lookup.data or {}).get("tool_result") or {}).get("returns") or []
        self.assertNotIn("CLAIM-LOCAL-STALE", [row.get("claim_id") for row in active_returns])

    def test_followup_stale_return_cleanup_uses_previous_stale_claim_list(self) -> None:
        now_ms = self._seed_ready_orders()
        self._insert_waiting_return(
            "CLAIM-FOLLOWUP-STALE",
            "ORD-FOLLOWUP-STALE",
            now_ms + 2_000,
            first_name="Followup",
            last_name="Stale",
        )
        self._insert_waiting_return(
            "CLAIM-ACTIVE-KEEP",
            "ORD-ACTIVE-KEEP",
            now_ms + 3_000,
            first_name="Active",
            last_name="Keep",
        )

        history = [
            {"role": "user", "content": "UI'da iade var ama Trendyol webde gorunmuyor, stale kayitlari kontrol et"},
            {"role": "assistant", "content": "Return reconciliation tamamlandi: 1 stale lokal iade. 1. stale CLAIM-FOLLOWUP-STALE | siparis ORD-FOLLOWUP-STALE"},
        ]
        prepared = run_agent_turn("bu stale iadeleri temizle", db_name=self.db_path, history=history)
        self._assert_prepare_answer_is_safe(prepared, expected_terms=("1 stale", "iade"))
        self.assertEqual(((prepared.data or {}).get("tool_call") or {}).get("tool"), "returns.mark_stale.prepare")
        action = self._action_from(prepared)
        payload = ((action.get("arguments") or {}).get("payload") or {})
        self.assertEqual(payload.get("claim_ids"), ["CLAIM-FOLLOWUP-STALE"])

        with tempfile.TemporaryDirectory() as temp_dir:
            executed = operation_runtime.execute_approved_operation(
                action,
                approved_by="behavioral-test",
                idempotency_key="followup-stale-return-cleanup",
                audit_path=Path(temp_dir) / "orders-agent-audit.jsonl",
            )

        self._assert_execute_verified(executed)
        with Session(get_engine(self.db_path)) as session:
            stale = session.exec(select(ReturnClaim).where(ReturnClaim.claimId == "CLAIM-FOLLOWUP-STALE")).one()
            active = session.exec(select(ReturnClaim).where(ReturnClaim.claimId == "CLAIM-ACTIVE-KEEP")).one()
            self.assertTrue(str(stale.sync_deleted_at or "").strip())
            self.assertFalse(str(active.sync_deleted_at or "").strip())

    def test_stale_return_review_prompt_lists_stale_claims_without_writing(self) -> None:
        now_ms = self._seed_ready_orders()
        with Session(get_engine(self.db_path)) as session:
            session.add(ApiAccount(pk=1, account_id="seller-stale-review", comp_name="Stale Review Seller", platform="trendyol"))
            session.add(
                ReturnClaim(
                    claimId="CLAIM-REVIEW-STALE",
                    orderNumber="ORD-REVIEW-STALE",
                    api_account_id=1,
                    claimDate=now_ms + 2_000,
                    orderDate=now_ms + 2_000,
                    lastModifiedDate=now_ms + 3_000,
                    claimItemStatus="WaitingInAction",
                    customerFirstName="Review",
                    customerLastName="Stale",
                    totalPrice=55.0,
                    itemCount=1,
                    waitingClaimLineItemIds=["line-1"],
                    sync_deleted_at="2026-01-01T00:00:00+03:00",
                )
            )
            session.commit()

        result = run_agent_turn("stale iadeleri goster", db_name=self.db_path)

        self.assertTrue(result.success, result.message)
        self.assertEqual(((result.data or {}).get("tool_call") or {}).get("tool"), "operations.manual_action.route")
        tool_result = (result.data or {}).get("tool_result") or {}
        self.assertEqual(tool_result.get("route_status"), "read_executed")
        self.assertEqual(((tool_result.get("match") or {}).get("capability_id")), "orders.get_stale_return_claims")
        claims = ((tool_result.get("result") or {}).get("claims") or [])
        claim_ids = [
            row.get("claimId") if isinstance(row, dict) else getattr(row, "claimId", "")
            for row in claims
        ]
        self.assertEqual(claim_ids, ["CLAIM-REVIEW-STALE"])
        answer = self._answer_from(result)
        self.assertIn("CLAIM-REVIEW-STALE", answer)
        self.assertNotIn("Onay gerekiyor", answer)
        self.assertNotIn("temizlendi", answer.casefold())

    def test_stale_return_restore_executes_and_returns_claim_to_active_lookup(self) -> None:
        now_ms = self._seed_ready_orders()
        with Session(get_engine(self.db_path)) as session:
            session.add(ApiAccount(pk=1, account_id="seller-stale", comp_name="Stale Seller", platform="trendyol"))
            session.add(
                ReturnClaim(
                    claimId="CLAIM-MAHMUT-STALE",
                    orderNumber="ORD-MAHMUT-READY",
                    api_account_id=1,
                    claimDate=now_ms + 2_000,
                    orderDate=now_ms + 2_000,
                    lastModifiedDate=now_ms + 3_000,
                    claimItemStatus="WaitingInAction",
                    customerFirstName="Mahmut",
                    customerLastName="Şevket",
                    totalPrice=55.0,
                    itemCount=1,
                    waitingClaimLineItemIds=["line-1"],
                    sync_deleted_at="2026-01-01T00:00:00+03:00",
                )
            )
            session.commit()

        prepared = run_agent_turn("CLAIM-MAHMUT-STALE stale iadesini geri al", db_name=self.db_path)
        action = self._action_from(prepared)
        self.assertEqual(action.get("action_type"), "orders.restore_return_claims_from_stale")

        with tempfile.TemporaryDirectory() as temp_dir:
            audit_path = Path(temp_dir) / "orders-agent-audit.jsonl"
            executed = operation_runtime.execute_approved_operation(
                action,
                approved_by="behavioral-test",
                idempotency_key="restore-mahmut-stale",
                audit_path=audit_path,
            )
            self.assertTrue(executed.success, executed.message)
            self.assertIn("orders.restore_return_claims_from_stale", audit_path.read_text(encoding="utf-8"))

        with Session(get_engine(self.db_path)) as session:
            claim = session.exec(select(ReturnClaim).where(ReturnClaim.claimId == "CLAIM-MAHMUT-STALE")).one()
            self.assertIsNone(claim.sync_deleted_at)
        lookup = run_agent_turn("CLAIM-MAHMUT-STALE iadesini göster", db_name=self.db_path)
        self.assertTrue(lookup.success, lookup.message)
        self.assertIn("CLAIM-MAHMUT-STALE", (lookup.data or {}).get("answer") or lookup.message)


    def test_return_lookup_rejects_unsupported_action_completion_claims(self) -> None:
        now_ms = self._seed_ready_orders()
        self._insert_waiting_return(
            "CLAIM-LOOKUP-ACTIVE",
            "ORD-MAHMUT-READY",
            now_ms + 2_000,
            first_name="Mahmut",
            last_name="Sevket",
        )

        result = run_agent_turn(
            "CLAIM-LOOKUP-ACTIVE iadesini goster",
            planner_client=lambda _prompt: '{"mode":"tool","tool":"returns.lookup","arguments":{"query":"CLAIM-LOOKUP-ACTIVE","limit":10}}',
            chat_client=lambda _prompt: "CLAIM-LOOKUP-ACTIVE iadesi beklemede, geri alma yapilmadi.",
            db_name=self.db_path,
        )

        self.assertTrue(result.success, result.message)
        self.assertEqual(((result.data or {}).get("tool_call") or {}).get("tool"), "returns.lookup")
        answer = self._answer_from(result)
        self.assertIn("CLAIM-LOOKUP-ACTIVE", answer)
        self.assertIn("WaitingInAction", answer)
        self.assertNotIn("geri alma yapilmadi", answer.casefold())


class TestOrdersBehavioralLiveLLME2E(TestOrdersBehavioralE2E):
    _provider_record: dict
    _provider_source: str

    @classmethod
    def setUpClass(cls) -> None:
        super().setUpClass()
        if not _live_eval_enabled():
            raise unittest.SkipTest("Set ORDERSCOUT_LIVE_LLM_EVAL=1 to run real cloud LLM behavioral Orders evals.")
        cls._provider_record, cls._provider_source = _live_provider_record()

    def _client(self) -> _CountingClient:
        record = dict(self._provider_record)
        missing = [key for key in ("provider_code", "api_key", "default_model") if not str(record.get(key) or "").strip()]
        if missing:
            self.fail("Live LLM provider record is missing required fields: " + ", ".join(missing))
        client = build_provider_chat_client(record, timeout=120)
        self.assertIsNotNone(client)
        return _CountingClient(client)

    def test_real_cloud_llm_filtered_lists_and_return_pagination(self) -> None:
        now_ms = int(datetime.now(TZ_GMT3).timestamp() * 1000)
        yesterday_ms = int((datetime.now(TZ_GMT3) - timedelta(days=1)).timestamp() * 1000)
        with Session(get_engine(self.db_path)) as session:
            session.add(ApiAccount(pk=10, account_id="live-ty-agent", comp_name="Live TY Agent", platform="trendyol"))
            session.add(ApiAccount(pk=20, account_id="live-amz-agent", comp_name="Live AMZ Agent", platform="amazon"))
            session.commit()
        self._insert_order(
            "ORD-LIVE-ARAS-1",
            now_ms,
            package_id=2_101,
            first_name="Live",
            last_name="ArasOne",
            status="ReadyToShip",
            cargo_provider_name="Aras",
        )
        self._insert_order(
            "ORD-LIVE-TEX-1",
            now_ms + 1_000,
            package_id=2_102,
            first_name="Live",
            last_name="TexOne",
            status="ReadyToShip",
            cargo_provider_name="Trendyol Express",
        )
        self._insert_order(
            "ORD-LIVE-ARAS-2",
            now_ms + 2_000,
            package_id=2_103,
            first_name="Live",
            last_name="ArasTwo",
            status="Created",
            cargo_provider_name="Aras",
        )
        self._insert_order(
            "ORD-LIVE-TODAY-CREATED",
            now_ms + 3_000,
            package_id=2_104,
            first_name="Live",
            last_name="CreatedToday",
            status="Created",
            cargo_provider_name="MNG",
        )
        self._insert_order(
            "ORD-LIVE-TY-PLATFORM",
            now_ms + 4_000,
            package_id=2_106,
            first_name="Live",
            last_name="TYPlatform",
            status="Created",
            cargo_provider_name="MNG",
            api_account_id=10,
        )
        self._insert_order(
            "ORD-LIVE-AMZ-PLATFORM",
            now_ms + 5_000,
            package_id=2_107,
            first_name="Live",
            last_name="AMZPlatform",
            status="Created",
            cargo_provider_name="MNG",
            api_account_id=20,
        )
        self._insert_order(
            "ORD-LIVE-YESTERDAY-CREATED",
            yesterday_ms,
            package_id=2_105,
            first_name="Live",
            last_name="CreatedYesterday",
            status="Created",
            cargo_provider_name="MNG",
        )
        for idx in range(1, 6):
            self._insert_waiting_return(
                f"CLAIM-LIVE-PAGE-{idx}",
                f"ORD-LIVE-RETURN-PAGE-{idx}",
                now_ms + (idx * 1_000),
                first_name=f"LiveReturn{idx}",
                last_name="Customer",
            )
        self._insert_waiting_return(
            "CLAIM-LIVE-TODAY-ACCEPTED",
            "ORD-LIVE-TODAY-ACCEPTED",
            now_ms + 6_000,
            first_name="LiveToday",
            last_name="Accepted",
        )
        self._insert_waiting_return(
            "CLAIM-LIVE-YESTERDAY-ACCEPTED",
            "ORD-LIVE-YESTERDAY-ACCEPTED",
            yesterday_ms,
            first_name="LiveYesterday",
            last_name="Accepted",
        )
        self._insert_waiting_return(
            "CLAIM-LIVE-TY-PLATFORM",
            "ORD-LIVE-TY-RETURN-PLATFORM",
            now_ms + 7_000,
            first_name="LiveTY",
            last_name="Return",
            api_account_id=10,
        )
        self._insert_waiting_return(
            "CLAIM-LIVE-AMZ-PLATFORM",
            "ORD-LIVE-AMZ-RETURN-PLATFORM",
            now_ms + 8_000,
            first_name="LiveAMZ",
            last_name="Return",
            api_account_id=20,
        )
        with Session(get_engine(self.db_path)) as session:
            for claim_id in ("CLAIM-LIVE-TODAY-ACCEPTED", "CLAIM-LIVE-YESTERDAY-ACCEPTED"):
                claim = session.exec(select(ReturnClaim).where(ReturnClaim.claimId == claim_id)).one()
                claim.claimItemStatus = "Accepted"
                claim.waitingClaimLineItemIds = []
                session.add(claim)
            for claim_id in ("CLAIM-LIVE-TY-PLATFORM", "CLAIM-LIVE-AMZ-PLATFORM"):
                claim = session.exec(select(ReturnClaim).where(ReturnClaim.claimId == claim_id)).one()
                claim.claimItemStatus = "PlatformOnly"
                claim.waitingClaimLineItemIds = []
                session.add(claim)
            session.commit()

        client = self._client()
        provider_label = _safe_provider_label(self._provider_record, self._provider_source)

        before_calls = len(client.prompts)
        orders_result = run_agent_turn(
            "Aras kargolu siparisleri getir",
            planner_client=client,
            chat_client=client,
            db_name=self.db_path,
        )
        self.assertGreater(len(client.prompts) - before_calls, 0, provider_label)
        self.assertTrue(orders_result.success, orders_result.message)
        self.assertEqual(((orders_result.data or {}).get("tool_call") or {}).get("tool"), "orders.query")
        orders = ((orders_result.data or {}).get("tool_result") or {}).get("orders") or []
        self.assertEqual([row.get("order_number") for row in orders], ["ORD-LIVE-ARAS-2", "ORD-LIVE-ARAS-1"])

        status_orders_result = run_agent_turn(
            "Bugunku durumu Created olan siparisleri getir",
            planner_client=client,
            chat_client=client,
            db_name=self.db_path,
        )
        self.assertTrue(status_orders_result.success, status_orders_result.message)
        self.assertEqual(((status_orders_result.data or {}).get("tool_call") or {}).get("tool"), "orders.query")
        status_orders = ((status_orders_result.data or {}).get("tool_result") or {}).get("orders") or []
        self.assertEqual(
            [row.get("order_number") for row in status_orders],
            ["ORD-LIVE-AMZ-PLATFORM", "ORD-LIVE-TY-PLATFORM", "ORD-LIVE-TODAY-CREATED", "ORD-LIVE-ARAS-2"],
        )

        platform_orders_result = run_agent_turn(
            "Trendyol pazaryeri siparislerini getir",
            planner_client=client,
            chat_client=client,
            db_name=self.db_path,
        )
        self.assertTrue(platform_orders_result.success, platform_orders_result.message)
        self.assertEqual(((platform_orders_result.data or {}).get("tool_call") or {}).get("tool"), "orders.query")
        platform_orders = ((platform_orders_result.data or {}).get("tool_result") or {}).get("orders") or []
        self.assertEqual([row.get("order_number") for row in platform_orders], ["ORD-LIVE-TY-PLATFORM"])

        before_calls = len(client.prompts)
        returns_result = run_agent_turn(
            "Aksiyon bekleyen iadelerin ikinci sayfasini getir, sayfa basi 2 kayit",
            planner_client=client,
            chat_client=client,
            db_name=self.db_path,
        )
        self.assertGreater(len(client.prompts) - before_calls, 0, provider_label)
        self.assertTrue(returns_result.success, returns_result.message)
        self.assertEqual(((returns_result.data or {}).get("tool_call") or {}).get("tool"), "returns.query")
        returns = ((returns_result.data or {}).get("tool_result") or {}).get("returns") or []
        self.assertEqual([row.get("claim_id") for row in returns], ["CLAIM-LIVE-PAGE-3", "CLAIM-LIVE-PAGE-2"])

        accepted_result = run_agent_turn(
            "Bugunku durumu Accepted olan iadeleri getir",
            planner_client=client,
            chat_client=client,
            db_name=self.db_path,
        )
        self.assertTrue(accepted_result.success, accepted_result.message)
        self.assertEqual(((accepted_result.data or {}).get("tool_call") or {}).get("tool"), "returns.query")
        accepted = ((accepted_result.data or {}).get("tool_result") or {}).get("returns") or []
        self.assertEqual([row.get("claim_id") for row in accepted], ["CLAIM-LIVE-TODAY-ACCEPTED"])

        platform_returns_result = run_agent_turn(
            "Amazon pazaryeri iadelerini getir",
            planner_client=client,
            chat_client=client,
            db_name=self.db_path,
        )
        self.assertTrue(platform_returns_result.success, platform_returns_result.message)
        self.assertEqual(((platform_returns_result.data or {}).get("tool_call") or {}).get("tool"), "returns.query")
        platform_returns = ((platform_returns_result.data or {}).get("tool_result") or {}).get("returns") or []
        self.assertEqual([row.get("claim_id") for row in platform_returns], ["CLAIM-LIVE-AMZ-PLATFORM"])

        amount_result = run_agent_turn(
            "Bugunku iadelerin toplam tutari ne kadar?",
            planner_client=client,
            chat_client=client,
            db_name=self.db_path,
        )
        self.assertTrue(amount_result.success, amount_result.message)
        self.assertEqual(((amount_result.data or {}).get("tool_call") or {}).get("tool"), "returns.amount_total")
        self.assertEqual(((amount_result.data or {}).get("tool_result") or {}).get("total_amount"), 440.0)

        before_calls = len(client.prompts)
        last_used_prepare = run_agent_turn(
            "Trendyol hesaplarinin son kullanim zamanini guncelle",
            planner_client=client,
            chat_client=client,
            db_name=self.db_path,
        )
        self.assertGreater(len(client.prompts) - before_calls, 0, provider_label)
        self._assert_prepare_answer_is_safe(last_used_prepare, expected_terms=("Hesap kullanim zamani",))
        self.assertNotIn("orders.update_last_used_at_for_accounts", self._answer_from(last_used_prepare))
        last_used_action = self._action_from(last_used_prepare)
        self.assertEqual(last_used_action.get("action_type"), "orders.update_last_used_at_for_accounts")
        last_used_payload = ((last_used_action.get("arguments") or {}).get("payload") or {})
        self.assertEqual(last_used_payload.get("api_account_ids"), [10])
        self.assertNotIn("accounts", last_used_payload)

        with tempfile.TemporaryDirectory() as temp_dir:
            audit_path = Path(temp_dir) / "orders-agent-audit.jsonl"
            last_used_execute = operation_runtime.execute_approved_operation(
                last_used_action,
                approved_by="live-behavioral-test",
                idempotency_key="live-last-used-trendyol",
                audit_path=audit_path,
            )
            audit_text = audit_path.read_text(encoding="utf-8")
        self.assertTrue(last_used_execute.success, last_used_execute.message)
        last_used_verify = ((last_used_execute.data or {}).get("verify") or {})
        self.assertTrue(last_used_verify.get("verified"), last_used_execute.data)
        self.assertEqual(last_used_verify.get("updated_pks"), [10])
        self.assertNotIn("api-key-live", audit_text)
        self.assertNotIn("api-secret-live", audit_text)
        with Session(get_engine(self.db_path)) as session:
            self.assertIsNotNone(getattr(session.get(ApiAccount, 10), "last_used_at", None))
            self.assertIsNone(getattr(session.get(ApiAccount, 20), "last_used_at", None))

    def test_real_cloud_llm_answers_and_approved_operations_finish(self) -> None:
        now_ms = self._seed_ready_orders()
        self._insert_order_item("ORD-MAHMUT-READY", now_ms + 2_000, product_name="Live Mahmut Test Urunu")
        self._insert_waiting_return(
            "CLAIM-MAHMUT-RETURN",
            "ORD-MAHMUT-READY",
            now_ms + 2_000,
            first_name="Mahmut",
            last_name="Şevket",
        )
        self._insert_waiting_return(
            "CLAIM-AYSE-RETURN",
            "ORD-AYSE-READY",
            now_ms + 3_000,
            first_name="Ayse",
            last_name="Yilmaz",
        )
        client = self._client()
        provider_label = _safe_provider_label(self._provider_record, self._provider_source)

        before_calls = len(client.prompts)
        advice_result = run_agent_turn(
            "valla bilmiyorum ki sence?",
            planner_client=client,
            chat_client=client,
            db_name=self.db_path,
            history=[
                {"role": "user", "content": "naber dayi napiyon"},
                {"role": "assistant", "content": "Iyiyim, buradayim."},
            ],
        )
        self.assertGreater(len(client.prompts) - before_calls, 0, provider_label)
        self.assertTrue(advice_result.success, advice_result.message)
        self.assertEqual((advice_result.data or {}).get("mode"), "answer")
        self.assertIsNone((advice_result.data or {}).get("tool_call"))
        self.assertNotIn("siparis var", self._answer_from(advice_result).casefold())

        analysis_result = run_agent_turn(
            "siparislerim nasil sence",
            planner_client=client,
            chat_client=client,
            db_name=self.db_path,
        )
        self.assertTrue(analysis_result.success, analysis_result.message)
        self.assertEqual((analysis_result.data or {}).get("mode"), "tools")
        analysis_tools = [row.get("tool") for row in ((analysis_result.data or {}).get("tool_results") or [])]
        self.assertIn("orders.count", analysis_tools)
        self.assertIn("orders.revenue_total", analysis_tools)
        self.assertIn("returns.count", analysis_tools)
        self.assertIn("orders.ready_to_ship_count", analysis_tools)

        before_calls = len(client.prompts)
        chat_result = run_agent_turn(
            "adin ne",
            planner_client=client,
            chat_client=client,
            db_name=self.db_path,
        )
        self.assertGreater(len(client.prompts) - before_calls, 0, provider_label)
        self.assertTrue(chat_result.success, chat_result.message)
        self.assertEqual((chat_result.data or {}).get("mode"), "answer")
        self.assertIsNone((chat_result.data or {}).get("tool_call"))
        chat_answer = self._answer_from(chat_result)
        self.assertIn("OrderScout", chat_answer)
        self.assertNotEqual(chat_answer, "Buradayim. Ne yapalim?")

        help_result = run_agent_turn(
            "siparisler tarafinda hangi islemleri yapabiliyorsun",
            planner_client=client,
            chat_client=client,
            db_name=self.db_path,
        )
        self.assertTrue(help_result.success, help_result.message)
        self.assertEqual((help_result.data or {}).get("mode"), "answer")
        self.assertIsNone((help_result.data or {}).get("tool_call"))
        help_answer = self._answer_from(help_result)
        self.assertIn("etiket", help_answer.casefold())
        self.assertIn("iade", help_answer.casefold())
        self.assertIn("kargo", help_answer.casefold())
        self.assertNotIn("Lokal kayitta", help_answer)

        before_calls = len(client.prompts)
        lookup = run_agent_turn(
            "Mahmut Şevket isimli müşterinin siparişini göster",
            planner_client=client,
            chat_client=client,
            db_name=self.db_path,
        )
        self.assertGreater(len(client.prompts) - before_calls, 0, provider_label)
        self.assertTrue(lookup.success, lookup.message)
        lookup_answer = (lookup.data or {}).get("answer") or lookup.message
        self.assertIn("ORD-MAHMUT-READY", lookup_answer)
        self.assertIn("Mahmut", lookup_answer)
        self.assertNotIn("ORD-AYSE-READY", lookup_answer)

        history = [
            {"role": "user", "content": "Kargoya hazır siparişleri getir"},
            {
                "role": "assistant",
                "content": "Kargoya hazir siparisler: ORD-MAHMUT-READY Mahmut Şevket, ORD-AYSE-READY Ayse Yilmaz.",
            },
        ]
        before_calls = len(client.prompts)
        followup = run_agent_turn(
            "Mahmut olani birak",
            planner_client=client,
            chat_client=client,
            db_name=self.db_path,
            history=history,
        )
        self.assertGreater(len(client.prompts) - before_calls, 0, provider_label)
        self.assertTrue(followup.success, followup.message)
        followup_answer = (followup.data or {}).get("answer") or followup.message
        self.assertIn("ORD-AYSE-READY", followup_answer)
        self.assertNotIn("ORD-MAHMUT-READY", followup_answer)

        before_calls = len(client.prompts)
        return_list = run_agent_turn(
            "Bekleyen iadeleri getir",
            planner_client=client,
            chat_client=client,
            db_name=self.db_path,
        )
        self.assertGreater(len(client.prompts) - before_calls, 0, provider_label)
        self.assertTrue(return_list.success, return_list.message)
        self.assertEqual(((return_list.data or {}).get("tool_call") or {}).get("tool"), "returns.query")
        return_list_answer = (return_list.data or {}).get("answer") or return_list.message
        self.assertIn("CLAIM-MAHMUT-RETURN", return_list_answer)
        self.assertIn("CLAIM-AYSE-RETURN", return_list_answer)

        before_calls = len(client.prompts)
        return_lookup = run_agent_turn(
            "Mahmut olanin iade detayini getir",
            planner_client=client,
            chat_client=client,
            db_name=self.db_path,
            history=[
                {"role": "user", "content": "Bekleyen iadeleri getir"},
                {"role": "assistant", "content": return_list_answer},
            ],
        )
        self.assertGreater(len(client.prompts) - before_calls, 0, provider_label)
        self.assertTrue(return_lookup.success, return_lookup.message)
        self.assertEqual(((return_lookup.data or {}).get("tool_call") or {}).get("tool"), "returns.lookup")
        self.assertEqual((((return_lookup.data or {}).get("tool_call") or {}).get("arguments") or {}).get("query"), "CLAIM-MAHMUT-RETURN")
        lookup_returns = ((return_lookup.data or {}).get("tool_result") or {}).get("returns") or []
        self.assertEqual([row.get("claim_id") for row in lookup_returns], ["CLAIM-MAHMUT-RETURN"])
        return_lookup_answer = (return_lookup.data or {}).get("answer") or return_lookup.message
        self.assertIn("CLAIM-MAHMUT-RETURN", return_lookup_answer)
        self.assertNotIn("CLAIM-AYSE-RETURN", return_lookup_answer)

        label_calls: list[list[str]] = []

        def fake_create_order_label_from_orders(order_numbers, *, brand_code="TANEX", model_code="TANEX_2736", db_name=None):
            label_calls.append(list(order_numbers))
            return Result.ok(
                "Etiketler olusturuldu.",
                close_dialog=False,
                data={"order_numbers": list(order_numbers), "brand_code": brand_code, "model_code": model_code},
            )

        label_prepare = run_agent_turn(
            "Kargoya hazır siparişleri yazdır",
            planner_client=client,
            chat_client=client,
            db_name=self.db_path,
        )
        label_action = self._action_from(label_prepare)
        with tempfile.TemporaryDirectory() as temp_dir, patch(
            "Labels.processors.pipeline.create_order_label_from_orders",
            side_effect=fake_create_order_label_from_orders,
        ):
            label_execute = operation_runtime.execute_approved_operation(
                label_action,
                approved_by="live-behavioral-test",
                idempotency_key="live-print-ready-orders",
                audit_path=Path(temp_dir) / "orders-agent-audit.jsonl",
            )
        self.assertTrue(label_execute.success, label_execute.message)
        self.assertEqual(label_calls, [["ORD-AYSE-READY", "ORD-MAHMUT-READY"]])

        before_calls = len(client.prompts)
        extracted_prepare = run_agent_turn(
            "bunlari fis cikarildi olarak isaretle",
            planner_client=client,
            chat_client=client,
            db_name=self.db_path,
            history=[
                {"role": "user", "content": "Kargoya hazir siparisleri getir"},
                {"role": "assistant", "content": "Kargoya hazir siparisler: ORD-MAHMUT-READY Mahmut Sevket, ORD-AYSE-READY Ayse Yilmaz."},
            ],
        )
        self.assertGreater(len(client.prompts) - before_calls, 0, provider_label)
        self._assert_prepare_answer_is_safe(extracted_prepare, expected_terms=("2 siparis",))
        extracted_action = self._action_from(extracted_prepare)
        self.assertEqual(extracted_action.get("action_type"), "labels.mark_orders_extracted_from_label_result")
        extracted_payload = ((extracted_action.get("arguments") or {}).get("payload") or {})
        self.assertEqual(set(extracted_payload.get("order_numbers") or []), {"ORD-MAHMUT-READY", "ORD-AYSE-READY"})
        with tempfile.TemporaryDirectory() as temp_dir:
            extracted_execute = operation_runtime.execute_approved_operation(
                extracted_action,
                approved_by="live-behavioral-test",
                idempotency_key="live-mark-label-extracted",
                audit_path=Path(temp_dir) / "orders-agent-audit.jsonl",
            )
        self.assertTrue(extracted_execute.success, extracted_execute.message)
        self.assertTrue(((extracted_execute.data or {}).get("verify") or {}).get("verified"), extracted_execute.data)
        with Session(get_engine(self.db_path)) as session:
            headers = session.exec(
                select(OrderHeader).where(OrderHeader.orderNumber.in_(["ORD-MAHMUT-READY", "ORD-AYSE-READY"]))
            ).all()
            self.assertEqual({row.orderNumber for row in headers if row.is_extracted}, {"ORD-MAHMUT-READY", "ORD-AYSE-READY"})

        explicit_label_calls: list[list[str]] = []

        def fake_create_label_for_explicit_order(order_numbers, *, brand_code="TANEX", model_code="TANEX_2736"):
            explicit_label_calls.append(list(order_numbers))
            return Result.ok(
                "Etiketler olusturuldu.",
                close_dialog=False,
                data={"order_numbers": list(order_numbers), "brand_code": brand_code, "model_code": model_code},
            )

        before_calls = len(client.prompts)
        explicit_label_prepare = run_agent_turn(
            "ORD-MAHMUT-READY etiketini yazdir",
            planner_client=client,
            chat_client=client,
            db_name=self.db_path,
        )
        self.assertGreater(len(client.prompts) - before_calls, 0, provider_label)
        explicit_label_action = self._action_from(explicit_label_prepare)
        explicit_payload = ((explicit_label_action.get("arguments") or {}).get("payload") or {})
        self.assertEqual(explicit_payload.get("order_numbers"), ["ORD-MAHMUT-READY"])
        with tempfile.TemporaryDirectory() as temp_dir, patch(
            "Labels.processors.pipeline.create_order_label_from_orders",
            side_effect=fake_create_label_for_explicit_order,
        ):
            explicit_label_execute = operation_runtime.execute_approved_operation(
                explicit_label_action,
                approved_by="live-behavioral-test",
                idempotency_key="live-print-explicit-order",
                audit_path=Path(temp_dir) / "orders-agent-audit.jsonl",
            )
        self.assertTrue(explicit_label_execute.success, explicit_label_execute.message)
        self.assertEqual(explicit_label_calls, [["ORD-MAHMUT-READY"]])

        with tempfile.TemporaryDirectory() as temp_dir:
            output_path = Path(temp_dir) / "live-mahmut-profit.xlsx"
            prompt_path = str(output_path).replace("\\", "/")
            export_prepare = run_agent_turn(
                f'ORD-MAHMUT-READY siparis kar Excel raporunu "{prompt_path}" dosyasina aktar',
                planner_client=client,
                chat_client=client,
                db_name=self.db_path,
            )
            export_action = self._action_from(export_prepare)
            self.assertEqual(export_action.get("action_type"), "orders.action_export_order_profit_excel")
            export_execute = operation_runtime.execute_approved_operation(
                export_action,
                approved_by="live-behavioral-test",
                idempotency_key="live-profit-export-mahmut",
                audit_path=Path(temp_dir) / "orders-agent-audit.jsonl",
            )
            self.assertTrue(export_execute.success, export_execute.message)
            self.assertTrue(output_path.exists(), export_execute.data)
            export_verify = ((export_execute.data or {}).get("verify") or {})
            self.assertTrue(export_verify.get("verified"), export_execute.data)

        bulk_calls: list[dict] = []

        def fake_change_orders_cargo_provider_bulk(orders, *, cargo_provider_code: str):
            order_numbers = [order.orderNumber for order in orders]
            bulk_calls.append({"orders": order_numbers, "cargo_provider_code": cargo_provider_code})
            with Session(get_engine(self.db_path)) as session:
                records = session.exec(select(OrderData).where(OrderData.orderNumber.in_(order_numbers))).all()
                for record in records:
                    record.cargoProviderName = "Aras"
                    session.add(record)
                session.commit()
            return Result.ok("Toplu kargo firmasi degistirildi.", close_dialog=False, data={"orders": order_numbers})

        bulk_prepare = run_agent_turn(
            "Kargoya hazır siparişlerin hepsinin kargo firmasını Aras yap",
            planner_client=client,
            chat_client=client,
            db_name=self.db_path,
        )
        bulk_action = self._action_from(bulk_prepare)
        self.assertEqual(bulk_action.get("action_type"), "orders.change_orders_cargo_provider_bulk")
        with tempfile.TemporaryDirectory() as temp_dir, patch(
            "Orders.views.actions.change_orders_cargo_provider_bulk",
            side_effect=fake_change_orders_cargo_provider_bulk,
        ):
            bulk_execute = operation_runtime.execute_approved_operation(
                bulk_action,
                approved_by="live-behavioral-test",
                idempotency_key="live-bulk-cargo-ready-aras",
                audit_path=Path(temp_dir) / "orders-agent-audit.jsonl",
            )
        self.assertTrue(bulk_execute.success, bulk_execute.message)
        self.assertEqual(len(bulk_calls), 1)
        self.assertEqual(set(bulk_calls[0].get("orders") or []), {"ORD-AYSE-READY", "ORD-MAHMUT-READY"})

        def fake_approve_return_claim(*, order, **_kwargs):
            with Session(get_engine(self.db_path)) as session:
                claim = session.exec(select(ReturnClaim).where(ReturnClaim.claimId == order.claimId)).one()
                claim.claimItemStatus = "Accepted"
                claim.waitingClaimLineItemIds = []
                session.add(claim)
                session.commit()
            return Result.ok("Iade onaylandi.", close_dialog=False, data={"claim_id": order.claimId})

        return_prepare = run_agent_turn(
            "Mahmut Şevket'in iadesini onayla",
            planner_client=client,
            chat_client=client,
            db_name=self.db_path,
        )
        return_action = self._action_from(return_prepare)
        with tempfile.TemporaryDirectory() as temp_dir, patch(
            "Orders.views.actions.approve_return_claim",
            side_effect=fake_approve_return_claim,
        ):
            return_execute = operation_runtime.execute_approved_operation(
                return_action,
                approved_by="live-behavioral-test",
                idempotency_key="live-return-mahmut-approve",
                audit_path=Path(temp_dir) / "orders-agent-audit.jsonl",
            )
        self.assertTrue(return_execute.success, return_execute.message)
        with Session(get_engine(self.db_path)) as session:
            claim = session.exec(select(ReturnClaim).where(ReturnClaim.claimId == "CLAIM-MAHMUT-RETURN")).one()
            self.assertEqual(claim.claimItemStatus, "Accepted")
            self.assertEqual(claim.waitingClaimLineItemIds, [])

        with Session(get_engine(self.db_path)) as session:
            session.add(ApiAccount(pk=1, account_id="seller-live-stale", comp_name="Live Stale Seller", platform="trendyol"))
            session.add(
                ReturnClaim(
                    claimId="CLAIM-LIVE-STALE",
                    orderNumber="ORD-MAHMUT-READY",
                    api_account_id=1,
                    claimDate=now_ms + 4_000,
                    orderDate=now_ms + 4_000,
                    lastModifiedDate=now_ms + 5_000,
                    claimItemStatus="WaitingInAction",
                    customerFirstName="Mahmut",
                    customerLastName="Şevket",
                    totalPrice=55.0,
                    itemCount=1,
                    waitingClaimLineItemIds=["line-stale"],
                    sync_deleted_at="2026-01-01T00:00:00+03:00",
                )
            )
            session.commit()

        stale_review = run_agent_turn(
            "stale iadeleri goster",
            planner_client=client,
            chat_client=client,
            db_name=self.db_path,
        )
        self.assertTrue(stale_review.success, stale_review.message)
        self.assertEqual(((stale_review.data or {}).get("tool_call") or {}).get("tool"), "operations.manual_action.route")
        stale_review_tool = (stale_review.data or {}).get("tool_result") or {}
        self.assertEqual(stale_review_tool.get("route_status"), "read_executed")
        self.assertEqual(((stale_review_tool.get("match") or {}).get("capability_id")), "orders.get_stale_return_claims")
        self.assertIn("CLAIM-LIVE-STALE", self._answer_from(stale_review))
        with Session(get_engine(self.db_path)) as session:
            stale_claim = session.exec(select(ReturnClaim).where(ReturnClaim.claimId == "CLAIM-LIVE-STALE")).one()
            self.assertEqual(stale_claim.sync_deleted_at, "2026-01-01T00:00:00+03:00")

        restore_prepare = run_agent_turn(
            "CLAIM-LIVE-STALE stale iadesini geri al",
            planner_client=client,
            chat_client=client,
            db_name=self.db_path,
        )
        restore_action = self._action_from(restore_prepare)
        self.assertEqual(restore_action.get("action_type"), "orders.restore_return_claims_from_stale")
        with tempfile.TemporaryDirectory() as temp_dir:
            restore_execute = operation_runtime.execute_approved_operation(
                restore_action,
                approved_by="live-behavioral-test",
                idempotency_key="live-restore-stale-claim",
                audit_path=Path(temp_dir) / "orders-agent-audit.jsonl",
            )
        self.assertTrue(restore_execute.success, restore_execute.message)
        with Session(get_engine(self.db_path)) as session:
            stale_claim = session.exec(select(ReturnClaim).where(ReturnClaim.claimId == "CLAIM-LIVE-STALE")).one()
            self.assertIsNone(stale_claim.sync_deleted_at)


if __name__ == "__main__":
    unittest.main()
