from __future__ import annotations

from copy import deepcopy
from dataclasses import replace
import json
from pathlib import Path
from unittest.mock import patch

import pytest

from AI.processors import file_workspace, observation_artifacts, task_runtime, tools
from AI.processors.observation_projection import source_projection_hash
from AI.processors.session_state import AgentSessionState, AgentTaskStep
from AI.processors.task_store import load_agent_task_state
from Core.processors.report_artifacts import verify_report_snapshot_xlsx


@pytest.fixture
def isolated_workspace(tmp_path, monkeypatch):
    monkeypatch.setattr(file_workspace, "ATTACHMENTS_ROOT", tmp_path / "imports")
    monkeypatch.setattr(observation_artifacts, "OBSERVATION_ARTIFACTS_ROOT", tmp_path / "observations")


def _interpretation(mode="inspect", fields=None):
    return {"mode": mode, "domain": "general", "intent_kind": "action" if mode == "act" else "data_read",
            "valid": True, "requested_fields": fields or ["records"], "requires_clarification": False}


def _planner(plan, captured=None):
    def planner(prompt):
        if captured is not None:
            captured.append(prompt)
        if prompt.startswith("OrderScout semantic action verifier"):
            return json.dumps({"verdict": "allow", "objective_aligned": True, "domain_aligned": True,
                               "unwanted_effect": False, "reason": "User requested this observation report"})
        if prompt.startswith("GoalPlan observation"):
            return json.dumps({"decision": "complete" if "Kalan plan: []\n" in prompt else "continue",
                               "reason": "The requested read graph is observed"})
        return json.dumps(plan)
    return planner


def _read_attachment(tmp_path, *, persist=False, callback=None, max_replans=None, records=None):
    if records is None:
        records = [{"sku": "001", "quantity": 6, "unit": "adet", "amount": 10.5, "currency": "TRY", "unknown": None},
                   {"sku": "002", "quantity": 4, "unit": "adet", "amount": 9, "currency": "USD", "unknown": None}]
    selected = tmp_path / "chosen.json"
    selected.write_text(json.dumps(records), encoding="utf-8")
    previous = AgentSessionState(goal="Import the selected file", active_domain="general")
    metadata = file_workspace.register_attachment(selected, previous.session_id)
    previous.active_entities["file_workspace"] = {"session_id": previous.session_id, "attachments": [metadata]}
    plan = {"steps": [{"id": "read_selected", "operation_id": "files.read", "arguments": {"attachment_id": metadata["attachment_id"]},
                       "depends_on": [], "objective": "Read the actual selected JSON source", "expected_evidence": ["rows", "source_hash"]}]}
    prompts = []
    result = task_runtime.run_agent_task(
        "Read the selected JSON file", domain="general", session_state=previous,
        planner_client=_planner(plan, prompts), turn_interpretation=_interpretation(),
        persist=persist, task_store_root=str(tmp_path / "tasks"), task_created_callback=callback, max_replans=max_replans,
    )
    return result, records, metadata, prompts


def test_granted_dictionary_can_become_exact_excel_rows_and_keep_delivery_proof(isolated_workspace, tmp_path):
    from AI.processors.observation_context import verified_export_report_answer
    from openpyxl import load_workbook
    first, _, _, _ = _read_attachment(tmp_path, records=[{"counts": {"Picking": 238, "Cancelled": 86, "Empty": 0}}])
    assert first.success
    previous = AgentSessionState.from_dict(first.data["state"])
    previous.record_turn(role="assistant", content="Picking 238; Cancelled 86; Empty 0.")
    ref = observation_artifacts.observation_artifact_refs(previous)[0]
    output = tmp_path / "dictionary.xlsx"
    source = {"source_id": "prior", "source_step_id": "read", "operation_id": "artifacts.read",
              "result": {"$from_step": "read", "path": "tool_result"}}
    plan = {"steps": [
        {"id": "read", "operation_id": "artifacts.read", "arguments": {"artifact_id": ref["artifact_id"]},
         "depends_on": [], "objective": "Read granted source", "expected_evidence": ["result"]},
        {"id": "report", "operation_id": "reports.build_snapshot", "depends_on": ["read"],
         "objective": "Project original dictionary entries", "expected_evidence": ["report_snapshot"],
         "arguments": {"title": "Original distribution", "sources": [source], "sheets": [
             {"name": "Distribution", "source_id": "prior", "rows_path": "result.rows.0.counts", "rows_mode": "object_entries",
              "columns": [{"key": "status", "label": "Status", "source_path": "key", "data_type": "string"},
                          {"key": "count", "label": "Count", "source_path": "value", "data_type": "integer"}]}]}},
        {"id": "prepare", "operation_id": "reports.xlsx.prepare", "depends_on": ["report"],
         "objective": "Prepare Excel", "expected_evidence": ["action"],
         "arguments": {"report_snapshot": {"$from_step": "report", "path": "tool_result.report_snapshot"}, "output_path": str(output)}}]}
    interpretation = {**_interpretation("act", ["artifact"]), "operation_policy": "allow_new_operations",
                      "source_policy": "existing_verified_sources_only"}
    prompts = []
    started = task_runtime.run_agent_task("Create Excel from those numbers", domain="general", session_state=previous,
        planner_client=_planner(plan, prompts), turn_interpretation=interpretation)
    assert started.success and started.data["state"]["status"] == "pending_approval", started.message
    assert any("Picking 238" in prompt and ref["artifact_id"] in prompt for prompt in prompts)
    assert any("Existing source reference (DATA):" in prompt and ref["artifact_id"] in prompt for prompt in prompts)
    assert not output.exists()
    approved = task_runtime.resume_agent_task_with_approval(started.data["state"], approved_by="test",
        audit_path=str(tmp_path / "audit.jsonl"))
    assert approved.success, approved.message
    book = load_workbook(output, read_only=True, data_only=True)
    try:
        assert list(book["Distribution"].values) == [("Status", "Count"), ("Cancelled", 86), ("Empty", 0), ("Picking", 238)]
    finally:
        book.close()
    assert verified_export_report_answer(AgentSessionState.from_dict(approved.data["state"]))
    from AI.processors.response_verifier import project_fallback_response, verify_response_projection
    final_state = AgentSessionState.from_dict(approved.data["state"])
    # These are workbook dimensions, not a demand to repeat all item rows in chat.
    final_state.active_entities["turn_interpretation"].update(request_shape="artifact_report", requested_fields=["artifact", "item_sales_rows"])
    projected = project_fallback_response(final_state, "Report produced")
    assert projected == verified_export_report_answer(final_state)
    assert verify_response_projection(projected, final_state).success
    answer = task_runtime._finalize_task_answer(final_state, chat_client=lambda _: "The Excel file has been produced.")
    assert verify_response_projection(answer, final_state).success
    assert final_state.active_entities["response_synthesis"]["final_projection_verified"] is True
    changed = AgentSessionState.from_dict(final_state.to_dict())
    changed.action_history[-1]["verify"]["verified"] = False
    assert not verify_response_projection(projected, changed).success


@pytest.mark.parametrize("restricted_sources", [False, True])
def test_actual_completed_file_read_is_published_and_new_task_reports_historical_source_after_current_approval(isolated_workspace, tmp_path, restricted_sources):
    from openpyxl import load_workbook

    first, records, metadata, first_prompts = _read_attachment(tmp_path, persist=True)
    assert first.success, first.message
    previous = AgentSessionState.from_dict(first.data["state"])
    assert previous.status == "complete"
    refs = observation_artifacts.observation_artifact_refs(previous)
    assert len(refs) == 1
    ref = refs[0]
    original = previous.completed_steps[0].result_data["tool_result"]
    assert ref["origin"]["operation_id"] == "files.read"
    assert ref["result_hash"] == source_projection_hash(original)
    assert original["rows"] == records and original["source_hash"] == metadata["sha256"]
    restored = load_agent_task_state(previous.session_id, root=tmp_path / "tasks")
    assert observation_artifacts.observation_artifact_refs(restored) == refs
    output = tmp_path / "historical-report.xlsx"
    source = {"source_id": "previous_file", "source_step_id": "historical_read", "operation_id": "artifacts.read",
              "result": {"$from_step": "historical_read", "path": "tool_result"}}
    plan = {"steps": [
        {"id": "historical_read", "operation_id": "artifacts.read", "arguments": {"artifact_id": ref["artifact_id"]},
         "depends_on": [], "objective": "Read the granted historical observation", "expected_evidence": ["result", "observed_at"]},
        {"id": "report", "operation_id": "reports.build_snapshot",
         "arguments": {"title": "Historical file observation", "sources": [source], "sheets": [
             {"name": "Previous source", "source_id": "previous_file", "rows_path": "result.rows", "columns": [
                 {"key": key, "label": key, "source_path": key, "data_type": kind}
                 for key, kind in [("sku", "string"), ("quantity", "integer"), ("unit", "string"),
                                   ("amount", "number"), ("currency", "string"), ("unknown", "number")]]}]},
         "depends_on": ["historical_read"], "objective": "Project the previous records without a new source query",
         "expected_evidence": ["report_snapshot", "source_receipts"]},
        {"id": "prepare", "operation_id": "reports.xlsx.prepare",
         "arguments": {"report_snapshot": {"$from_step": "report", "path": "tool_result.report_snapshot"}, "output_path": str(output)},
         "depends_on": ["report"], "objective": "Prepare the user requested Excel file", "expected_evidence": ["action", "snapshot_id"]},
    ]}
    prompts, progress = [], []
    registry = tools.available_tools()
    actual_file_reader, fresh_source_calls = registry["files.read"].handler, []
    registry["files.read"] = replace(registry["files.read"], handler=lambda **kwargs: fresh_source_calls.append(kwargs) or actual_file_reader(**kwargs))
    interpretation = _interpretation("act", ["artifact"])
    if restricted_sources:
        interpretation.update(source_policy="existing_verified_sources_only", operation_policy="allow_new_operations")
    with patch("AI.processors.tools.available_tools", return_value=registry):
        started = task_runtime.run_agent_task(
            "Create an Excel report from the previous file observation; retain its historical time and source",
            domain="general", session_state=restored, planner_client=_planner(plan, prompts), persist=False,
            turn_interpretation=interpretation,
            progress_callback=lambda percent, message: progress.append({"percent": percent, "message": message}),
        )
        assert started.success, started.message
        assert fresh_source_calls == []
        current = AgentSessionState.from_dict(started.data["state"])
        if restricted_sources:
            assert current.active_entities["source_access_policy"] == {
                "version": 1, "source_policy": "existing_verified_sources_only"}
            current = AgentSessionState.from_dict(current.to_dict())
        assert current.session_id != previous.session_id
        assert current.status == "pending_approval", [(s.step_id, s.status, s.result_message) for s in current.completed_steps]
        assert [s.tool_names for s in current.completed_steps] == [["artifacts.read"], ["reports.build_snapshot"], ["reports.xlsx.prepare"]]
        historical = current.completed_steps[0].result_data["tool_result"]
        assert historical["result"] == original
        assert historical["fresh_read"] is False and historical["freshness"] == "historical_snapshot"
        assert historical["observed_at"] == ref["observed_at"]
        assert historical["artifact"]["origin"]["task_id"] == previous.session_id
        assert not output.exists()
        snapshot = current.completed_steps[1].result_data["tool_result"]["report_snapshot"]
        audit_path = tmp_path / "audit.jsonl"
        approved = task_runtime.resume_agent_task_with_approval(
            current.to_dict(), approved_by="isolated-historical-report-test", persist=False,
            idempotency_key="historical-report-exact", audit_path=str(audit_path),
        )
        assert approved.success, approved.message
    assert approved.data["state"]["status"] == "complete"
    assert verify_report_snapshot_xlsx(snapshot, output)["verified"]
    assert approved.data["state"]["action_history"][-1]["verify"]["verified_by"] == "xlsx_exact_snapshot_readback"
    from AI.processors.observation_context import verified_export_report_answer
    explanation = verified_export_report_answer(AgentSessionState.from_dict(approved.data["state"]))
    assert "chosen.json" in explanation and ref["observed_at"] in explanation
    assert "geçmiş kayda dayanıyor" in explanation and "yeni kaynak sorgusu yapılmadı" in explanation
    assert "şu anki varlığı yeniden kontrol edilmedi" in explanation
    workbook = load_workbook(output, read_only=True, data_only=True)
    try:
        assert list(workbook["Previous source"].values)[1:] == [("001", 6, "adet", 10.5, "TRY", None), ("002", 4, "adet", 9, "USD", None)]
        proof_rows = list(workbook["Kaynak Kanıtı"].values)
        proof = dict(zip(proof_rows[0], proof_rows[1]))
        assert proof["Kaynak güncelliği"] == "historical_snapshot"
        assert proof["Asıl gözlem zamanı"] == ref["observed_at"]
        assert proof["Asıl kaynak operasyonu"] == "files.read"
        assert proof["Yeni kaynak sorgusu"] is False
    finally:
        workbook.close()
    all_caps = previous.active_entities["observation_artifacts"]["capabilities"] + current.active_entities["observation_artifacts"]["capabilities"]
    public = json.dumps({"prompts": first_prompts + prompts, "progress": progress,
                         "steps": [s.to_dict() for s in current.completed_steps], "actions": approved.data["state"]["action_history"]})
    public += audit_path.read_text(encoding="utf-8")
    assert "authorized_capabilities" not in public and "capability_token" not in public
    for capability in all_caps:
        assert capability["capability_token"] not in public
    assert any(ref["artifact_id"] in prompt and "historical" in prompt for prompt in prompts)


@pytest.mark.parametrize("kind", ["foreign", "token", "hash", "session", "model_private_kwargs"])
def test_runtime_rejects_ungranted_or_forged_artifact_before_handler(isolated_workspace, tmp_path, kind):
    first, _, _, _ = _read_attachment(tmp_path)
    assert first.success
    previous = AgentSessionState.from_dict(first.data["state"])
    ref = observation_artifacts.observation_artifact_refs(previous)[0]
    if kind == "foreign":
        current = AgentSessionState(goal="Different conversation", active_domain="general")
    else:
        current = task_runtime._fresh_task_state_from_conversation(goal="Use the prior observation", domain="general", session_state=previous)
        capability = current.active_entities["observation_artifacts"]["capabilities"][0]
        if kind != "model_private_kwargs":
            capability[{"token": "capability_token", "hash": "artifact_hash", "session": "session_id"}[kind]] = "0" * 64
    arguments = {"artifact_id": ref["artifact_id"]}
    if kind == "model_private_kwargs":
        arguments["authorized_capabilities"] = previous.active_entities["observation_artifacts"]["capabilities"]
        arguments["session_id"] = previous.session_id
    registry = tools.available_tools()
    calls = []
    actual = registry["artifacts.read"]
    registry["artifacts.read"] = replace(actual, handler=lambda **kwargs: calls.append(kwargs) or actual.handler(**kwargs))
    step = AgentTaskStep(index=1, step_id="historical_read", prompt="Read the stored observation", operation_domain="artifacts",
                         operation_candidates=["artifacts.read"], arguments=arguments)
    with patch("AI.processors.tools.available_tools", return_value=registry):
        result = task_runtime._run_bound_domain_read_step(current, step, db_name=None)
    assert result is not None and not result.success
    assert calls == []
    assert "capability_token" not in json.dumps(result.data)


def test_fresh_workflow_invocation_preserves_trusted_metadata_but_inherits_no_previous_grants(isolated_workspace, tmp_path):
    first, _, _, _ = _read_attachment(tmp_path)
    previous = AgentSessionState.from_dict(first.data["state"])
    previous.completed_steps.clear()
    metadata = {"workflow_id": "wf_example", "revision_id": "rev_example", "spec_hash": "a" * 64,
                "fresh_model_plan_required": True, "requires_current_sources": True, "requires_current_approval": True,
                "untrusted_extra": {"approval_token": "fake"}}
    previous.active_entities["workflow_invocation"] = metadata
    current = task_runtime._fresh_task_state_from_conversation(goal="Run this saved workflow now", domain="general", session_state=previous)
    assert current.active_entities["workflow_invocation"] == {k: v for k, v in metadata.items() if k != "untrusted_extra"}
    assert "observation_artifacts" not in current.active_entities
    assert current.completed_steps == [] and current.pending_actions == []


def test_workflow_action_payload_is_not_mutated_by_database_injection():
    action = {"action_type": "ai.action_save_workflow_spec", "arguments": {"payload": {"spec": {"name": "Daily research"}}}}
    assert task_runtime._action_with_runtime_db(action, "a-db.sqlite") == action


def test_actual_task_creation_callback_gets_persisted_identity_before_any_plan_and_cannot_inject_execution(isolated_workspace, tmp_path):
    seen = []
    def callback(state):
        saved = load_agent_task_state(state.session_id, root=tmp_path / "tasks")
        assert saved.active_entities["task_checkpoint"]["phase"] == "initialized"
        assert saved.active_entities["planning_not_started"] is True
        assert saved.completed_steps == [] and saved.remaining_steps == []
        seen.append(state.session_id)
        state.session_id = "forged-by-callback"
        state.goal = "Run a different task"
        state.pending_actions.append({"approval_token": "forged"})
    result, _, _, prompts = _read_attachment(tmp_path, persist=True, callback=callback, max_replans=0)
    assert result.success, result.message
    assert seen == [result.data["state"]["session_id"]] and prompts
    assert result.data["state"]["pending_actions"] == []
    assert result.data["state"]["goal"] == "Read the selected JSON file"
    assert result.data["state"]["active_entities"]["execution_budget"]["max_replans"] == 0
    assert "planning_not_started" not in result.data["state"]["active_entities"]


def test_callback_failure_stops_before_model_or_tools_and_saved_initial_task_cannot_false_complete(isolated_workspace, tmp_path):
    initialized = []
    def callback(state):
        initialized.append(deepcopy(state.to_dict()))
        raise RuntimeError("Do not reveal callback details")
    result, _, _, prompts = _read_attachment(tmp_path, persist=True, callback=callback)
    assert not result.success and result.data["task_created_callback_failed"]
    assert prompts == [] and result.data["state"]["completed_steps"] == []
    state = load_agent_task_state(result.data["state"]["session_id"], root=tmp_path / "tasks")
    assert state.status == "failed" and state.active_entities["task_checkpoint"]["phase"] == "task_creation_failed"
    assert "Do not reveal callback details" not in json.dumps(result.data)
    interrupted = AgentSessionState.from_dict(initialized[0])
    interrupted.status = "interrupted"
    resumed = task_runtime.resume_read_task(interrupted, persist=False)
    assert not resumed.success and resumed.data["requires_fresh_plan"] is True
    assert resumed.data["state"]["status"] == "interrupted"


@pytest.mark.parametrize("budget", [-1, 65, True, "2", 1.5])
def test_invalid_explicit_replan_budget_does_not_invoke_provider(isolated_workspace, tmp_path, budget):
    result, _, _, prompts = _read_attachment(tmp_path, max_replans=budget)
    assert not result.success and result.data["invalid_task_replan_budget"]
    assert prompts == [] and result.data["state"]["completed_steps"] == []
