from __future__ import annotations

import json
import os
import re
import tempfile
import unittest
import unicodedata
from datetime import datetime
from pathlib import Path
from typing import Any, Callable
from unittest.mock import patch

from sqlmodel import Session, select

from Account.models.models import ApiAccount
from AI.processors import operation_runtime
from AI.processors.agent import run_agent_turn
from Core.utils.model_utils import get_engine
from Feedback.processors.pipeline import Result
from Orders.constants import TZ_GMT3
from Orders.models.trendyol.models import OrderData, OrderHeader, OrderItem, ReturnClaim
from tests.domain_ai.test_orders_agentic_live_eval import (
    _live_eval_enabled,
    _live_provider_record,
    _safe_provider_label,
    _write_live_artifact,
)
from tests.support.db_case import IsolatedDbTestCase


def _text(value: Any) -> str:
    return str(value or "").strip()


def _fold_text(value: Any) -> str:
    folded = unicodedata.normalize("NFKD", _text(value).casefold())
    return "".join(char for char in folded if not unicodedata.combining(char))


def _action_payload(action: dict[str, Any]) -> dict[str, Any]:
    arguments = action.get("arguments") if isinstance(action.get("arguments"), dict) else {}
    payload = arguments.get("payload")
    if isinstance(payload, dict):
        return dict(payload)
    rows = action.get("rows") if isinstance(action.get("rows"), list) else []
    first = rows[0] if rows and isinstance(rows[0], dict) else {}
    payload = first.get("payload")
    return dict(payload or {}) if isinstance(payload, dict) else {}


def _assert_no_raw_action_id(testcase: unittest.TestCase, answer: str, *action_ids: str) -> None:
    for action_id in action_ids:
        testcase.assertNotIn(action_id, answer, answer)


class _RecordingClient:
    def __init__(self, client: Callable[[str], str]):
        self._client = client
        self.prompts: list[str] = []
        self.responses: list[str] = []
        self.errors: list[str] = []

    def __call__(self, prompt: str) -> str:
        self.prompts.append(str(prompt or ""))
        try:
            response = self._client(prompt)
        except Exception as exc:
            self.errors.append(repr(exc))
            raise
        self.responses.append(str(response or ""))
        return response


class TestOrdersLiveConversationQuality(IsolatedDbTestCase):
    _provider_record: dict[str, Any]
    _provider_source: str

    @classmethod
    def setUpClass(cls) -> None:
        super().setUpClass()
        if not _live_eval_enabled():
            raise unittest.SkipTest("Set ORDERSCOUT_LIVE_LLM_EVAL=1 to run real cloud LLM conversation quality evals.")
        cls._provider_record, cls._provider_source = _live_provider_record()

    def _client(self) -> _RecordingClient:
        from AI.processors.cloud_runtime import build_provider_chat_client

        record = dict(self._provider_record)
        missing = [key for key in ("provider_code", "api_key", "default_model") if not _text(record.get(key))]
        if missing:
            self.fail("Live LLM provider record is missing required fields: " + ", ".join(missing))
        client = build_provider_chat_client(record, timeout=float(os.environ.get("ORDERSCOUT_LIVE_LLM_TIMEOUT", "120")))
        self.assertIsNotNone(client)
        return _RecordingClient(client)

    def _insert_order(
        self,
        order_number: str,
        order_date_ms: int,
        *,
        package_id: int,
        first_name: str,
        last_name: str,
        status: str = "Created",
        cargo_provider_name: str | None = "Trendyol Express",
        cargo_tracking_number: str | None = None,
        total_price: float = 100.0,
        printed: bool = False,
    ) -> None:
        with Session(get_engine(self.db_path)) as session:
            header = OrderHeader(orderNumber=order_number, is_printed=printed, is_extracted=False)
            session.add(header)
            session.flush()
            session.add(
                OrderData(
                    order_header_id=header.pk,
                    orderNumber=order_number,
                    id=package_id,
                    shipmentPackageId=package_id,
                    orderDate=order_date_ms,
                    lastModifiedDate=order_date_ms + package_id,
                    status=status,
                    shipmentPackageStatus=status,
                    customerFirstName=first_name,
                    customerLastName=last_name,
                    cargoProviderName=cargo_provider_name,
                    cargoTrackingNumber=f"TRK-{package_id}" if cargo_tracking_number is None else cargo_tracking_number,
                    totalPrice=total_price,
                    packageTotalPrice=total_price,
                )
            )
            session.commit()

    def _insert_waiting_return(
        self,
        claim_id: str,
        order_number: str,
        now_ms: int,
        *,
        first_name: str,
        last_name: str,
    ) -> None:
        with Session(get_engine(self.db_path)) as session:
            session.add(
                ReturnClaim(
                    claimId=claim_id,
                    orderNumber=order_number,
                    claimDate=now_ms,
                    orderDate=now_ms,
                    lastModifiedDate=now_ms + 1_000,
                    claimItemStatus="WaitingInAction",
                    customerFirstName=first_name,
                    customerLastName=last_name,
                    totalPrice=55.0,
                    itemCount=1,
                    waitingClaimLineItemIds=["line-1"],
                )
            )
            session.commit()

    def _insert_order_item(
        self,
        order_number: str,
        now_ms: int,
        *,
        line_id: int = 1,
        product_name: str = "Demo Urun",
    ) -> None:
        with Session(get_engine(self.db_path)) as session:
            session.add(
                OrderItem(
                    orderNumber=order_number,
                    order_data_id=1,
                    snapshotLastModifiedDate=now_ms,
                    quantity=2,
                    productName=product_name,
                    merchantSku="SKU-LIVE-CONVO",
                    barcode="BAR-LIVE-CONVO",
                    price=25.0,
                    lineId=line_id,
                )
            )
            session.commit()

    def _seed_ready_orders(self) -> int:
        now_ms = int(datetime.now(TZ_GMT3).timestamp() * 1000)
        self._insert_order(
            "ORD-MAHMUT-READY",
            now_ms,
            package_id=901,
            first_name="Mahmut",
            last_name="Sevket",
            status="ReadyToShip",
        )
        self._insert_order(
            "ORD-AYSE-READY",
            now_ms + 1_000,
            package_id=902,
            first_name="Ayse",
            last_name="Yilmaz",
            status="ReadyToShip",
        )
        return now_ms

    def _append_turn(self, transcript: list[dict[str, Any]], prompt: str, result) -> None:
        data = result.data or {}
        tool_call = data.get("tool_call") if isinstance(data.get("tool_call"), dict) else None
        tool_calls = data.get("tool_calls") if isinstance(data.get("tool_calls"), list) else []
        action = data.get("action_proposal") if isinstance(data.get("action_proposal"), dict) else {}
        payload = _action_payload(action)
        transcript.append(
            {
                "user": prompt,
                "success": bool(result.success),
                "mode": data.get("mode"),
                "approval_required": bool(data.get("approval_required")),
                "tool": (tool_call or {}).get("tool"),
                "tools": [str((row or {}).get("tool") or "") for row in tool_calls],
                "action_type": action.get("action_type"),
                "execute_enabled": action.get("execute_enabled"),
                "payload": {
                    key: payload.get(key)
                    for key in (
                        "order",
                        "orders",
                        "order_number",
                        "order_numbers",
                        "claim_ids",
                        "cargo_provider_code",
                        "api_account_ids",
                        "output_path",
                        "save_to_db",
                        "fetch_returns",
                    )
                    if key in payload
                },
                "answer": _text(data.get("answer") or result.message),
            }
        )
        if isinstance(payload.get("accounts"), list):
            transcript[-1]["payload"]["account_count"] = len(payload.get("accounts") or [])
        elif isinstance(payload.get("api_account_ids"), list):
            transcript[-1]["payload"]["account_count"] = len(payload.get("api_account_ids") or [])

    def _append_execute(self, transcript: list[dict[str, Any]], operation: str, result, *, extra: dict[str, Any] | None = None) -> None:
        data = result.data or {}
        verify = data.get("verify") if isinstance(data.get("verify"), dict) else {}
        row = {
            "phase": "execute",
            "operation": operation,
            "success": bool(result.success),
            "message": _text(result.message),
            "verify": {
                key: verify.get(key)
                for key in (
                    "verified",
                    "operation",
                    "order_number",
                    "order_numbers",
                    "claim_id",
                    "claim_ids",
                    "output_path",
                    "save_counts",
                    "restored_count",
                    "stale_count",
                    "updated_count",
                )
                if key in verify
            },
        }
        if extra:
            row["extra"] = dict(extra)
        transcript.append(row)

    def test_real_cloud_conversation_transcript_quality(self) -> None:
        now_ms = self._seed_ready_orders()
        self._insert_order(
            "ORD-LIVE-CONVO-EXTRA",
            now_ms + 2_000,
            package_id=2_222,
            first_name="Live",
            last_name="Conversation",
            status="ReadyToShip",
            total_price=75.0,
        )
        self._insert_waiting_return(
            "CLAIM-LIVE-CONVO-WAIT",
            "ORD-MAHMUT-READY",
            int(datetime.now(TZ_GMT3).timestamp() * 1000),
            first_name="Mahmut",
            last_name="Sevket",
        )

        client = self._client()
        provider_label = _safe_provider_label(self._provider_record, self._provider_source)
        transcript: list[dict[str, Any]] = []
        history: list[dict[str, str]] = []

        def turn(prompt: str):
            before = len(client.prompts)
            before_responses = len(client.responses)
            before_errors = len(client.errors)
            result = run_agent_turn(prompt, planner_client=client, chat_client=client, db_name=self.db_path, history=history)
            self.assertGreater(len(client.prompts) - before, 0, provider_label + f"\nPrompt did not call live LLM: {prompt}")
            self.assertEqual(
                len(client.errors),
                before_errors,
                provider_label + f"\nLive LLM error during prompt: {prompt}\n" + "\n".join(client.errors[before_errors:]),
            )
            self.assertTrue(result.success, result.message)
            answer = _text((result.data or {}).get("answer") or result.message)
            history.extend([{"role": "user", "content": prompt}, {"role": "assistant", "content": answer}])
            self._append_turn(transcript, prompt, result)
            transcript[-1]["llm_responses"] = client.responses[before_responses:]
            return result

        greeting = turn("naber dayi napiyon")
        self.assertEqual((greeting.data or {}).get("mode"), "answer")
        self.assertIsNone((greeting.data or {}).get("tool_call"))
        greeting_answer = _text((greeting.data or {}).get("answer"))
        self.assertNotIn("nasil yardimci olabilirim", greeting_answer.casefold())

        advice = turn("valla bilmiyorum ki sence?")
        self.assertEqual((advice.data or {}).get("mode"), "answer")
        self.assertIsNone((advice.data or {}).get("tool_call"))
        advice_answer = _text((advice.data or {}).get("answer"))
        self.assertGreater(len(advice_answer), 20)
        self.assertNotIn("siparis var", advice_answer.casefold())
        self.assertNotIn("bilemiyorum", advice_answer.casefold())
        self.assertNotIn("bilemem", advice_answer.casefold())
        self.assertNotIn("ne bileyim", advice_answer.casefold())
        self.assertNotIn("ne yapmak istersin", advice_answer.casefold())
        self.assertNotIn("sorabilirsin", advice_answer.casefold())
        self.assertNotIn("maalesef", advice_answer.casefold())
        self.assertNotIn("detay verirseniz", advice_answer.casefold())
        self.assertNotIn("yemek siparisi", advice_answer.casefold())
        self.assertNotIn("rota", advice_answer.casefold())
        self.assertNotIn("canin ne istiyor", advice_answer.casefold())

        ecommerce = turn("e ticaret nasil bisey sence iyi mi")
        self.assertEqual((ecommerce.data or {}).get("mode"), "answer")
        self.assertIsNone((ecommerce.data or {}).get("tool_call"))
        ecommerce_answer = _text((ecommerce.data or {}).get("answer"))
        self.assertGreater(len(ecommerce_answer), 35)
        self.assertFalse(ecommerce_answer.lstrip().startswith("{"))
        self.assertNotIn('"reply"', ecommerce_answer.casefold())

        analysis = turn("siparislerim nasil sence")
        self.assertEqual((analysis.data or {}).get("mode"), "tools")
        analysis_tools = [row.get("tool") for row in ((analysis.data or {}).get("tool_results") or [])]
        self.assertIn("orders.count", analysis_tools)
        self.assertIn("orders.revenue_total", analysis_tools)
        self.assertIn("returns.count", analysis_tools)
        self.assertIn("orders.ready_to_ship_count", analysis_tools)
        self.assertGreaterEqual(len(set(analysis_tools)), 5)
        analysis_answer = _text((analysis.data or {}).get("answer"))
        self.assertNotIn("dun", analysis_answer.casefold())
        self.assertNotIn("dün", analysis_answer.casefold())
        self.assertNotIn("yesterday", analysis_answer.casefold())
        self.assertIsNone(re.search(r"\b\d{12,14}\b", analysis_answer))

        help_result = turn("siparislerimle ilgili neler yapabiliriz")
        self.assertEqual((help_result.data or {}).get("mode"), "answer")
        self.assertIsNone((help_result.data or {}).get("tool_call"))
        help_answer = _text((help_result.data or {}).get("answer"))
        for term in ("etiket", "iade", "kargo"):
            self.assertIn(term, _fold_text(help_answer))

        listed = turn("son gelen 2 siparisi soyle")
        self.assertEqual(((listed.data or {}).get("tool_call") or {}).get("tool"), "orders.recent")
        listed_answer = _text((listed.data or {}).get("answer"))
        self.assertIn("ORD-LIVE-CONVO-EXTRA", listed_answer)

        prepared = turn("simdi bu iki siparisi yazdir")
        self.assertEqual((prepared.data or {}).get("mode"), "approval_required")
        self.assertEqual(((prepared.data or {}).get("tool_call") or {}).get("tool"), "labels.create_from_orders.prepare")
        action = (prepared.data or {}).get("action_proposal") or {}
        payload = ((action.get("arguments") or {}).get("payload") or {})
        self.assertEqual(set(payload.get("order_numbers") or []), {"ORD-LIVE-CONVO-EXTRA", "ORD-AYSE-READY"})

        print("\nLIVE_CONVERSATION_TRANSCRIPT " + provider_label)
        print(json.dumps(transcript, ensure_ascii=True, indent=2))
        artifact_path = _write_live_artifact(
            "live-conversation-transcript-quality",
            provider_label=provider_label,
            transcript=transcript,
        )
        if artifact_path is not None:
            print(f"LIVE_CONVERSATION_TRANSCRIPT_ARTIFACT {artifact_path}")

    def test_real_cloud_action_followup_transcript_quality(self) -> None:
        now_ms = self._seed_ready_orders()
        self._insert_waiting_return(
            "CLAIM-MAHMUT-CONVO",
            "ORD-MAHMUT-READY",
            now_ms + 3_000,
            first_name="Mahmut",
            last_name="Sevket",
        )
        self._insert_waiting_return(
            "CLAIM-AYSE-CONVO",
            "ORD-AYSE-READY",
            now_ms + 4_000,
            first_name="Ayse",
            last_name="Yilmaz",
        )

        client = self._client()
        provider_label = _safe_provider_label(self._provider_record, self._provider_source)
        transcript: list[dict[str, Any]] = []
        history: list[dict[str, str]] = []

        def turn(prompt: str):
            before = len(client.prompts)
            before_responses = len(client.responses)
            before_errors = len(client.errors)
            result = run_agent_turn(prompt, planner_client=client, chat_client=client, db_name=self.db_path, history=history)
            self.assertGreater(len(client.prompts) - before, 0, provider_label + f"\nPrompt did not call live LLM: {prompt}")
            self.assertEqual(
                len(client.errors),
                before_errors,
                provider_label + f"\nLive LLM error during prompt: {prompt}\n" + "\n".join(client.errors[before_errors:]),
            )
            self.assertTrue(result.success, result.message)
            answer = _text((result.data or {}).get("answer") or result.message)
            history.extend([{"role": "user", "content": prompt}, {"role": "assistant", "content": answer}])
            self._append_turn(transcript, prompt, result)
            transcript[-1]["llm_responses"] = client.responses[before_responses:]
            return result

        ready_list = turn("kargoya hazir siparisleri getir")
        self.assertEqual(((ready_list.data or {}).get("tool_call") or {}).get("tool"), "orders.ready_to_ship")
        ready_answer = _text((ready_list.data or {}).get("answer"))
        self.assertIn("ORD-MAHMUT-READY", ready_answer)
        self.assertIn("ORD-AYSE-READY", ready_answer)
        self.assertNotIn("gore icin", ready_answer.casefold())

        extracted_prepare = turn("bunlari fis cikarildi olarak isaretle")
        self.assertEqual((extracted_prepare.data or {}).get("mode"), "approval_required")
        self.assertEqual(((extracted_prepare.data or {}).get("tool_call") or {}).get("tool"), "operations.manual_action.route")
        extracted_action = (extracted_prepare.data or {}).get("action_proposal") or {}
        extracted_payload = _action_payload(extracted_action)
        self.assertEqual(extracted_action.get("action_type"), "labels.mark_orders_extracted_from_label_result")
        self.assertEqual(set(extracted_payload.get("order_numbers") or []), {"ORD-MAHMUT-READY", "ORD-AYSE-READY"})
        extracted_answer = _text((extracted_prepare.data or {}).get("answer") or extracted_prepare.message)
        self.assertIn("Onay", extracted_answer)
        self.assertNotIn("db_name", extracted_answer)
        self.assertNotIn("test.sqlite3", extracted_answer)
        _assert_no_raw_action_id(self, extracted_answer, "labels.mark_orders_extracted_from_label_result")
        with tempfile.TemporaryDirectory() as temp_dir:
            extracted_execute = operation_runtime.execute_approved_operation(
                extracted_action,
                approved_by="live-conversation-quality",
                idempotency_key="live-convo-mark-extracted-ready",
                audit_path=Path(temp_dir) / "orders-agent-audit.jsonl",
            )
        self.assertTrue(extracted_execute.success, extracted_execute.message)
        self.assertTrue(((extracted_execute.data or {}).get("verify") or {}).get("verified"), extracted_execute.data)
        with Session(get_engine(self.db_path)) as session:
            headers = session.exec(
                select(OrderHeader).where(OrderHeader.orderNumber.in_(["ORD-MAHMUT-READY", "ORD-AYSE-READY"]))
            ).all()
            self.assertEqual({row.orderNumber for row in headers if row.is_extracted}, {"ORD-MAHMUT-READY", "ORD-AYSE-READY"})
        self._append_execute(
            transcript,
            "labels.mark_orders_extracted_from_label_result",
            extracted_execute,
            extra={"db_extracted_orders": ["ORD-MAHMUT-READY", "ORD-AYSE-READY"]},
        )

        cargo_prepare = turn("Mahmut olanin kargo firmasini Aras yap")
        self.assertEqual((cargo_prepare.data or {}).get("mode"), "approval_required")
        self.assertEqual(((cargo_prepare.data or {}).get("tool_call") or {}).get("tool"), "orders.change_cargo_provider.prepare")
        cargo_action = (cargo_prepare.data or {}).get("action_proposal") or {}
        cargo_payload = _action_payload(cargo_action)
        self.assertEqual(cargo_action.get("action_type"), "orders.change_order_cargo_provider")
        self.assertEqual(cargo_payload.get("order"), "ORD-MAHMUT-READY")
        self.assertEqual(cargo_payload.get("cargo_provider_code"), "ARAS")
        cargo_answer = _text((cargo_prepare.data or {}).get("answer") or cargo_prepare.message)
        self.assertIn("Onay", cargo_answer)
        self.assertNotIn("db_name", cargo_answer)
        self.assertNotIn("test.sqlite3", cargo_answer)
        _assert_no_raw_action_id(self, cargo_answer, "orders.change_order_cargo_provider")

        def fake_change_order_cargo_provider(order, *, cargo_provider_code: str):
            with Session(get_engine(self.db_path)) as session:
                record = session.exec(select(OrderData).where(OrderData.orderNumber == order.orderNumber)).first()
                self.assertIsNotNone(record)
                record.cargoProviderName = "Aras"
                session.add(record)
                session.commit()
            return Result.ok("Kargo firmasi degistirildi.", close_dialog=False, data={"order_number": order.orderNumber, "cargo_provider_code": cargo_provider_code})

        with tempfile.TemporaryDirectory() as temp_dir, patch(
            "Orders.views.actions.change_order_cargo_provider",
            side_effect=fake_change_order_cargo_provider,
        ):
            cargo_execute = operation_runtime.execute_approved_operation(
                cargo_action,
                approved_by="live-conversation-quality",
                idempotency_key="live-convo-cargo-mahmut-aras",
                audit_path=Path(temp_dir) / "orders-agent-audit.jsonl",
            )
        self.assertTrue(cargo_execute.success, cargo_execute.message)
        self.assertTrue(((cargo_execute.data or {}).get("verify") or {}).get("verified"), cargo_execute.data)
        with Session(get_engine(self.db_path)) as session:
            mahmut_order = session.exec(select(OrderData).where(OrderData.orderNumber == "ORD-MAHMUT-READY")).first()
            ayse_order = session.exec(select(OrderData).where(OrderData.orderNumber == "ORD-AYSE-READY")).first()
            self.assertEqual(getattr(mahmut_order, "cargoProviderName", ""), "Aras")
            self.assertNotEqual(getattr(ayse_order, "cargoProviderName", ""), "Aras")
        self._append_execute(
            transcript,
            "orders.change_order_cargo_provider",
            cargo_execute,
            extra={
                "db_cargo_provider_by_order": {
                    "ORD-MAHMUT-READY": getattr(mahmut_order, "cargoProviderName", ""),
                    "ORD-AYSE-READY": getattr(ayse_order, "cargoProviderName", ""),
                }
            },
        )

        return_list = turn("bekleyen iadeleri getir")
        self.assertEqual(((return_list.data or {}).get("tool_call") or {}).get("tool"), "returns.query")
        return_answer = _text((return_list.data or {}).get("answer"))
        self.assertIn("CLAIM-MAHMUT-CONVO", return_answer)
        self.assertIn("CLAIM-AYSE-CONVO", return_answer)
        self.assertNotIn("gore icin", return_answer.casefold())

        return_prepare = turn("Mahmut olanin iadesini onayla")
        self.assertEqual((return_prepare.data or {}).get("mode"), "approval_required")
        self.assertEqual(((return_prepare.data or {}).get("tool_call") or {}).get("tool"), "returns.approve.prepare")
        return_action = (return_prepare.data or {}).get("action_proposal") or {}
        return_payload = _action_payload(return_action)
        self.assertEqual(return_action.get("action_type"), "orders.approve_return_claim")
        self.assertEqual(return_payload.get("order"), "CLAIM-MAHMUT-CONVO")
        return_answer = _text((return_prepare.data or {}).get("answer") or return_prepare.message)
        self.assertIn("Onay", return_answer)
        self.assertNotIn("db_name", return_answer)
        self.assertNotIn("test.sqlite3", return_answer)
        _assert_no_raw_action_id(self, return_answer, "orders.approve_return_claim")

        def fake_approve_return_claim(*, order, **_kwargs):
            with Session(get_engine(self.db_path)) as session:
                claim = session.exec(select(ReturnClaim).where(ReturnClaim.claimId == order.claimId)).one()
                claim.claimItemStatus = "Accepted"
                claim.waitingClaimLineItemIds = []
                session.add(claim)
                session.commit()
            return Result.ok("Iade onaylandi.", close_dialog=False, data={"claim_id": order.claimId})

        with tempfile.TemporaryDirectory() as temp_dir, patch(
            "Orders.views.actions.approve_return_claim",
            side_effect=fake_approve_return_claim,
        ):
            return_execute = operation_runtime.execute_approved_operation(
                return_action,
                approved_by="live-conversation-quality",
                idempotency_key="live-convo-return-mahmut-approve",
                audit_path=Path(temp_dir) / "orders-agent-audit.jsonl",
            )
        self.assertTrue(return_execute.success, return_execute.message)
        self.assertTrue(((return_execute.data or {}).get("verify") or {}).get("verified"), return_execute.data)
        with Session(get_engine(self.db_path)) as session:
            mahmut_claim = session.exec(select(ReturnClaim).where(ReturnClaim.claimId == "CLAIM-MAHMUT-CONVO")).one()
            ayse_claim = session.exec(select(ReturnClaim).where(ReturnClaim.claimId == "CLAIM-AYSE-CONVO")).one()
            self.assertEqual(mahmut_claim.claimItemStatus, "Accepted")
            self.assertEqual(mahmut_claim.waitingClaimLineItemIds, [])
            self.assertEqual(ayse_claim.claimItemStatus, "WaitingInAction")
        self._append_execute(
            transcript,
            "orders.approve_return_claim",
            return_execute,
            extra={
                "db_return_status_by_claim": {
                    "CLAIM-MAHMUT-CONVO": mahmut_claim.claimItemStatus,
                    "CLAIM-AYSE-CONVO": ayse_claim.claimItemStatus,
                }
            },
        )

        print("\nLIVE_ACTION_FOLLOWUP_TRANSCRIPT " + provider_label)
        print(json.dumps(transcript, ensure_ascii=True, indent=2))
        artifact_path = _write_live_artifact(
            "live-action-followup-transcript-quality",
            provider_label=provider_label,
            transcript=transcript,
            extra={
                "verified_side_effects": [
                    "label-extracted DB verify",
                    "Mahmut cargo change DB verify",
                    "Mahmut return approval DB verify",
                ]
            },
        )
        if artifact_path is not None:
            print(f"LIVE_ACTION_FOLLOWUP_TRANSCRIPT_ARTIFACT {artifact_path}")

    def test_real_cloud_fetch_sync_transcript_quality(self) -> None:
        with Session(get_engine(self.db_path)) as session:
            session.add(
                ApiAccount(
                    account_id="seller-live-fetch",
                    comp_name="Live Fetch Seller",
                    platform="trendyol",
                    api_key="live-fetch-key",
                    api_secret="live-fetch-secret",
                    is_active=True,
                )
            )
            session.add(
                ApiAccount(
                    account_id="seller-live-other",
                    comp_name="Other Marketplace Seller",
                    platform="amazon",
                    api_key="other-key",
                    api_secret="other-secret",
                    is_active=True,
                )
            )
            session.commit()
            trendyol_pk = int(session.exec(select(ApiAccount.pk).where(ApiAccount.account_id == "seller-live-fetch")).one())

        client = self._client()
        provider_label = _safe_provider_label(self._provider_record, self._provider_source)
        transcript: list[dict[str, Any]] = []
        history: list[dict[str, str]] = []

        def turn(prompt: str):
            before = len(client.prompts)
            before_responses = len(client.responses)
            before_errors = len(client.errors)
            result = run_agent_turn(prompt, planner_client=client, chat_client=client, db_name=self.db_path, history=history)
            self.assertGreater(len(client.prompts) - before, 0, provider_label + f"\nPrompt did not call live LLM: {prompt}")
            self.assertEqual(
                len(client.errors),
                before_errors,
                provider_label + f"\nLive LLM error during prompt: {prompt}\n" + "\n".join(client.errors[before_errors:]),
            )
            self.assertTrue(result.success, result.message)
            answer = _text((result.data or {}).get("answer") or result.message)
            history.extend([{"role": "user", "content": prompt}, {"role": "assistant", "content": answer}])
            self._append_turn(transcript, prompt, result)
            transcript[-1]["llm_responses"] = client.responses[before_responses:]
            return result

        prepared = turn("Trendyol siparis verilerini cek")
        self.assertEqual((prepared.data or {}).get("mode"), "approval_required")
        self.assertEqual(((prepared.data or {}).get("tool_call") or {}).get("tool"), "operations.manual_action.route")
        action = (prepared.data or {}).get("action_proposal") or {}
        payload = _action_payload(action)
        self.assertEqual(action.get("action_type"), "orders.get_orders_from_companies")
        self.assertTrue(payload.get("save_to_db"))
        self.assertFalse(payload.get("fetch_returns"))
        self.assertEqual(payload.get("api_account_ids"), [trendyol_pk])
        self.assertNotIn("accounts", payload)
        answer = _text((prepared.data or {}).get("answer") or prepared.message)
        self.assertIn("Onay", answer)
        self.assertNotIn("live-fetch-key", answer)
        self.assertNotIn("live-fetch-secret", answer)
        self.assertNotIn("db_name", answer)
        _assert_no_raw_action_id(self, answer, "orders.get_orders_from_companies")

        fetch_calls: list[dict[str, Any]] = []
        save_calls: list[dict[str, Any]] = []

        async def fake_fetch_orders_all(**kwargs):
            fetch_calls.append(dict(kwargs))
            return Result.ok(
                "Orders fake cekildi.",
                close_dialog=False,
                data={"order_data_list": [{"orderNumber": "ORD-LIVE-FETCH-1"}], "order_item_list": []},
            )

        def fake_save_orders(result, db_name=None):
            rows = list((result.data or {}).get("order_data_list") or [])
            save_calls.append({"rows": rows, "db_name": db_name})
            return Result.ok(
                "Orders fake kaydedildi.",
                close_dialog=False,
                data={"rows": rows, "counts": {"orders_upserted": len(rows)}, "db_name": db_name},
            )

        with tempfile.TemporaryDirectory() as temp_dir, patch(
            "Orders.processors.trendyol.pipeline.fetch_orders_all",
            side_effect=fake_fetch_orders_all,
        ), patch(
            "Orders.processors.trendyol.pipeline.save_orders_to_db",
            side_effect=fake_save_orders,
        ):
            executed = operation_runtime.execute_approved_operation(
                action,
                approved_by="live-conversation-quality",
                idempotency_key="live-convo-fetch-trendyol-orders",
                audit_path=Path(temp_dir) / "orders-agent-audit.jsonl",
            )

        self.assertTrue(executed.success, executed.message)
        verify = (executed.data or {}).get("verify") or {}
        self.assertTrue(verify.get("verified"), executed.data)
        self.assertEqual(verify.get("save_counts"), {"orders_upserted": 1})
        self.assertEqual(len(fetch_calls), 1)
        account_row = (fetch_calls[0].get("comp_api_account_list") or [])[0]
        self.assertEqual(account_row[0], trendyol_pk)
        self.assertGreaterEqual(len(account_row), 4)
        self.assertEqual(save_calls, [{"rows": [{"orderNumber": "ORD-LIVE-FETCH-1"}], "db_name": self.db_path}])
        self._append_execute(
            transcript,
            "orders.get_orders_from_companies",
            executed,
            extra={
                "fetch_call_count": len(fetch_calls),
                "saved_order_numbers": ["ORD-LIVE-FETCH-1"],
            },
        )
        joined_llm_responses = "\n".join(client.responses)
        self.assertNotIn("live-fetch-key", joined_llm_responses)
        self.assertNotIn("live-fetch-secret", joined_llm_responses)
        self.assertNotIn(self.db_path, joined_llm_responses)
        self.assertNotIn("test.sqlite3", joined_llm_responses)

        print("\nLIVE_FETCH_SYNC_TRANSCRIPT " + provider_label)
        print(json.dumps(transcript, ensure_ascii=True, indent=2))
        artifact_path = _write_live_artifact(
            "live-fetch-sync-transcript-quality",
            provider_label=provider_label,
            transcript=transcript,
            extra={"verified_side_effects": ["orders fetch adapter called", "save_orders_to_db verify"]},
        )
        if artifact_path is not None:
            print(f"LIVE_FETCH_SYNC_TRANSCRIPT_ARTIFACT {artifact_path}")

    def test_real_cloud_profit_export_followup_transcript_quality(self) -> None:
        now_ms = self._seed_ready_orders()
        self._insert_order_item("ORD-MAHMUT-READY", now_ms + 2_000, product_name="Mahmut Live Profit Urunu")
        self._insert_order_item("ORD-AYSE-READY", now_ms + 3_000, product_name="Ayse Live Profit Urunu")

        client = self._client()
        provider_label = _safe_provider_label(self._provider_record, self._provider_source)
        transcript: list[dict[str, Any]] = []
        history: list[dict[str, str]] = []

        def turn(prompt: str):
            before = len(client.prompts)
            before_responses = len(client.responses)
            before_errors = len(client.errors)
            result = run_agent_turn(prompt, planner_client=client, chat_client=client, db_name=self.db_path, history=history)
            self.assertGreater(len(client.prompts) - before, 0, provider_label + f"\nPrompt did not call live LLM: {prompt}")
            self.assertEqual(
                len(client.errors),
                before_errors,
                provider_label + f"\nLive LLM error during prompt: {prompt}\n" + "\n".join(client.errors[before_errors:]),
            )
            self.assertTrue(result.success, result.message)
            answer = _text((result.data or {}).get("answer") or result.message)
            history.extend([{"role": "user", "content": prompt}, {"role": "assistant", "content": answer}])
            self._append_turn(transcript, prompt, result)
            transcript[-1]["llm_responses"] = client.responses[before_responses:]
            return result

        ready_list = turn("kargoya hazir siparisleri getir")
        self.assertEqual(((ready_list.data or {}).get("tool_call") or {}).get("tool"), "orders.ready_to_ship")
        ready_answer = _text((ready_list.data or {}).get("answer"))
        self.assertIn("ORD-MAHMUT-READY", ready_answer)
        self.assertIn("ORD-AYSE-READY", ready_answer)

        interlude = turn("bu listeyle neler yapabiliriz?")
        self.assertEqual((interlude.data or {}).get("mode"), "answer")
        self.assertIsNone((interlude.data or {}).get("tool_call"))
        interlude_answer = _text((interlude.data or {}).get("answer"))
        self.assertIn("etiket", interlude_answer.casefold())
        self.assertIn("kargo", interlude_answer.casefold())
        self.assertTrue(
            any(term in interlude_answer.casefold() for term in ("excel", "detay", "kalem", "fis", "fiş", "iade")),
            interlude_answer,
        )

        with tempfile.TemporaryDirectory() as temp_dir:
            output_path = Path(temp_dir) / "mahmut-live-profit.xlsx"
            prompt_path = str(output_path).replace("\\", "/")
            prepared = turn(f'Mahmut olanin kar Excel raporunu "{prompt_path}" dosyasina aktar')
            self.assertEqual((prepared.data or {}).get("mode"), "approval_required")
            self.assertEqual(((prepared.data or {}).get("tool_call") or {}).get("tool"), "orders.export_profit_excel.prepare")
            action = (prepared.data or {}).get("action_proposal") or {}
            payload = _action_payload(action)
            self.assertEqual(action.get("action_type"), "orders.action_export_order_profit_excel")
            self.assertEqual(payload.get("order_number"), "ORD-MAHMUT-READY")
            self.assertEqual(Path(payload.get("output_path") or "").name, "mahmut-live-profit.xlsx")
            prepare_answer = _text((prepared.data or {}).get("answer") or prepared.message)
            self.assertIn("Onay", prepare_answer)
            self.assertIn("ORD-MAHMUT-READY", prepare_answer)
            self.assertNotIn("ORD-AYSE-READY", prepare_answer)
            self.assertNotIn("db_name", prepare_answer)
            self.assertNotIn("test.sqlite3", prepare_answer)
            _assert_no_raw_action_id(self, prepare_answer, "orders.action_export_order_profit_excel")

            executed = operation_runtime.execute_approved_operation(
                action,
                approved_by="live-conversation-quality",
                idempotency_key="live-convo-profit-export-mahmut",
                audit_path=Path(temp_dir) / "orders-agent-audit.jsonl",
            )

            self.assertTrue(executed.success, executed.message)
            self.assertTrue(output_path.exists(), executed.data)
            verify = (executed.data or {}).get("verify") or {}
            self.assertTrue(verify.get("verified"), executed.data)
            self.assertEqual(Path(verify.get("output_path") or "").name, "mahmut-live-profit.xlsx")

            from openpyxl import load_workbook

            workbook = load_workbook(output_path, read_only=True)
            try:
                self.assertEqual(set(workbook.sheetnames), {"Ozet", "Satis", "Maliyetler"})
                summary_rows = list(workbook["Ozet"].iter_rows(values_only=True))
                self.assertIn(("Sipariş", "ORD-MAHMUT-READY"), summary_rows)
                sales_rows = list(workbook["Satis"].iter_rows(values_only=True))
                self.assertTrue(
                    any("Mahmut Live Profit Urunu" in str(cell) for row in sales_rows for cell in row),
                    sales_rows,
                )
                self.assertFalse(any("Ayse Live Profit Urunu" in str(cell) for row in sales_rows for cell in row))
            finally:
                workbook.close()
            self._append_execute(
                transcript,
                "orders.action_export_order_profit_excel",
                executed,
                extra={
                    "workbook_exists": output_path.exists(),
                    "workbook_name": output_path.name,
                    "verified_excludes": ["Ayse Live Profit Urunu"],
                },
            )

        print("\nLIVE_PROFIT_EXPORT_TRANSCRIPT " + provider_label)
        print(json.dumps(transcript, ensure_ascii=True, indent=2))
        artifact_path = _write_live_artifact(
            "live-profit-export-transcript-quality",
            provider_label=provider_label,
            transcript=transcript,
            extra={"verified_side_effects": ["Mahmut workbook exists", "workbook excludes Ayse item"]},
        )
        if artifact_path is not None:
            print(f"LIVE_PROFIT_EXPORT_TRANSCRIPT_ARTIFACT {artifact_path}")

    def test_real_cloud_stale_return_restore_transcript_quality(self) -> None:
        now_ms = self._seed_ready_orders()
        with Session(get_engine(self.db_path)) as session:
            session.add(ApiAccount(pk=1, account_id="seller-live-stale-restore", comp_name="Live Stale Seller", platform="trendyol"))
            session.add(
                ReturnClaim(
                    claimId="CLAIM-LIVE-STALE-RESTORE",
                    orderNumber="ORD-MAHMUT-READY",
                    api_account_id=1,
                    claimDate=now_ms + 2_000,
                    orderDate=now_ms + 2_000,
                    lastModifiedDate=now_ms + 3_000,
                    claimItemStatus="WaitingInAction",
                    customerFirstName="Mahmut",
                    customerLastName="Sevket",
                    totalPrice=55.0,
                    itemCount=1,
                    waitingClaimLineItemIds=["line-1"],
                    sync_deleted_at="2026-01-01T00:00:00+03:00",
                )
            )
            session.commit()

        client = self._client()
        provider_label = _safe_provider_label(self._provider_record, self._provider_source)
        transcript: list[dict[str, Any]] = []
        history: list[dict[str, str]] = []

        def turn(prompt: str):
            before = len(client.prompts)
            before_responses = len(client.responses)
            before_errors = len(client.errors)
            result = run_agent_turn(prompt, planner_client=client, chat_client=client, db_name=self.db_path, history=history)
            self.assertGreater(len(client.prompts) - before, 0, provider_label + f"\nPrompt did not call live LLM: {prompt}")
            self.assertEqual(
                len(client.errors),
                before_errors,
                provider_label + f"\nLive LLM error during prompt: {prompt}\n" + "\n".join(client.errors[before_errors:]),
            )
            self.assertTrue(result.success, result.message)
            answer = _text((result.data or {}).get("answer") or result.message)
            history.extend([{"role": "user", "content": prompt}, {"role": "assistant", "content": answer}])
            self._append_turn(transcript, prompt, result)
            transcript[-1]["llm_responses"] = client.responses[before_responses:]
            return result

        review = turn("stale iadeleri goster")
        self.assertEqual((review.data or {}).get("mode"), "tool")
        self.assertEqual(((review.data or {}).get("tool_call") or {}).get("tool"), "operations.manual_action.route")
        review_tool = (review.data or {}).get("tool_result") or {}
        self.assertEqual(review_tool.get("route_status"), "read_executed")
        self.assertEqual(((review_tool.get("match") or {}).get("capability_id")), "orders.get_stale_return_claims")
        review_answer = _text((review.data or {}).get("answer"))
        self.assertIn("CLAIM-LIVE-STALE-RESTORE", review_answer)
        self.assertNotIn("Onay gerekiyor", review_answer)
        self.assertNotIn("UTC", review_answer)
        self.assertNotIn("Talep Tarihi", review_answer)

        prepared = turn("bu stale iadeyi geri al")
        self.assertEqual((prepared.data or {}).get("mode"), "approval_required")
        self.assertEqual(((prepared.data or {}).get("tool_call") or {}).get("tool"), "returns.restore_stale.prepare")
        action = (prepared.data or {}).get("action_proposal") or {}
        payload = _action_payload(action)
        self.assertEqual(action.get("action_type"), "orders.restore_return_claims_from_stale")
        self.assertEqual(payload.get("claim_ids"), ["CLAIM-LIVE-STALE-RESTORE"])
        prepare_answer = _text((prepared.data or {}).get("answer") or prepared.message)
        self.assertIn("Onay", prepare_answer)
        self.assertIn("CLAIM-LIVE-STALE-RESTORE", prepare_answer)
        self.assertNotIn("db_name", prepare_answer)
        self.assertNotIn("test.sqlite3", prepare_answer)
        _assert_no_raw_action_id(self, prepare_answer, "orders.restore_return_claims_from_stale")

        with tempfile.TemporaryDirectory() as temp_dir:
            executed = operation_runtime.execute_approved_operation(
                action,
                approved_by="live-conversation-quality",
                idempotency_key="live-convo-stale-return-restore",
                audit_path=Path(temp_dir) / "orders-agent-audit.jsonl",
            )

        self.assertTrue(executed.success, executed.message)
        verify = (executed.data or {}).get("verify") or {}
        self.assertTrue(verify.get("verified"), executed.data)
        self.assertEqual(verify.get("claim_ids"), ["CLAIM-LIVE-STALE-RESTORE"])
        with Session(get_engine(self.db_path)) as session:
            claim = session.exec(select(ReturnClaim).where(ReturnClaim.claimId == "CLAIM-LIVE-STALE-RESTORE")).one()
            self.assertIsNone(claim.sync_deleted_at)
        self._append_execute(
            transcript,
            "orders.restore_return_claims_from_stale",
            executed,
            extra={"db_sync_deleted_at": {"CLAIM-LIVE-STALE-RESTORE": claim.sync_deleted_at}},
        )

        lookup = turn("CLAIM-LIVE-STALE-RESTORE iadesini goster")
        self.assertEqual(((lookup.data or {}).get("tool_call") or {}).get("tool"), "returns.lookup")
        lookup_answer = _text((lookup.data or {}).get("answer"))
        self.assertIn("CLAIM-LIVE-STALE-RESTORE", lookup_answer)
        self.assertNotIn("geri alma yapilmadi", lookup_answer.casefold())
        self.assertNotIn("islem yapilmadi", lookup_answer.casefold())

        print("\nLIVE_STALE_RETURN_RESTORE_TRANSCRIPT " + provider_label)
        print(json.dumps(transcript, ensure_ascii=True, indent=2))
        artifact_path = _write_live_artifact(
            "live-stale-return-restore-transcript-quality",
            provider_label=provider_label,
            transcript=transcript,
            extra={"verified_side_effects": ["stale claim restored", "active lookup shows restored claim"]},
        )
        if artifact_path is not None:
            print(f"LIVE_STALE_RETURN_RESTORE_TRANSCRIPT_ARTIFACT {artifact_path}")


if __name__ == "__main__":
    unittest.main()
