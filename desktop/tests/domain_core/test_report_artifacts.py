from __future__ import annotations

import tempfile
import unittest
from pathlib import Path
from zipfile import ZipFile

from openpyxl import load_workbook

from Core.processors.report_artifacts import (
    ReportColumn,
    ReportSnapshot,
    make_report_sheet,
    verify_report_snapshot_xlsx,
    write_report_snapshot_xlsx,
)


class TestReportArtifacts(unittest.TestCase):
    def test_nullable_text_date_and_datetime_keep_source_nulls_and_verify_native_blanks(self):
        columns = tuple(ReportColumn(kind, kind, data_type=kind) for kind in ("string", "date", "datetime"))
        snapshot = ReportSnapshot.create(title="Nullable source", source_domain="files", report_type="fixture",
            filters={}, metadata={}, sheets=[make_report_sheet(name="Blank cells", columns=columns,
                rows=[{kind: None for kind in ("string", "date", "datetime")},
                      {kind: "" for kind in ("string", "date", "datetime")}])])
        self.assertEqual(snapshot.sheets[0].rows, ((None, None, None), ("", "", "")))
        self.assertEqual(ReportSnapshot.from_dict(snapshot.to_dict()), snapshot)
        with tempfile.TemporaryDirectory() as isolated:
            output = Path(isolated) / "nullable.xlsx"
            self.assertTrue(write_report_snapshot_xlsx(snapshot, output)["verified"])
            book = load_workbook(output)
            self.assertEqual(list(book["Blank cells"].values)[1:], [(None, None, None), (None, None, None)])
            book["Blank cells"]["A2"] = "0"
            book.save(output)
            book.close()
            self.assertFalse(verify_report_snapshot_xlsx(snapshot, output)["verified"])

    def _snapshot(self, *, created_at: str = "2026-07-21T10:00:00+00:00") -> ReportSnapshot:
        return ReportSnapshot.create(
            title="Son 3 Hafta En Çok Satanlar",
            source_domain="orders",
            report_type="orders.item_sales",
            filters={"period": "last_21_days", "group_by": "product"},
            metadata={"row_count": 2, "contains_customer_pii": False},
            sheets=(
                make_report_sheet(
                    name="En Cok Satanlar",
                    columns=(
                        ReportColumn("product", "Ürün", width=36),
                        ReportColumn("quantity", "Adet", data_type="integer", number_format="0"),
                        ReportColumn("gross", "Brüt", data_type="number", number_format="#,#0.00"),
                        ReportColumn("barcode", "Barkod", number_format="@", width=22),
                    ),
                    rows=(
                        {"product": "Musluk Başlığı", "quantity": 7, "gross": 699.93, "barcode": "869000000001"},
                        {"product": "=FORMULA-OLMASIN", "quantity": 3, "gross": 299.97, "barcode": "869000000002"},
                    ),
                ),
            ),
            created_at=created_at,
        )

    def test_snapshot_identity_is_content_stable_and_tamper_evident(self) -> None:
        first = self._snapshot(created_at="2026-07-21T10:00:00+00:00")
        later = self._snapshot(created_at="2026-07-22T10:00:00+00:00")
        self.assertEqual(first.snapshot_id, later.snapshot_id)
        self.assertEqual(first.content_hash, later.content_hash)
        self.assertEqual(ReportSnapshot.from_dict(first.to_dict()), first)

        tampered = first.to_dict()
        tampered["sheets"][0]["rows"][0][1] = 99
        with self.assertRaisesRegex(ValueError, "hash mismatch"):
            ReportSnapshot.from_dict(tampered)

    def test_xlsx_write_is_exactly_read_back_and_detects_cell_tamper(self) -> None:
        snapshot = self._snapshot()
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "sales.xlsx"
            written = write_report_snapshot_xlsx(snapshot, path)

            self.assertTrue(written.get("verified"), written)
            self.assertEqual(written.get("snapshot_id"), snapshot.snapshot_id)
            self.assertGreater(int(written.get("size_bytes") or 0), 0)

            workbook = load_workbook(path)
            self.assertEqual(workbook["En Cok Satanlar"]["A3"].value, "'=FORMULA-OLMASIN")
            self.assertEqual(workbook["En Cok Satanlar"]["D2"].value, "869000000001")
            self.assertEqual(workbook["En Cok Satanlar"]["D2"].data_type, "s")
            self.assertEqual(workbook["En Cok Satanlar"]["D2"].number_format, "@")
            with ZipFile(path) as archive:
                sheet_xml = archive.read("xl/worksheets/sheet1.xml")
            self.assertIn(b"<r><rPr/>", sheet_xml)
            self.assertIn(b"869000000001", sheet_xml)
            workbook["En Cok Satanlar"]["B2"] = 999
            workbook.save(path)
            workbook.close()

            verification = verify_report_snapshot_xlsx(snapshot, path)
            self.assertFalse(verification.get("verified"), verification)
            self.assertIn("rows_mismatch:En Cok Satanlar", verification.get("problems") or [])


if __name__ == "__main__":
    unittest.main()
